"""Messwerkzeug: python -m src.cli messen --korpus entwicklung|ex_test|test|nachbau_word [--paar NAME] [--ordner PFAD]
                                         [--mit-vorschlaegen] [--vergleich [--gegen NNN]] [--legende L] [--nachher-format F]
              Rueckbau-Messung (nur PDF + nachher.xlsx):  ... messen --nachher [--ordner PFAD] [--paar NAME]

Ein Paar = aenderungen.pdf + vorher.xlsx + nachher.xlsx (Format: docs/paar_format.md).
Das Tool laeuft auf PDF + vorher; das Ergebnis wird mit "nachher" verglichen.
Das Soll kommt NUR aus vorher <-> nachher (Zellvergleich, Zeilenausrichtung), nie aus dem PDF-Leser.

Regel 16: Konsole und Bericht enthalten NIE Text aus Dokumenten -- nur Paar-Namen, Blattnummer,
Zeile, Spalte, Aenderungs-ID, Zaehler und Fehlerart. Texte werden nur im Speicher verglichen.
Regel 29: Berichte sind append-only (NNN_<korpus>), es wird nie etwas ueberschrieben.

Aufbau
  1. Dateien lesen und Ablauf je Paar (ohne runs-Ordner, im Temp-Ordner)
  2. Soll/Ist: Zeilen und Zellen ausrichten (difflib), Einheiten bilden
  3. Kennzahlen, stille Fehler, Fehlerart-Kaskade
  4. Bericht (json + md) und Konsole
"""
from __future__ import annotations

import argparse
import bisect
import contextlib
import copy
import datetime
import hashlib
import json
import re
import shutil
import sys
import tempfile
import time
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from itertools import groupby
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string

from . import config, konsens
from .applier import uebernehmen
from .excel_model import ExcelModell
from .models import ADD, DEL, MOD, NICHT, PRUEFEN, SCHEMA_VERSION, UEBERNOMMEN
from .pdf_reader import analysiere_pdf
from .planner import plane
from .protokoll import bilanz

BERICHT_FORMAT = 1
FEHLERARTEN = ("pdf_falsch_gelesen", "layout_falsch", "zuordnung_falsch", "plan_falsch", "bilanz_kaputt", "unklar")
STILLE_ARTEN = ("uebernommen_aber_falsch", "kollateral", "bilanz", "nachkontrolle_ok_aber_falsch",
                "keine_wirkung", "luecke_nicht_gemeldet", "regel27_verstoss", "datentyp_gewechselt", "format_falsch",
                "fremde_spec")
STRUKTUR_ARTEN = ("blatt", "sichtbarkeit", "zeile_versteckt", "spalte_versteckt", "verbunden", "validierung",
                  "zahlenformat", "schrift", "datentyp", "bedingte_formatierung", "hyperlink", "kommentar",
                  "benannter_bereich", "autofilter", "tabelle", "fixierung", "spaltenbreite", "zeilenhoehe")
# Was die Ampel "Struktur unveraendert" prueft und was NICHT (K-MESS2: ehrlich benennen)
STRUKTUR_GEPRUEFT = ("Blätter (Name, Reihenfolge, Sichtbarkeit), ausgeblendete Zeilen und Spalten, verbundene Zellen, Datenvalidierung, "
                     "bedingte Formatierung, Hyperlink-Adressen, Kommentare, benannte Bereiche, Autofilter, Excel-Tabellen, fixierte Zeilen/Spalten, "
                     "Spaltenbreiten, Zeilenhöhen unveränderter Zeilen, Zahlenformat, Schrift und Füllung unveränderter Zellen, "
                     "Datentyp (Zahl/Text/Formel/Datum) unveränderter Zellen und von Zellen, die das Tool richtig befüllt hat")
STRUKTUR_NICHT_GEPRUEFT = ("Rahmen, Ausrichtung und Zellschutz, Blattschutz, Druckbereich und Seitenlayout, Kopf- und Fußzeilen, Bilder, "
                           "Diagramme und Formen, Gliederungsebenen, Makros, Dokumenteigenschaften, Teilformatierung innerhalb einer Zelle, "
                           "ob Formeln nach dem Verschieben noch auf die richtigen Zellen zeigen")
# Woran man in den Hinweisen des Plans erkennt, dass ein Strukturrisiko VORAB gemeldet war (Text aus excel_model/planner/applier)
STRUKTUR_RISIKO_STICHWORT = {
    "verbunden": "verbundene zellen", "validierung": "datenüberprüfung", "bedingte_formatierung": "bedingte formatierung",
    "autofilter": "autofilter", "tabelle": "excel-tabelle", "benannter_bereich": "benannte", "hyperlink": "hyperlink",
    "zeile_versteckt": "ausgeblendete", "zahlenformat": "zahlenformat", "datentyp": "zahlenformat",
}
SOLL_KLASSEN = ("haupt", "fuell_konstant", "formel", "nicht_ableitbar")     # Messung v2 (Etappe 2), siehe _Auswertung._klasse2
ABLEITBAR = ("haupt", "fuell_konstant", "formel")
OHNE_CODE = "ohne_code"          # gesperrte Aenderung ohne Pruefcode
RESTARBEIT_TEILE = ("pruefen", "nicht_uebernommen", "leere_zellen_neue_zeilen", "luecken_woerter")
RESTARBEIT_RICHTWERT = 12        # docs/ziel_v2.md: berichten, Richtwert hoechstens 12 je 100 Aenderungen
ABLEITBAR_ZIEL = 0.85            # docs/ziel_v2.md (Etappe 2): Abdeckung ableitbarer Zellen mindestens 85 % (bis dahin nur Hinweis)
ERZEUGER_KENNUNGEN = ("writer", "libreoffice")
KORPORA = ("entwicklung", "ex_test", "test", "nachbau_word", "nachbau_test")
UNGESEHEN = ("test", "nachbau_test")      # reservierte Endmess-Korpora (Regel 11): nie zum Einstellen, genau einmal messen
# Legende und nachher-Format je Korpus (KORPUS.json, Auftrag 003): fehlt die Angabe, gilt der Standard der Korpora v1/v2
LEGENDEN = ("alt", "auto", "linie", "farbe")
NACHHER_FORMATE = ("sauber", "markiert")
KORPUS_STANDARD = {"legende": "alt", "nachher_format": "sauber"}
KORPUS_VORGABE = {"entwicklung": KORPUS_STANDARD, "ex_test": KORPUS_STANDARD, "test": KORPUS_STANDARD,
                  "nachbau_word": {"legende": "auto", "nachher_format": "markiert"},
                  "nachbau_test": {"legende": "auto", "nachher_format": "markiert"}}
# Rueckbau-Messung (messen --nachher): Kategorien je Zelle
RUECKBAU_KATEGORIEN = ("gleich", "gleich_bis_leerraum", "tool_fehlt_vorschlag_offen", "tool_fehlt_gar_nicht", "tool_extra",
                       "verschieden")
RESTARBEIT_FORMEL = ("Restarbeit je 100 Änderungen = 100 × (Änderungen „prüfen“ + Änderungen „nicht übernommen“ + leere Zellen in "
                     "neuen Zeilen, die laut nachher gefüllt sein müssten + Wörter in gemeldeten Lücken ÷ 10) ÷ Änderungen. "
                     "Sie schätzt, wie viel Handarbeit nach dem Lauf bleibt (10 Lücken-Wörter ≈ eine Änderung).")
RESTARBEIT_LUECKEN_TEILER = 10   # Restarbeit: 10 Woerter in einer gemeldeten Luecke zaehlen wie eine Aenderung
FUELL_MAX_WOERTER = 3       # Zelle einer neuen/entfernten Zeile in einer nicht verwalteten Spalte: bis hierhin "harmlose Fuellung"
_LEER_SPEZIAL = frozenset("               　")
_PDF_BLIND = _LEER_SPEZIAL | frozenset("­​‌‍⁠﻿")     # Zeichen, die ein PDF meist nicht traegt
PLAUSI_BILD_ANTEIL = 0.05  # ein Bild zaehlt als Hinweis, wenn es mindestens so viel der Seitenflaeche deckt (kein Logo)
ZEILEN_AEHNLICH = 0.6       # ab hier gelten zwei Zeilen als "dieselbe Zeile, veraendert"
BLOCK_MAX = 4000            # groessere Austausch-Bloecke werden nur nach Position gepaart
EXCEL_MAX_ZEILE = 1048576   # letzte Zeile eines Excel-Blatts (Bereiche bis dorthin wachsen beim Einfuegen nicht weiter)
SPEC_MIN_BLOCK = 3          # so viele aufeinanderfolgende Zeilen mit derselben fremden Spec-Nr. bilden einen fremden Spec-Block
SPEC_SEITEN = 8             # Spec-Nummern werden auf so vielen ersten Seiten des PDF gesucht (Kopf, Fuss, SECTION-Zeile)
_SPEC_NR = re.compile(r"(?<!\d)(\d{2})[ _.]?(\d{2})[ _.]?(\d{2})(?!\d)")
DATEIEN = ("aenderungen.pdf", "vorher.xlsx", "nachher.xlsx")

_UNSICHTBAR = dict.fromkeys(map(ord, "­​‌‍﻿"), None)


# ======================================================================================
# 1. Dateien lesen
# ======================================================================================
def _k(text: str) -> str:
    """LOCKERE Vergleichsform einer Zelle: Unicode-normalisiert, ohne unsichtbare Zeichen, Leerraum zusammengefasst.
    Nur noch fuer "ist die Zelle ueberhaupt leer?" und zur Fehlerart-Bestimmung. Der Zellvergleich selbst ist streng (`_s`)."""
    return " ".join(unicodedata.normalize("NFKC", text).translate(_UNSICHTBAR).split())


def _s(text: str) -> str:
    """STRENGE Vergleichsform einer Zelle (RT-135): jedes Zeichen zaehlt -- doppelte und geschuetzte Leerzeichen, weiche und
    unsichtbare Zeichen, Ligaturen, Absatzgrenzen. Nur die Zeilenenden werden vereinheitlicht. Eine Zelle ohne sichtbaren
    Inhalt gilt als leer. `textutil.norm` ist nur fuer Wort-Vergleiche da (Zuordnung Aenderung <-> Zelle).
    Markierte Zelle (`_FText`, nur bei nachher_format "markiert"): die Formatklassen je Zeichen gehoeren zur Vergleichsform --
    gleicher Text mit anderer Markierung ist eine andere Zelle."""
    t = text.replace("\r\n", "\n").replace("\r", "\n") if _k(text) else ""
    kl = getattr(text, "klassen", None)
    return t + "\x00" + _lauflaengen(kl) if t and kl else t


def _s_text(text: Any) -> str:
    """Wie `_s`, aber nur der Text (ohne Formatklassen) -- fuer den Vergleich mit Plan-Werten des Tools."""
    t = text if isinstance(text, str) else _zelltext(text)
    return str(t).replace("\r\n", "\n").replace("\r", "\n") if _k(str(t)) else ""


def _ohne_format(key: str) -> str:
    """Vergleichsform ohne angehaengte Formatklassen."""
    return key.split("\x00", 1)[0]


class _FText(str):
    """Zelltext mit Formatklasse je Zeichen (nur nachher_format "markiert", ADR 0005): '-' normal, 'a' alt (durchgestrichen,
    jede Farbe), 'n' neu (Schrift in der Neu-Farbe der Markierung, ohne Strich). Nur fuer Zellen mit mindestens einem markierten
    Zeichen; der Text selbst bleibt ein gewoehnlicher str (Woerter, Laenge, Vergleich ohne Format)."""
    klassen: str

    def __new__(cls, text: str, klassen: str) -> "_FText":
        o = super().__new__(cls, text)
        o.klassen = klassen
        return o


def _lauflaengen(klassen: str) -> str:
    """'--aan' -> '-2a2n1' (kurz, damit die Zeilen-Aehnlichkeit vom Text bestimmt bleibt)."""
    return "".join(f"{k}{sum(1 for _ in g)}" for k, g in groupby(klassen))


def _klassen(text: Any) -> str:
    """Formatklassen eines Zelltexts ('-' je Zeichen, wenn er keine Markierung traegt)."""
    kl = getattr(text, "klassen", None)
    return kl if kl else "-" * len(str(text))


def _ansicht(text: Any, welche: str) -> str:
    """'alt' = normal + alt (Original), 'neu' = normal + neu (Stand nach der Aenderung); ohne Markierung der Text selbst."""
    kl = getattr(text, "klassen", None)
    if not kl:
        return str(text)
    weg = "n" if welche == "alt" else "a"
    return "".join(c for c, k in zip(text, kl) if k != weg)


def _stuecke(text: Any) -> list[tuple[str, str]]:
    """[(normal|alt|neu, Text)] eines Zelltexts (fuer die Formen in Schablonen)."""
    namen = {"-": "normal", "a": "alt", "n": "neu"}
    out: list[tuple[str, str]] = []
    for c, k in zip(str(text), _klassen(text)):
        if out and out[-1][0] == namen[k]:
            out[-1] = (out[-1][0], out[-1][1] + c)
        else:
            out.append((namen[k], c))
    return out


def _gleich_bis_leerraum(ist: Any, soll: Any) -> bool:
    """Gleiche sichtbare Zeichen mit gleicher Formatklasse; nur Leerraum (Anzahl, Art, Zeilenumbruch) darf abweichen."""
    def form(t: Any) -> list[tuple[str, str]]:
        return [(c, k) for c, k in zip(unicodedata.normalize("NFC", str(t)), _klassen(t)) if not c.isspace()
                and c not in "­​‌‍﻿"]
    return form(ist) == form(soll)


def _farb_rgb(farbe: Any) -> str | None:
    if farbe is None or getattr(farbe, "type", "rgb") != "rgb":
        return None
    rgb = getattr(farbe, "rgb", None)
    return rgb[-6:].upper() if isinstance(rgb, str) and len(rgb) in (6, 8) else None


def _markierungsfarben() -> tuple[str, str]:
    """(alt, neu) der Markierung dieser Runde laut config/zuordnung.yaml (ADR 0005: FF00B0F0 + Strich, FFFF0000)."""
    m = config.lade("zuordnung").get("markierung", {}) or {}
    return str(m.get("alt_farbe", "FF00B0F0"))[-6:].upper(), str(m.get("neu_farbe", "FFFF0000"))[-6:].upper()


def _formatklasse(schrift: Any, neu_rgb: str) -> str:
    """Formatklasse einer Schrift fuer die Format-Treue: durchgestrichen = alt (jede Farbe), Neu-Farbe ohne Strich = neu, sonst
    normal (auch Markierungen frueherer Runden in anderen Farben)."""
    if schrift is None:
        return "-"
    if getattr(schrift, "strike", None):
        return "a"
    return "n" if _farb_rgb(getattr(schrift, "color", None)) == neu_rgb else "-"


def _zelltext_formatiert(zelle: Any, neu_rgb: str) -> str:
    """Zelltext mit Formatklassen (`_FText`), wenn ein Zeichen markiert ist; sonst wie `_zelltext`."""
    v = zelle.value
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        teile = [(b.text, _formatklasse(b.font, neu_rgb)) if isinstance(b, TextBlock) else (str(b), _formatklasse(zelle.font, neu_rgb))
                 for b in v]
    else:
        teile = [(_zelltext(v), _formatklasse(zelle.font, neu_rgb))]
    if all(k == "-" for _t, k in teile):
        return "".join(t for t, _k in teile)
    text, kl = [], []
    for t, k in teile:
        t = t.replace("\r\n", "\n").replace("\r", "\n")
        text.append(t)
        kl.append(k * len(t))
    return _FText("".join(text), "".join(kl))


def _zelltext(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        return str(v)
    if isinstance(v, str):
        return v
    if isinstance(v, float):
        return format(v, ".15g")
    return str(v)


def _typ_von(zelle: Any) -> str:
    """Datentyp einer Zelle: leer | text | zahl | formel | datum | bool | fehler. Der Zelltext (`_zelltext`) macht aus der Zahl 31
    und dem Text "31" dasselbe; der Typ nicht (K-MESS2)."""
    v = zelle.value
    if v is None or (isinstance(v, str) and v == ""):
        return "leer"
    dt = zelle.data_type
    if dt == "f":
        return "formel"
    if dt == "b":
        return "bool"
    if dt == "e":
        return "fehler"
    if dt == "d" or isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return "datum"
    if dt == "n":
        return "datum" if zelle.is_date else "zahl"
    return "text"


def _lese_excel_mit_typen(pfad: Path, formate: bool = False) -> tuple[list[list[list[str]]], list[list[list[str]]]]:
    """Alle Blaetter als Zeilen: (Zelltexte, Datentypen) -- beide gleich geformt. Formeln stehen als Text.
    formate=True (nachher_format "markiert"): markierte Zellen kommen als `_FText` mit Formatklasse je Zeichen."""
    wb = load_workbook(str(pfad), rich_text=True)
    neu_rgb = _markierungsfarben()[1]
    try:
        texte: list[list[list[str]]] = []
        typen: list[list[list[str]]] = []
        for ws in wb.worksheets:
            t_zeilen: list[list[str]] = []
            y_zeilen: list[list[str]] = []
            for row in ws.iter_rows():
                t_zeilen.append([_zelltext_formatiert(c, neu_rgb) if formate else _zelltext(c.value) for c in row])
                y_zeilen.append([_typ_von(c) for c in row])
            texte.append(t_zeilen)
            typen.append(y_zeilen)
        return texte, typen
    finally:
        wb.close()


def _lese_excel(pfad: Path, formate: bool = False) -> list[list[list[str]]]:
    """Alle Blaetter als Zeilen mit Zelltexten (Formeln als Text)."""
    return _lese_excel_mit_typen(pfad, formate)[0]


def _mit_format(blaetter: list[list[list[str]]]) -> list[list[list[tuple[str, str | None]]]]:
    """Vergleichbare Form einer gelesenen Excel samt Formatklassen (Idempotenz im Modus "markiert")."""
    return [[[(str(c), getattr(c, "klassen", None)) for c in z] for z in b] for b in blaetter]


def _blattnamen(pfad: Path) -> list[str]:
    wb = load_workbook(str(pfad), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _sha(pfad: Path) -> str:
    h = hashlib.sha256()
    with open(pfad, "rb") as f:
        for stueck in iter(lambda: f.read(1 << 20), b""):
            h.update(stueck)
    return h.hexdigest()


@dataclass
class Zeile:
    nr: int                       # Zeilennummer im Blatt (1 = erste)
    zellen: tuple[str, ...]       # Rohtexte (hinten ohne leere Zellen)
    keys: tuple[str, ...]         # Vergleichsformen


def _zeilen(rows: list[list[str]]) -> list[Zeile]:
    """Nur nicht-leere Zeilen (leere Zeilen zaehlen nicht; die Nummern bleiben die des Originals)."""
    out: list[Zeile] = []
    for i, row in enumerate(rows, 1):
        keys = [_s(c) for c in row]
        while keys and not keys[-1]:
            keys.pop()
        if keys:
            out.append(Zeile(i, tuple(row[:len(keys)]), tuple(keys)))
    return out


# ======================================================================================
# 2. Ausrichten: Diff vorher <-> X (X = nachher oder Ergebnis)
# ======================================================================================
@dataclass
class Diff:
    zellen: dict[tuple[int, int], str] = field(default_factory=dict)   # (Zeile, Spalte) -> neuer Rohtext ("" = leer)
    weg: set[int] = field(default_factory=set)                          # vorher-Zeilen ohne Gegenstueck
    neue: list[tuple[int, tuple[str, ...], tuple[str, ...]]] = field(default_factory=list)   # (Anker, Rohtexte, Keys)
    neue_nr: list[int] = field(default_factory=list)                    # Zeilennummer jeder neuen Zeile in X (parallel zu `neue`)
    blind: set = field(default_factory=set)                             # Soll-Zellen, die nur "PDF-blind" verglichen werden (siehe _blind)
    neue_vi: list[int] = field(default_factory=list)                    # Index der vorher-Zeile direkt vor jeder neuen Zeile (roh)
    # Anker = Nummer der letzten behaltenen vorher-Zeile VOR der neuen Zeile (0 = ganz oben).
    # "Behalten" heisst: nicht weg UND nicht ausgeleert (eine Zeile, in der nur noch unverwaltete Spalten wie
    # Bemerkung stehen, zaehlt wie eine geloeschte Zeile -- im Soll genauso wie im Ergebnis).
    zuordnung: dict[int, int] = field(default_factory=dict)             # vorher-Zeile -> Zeile in X (alle gepaarten)
    geleert: set[int] = field(default_factory=set)                      # gepaarte vorher-Zeilen, die in X ausgeleert sind


def _woerter_zaehler(text: str) -> Counter:
    return Counter(text.split())


def _enthalten(klein: str, gross: str) -> bool:
    """Steht der kuerzere Text (mindestens drei Woerter) im Wesentlichen im laengeren? (Teilstring oder Woerter-Teilmenge)"""
    if len(klein.split()) < 3 or len(klein) >= len(gross):
        return False
    if klein in gross:
        return True
    return not (_woerter_zaehler(klein) - _woerter_zaehler(gross))


def _zell_aehnlich(x: str, y: str) -> float:
    """Aehnlichkeit zweier Zelltexte (0..1). Ein gekuerzter oder verlaengerter Text (der eine steckt im anderen) zaehlt
    als dieselbe Zelle, veraendert -- so bleibt eine Zeile erkennbar, wenn ein Absatz hinzukommt oder wegfaellt."""
    if x == y:
        return 1.0
    if not x or not y:
        return 0.0
    kurz, lang = (x, y) if len(x) <= len(y) else (y, x)
    best = 0.8 if _enthalten(kurz, lang) else 0.0
    sm = SequenceMatcher(None, x, y, autojunk=False)
    if sm.real_quick_ratio() >= 0.5 and sm.quick_ratio() >= 0.5:
        r = sm.ratio()
        if r >= 0.5:
            best = max(best, r)
    return best


def _aehnlich(a: tuple[str, ...], b: tuple[str, ...]) -> float:
    """Aehnlichkeit zweier Zeilen (0..1): Anteil gleicher bzw. aehnlicher Zellen. Eine Zeile, die nur "ausgeleert" wurde
    (die belegten Zellen der einen Zeile stehen -- gleich oder gekuerzt -- in der anderen), zaehlt als dieselbe Zeile."""
    a_ne = {c: v for c, v in enumerate(a) if v}
    b_ne = {c: v for c, v in enumerate(b) if v}
    teil = 0.0
    for klein, gross in ((a_ne, b_ne), (b_ne, a_ne)):
        if klein and klein.keys() <= gross.keys():
            werte = [_zell_aehnlich(v, gross[c]) for c, v in klein.items()]
            if all(w >= 0.5 for w in werte):
                teil = max(teil, sum(werte) / len(werte))
    gewicht = summe = 0.0
    for c in range(max(len(a), len(b))):
        x = a[c] if c < len(a) else ""
        y = b[c] if c < len(b) else ""
        if not x and not y:
            continue
        gewicht += 1.0                               # jede belegte Zelle zaehlt gleich (lange Texte erdruecken Nummer/Titel nicht)
        summe += _zell_aehnlich(x, y)
    return max(teil, summe / gewicht if gewicht else 0.0)


def _proj(keys: tuple[str, ...], spalten: Any) -> tuple[str, ...]:
    """Nur die Haupt-Spalten einer Zeile (die anderen leer)."""
    return tuple(k if c in spalten else "" for c, k in enumerate(keys, 1))


def _aehnlich_haupt(a: tuple[str, ...], b: tuple[str, ...], spalten: Any = None) -> float:
    """Aehnlichkeit zweier Zeilen nur ueber die Haupt-Spalten (Nacharbeit Welle 2): eine Fuellspalte (Bemerkung, Prio ...) darf
    die Zeilen-Paarung nicht verfaelschen. Ist eine Seite in den Haupt-Spalten leer (ausgeleerte Zeile, Zeile ohne Haupt-Inhalt),
    zaehlt wie bisher die ganze Zeile."""
    if spalten:
        pa, pb = _proj(a, spalten), _proj(b, spalten)
        if any(pa) and any(pb):
            return _aehnlich(pa, pb)
    return _aehnlich(a, b)


def _paare(va: list[Zeile], xb: list[Zeile], spalten: Any = None) -> list[tuple[int, int]]:
    """Ordnungstreue Zuordnung zwischen zwei Zeilenbloecken (Austausch-Block), beste Gesamtaehnlichkeit.
    `spalten`: Haupt-Spalten (nur sie zaehlen fuer die Aehnlichkeit), None = alle."""
    n, m = len(va), len(xb)
    if not n or not m:
        return []
    if n * m > BLOCK_MAX:
        return [(i, i) for i in range(min(n, m)) if _aehnlich_haupt(va[i].keys, xb[i].keys, spalten) >= ZEILEN_AEHNLICH]
    s = [[_aehnlich_haupt(va[i].keys, xb[j].keys, spalten) for j in range(m)] for i in range(n)]
    best = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            b = max(best[i - 1][j], best[i][j - 1])
            if s[i - 1][j - 1] >= ZEILEN_AEHNLICH:
                b = max(b, best[i - 1][j - 1] + s[i - 1][j - 1])
            best[i][j] = b
    out: list[tuple[int, int]] = []
    i, j = n, m
    while i > 0 and j > 0:
        if s[i - 1][j - 1] >= ZEILEN_AEHNLICH and best[i][j] == best[i - 1][j - 1] + s[i - 1][j - 1]:
            out.append((i - 1, j - 1))
            i, j = i - 1, j - 1
        elif best[i][j] == best[i - 1][j]:
            i -= 1
        else:
            j -= 1
    return out[::-1]


def _diff(v: list[Zeile], x: list[Zeile], leer: Callable[[Zeile], bool] | None = None, spalten: Any = None) -> Diff:
    """Zellen-, Zeilen-Unterschiede vorher -> X. `leer(zeile)` sagt, ob eine Zeile in den verwalteten Spalten nichts mehr
    enthaelt (dann ist sie als Anker so gut wie geloescht)."""
    d = Diff()
    sm = SequenceMatcher(None, [z.keys for z in v], [z.keys for z in x], autojunk=False)
    weg_idx: set[int] = set()
    leer_idx: set[int] = set()
    roh_neu: list[tuple[int, Zeile]] = []

    def weg(i: int) -> None:
        z = v[i]
        weg_idx.add(i)
        d.weg.add(z.nr)
        for c, k in enumerate(z.keys, 1):
            if k:
                d.zellen[(z.nr, c)] = ""

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                d.zuordnung[v[i1 + k].nr] = x[j1 + k].nr
            continue
        paare = _paare(v[i1:i2], x[j1:j2], spalten) if tag == "replace" else []
        zu_x = {j1 + b: i1 + a for a, b in paare}
        zeiger = i1
        for j in range(j1, j2):
            if j in zu_x:
                i = zu_x[j]
                while zeiger < i:
                    weg(zeiger)
                    zeiger += 1
                a, b = v[i], x[j]
                d.zuordnung[a.nr] = b.nr
                if leer is not None and leer(b):
                    leer_idx.add(i)
                    d.geleert.add(a.nr)
                for c in range(max(len(a.keys), len(b.keys))):
                    ka = a.keys[c] if c < len(a.keys) else ""
                    kb = b.keys[c] if c < len(b.keys) else ""
                    if ka != kb:
                        d.zellen[(a.nr, c + 1)] = b.zellen[c] if c < len(b.zellen) else ""
                zeiger = i + 1
            else:
                roh_neu.append((zeiger - 1, x[j]))
        while zeiger < i2:
            weg(zeiger)
            zeiger += 1
    for vi, z in roh_neu:
        d.neue_vi.append(vi)
        while vi >= 0 and (vi in weg_idx or vi in leer_idx):
            vi -= 1
        d.neue.append((v[vi].nr if vi >= 0 else 0, z.zellen, z.keys))
        d.neue_nr.append(z.nr)
    return d


def _anker_neu(d: Diff, v: list[Zeile], tot: set[int]) -> None:
    """Setzt die Anker der neuen Zeilen neu: Zeilen, die Soll ODER Ergebnis als entfernt/ausgeleert sehen (`tot`), zaehlen auf
    beiden Seiten gleich nicht als Anker. So hat eine geleerte (statt geloeschte) Zeile keinen Einfluss auf den Ort (MA-3)."""
    for k, vi in enumerate(d.neue_vi):
        while vi >= 0 and v[vi].nr in tot:
            vi -= 1
        _alt, zellen, keys = d.neue[k]
        d.neue[k] = (v[vi].nr if vi >= 0 else 0, zellen, keys)


def _neu_abgleich(soll: list, ist: list, spalten: Any = None) -> tuple[dict[int, int], dict[int, bool]]:
    """Ordnet neue Zeilen des Ist den neuen Zeilen des Soll zu.
    -> (soll_index -> ist_index, soll_index -> Ort stimmt).
    `spalten` (Nacharbeit Welle 2): "gleicher Inhalt" heisst gleiche Haupt-Spalten -- eine anders gefuellte Fuellspalte darf die
    Paarung nicht verfaelschen (ob die Fuellzelle stimmt, prueft die Messung je Zelle trotzdem). None = alle Spalten."""
    def gleich(a: tuple[str, ...], b: tuple[str, ...]) -> bool:
        if spalten:
            pa, pb = _proj(a, spalten), _proj(b, spalten)
            if any(pa) or any(pb):
                return pa == pb
        return a == b

    paar: dict[int, int] = {}
    ort: dict[int, bool] = {}
    frei = list(range(len(ist)))
    for si, (anker, _, keys) in enumerate(soll):                       # 1. gleicher Anker, gleicher Inhalt
        for ii in frei:
            if ist[ii][0] == anker and gleich(ist[ii][2], keys):
                paar[si], ort[si] = ii, True
                frei.remove(ii)
                break
    for si, (anker, _, keys) in enumerate(soll):                       # 2. gleicher Inhalt, anderer Ort
        if si in paar:
            continue
        for ii in frei:
            if gleich(ist[ii][2], keys):
                paar[si], ort[si] = ii, False
                frei.remove(ii)
                break
    for si, (anker, _, keys) in enumerate(soll):                       # 3. gleicher Anker, anderer Inhalt
        if si in paar:
            continue
        for ii in frei:
            if ist[ii][0] == anker:
                paar[si], ort[si] = ii, True
                frei.remove(ii)
                break
    return paar, ort


# ======================================================================================
# 3. Einheiten und Auswertung
# ======================================================================================
_TOKEN = re.compile(r"\s+|\S+")


_CODE_ERLAUBT = re.compile(r"^[a-z0-9_:.\-]{1,60}$")


def _code(c: Any) -> str:
    """Pruefcode fuer den Bericht: nur Maschinen-Kennungen (Kleinbuchstaben, Ziffern, _ : . -), nie Text (Regel 16)."""
    c = str(c).strip().lower()
    return c if _CODE_ERLAUBT.match(c) else "sonstiger_code"


def _wort_norm(wort: str) -> str:
    from .textutil import norm
    return norm(wort)


def _wort_liste(text: str) -> list[str]:
    return [w for w in (_wort_norm(x) for x in text.split()) if w]


def _wort_diff(alt: str, neu: str) -> tuple[list[str], list[str]]:
    """-> (entfallene Woerter, neue Woerter) als normalisierte Woerter."""
    a, b = _wort_liste(alt), _wort_liste(neu)
    entf: list[str] = []
    dazu: list[str] = []
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag in ("replace", "delete"):
            entf += a[i1:i2]
        if tag in ("replace", "insert"):
            dazu += b[j1:j2]
    return entf, dazu


@dataclass(frozen=True)
class Ziel:
    bi: int
    zeile: int | None
    spalte: int | None


@dataclass
class Fehler:
    """Ein Eintrag im Bericht (nur Positionen, nie Text)."""
    bi: int
    zeile: int
    spalte: int
    aenderungen: list[str] = field(default_factory=list)
    art: str = ""
    stumm: bool = False
    stille_arten: set[str] = field(default_factory=set)
    quelle: str = "soll"      # soll = Soll-Zelle verfehlt | auto = Ergebnis falsch | paar = ganzes Paar
    key: Any = None           # Schluessel der Einheit (fuer die Orte in nachher und Ergebnis)


_ZIEL_ZELLE = re.compile(r"^(.+)!([A-Z]+)(\d+)$")
_ZIEL_NEU = re.compile(r"^(.+)!Zeile (\d+) \(neu\)$")


class _Auswertung:
    """Vergleicht Ergebnis (E_auto) mit nachher und sammelt alles, was der Bericht braucht."""

    def __init__(self, v_blaetter, n_blaetter, e_blaetter, e_alle_blaetter, namen: list[str], aenderungen, plan,
                 pdf_analyse, bil_auto, bil_plan, idempotent: bool | None, luecken_erwartet: list[dict] | None,
                 vor_status: dict[str, str] | None = None, pdf_zeichen: set[str] | None = None,
                 typen: tuple[list, list] | None = None, markiert: bool = False,
                 pdf_specs: set[str] | None = None) -> None:
        self.namen = namen
        # nachher_format "markiert": die Zellen tragen Formatklassen (alt/neu/normal); Woerter zaehlen je Ansicht (ADR 0005)
        self.markiert = markiert
        self._polster = [self._polster_muster(b) for b in v_blaetter]
        # Datentypen (nachher, Ergebnis) je Blatt/Zeile/Spalte; None = nicht gemessen (Tests, die nur Texte liefern)
        self.TN, self.TE = typen if typen is not None else (None, None)
        self.typ_falsch_ist: set[Any] = set()     # Ist-Schluessel verwalteter Zellen, deren Datentyp vom Soll abweicht
        self.typ_stille: list[tuple[int, int, int]] = []    # (Blatt-Index, Zeile, Spalte): Typwechsel an unveraendertem Text, verwaltet
        self.pdf_zeichen = pdf_zeichen            # alle Zeichen, die im PDF vorkommen (None = unbekannt -> streng)
        self.vor_status = vor_status or {}      # Status laut Plan, VOR dem Anwenden (fuer Regel 27)
        self.bi_von_name = {n: i for i, n in enumerate(namen)}
        self.V = [_zeilen(b) for b in v_blaetter]
        self.N = [_zeilen(b) for b in n_blaetter] + [[] for _ in range(max(0, len(self.V) - len(n_blaetter)))]
        self.E = [_zeilen(b) for b in e_blaetter] + [[] for _ in range(max(0, len(self.V) - len(e_blaetter)))]
        self.vtext = [{z.nr: z.zellen for z in blatt} for blatt in self.V]
        self.aend = aenderungen
        self.plan = plan
        self._inhaltsspalte: dict[int, int | None] = {}
        self._konst_cache: dict[tuple, str | None] = {}
        self.e_alle = None
        self.e_alle_diffs: list[Diff] | None = None
        self.S: list[Diff] = []
        # Die Inhaltsspalte des Solls steht fest, BEVOR "leer" (verwaltete Spalten) gebraucht wird: erst ein Diff ohne `leer`,
        # daraus die Inhaltsspalte; dann das eigentliche Soll-Diff (self.S, mit `leer`).
        roh = [_diff(self.V[i], self.N[i]) for i in range(len(self.V))]
        self._inhaltsspalte = {i: self._soll_inhaltsspalte_berechnen(i, roh[i]) for i in range(len(self.V))}
        self.S = [_diff(self.V[i], self.N[i], self._leer_fn(i), self._paar_spalten(i)) for i in range(len(self.V))]
        self.I = [_diff(self.V[i], self.E[i], self._leer_fn(i), self._paar_spalten(i)) for i in range(len(self.V))]
        self._tot = [self.S[i].weg | self.S[i].geleert | self.I[i].weg | self.I[i].geleert for i in range(len(self.V))]
        for i in range(len(self.V)):         # Anker neuer Zeilen: entfernte/ausgeleerte Zeilen zaehlen auf beiden Seiten nicht (MA-3)
            _anker_neu(self.S[i], self.V[i], self._tot[i])
            _anker_neu(self.I[i], self.V[i], self._tot[i])
        for bi, s in enumerate(self.S):
            self._soll_pdf_blind(bi, s)
        # Zeilen fremder Specs (mehrere Specs in einem Register; nur die Spec des PDF darf sich aendern)
        self.fremd: list[set[int]] = [self._fremde_zeilen(bi, pdf_specs) for bi in range(len(self.V))]
        self.bil_auto, self.bil_plan = bil_auto, bil_plan
        self.idempotent = idempotent
        self.fehler: dict[Any, Fehler] = {}
        self.hinweise: list[str] = []
        if len(n_blaetter) != len(v_blaetter):
            self.hinweise.append("blattanzahl_vorher_nachher_verschieden")
        if len(e_blaetter) != len(v_blaetter):
            self.hinweise.append("blattanzahl_ergebnis_verschieden")

        # PDF-erkannte Woerter (nur im Speicher)
        self.pdf_neu: Counter[str] = Counter()
        self.pdf_alt: Counter[str] = Counter()
        for b in pdf_analyse.blocks:
            for t in b.tokens:
                if t.status in (ADD, MOD) and t.new:
                    self.pdf_neu.update(_wort_liste(t.new))
                if t.status in (DEL, MOD) and t.old:
                    self.pdf_alt.update(_wort_liste(t.old))
        # Woerter in unbekannter Farbe auf Seiten, fuer die das Tool eine Luecke "unbekannte_farbe" GEMELDET hat (nur im Speicher)
        gemeldet_seiten = {lk.get("seite") for lk in plan.luecken if lk.get("art") == "unbekannte_farbe"}
        self.farbe_gemeldet: Counter[str] = Counter()
        for b in pdf_analyse.blocks:
            for t in b.tokens:
                if "unbekannte_farbe" in (t.hinweis or "") and t.page in gemeldet_seiten:
                    self.farbe_gemeldet.update(_wort_liste(t.new or t.old))
        self.ae_woerter = {a.id: (Counter(_wort_liste(a.neu)), Counter(_wort_liste(a.alt))) for a in aenderungen}
        self.ziele_von = {a.id: self._ziele(a) for a in aenderungen}

        self._bauen()
        self.luecken_erwartet = luecken_erwartet
        if e_alle_blaetter is not None:
            alle = [_diff(self.V[i], _zeilen(e_alle_blaetter[i]) if i < len(e_alle_blaetter) else [], self._leer_fn(i),
                          self._paar_spalten(i))
                    for i in range(len(self.V))]
            for i, d in enumerate(alle):
                _anker_neu(d, self.V[i], self._tot[i] | d.weg | d.geleert)
            self.e_alle = self._ist_liste(alle)
            self.e_alle_diffs = alle

    # ------------------------------------------------------------------ Hilfen
    def _ziele(self, ae) -> list[Ziel]:
        out: list[Ziel] = []
        for op in ae.ops:
            bi = self.bi_von_name.get(op.blatt)
            if bi is None:
                continue
            if op.art in ("ersetzen", "einfuegen", "zelle_leeren"):
                out.append(Ziel(bi, op.zeile, op.spalte))
            elif op.art == "zeile_einfuegen":
                out.append(Ziel(bi, None, None))
            else:
                out.append(Ziel(bi, op.zeile, None))
        if not out and ae.ziel:
            m = _ZIEL_ZELLE.match(ae.ziel)
            if m and m.group(1) in self.bi_von_name:
                out.append(Ziel(self.bi_von_name[m.group(1)], int(m.group(3)), column_index_from_string(m.group(2))))
            else:
                m = _ZIEL_NEU.match(ae.ziel)
                if m and m.group(1) in self.bi_von_name:
                    out.append(Ziel(self.bi_von_name[m.group(1)], None, None))
        return out

    def _vorher(self, bi: int, zeile: int, spalte: int) -> str:
        z = self.vtext[bi].get(zeile, ())
        return z[spalte - 1] if spalte - 1 < len(z) else ""

    def _erwartet(self, bi: int, zeile: int, spalte: int) -> str:
        """Inhalt der Zelle laut nachher."""
        return self.S[bi].zellen.get((zeile, spalte), self._vorher(bi, zeile, spalte))

    def _tatsaechlich(self, bi: int, zeile: int, spalte: int) -> str:
        return self.I[bi].zellen.get((zeile, spalte), self._vorher(bi, zeile, spalte))

    def _typ_an(self, matrix: list | None, bi: int, zeile: int | None, spalte: int) -> str:
        if matrix is None or zeile is None or bi >= len(matrix):
            return "leer"
        try:
            return matrix[bi][zeile - 1][spalte - 1]
        except IndexError:
            return "leer"

    def _typ_abweichung(self, bi: int, zeile_nachher: int | None, zeile_ergebnis: int | None, spalte: int) -> bool:
        """Hat die Zelle im Ergebnis einen anderen Datentyp als im Soll (nachher)? Ohne Typdaten oder ohne Zeilenpaar: nein."""
        if self.TN is None or self.TE is None or zeile_nachher is None or zeile_ergebnis is None:
            return False
        return self._typ_an(self.TN, bi, zeile_nachher, spalte) != self._typ_an(self.TE, bi, zeile_ergebnis, spalte)

    # ------------------------------------------------------------------ Einheiten bauen
    def _bauen(self) -> None:
        # Soll-Einheiten: (Schluessel, bi, zeile, spalte, soll_key, art)
        self.einheiten: list[dict[str, Any]] = []
        self.soll_woerter: dict[Any, tuple[list[str], list[str]]] = {}
        self.neu_abgleich: list[tuple[dict[int, int], dict[int, bool]]] = []
        for bi, s in enumerate(self.S):
            paar, ort = _neu_abgleich(s.neue, self.I[bi].neue, self._paar_spalten(bi))
            self.neu_abgleich.append((paar, ort))
            for (z, c), inhalt in sorted(s.zellen.items()):
                key = ("z", bi, z, c)
                ist_da = (z, c) in self.I[bi].zellen
                ist = self.I[bi].zellen.get((z, c))
                erreicht = ist_da and _s(ist) == _s(inhalt)
                v = self._vorher(bi, z, c)
                art = "weg" if z in s.weg else "zelle"
                klasse = self._klasse(bi, c, art, v, inhalt)
                # gleicher Text, aber anderer Datentyp (Zahl 31 -> Text "31"): nicht erreicht (K-MESS2)
                typ_falsch = bool(erreicht and _s(inhalt) and self._typ_abweichung(bi, s.zuordnung.get(z), self.I[bi].zuordnung.get(z), c))
                if typ_falsch:
                    erreicht = False
                    if klasse == "haupt":
                        self.typ_falsch_ist.add(key)
                self.einheiten.append({"key": key, "bi": bi, "zeile": z, "spalte": c, "soll": inhalt, "ist_da": ist_da,
                                       "ist": ist, "erreicht": erreicht, "art": art, "typ_falsch": typ_falsch, "ist_key": key,
                                       "vorher": v, "verwaltet": self._verwaltet(bi, c),
                                       "klasse": klasse, "blind": ("z", z, c) in s.blind,
                                       "klasse2": self._klasse2(bi, c, art, klasse, v, inhalt, self._vorher_keys(bi, z)),
                                       **self._format_vergleich(ist if ist_da else None, inhalt)})
                self.soll_woerter[key] = _wort_diff(_ansicht(v, "neu"), _ansicht(inhalt, "neu"))
            for si, (anker, zellen, keys) in enumerate(s.neue):
                for c, k in enumerate(keys, 1):
                    if not k:
                        continue
                    key = ("n", bi, si, c)
                    ii = paar.get(si)
                    ist = None
                    erreicht = False
                    typ_falsch = False
                    ist_key = None
                    klasse = self._klasse(bi, c, "neu", "", zellen[c - 1])
                    if ii is not None:
                        ik = self.I[bi].neue[ii][2]
                        ist = self.I[bi].neue[ii][1][c - 1] if c - 1 < len(self.I[bi].neue[ii][1]) else ""
                        erreicht = ort[si] and c - 1 < len(ik) and ik[c - 1] == k
                        ist_key = ("n", bi, self.I[bi].neue[ii][0], ik, c)
                        if erreicht and self._typ_abweichung(bi, s.neue_nr[si], self.I[bi].neue_nr[ii], c):
                            typ_falsch, erreicht = True, False
                            if klasse == "haupt":
                                self.typ_falsch_ist.add(ist_key)
                    self.einheiten.append({"key": key, "bi": bi, "zeile": anker, "spalte": c, "soll": zellen[c - 1],
                                           "ist_da": ii is not None, "ist": ist, "erreicht": erreicht, "art": "neu",
                                           "typ_falsch": typ_falsch, "ist_key": ist_key,
                                           "vorher": "", "verwaltet": self._verwaltet(bi, c),
                                           "klasse": klasse, "si": si,
                                           "blind": ("n", si, c) in s.blind,
                                           "klasse2": self._klasse2(bi, c, "neu", klasse, "", zellen[c - 1], keys),
                                           **self._format_vergleich(ist if (ii is not None and ort[si]) else None, zellen[c - 1])})
                    self.soll_woerter[key] = ([], _wort_liste(_ansicht(zellen[c - 1], "neu")))

        # Ist-Aenderungen (Auto): jede veraenderte Zelle bzw. neue Zeile
        self.ist = self._ist_liste(self.I)
        for it in self.ist:                    # Text stimmt, Datentyp nicht (verwaltete Zelle): kein Treffer
            if it["key"] in self.typ_falsch_ist:
                it["richtig"] = False
        self.ist_woerter: dict[Any, tuple[Counter, Counter]] = {}
        for it in self.ist:
            if it["art"] == "neu":
                self.ist_woerter[it["key"]] = (Counter(), Counter(_wort_liste(_ansicht(it["inhalt"], "neu"))))
            else:
                entf, dazu = _wort_diff(_ansicht(self._vorher(it["bi"], it["zeile"], it["spalte"]), "neu"),
                                        _ansicht(it["inhalt"], "neu"))
                self.ist_woerter[it["key"]] = (Counter(entf), Counter(dazu))

    def _format_vergleich(self, ist: Any, soll: Any) -> dict[str, Any]:
        """Format-Treue einer Soll-Zelle (nur "markiert"): stimmt der Text, wird Zeichen fuer Zeichen die Formatklasse
        (normal/alt/neu) verglichen. -> text_gleich, format_gleich, zeichen, zeichen_gleich."""
        if not self.markiert or ist is None or not _s_text(soll) or _s_text(ist) != _s_text(soll):
            return {"text_gleich": False, "format_gleich": False, "zeichen": 0, "zeichen_gleich": 0}
        a, b = _klassen(ist), _klassen(soll)
        gleich = sum(1 for x, y in zip(a, b) if x == y)
        return {"text_gleich": True, "format_gleich": a == b, "zeichen": len(b), "zeichen_gleich": gleich}

    def _ist_liste(self, diffs: list[Diff]) -> list[dict[str, Any]]:
        """Alle Zellen, die im Ergebnis anders sind als vorher, mit der Angabe, ob sie laut nachher richtig sind
        (`richtig`: Text und -- bei "markiert" -- Formatklassen; `text_richtig`: nur der Text)."""
        out: list[dict[str, Any]] = []
        for bi, d in enumerate(diffs):
            s = self.S[bi]
            for (z, c), inhalt in sorted(d.zellen.items()):
                soll = s.zellen.get((z, c))
                richtig = soll is not None and _s(soll) == _s(inhalt)
                out.append({"key": ("z", bi, z, c), "bi": bi, "zeile": z, "spalte": c, "inhalt": inhalt,
                            "richtig": richtig, "text_richtig": soll is not None and _s_text(soll) == _s_text(inhalt),
                            "art": "weg" if z in d.weg else "zelle"})
            paar, ort = _neu_abgleich(s.neue, d.neue, self._paar_spalten(bi))
            rev = {ii: si for si, ii in paar.items()}
            for ii, (anker, zellen, keys) in enumerate(d.neue):
                si = rev.get(ii)
                for c, k in enumerate(keys, 1):
                    if not k:
                        continue
                    passt = si is not None and ort[si] and c - 1 < len(s.neue[si][2])
                    richtig = passt and s.neue[si][2][c - 1] == k
                    out.append({"key": ("n", bi, anker, keys, c), "bi": bi, "zeile": anker, "spalte": c,
                                "inhalt": zellen[c - 1], "richtig": richtig, "art": "neu", "ist_index": ii,
                                "text_richtig": passt and _ohne_format(s.neue[si][2][c - 1]) == _ohne_format(k)})
        return out

    # ------------------------------------------------------------------ Fehler eintragen
    def _fehler(self, key: Any, bi: int, zeile: int, spalte: int, quelle: str) -> Fehler:
        f = self.fehler.get(key)
        if f is None:
            f = self.fehler[key] = Fehler(bi, zeile, spalte, quelle=quelle, key=key)
        return f

    def orte(self, key: Any, zeile: int = 0) -> tuple[int | None, int | None, bool]:
        """(Zeile in nachher, Zeile im Ergebnis, neue Zeile?) zu einem Fehler-Schluessel. `zeile` im Bericht ist die Zeile in
        vorher (bei neuen Zeilen die Zeile davor); nach eingefuegten Zeilen liegen Soll und Ist woanders -- diese Angaben nennen
        die Zeilen, an denen man Soll und Ist direkt vergleichen kann (None = gibt es dort nicht)."""
        if not isinstance(key, tuple) or not key:
            return None, None, False
        if key[0] in ("z", "typ") and len(key) == 4:
            _, bi, z, _c = key
            return self.S[bi].zuordnung.get(z), self.I[bi].zuordnung.get(z), False
        if key[0] == "n" and len(key) == 4:                        # Soll-Einheit einer neuen Zeile
            _, bi, si, _c = key
            ii = self.neu_abgleich[bi][0].get(si)
            return self.S[bi].neue_nr[si], (self.I[bi].neue_nr[ii] if ii is not None else None), True
        if key[0] == "n" and len(key) == 5:                        # neue Zeile im Ergebnis
            bi = key[1]
            it = next((i for i in self.ist if i["key"] == key), None)
            if it is None:
                return None, None, True
            rev = {ii: si for si, ii in self.neu_abgleich[bi][0].items()}
            si = rev.get(it["ist_index"])
            return (self.S[bi].neue_nr[si] if si is not None else None), self.I[bi].neue_nr[it["ist_index"]], True
        if key[0] == "r27" and zeile > 0:                          # Regel 27: Ort der Wirkung (Zeile in vorher)
            f = self.fehler.get(key)
            bi = f.bi if f is not None else 0
            return self.S[bi].zuordnung.get(zeile), self.I[bi].zuordnung.get(zeile), False
        return None, None, False

    # ------------------------------------------------------------------ Zuordnung Aenderung <-> Zelle
    def _zugehoerig(self, W: Counter, R: Counter) -> list:
        out = []
        for ae in self.aend:
            an, aa = self.ae_woerter[ae.id]
            gesamt = sum(an.values()) + sum(aa.values())
            ueber = sum((an & W).values()) + sum((aa & R).values())
            if gesamt and ueber and ueber / gesamt >= 0.4:
                out.append(ae)
        return out

    @staticmethod
    def _deckung(unit_dazu: Counter, unit_entf: Counter, ist_entf: Counter, ist_dazu: Counter) -> float:
        """Wie viel vom Soll-Inhalt einer Einheit steckt in einer Ist-Aenderung? Massgeblich sind die NEUEN Woerter;
        gibt es keine (reine Streichung), zaehlen die entfallenen."""
        if unit_dazu:
            return sum((unit_dazu & ist_dazu).values()) / sum(unit_dazu.values())
        if unit_entf:
            return sum((unit_entf & ist_entf).values()) / sum(unit_entf.values())
        return 0.0

    def _landung(self, W: Counter, R: Counter, ausser: Any):
        """Wo im Ergebnis sind die Woerter gelandet (veraenderte Zelle, die den Soll-Inhalt enthaelt)?"""
        best, best_w = None, (0.0, 0)
        for it in self.ist:
            if it["key"] == ausser:
                continue
            entf, dazu = self.ist_woerter[it["key"]]
            d = self._deckung(W, R, entf, dazu)
            wert = (d, -(sum(dazu.values()) + sum(entf.values())))
            if d >= 0.6 and wert > best_w:
                best, best_w = it, wert
        return (best["bi"], best["zeile"], best["spalte"]) if best else None

    def _soll_ort(self, ist_dazu: Counter, ist_entf: Counter):
        """Welche Soll-Einheit passt zu dieser veraenderten Ist-Zelle (deren Soll-Inhalt steckt in ihr)?"""
        best, best_w = None, (0.0, 0)
        for e in self.einheiten:
            entf, dazu = self.soll_woerter[e["key"]]
            d = self._deckung(Counter(dazu), Counter(entf), ist_entf, ist_dazu)
            wert = (d, -(len(dazu) + len(entf)))
            if d >= 0.6 and wert > best_w:
                best, best_w = e, wert
        return (best["bi"], best["zeile"], best["spalte"]) if best else None

    def _leer_fn(self, bi: int) -> Callable[[Zeile], bool]:
        """Eine Zeile ist "leer", wenn in den verwalteten Spalten nichts steht (nur Bemerkung/Prio o. ae. uebrig)."""
        def leer(z: Zeile) -> bool:
            return not any(k for c, k in enumerate(z.keys, 1) if k and self._verwaltet(bi, c))
        return leer

    def _soll_inhaltsspalte(self, bi: int) -> int | None:
        return self._inhaltsspalte.get(bi)

    def _soll_inhaltsspalte_berechnen(self, bi: int, s: "Diff") -> int | None:
        """Spalte mit den meisten veraenderten Zeichen laut Soll (Inhaltsspalte im Soll)."""
        gewicht: Counter[int] = Counter()
        for (z, c), inhalt in s.zellen.items():
            gewicht[c] += len(inhalt) + len(self._vorher(bi, z, c))
        for _anker, zellen, _keys in s.neue:
            for c, t in enumerate(zellen, 1):
                gewicht[c] += len(t)
        return gewicht.most_common(1)[0][0] if gewicht else None

    def _layout_von(self, bi: int) -> dict | None:
        name = self.namen[bi] if bi < len(self.namen) else ""
        return next((l for l in self.plan.layout if l.get("blatt") == name), None)

    def _verwaltet(self, bi: int, spalte: int) -> bool:
        """Soll die Spalte vom Tool bearbeitet werden? Nur wenn das erkannte Layout zum Soll passt (Inhaltsspalte stimmt),
        gelten die uebrigen Spalten (z.B. Bemerkung) als "nicht verwaltet" und zaehlen nicht als Soll-Zelle.
        Bei unklarem oder abweichendem Layout zaehlt jede Spalte (nichts wird versteckt)."""
        lay = self._layout_von(bi)
        if not lay or not lay.get("inhalt_spalte"):
            return True
        if column_index_from_string(lay["inhalt_spalte"]) != self._soll_inhaltsspalte(bi):
            return True
        spalten = {column_index_from_string(lay[k]) for k in ("section_spalte", "titel_spalte", "inhalt_spalte") if lay.get(k)}
        return spalte in spalten

    def _fremde_zeilen(self, bi: int, pdf_specs: set[str] | None) -> set[int]:
        """vorher-Zeilen, die zu einer FREMDEN Spec gehoeren -- nur aus vorher <-> nachher und dem PDF abgeleitet (auch fuer echte
        Paare): Spec-Spalte ist die erste Spalte, deren Wert in allen laut Soll geaenderten Zeilen dieselbe Spec-Nummer des PDF
        ist (Schreibweise vereinheitlicht: "11_22_33", "11 22 33 "). Fremd sind Bloecke von mindestens SPEC_MIN_BLOCK
        aufeinanderfolgenden Zeilen mit einer anderen Spec-Nummer in dieser Spalte. Ohne Spec-Nummer im PDF: keine."""
        if not pdf_specs:
            return set()
        geaendert = {z for (z, _c) in self.S[bi].zellen}
        zeilen = self.V[bi]
        if not geaendert or not zeilen:
            return set()
        breite = max(len(z.zellen) for z in zeilen)
        for c in range(breite):
            werte = {z.nr: (_spec_norm(z.zellen[c]) if c < len(z.zellen) else "") for z in zeilen}
            ziel = {werte.get(z, "") for z in geaendert}
            if len(ziel) != 1 or next(iter(ziel)) not in pdf_specs:
                continue
            v = next(iter(ziel))
            fremd: set[int] = set()
            lauf: list[int] = []
            letzter = None
            for z in zeilen:                                      # Laeufe gleicher fremder Nummer
                w = werte[z.nr]
                if lauf and w == letzter:
                    lauf.append(z.nr)
                else:
                    if len(lauf) >= SPEC_MIN_BLOCK:
                        fremd.update(lauf)
                    lauf = [z.nr] if (w and w != v and _SPEC_NR.fullmatch(w)) else []
                letzter = w
            if len(lauf) >= SPEC_MIN_BLOCK:
                fremd.update(lauf)
            return fremd
        return set()

    _POLSTER = re.compile(r"^(?:\d{1,3}|[A-Za-z]{1,3})[.)]([\t  ]* [\t  ]*)")

    @classmethod
    def _polster_muster(cls, blatt: list[list[str]]) -> set[str]:
        """Polster zwischen Listennummer und Text (mit geschuetzten Leerzeichen), die in vorher am Zellanfang vorkommen
        (z. B. "a." + 15 geschuetzte Leerzeichen + " "). Ein Mensch uebernimmt sie fuer eine neue Zeile von der Nachbarzeile;
        sie sind also aus der Excel ableitbar, auch wenn das PDF sie nicht traegt."""
        out: set[str] = set()
        for zeile in blatt:
            for c in zeile:
                m = cls._POLSTER.match(str(c)) if c else None
                if m:
                    out.add(m.group(1))
        return out

    def _blind(self, soll: str, vorher: str, polster: set[str] | None = None) -> str:
        """Das Soll so, wie ein PDF-Leser es hoechstens sehen kann: geschuetzte Leerzeichen und unsichtbare Zeichen, die im
        PDF GAR NICHT vorkommen, koennen in NEUEM Text nicht uebernommen werden (Text, der schon in `vorher` stand, bleibt).
        Nur diese Richtung: steht so ein Zeichen im ERGEBNIS, aber nicht im Soll, bleibt es ein Fehler (RT-135).
        Ausnahme: ein Listen-Polster am Zellanfang, das genau so schon in vorher vorkommt (`_polster_muster`), bleibt streng.
        Formatklassen einer markierten Zelle bleiben erhalten."""
        if self.pdf_zeichen is None or not any(ch in _PDF_BLIND for ch in soll):
            return soll
        blind = {ch for ch in set(soll) if ch in _PDF_BLIND and ch not in self.pdf_zeichen}
        if not blind:
            return soll
        schutz = range(0)
        m = self._POLSTER.match(str(soll))
        if m and polster and m.group(1) in polster:
            schutz = range(m.start(1), m.end(1))
        kl = getattr(soll, "klassen", None)
        text: list[str] = []
        klassen: list[str] = []
        for tag, _i1, _i2, j1, j2 in SequenceMatcher(None, str(vorher), str(soll), autojunk=False).get_opcodes():
            for j in range(j1, j2):
                ch = soll[j]
                if tag != "equal" and ch in blind and j not in schutz:
                    ch = " " if ch in _LEER_SPEZIAL else ""
                if ch:
                    text.append(ch)
                    klassen.append(kl[j] if kl else "-")
        neu = "".join(text)
        if neu == str(soll):
            return soll
        return _FText(neu, "".join(klassen)) if kl else neu

    def _soll_pdf_blind(self, bi: int, s: "Diff") -> None:
        polster = self._polster[bi] if bi < len(self._polster) else set()
        for (z, c), t in list(s.zellen.items()):
            p = self._blind(t, self._vorher(bi, z, c), polster)
            if p is not t:
                s.zellen[(z, c)] = p
                s.blind.add(("z", z, c))
        for si, (anker, zellen, keys) in enumerate(list(s.neue)):
            neu = tuple(self._blind(t, "", polster) for t in zellen)
            if any(a is not b for a, b in zip(zellen, neu)):
                s.neue[si] = (anker, neu, tuple(_s(t) for t in neu))
                s.blind.update(("n", si, c) for c, (a, b) in enumerate(zip(zellen, neu), 1) if a is not b)

    def _klasse(self, bi: int, spalte: int, art: str, vorher: str, soll: str) -> str:
        """Einordnung einer Soll-Zelle (RT-142):
        formel    = Formelzelle (steht nie im PDF; verschobene Formeln sind eine echte Gefahr),
        haupt     = verwaltete Spalte (Nummer/Titel/Inhalt): zaehlt bei Abdeckung, Fehlern und Lese-Luecke,
        harmlos   = nicht verwaltete Fuellung einer neuen/entfernten Zeile (kurze Bemerkung/Prio o. ae.),
        sonstige  = alles andere in nicht verwalteten Spalten (z. B. neu nummerierte Spalte, lange Kommentare)."""
        if soll.lstrip().startswith("=") or vorher.lstrip().startswith("="):
            return "formel"
        if self._verwaltet(bi, spalte):
            return "haupt"
        if art in ("neu", "weg") and max(len(soll.split()), len(vorher.split())) <= FUELL_MAX_WOERTER:
            return "harmlos"
        return "sonstige"

    # ---- Klassen der Soll-Zellen fuer Etappe 2 (Messung v2) ----------------------------------------------------
    def _verwaltete_spalten(self, bi: int) -> list[int] | None:
        """Haupt-Spalten (Nummer/Titel/Inhalt) laut erkanntem Aufbau -- nur wenn er zum Soll passt; sonst None."""
        lay = self._layout_von(bi)
        if not lay or not lay.get("inhalt_spalte"):
            return None
        if column_index_from_string(lay["inhalt_spalte"]) != self._soll_inhaltsspalte(bi):
            return None
        return sorted({column_index_from_string(lay[k]) for k in ("section_spalte", "titel_spalte", "inhalt_spalte") if lay.get(k)})

    def _paar_spalten(self, bi: int) -> set[int] | None:
        """Haupt-Spalten fuer die Zeilen-Paarung (None = alle Spalten, wenn der Aufbau unklar ist oder nicht zum Soll passt)."""
        sp = self._verwaltete_spalten(bi)
        return set(sp) if sp else None

    def _vorher_keys(self, bi: int, zeile: int) -> tuple[str, ...]:
        return tuple(_s(c) for c in self.vtext[bi].get(zeile, ()))

    def _zeilentyp(self, bi: int, keys: tuple[str, ...]) -> tuple[bool, ...] | None:
        """Typ einer Zeile = welche Haupt-Spalten belegt sind (z. B. Section-Kopf: Nummer+Titel(+Inhalt), Absatz: nur Inhalt)."""
        sp = self._verwaltete_spalten(bi)
        if sp is None:
            return None
        return tuple(bool(keys[c - 1]) if c - 1 < len(keys) else False for c in sp)

    def _konstante(self, bi: int, spalte: int, typ: tuple[bool, ...]) -> str | None:
        """Hat die Spalte in vorher fuer ALLE Zeilen dieses Typs denselben, nicht leeren Wert (mindestens 2 Zeilen)?
        Die Kopfzeile der Tabelle (erste Zeile mit mindestens zwei belegten Haupt-Spalten) zaehlt nicht mit."""
        k = (bi, spalte, typ)
        if k in self._konst_cache:
            return self._konst_cache[k]
        sp = self._verwaltete_spalten(bi) or []
        zeilen = [z for z in self.V[bi] if any(c - 1 < len(z.keys) and z.keys[c - 1] for c in sp)]
        kopf = next((z.nr for z in zeilen if sum(1 for c in sp if c - 1 < len(z.keys) and z.keys[c - 1]) >= 2), None)
        werte = [(z.keys[spalte - 1] if spalte - 1 < len(z.keys) else "") for z in zeilen
                 if z.nr != kopf and self._zeilentyp(bi, z.keys) == typ]
        wert = werte[0] if len(werte) >= 2 and len(set(werte)) == 1 and werte[0] else None
        self._konst_cache[k] = wert
        return wert

    def _klasse2(self, bi: int, spalte: int, art: str, klasse: str, vorher: str, soll: str, zeilen_keys: tuple[str, ...]) -> str:
        """Klassen der Soll-Zellen fuer Etappe 2 (Messung v2):
        haupt           = verwaltete Spalte (Nummer/Titel/Inhalt),
        formel          = Formelzelle,
        fuell_konstant  = Zelle einer neuen/entfernten Zeile in einer Spalte, die in vorher fuer alle Zeilen desselben Typs denselben
                          Wert hat (prueft die Messung selbst) -- ein Mensch uebernimmt den Wert von der Nachbarzeile,
        nicht_ableitbar = alles andere (Kommentar, laufende Nummer, geaenderte Zelle ausserhalb der Haupt-Spalten ...).
        Ableitbar (aus PDF + Excel herleitbar) sind haupt + fuell_konstant + formel."""
        if klasse in ("formel", "haupt"):
            return klasse
        if art not in ("neu", "weg"):
            return "nicht_ableitbar"
        typ = self._zeilentyp(bi, zeilen_keys)
        if typ is None:
            return "nicht_ableitbar"
        konst = self._konstante(bi, spalte, typ)
        if konst is None:
            return "nicht_ableitbar"
        wert = _s(soll) if art == "neu" else _s(vorher)
        return "fuell_konstant" if wert == konst else "nicht_ableitbar"

    def richtig_nach_freigabe(self, ae: Any) -> bool | None:
        """Waere diese (gesperrte) Aenderung nach der Freigabe richtig gewesen? Grundlage: die Excel mit ALLEN freigegebenen
        Vorschlaegen (--mit-vorschlaegen). None = nicht messbar (ohne Vorschlags-Excel oder ohne Ziel)."""
        if self.e_alle_diffs is None or self.e_alle is None or not ae.ops:
            return None
        for op in ae.ops:
            bi = self.bi_von_name.get(op.blatt)
            if bi is None:
                return None
            if op.art in ("ersetzen", "einfuegen", "zelle_leeren"):
                ist = self.e_alle_diffs[bi].zellen.get((op.zeile, op.spalte), self._vorher(bi, op.zeile, op.spalte))
                if _s(ist) != _s(self._erwartet(bi, op.zeile, op.spalte)):
                    return False
            elif op.art == "zeile_einfuegen":
                gesucht = {c: _s_text(t) for c, t in op.zellen.items() if _s_text(t)}
                treffer = [it for it in self.e_alle if it["art"] == "neu" and it["bi"] == bi and it["key"][3] and
                           all(c - 1 < len(it["key"][3]) and _ohne_format(it["key"][3][c - 1]) == t for c, t in gesucht.items())]
                if not treffer or not all(it["richtig"] for it in treffer):
                    return False
            else:                                   # Zeile loeschen/leeren: laut Soll muss die Zeile weg bzw. leer sein
                s = self.S[bi]
                if op.zeile not in s.weg and op.zeile not in s.geleert:
                    return False
        return True

    @staticmethod
    def _unerklaert(ist: str, soll: str, erlaubt: Counter) -> Counter:
        """Woerter, in denen sich `ist` und `soll` unterscheiden und die NICHT zu den erlaubten (offenen) Aenderungen gehoeren.
        Leer = die Zelle weicht nur um offene Aenderungen ab. Ein Unterschied ohne Wortunterschied (Leerraum, unsichtbare
        Zeichen, Satzzeichen am Rand) ist nie "erklaert".
        Markierte Zellen: verglichen werden beide Ansichten (alt = normal + alt, neu = normal + neu) je fuer sich -- eine offene
        Aenderung fehlt dann in der Neu-Ansicht als neues Wort und in der Alt-Ansicht gar nicht. Stimmen beide Ansichten und weicht
        die Zelle trotzdem ab (Reihenfolge oder Markierung an gleichem Text), ist das nie erklaert."""
        if _s(ist) == _s(soll):
            return Counter()
        if getattr(ist, "klassen", None) or getattr(soll, "klassen", None):
            out: Counter = Counter()
            for welche in ("neu", "alt"):
                out.update(_Auswertung._unerklaert_text(_ansicht(ist, welche), _ansicht(soll, welche), erlaubt))
            if not out and _s(_ansicht(ist, "neu")) == _s(_ansicht(soll, "neu")) \
                    and _s(_ansicht(ist, "alt")) == _s(_ansicht(soll, "alt")):
                out["\x00format"] += 1
            return out
        return _Auswertung._unerklaert_text(ist, soll, erlaubt)

    @staticmethod
    def _unerklaert_text(ist: str, soll: str, erlaubt: Counter) -> Counter:
        """`_unerklaert` fuer Texte ohne Formatklassen."""
        if _s(ist) == _s(soll):
            return Counter()
        a, b = _TOKEN.findall(_s(ist)), _TOKEN.findall(_s(soll))      # Woerter UND Leerraum als eigene Stuecke
        rest = Counter(erlaubt)
        out: Counter = Counter()
        for tag, i1, i2, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
            if tag == "equal":
                continue
            if any(t.isspace() and t not in (" ", "\n") for t in a[i1:i2] + b[j1:j2]):       # doppelter/geschuetzter Leerraum
                out["\x00format"] += 1
            woerter = Counter(w for w in (_wort_norm(t) for t in a[i1:i2] + b[j1:j2] if not t.isspace()) if w)
            if not woerter:                                  # nur Leerraum/Zeichen verschieden: nie durch eine Aenderung erklaert
                out["\x00format"] += 1
                # die Nachbarwoerter nennen, damit die schuldige Aenderung gefunden wird
                nachbarn = [a[i1 - 1]] if i1 > 0 else []
                nachbarn += [a[i2]] if i2 < len(a) else []
                out.update(w for w in (_wort_norm(t) for t in nachbarn if not t.isspace()) if w)
                continue
            out.update(woerter - rest)
            rest -= woerter
        return out

    @staticmethod
    def _locker_gleich(a: Any, b: Any) -> bool:
        """Lockerer Vergleich; markierte Zellen je Ansicht (eine andere Markierung ist nie nur "Textform")."""
        if getattr(a, "klassen", None) or getattr(b, "klassen", None):
            return all(_k(_ansicht(a, w)) == _k(_ansicht(b, w)) for w in ("alt", "neu"))
        return _k(a) == _k(b)

    def _nur_textform_key(self, key: Any) -> bool:
        """Unterscheiden sich Ist und Soll NUR in Leerraum/unsichtbaren Zeichen/Absatzgrenzen (lockerer Vergleich waere gleich)?"""
        if key[0] == "z":
            _, bi, z, c = key
            return self._locker_gleich(self._tatsaechlich(bi, z, c), self._erwartet(bi, z, c))
        if key[0] == "n" and len(key) == 5:
            _, bi, _anker, _keys, c = key
            it = next((i for i in self.ist if i["key"] == key), None)
            if it is None:
                return False
            paar, _ort = _neu_abgleich(self.S[bi].neue, self.I[bi].neue, self._paar_spalten(bi))
            si = {ii: s for s, ii in paar.items()}.get(it["ist_index"])
            if si is None:
                return False
            soll = self.S[bi].neue[si][1]
            return c - 1 < len(soll) and self._locker_gleich(it["inhalt"], soll[c - 1])
        return False

    def _wirkung_vorhanden(self, ae: Any, vorher: str, ergebnis: str, angewendet: list[str] | tuple[str, ...] = ()) -> bool:
        """Ist die Wirkung dieser (offenen) Aenderung im Ergebnistext der Zelle schon zu sehen? (neue Woerter stehen da bzw.
        gestrichene Woerter sind weg.) Nur sicher erkennbar, wenn die Aenderung Woerter bringt oder nimmt, die vorher anders waren.
        Nacharbeit Welle 2: Gemessen wird nur, was UEBER die Wirkung der angewendeten Aenderungen derselben Zelle (`angewendet`)
        hinausgeht -- fuegt eine angewendete Nachbar-Aenderung zufaellig dasselbe Wort ein, ist das keine Wirkung der offenen.
        Bringt die offene Aenderung neue Woerter UND nimmt alte weg, muessen beide Seiten zu sehen sein."""
        neu, alt = self.ae_woerter[ae.id]
        vw, ew = Counter(_wort_liste(_ansicht(vorher, "neu"))), Counter(_wort_liste(_ansicht(ergebnis, "neu")))
        ohne = Counter(vw)                       # Wortbestand, wenn NUR die angewendeten Aenderungen gewirkt haetten
        for i in angewendet:
            if i == ae.id or i not in self.ae_woerter:
                continue
            a_neu, a_alt = self.ae_woerter[i]
            ohne = (ohne - (a_alt & vw)) + a_neu
        zusatz, fehlt = ew - ohne, ohne - ew
        gebraucht = neu - vw
        weg = (alt - neu) & vw
        if not gebraucht and not weg:
            return False
        if gebraucht and (gebraucht - zusatz):
            return False
        if weg and (weg - fehlt):
            return False
        return True

    def _nur_gemeldete_farbe(self, ist: str, soll: str, vorher: str, offene: list) -> bool:
        """Nacharbeit Welle 2 (eng gefasst): Die Zelle weicht vom Soll NUR darin ab, dass NEUE Woerter fehlen, die das PDF in einer
        unbekannten Farbe zeigt (z. B. gruene Schrift), und fuer deren Seite das Tool die Luecke "unbekannte_farbe" gemeldet hat --
        sonst nur um offene ("pruefen") Aenderungen. Dann hat das Tool die Abweichung gemeldet (Luecke bzw. gesperrter Absatz); sie ist
        kein stiller Fehler. Nie entlastet werden: Leerraum-/Zeichenform-Unterschiede, Woerter, die zu viel im Ergebnis stehen
        (ausser alten Woertern offener Aenderungen), fehlende Woerter, die schon in vorher standen (das Tool hat etwas geloescht),
        fehlende Woerter ohne gemeldete Farb-Luecke und falsche Reihenfolge."""
        if not self.farbe_gemeldet or _s(ist) == _s(soll):
            return False
        erlaubt = self._offene_woerter(offene)
        gesperrt = any(a.status == PRUEFEN and any("unbekannte_farbe" in str(c) for c in (a.pruefcodes or [])) for a in offene)
        if "\x00format" in self._unerklaert(ist, soll, Counter(erlaubt) + self.farbe_gemeldet):
            return False
        a, b = _wort_liste(_ansicht(ist, "neu")), _wort_liste(_ansicht(soll, "neu"))
        neu_im_soll = Counter(b) - Counter(_wort_liste(_ansicht(vorher, "neu")))
        rest_offen = Counter(erlaubt)
        rest_farbe = Counter(self.farbe_gemeldet) & neu_im_soll
        farbe_fehlt = 0
        for tag, i1, i2, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
            if tag == "equal":
                continue
            for w in a[i1:i2]:                    # steht zu viel im Ergebnis: nur alte Woerter offener Aenderungen
                if rest_offen[w] <= 0:
                    return False
                rest_offen[w] -= 1
            stueck_farbe: list[str] = []
            for w in b[j1:j2]:                    # fehlt im Ergebnis: offene Aenderung oder neues Wort in gemeldeter Farbe
                if rest_farbe[w] > 0:
                    rest_farbe[w] -= 1
                    stueck_farbe.append(w)
                elif rest_offen[w] > 0:
                    rest_offen[w] -= 1
                else:
                    return False
            # Ist der Absatz NICHT wegen der unbekannten Farbe gesperrt, zaehlt ein fehlendes Stueck nur als "Farbe", wenn mindestens
            # ein Wort darin NICHT auch blau (neu) im PDF steht -- sonst koennte es ein verpasstes Wort einer blauen Aenderung sein.
            if stueck_farbe and not gesperrt and all(self.pdf_neu[w] > 0 for w in stueck_farbe):
                return False
            farbe_fehlt += len(stueck_farbe)
        return farbe_fehlt > 0

    def _neu_als_gepaarte_zeile(self, bi: int, gesucht: dict[int, str]) -> bool:
        """Nacharbeit Welle 2 (Paar 0136): Eine eingefuegte Zeile kann von der Zeilen-Ausrichtung (difflib + Aehnlichkeit) als
        VERAENDERTE alte Zeile gepaart werden, wenn sie einer alten Zeile aehnelt. Steht ihr Inhalt in Ergebnis UND Soll an derselben
        gepaarten Stelle (Soll und Ergebnis wurden gleich ausgerichtet), ist das keine fehlende Wirkung. Eng: alle Zellen der Operation
        muessen dort im Ergebnis stehen UND laut Soll genau so stehen; sonst bleibt es ein stiller Fehler."""
        s, i = self.S[bi], self.I[bi]
        for vz in set(i.zuordnung) & set(s.zuordnung):
            if all((vz, c) in i.zellen and _s_text(i.zellen[(vz, c)]) == t and (vz, c) in s.zellen
                   and _s_text(s.zellen[(vz, c)]) == t for c, t in gesucht.items()):
                return True
        return False

    def _offene_woerter(self, aend_liste: list) -> Counter:
        erlaubt: Counter = Counter()
        for ae in aend_liste:
            neu, alt = self.ae_woerter[ae.id]
            erlaubt.update(neu)
            erlaubt.update(alt)
        return erlaubt

    def lese_luecke_fehlend(self) -> Counter:
        """Woerter, die laut Soll im PDF stehen muessten (nur Haupt-Zellen, keine Formeln), aber nicht erkannt wurden."""
        fehlend: Counter = Counter()
        for e in self.einheiten:
            if e["klasse"] != "haupt":
                continue
            entf, dazu = self.soll_woerter[e["key"]]
            fehlend.update((Counter(dazu) - self.pdf_neu) + (Counter(entf) - self.pdf_alt))
        return fehlend

    def _layout_abweichung(self, bi: int, spalte: int) -> bool:
        """Erkannte Inhaltsspalte weicht von der Soll-Inhaltsspalte ab, und die Zelle liegt in der Soll-Inhaltsspalte."""
        lay = self._layout_von(bi)
        if lay is None:
            return False
        soll_sp = self._soll_inhaltsspalte(bi)
        if soll_sp is None or spalte != soll_sp:
            return False
        inh = lay.get("inhalt_spalte")
        return (not inh) or column_index_from_string(inh) != soll_sp

    # ------------------------------------------------------------------ Fehlerart-Kaskade
    def _art(self, bi: int, zeile: int, spalte: int, W: Counter, R: Counter, e_text: str | None, n_text: str | None,
             in_e_geaendert: bool, ausser: Any, ist_neu: bool = False) -> str:
        if (W - self.pdf_neu) or (R - self.pdf_alt):
            return "pdf_falsch_gelesen"
        zu = self._zugehoerig(W, R)
        ziele = [z for ae in zu for z in self.ziele_von[ae.id]]
        landung = self._landung(W, R, ausser)
        # Layout: falsches Blatt/falsche Spalte
        if self._layout_abweichung(bi, spalte):
            return "layout_falsch"
        if landung and (landung[0] != bi or landung[2] != spalte):
            return "layout_falsch"
        if ziele and not any(z.bi == bi and z.spalte in (None, spalte) for z in ziele):
            return "layout_falsch"
        # Zuordnung: richtige Spalte, falsche Zeile/falscher Absatz
        if landung and landung[1] != zeile:
            return "zuordnung_falsch"
        if ziele and not any(z.bi == bi and z.spalte in (None, spalte) and z.zeile in (None, zeile) for z in ziele):
            return "zuordnung_falsch"
        if e_text is not None and n_text is not None and _k(e_text) != _k(n_text) \
                and sorted(_wort_liste(e_text)) == sorted(_wort_liste(n_text)):
            return "zuordnung_falsch"          # dieselben Woerter, aber an anderer Stelle (Absatz vertauscht)
        # Plan: Ziel stimmt, Ergebnis nicht
        if in_e_geaendert or any(z.bi == bi and z.zeile in (None, zeile) and z.spalte in (None, spalte) for z in ziele):
            return "plan_falsch"
        return "unklar"

    def _art_fremd(self, it: dict[str, Any], ae_ids: list[str]) -> str:
        """Fehlerart fuer eine Zelle im Ergebnis, die laut Soll gar nicht (so) veraendert sein sollte."""
        entf, dazu = self.ist_woerter[it["key"]]
        ort = self._soll_ort(dazu, entf)      # (Ist-Woerter: neu, entfallen)
        if ort is None:
            return "pdf_falsch_gelesen" if ae_ids else "unklar"
        if ort[0] != it["bi"] or ort[2] != it["spalte"]:
            return "layout_falsch"
        if ort[1] != it["zeile"]:
            return "zuordnung_falsch"
        return "plan_falsch"

    # ------------------------------------------------------------------ Hauptauswertung
    def _schuldige(self, ids: list[str], unerklaert: Counter) -> list[str]:
        """Welche der angewendeten Aenderungen einer Zelle haengen an den unerklaerten Woertern? (Je Aenderung bewerten.)
        Findet sich keine, bleiben alle verdaechtig."""
        treffer = [i for i in ids if any(w in unerklaert for w in (*self.ae_woerter[i][0], *self.ae_woerter[i][1]))]
        return treffer or list(ids)

    def datentyp_orte_verwaltet(self, orte: list[tuple[int, int, int]]) -> list[tuple[int, int, int]]:
        """Welche Zellen mit Datentyp-Wechsel (unveraenderter Text, siehe `_struktur_befund`) liegen in einer VERWALTETEN Spalte
        (und sind keine Formelzellen)? Dort hat das Tool gearbeitet: stiller Fehler. Alle anderen sind Strukturbefunde."""
        return [(bi, r, c) for bi, r, c in orte
                if self._verwaltet(bi, c) and not self._vorher(bi, r, c).lstrip().startswith("=")]

    def auswerten(self) -> dict[str, Any]:
        typ_stille = self.typ_stille
        nach_id = {a.id: a for a in self.aend}
        # --- offene Aenderungen (pruefen / nicht uebernommen) je Ziel-Zelle bzw. Zeile ----------------
        offen_ziel: dict[tuple[int, int, int], list] = {}
        offen_zeile: dict[tuple[int, int], list] = {}
        for ae in self.aend:
            if ae.status == UEBERNOMMEN:
                continue
            for z in self.ziele_von[ae.id]:
                if z.zeile is None:
                    continue
                if z.spalte is not None:
                    offen_ziel.setdefault((z.bi, z.zeile, z.spalte), []).append(ae)
                else:
                    offen_zeile.setdefault((z.bi, z.zeile), []).append(ae)

        # --- Ziele der "uebernommen"-Aenderungen (nach dem Anwenden) -------------------------------
        ziel_zellen: dict[tuple[int, int, int], list[str]] = {}
        neu_ops: list[tuple[str, int, dict[int, str], str]] = []     # (Aenderung, bi, Zellen, nachkontrolle)
        zeilen_ziel_u: set[tuple[int, int]] = set()                  # Zeilen, auf die eine uebernommene Aenderung zielt
        uebernommen = [a for a in self.aend if a.status == UEBERNOMMEN]
        for ae in uebernommen:
            for op in ae.ops:
                bi = self.bi_von_name.get(op.blatt)
                if bi is None:
                    continue
                if op.art in ("ersetzen", "einfuegen", "zelle_leeren"):
                    ziel_zellen.setdefault((bi, op.zeile, op.spalte), []).append(ae.id)
                    zeilen_ziel_u.add((bi, op.zeile))
                elif op.art == "zeile_einfuegen":
                    neu_ops.append((ae.id, bi, dict(op.zellen), ae.nachkontrolle))
                else:
                    zeilen_ziel_u.add((bi, op.zeile))

        def stumm(key, bi, zeile, spalte, art, ids):
            f = self._fehler(key, bi, zeile, spalte, "auto")
            f.stumm = True
            f.stille_arten.add(art)
            for i in ids:
                if i not in f.aenderungen:
                    f.aenderungen.append(i)

        # Datentyp-Wechsel an Zellen, deren Text unveraendert blieb (Zahl wurde Text o. ae.), in verwalteten Spalten: stiller Fehler
        for bi_t, z_t, c_t in typ_stille or []:
            key_t = ("typ", bi_t, z_t, c_t)
            stumm(key_t, bi_t, z_t, c_t, "datentyp_gewechselt", [])
            self.fehler[key_t].art = "plan_falsch"

        # Regel 27 (1): Der Plan hatte die Aenderung auf "pruefen"/"nicht uebernommen", ohne Freigabe steht sie danach auf "uebernommen"
        r27: dict[str, tuple[int, int, int]] = {}
        for ae in self.aend:
            if self.vor_status.get(ae.id) in (PRUEFEN, NICHT) and ae.status == UEBERNOMMEN:
                ort = next(((z.bi, z.zeile or 0, z.spalte or 0) for z in self.ziele_von[ae.id] if z.zeile is not None), (0, 0, 0))
                r27[ae.id] = ort

        # (1)(4)(5) je Ziel-Zelle der Aenderungen. Bewertet wird je AENDERUNG: weicht die Zelle vom Soll nur um die Woerter
        # einer offenen ("pruefen") Aenderung ab, wartet sie auf die Freigabe und ist kein Fehler (MA-1).
        wartet: set[Any] = set()
        gemeldet_farbe: set[Any] = set()
        r27_wirkung: dict[str, tuple[int, int, int]] = {}
        for (bi, z, c), ids in sorted(ziel_zellen.items()):
            erw, tat = self._erwartet(bi, z, c), self._tatsaechlich(bi, z, c)
            key = ("z", bi, z, c)
            for ae in offen_ziel.get((bi, z, c), []):          # Regel 27 an der Wirkung: die offene Aenderung steht schon im Text
                if ae.id not in r27 and self._wirkung_vorhanden(ae, self._vorher(bi, z, c), tat, ids):
                    r27_wirkung.setdefault(ae.id, (bi, z, c))
            if _s(tat) != _s(erw):
                offene = offen_ziel.get((bi, z, c), [])
                unerkl = self._unerklaert(tat, erw, self._offene_woerter(offene))
                if offene and not unerkl:
                    wartet.add(key)
                elif self._nur_gemeldete_farbe(tat, erw, self._vorher(bi, z, c), offene):
                    gemeldet_farbe.add(key)              # Abweichung steckt ganz in einer gemeldeten Luecke: nicht still
                else:
                    schuld = self._schuldige(ids, unerkl)
                    stumm(key, bi, z, c, "uebernommen_aber_falsch", schuld)
                    if any(nach_id[i].nachkontrolle == "ok" for i in schuld):
                        stumm(key, bi, z, c, "nachkontrolle_ok_aber_falsch", schuld)
                    if self.markiert and _s_text(tat) == _s_text(erw):
                        stumm(key, bi, z, c, "format_falsch", schuld)      # Text richtig, Markierung (Formatklasse) falsch
            if _s(tat) == _s(self._vorher(bi, z, c)) and _s(erw) != _s(tat):
                stumm(key, bi, z, c, "keine_wirkung", ids)            # nur wenn die Zelle laut Soll anders sein muesste

        # neue Zeilen: gibt es im Ergebnis eine Zeile mit dem Inhalt der Operation, und ist sie laut Soll richtig?
        ist_neu_zellen = [it for it in self.ist if it["art"] == "neu"]
        erklaert_neu: set[Any] = set()
        r27_neu: set[tuple[int, int]] = set()                          # (Blatt, Index der neuen Zeile im Ergebnis)
        for aid, bi, zellen, nk in neu_ops:
            gesucht = {c: _s_text(t) for c, t in zellen.items() if _s_text(t)}
            treffer = [it for it in ist_neu_zellen if it["bi"] == bi and it["key"][3] and
                       all(c - 1 < len(it["key"][3]) and _ohne_format(it["key"][3][c - 1]) == t for c, t in gesucht.items())]
            ref = (bi, min(zellen) if zellen else 0)
            if not gesucht:
                continue
            if not treffer and self._neu_als_gepaarte_zeile(bi, gesucht):
                continue          # die Zeilen-Ausrichtung hat die neue Zeile als veraenderte alte Zeile gepaart -- Inhalt steht dort RICHTIG
            if not treffer:
                key = ("nop", bi, aid)
                stumm(key, bi, 0, ref[1], "keine_wirkung", [aid])
                if nk == "ok":
                    stumm(key, bi, 0, ref[1], "nachkontrolle_ok_aber_falsch", [aid])
                continue
            for it in treffer:
                erklaert_neu.add(it["key"])
                if aid in r27:
                    r27_neu.add((bi, it["ist_index"]))
                if not it["richtig"]:
                    stumm(it["key"], bi, it["zeile"], it["spalte"], "uebernommen_aber_falsch", [aid])
                    if nk == "ok":
                        stumm(it["key"], bi, it["zeile"], it["spalte"], "nachkontrolle_ok_aber_falsch", [aid])
                    if self.markiert and it.get("text_richtig"):
                        stumm(it["key"], bi, it["zeile"], it["spalte"], "format_falsch", [aid])

        # Regel 27 (2), an der WIRKUNG erkannt: eine richtige Aenderung steht im Ergebnis, obwohl keine uebernommene Aenderung
        # darauf zielt, wohl aber eine offene -- auch wenn der Status das nicht verraet.
        for ae in self.aend:
            if ae.status == UEBERNOMMEN or ae.id in r27 or ae.id in r27_wirkung:
                continue
            for op in ae.ops:
                bi = self.bi_von_name.get(op.blatt)
                if bi is None or op.art != "zeile_einfuegen":
                    continue
                gesucht = {c: _s_text(t) for c, t in op.zellen.items() if _s_text(t)}
                if not gesucht:
                    continue
                for it in ist_neu_zellen:
                    if it["bi"] == bi and it["key"] not in erklaert_neu and it["key"][3] and \
                            all(c - 1 < len(it["key"][3]) and _ohne_format(it["key"][3][c - 1]) == t for c, t in gesucht.items()):
                        r27_wirkung[ae.id] = (bi, it["zeile"], it["spalte"])
                        r27_neu.add((bi, it["ist_index"]))
                        erklaert_neu.add(it["key"])

        # (2) Kollateral: veraenderte Zelle, laut Soll falsch, und keine Aenderung zielt darauf
        zeilen_mit_ziel = {(bi, z) for (bi, z, _c) in ziel_zellen}
        kollateral = 0
        kollateral_keys: set[Any] = set()
        for it in self.ist:
            key = it["key"]
            if it["richtig"]:
                if it["art"] != "neu":
                    pos = (it["bi"], it["zeile"], it["spalte"])
                    if pos not in ziel_zellen and (it["bi"], it["zeile"]) not in zeilen_ziel_u:
                        for ae in offen_ziel.get(pos, []) + offen_zeile.get((it["bi"], it["zeile"]), []):
                            if ae.id not in r27:
                                r27_wirkung.setdefault(ae.id, pos)
                continue
            if key in wartet or key in gemeldet_farbe:
                continue
            ids = ziel_zellen.get((it["bi"], it["zeile"], it["spalte"]), []) if it["art"] != "neu" else []
            if it["art"] == "neu":
                if key in erklaert_neu:
                    continue
            elif ids or (it["art"] == "weg" and (it["bi"], it["zeile"]) in zeilen_mit_ziel):
                if not ids:      # Zeile wurde entfernt/geleert, weil die Aenderung eine andere Zelle der Zeile traf
                    ids = next(i for (b, z, c), i in ziel_zellen.items() if b == it["bi"] and z == it["zeile"])
                stumm(key, it["bi"], it["zeile"], it["spalte"], "uebernommen_aber_falsch", ids)
                continue
            stumm(key, it["bi"], it["zeile"], it["spalte"], "kollateral", [])
            kollateral += 1
            kollateral_keys.add(key)

        # Fremde Spec: jede Aenderung in einer fremden Spec ist Kollateralschaden (laut Soll bleibt sie unveraendert)
        fremd_keys: set[Any] = set()
        for it in self.ist:
            if it["richtig"] or it["zeile"] not in self.fremd[it["bi"]]:
                continue
            key = it["key"]
            if key not in kollateral_keys:
                kollateral += 1
                kollateral_keys.add(key)
            stumm(key, it["bi"], it["zeile"], it["spalte"], "kollateral", [])
            stumm(key, it["bi"], it["zeile"], it["spalte"], "fremde_spec", [])
            fremd_keys.add(key)

        # Regel-27-Verstoesse eintragen (zaehlen als stille Fehler, nie als Treffer)
        r27_alle = {**r27_wirkung, **r27}
        for aid, (bi, z, c) in r27_alle.items():
            stumm(("r27", aid), bi, z, c, "regel27_verstoss", [aid])
            self.fehler[("r27", aid)].art = "plan_falsch"
        r27_zellen: set[tuple[int, int, int]] = set()
        r27_zeilen: set[tuple[int, int]] = set()
        for aid in r27_alle:
            for zl in self.ziele_von[aid]:
                if zl.zeile is None:
                    continue
                (r27_zellen.add((zl.bi, zl.zeile, zl.spalte)) if zl.spalte is not None else r27_zeilen.add((zl.bi, zl.zeile)))
        for aid, pos in r27_wirkung.items():
            r27_zellen.add(pos)

        def ist_r27(bi: int, zeile: int, spalte: int) -> bool:
            return (bi, zeile, spalte) in r27_zellen or (bi, zeile) in r27_zeilen

        def einheit_r27(e: dict[str, Any]) -> bool:
            if e["art"] == "neu":
                ii = self.neu_abgleich[e["bi"]][0].get(e["si"])
                return ii is not None and (e["bi"], ii) in r27_neu
            return ist_r27(e["bi"], e["zeile"], e["spalte"])

        # --- Soll-Zellen: erreicht / Grund (nur Haupt-Zellen) ----------------------------------------
        haupt = [e for e in self.einheiten if e["klasse"] == "haupt"]
        soll_n = len(haupt)
        erreicht = 0
        pruefen_n = nicht_n = ohne_n = 0
        for e in haupt:
            if einheit_r27(e):
                continue                       # steht schon als Regel-27-Verstoss im Bericht; zaehlt nicht als "ohne Nachfrage erreicht"
            if e["erreicht"]:
                erreicht += 1
                continue
            entf, dazu = self.soll_woerter[e["key"]]
            W, R = Counter(dazu), Counter(entf)
            zu = self._zugehoerig(W, R)
            ziel_hier = [a for a in self.aend if any(
                z.bi == e["bi"] and z.zeile == e["zeile"] and z.spalte in (None, e["spalte"]) for z in self.ziele_von[a.id])]
            verwandt = {a.id: a for a in zu + ziel_hier}
            stati = {a.status for a in verwandt.values()}
            key = e["key"]
            if UEBERNOMMEN not in stati and PRUEFEN in stati:
                pruefen_n += 1
                continue
            if UEBERNOMMEN not in stati and NICHT in stati:
                nicht_n += 1
                continue
            if UEBERNOMMEN in stati:
                # Neben einer angewendeten Aenderung steht eine offene in der Zelle: weicht die Zelle nur um deren Woerter ab,
                # wartet sie auf die Freigabe (MA-1) und ist weder Fehler noch "erreicht".
                offene = [a for a in verwandt.values() if a.status != UEBERNOMMEN]
                if e["art"] != "neu":
                    offene += [a for a in offen_ziel.get((e["bi"], e["zeile"], e["spalte"]), []) if a.id not in verwandt]
                if offene:
                    ist_text = e["ist"] if (e["ist_da"] and e["ist"] is not None) else e["vorher"]
                    if not self._unerklaert(ist_text, e["soll"], self._offene_woerter(offene)):
                        if any(a.status == PRUEFEN for a in offene):
                            pruefen_n += 1
                        else:
                            nicht_n += 1
                        continue
            if not verwandt:
                ohne_n += 1
            if e.get("typ_falsch") and e.get("ist_key") in self.fehler and self.fehler[e["ist_key"]].stumm:
                continue                      # Text stimmt, Datentyp nicht: steht schon als stiller Fehler im Bericht (nicht doppelt)
            f = self._fehler(key, e["bi"], e["zeile"], e["spalte"], "soll")
            for i in sorted(verwandt):
                if i not in f.aenderungen:
                    f.aenderungen.append(i)
            if e["art"] == "neu":
                art = self._art(e["bi"], e["zeile"], e["spalte"], W, R, None, None, e["ist_da"], None, True)
            else:
                en = self._erwartet(e["bi"], e["zeile"], e["spalte"])
                et = self._tatsaechlich(e["bi"], e["zeile"], e["spalte"])
                art = self._art(e["bi"], e["zeile"], e["spalte"], W, R, et, en, e["ist_da"], key)
            f.art = f.art or art

        # --- nicht verwaltete Soll-Zellen (RT-142): Formeln und Sonstiges mit eigener Fehlerzaehlung -----------------
        formeln = [e for e in self.einheiten if e["klasse"] == "formel"]
        sonstige = [e for e in self.einheiten if e["klasse"] == "sonstige"]
        harmlose = [e for e in self.einheiten if e["klasse"] == "harmlos"]
        verfehlt_orte = [{"blatt": e["bi"] + 1, "zeile": e["zeile"], "spalte": e["spalte"], "klasse": e["klasse"]}
                         for e in formeln + sonstige if not e["erreicht"]][:200]

        # Fehlerart der Eintraege, die nur als "Ergebnis falsch" entstanden sind
        for key, f in self.fehler.items():
            if f.art:
                continue
            it = next((i for i in self.ist if i["key"] == key), None)
            if it is not None and "kollateral" in f.stille_arten:
                f.art = self._art_fremd(it, [a.id for a in self.aend if self._trifft(a, it)])
            elif key[0] == "z":
                bi, z, c = key[1:]
                e = next((x for x in self.einheiten if x["key"] == key), None)
                if e is not None:
                    entf, dazu = self.soll_woerter[key]
                    f.art = self._art(bi, z, c, Counter(dazu), Counter(entf), self._tatsaechlich(bi, z, c),
                                      self._erwartet(bi, z, c), any(i["key"] == key for i in self.ist), key)
                elif it is not None:
                    f.art = self._art_fremd(it, f.aenderungen)
                else:
                    f.art = "plan_falsch"
            elif it is not None:
                f.art = self._art_fremd(it, f.aenderungen)
            else:
                f.art = "plan_falsch"

        # Soll-Zellen mit richtigem Text, aber falschem Datentyp: verwaltete = stiller Fehler (Art "datentyp_gewechselt"), die uebrigen
        # (Fuellung, Formeln, sonstige) zaehlt `messe_paar` als Strukturbefund "datentyp".
        for e in self.einheiten:
            if e["klasse"] == "haupt" and e.get("typ_falsch"):
                f = self.fehler.get(e["ist_key"]) or self.fehler.get(e["key"])
                if f is not None and f.stumm:
                    f.stille_arten.add("datentyp_gewechselt")
                    f.art = f.art or "plan_falsch"
        datentyp_soll_haupt = sum(1 for e in haupt if e.get("typ_falsch"))
        datentyp_soll_sonst = sum(1 for e in self.einheiten if e["klasse"] != "haupt" and e.get("typ_falsch"))

        # (3) Bilanz
        bilanz_ok = bool(self.bil_auto["stimmt"]) and bool(self.bil_plan["stimmt"])
        if not bilanz_ok:
            f = self._fehler(("paar", "bilanz"), 0, 0, 0, "paar")
            f.stumm, f.art = True, "bilanz_kaputt"
            f.stille_arten.add("bilanz")
        # Erwartete Luecken muessen gemeldet sein
        luecke_fehlt = 0
        if self.luecken_erwartet:
            gemeldet = Counter(l["art"] for l in self.plan.luecken)
            for art, n in Counter(l.get("art") for l in self.luecken_erwartet).items():
                if gemeldet.get(art, 0) < n:
                    luecke_fehlt += n - gemeldet.get(art, 0)
            if luecke_fehlt:
                f = self._fehler(("paar", "luecke"), 0, 0, 0, "paar")
                f.stumm, f.art = True, "bilanz_kaputt"
                f.stille_arten.add("luecke_nicht_gemeldet")

        # --- Treffergenauigkeit ---------------------------------------------------------------
        auto_gesamt = len(self.ist)

        def item_r27(i: dict[str, Any]) -> bool:
            if i["art"] == "neu":
                return (i["bi"], i["ist_index"]) in r27_neu
            return ist_r27(i["bi"], i["zeile"], i["spalte"])

        auto_richtig = sum(1 for i in self.ist if i["richtig"] and not item_r27(i))
        v_gesamt = v_richtig = None
        if self.e_alle is not None:
            vorhanden = {(i["key"], _s(i["inhalt"])) for i in self.ist}
            neu = [i for i in self.e_alle if (i["key"], _s(i["inhalt"])) not in vorhanden]
            v_gesamt = len(neu)
            v_richtig = sum(1 for i in neu if i["richtig"])

        # --- Lese-Luecke (Woerter, die laut Soll im PDF stehen muessten, aber nicht erkannt wurden) ----
        soll_woerter_n = sum(len(self.soll_woerter[e["key"]][0]) + len(self.soll_woerter[e["key"]][1]) for e in haupt)
        lese_luecke = sum(self.lese_luecke_fehlend().values())
        zaehlung = max(0, soll_woerter_n - self.bil_plan["woerter_im_pdf"])

        # --- Zusammenfassen -------------------------------------------------------------------
        stille = [f for f in self.fehler.values() if f.stumm]
        art_zaehler = Counter()
        for f in stille:
            for a in f.stille_arten:
                art_zaehler[a] += 1
        nicht_ohne_grund = sum(1 for a in self.aend if a.status == NICHT and not a.grund)

        # --- Kennzahlen je AENDERUNG (nicht je Zelle) ---------------------------------------------------
        # "falsch" = die Aenderung steht an einem stillen Fehler (Tool sagt "uebernommen", Ergebnis stimmt nicht) oder verstiess gegen
        # Regel 27. Eine Soll-Zelle, die nur nicht erreicht wurde (z. B. weil der Leser Woerter verpasste), macht eine angewendete
        # Aenderung nicht falsch.
        falsch_ids: set[str] = set(r27_alle)
        for f in stille:
            if f.quelle != "paar":
                falsch_ids.update(i for i in f.aenderungen if nach_id.get(i) is not None and nach_id[i].status == UEBERNOMMEN)
        uebernommen_ids = {a.id for a in uebernommen}
        stille_ids = {i for f in stille for i in f.aenderungen if nach_id.get(i) is not None and nach_id[i].status == UEBERNOMMEN}

        # --- Messung v2: Klassen der Soll-Zellen, Restarbeit, Pruefcodes, Sperr-Quote (nur Zahlen und Kennungen) ----------
        klassen2 = {k: 0 for k in SOLL_KLASSEN}
        klassen2_erreicht = {k: 0 for k in SOLL_KLASSEN}
        for e in self.einheiten:
            klassen2[e["klasse2"]] += 1
            if e["klasse2"] != "haupt" and e["erreicht"]:
                klassen2_erreicht[e["klasse2"]] += 1
        klassen2_erreicht["haupt"] = erreicht           # wie "erreicht_auto" (ohne Regel-27-Zellen)
        leer_neu = sum(1 for e in self.einheiten if e["art"] == "neu" and e["ist_da"] and not _s(e["ist"] or ""))
        codes: Counter = Counter()
        sperre: dict[str, dict[str, int]] = {}
        for ae in self.aend:
            eigene = sorted({_code(c) for c in (ae.pruefcodes or [])})
            codes.update(eigene)
            if self.vor_status.get(ae.id, ae.status) != PRUEFEN:
                continue
            richtig = self.richtig_nach_freigabe(ae)
            for c in eigene or [OHNE_CODE]:
                z = sperre.setdefault(c, {"gesperrt": 0, "messbar": 0, "richtig_nach_freigabe": 0})
                z["gesperrt"] += 1
                if richtig is not None:
                    z["messbar"] += 1
                    z["richtig_nach_freigabe"] += int(richtig)
        # --- Format-Treue (nur "markiert"): Soll-Zellen der Haupt-Spalten, deren Text das Tool richtig geschrieben hat -------------
        fmt = [e for e in haupt if e.get("text_gleich")]
        return {
            "fremde_spec_zeilen": sum(len(f) for f in self.fremd),
            "fremde_spec_beruehrt": len(fremd_keys),
            "markiert": self.markiert,
            "format_treue_zellen_geprueft": len(fmt),
            "format_treue_zellen_gleich": sum(1 for e in fmt if e["format_gleich"]),
            "format_treue_zeichen": sum(e["zeichen"] for e in fmt),
            "format_treue_zeichen_gleich": sum(e["zeichen_gleich"] for e in fmt),
            "soll_klassen": klassen2,
            "soll_klassen_erreicht": klassen2_erreicht,
            "restarbeit": {"pruefen": sum(1 for a in self.aend if a.status == PRUEFEN),
                           "nicht_uebernommen": sum(1 for a in self.aend if a.status == NICHT),
                           "leere_zellen_neue_zeilen": leer_neu,
                           "luecken_woerter": sum(int(lk.get("woerter") or 0) for lk in self.plan.luecken)},
            "pruefcodes": dict(codes),
            "sperr_quote": sperre,
            "bilanz_stimmt": bilanz_ok,
            "vollstaendig": bool(self.bil_auto.get("vollstaendig")),
            "stille_fehler": len(stille),
            "stille_fehler_zellen": sum(1 for f in stille if f.quelle != "paar"),
            "stille_fehler_paar_befunde": sum(1 for f in stille if f.quelle == "paar"),
            "stille_fehler_aenderungen": len(stille_ids),
            "stille_fehler_nach_art": dict(art_zaehler),
            "regel27_verstoesse": len(r27_alle),
            "soll_zellen": soll_n,
            "erreicht_auto": erreicht,
            "soll_zellen_alle": len(self.einheiten),
            "erreicht_alle": erreicht + sum(1 for e in self.einheiten if e["klasse"] != "haupt" and e["erreicht"]),
            "erreicht_formel": sum(1 for e in formeln if e["erreicht"]),
            "soll_zellen_unverwaltet_erreicht": sum(1 for e in harmlose if e["erreicht"]),
            "soll_zellen_unverwaltet_sonstige_erreicht": sum(1 for e in sonstige if e["erreicht"]),
            "datentyp_soll_verwaltet": datentyp_soll_haupt,
            "datentyp_soll_unverwaltet": datentyp_soll_sonst,
            "datentyp_stille_unveraendert": len(typ_stille or []),
            "pruefen_zellen": pruefen_n,
            "nicht_uebernommen_zellen": nicht_n,
            "ohne_aenderung_zellen": ohne_n,
            "zellen_warten_auf_freigabe": len(wartet),
            "zellen_luecke_gemeldet": len(gemeldet_farbe),
            "soll_zellen_pdf_blind": sum(1 for e in haupt if e["blind"]),
            "soll_zellen_pdf_blind_erreicht": sum(1 for e in haupt if e["blind"] and e["erreicht"]),
            "stille_fehler_nur_textform": sum(1 for key, f in self.fehler.items()
                                              if f.stumm and "uebernommen_aber_falsch" in f.stille_arten
                                              and self._nur_textform_key(key)),
            "soll_zellen_unverwaltet": len(harmlose),
            "soll_zellen_unverwaltet_aufschluesselung": {
                "neue_zeile": sum(1 for e in harmlose if e["art"] == "neu"),
                "entfernte_zeile": sum(1 for e in harmlose if e["art"] == "weg")},
            "soll_zellen_formel": len(formeln),
            "soll_zellen_formel_verfehlt": sum(1 for e in formeln if not e["erreicht"]),
            "soll_zellen_unverwaltet_sonstige": len(sonstige),
            "soll_zellen_unverwaltet_sonstige_verfehlt": sum(1 for e in sonstige if not e["erreicht"]),
            "nicht_verwaltet_verfehlt_orte": verfehlt_orte,
            "auto_richtig": auto_richtig,
            "auto_gesamt": auto_gesamt,
            "vorschlaege_richtig": v_richtig,
            "vorschlaege_gesamt": v_gesamt,
            "kollateral": kollateral,
            "idempotent": self.idempotent,
            "nicht_uebernommen_ohne_grund": nicht_ohne_grund,
            "soll_woerter": soll_woerter_n,
            "lese_luecke_woerter": lese_luecke,
            "lese_luecke_zaehlung": zaehlung,
            "luecken_nicht_gemeldet": luecke_fehlt,
            "luecken_gemeldet": len(self.plan.luecken),
            "aenderungen": len(self.aend),
            "aenderungen_uebernommen_korrekt": len(uebernommen_ids - falsch_ids),
            "aenderungen_uebernommen_falsch": len(uebernommen_ids & falsch_ids),
            "aenderungen_zu_pruefen": sum(1 for a in self.aend if a.status == PRUEFEN),
            "aenderungen_nicht_uebernommen": sum(1 for a in self.aend if a.status == NICHT),
            "bilanz": {k: v for k, v in self.bil_auto.items() if isinstance(v, (int, bool))},
            "aenderungen_nach_status": {s: sum(1 for a in self.aend if a.status == s) for s in (UEBERNOMMEN, PRUEFEN, NICHT)},
            "hinweise": self.hinweise,
            "fehler": [
                {"blatt": f.bi + 1 if f.quelle != "paar" else 0, "zeile": f.zeile, "spalte": f.spalte,
                 "aenderung": f.aenderungen[0] if f.aenderungen else "", "art": f.art, "stumm": f.stumm,
                 "stille_arten": sorted(f.stille_arten),
                 **dict(zip(("zeile_nachher", "zeile_ergebnis", "neue_zeile"), self.orte(f.key, f.zeile)))}
                for f in sorted(self.fehler.values(), key=lambda f: (f.bi, f.zeile, f.spalte, f.art))
            ],
        }

    def _trifft(self, ae, it: dict[str, Any]) -> bool:
        return any(z.bi == it["bi"] and z.zeile in (None, it["zeile"]) and z.spalte in (None, it["spalte"])
                   for z in self.ziele_von[ae.id] if ae.status == UEBERNOMMEN)

    def _trifft_offen(self, ae, it: dict[str, Any]) -> bool:
        """Wie `_trifft`, aber unabhaengig vom Status (fuer die Fehler-Schablonen der freigegebenen Vorschlaege)."""
        return any(z.bi == it["bi"] and z.zeile in (None, it["zeile"]) and z.spalte in (None, it["spalte"])
                   for z in self.ziele_von[ae.id])


# ======================================================================================
# 4. Ablauf je Paar
# ======================================================================================
def _lade_erwartet(ordner: Path) -> dict[str, Any] | None:
    p = ordner / "erwartet.json"
    if not p.is_file():
        return None
    try:
        daten = json.loads(p.read_text(encoding="utf-8"))
        return daten if isinstance(daten, dict) else None
    except (OSError, ValueError):
        return None


def _pdf_pruefung(pdf: Path) -> tuple[list[dict[str, Any]] | None, set[str] | None]:
    """Unabhaengige Plausibilitaet (RT-141), direkt mit pdfplumber und NICHT mit dem Code des Tools: Seiten ohne Textzeichen
    (Scan) und Seiten mit einem grossen eingebetteten Bild. Nur ein Hinweis, nie ein Fehler. Liefert ausserdem die Menge aller
    Zeichen des PDF (fuer `_Auswertung._blind`). None = konnte nicht geprueft werden."""
    try:
        import pdfplumber
        out: list[dict[str, Any]] = []
        zeichen: set[str] = set()
        with pdfplumber.open(str(pdf)) as doc:
            for nr, seite in enumerate(doc.pages, 1):
                if not seite.chars:
                    out.append({"seite": nr, "art": "keine_textzeichen"})
                for c in seite.chars:
                    zeichen.update(c["text"])
                flaeche = float(seite.width * seite.height) or 1.0
                if any(abs((im["x1"] - im["x0"]) * (im["bottom"] - im["top"])) / flaeche >= PLAUSI_BILD_ANTEIL
                       for im in seite.images):
                    out.append({"seite": nr, "art": "bild"})
        return out, zeichen
    except Exception:
        return None, None


def _spec_norm(text: Any) -> str:
    """Vereinheitlichte Spec-Nummer ("11_22_33", "11 22 33 ", "11.22.33" -> "112233"); sonst der vereinheitlichte Text."""
    return re.sub(r"[\s_.\-]", "", str(text or "")).lower()


def _pdf_specs(pdf: Path) -> set[str] | None:
    """Spec-Nummern (Form "11 22 33") auf den ersten Seiten des PDF -- unabhaengig vom Tool. None = nicht lesbar."""
    try:
        import pdfplumber
        out: set[str] = set()
        with pdfplumber.open(str(pdf)) as doc:
            for seite in doc.pages[:SPEC_SEITEN]:
                for m in _SPEC_NR.finditer(seite.extract_text() or ""):
                    out.add("".join(m.groups()))
        return out
    except Exception:
        return None


# ---- Struktur (RT-136): alles, was keine Zelle mit Text ist -----------------------------------------------------------
def _zeilenabbildung(d: Diff, max_zeile: int) -> dict[int, int | None]:
    """vorher-Zeile -> Zeile im Ergebnis (None = entfernt oder nicht eindeutig). Zeilen mit Inhalt kommen aus der Ausrichtung,
    leere Zeilen dazwischen nur, wenn die Zeile davor und die Zeile danach dieselbe Verschiebung ergeben."""
    abb: dict[int, int | None] = {}
    gepaart = sorted(d.zuordnung)
    for r in range(1, max_zeile + 1):
        if r in d.zuordnung:
            abb[r] = d.zuordnung[r]
        elif r in d.weg:
            abb[r] = None
        else:
            k = bisect.bisect_left(gepaart, r)                  # gepaart[k-1] < r < gepaart[k] (r selbst ist nicht gepaart)
            a = d.zuordnung[gepaart[k - 1]] + (r - gepaart[k - 1]) if k > 0 else None
            b = d.zuordnung[gepaart[k]] - (gepaart[k] - r) if k < len(gepaart) else None
            if a is None and b is None:
                abb[r] = r
            elif a is not None and b is not None:
                abb[r] = a if a == b else None
            else:
                abb[r] = a if a is not None else b
    return abb


def _farbe(c: Any) -> Any:
    return None if c is None else (c.type, c.value)


def _schrift(z: Any, nur_form: bool) -> tuple:
    f = z.font
    form = (bool(f.b), bool(f.i), f.sz, f.name)
    return form if nur_form else form + (f.u, _farbe(f.color), z.fill.patternType, _farbe(z.fill.fgColor))


def _spalten_versteckt(ws: Any, grenze: int) -> set[int]:
    out: set[int] = set()
    for buchstabe, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(buchstabe)
            out.update(range(lo, min(dim.max or lo, max(grenze, lo)) + 1))
    return out


def _validierungen_liste(ws: Any, abb: dict[int, int | None] | None, ende: dict[int, int | None] | None = None) -> list:
    """Datenvalidierungen als Liste (None = nicht abbildbar); `ende` bildet nur die LETZTE Zeile eines Bereichs ab (siehe `_ende_abb`)."""
    from openpyxl.worksheet.cell_range import CellRange
    out: list = []
    for dv in ws.data_validations.dataValidation:
        bereiche = []
        ok = True
        for rng in str(dv.sqref).split():
            cr = CellRange(rng)
            if abb is None:
                lo, hi = cr.min_row, cr.max_row
            else:
                lo, hi = abb.get(cr.min_row), (ende if ende is not None else abb).get(cr.max_row)
            if lo is None or hi is None:
                ok = False
                break
            bereiche.append((cr.min_col, lo, cr.max_col, hi))
        out.append((dv.type, dv.operator, dv.formula1, dv.formula2, bool(dv.allow_blank), tuple(sorted(bereiche))) if ok else None)
    return out


def _validierungen(ws: Any, abb: dict[int, int | None] | None) -> tuple[set, int]:
    """Datenvalidierungen als Menge; mit `abb` werden die Zeilen auf das Ergebnis abgebildet. -> (Menge, nicht abbildbare)"""
    liste = _validierungen_liste(ws, abb)
    return {x for x in liste if x is not None}, sum(1 for x in liste if x is None)


def _ende_abb(abb: dict[int, int | None], neue_nr: list[int]) -> dict[int, int | None] | None:
    """Abbildung nur fuer das ENDE von Bereichen (Datenpruefung, bedingte Formatierung, Autofilter): endet ein Bereich an einer Zeile,
    hinter der im Ergebnis neue Zeilen eingefuegt wurden, darf er bis zur letzten dieser direkt folgenden neuen Zeilen reichen (wie
    "kopierte Zeile einfuegen" in Excel) ODER beim alten Ende bleiben -- beides ist richtig. Anfaenge bleiben unveraendert; ein
    verkuerzter Bereich oder ein verschobener Anfang ist weiter ein Befund. None = keine neuen Zeilen."""
    neu = set(neue_nr)
    if not neu:
        return None
    ende = dict(abb)
    gewachsen = False
    for r, x in abb.items():
        if x is None:
            continue
        k = 0
        while x + k + 1 in neu:
            k += 1
        if k:
            ende[r] = x + k
            gewachsen = True
    return ende if gewachsen else None


def _tolerant_diff(erw_a: list, erw_g: list | None, da: set) -> int:
    """Unterschiede zwischen erwarteten Elementen und vorhandenen (`da`): je Element zaehlt Variante a (altes Ende) ODER g (Ende bis
    zur letzten neuen Zeile). Nicht abbildbare Elemente (None) werden nicht vermutet: dann zaehlen nur fehlende Eintraege."""
    erw_g = erw_g if erw_g is not None else erw_a
    offen = any(a is None for a in erw_a)
    rest = set(da)
    fehlt = 0
    for a, g in dict.fromkeys(zip(erw_a, erw_g)):             # wie bisher als Menge: gleiche Eintraege nur einmal
        if a is None:
            continue
        if a in rest:
            rest.discard(a)
        elif g in rest:
            rest.discard(g)
        else:
            fehlt += 1
    return max(fehlt, 0 if offen else len(rest))              # wie `_menge_diff`: eine geaenderte Regel zaehlt einmal


def _zeile_abb(abb: dict[int, int | None], r: int, vmax: int) -> int | None:
    """Zeile `r` der Arbeitsmappe vorher -> Zeile im Ergebnis. Unterhalb der letzten belegten Zeile gilt deren Verschiebung;
    ein Bereich bis zur letzten Excel-Zeile ("R8:R1048576") bleibt dort (Excel deckelt beim Einfuegen)."""
    if r in abb:
        return abb[r]
    if r >= EXCEL_MAX_ZEILE:
        return r
    if r > vmax:
        basis = abb.get(vmax)
        return None if basis is None else min(EXCEL_MAX_ZEILE, r + (basis - vmax))
    return None


def _bereich_abb(rng: str, abb: dict[int, int | None] | None, vmax: int, ende: dict[int, int | None] | None = None) -> tuple | None:
    """'A2:C9' -> (Spalte von, Zeile von, Spalte bis, Zeile bis). Mit `abb` werden die Zeilen auf das Ergebnis abgebildet.
    None = nicht abbildbar (Zeile entfernt oder nicht eindeutig). Ganze Spalten/Zeilen und Unlesbares bleiben unveraendert."""
    try:
        c1, r1, c2, r2 = range_boundaries(str(rng).replace("$", ""))
    except Exception:
        return ("roh", str(rng))
    if abb is None or r1 is None or r2 is None:
        return (c1, r1, c2, r2)
    lo, hi = _zeile_abb(abb, r1, vmax), _zeile_abb(abb if ende is None else ende, r2, vmax)
    return None if lo is None or hi is None else (c1, lo, c2, hi)


def _menge_diff(erwartet: set, da: set, offen: int) -> int:
    """Anzahl Unterschiede zwischen dem, was nach dem Verschieben erwartet wird, und dem, was da ist. Ist etwas nicht abbildbar
    (`offen`), zaehlen nur die FEHLENDEN Eintraege (nie eine Vermutung ueber zusaetzliche)."""
    return max(len(erwartet - da), 0 if offen else len(da - erwartet))


_ZELLBEZUG = re.compile(r"(?<![A-Za-z0-9_.\"])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_(])")


def _relativ(formel: str, anker: tuple[int, int] | None) -> str:
    """Formel einer bedingten Formatierung mit Bezuegen relativ zur linken oberen Zelle ihres Bereichs (wie Excel sie beim
    Einfuegen von Zeilen mitzieht: "LEN(F62)>2" auf F62:F107 == "LEN(F63)>2" auf F63:F108). Absolute Bezuege bleiben."""
    if anker is None:
        return formel

    def ers(m: re.Match) -> str:
        sp = column_index_from_string(m.group(2))
        z = int(m.group(4))
        c_ = m.group(2) if m.group(1) else f"C[{sp - anker[1]}]"
        r_ = m.group(4) if m.group(3) else f"R[{z - anker[0]}]"
        return f"{m.group(1)}{c_}{m.group(3)}{r_}"
    return _ZELLBEZUG.sub(ers, formel)


def _bedingte_formatierung_liste(ws: Any, abb: dict[int, int | None] | None, vmax: int,
                                 ende: dict[int, int | None] | None = None) -> list:
    out: list = []
    for cf in ws.conditional_formatting:
        try:
            c1, r1, _c2, _r2 = range_boundaries(str(cf.sqref).split()[0].replace("$", ""))
            anker = (r1, c1) if r1 is not None and c1 is not None else None
        except Exception:
            anker = None
        regeln = tuple(sorted(
            (str(r.type), str(r.operator), tuple(_relativ(str(f), anker) for f in (r.formula or [])), str(r.text), str(r.rank),
             str(r.percent), str(r.bottom), str(r.aboveAverage), str(r.stopIfTrue)) for r in cf.rules))
        for rng in str(cf.sqref).split():
            b = _bereich_abb(rng, abb, vmax, ende)
            out.append(None if b is None else (b, regeln))
    return out


def _bedingte_formatierung(ws: Any, abb: dict[int, int | None] | None, vmax: int) -> tuple[set, int]:
    liste = _bedingte_formatierung_liste(ws, abb, vmax)
    return {x for x in liste if x is not None}, sum(1 for x in liste if x is None)


def _hyperlinks(ws: Any, abb: dict[int, int | None] | None, vmax: int) -> tuple[set, int]:
    out: set = set()
    offen = 0
    for (r, c), z in ws._cells.items():
        link = getattr(z, "_hyperlink", None)          # verbundene Zellen (MergedCell) haben keinen Link
        if link is None:
            continue
        x = r if abb is None else _zeile_abb(abb, r, vmax)
        if x is None:
            offen += 1
        else:
            out.add((x, c, link.target, link.location))
    return out, offen


def _tabellen(ws: Any, abb: dict[int, int | None] | None, vmax: int) -> tuple[set, int]:
    out: set = set()
    offen = 0
    for t in ws.tables.values():
        b = _bereich_abb(t.ref, abb, vmax)
        if b is None:
            offen += 1
        else:
            out.add((t.name, b))
    return out, offen


def _autofilter(ws: Any, abb: dict[int, int | None] | None, vmax: int, ende: dict[int, int | None] | None = None) -> tuple[Any, bool]:
    """-> (Bereich und Zahl der Filterspalten oder None, nicht abbildbar)"""
    ref = ws.auto_filter.ref
    if not ref:
        return None, False
    b = _bereich_abb(ref, abb, vmax, ende)
    return (None, True) if b is None else ((b, len(ws.auto_filter.filterColumn)), False)


def _fixierung(ws: Any, abb: dict[int, int | None] | None, vmax: int) -> tuple[Any, bool]:
    """Fixierte Zeilen/Spalten (Freeze Panes). Abgebildet wird die LETZTE fixierte Zeile (eine Einfuegung direkt darunter
    verschiebt die Fixierung nicht). -> ((Spalte, erste nicht fixierte Zeile) oder None, nicht abbildbar)"""
    fp = ws.freeze_panes
    if not fp:
        return None, False
    spalte, zeile = coordinate_from_string(fp)
    if zeile <= 1 or abb is None:
        return (spalte, zeile), False
    letzte = _zeile_abb(abb, zeile - 1, vmax)
    return (None, True) if letzte is None else ((spalte, letzte + 1), False)


def _spaltenbreiten(ws: Any, grenze: int) -> dict[int, float]:
    out: dict[int, float] = {}
    for buchstabe, dim in ws.column_dimensions.items():
        if dim.width is None:
            continue
        lo = dim.min or column_index_from_string(buchstabe)
        for c in range(lo, min(dim.max or lo, max(grenze, lo)) + 1):
            out[c] = round(float(dim.width), 2)
    return out


def _kommentar_text(z: Any) -> tuple[str, str] | None:
    k = getattr(z, "comment", None)
    return None if k is None else (str(k.author or ""), str(k.text or ""))


_NAMEN_REF = re.compile(r"(?<![A-Za-z0-9_.])(?:(?P<blatt>'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?(?P<d1>\$?)(?P<c1>[A-Z]{1,3})(?P<d2>\$?)"
                        r"(?P<r1>\d+)(?::(?P<d3>\$?)(?P<c2>[A-Z]{1,3})(?P<d4>\$?)(?P<r2>\d+))?(?![A-Za-z0-9_(])")


def _namen(wb: Any, abbs: list[dict[int, int | None]] | None, vmaxs: list[int]) -> tuple[dict, set]:
    """Benannte Bereiche: (Geltungsbereich, Name) -> Verweistext. Mit `abbs` werden die Zeilen der Verweise auf das Ergebnis
    abgebildet. -> (Namen, Namen mit nicht abbildbarem Verweis)"""
    titel = [ws.title for ws in wb.worksheets]
    offen: set = set()

    def text(schluessel: Any, roh: str, standard: str | None) -> str:
        if abbs is None or not roh:
            return roh or ""
        fehlt = False

        def ersetze(m: "re.Match[str]") -> str:
            nonlocal fehlt
            blatt = m.group("blatt")
            if blatt is None:
                name = standard
            else:
                name = blatt[1:-1].replace("''", "'") if blatt.startswith("'") else blatt
            if name not in titel:
                return m.group(0)
            i = titel.index(name)
            if i >= len(abbs):
                return m.group(0)
            r1 = _zeile_abb(abbs[i], int(m.group("r1")), vmaxs[i])
            r2 = _zeile_abb(abbs[i], int(m.group("r2")), vmaxs[i]) if m.group("r2") else None
            if r1 is None or (m.group("r2") and r2 is None):
                fehlt = True
                return m.group(0)
            aus = (blatt + "!" if blatt else "") + m.group("d1") + m.group("c1") + m.group("d2") + str(r1)
            if m.group("r2"):
                aus += ":" + m.group("d3") + m.group("c2") + m.group("d4") + str(r2)
            return aus
        neu = _NAMEN_REF.sub(ersetze, roh)
        if fehlt:
            offen.add(schluessel)
        return neu

    out: dict = {}
    for n, dn in wb.defined_names.items():
        out[(None, n)] = text((None, n), dn.attr_text, None)
    for i, ws in enumerate(wb.worksheets):
        for n, dn in ws.defined_names.items():
            out[(i, n)] = text((i, n), dn.attr_text, ws.title)
    return out, offen


def _struktur_befund(vorher: Path, ergebnis: Path, diffs: list[Diff]) -> dict[str, Any]:
    """Vergleicht Nicht-Zell-Eigenschaften vorher <-> Ergebnis: Blaetter (Name, Reihenfolge, Sichtbarkeit), ausgeblendete Zeilen
    und Spalten, verbundene Zellen, Datenvalidierung, bedingte Formatierung, Hyperlinks, Kommentare, benannte Bereiche, Autofilter,
    Excel-Tabellen, fixierte Zeilen/Spalten, Spaltenbreiten, Zeilenhoehen unveraenderter Zeilen, Zahlenformat, Schrift/Fuellung
    und Datentyp unveraenderter Zellen. Zeilen werden ueber die Ausrichtung abgebildet (Einfuegen/Loeschen verschiebt legitim);
    nur UNERWARTETE Abweichungen zaehlen. Nie Text, nur Zaehler. `datentyp_orte` (Blatt-Index, Zeile, Spalte) nennt die Zellen mit
    Datentyp-Wechsel, damit `messe_paar` die in verwalteten Spalten als stillen Fehler zaehlen kann."""
    n = Counter()
    orte: list[tuple[int, int, int]] = []
    wv = load_workbook(str(vorher))
    we = load_workbook(str(ergebnis))
    try:
        sv, se = wv.worksheets, we.worksheets
        for i in range(max(len(sv), len(se))):
            if i >= len(sv) or i >= len(se) or sv[i].title != se[i].title:
                n["blatt"] += 1
            elif sv[i].sheet_state != se[i].sheet_state:
                n["sichtbarkeit"] += 1
        gemeinsam = min(len(sv), len(se))
        vmaxs = [max(sv[i].max_row, 1) for i in range(gemeinsam)]
        abbs = [_zeilenabbildung(diffs[i] if i < len(diffs) else Diff(), vmaxs[i]) for i in range(gemeinsam)]
        # benannte Bereiche (Arbeitsmappe und Blatt)
        nv, nv_offen = _namen(wv, abbs, vmaxs)
        ne, _ = _namen(we, None, [])
        n["benannter_bereich"] += sum(1 for k, t in nv.items() if k not in nv_offen and ne.get(k) != t) \
            + (0 if nv_offen else sum(1 for k in ne if k not in nv))
        for i in range(gemeinsam):
            v, e = sv[i], se[i]
            d = diffs[i] if i < len(diffs) else Diff()
            abb, vmax = abbs[i], vmaxs[i]
            geaendert = set(d.zellen)
            geaenderte_zeilen = {r for r, _c in geaendert}
            bilder_e = {nr for nr in d.neue_nr}
            rueck = {x: r for r, x in abb.items() if x is not None}            # Zeile im Ergebnis -> Zeile vorher
            # ausgeblendete Zeilen: gleiche Zeile -> gleicher Zustand; neue Zeilen nie ausgeblendet
            for r, x in abb.items():
                if x is None:
                    continue
                hv = bool(v.row_dimensions[r].hidden) if r in v.row_dimensions else False
                he = bool(e.row_dimensions[x].hidden) if x in e.row_dimensions else False
                n["zeile_versteckt"] += hv != he
            for r in bilder_e:
                n["zeile_versteckt"] += bool(e.row_dimensions[r].hidden) if r in e.row_dimensions else False
            grenze = max(v.max_column, e.max_column)
            n["spalte_versteckt"] += len(_spalten_versteckt(v, grenze) ^ _spalten_versteckt(e, grenze))
            # Spaltenbreiten
            bv, be = _spaltenbreiten(v, grenze), _spaltenbreiten(e, grenze)
            n["spaltenbreite"] += sum(1 for c in set(bv) | set(be) if bv.get(c) != be.get(c))
            # Zeilenhoehen unveraenderter Zeilen (veraenderte Zeilen und neue Zeilen: keine Aussage)
            hoehen_v = {r for r in v.row_dimensions if v.row_dimensions[r].height is not None}
            hoehen_e = {rueck[x] for x in e.row_dimensions if e.row_dimensions[x].height is not None and x in rueck}
            for r in hoehen_v | hoehen_e:
                x = abb.get(r)
                if x is None or r in geaenderte_zeilen:
                    continue
                hv = v.row_dimensions[r].height if r in v.row_dimensions else None
                he = e.row_dimensions[x].height if x in e.row_dimensions else None
                n["zeilenhoehe"] += hv != he
            # fixierte Zeilen/Spalten
            fv, fv_offen = _fixierung(v, abb, vmax)
            fe, _ = _fixierung(e, None, vmax)
            n["fixierung"] += (not fv_offen) and fv != fe
            # verbundene Zellen
            erwartet, offen_m = set(), 0
            for rng in v.merged_cells.ranges:
                lo, hi = abb.get(rng.min_row), abb.get(rng.max_row)
                if lo is None or hi is None:
                    offen_m += 1
                else:
                    erwartet.add((rng.min_col, lo, rng.max_col, hi))
            da = {(rng.min_col, rng.min_row, rng.max_col, rng.max_row) for rng in e.merged_cells.ranges}
            n["verbunden"] += max(len(erwartet - da), 0 if offen_m else len(da - erwartet))
            # Neue Zeilen direkt hinter dem Ende eines Bereichs (Blattende wie Blattmitte): ein Bereich, der an der Vorlagezeile
            # endet, darf bis zur letzten direkt folgenden neuen Zeile reichen ODER beim alten Ende bleiben (siehe `_ende_abb`).
            # Anfaenge bleiben unveraendert; Verkuerzen und Verschieben des Starts ist weiter ein Befund.
            ende = _ende_abb(abb, d.neue_nr)
            # Datenvalidierung
            dv_da = {x for x in _validierungen_liste(e, None) if x is not None}
            n["validierung"] += _tolerant_diff(_validierungen_liste(v, abb), _validierungen_liste(v, abb, ende) if ende else None, dv_da)
            # bedingte Formatierung, Hyperlinks, Excel-Tabellen
            cf_da = {x for x in _bedingte_formatierung_liste(e, None, vmax) if x is not None}
            n["bedingte_formatierung"] += _tolerant_diff(
                _bedingte_formatierung_liste(v, abb, vmax),
                _bedingte_formatierung_liste(v, abb, vmax, ende) if ende else None, cf_da)
            lv, lv_offen = _hyperlinks(v, abb, vmax)
            n["hyperlink"] += _menge_diff(lv, _hyperlinks(e, None, vmax)[0], lv_offen)
            tv, tv_offen = _tabellen(v, abb, vmax)
            n["tabelle"] += _menge_diff(tv, _tabellen(e, None, vmax)[0], tv_offen)
            # Autofilter
            ae_, _ = _autofilter(e, None, vmax)
            varianten = [(av, av_offen) for av, av_offen in (_autofilter(v, abb, vmax), *([_autofilter(v, abb, vmax, ende)] if ende else []))]
            n["autofilter"] += min(int((not av_offen) and av != ae_) for av, av_offen in varianten)
            # Kommentare (an veraenderten Zellen und in neuen Zeilen darf das Tool welche anlegen oder ergaenzen)
            ev_kommentare = {(r, c): _kommentar_text(z) for (r, c), z in v._cells.items() if _kommentar_text(z) is not None}
            ee_kommentare = {(r, c): _kommentar_text(z) for (r, c), z in e._cells.items() if _kommentar_text(z) is not None}
            for (r, c), inhalt in ev_kommentare.items():
                x = abb.get(r)
                if x is None:
                    continue
                gefunden = ee_kommentare.get((x, c))
                n["kommentar"] += gefunden is None or ((r, c) not in geaendert and gefunden != inhalt)
            for (x, c) in ee_kommentare:
                r = rueck.get(x)
                if r is not None and (r, c) not in ev_kommentare and (r, c) not in geaendert:
                    n["kommentar"] += 1
            # Zahlenformat, Schrift und Datentyp der Zellen
            e_zellen = e._cells
            for (r, c), zv in list(v._cells.items()):
                x = abb.get(r)
                if x is None:
                    continue
                ze = e_zellen.get((x, c))
                tv_ = _typ_von(zv)
                if tv_ != "leer" and (r, c) not in geaendert and tv_ != (_typ_von(ze) if ze is not None else "leer"):
                    n["datentyp"] += 1                              # Zahl <-> Text <-> Formel <-> Datum an unveraendertem Text
                    orte.append((i, r, c))
                if not zv.has_style and (ze is None or not ze.has_style):
                    continue                                        # beide ohne eigenes Format: nichts zu vergleichen
                if ze is None:
                    ze = e.cell(x, c)
                if zv.number_format != ze.number_format:
                    n["zahlenformat"] += 1
                if _schrift(zv, (r, c) in geaendert) != _schrift(ze, (r, c) in geaendert):
                    n["schrift"] += 1
    finally:
        wv.close()
        we.close()
    return {"struktur_geaendert": sum(n.values()), "struktur_nach_art": {a: int(n.get(a, 0)) for a in STRUKTUR_ARTEN},
            "datentyp_orte": orte}


# ---- Legende und nachher-Format je Korpus (Auftrag 003) ---------------------------------------------------------------------
def korpus_einstellungen(eltern: Path, korpus: str | None = None) -> dict[str, str]:
    """Legende und nachher-Format eines Korpus aus seiner KORPUS.json. Fehlt dort eine Angabe, gilt der Standard der Korpora
    v1/v2 (legende "alt", nachher_format "sauber"). Ohne KORPUS.json: Vorgabe des genannten Korpus (KORPUS_VORGABE), sonst
    {} -- dann gilt die Konfiguration wie sie ist (z. B. echte Paare unter --ordner)."""
    p = eltern / "KORPUS.json"
    daten: dict[str, Any] | None = None
    if p.is_file():
        try:
            roh = json.loads(p.read_text(encoding="utf-8"))
            daten = roh if isinstance(roh, dict) else {}
        except (OSError, ValueError):
            daten = {}
    if daten is None:
        return dict(KORPUS_VORGABE.get(korpus or "", {}))
    legende = daten.get("legende") if daten.get("legende") in LEGENDEN else KORPUS_STANDARD["legende"]
    fmt = daten.get("nachher_format") if daten.get("nachher_format") in NACHHER_FORMATE else KORPUS_STANDARD["nachher_format"]
    return {"legende": legende, "nachher_format": fmt}


_FEHLT = object()


@contextlib.contextmanager
def konfig_ueberschreibung(einstellungen: dict[str, str] | None):
    """Setzt `legende.modus` (config/farben.yaml) und `markierung.aenderungen` (config/zuordnung.yaml) NUR fuer die Dauer der
    Messung eines Paars: markiert -> farbe, sauber -> keine. Danach gilt wieder der vorige Wert (nichts wird global gesetzt)."""
    alt: list[tuple[dict, str, Any]] = []
    try:
        e = einstellungen or {}
        if e.get("legende"):
            d = config.lade("farben").setdefault("legende", {})
            alt.append((d, "modus", d.get("modus", _FEHLT)))
            d["modus"] = e["legende"]
        if e.get("nachher_format"):
            d = config.lade("zuordnung").setdefault("markierung", {})
            alt.append((d, "aenderungen", d.get("aenderungen", _FEHLT)))
            d["aenderungen"] = "farbe" if e["nachher_format"] == "markiert" else "keine"
        yield
    finally:
        for d, k, v in reversed(alt):
            if v is _FEHLT:
                d.pop(k, None)
            else:
                d[k] = v


def _wirksam() -> dict[str, str]:
    """Legende und nachher-Format, die gerade gelten (fuer den Bericht)."""
    leg = str((config.lade("farben").get("legende") or {}).get("modus", "alt"))
    farbe = str((config.lade("zuordnung").get("markierung") or {}).get("aenderungen", "keine")) == "farbe"
    return {"legende": leg, "nachher_format": "markiert" if farbe else "sauber"}


def messe_paar(ordner: Path, mit_vorschlaegen: bool = False, ergebnis_kopie: Path | None = None,
               schablonen: bool = False, einstellungen: dict[str, str] | None = None) -> dict[str, Any]:
    """Misst ein Paar. Gibt die Kennzahlen (ohne Dokumentinhalt) zurueck; wirft bei kaputter Eingabe.
    Faellt nur ein Teil der Messung aus (Wiederholungslauf, Vorschlaege, Struktur), bleiben die Funde des ersten Laufs erhalten:
    das Paar gilt dann als "abgestuerzt" (teilweise gemessen), die Funde zaehlen weiter.
    ergebnis_kopie: nur fuer Tests -- kopiert die Ergebnis-Excel (nur Automatik) dorthin.
    schablonen: ergaenzt `fehler_schablonen` (Formen Soll/Ist je stillem Fehler und falschem Vorschlag, ohne Klartext).
    einstellungen: {"legende", "nachher_format"} fuer diesen Lauf (None = aus KORPUS.json des Eltern-Ordners, siehe
    `korpus_einstellungen`; {} = Konfiguration wie sie ist). "markiert" vergleicht zusaetzlich die Formatklassen je Zeichen."""
    einst = korpus_einstellungen(ordner.parent) if einstellungen is None else einstellungen
    with konfig_ueberschreibung(einst):
        wirksam = _wirksam()
        res = _messe_paar(ordner, mit_vorschlaegen, ergebnis_kopie, schablonen, wirksam["nachher_format"] == "markiert")
    res.update(wirksam)
    return res


def _messe_paar(ordner: Path, mit_vorschlaegen: bool, ergebnis_kopie: Path | None, schablonen: bool,
                markiert: bool) -> dict[str, Any]:
    pdf, vorher, nachher = (ordner / n for n in DATEIEN)
    for p in (pdf, vorher, nachher):
        if not p.is_file():
            raise FileNotFoundError(p.name)
    erwartet = _lade_erwartet(ordner)
    hashes_vorher = {p.name: _sha(p) for p in (pdf, vorher, nachher)}
    ausfaelle: list[tuple[str, str]] = []              # (Messteil, Name der Ausnahme-Klasse)
    t0 = time.monotonic()
    with tempfile.TemporaryDirectory(prefix="cde_messen_") as tmp_name:
        tmp = Path(tmp_name)
        # A: nur Automatik
        pa = analysiere_pdf(pdf)
        xl = ExcelModell(vorher)
        plan = plane(pa, xl)
        konsens.bildprobe_anwenden(plan.aenderungen, pa, pdf=pdf)
        bil_plan = bilanz(plan.aenderungen, plan.markierte_woerter, plan.ignoriert, plan.luecken)
        vor_status = {a.id: a.status for a in plan.aenderungen}
        aend_auto = copy.deepcopy(plan.aenderungen)
        e_auto = tmp / "ergebnis_auto.xlsx"
        uebernehmen(aend_auto, vorher, e_auto, None, absaetze=plan.absaetze)
        bil_auto = bilanz(aend_auto, plan.markierte_woerter, plan.ignoriert, plan.luecken)
        # B: alle "pruefen" freigegeben (Qualitaet der Vorschlaege)
        e_alle = None
        if mit_vorschlaegen:
            try:
                aend_alle = copy.deepcopy(plan.aenderungen)
                e_alle = tmp / "ergebnis_alle.xlsx"
                uebernehmen(aend_alle, vorher, e_alle, {a.id for a in aend_alle if a.status == PRUEFEN},
                            absaetze=plan.absaetze)
            except Exception as exc:
                e_alle = None
                ausfaelle.append(("vorschlaege", type(exc).__name__))
        # Idempotenz: Ergebnis nochmals als "vorher" mit demselben PDF. Ein Absturz hier darf die Funde des ersten Laufs nicht loeschen.
        idempotent: bool | None
        try:
            e_auto2 = tmp / "ergebnis_auto2.xlsx"
            plan2 = plane(analysiere_pdf(pdf), ExcelModell(e_auto))
            uebernehmen(plan2.aenderungen, e_auto, e_auto2, None, absaetze=plan2.absaetze)
            if markiert:                               # auch die Markierung darf sich beim zweiten Lauf nicht aendern
                idempotent = _mit_format(_lese_excel(e_auto2, True)) == _mit_format(_lese_excel(e_auto, True))
            else:
                idempotent = _lese_excel(e_auto2) == _lese_excel(e_auto)
        except Exception as exc:
            idempotent = None
            ausfaelle.append(("idempotenz", type(exc).__name__))

        if ergebnis_kopie is not None:
            shutil.copy2(e_auto, ergebnis_kopie)
        v_bl, (n_bl, n_typ), (e_bl, e_typ) = (_lese_excel(vorher, markiert), _lese_excel_mit_typen(nachher, markiert),
                                              _lese_excel_mit_typen(e_auto, markiert))
        e_alle_bl = _lese_excel(e_alle, markiert) if e_alle else None
        plausi, pdf_zeichen = _pdf_pruefung(pdf)
        auswertung = _Auswertung(v_bl, n_bl, e_bl, e_alle_bl, _blattnamen(vorher), aend_auto, plan, pa, bil_auto,
                                 bil_plan, idempotent, (erwartet or {}).get("luecken"), vor_status, pdf_zeichen,
                                 (n_typ, e_typ), markiert=markiert, pdf_specs=_pdf_specs(pdf))
        # Struktur zuerst (liefert die Zellen mit Datentyp-Wechsel), dann die Auswertung
        struktur: dict[str, Any] | None
        try:
            struktur = _struktur_befund(vorher, e_auto, auswertung.I)
        except Exception as exc:
            struktur = None
            ausfaelle.append(("struktur", type(exc).__name__))
        orte = struktur.pop("datentyp_orte") if struktur else []
        auswertung.typ_stille = typ_stille = auswertung.datentyp_orte_verwaltet(orte)
        res = auswertung.auswerten()
        if struktur is None:
            res.update({"struktur_geaendert": None, "struktur_nach_art": {a: 0 for a in STRUKTUR_ARTEN}})
        else:
            # Datentyp-Wechsel in verwalteten Spalten sind schon stille Fehler (nicht doppelt zaehlen); Wechsel an richtig befuellten
            # Zellen ausserhalb der Haupt-Spalten (Fuellung, Formeln, sonstige) kommen als Strukturbefund dazu
            mehr = res["datentyp_soll_unverwaltet"] - len(typ_stille)
            struktur["struktur_nach_art"]["datentyp"] += mehr
            struktur["struktur_geaendert"] += mehr
            res.update(struktur)
        res["struktur_vorschlaege_geaendert"] = res["struktur_vorschlaege_angekuendigt"] = None
        res["struktur_vorschlaege_nach_art"] = res["struktur_vorschlaege_angekuendigt_nach_art"] = {a: 0 for a in STRUKTUR_ARTEN}
        if e_alle is not None and auswertung.e_alle_diffs is not None:
            # Auch die Excel der Variante "alle Vorschlaege freigegeben" wird auf Struktur geprueft. Weicht sie ab, ist das erlaubt,
            # wenn der Plan das Risiko vorab gemeldet hat (Hinweis in einer Aenderung); sonst ist es ein Befund.
            try:
                sv = _struktur_befund(vorher, e_alle, auswertung.e_alle_diffs)
                sv.pop("datentyp_orte", None)
                gewarnt = " ".join(t.lower() for a in plan.aenderungen for t in (*a.hinweise, a.grund))
                ang = {art: n for art, n in sv["struktur_nach_art"].items()
                       if n and STRUKTUR_RISIKO_STICHWORT.get(art) and STRUKTUR_RISIKO_STICHWORT[art] in gewarnt}
                unerwartet = {art: (0 if art in ang else n) for art, n in sv["struktur_nach_art"].items()}
                res["struktur_vorschlaege_geaendert"] = sum(unerwartet.values())
                res["struktur_vorschlaege_nach_art"] = unerwartet
                res["struktur_vorschlaege_angekuendigt"] = sum(ang.values())
                res["struktur_vorschlaege_angekuendigt_nach_art"] = {a: ang.get(a, 0) for a in STRUKTUR_ARTEN}
            except Exception as exc:
                ausfaelle.append(("struktur_vorschlaege", type(exc).__name__))

    # Plausibilitaet direkt aus dem PDF (unabhaengig vom Tool); nur ein Hinweis
    seiten_mit_luecke = {l.get("seite") for l in plan.luecken}
    if plausi is None:
        res["plausibilitaet_hinweise"] = None
        res["plausibilitaet_ungemeldet"] = None
    else:
        res["plausibilitaet_hinweise"] = [dict(h, gemeldet=h["seite"] in seiten_mit_luecke) for h in plausi]
        res["plausibilitaet_ungemeldet"] = sum(1 for h in plausi if h["seite"] not in seiten_mit_luecke)
    res["luecken_pruefbar"] = bool(erwartet) and isinstance(erwartet.get("luecken"), list)

    res["original_unveraendert"] = all(_sha(p) == hashes_vorher[p.name] for p in (pdf, vorher, nachher))
    # vom Leser gewaehlte Legende (bei "auto" je Dokument entschieden) und Markup-Leiste -- nur feste Kennungen (Regel 16)
    erkannt = str(getattr(pa, "legende_modus", "") or "")
    res["legende_erkannt"] = erkannt if erkannt in LEGENDEN else ("unbekannt" if not erkannt else "sonstige")
    res["markup_leiste_erkannt"] = bool(getattr(pa, "markup_leiste", False))
    res["dauer_sekunden"] = round(time.monotonic() - t0, 2)
    res["profil_angenommen"] = True                      # ohne erwartet.json/Profil: Standard-Aufbau angenommen (siehe fasse_zusammen)
    res["erzeuger"] = _erzeuger(erwartet)
    if schablonen:
        try:
            res["fehler_schablonen"], res["fehler_schablonen_verworfen"] = _fehler_schablonen(auswertung, res["erzeuger"])
        except Exception:                              # reine Zusatzangabe: faellt sie aus, ist das Paar trotzdem gemessen
            res["fehler_schablonen"], res["fehler_schablonen_verworfen"] = [], -1
    if erwartet:
        res["stoerer"] = [s for s in erwartet.get("stoerer", []) if isinstance(s, str)]
        res["profil"] = erwartet.get("profil", "") if isinstance(erwartet.get("profil", ""), str) else ""
        res["profil_angenommen"] = not res["profil"]
        wm = erwartet.get("woerter_markiert")
        if isinstance(wm, int):
            res["lese_luecke_erwartet"] = max(0, wm - bil_plan["woerter_im_pdf"])
    if ausfaelle:                                       # teilweise gemessen: Funde bleiben, aber das Paar ist nicht vollstaendig gemessen
        res["abgestuerzt"] = ausfaelle[0][1]
        res["teilweise_gemessen"] = True
        res["ausgefallene_messteile"] = [{"teil": t, "grund": g} for t, g in ausfaelle]
        res["fehler"] = list(res["fehler"]) + [
            {"blatt": 0, "zeile": 0, "spalte": 0, "aenderung": "", "art": "unklar", "stumm": False, "grund": g}
            for _t, g in ausfaelle]
    return res


def _soll_von_item(aw: "_Auswertung", it: dict[str, Any], diffs: list[Diff]) -> str:
    """Soll-Text zu einer veraenderten Ist-Zelle (nur im Speicher)."""
    if it["art"] != "neu":
        return aw._erwartet(it["bi"], it["zeile"], it["spalte"])
    bi = it["bi"]
    paar, _ort = _neu_abgleich(aw.S[bi].neue, diffs[bi].neue, aw._paar_spalten(bi))
    si = {ii: s for s, ii in paar.items()}.get(it.get("ist_index"))
    zellen = aw.S[bi].neue[si][1] if si is not None else ()
    return zellen[it["spalte"] - 1] if it["spalte"] - 1 < len(zellen) else ""


def _fehler_schablonen(aw: "_Auswertung", erzeuger: str, maximal: int = 40) -> tuple[list[dict[str, Any]], int]:
    """`messen --schablonen`: fuer jeden stillen Fehler und jeden falschen Vorschlag (hoechstens `maximal`) die FORMEN von Soll und Ist
    der betroffenen Zelle (Aufbau und Zeichensatz siehe src/steckbrief.py, `schablone`), dazu Fehlerart, Aenderungstyp und Erzeuger.
    Kein Klartext. -> (Eintraege, Zahl der Eintraege, die die Positivliste nicht bestanden haben und verworfen wurden)."""
    from . import steckbrief as sb
    nach_id = {a.id: a for a in aw.aend}
    stille: list[dict[str, Any]] = []
    for key, f in aw.fehler.items():
        if not f.stumm or f.quelle == "paar":
            continue
        k0, soll, ist = (key[0] if isinstance(key, tuple) and key else ""), "", ""
        if k0 in ("z", "typ", "r27") and f.zeile > 0:
            soll, ist = aw._erwartet(f.bi, f.zeile, f.spalte), aw._tatsaechlich(f.bi, f.zeile, f.spalte)
        elif k0 == "n" and len(key) == 5:
            it = next((i for i in aw.ist if i["key"] == key), None)
            if it is not None:
                soll, ist = _soll_von_item(aw, it, aw.I), it["inhalt"]
        elif k0 == "n":
            e = next((x for x in aw.einheiten if x["key"] == key), None)
            if e is not None:
                soll, ist = e["soll"], e["ist"] or ""
        typ = next((nach_id[i].typ for i in f.aenderungen if i in nach_id), "")
        stille.append(_schablonen_eintrag("stiller_fehler", f.art, typ, erzeuger, soll, ist))
    falsche: list[dict[str, Any]] = []
    if aw.e_alle is not None and aw.e_alle_diffs is not None:
        vorhanden = {(i["key"], _s(i["inhalt"])) for i in aw.ist}
        for it in aw.e_alle:
            if it["richtig"] or (it["key"], _s(it["inhalt"])) in vorhanden:
                continue
            ae = next((a for a in aw.aend if aw._trifft_offen(a, it)), None)
            falsche.append(_schablonen_eintrag("falscher_vorschlag", "", ae.typ if ae else "", erzeuger,
                                               _soll_von_item(aw, it, aw.e_alle_diffs), it["inhalt"]))
    auswahl: list[dict[str, Any]] = []
    for i in range(max(len(stille), len(falsche))):           # abwechselnd, damit beide Quellen vorkommen
        auswahl += [lst[i] for lst in (stille, falsche) if i < len(lst)]
    auswahl = auswahl[:maximal]
    ok = [e for e in auswahl if not sb.pruefe_eintrag(e)]
    return ok, len(auswahl) - len(ok)


def _formen(text: Any, pos: int) -> list[dict[str, str]]:
    """Formen eines Zelltexts rund um die Stelle `pos`, je Lauf mit Formatklasse (kein Klartext)."""
    from . import steckbrief as sb
    return [{"f": k, "form": f} for k, f in sb.fenster(_stuecke(text), pos)]


def _erste_abweichung(soll: Any, ist: Any) -> int:
    a, b = list(zip(str(soll), _klassen(soll))), list(zip(str(ist), _klassen(ist)))
    return next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), min(len(a), len(b)))


def _schablonen_eintrag(quelle: str, art: str, typ: str, erzeuger: str, soll: Any, ist: Any) -> dict[str, Any]:
    """Wie `steckbrief.fehler_eintrag`, aber markierte Zellen behalten ihre Formatklassen je Lauf ("-", "a", "n")."""
    from . import steckbrief as sb
    if not (getattr(soll, "klassen", None) or getattr(ist, "klassen", None)):
        return sb.fehler_eintrag(quelle, art, typ, erzeuger, soll, ist)
    e = sb.fehler_eintrag(quelle, art, typ, erzeuger, "", "")
    pos = _erste_abweichung(soll, ist)
    e["soll"], e["ist"] = _formen(soll, pos), _formen(ist, pos)
    return e


def _erzeuger(erwartet: dict[str, Any] | None) -> str:
    """PDF-Erzeuger laut erwartet.json: writer | libreoffice; Generator v1 kannte nur den eigenen Schreiber; ohne erwartet.json
    (echte Paare): unbekannt. Nur feste Kennungen (Regel 16)."""
    if not erwartet:
        return "unbekannt"
    e = erwartet.get("erzeuger", "writer")
    return e if e in ERZEUGER_KENNUNGEN else "sonstiger"


def _fehlgeschlagen(name: str, exc: BaseException) -> dict[str, Any]:
    """Paar konnte nicht gemessen werden: Fehlerart "unklar", Grund nur als Name der Ausnahme-Klasse."""
    return {"name": name, "abgestuerzt": type(exc).__name__, "bilanz_stimmt": None, "vollstaendig": None,
            "stille_fehler": 0, "stille_fehler_zellen": 0, "stille_fehler_paar_befunde": 0, "stille_fehler_aenderungen": 0,
            "stille_fehler_nach_art": {}, "regel27_verstoesse": 0, "soll_zellen": 0, "erreicht_auto": 0, "pruefen_zellen": 0,
            "nicht_uebernommen_zellen": 0, "ohne_aenderung_zellen": 0, "zellen_warten_auf_freigabe": 0,
            "soll_zellen_unverwaltet": 0, "soll_zellen_unverwaltet_aufschluesselung": {"neue_zeile": 0, "entfernte_zeile": 0},
            "soll_zellen_formel": 0, "soll_zellen_formel_verfehlt": 0, "soll_zellen_unverwaltet_sonstige": 0,
            "soll_zellen_unverwaltet_sonstige_verfehlt": 0, "nicht_verwaltet_verfehlt_orte": [], "auto_richtig": 0,
            "auto_gesamt": 0, "vorschlaege_richtig": None, "vorschlaege_gesamt": None, "kollateral": 0,
            "idempotent": None, "nicht_uebernommen_ohne_grund": 0, "soll_woerter": 0, "lese_luecke_woerter": 0,
            "lese_luecke_zaehlung": 0, "luecken_nicht_gemeldet": 0, "luecken_gemeldet": 0, "aenderungen": 0,
            "aenderungen_uebernommen_korrekt": 0, "aenderungen_uebernommen_falsch": 0, "aenderungen_zu_pruefen": 0,
            "aenderungen_nicht_uebernommen": 0, "struktur_geaendert": None,
            "soll_zellen_alle": 0, "erreicht_alle": 0, "erreicht_formel": 0, "soll_zellen_unverwaltet_erreicht": 0,
            "soll_zellen_unverwaltet_sonstige_erreicht": 0, "datentyp_soll_verwaltet": 0, "datentyp_soll_unverwaltet": 0,
            "datentyp_stille_unveraendert": 0, "struktur_vorschlaege_geaendert": None,
            "struktur_vorschlaege_angekuendigt": None, "profil_angenommen": True,
            "struktur_nach_art": {a: 0 for a in STRUKTUR_ARTEN}, "plausibilitaet_hinweise": None,
            "plausibilitaet_ungemeldet": None, "luecken_pruefbar": False,
            "aenderungen_nach_status": {}, "hinweise": [], "original_unveraendert": None,
            "soll_klassen": {k: 0 for k in SOLL_KLASSEN}, "soll_klassen_erreicht": {k: 0 for k in SOLL_KLASSEN},
            "restarbeit": {k: 0 for k in RESTARBEIT_TEILE}, "pruefcodes": {}, "sperr_quote": {}, "erzeuger": "unbekannt",
            "fehler": [{"blatt": 0, "zeile": 0, "spalte": 0, "aenderung": "", "art": "unklar", "stumm": False,
                        "grund": type(exc).__name__}]}

# ======================================================================================
# 5. Bericht
# ======================================================================================
def _anteil(zaehler: float, nenner: float) -> float | None:
    return round(zaehler / nenner, 4) if nenner else None


def _ampel(ok: bool | None, leer: str = "nicht gemessen", luecke: bool = False) -> str:
    """ok=None: nicht messbar; ok=False: Fehler (zaehlt immer); ok=True: nur dann "ok", wenn alle Paare gemessen wurden
    (`luecke` = es gab abgestuerzte Paare)."""
    if ok is None:
        return leer
    if not ok:
        return "NICHT ok"
    return "nicht vollständig gemessen" if luecke else "ok"


# Zaehleinheit je Kennzahl (damit "97 % Zellen" nicht groesser wirkt als es ist)
ZAEHLEINHEITEN = {
    "stille_fehler": "Zellen und Befunde (Zellen + Paar-Befunde + Regel-27-Verstöße)",
    "stille_fehler_aenderungen": "Änderungen",
    "regel27_verstoesse": "Änderungen",
    "treffergenauigkeit_auto": "Zellen",
    "vorschlags_genauigkeit": "Zellen",
    "abdeckung_ohne_nachfrage": "Zellen",
    "aenderungen_uebernommen_korrekt": "Änderungen",
    "lese_luecke_anteil": "Wörter",
    "kollateral": "Zellen",
    "idempotent": "Paare",
    "original_unveraendert": "Paare",
    "bilanz_stimmig": "Paare",
    "luecken_nicht_gemeldet": "Lücken (Paar-Befunde)",
    "struktur_geaendert": "Befunde (Blätter, Zeilen, Spalten, Bereiche, Zellen)",
    "plausibilitaet_ungemeldet": "Seiten",
    "soll_zellen_formel": "Zellen",
    "abdeckung_alle": "Zellen",
    "struktur_vorschlaege_geaendert": "Befunde",
}


def fasse_zusammen(paare: list[dict[str, Any]], korpus: str, generator_version: str, vorschlaege: bool) -> dict[str, Any]:
    gemessen = [p for p in paare if "abgestuerzt" not in p]
    teilweise = [p for p in paare if p.get("teilweise_gemessen")]

    def s(k: str) -> int:
        return sum(p.get(k) or 0 for p in paare)

    def s_dict(k: str, schluessel: tuple[str, ...]) -> dict[str, int]:
        return {a: sum((p.get(k) or {}).get(a, 0) for p in paare) for a in schluessel}

    fehlerarten = {a: 0 for a in FEHLERARTEN}
    for p in paare:
        for f in p["fehler"]:
            fehlerarten[f["art"]] = fehlerarten.get(f["art"], 0) + 1
    stille_art = Counter()
    for p in paare:
        stille_art.update(p.get("stille_fehler_nach_art", {}))
    soll = s("soll_zellen")
    pruefbar = sum(1 for p in paare if p.get("luecken_pruefbar"))
    struktur_paare = [p for p in paare if p.get("struktur_geaendert") is not None]
    plausi_paare = [p for p in paare if p.get("plausibilitaet_hinweise") is not None]
    ae_korrekt, ae_falsch = s("aenderungen_uebernommen_korrekt"), s("aenderungen_uebernommen_falsch")
    unverw = s("soll_zellen_unverwaltet")
    # Fehlt "erreicht_alle" (altes Format), gilt die Abdeckung der Haupt-Spalten
    erreicht_alle = sum(p["erreicht_alle"] if p.get("erreicht_alle") is not None else (p.get("erreicht_auto") or 0) for p in paare)
    vorschlag_paare = [p for p in paare if p.get("struktur_vorschlaege_geaendert") is not None]
    gesamt = {
        "paare": len(paare), "paare_gemessen": len(gemessen), "paare_abgestuerzt": len(paare) - len(gemessen),
        "paare_teilweise_gemessen": len(teilweise),
        "soll_zellen": soll,
        "bilanz_stimmig": sum(1 for p in paare if p.get("bilanz_stimmt")),
        "vollstaendig": sum(1 for p in gemessen if p["vollstaendig"]),
        "stille_fehler": s("stille_fehler"),
        "stille_fehler_zellen": s("stille_fehler_zellen"),
        "stille_fehler_paar_befunde": s("stille_fehler_paar_befunde"),
        "stille_fehler_aenderungen": s("stille_fehler_aenderungen"),
        "stille_fehler_nur_textform": s("stille_fehler_nur_textform"),
        "stille_fehler_nach_art": {a: stille_art.get(a, 0) for a in STILLE_ARTEN},
        "regel27_verstoesse": s("regel27_verstoesse"),
        "auto_richtig": s("auto_richtig"), "auto_gesamt": s("auto_gesamt"),
        "vorschlaege_gemessen": vorschlaege,
        "vorschlaege_richtig": s("vorschlaege_richtig") if vorschlaege else None,
        "vorschlaege_gesamt": s("vorschlaege_gesamt") if vorschlaege else None,
        "treffergenauigkeit_auto": _anteil(s("auto_richtig"), s("auto_gesamt")),
        "vorschlags_genauigkeit": _anteil(s("vorschlaege_richtig"), s("vorschlaege_gesamt")) if vorschlaege else None,
        "abdeckung_ohne_nachfrage": _anteil(s("erreicht_auto"), soll),
        "anteil_pruefen": _anteil(s("pruefen_zellen"), soll),
        "anteil_nicht_uebernommen": _anteil(s("nicht_uebernommen_zellen"), soll),
        "ohne_aenderung_zellen": s("ohne_aenderung_zellen"),
        "zellen_warten_auf_freigabe": s("zellen_warten_auf_freigabe"),
        # je Aenderung (statt je Zelle)
        "aenderungen": s("aenderungen"),
        "aenderungen_uebernommen_korrekt": ae_korrekt,
        "aenderungen_uebernommen_falsch": ae_falsch,
        "aenderungen_zu_pruefen": s("aenderungen_zu_pruefen"),
        "aenderungen_nicht_uebernommen": s("aenderungen_nicht_uebernommen"),
        "treffergenauigkeit_aenderungen": _anteil(ae_korrekt, ae_korrekt + ae_falsch),
        # Soll-Zellen ausserhalb der Haupt-Spalten (RT-142)
        "soll_zellen_unverwaltet": unverw,
        "soll_zellen_unverwaltet_aufschluesselung": s_dict("soll_zellen_unverwaltet_aufschluesselung",
                                                          ("neue_zeile", "entfernte_zeile")),
        "soll_zellen_formel": s("soll_zellen_formel"),
        "soll_zellen_formel_verfehlt": s("soll_zellen_formel_verfehlt"),
        "soll_zellen_unverwaltet_sonstige": s("soll_zellen_unverwaltet_sonstige"),
        "soll_zellen_unverwaltet_sonstige_verfehlt": s("soll_zellen_unverwaltet_sonstige_verfehlt"),
        "soll_zellen_alle": soll + unverw + s("soll_zellen_formel") + s("soll_zellen_unverwaltet_sonstige"),
        "erreicht_alle": erreicht_alle,
        "abdeckung_alle": _anteil(erreicht_alle, soll + unverw + s("soll_zellen_formel") + s("soll_zellen_unverwaltet_sonstige")),
        "erreicht_formel": s("erreicht_formel"),
        "abdeckung_alle_ohne_formeln": _anteil(
            erreicht_alle - s("erreicht_formel"),
            soll + unverw + s("soll_zellen_unverwaltet_sonstige")),
        "soll_zellen_unverwaltet_erreicht": s("soll_zellen_unverwaltet_erreicht"),
        "soll_zellen_unverwaltet_sonstige_erreicht": s("soll_zellen_unverwaltet_sonstige_erreicht"),
        "datentyp_stille_fehler": s("datentyp_stille_unveraendert") + s("datentyp_soll_verwaltet"),
        "soll_zellen_pdf_blind": s("soll_zellen_pdf_blind"),
        "soll_zellen_pdf_blind_erreicht": s("soll_zellen_pdf_blind_erreicht"),
        "nicht_uebernommen_ohne_grund": s("nicht_uebernommen_ohne_grund"),
        "soll_woerter": s("soll_woerter"),
        "lese_luecke_woerter": s("lese_luecke_woerter"),
        "lese_luecke_anteil": _anteil(s("lese_luecke_woerter"), s("soll_woerter")),
        "idempotent": sum(1 for p in gemessen if p["idempotent"]),
        "kollateral": s("kollateral"),
        "original_unveraendert": sum(1 for p in gemessen if p["original_unveraendert"]),
        "luecken_nicht_gemeldet": s("luecken_nicht_gemeldet"),
        "luecken_pruefbar_paare": pruefbar,
        "luecken_nicht_pruefbar_paare": len(paare) - pruefbar,
        # Struktur (RT-136)
        "struktur_geaendert": sum(p["struktur_geaendert"] for p in struktur_paare),
        "struktur_nach_art": s_dict("struktur_nach_art", STRUKTUR_ARTEN),
        "struktur_paare_geprueft": len(struktur_paare),
        "struktur_geprueft": STRUKTUR_GEPRUEFT,
        "struktur_nicht_geprueft": STRUKTUR_NICHT_GEPRUEFT,
        # Struktur der Excel aus der Variante "alle Vorschlaege freigegeben" (nur mit --mit-vorschlaegen)
        "struktur_vorschlaege_paare_geprueft": len(vorschlag_paare),
        "struktur_vorschlaege_geaendert": sum(p["struktur_vorschlaege_geaendert"] for p in vorschlag_paare),
        "struktur_vorschlaege_angekuendigt": sum(p.get("struktur_vorschlaege_angekuendigt") or 0 for p in vorschlag_paare),
        "struktur_vorschlaege_nach_art": s_dict("struktur_vorschlaege_nach_art", STRUKTUR_ARTEN),
        # Plausibilitaet direkt aus dem PDF (RT-141), nur Hinweis
        "plausibilitaet_seiten": sum(len(p["plausibilitaet_hinweise"]) for p in plausi_paare),
        "plausibilitaet_ungemeldet": sum(p["plausibilitaet_ungemeldet"] or 0 for p in plausi_paare),
        "plausibilitaet_paare_geprueft": len(plausi_paare),
        "zaehleinheiten": ZAEHLEINHEITEN,
    }
    # --- Messung v2: Klassen, Restarbeit, Pruefcodes, Sperr-Quote, je Erzeuger -------------------------------------------
    klassen = s_dict("soll_klassen", SOLL_KLASSEN)
    klassen_err = s_dict("soll_klassen_erreicht", SOLL_KLASSEN)
    gesamt["soll_klassen"] = klassen
    gesamt["soll_klassen_erreicht"] = klassen_err
    gesamt["soll_zellen_ableitbar"] = sum(klassen[k] for k in ABLEITBAR)
    gesamt["erreicht_ableitbar"] = sum(klassen_err[k] for k in ABLEITBAR)
    gesamt["abdeckung_ableitbar"] = _anteil(gesamt["erreicht_ableitbar"], gesamt["soll_zellen_ableitbar"])
    gesamt["abdeckung_alle_klassen"] = _anteil(sum(klassen_err.values()), sum(klassen.values()))
    rest = s_dict("restarbeit", RESTARBEIT_TEILE)
    gesamt["restarbeit"] = rest
    gesamt["restarbeit_je_100"] = _restarbeit(rest, gesamt["aenderungen"])
    gesamt["restarbeit_formel"] = RESTARBEIT_FORMEL
    codes: Counter = Counter()
    for p in paare:
        codes.update(p.get("pruefcodes") or {})
    gesamt["pruefcodes"] = dict(sorted(codes.items()))
    sperre: dict[str, dict[str, Any]] = {}
    for p in paare:
        for code, z in (p.get("sperr_quote") or {}).items():
            e = sperre.setdefault(code, {"gesperrt": 0, "messbar": 0, "richtig_nach_freigabe": 0})
            for k in e:
                e[k] += int(z.get(k, 0))
    for e in sperre.values():
        e["sperr_quote"] = _anteil(e["richtig_nach_freigabe"], e["messbar"])
    gesamt["sperr_quote"] = dict(sorted(sperre.items()))
    je_erzeuger: dict[str, dict[str, Any]] = {}
    for p in paare:
        e = je_erzeuger.setdefault(p.get("erzeuger") or "unbekannt",
                                   {"paare": 0, "stille_fehler": 0, "regel27_verstoesse": 0, "kollateral": 0, "soll_zellen": 0,
                                    "erreicht_auto": 0, "auto_richtig": 0, "auto_gesamt": 0, "soll_woerter": 0,
                                    "lese_luecke_woerter": 0, "idempotent": 0, "aenderungen": 0,
                                    "restarbeit": {k: 0 for k in RESTARBEIT_TEILE}})
        e["paare"] += 1
        for k in ("stille_fehler", "regel27_verstoesse", "kollateral", "soll_zellen", "erreicht_auto", "auto_richtig",
                  "auto_gesamt", "soll_woerter", "lese_luecke_woerter", "aenderungen"):
            e[k] += p.get(k) or 0
        e["idempotent"] += 1 if p.get("idempotent") else 0
        for k in RESTARBEIT_TEILE:
            e["restarbeit"][k] += (p.get("restarbeit") or {}).get(k, 0)
    for e in je_erzeuger.values():
        e["treffergenauigkeit_auto"] = _anteil(e["auto_richtig"], e["auto_gesamt"])
        e["abdeckung_ohne_nachfrage"] = _anteil(e["erreicht_auto"], e["soll_zellen"])
        e["lese_luecke_anteil"] = _anteil(e["lese_luecke_woerter"], e["soll_woerter"])
        e["restarbeit_je_100"] = _restarbeit(e["restarbeit"], e["aenderungen"])

    je_stoerer: dict[str, dict[str, Any]] = {}
    for p in paare:
        for st in p.get("stoerer", []):
            e = je_stoerer.setdefault(st, {"paare": 0, "stille_fehler": 0, "soll_zellen": 0, "erreicht_auto": 0})
            e["paare"] += 1
            e["stille_fehler"] += p["stille_fehler"]
            e["soll_zellen"] += p["soll_zellen"]
            e["erreicht_auto"] += p["erreicht_auto"]
    for e in je_stoerer.values():
        e["abdeckung_ohne_nachfrage"] = _anteil(e.pop("erreicht_auto"), e.pop("soll_zellen"))
    # --- Format-Treue (Auftrag 003): nur Paare mit nachher_format "markiert" ---------------------------------------------------
    f_zellen, f_gleich = s("format_treue_zellen_geprueft"), s("format_treue_zellen_gleich")
    gesamt["paare_markiert"] = sum(1 for p in paare if p.get("markiert"))
    gesamt["format_treue_zellen_geprueft"] = f_zellen
    gesamt["format_treue_zellen_gleich"] = f_gleich
    gesamt["format_treue_zeichen"] = s("format_treue_zeichen")
    gesamt["format_treue_zeichen_gleich"] = s("format_treue_zeichen_gleich")
    gesamt["format_treue"] = _anteil(f_gleich, f_zellen)
    gesamt["format_treue_zeichen_anteil"] = _anteil(gesamt["format_treue_zeichen_gleich"], gesamt["format_treue_zeichen"])
    gesamt["stille_fehler_format"] = stille_art.get("format_falsch", 0)
    gesamt["fremde_spec_beruehrt"] = s("fremde_spec_beruehrt")
    gesamt["paare_mit_fremder_spec"] = sum(1 for p in paare if p.get("fremde_spec_zeilen"))
    gesamt["paare_fremde_spec_beruehrt"] = sum(1 for p in paare if p.get("fremde_spec_beruehrt"))
    gesamt["legende_erkannt"] = dict(sorted(Counter(p.get("legende_erkannt") or "unbekannt" for p in gemessen).items()))
    gesamt["markup_leiste_erkannt"] = sum(1 for p in gemessen if p.get("markup_leiste_erkannt"))
    n = len(paare)                       # alle Paare
    m = len(gemessen)                    # davon voll messbar
    voll = n > 0 and not gesamt["paare_abgestuerzt"]
    lue = not voll                       # es gab Paare, die nicht (vollstaendig) gemessen werden konnten
    # Die 90-%-Regel gilt fuer Standard-Aufbauten. Paare OHNE Profil (alle echten Paare, keine erwartet.json) werden wie Standard-Paare
    # behandelt ("Standard-Aufbau angenommen"); sonst waere die Messung bei echten Paaren milder als bei erfundenen.
    standard = [p for p in gemessen if p.get("profil") == "standard" or not p.get("profil")]
    standard_angenommen = [p for p in standard if not p.get("profil")]
    standard_abdeckung = _anteil(sum(p["erreicht_auto"] for p in standard), sum(p["soll_zellen"] for p in standard))
    gesamt["standard_paare"] = len(standard)
    gesamt["standard_paare_angenommen"] = len(standard_angenommen)
    gesamt["standard_abdeckung"] = standard_abdeckung
    nicht_verfehlt = gesamt["soll_zellen_formel_verfehlt"] + gesamt["soll_zellen_unverwaltet_sonstige_verfehlt"]

    bilanz_fehl = any(p.get("bilanz_stimmt") is False for p in paare) or gesamt["luecken_nicht_gemeldet"] > 0
    k1 = _ampel(None if not n else not bilanz_fehl, luecke=lue)
    k1_wert = f"{gesamt['bilanz_stimmig']} von {n} Paaren"
    if k1 == "ok" and gesamt["luecken_nicht_pruefbar_paare"]:
        k1 = "nicht gemessen"            # RT-141: ohne erwartet.json kann "Luecke nicht gemeldet" nie auffallen
    if gesamt["luecken_nicht_pruefbar_paare"]:
        k1_wert += f"; Lücken: nur {pruefbar} von {n} prüfbar (erwartet.json)"
    k_plausi_ampel = ("nicht gemessen" if not plausi_paare else
                      ("ok" if not gesamt["plausibilitaet_ungemeldet"] else "Hinweis"))
    kriterien = [
        {"nr": 1, "name": "Bilanz stimmig", "wert": k1_wert, "ziel": "alle Paare", "einheit": "Paare", "ampel": k1,
         "detail": "Die Lückenprüfung (Scan, Zwei-Spalten, Tabelle, Bild nicht gemeldet) braucht erwartet.json"},
        {"nr": 1, "name": "Plausibilität direkt aus dem PDF (nur Hinweis)",
         "wert": f"{gesamt['plausibilitaet_ungemeldet']} von {gesamt['plausibilitaet_seiten']}",
         "ziel": "0 ungemeldet", "einheit": "Seiten", "ampel": k_plausi_ampel,
         "detail": "Seiten ohne Textzeichen oder mit großem Bild, zu denen das Tool keine Lücke meldet (Hinweis, kein Fehler)"},
        {"nr": 2, "name": "Stille Fehler", "wert": str(gesamt["stille_fehler"]), "ziel": "0", "einheit": "Zellen und Befunde",
         "detail": f"{gesamt['stille_fehler_zellen']} Zellen, {gesamt['stille_fehler_paar_befunde']} Paar-Befunde, "
                   f"{gesamt['regel27_verstoesse']} Regel-27-Verstöße; {gesamt['stille_fehler_aenderungen']} Änderungen betroffen; "
                   f"davon nur Leerraum/Zeichenform: {gesamt['stille_fehler_nur_textform']}",
         "ampel": _ampel(gesamt["stille_fehler"] == 0, luecke=lue)},
        {"nr": 2, "name": "Regel 27: Vorschlag ohne Freigabe angewendet", "wert": str(gesamt["regel27_verstoesse"]), "ziel": "0",
         "einheit": "Änderungen", "ampel": _ampel(gesamt["regel27_verstoesse"] == 0, luecke=lue),
         "detail": "Eine Änderung stand im Plan auf „prüfen“ und wurde ohne Freigabe angewendet"},
        {"nr": 3, "name": "Treffergenauigkeit automatisch",
         "wert": _prozent(gesamt["treffergenauigkeit_auto"]), "ziel": "100 %", "einheit": "Zellen",
         "ampel": _ampel(None if not gesamt["auto_gesamt"] else gesamt["auto_richtig"] == gesamt["auto_gesamt"], luecke=lue)},
        {"nr": 3, "name": "Treffergenauigkeit je Änderung (automatisch übernommen)",
         "wert": f"{_prozent(gesamt['treffergenauigkeit_aenderungen'])} ({ae_korrekt} von {ae_korrekt + ae_falsch})",
         "ziel": "100 %", "einheit": "Änderungen",
         "ampel": _ampel(None if not (ae_korrekt + ae_falsch) else ae_falsch == 0, luecke=lue)},
        {"nr": 3, "name": "Format-Treue (markierte Zellen: Zeichen für Zeichen normal/alt/neu)",
         "wert": f"{_prozent(gesamt['format_treue'])} ({f_gleich} von {f_zellen})", "ziel": "100 %", "einheit": "Zellen",
         "ampel": ("nicht gemessen" if not gesamt["paare_markiert"] else _hinweis(gesamt["format_treue"], lambda x: x >= 1.0)),
         "detail": f"Nur Paare mit nachher_format „markiert“ ({gesamt['paare_markiert']} Paare): Soll-Zellen der Haupt-Spalten, deren "
                   f"Text das Tool richtig geschrieben hat; je Zeichen zählt die Klasse (alt = durchgestrichen, neu = Neu-Farbe ohne "
                   f"Strich, sonst normal). Zeichen: {_prozent(gesamt['format_treue_zeichen_anteil'])}. Eine automatisch geschriebene "
                   f"Zelle mit falscher Klasse ist zusätzlich ein stiller Fehler (Art format_falsch: {gesamt['stille_fehler_format']})"},
        {"nr": 3, "name": "Genauigkeit der Vorschläge (nach Freigabe)", "wert": _prozent(gesamt["vorschlags_genauigkeit"]),
         "ziel": "mindestens 95 %", "einheit": "Zellen",
         "ampel": _ampel(None if gesamt["vorschlags_genauigkeit"] is None else gesamt["vorschlags_genauigkeit"] >= 0.95,
                         "keine Vorschläge" if vorschlaege else "nicht gemessen", luecke=lue)},
        {"nr": 4, "name": "Abdeckung ohne Nachfrage (nur Haupt-Spalten)", "wert": _prozent(gesamt["abdeckung_ohne_nachfrage"]),
         "ziel": "mindestens 70 % (Standard-Aufbauten 90 %)", "einheit": "Zellen",
         "ampel": _ampel(None if gesamt["abdeckung_ohne_nachfrage"] is None else
                         gesamt["abdeckung_ohne_nachfrage"] >= 0.7 and (standard_abdeckung is None or standard_abdeckung >= 0.9),
                         luecke=lue),
         "detail": f"Zählt nur die {soll} Soll-Zellen der Haupt-Spalten (Nummer, Titel, Inhalt; ohne Formeln und Füllspalten), "
                   f"nicht alle {gesamt['soll_zellen_alle']}. Standard-Aufbauten: {len(standard)} Paare, davon "
                   f"{len(standard_angenommen)} ohne Profil (Standard-Aufbau angenommen), Abdeckung dort {_prozent(standard_abdeckung)} "
                   f"(Ziel 90 %). Die Zeile darunter zählt ALLE Soll-Zellen"},
        {"nr": 4, "name": "Abdeckung über ALLE Soll-Zellen (inkl. Füllspalten, Formeln)",
         "wert": f"{_prozent(gesamt['abdeckung_alle'])} ({erreicht_alle} von {gesamt['soll_zellen_alle']})",
         "ziel": "mindestens 70 %", "einheit": "Zellen",
         "ampel": _ampel(None if gesamt["abdeckung_alle"] is None else gesamt["abdeckung_alle"] >= 0.7, luecke=lue),
         "detail": f"Haupt-Spalten {s('erreicht_auto')} von {soll}; harmlose Füllung {gesamt['soll_zellen_unverwaltet_erreicht']} von "
                   f"{unverw}; sonstige nicht verwaltete {gesamt['soll_zellen_unverwaltet_sonstige_erreicht']} von "
                   f"{gesamt['soll_zellen_unverwaltet_sonstige']}; Formelzellen (getrennt ausgewiesen, v1 passt sie bewusst nicht an) "
                   f"{gesamt['erreicht_formel']} von {gesamt['soll_zellen_formel']}. Ohne Formelzellen: "
                   f"{_prozent(gesamt['abdeckung_alle_ohne_formeln'])}. „Erreicht“ heißt: Text UND Datentyp stimmen mit nachher überein"},
        {"nr": 4, "name": "Nicht-Übernahmen mit Grund", "wert": f"{gesamt['nicht_uebernommen_ohne_grund']} ohne Grund",
         "ziel": "0 ohne Grund", "einheit": "Änderungen",
         "ampel": _ampel(gesamt["nicht_uebernommen_ohne_grund"] == 0, luecke=lue)},
        {"nr": 4, "name": "Nicht verwaltete Spalten: Formeln und Sonstiges",
         "wert": f"{nicht_verfehlt} von {gesamt['soll_zellen_formel'] + gesamt['soll_zellen_unverwaltet_sonstige']} verfehlt",
         "ziel": "0 verfehlt", "einheit": "Zellen", "ampel": _ampel(nicht_verfehlt == 0, luecke=lue),
         "detail": f"Formeln {gesamt['soll_zellen_formel_verfehlt']} von {gesamt['soll_zellen_formel']}, "
                   f"sonstige {gesamt['soll_zellen_unverwaltet_sonstige_verfehlt']} von {gesamt['soll_zellen_unverwaltet_sonstige']}; "
                   f"harmlose Füllung nicht gezählt: {unverw}. v1 passt Formeln und nicht verwaltete Spalten bewusst nicht an; "
                   f"jede Abweichung zum Soll ist eine offene Aufgabe"},
        {"nr": 5, "name": "Lese-Lücke", "wert": _prozent(gesamt["lese_luecke_anteil"]), "ziel": "höchstens 2 %", "einheit": "Wörter",
         "ampel": _ampel(None if gesamt["lese_luecke_anteil"] is None else gesamt["lese_luecke_anteil"] <= 0.02, luecke=lue)},
        {"nr": 6, "name": "Zweiter Lauf ändert nichts (Idempotenz)", "wert": f"{gesamt['idempotent']} von {m}",
         "ziel": "alle Paare", "einheit": "Paare",
         "ampel": _ampel(None if not m else gesamt["idempotent"] == m, luecke=lue)},
        {"nr": 6, "name": "Kollateralschaden", "wert": str(gesamt["kollateral"]), "ziel": "0", "einheit": "Zellen",
         "ampel": _ampel(gesamt["kollateral"] == 0, luecke=lue)},
        {"nr": 6, "name": "Fremde Spec berührt", "wert": f"{gesamt['fremde_spec_beruehrt']} Zellen in "
                                                         f"{gesamt['paare_fremde_spec_beruehrt']} Paaren",
         "ziel": "0", "einheit": "Zellen",
         "ampel": _ampel(None if not gesamt["paare_mit_fremder_spec"] else gesamt["fremde_spec_beruehrt"] == 0, luecke=lue),
         "detail": f"{gesamt['paare_mit_fremder_spec']} Paare haben mehrere Specs im Register (Spec-Nr. des PDF, Blöcke mit anderer "
                   f"Spec-Nr.). Jede Änderung dort zählt auch als Kollateralschaden und stiller Fehler (Art fremde_spec)"},
        {"nr": 6, "name": "Struktur unverändert (nur die genannten Eigenschaften)", "wert": str(gesamt["struktur_geaendert"]),
         "ziel": "0", "einheit": "Befunde",
         "ampel": _ampel(None if not struktur_paare else gesamt["struktur_geaendert"] == 0, luecke=lue),
         "detail": f"geprüft: {STRUKTUR_GEPRUEFT}; nicht geprüft: {STRUKTUR_NICHT_GEPRUEFT}. „Ok“ heißt NICHT, dass die Excel "
                   f"sonst unverändert ist"},
        {"nr": 6, "name": "Vorschlags-Excel (alle freigegeben): Struktur unverändert",
         "wert": str(gesamt["struktur_vorschlaege_geaendert"]) if vorschlag_paare else "-", "ziel": "0 unerwartete",
         "einheit": "Befunde",
         "ampel": _ampel(None if not vorschlag_paare else gesamt["struktur_vorschlaege_geaendert"] == 0,
                         "keine Vorschläge" if vorschlaege else "nicht gemessen", luecke=lue),
         "detail": f"Dieselben Eigenschaften wie oben, an der Excel nach Freigabe aller Vorschläge. Eine Abweichung ist dann zulässig, wenn der "
                   f"Plan das Risiko vorab gemeldet hat (angekündigt: {gesamt['struktur_vorschlaege_angekuendigt']}); sonst ist sie ein Befund"},
        {"nr": 6, "name": "Originale unverändert", "wert": f"{gesamt['original_unveraendert']} von {m}", "ziel": "alle Paare",
         "einheit": "Paare", "ampel": _ampel(None if not m else gesamt["original_unveraendert"] == m, luecke=lue)},
        # Messung v2 (docs/ziel_v2.md): nur "ok" oder "Hinweis" -- aendert den Exit-Code nicht (Ziele gelten erst ab Etappe 2)
        {"nr": 7, "name": "Abdeckung ableitbare Zellen (Haupt + Füll-Konstante + Formel)",
         "wert": f"{_prozent(gesamt['abdeckung_ableitbar'])} ({gesamt['erreicht_ableitbar']} von {gesamt['soll_zellen_ableitbar']})",
         "ziel": "mindestens 85 % (Etappe 2)", "einheit": "Zellen",
         "ampel": _hinweis(gesamt["abdeckung_ableitbar"], lambda x: x >= ABLEITBAR_ZIEL),
         "detail": "Klassen: " + ", ".join(f"{k} {klassen_err[k]} von {klassen[k]}" for k in SOLL_KLASSEN)
                   + ". „fuell_konstant“ = Spalte hat in vorher für alle Zeilen desselben Typs denselben Wert (prüft die Messung "
                     "selbst); „nicht_ableitbar“ (z. B. Kommentar, laufende Nummer) zählt hier nicht mit"},
        {"nr": 7, "name": "Restarbeit je 100 Änderungen", "wert": _zahl(gesamt["restarbeit_je_100"]),
         "ziel": f"berichten (Richtwert höchstens {RESTARBEIT_RICHTWERT})", "einheit": "Änderungen (gewichtet)",
         "ampel": _hinweis(gesamt["restarbeit_je_100"], lambda x: x <= RESTARBEIT_RICHTWERT),
         "detail": RESTARBEIT_FORMEL + " Teile: " + ", ".join(f"{k} {rest[k]}" for k in RESTARBEIT_TEILE)},
    ]

    gesamt["kriterien_nicht_ok"] = sum(1 for k in kriterien if k["ampel"] == "NICHT ok")
    gesamt["ziel_verfehlt"] = gesamt["kriterien_nicht_ok"] > 0

    def einheitlich(feld: str) -> str:
        werte = sorted({str(p[feld]) for p in paare if p.get(feld)})
        return werte[0] if len(werte) == 1 else ("gemischt" if werte else "unbekannt")

    return {"format": BERICHT_FORMAT, "korpus": korpus, "generator_version": generator_version,
            "schema_version": SCHEMA_VERSION, "legende": einheitlich("legende"), "nachher_format": einheitlich("nachher_format"),
            "gesamt": gesamt, "fehlerarten": fehlerarten, "je_stoerer": je_stoerer,
            "je_erzeuger": je_erzeuger, "kriterien": kriterien, "paare": paare}


def _restarbeit(rest: dict[str, int], aenderungen: int) -> float | None:
    if not aenderungen:
        return None
    last = rest["pruefen"] + rest["nicht_uebernommen"] + rest["leere_zellen_neue_zeilen"] + rest["luecken_woerter"] / RESTARBEIT_LUECKEN_TEILER
    return round(100.0 * last / aenderungen, 1)


def _hinweis(wert: float | None, ok: Callable[[float], bool]) -> str:
    """Ampel fuer Kennzahlen, die nur berichtet werden: nie "NICHT ok" (Exit-Code bleibt unberuehrt)."""
    if wert is None:
        return "nicht gemessen"
    return "ok" if ok(wert) else "Hinweis"


def _zahl(x: float | None) -> str:
    return "-" if x is None else f"{x:.1f}".replace(".", ",")


# ---------------------------------------------------------------- Vergleich mit einem frueheren Bericht (--vergleich)
VERGLEICH_HAUPT = ("stille_fehler", "regel27_verstoesse", "kollateral", "treffergenauigkeit_auto", "vorschlags_genauigkeit",
                   "abdeckung_ohne_nachfrage", "standard_abdeckung", "abdeckung_alle", "abdeckung_ableitbar", "lese_luecke_anteil",
                   "idempotent", "restarbeit_je_100", "aenderungen_zu_pruefen", "aenderungen_nicht_uebernommen", "format_treue")


def _numerisch(x: Any) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _delta(a: Any, b: Any) -> dict[str, Any]:
    return {"vorher": a, "jetzt": b, "delta": round(b - a, 4) if _numerisch(a) and _numerisch(b) else None}


def vergleiche(alt: dict[str, Any], neu: dict[str, Any], gegen: str) -> dict[str, Any]:
    """Deltas jeder Kennzahl (gesamt), je Stoerer-Familie und je Erzeuger -- nur Zahlen."""
    ga, gn = alt.get("gesamt") or {}, neu.get("gesamt") or {}
    kennzahlen = {k: _delta(ga.get(k), v) for k, v in gn.items()
                  if _numerisch(v) and (k in ga and (_numerisch(ga[k]) or ga[k] is None))}
    for k in ("vorschlags_genauigkeit", "abdeckung_ableitbar", "restarbeit_je_100"):
        if k in gn and k not in kennzahlen:
            kennzahlen[k] = _delta(ga.get(k), gn.get(k))

    def je(name: str, felder: tuple[str, ...]) -> dict[str, dict[str, Any]]:
        a, b = alt.get(name) or {}, neu.get(name) or {}
        out: dict[str, dict[str, Any]] = {}
        for fam in sorted(set(a) | set(b)):
            out[fam] = {f: _delta((a.get(fam) or {}).get(f), (b.get(fam) or {}).get(f)) for f in felder}
        return out

    return {"gegen": gegen, "korpus_vorher": alt.get("korpus"),
            "generator_version_vorher": alt.get("generator_version"),
            "vergleichbar": alt.get("generator_version") == neu.get("generator_version") and alt.get("korpus") == neu.get("korpus"),
            "kennzahlen": kennzahlen,
            "je_stoerer": je("je_stoerer", ("paare", "stille_fehler", "abdeckung_ohne_nachfrage")),
            "je_erzeuger": je("je_erzeuger", ("paare", "stille_fehler", "treffergenauigkeit_auto", "abdeckung_ohne_nachfrage",
                                              "restarbeit_je_100"))}


def finde_bericht(korpus: str, gegen: str | None = None) -> Path | None:
    """Ordner des letzten Berichts desselben Korpus (oder des Berichts NNN bei --gegen) unter messung_dir."""
    basis = config.pfad("messung_dir")
    if not basis.is_dir():
        return None
    ordner = [(int(m.group(1)), d) for d in basis.iterdir()
              if d.is_dir() and (m := re.match(r"^(\d{3,})_(.+)$", d.name)) and (d / "bericht.json").is_file()]
    if gegen is not None:
        try:
            nr = int(gegen)
        except ValueError:
            return None
        return next((d for n, d in sorted(ordner) if n == nr), None)
    passend = [(n, d) for n, d in ordner if d.name.split("_", 1)[1] == korpus]
    return max(passend)[1] if passend else None


def _prozent(x: float | None) -> str:
    return "-" if x is None else f"{x * 100:.1f} %".replace(".", ",")


def _md(b: dict[str, Any]) -> str:
    g = b["gesamt"]
    z = [f"# Messbericht {b['korpus']}", "",
         f"Paare: {g['paare']} (davon nicht oder nicht vollständig messbar: {g['paare_abgestuerzt']}) · Soll-Zellen (Haupt-Spalten): "
         f"{g['soll_zellen']} · Schema {b['schema_version']} · Generator-Version: {b['generator_version'] or 'unbekannt'}", "",
         f"Legende: {b.get('legende', 'unbekannt')} · nachher-Format: {b.get('nachher_format', 'unbekannt')} (aus KORPUS.json bzw. "
         f"Aufruf; „markiert“ = nachher.xlsx trägt die Markierung, die Format-Treue wird mitgemessen) · vom Leser gewählte Legende "
         f"(Paare): {', '.join(f'{k} {n}' for k, n in (g.get('legende_erkannt') or {}).items()) or '-'} · Markup-Leiste erkannt: "
         f"{g.get('markup_leiste_erkannt', 0)} Paare", "",
         "Dieser Bericht enthält keine Texte aus Dokumenten, nur Zahlen, Positionen und Fehlerarten.", "",
         "## Kriterien (docs/ziel_v1.md)", "",
         "Die Spalte „Einheit“ sagt, WAS gezählt wird (Zellen, Änderungen, Paare, Wörter ...): Zahlen verschiedener Einheiten "
         "sind nicht miteinander vergleichbar.", "",
         "| Nr | Kriterium | Wert | Einheit | Ziel | Ergebnis |", "|---|---|---|---|---|---|"]
    for k in b["kriterien"]:
        z.append(f"| {k['nr']} | {k['name']} | {k['wert']} | {k.get('einheit', '')} | {k['ziel']} | {k['ampel']} |")
    z += ["", "Erläuterungen:", ""] + [f"- {k['name']}: {k['detail']}" for k in b["kriterien"] if k.get("detail")]
    z += ["", f"**Struktur ok heißt nur:** Es wurde keine unerwartete Abweichung bei diesen Eigenschaften gefunden: "
              f"{g['struktur_geprueft']}. Nicht geprüft wurde: {g['struktur_nicht_geprueft']}. Eine Ampel „ok“ bei Struktur sagt also "
              f"nicht, dass die Excel sonst unverändert ist.", ""]
    z += [f"Abdeckung: Die Zeile „ohne Nachfrage (nur Haupt-Spalten)“ zählt {g['soll_zellen']} von {g['soll_zellen_alle']} Soll-Zellen; "
          f"die Zeile „ALLE Soll-Zellen“ zählt alle ({_prozent(g['abdeckung_alle'])}). Standard-Aufbau: {g['standard_paare']} Paare, davon "
          f"{g['standard_paare_angenommen']} ohne Profil (Standard-Aufbau angenommen, damit echte Paare nicht milder gemessen werden "
          f"als erfundene).", ""]
    z += [f"Ergebnis für den Exit-Code: {g['kriterien_nicht_ok']} Kriterium/Kriterien auf „NICHT ok“ "
          f"({'Code 1: Ziel verfehlt' if g['kriterien_nicht_ok'] else 'Code 0: kein Kriterium verfehlt'}); Code 2 bedeutet kaputte Eingabe.", ""]
    ae_ges = g["aenderungen_uebernommen_korrekt"] + g["aenderungen_uebernommen_falsch"]
    z += ["", "## Zahlen (mit Zähleinheit)", "",
          f"- Automatisch richtig: {g['auto_richtig']} von {g['auto_gesamt']} (Einheit: Zellen)",
          f"- Änderungen: {g['aenderungen']} insgesamt (Einheit: Änderungen) · übernommen und korrekt: "
          f"{g['aenderungen_uebernommen_korrekt']} · übernommen und falsch: {g['aenderungen_uebernommen_falsch']} · "
          f"zu prüfen: {g['aenderungen_zu_pruefen']} · nicht übernommen: {g['aenderungen_nicht_uebernommen']} "
          f"(von den übernommenen sind {_prozent(g['treffergenauigkeit_aenderungen'])} korrekt, {ae_ges} übernommene)",
          f"- Anteil „zu prüfen“: {_prozent(g['anteil_pruefen'])} · Anteil „nicht übernommen“: {_prozent(g['anteil_nicht_uebernommen'])}"
          f" · Soll-Zellen ohne passende Änderung: {g['ohne_aenderung_zellen']} (Einheit: Zellen)",
          f"- Zellen, die nur auf die Freigabe einer offenen Änderung warten (kein Fehler): {g['zellen_warten_auf_freigabe']} (Einheit: Zellen)",
          f"- Soll-Zellen außerhalb der Haupt-Spalten (Einheit: Zellen): Formeln {g['soll_zellen_formel']} "
          f"(verfehlt {g['soll_zellen_formel_verfehlt']}) · sonstige nicht verwaltete {g['soll_zellen_unverwaltet_sonstige']} "
          f"(verfehlt {g['soll_zellen_unverwaltet_sonstige_verfehlt']}) · harmlose Füllung {g['soll_zellen_unverwaltet']} "
          f"(neue Zeilen {g['soll_zellen_unverwaltet_aufschluesselung']['neue_zeile']}, "
          f"entfernte Zeilen {g['soll_zellen_unverwaltet_aufschluesselung']['entfernte_zeile']}) · alle Soll-Zellen zusammen: "
          f"{g['soll_zellen_alle']}",
          f"- Soll-Zellen mit Zeichen, die ein PDF nicht tragen kann (geschützte Leerzeichen u. ä.; nur so verglichen): "
          f"{g['soll_zellen_pdf_blind']}, davon erreicht {g['soll_zellen_pdf_blind_erreicht']} (Einheit: Zellen)",
          f"- Bilanz „vollständig“: {g['vollstaendig']} von {g['paare']} Paaren (Einheit: Paare)",
          f"- Soll-Wörter: {g['soll_woerter']} · nicht erkannt (Lese-Lücke): {g['lese_luecke_woerter']} (Einheit: Wörter)",
          f"- Lücken-Prüfung (erwartet.json): bei {g['luecken_pruefbar_paare']} von {g['paare']} Paaren möglich; "
          f"ohne erwartet.json kann eine nicht gemeldete Lücke NICHT auffallen",
          f"- Struktur ({g['struktur_paare_geprueft']} Paare geprüft, Einheit: Befunde): {g['struktur_geaendert']} unerwartet verändert "
          + ", ".join(f"{a} {n}" for a, n in g["struktur_nach_art"].items() if n),
          "", "## Stille Fehler nach Art", "",
          f"Zusammensetzung: {g['stille_fehler_zellen']} Zellen · {g['stille_fehler_paar_befunde']} Paar-Befunde · "
          f"{g['regel27_verstoesse']} Regel-27-Verstöße (alle zusammen {g['stille_fehler']}); betroffene Änderungen: "
          f"{g['stille_fehler_aenderungen']}; davon Zellen, die sich nur in Leerraum/Zeichenform/Absatzgrenzen unterscheiden: "
          f"{g['stille_fehler_nur_textform']}", ""]
    z += [f"- {a}: {n}" for a, n in g["stille_fehler_nach_art"].items()]
    z += ["", "## Fehlerarten", ""] + [f"- {a}: {n}" for a, n in b["fehlerarten"].items()]
    z += ["", "## Messung v2: Klassen, Restarbeit, Prüfcodes", "",
          f"- Soll-Zellen nach Klasse (erreicht von gesamt, Einheit: Zellen): "
          + ", ".join(f"{k} {g['soll_klassen_erreicht'][k]} von {g['soll_klassen'][k]}" for k in SOLL_KLASSEN)
          + f" · ableitbar (haupt + fuell_konstant + formel): {_prozent(g['abdeckung_ableitbar'])}",
          f"- Restarbeit je 100 Änderungen: {_zahl(g['restarbeit_je_100'])} (Teile: "
          + ", ".join(f"{k} {g['restarbeit'][k]}" for k in RESTARBEIT_TEILE) + ")",
          f"- Formel: {g['restarbeit_formel']}",
          "- Prüfcodes (Anzahl Änderungen je Code): " + (", ".join(f"{c} {n}" for c, n in g["pruefcodes"].items()) or "keine"),
          "- Sperr-Quote je Code (gesperrte Änderungen, die nach Freigabe richtig gewesen wären; nur mit --mit-vorschlaegen messbar): "
          + (", ".join(f"{c} {e['richtig_nach_freigabe']} von {e['messbar']} ({_prozent(e['sperr_quote'])}; gesperrt {e['gesperrt']})"
                       for c, e in g["sperr_quote"].items()) or "keine gesperrten Änderungen")]
    if b.get("je_erzeuger"):
        z += ["", "## Je PDF-Erzeuger", "",
              "| Erzeuger | Paare | Stille Fehler | Regel 27 | Treffergenauigkeit | Abdeckung | Lese-Lücke | Idempotent | Restarbeit/100 |",
              "|---|---|---|---|---|---|---|---|---|"]
        for name, e in sorted(b["je_erzeuger"].items()):
            z.append(f"| {name} | {e['paare']} | {e['stille_fehler']} | {e['regel27_verstoesse']} | "
                     f"{_prozent(e['treffergenauigkeit_auto'])} | {_prozent(e['abdeckung_ohne_nachfrage'])} | "
                     f"{_prozent(e['lese_luecke_anteil'])} | {e['idempotent']} | {_zahl(e['restarbeit_je_100'])} |")
    v = b.get("vergleich")
    if v:
        z += ["", f"## Vergleich mit Bericht {v['gegen']}", ""]
        if not v["vergleichbar"]:
            z.append(f"Achtung: anderer Korpus oder andere Generator-Version (vorher {v['korpus_vorher']}, Generator "
                     f"{v['generator_version_vorher']}) -- die Zahlen sind nur bedingt vergleichbar.")
            z.append("")
        z += ["| Kennzahl | vorher | jetzt | Delta |", "|---|---|---|---|"]
        for k in VERGLEICH_HAUPT:
            d = v["kennzahlen"].get(k)
            if d:
                z.append(f"| {k} | {d['vorher']} | {d['jetzt']} | {d['delta']} |")
        geaendert = {f: d for f, d in v["je_stoerer"].items() if any(x["delta"] for x in d.values())}
        if geaendert:
            z += ["", "| Störer | Δ Paare | Δ stille Fehler | Δ Abdeckung |", "|---|---|---|---|"]
            for f, d in geaendert.items():
                z.append(f"| {f} | {d['paare']['delta']} | {d['stille_fehler']['delta']} | {d['abdeckung_ohne_nachfrage']['delta']} |")
    if b["je_stoerer"]:
        z += ["", "## Je Störer-Familie", "", "| Störer | Paare | Stille Fehler | Abdeckung |", "|---|---|---|---|"]
        for name, e in sorted(b["je_stoerer"].items()):
            z.append(f"| {name} | {e['paare']} | {e['stille_fehler']} | {_prozent(e['abdeckung_ohne_nachfrage'])} |")
    z += ["", "## Paare", "", "| Paar | Soll-Zellen | Abdeckung | Stille Fehler | Regel 27 | Struktur | Plausibilität | Bilanz | Fehler |",
          "|---|---|---|---|---|---|---|---|---|"]
    for p in b["paare"]:
        ab = _prozent(_anteil(p["erreicht_auto"], p["soll_zellen"]))
        bil = "nicht messbar" if p["bilanz_stimmt"] is None else ("stimmt" if p["bilanz_stimmt"] else "NICHT")
        if p.get("teilweise_gemessen"):
            bil += " (teilweise gemessen)"
        fe = "; ".join(" ".join(x for x in (f"Blatt {f['blatt']} Z{f['zeile']}{_ort_md(f)} S{f['spalte']}", f["aenderung"], f["art"],
                                            f"({f['grund']})" if f.get("grund") else "") if x) for f in p["fehler"])
        struktur = "-" if p.get("struktur_geaendert") is None else str(p["struktur_geaendert"])
        plausi = p.get("plausibilitaet_hinweise")
        plausi_t = "-" if plausi is None else (str(p.get("plausibilitaet_ungemeldet")) + " von " + str(len(plausi)) if plausi else "0")
        z.append(f"| {p['name']} | {p['soll_zellen']} | {ab} | {p['stille_fehler']} | {p.get('regel27_verstoesse', 0)} | "
                 f"{struktur} | {plausi_t} | {bil} | {fe or '-'} |")
    if any("fehler_schablonen" in p for p in b["paare"]):
        z += _md_schablonen(b["paare"])
    return "\n".join(z) + "\n"


def _ort_md(f: dict[str, Any]) -> str:
    """Zeile in nachher (N) und im Ergebnis (E), wenn sie von der vorher-Zeile abweicht oder es eine neue Zeile ist."""
    zn, ze = f.get("zeile_nachher"), f.get("zeile_ergebnis")
    if not f.get("neue_zeile") and zn == f.get("zeile") and ze == f.get("zeile"):
        return ""
    return f" (N{zn if zn is not None else '-'}/E{ze if ze is not None else '-'}{', neu' if f.get('neue_zeile') else ''})"


def _md_schablonen(paare: list[dict[str, Any]]) -> list[str]:
    """Abschnitt "Fehler-Schablonen" (nur mit --schablonen): Formen Soll/Ist, kein Klartext."""
    z = ["", "## Fehler-Schablonen", "",
         "Formen statt Text: A = Großbuchstabe, a = Kleinbuchstabe, 9 = Ziffer, · = geschütztes Leerzeichen, ¶ = Zeilenumbruch, ⇥ = Tab, "
         "W… = langes Wort, ? = anderes Zeichen. Vor jeder Form steht das Format: `-:` normal, `a:` alt (durchgestrichen), `n:` neu; die Läufe sind durch ‖ getrennt.", ""]
    laeufe = lambda ls: " ‖ ".join(f"{x['f']}:{x['form']}" for x in ls) or "(leer)"          # noqa: E731
    for p in paare:
        if not p.get("fehler_schablonen") and not p.get("fehler_schablonen_verworfen"):
            continue
        z += [f"### {p['name']}", "", "| Quelle | Fehlerart | Änderungstyp | Erzeuger | Soll | Ist |", "|---|---|---|---|---|---|"]
        for e in p["fehler_schablonen"]:
            z.append(f"| {e['quelle']} | {e['fehlerart']} | {e['aenderungstyp']} | {e['erzeuger']} | "
                     f"`{laeufe(e['soll'])}` | `{laeufe(e['ist'])}` |")
        verworfen = p.get("fehler_schablonen_verworfen") or 0
        if verworfen > 0:
            z += ["", f"Verworfen (Positivliste nicht bestanden): {verworfen}"]
        elif verworfen < 0:
            z += ["", "Die Fehler-Schablonen dieses Paars konnten nicht erstellt werden."]
        z.append("")
    return z


def schreibe_bericht(bericht: dict[str, Any], korpus: str) -> Path:
    """Legt einen NEUEN Ordner NNN_<korpus> an (append-only) und schreibt bericht.json + bericht.md."""
    basis = config.pfad("messung_dir")
    basis.mkdir(parents=True, exist_ok=True)
    nummern = [int(m.group(1)) for d in basis.iterdir() if (m := re.match(r"^(\d{3,})_", d.name))]
    nr = max(nummern, default=0) + 1
    while True:
        ziel = basis / f"{nr:03d}_{korpus}"
        try:
            ziel.mkdir(exist_ok=False)
            break
        except FileExistsError:
            nr += 1
    with open(ziel / "bericht.json", "x", encoding="utf-8") as f:
        json.dump(bericht, f, ensure_ascii=False, indent=1)
    with open(ziel / "bericht.md", "x", encoding="utf-8") as f:
        f.write(_md(bericht))
    return ziel


def _konsole(b: dict[str, Any], ziel: Path, vorschlaege: bool) -> None:
    g = b["gesamt"]
    print(f"Messung {b['korpus']}: {g['paare']} Paare, {g['soll_zellen']} Soll-Zellen (Haupt-Spalten)")
    print(f"Legende: {b.get('legende', 'unbekannt')} · nachher-Format: {b.get('nachher_format', 'unbekannt')} · vom Leser gewählt: "
          + (", ".join(f"{k} {n}" for k, n in (g.get("legende_erkannt") or {}).items()) or "-"))
    if g["paare_abgestuerzt"]:
        print(f"Achtung: {g['paare_abgestuerzt']} Paar(e) konnten nicht (vollständig) gemessen werden "
              f"({g['paare_teilweise_gemessen']} davon teilweise: ihre Funde zählen mit). Die Ampeln stehen dann nicht auf „ok“.")
    print("")
    breite = max(len(k["name"]) for k in b["kriterien"])
    for k in b["kriterien"]:
        print(f"  [{k['nr']}] {k['name']:<{breite}}  {k['wert']}  ({k.get('einheit', '')})  Ziel: {k['ziel']}  ->  {k['ampel']}")
    print("")
    print(f"  Automatisch richtig: {g['auto_richtig']} von {g['auto_gesamt']} Zellen  |  Änderungen: "
          f"{g['aenderungen_uebernommen_korrekt']} übernommen & korrekt, {g['aenderungen_uebernommen_falsch']} übernommen & falsch, "
          f"{g['aenderungen_zu_pruefen']} zu prüfen, {g['aenderungen_nicht_uebernommen']} nicht übernommen")
    print(f"  Anteil zu prüfen: {_prozent(g['anteil_pruefen'])}  |  nicht übernommen: {_prozent(g['anteil_nicht_uebernommen'])}"
          f"  |  ohne passende Änderung: {g['ohne_aenderung_zellen']}  |  warten nur auf Freigabe: {g['zellen_warten_auf_freigabe']}")
    print(f"  Nicht verwaltet: Formeln {g['soll_zellen_formel']} (verfehlt {g['soll_zellen_formel_verfehlt']}), "
          f"sonstige {g['soll_zellen_unverwaltet_sonstige']} (verfehlt {g['soll_zellen_unverwaltet_sonstige_verfehlt']}), "
          f"harmlose Füllung {g['soll_zellen_unverwaltet']}")
    print("  Fehlerarten: " + ", ".join(f"{a} {n}" for a, n in b["fehlerarten"].items()))
    if g["luecken_nicht_pruefbar_paare"]:
        print("")
        print(f"HINWEIS Kriterium 1: Bei {g['luecken_nicht_pruefbar_paare']} von {g['paare']} Paaren fehlt erwartet.json. "
              "Ohne sie kann eine nicht gemeldete Lücke (Scan, Zwei-Spalten, Tabelle, Bild) NICHT auffallen: "
              "Kriterium 1 ist dort nicht prüfbar (Ampel „nicht gemessen“).")
    if g["plausibilitaet_ungemeldet"]:
        print(f"HINWEIS Plausibilität (direkt aus dem PDF, nur Hinweis): {g['plausibilitaet_ungemeldet']} auffällige Seite(n) "
              f"(ohne Textzeichen oder mit großem Bild) ohne gemeldete Lücke. Bitte von Hand ansehen.")
    if g["struktur_geaendert"]:
        print(f"ACHTUNG Struktur: {g['struktur_geaendert']} unerwartet veränderte Eigenschaft(en) "
              + ", ".join(f"{a} {n}" for a, n in g["struktur_nach_art"].items() if n))
    if g["struktur_vorschlaege_geaendert"]:
        print(f"ACHTUNG Struktur (Excel mit allen freigegebenen Vorschlägen): {g['struktur_vorschlaege_geaendert']} unerwartete "
              f"Abweichung(en) ohne Vorab-Warnung des Plans: "
              + ", ".join(f"{a} {n}" for a, n in g["struktur_vorschlaege_nach_art"].items() if n))
    print("")
    print("Struktur „ok“ heißt nur: keine unerwartete Abweichung bei den geprüften Eigenschaften (Blätter, versteckte Zeilen/Spalten, "
          "Verbünde, Validierung, bedingte Formatierung, Hyperlink-Adressen, Kommentare, benannte Bereiche, Filter, Tabellen, Fixierung, Breiten, "
          "Höhen, Zahlenformat, Schrift, Datentyp). Nicht geprüft: Rahmen, Ausrichtung, Schutz, Druckeinstellungen, Bilder, Diagramme u. a.")
    if g["standard_paare_angenommen"]:
        print(f"Hinweis: Bei {g['standard_paare_angenommen']} Paar(en) ohne Profil wurde der Standard-Aufbau angenommen "
              f"(90-%-Regel der Abdeckung gilt auch dort).")
    if g["stille_fehler"]:
        print("")
        print(f"ACHTUNG: {g['stille_fehler']} stille(r) Fehler ({g['stille_fehler_zellen']} Zellen, "
              f"{g['stille_fehler_paar_befunde']} Paar-Befunde, {g['regel27_verstoesse']} Regel-27-Verstöße). "
              "Das Tool hat etwas als erledigt gemeldet, das nicht stimmt.")
        print("  Einzelheiten stehen im Bericht (Paar-Name, Blatt, Zeile, Spalte, Änderungs-ID).")
    if not vorschlaege:
        print("")
        print("Hinweis: Die Qualität der Vorschläge wurde nicht gemessen (Option --mit-vorschlaegen).")
    print(f"  Abdeckung ableitbare Zellen: {_prozent(g['abdeckung_ableitbar'])}  |  Restarbeit je 100 Änderungen: "
          f"{_zahl(g['restarbeit_je_100'])}  |  Prüfcodes: {sum(g['pruefcodes'].values())}")
    if len(b.get("je_erzeuger") or {}) > 1 or "libreoffice" in (b.get("je_erzeuger") or {}):
        for name, e in sorted(b["je_erzeuger"].items()):
            print(f"  Erzeuger {name}: {e['paare']} Paare, {e['stille_fehler']} stille Fehler, Treffergenauigkeit "
                  f"{_prozent(e['treffergenauigkeit_auto'])}, Abdeckung {_prozent(e['abdeckung_ohne_nachfrage'])}, "
                  f"Lese-Lücke {_prozent(e['lese_luecke_anteil'])}")
    v = b.get("vergleich")
    if v:
        print("")
        print(f"Vergleich mit {v['gegen']}" + ("" if v["vergleichbar"] else " (anderer Korpus/Generator: nur bedingt vergleichbar)") + ":")
        for k in VERGLEICH_HAUPT:
            d = v["kennzahlen"].get(k)
            if d and d["delta"]:
                print(f"  {k}: {d['vorher']} -> {d['jetzt']} (Delta {d['delta']})")
        d = v["kennzahlen"].get("stille_fehler")
        if d and d["delta"] and d["delta"] > 0:
            print(f"ACHTUNG Stop-Regel: Die stillen Fehler sind gestiegen ({d['vorher']} -> {d['jetzt']}).")
    if b["korpus"] in UNGESEHEN:
        print(f"Hinweis (Regel 11): Den Korpus {b['korpus']} nie zum Einstellen des Tools benutzen (genau einmal messen).")
    if b["korpus"] == "ex_test":
        print("Hinweis: ex_test ist der frühere Testkorpus (Seeds 1001–1030). Er ist offen als Entwicklungskorpus erklärt "
              "und nicht mehr ungesehen.")
    print("")
    print(f"Bericht: {ziel / 'bericht.md'}")
    if g["kriterien_nicht_ok"]:
        print(f"Ergebnis: Ziel verfehlt, {g['kriterien_nicht_ok']} Kriterium/Kriterien stehen auf „NICHT ok“ (Exit-Code 1).")
    else:
        print("Ergebnis: kein Kriterium verfehlt (Exit-Code 0).")

# ======================================================================================
# 5b. Rueckbau-Messung (messen --nachher): Paare mit nur PDF + nachher.xlsx (z. B. Nicos echte Paare)
# ======================================================================================
def _rueck_art(schrift: Any, alt_rgb: str, neu_rgb: str) -> str:
    """Markierung der LETZTEN Runde: alt = Alt-Farbe + durchgestrichen, neu = Neu-Farbe ohne Strich; alles andere bleibt
    (auch Markierungen frueherer Runden in anderen Farben)."""
    if schrift is None:
        return "sonst"
    rgb = _farb_rgb(getattr(schrift, "color", None))
    if getattr(schrift, "strike", None) and rgb == alt_rgb:
        return "alt"
    if not getattr(schrift, "strike", None) and rgb == neu_rgb:
        return "neu"
    return "sonst"


def _rueck_teile(zelle: Any, alt_rgb: str, neu_rgb: str) -> list[tuple[str, str, Any]] | None:
    """[(Text, alt|neu|sonst, Laufschrift oder None)] einer Zelle; None = leer."""
    v = zelle.value
    if v is None:
        return None
    if isinstance(v, CellRichText):
        return [(b.text, _rueck_art(b.font, alt_rgb, neu_rgb), b.font) if isinstance(b, TextBlock)
                else (str(b), _rueck_art(zelle.font, alt_rgb, neu_rgb), None) for b in v]
    return [(v if isinstance(v, str) else _zelltext(v), _rueck_art(zelle.font, alt_rgb, neu_rgb), None)]


def _ohne_markierung(f: Any) -> Any:
    """Dieselbe Schrift ohne Farbe und Strich (Zellschrift bzw. Laufschrift)."""
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font
    if f is None:
        return None
    if isinstance(f, InlineFont):
        return InlineFont(rFont=f.rFont, sz=f.sz, b=f.b, i=f.i, u=f.u, family=f.family, charset=f.charset, scheme=f.scheme,
                          vertAlign=f.vertAlign)
    return Font(name=f.name, sz=f.sz, b=f.b, i=f.i, u=f.u, family=f.family, charset=f.charset, scheme=f.scheme,
                vertAlign=f.vertAlign)


def _zelle_zurueck(zelle: Any, teile: list[tuple[str, str, Any]]) -> None:
    """Rueckbau einer Zelle: Alt-Laeufe werden normaler Text, Neu-Laeufe fallen weg, alles andere bleibt unangetastet."""
    if not isinstance(zelle.value, CellRichText):
        art = teile[0][1]
        if art == "neu":
            zelle.value = None
        zelle.font = _ohne_markierung(zelle.font)
        return
    rest = [(t, a, f) for t, a, f in teile if a != "neu" and t]
    if not rest:
        zelle.value = None
        return
    schriften = [_ohne_markierung(f) if a == "alt" else f for _t, a, f in rest]
    if all(f is None or (_farb_rgb(getattr(f, "color", None)) is None and not getattr(f, "strike", None)) for f in schriften):
        zelle.value = "".join(t for t, _a, _f in rest)               # nur noch Grundschrift: einfacher Text (wie vor der Runde)
        return
    zelle.value = CellRichText(*[TextBlock(f, t) if f is not None else t for (t, _a, _f), f in zip(rest, schriften)])


def baue_vorher_zurueck(nachher: Path, ziel: Path) -> dict[str, int]:
    """Baut aus einer markierten nachher.xlsx den Stand VOR dieser Runde zurueck (nur Markierungen der letzten Runde laut
    config/zuordnung.yaml markierung.alt_farbe/neu_farbe): alt (+ Strich) -> normaler Text, neu (ohne Strich) -> weg. Zeilen, deren
    laengste Zelle ganz neu ist und die nichts Altes tragen, sind neue Zeilen und entfallen. Zellen ohne diese Markierung bleiben
    unangetastet. Nur fuer die Zellen-Messung: Zeichnungen u. ae. gehen beim Speichern mit openpyxl verloren."""
    alt_rgb, neu_rgb = _markierungsfarben()
    wb = load_workbook(str(nachher), rich_text=True)
    zahl: Counter = Counter()
    try:
        for ws in wb.worksheets:
            neue_zeilen: list[int] = []
            for row in ws.iter_rows():
                laengste: tuple[int, bool] | None = None
                altes = False
                for c in row:
                    teile = _rueck_teile(c, alt_rgb, neu_rgb)
                    text = "".join(t for t, _a, _f in teile or [])
                    if not teile or not text.strip():
                        continue
                    arten = {a for t, a, _f in teile if t}
                    ganz_neu = arten == {"neu"}
                    if laengste is None or len(text) > laengste[0]:
                        laengste = (len(text), ganz_neu)
                    altes |= not ganz_neu and bool(arten & {"alt", "neu"})
                    if arten & {"alt", "neu"}:
                        _zelle_zurueck(c, teile)
                        zahl["zellen"] += 1
                if laengste and laengste[1] and not altes:
                    neue_zeilen.append(row[0].row)
            for r in sorted(neue_zeilen, reverse=True):
                ws.delete_rows(r)
                zahl["zeilen"] += 1
        wb.save(str(ziel))
    finally:
        wb.close()
    return {"zellen_zurueckgebaut": zahl["zellen"], "zeilen_entfernt": zahl["zeilen"]}


def _rueckbau_kategorie(aw: "_Auswertung", e: dict[str, Any]) -> str:
    if e["erreicht"]:
        return "gleich"
    platz = e["ist_da"] and e["ist"] is not None
    if e["art"] == "neu":
        platz = platz and bool(aw.neu_abgleich[e["bi"]][1].get(e["si"]))
    if platz and _gleich_bis_leerraum(e["ist"], e["soll"]):
        return "gleich_bis_leerraum"
    if not e["ist_da"]:
        entf, dazu = aw.soll_woerter[e["key"]]
        verwandt = aw._zugehoerig(Counter(dazu), Counter(entf)) + [a for a in aw.aend if any(
            z.bi == e["bi"] and z.zeile == e["zeile"] and z.spalte in (None, e["spalte"]) for z in aw.ziele_von[a.id])]
        return "tool_fehlt_vorschlag_offen" if any(a.status == PRUEFEN for a in verwandt) else "tool_fehlt_gar_nicht"
    return "verschieden"


def _rueckbau_form(text: Any, pos: int) -> list[dict[str, str]]:
    return _formen(text, pos) if text else []


def messe_nachher(ordner: Path, einstellungen: dict[str, str] | None = None, maximal: int = 40) -> dict[str, Any]:
    """Rueckbau-Messung eines Paars, das nur aenderungen.pdf + nachher.xlsx hat (ein vorhandenes vorher.xlsx wird ignoriert):
    vorher zurueckbauen, Tool laufen lassen, Zelle fuer Zelle mit nachher vergleichen. Ergebnis: nur Zahlen, Positionen und
    Formen (Schablonen), nie Text (Regel 16)."""
    from . import steckbrief as sb
    pdf, nachher = ordner / "aenderungen.pdf", ordner / "nachher.xlsx"
    for p in (pdf, nachher):
        if not p.is_file():
            raise FileNotFoundError(p.name)
    einst = dict(korpus_einstellungen(ordner.parent) if einstellungen is None else einstellungen)
    einst["nachher_format"] = "markiert"                 # Rueckbau setzt eine markierte nachher.xlsx voraus
    hash_vorher = {p.name: _sha(p) for p in (pdf, nachher)}
    t0 = time.monotonic()
    with konfig_ueberschreibung(einst), tempfile.TemporaryDirectory(prefix="cde_rueckbau_") as tmp_name:
        wirksam = _wirksam()
        tmp = Path(tmp_name)
        rueck = tmp / "vorher_rueckbau.xlsx"
        stats = baue_vorher_zurueck(nachher, rueck)
        pa = analysiere_pdf(pdf)
        plan = plane(pa, ExcelModell(rueck))
        konsens.bildprobe_anwenden(plan.aenderungen, pa, pdf=pdf)
        bil_plan = bilanz(plan.aenderungen, plan.markierte_woerter, plan.ignoriert, plan.luecken)
        vor_status = {a.id: a.status for a in plan.aenderungen}
        aend = copy.deepcopy(plan.aenderungen)
        ergebnis = tmp / "ergebnis.xlsx"
        uebernehmen(aend, rueck, ergebnis, None, absaetze=plan.absaetze)
        bil = bilanz(aend, plan.markierte_woerter, plan.ignoriert, plan.luecken)
        _plausi, pdf_zeichen = _pdf_pruefung(pdf)
        aw = _Auswertung(_lese_excel(rueck, True), _lese_excel(nachher, True), _lese_excel(ergebnis, True), None,
                         _blattnamen(rueck), aend, plan, pa, bil, bil_plan, None, None, vor_status, pdf_zeichen, markiert=True,
                         pdf_specs=_pdf_specs(pdf))
    kat: Counter = Counter()
    kat_haupt: Counter = Counter()
    schablonen: list[dict[str, Any]] = []
    belegt: set[Any] = set()
    for e in aw.einheiten:
        belegt.add(e["ist_key"] if e["art"] == "neu" else e["key"])
        k = _rueckbau_kategorie(aw, e)
        kat[k] += 1
        if e["klasse"] == "haupt":
            kat_haupt[k] += 1
        if k != "gleich":
            pos = _erste_abweichung(e["soll"], e["ist"] or "")
            schablonen.append({"kategorie": k, "blatt": e["bi"] + 1, "zeile": e["zeile"], "spalte": e["spalte"],
                               "neue_zeile": e["art"] == "neu", "soll": _rueckbau_form(e["soll"], pos),
                               "ist": _rueckbau_form(e["ist"] or "", pos)})
    fremd_extra = 0
    for it in aw.ist:
        if it["key"] in belegt:
            continue
        kat["tool_extra"] += 1
        fremd_extra += it["zeile"] in aw.fremd[it["bi"]]
        if aw._verwaltet(it["bi"], it["spalte"]):
            kat_haupt["tool_extra"] += 1
        schablonen.append({"kategorie": "tool_extra", "blatt": it["bi"] + 1, "zeile": it["zeile"], "spalte": it["spalte"],
                           "neue_zeile": it["art"] == "neu", "soll": [], "ist": _rueckbau_form(it["inhalt"], 0)})
    # reihum je Kategorie, damit jede Art von Abweichung vorkommt; Positivliste: nur Formen (wie --schablonen)
    je_kat = [[s for s in schablonen if s["kategorie"] == k] for k in RUECKBAU_KATEGORIEN]
    schablonen = [lst[i] for i in range(max((len(x) for x in je_kat), default=0)) for lst in je_kat if i < len(lst)]
    form_spec = sb.FEHLER_SCHABLONE.felder["soll"]
    gut: list[dict[str, Any]] = []
    for s in schablonen[:maximal]:
        fehler: list[str] = []
        form_spec.pruefe(s["soll"], "soll", fehler)
        form_spec.pruefe(s["ist"], "ist", fehler)
        if not fehler:
            gut.append(s)
    return {"name": ordner.name, "rueckbau": stats, "zellen": {k: int(kat.get(k, 0)) for k in RUECKBAU_KATEGORIEN},
            "tool_extra_fremde_spec": fremd_extra, "fremde_spec_zeilen": sum(len(f) for f in aw.fremd),
            "zellen_haupt": {k: int(kat_haupt.get(k, 0)) for k in RUECKBAU_KATEGORIEN},
            "aenderungen": len(aend), "aenderungen_nach_status": {s: sum(1 for a in aend if a.status == s)
                                                                 for s in (UEBERNOMMEN, PRUEFEN, NICHT)},
            "bilanz_stimmt": bool(bil.get("stimmt")), "schablonen": gut,
            "schablonen_verworfen": min(len(schablonen), maximal) - len(gut),
            "schablonen_weitere": max(0, len(schablonen) - maximal),
            "original_unveraendert": all(_sha(p) == hash_vorher[p.name] for p in (pdf, nachher)),
            "dauer_sekunden": round(time.monotonic() - t0, 2), **wirksam}


def fasse_rueckbau_zusammen(paare: list[dict[str, Any]], korpus: str) -> dict[str, Any]:
    gemessen = [p for p in paare if "abgestuerzt" not in p]
    gesamt = {"paare": len(paare), "paare_gemessen": len(gemessen),
              "zellen": {k: sum(p["zellen"][k] for p in gemessen) for k in RUECKBAU_KATEGORIEN},
              "zellen_haupt": {k: sum(p["zellen_haupt"][k] for p in gemessen) for k in RUECKBAU_KATEGORIEN},
              "zellen_zurueckgebaut": sum(p["rueckbau"]["zellen_zurueckgebaut"] for p in gemessen),
              "zeilen_entfernt": sum(p["rueckbau"]["zeilen_entfernt"] for p in gemessen),
              "aenderungen": sum(p["aenderungen"] for p in gemessen),
              "aenderungen_nach_status": {s: sum(p["aenderungen_nach_status"][s] for p in gemessen)
                                          for s in (UEBERNOMMEN, PRUEFEN, NICHT)},
              "original_unveraendert": sum(1 for p in gemessen if p["original_unveraendert"]),
              "tool_extra_fremde_spec": sum(p.get("tool_extra_fremde_spec", 0) for p in gemessen)}
    alle = sum(gesamt["zellen"].values())
    gesamt["anteil_gleich"] = _anteil(gesamt["zellen"]["gleich"] + gesamt["zellen"]["gleich_bis_leerraum"], alle)
    return {"format": BERICHT_FORMAT, "art": "rueckbau", "korpus": korpus, "schema_version": SCHEMA_VERSION,
            "legende": sorted({p.get("legende", "") for p in gemessen}) or [], "nachher_format": "markiert",
            "kategorien": list(RUECKBAU_KATEGORIEN), "gesamt": gesamt, "paare": paare}


def _md_rueckbau(b: dict[str, Any]) -> str:
    g = b["gesamt"]
    z = [f"# Rückbau-Messung {b['korpus']}", "",
         f"Paare: {g['paare']} (gemessen {g['paare_gemessen']}) · Legende: {', '.join(b['legende']) or 'unbekannt'} · "
         f"nachher-Format: markiert", "",
         "Vorher wurde aus nachher.xlsx zurückgebaut (Markierung der letzten Runde: alt → normaler Text, neu → weg). Das Tool lief auf "
         "PDF + zurückgebautem vorher; verglichen wird Zelle für Zelle mit nachher. Dieser Bericht enthält keine Texte, nur Zahlen, "
         "Positionen und Formen.", "",
         f"Rückbau: {g['zellen_zurueckgebaut']} Zellen zurückgebaut, {g['zeilen_entfernt']} neue Zeilen entfernt · Änderungen des Tools: "
         f"{g['aenderungen']} ({', '.join(f'{k} {n}' for k, n in g['aenderungen_nach_status'].items())})", "",
         "| Kategorie | alle Zellen | Haupt-Spalten |", "|---|---|---|"]
    for k in RUECKBAU_KATEGORIEN:
        z.append(f"| {k} | {g['zellen'][k]} | {g['zellen_haupt'][k]} |")
    z += ["", f"Anteil gleich (inkl. nur Leerraum verschieden): {_prozent(g['anteil_gleich'])}", "", "## Paare", "",
          "| Paar | " + " | ".join(RUECKBAU_KATEGORIEN) + " |", "|---" * (len(RUECKBAU_KATEGORIEN) + 1) + "|"]
    for p in b["paare"]:
        if "abgestuerzt" in p:
            z.append(f"| {p['name']} | nicht messbar ({p['abgestuerzt']}) |" + " |" * (len(RUECKBAU_KATEGORIEN) - 1))
            continue
        z.append(f"| {p['name']} | " + " | ".join(str(p["zellen"][k]) for k in RUECKBAU_KATEGORIEN) + " |")
    laeufe = lambda ls: " ‖ ".join(f"{x['f']}:{x['form']}" for x in ls) or "(leer)"          # noqa: E731
    for p in b["paare"]:
        if not p.get("schablonen"):
            continue
        z += ["", f"### Schablonen {p['name']}", "", "| Kategorie | Blatt | Zeile | Spalte | Soll | Ist |", "|---|---|---|---|---|---|"]
        for s in p["schablonen"]:
            z.append(f"| {s['kategorie']} | {s['blatt']} | {s['zeile']}{' (neu)' if s['neue_zeile'] else ''} | {s['spalte']} | "
                     f"`{laeufe(s['soll'])}` | `{laeufe(s['ist'])}` |")
        if p.get("schablonen_weitere"):
            z.append(f"\nWeitere Abweichungen ohne Schablone: {p['schablonen_weitere']}")
    return "\n".join(z) + "\n"


def _main_nachher(eltern: Path, kandidaten: list[Path], korpus: str, einstellungen: dict[str, str] | None) -> int:
    paare: list[dict[str, Any]] = []
    for d in kandidaten:
        try:
            p = messe_nachher(d, einstellungen)
        except Exception as exc:                  # ein kaputtes Paar bricht die Messung nicht ab
            p = {"name": d.name, "abgestuerzt": type(exc).__name__}
        paare.append(p)
        print(f"  Paar {d.name}: {'nicht messbar' if 'abgestuerzt' in p else 'gemessen'}", flush=True)
    b = fasse_rueckbau_zusammen(paare, korpus)
    basis = config.pfad("messung_dir")
    basis.mkdir(parents=True, exist_ok=True)
    nummern = [int(m.group(1)) for x in basis.iterdir() if (m := re.match(r"^(\d{3,})_", x.name))]
    nr = max(nummern, default=0) + 1
    while True:                                   # append-only (Regel 29)
        ziel = basis / f"{nr:03d}_{korpus}_rueckbau"
        try:
            ziel.mkdir(exist_ok=False)
            break
        except FileExistsError:
            nr += 1
    with open(ziel / "bericht.json", "x", encoding="utf-8") as f:
        json.dump(b, f, ensure_ascii=False, indent=1)
    with open(ziel / "bericht.md", "x", encoding="utf-8") as f:
        f.write(_md_rueckbau(b))
    g = b["gesamt"]
    print(f"Rückbau-Messung {korpus}: {g['paare']} Paare (gemessen {g['paare_gemessen']})")
    for k in RUECKBAU_KATEGORIEN:
        print(f"  {k:<28} {g['zellen'][k]:>6} Zellen  (Haupt-Spalten {g['zellen_haupt'][k]})")
    print(f"  Anteil gleich: {_prozent(g['anteil_gleich'])}  |  davon tool_extra in fremder Spec: {g['tool_extra_fremde_spec']}")
    print(f"Bericht: {ziel / 'bericht.md'}")
    return 0


# ======================================================================================
# 6. Einstieg
# ======================================================================================
def _lade_korpus_info(eltern: Path) -> str:
    p = eltern / "KORPUS.json"
    try:
        return str(json.loads(p.read_text(encoding="utf-8")).get("generator_version", "")) if p.is_file() else ""
    except (OSError, ValueError, AttributeError):
        return ""


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="cde messen", description="Misst das Tool an Paaren (PDF + Excel vorher + Excel nachher).")
    ap.add_argument("--korpus", choices=list(KORPORA), default="entwicklung")
    ap.add_argument("--paar", help="nur dieses Paar (Ordnername)")
    ap.add_argument("--ordner", help="anderer Eltern-Ordner mit Paaren (statt paare_dir/<korpus>)")
    ap.add_argument("--mit-vorschlaegen", action="store_true", help="auch die Qualität der Vorschläge (nach Freigabe) messen")
    ap.add_argument("--schablonen", action="store_true",
                    help="Bericht um den Abschnitt Fehler-Schablonen ergaenzen: Formen (kein Text) von Soll und Ist je stillem Fehler/falschem Vorschlag")
    ap.add_argument("--vergleich", action="store_true", help="Deltas gegen den letzten Bericht desselben Korpus ausgeben")
    ap.add_argument("--gegen", metavar="NNN", help="mit --vergleich: gegen den Bericht mit dieser Nummer vergleichen")
    ap.add_argument("--legende", choices=list(LEGENDEN),
                    help="Legende fuer diesen Lauf (statt KORPUS.json; ohne beides: Korpus-Vorgabe bzw. config/farben.yaml)")
    ap.add_argument("--nachher-format", choices=list(NACHHER_FORMATE),
                    help="markiert = nachher.xlsx traegt die Markierung (Lauf mit markierung.aenderungen: farbe, Format-Treue); "
                         "sauber = ohne (Lauf mit keine)")
    ap.add_argument("--nachher", action="store_true",
                    help="Rueckbau-Messung fuer Paare mit nur aenderungen.pdf + nachher.xlsx: vorher aus der Markierung zurueckbauen, "
                         "Tool laufen lassen, Zelle fuer Zelle vergleichen (nur Zahlen und Schablonen)")
    try:
        a = ap.parse_args(argv)
    except SystemExit as e:                       # argparse beendet selbst; wir geben den Code zurueck
        return 2 if e.code else 0
    eltern = Path(a.ordner) if a.ordner else config.pfad("paare_dir") / a.korpus
    if not eltern.is_dir():
        print("Der Ordner mit den Paaren wurde nicht gefunden. Bitte Paare dort ablegen (siehe docs/paar_format.md).")
        return 2
    if a.paar:
        kandidaten = [eltern / a.paar] if (eltern / a.paar).is_dir() and Path(a.paar).name == a.paar else []
    else:
        kandidaten = sorted(d for d in eltern.iterdir() if d.is_dir() and not d.name.startswith("."))
    # Legende und nachher-Format: Aufruf > KORPUS.json > Vorgabe des Korpus (nur ohne --ordner) > Konfiguration
    einst = korpus_einstellungen(eltern, None if a.ordner else a.korpus)
    if a.legende:
        einst["legende"] = a.legende
    if a.nachher_format:
        einst["nachher_format"] = a.nachher_format
    if a.nachher:
        kandidaten = [d for d in kandidaten if (d / "aenderungen.pdf").is_file() and (d / "nachher.xlsx").is_file()]
        if not kandidaten:
            print("Es wurden keine Paare gefunden. Fuer --nachher braucht jedes Paar aenderungen.pdf und nachher.xlsx.")
            return 2
        return _main_nachher(eltern, kandidaten, a.korpus, einst)
    if not kandidaten:
        print("Es wurden keine Paare gefunden. Jedes Paar ist ein Ordner mit aenderungen.pdf, vorher.xlsx und nachher.xlsx.")
        return 2
    if a.korpus in UNGESEHEN:
        print(f"Hinweis (Regel 11): Der Korpus {a.korpus} wird nur zur Endmessung benutzt, nie zum Einstellen.")
    elif a.korpus == "ex_test":
        print("Hinweis: ex_test = früherer Testkorpus (Seeds 1001–1030), offen als Entwicklung erklärt.")
    alt_ordner = None
    if a.vergleich or a.gegen:
        alt_ordner = finde_bericht(a.korpus, a.gegen)
        if alt_ordner is None:
            print("Kein früherer Bericht zum Vergleich gefunden" + (f" (Nummer {a.gegen})." if a.gegen else " für diesen Korpus."))
            if a.gegen:
                return 2
    paare: list[dict[str, Any]] = []
    for d in kandidaten:
        try:
            p = messe_paar(d, a.mit_vorschlaegen, schablonen=a.schablonen, einstellungen=einst)
        except Exception as exc:                  # ein kaputtes Paar bricht die Messung nicht ab
            p = _fehlgeschlagen(d.name, exc)
        p["name"] = d.name
        paare.append(p)
        stand = ("teilweise gemessen" if p.get("teilweise_gemessen") else "nicht messbar") if "abgestuerzt" in p else "gemessen"
        print(f"  Paar {d.name}: {stand}", flush=True)
    bericht = fasse_zusammen(paare, a.korpus, _lade_korpus_info(eltern), a.mit_vorschlaegen)
    if alt_ordner is not None:
        try:
            alt = json.loads((alt_ordner / "bericht.json").read_text(encoding="utf-8"))
            bericht["vergleich"] = vergleiche(alt, bericht, alt_ordner.name)
        except (OSError, ValueError):
            print("Der frühere Bericht konnte nicht gelesen werden; kein Vergleich.")
    ziel = schreibe_bericht(bericht, a.korpus)
    _konsole(bericht, ziel, a.mit_vorschlaegen)
    return 1 if bericht["gesamt"]["kriterien_nicht_ok"] else 0       # 0 = kein Kriterium "NICHT ok", 1 = Ziel verfehlt, 2 = kaputte Eingabe


if __name__ == "__main__":
    sys.exit(main())
