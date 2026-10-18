"""Traegt die geplanten Aenderungen in eine KOPIE der Excel ein.

Grundsaetze
  * Das Original wird nie veraendert; das Ergebnis ist eine neue Datei.
  * Nur Inhalt: Texte in Zellen aendern, Zeilen einfuegen, Inhalt leeren. Keine Formatierung ausser
    dem Kopieren des Zeilenformats fuer neue Zeilen.
  * Alle Operationen beziehen sich auf den Originalzustand. Zuerst die Zelltexte (von hinten nach
    vorn), danach Zeilen einfuegen/loeschen (von unten nach oben) -- so bleiben Koordinaten gueltig.
  * Stimmt der erwartete Text an einer Stelle nicht mehr, wird die Aenderung nicht eingetragen
    (Konflikt) und mit Grund gemeldet.
  * Nach dem Speichern wird die Datei neu gelesen. Jede Aenderung bekommt eine `wirkung`: "ok" nur, wenn die
    Zielzellen/-zeilen nachweislich so aussehen wie geplant UND die Aenderung sie gegenueber dem Original wirklich
    veraendert hat. Sonst wird sie auf "pruefen" bzw. "nicht_uebernommen" zurueckgestuft (nie "uebernommen" ohne Wirkung).
  * Eine neue Section und ihre Folgeabsaetze (stehen in derselben Excel-Zeile) gelten als Gruppe: Status und Freigabe
    gelten fuer alle gemeinsam.
  * Schreiber (ADR 0007, `anwenden.schreiber`): "patch" (Standard) schreibt verlustfrei mit src/xlsx_patch.py -- openpyxl
    ist dann nur noch Lese- und Planungsmodell; "openpyxl" speichert die Mappe wie frueher (nur fuer alte Tests und
    Notfaelle: Bilder, Formen, Kennzeichnung u. a. koennen dabei verloren gehen).
  * Markierung (ADR 0005, `markierung.aenderungen: farbe`): alter Text bleibt blau durchgestrichen stehen, neuer steht rot
    dahinter (markierung.markiere_ops). Jeder markierte Absatz wird VOR dem Schreiben gegen das PDF geprueft
    (gegenprobe.pruefe_absatz_markiert); was nicht besteht, kommt nicht in die Datei ("pruefen", Wirkung "abweichung").
"""
from __future__ import annotations

import re
import zipfile
from collections import Counter
from copy import copy
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from . import config
from . import markierung as mk
from .fehler import DateiVorhanden
from .excel_model import ExcelModell, zahl_als_text
from .gegenprobe import gleich, pruefe_absatz, pruefe_absatz_markiert
from .models import (MITTEL, NICHT, NIEDRIG, PRUEFEN, UEBERNOMMEN, WIRKUNG_ABWEICHUNG, WIRKUNG_KEINE, WIRKUNG_OK,
                     Absatz, Aenderung, Operation)
from .planner import neutext_probleme, zahl_aus_text
from .textutil import norm
from .xlsx_patch import NeueZeile, PatchFehler, ZellInhalt
from .xlsx_patch import hindernisse as patch_hindernisse
from .xlsx_patch import pruefe_treue
from .xlsx_patch import schreibe as patch_schreiben


@dataclass
class UebernahmeErgebnis:
    ausgabe: str
    warnungen: list[str] = field(default_factory=list)
    zeilen_verschoben: dict[str, str] = field(default_factory=dict)


def _text(zelle) -> str:
    v = zelle.value
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        return str(v)
    if isinstance(v, str):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):      # so wie das Excel-Modell (nie "%g": 2.10 bleibt 2.10)
        return zahl_als_text(v, zelle.number_format)
    return str(v)


def _ist_zahlzelle(zelle) -> bool:
    return zelle.data_type == "n" and isinstance(zelle.value, (int, float)) and not isinstance(zelle.value, bool)


def _ws(text: str) -> str:
    """Leerraum vereinheitlichen (zum Vergleichen)."""
    return " ".join(text.split())


def _hinweise(ae: Aenderung, *texte: str) -> None:
    for t in texte:
        if t and t not in ae.hinweise:
            ae.hinweise.append(t)


def _code(ae: Aenderung, code: str) -> None:
    """Maschinenlesbaren Pruefcode anhaengen (Schema 3), ohne Doppelte."""
    if code and code not in ae.pruefcodes:
        ae.pruefcodes.append(code)


def _konflikt(ae: Aenderung, grund: str, *hinweise: str) -> None:
    ae.status, ae.konfidenz, ae.ops = NICHT, NIEDRIG, []
    ae.grund = (ae.grund + "; " if ae.grund else "") + grund
    _hinweise(ae, *hinweise)


def _vorschlag(ae: Aenderung, grund: str, *hinweise: str) -> None:
    """Zurueck auf 'pruefen': ein Mensch muss entscheiden (die Operationen bleiben als Vorschlag stehen)."""
    ae.status, ae.konfidenz = PRUEFEN, MITTEL
    ae.grund = (ae.grund + "; " if ae.grund else "") + grund
    _hinweise(ae, *hinweise)


def _setze_wert(zelle, text: str) -> None:
    """Text in eine Zelle schreiben. Beginnt er mit '=', bleibt er Text (keine Formel)."""
    zelle.value = text if text != "" else None
    if text.startswith("="):
        zelle.data_type = "s"
        zelle.quotePrefix = True


def _kommentar(zelle, text: str, autor: str) -> None:
    """Kommentar anhaengen; ein vorhandener Kommentar wird nie ueberschrieben."""
    text = ILLEGAL_CHARACTERS_RE.sub("\ufffd", text)          # Steuerzeichen aus dem PDF duerfen das Speichern nicht abbrechen
    alt = zelle.comment
    if alt is not None and (alt.text or "").strip():
        text = f"{alt.text}\n---\n{text}"
        autor = alt.author or autor
    zelle.comment = Comment(text, autor)


def _wortgrenze(text: str, pos: int) -> bool:
    """Liegt `pos` zwischen zwei Woertern (oder am Rand)? Einfuegen mitten in einem Wort ist falsch."""
    return pos <= 0 or pos >= len(text) or text[pos - 1].isspace() or text[pos].isspace()


def _folge_gruppen(aenderungen: list[Aenderung]) -> dict[str, list[Aenderung]]:
    """Neue Section (Kopf-ID) -> ihre Folgeabsaetze. Die Folgeabsaetze stehen in der Zeile der Section (keine eigene Zeile)."""
    koepfe: dict[tuple, Aenderung] = {}
    for a in aenderungen:
        if a.typ == "section_neu":
            for o in a.ops:
                if o.art == "zeile_einfuegen" and not o.abgedeckt and o.zellen:
                    koepfe[(o.blatt, o.zeile, o.vor, o.seq)] = a
    gruppen: dict[str, list[Aenderung]] = {}
    for a in aenderungen:
        if a.typ != "absatz_neu":
            continue
        marker = next((o for o in a.ops if o.art == "zeile_einfuegen" and o.abgedeckt and not o.zellen), None)
        kopf = None
        if marker is not None:
            kopf = koepfe.get((marker.blatt, marker.zeile, marker.vor, marker.seq))
        elif not a.ops and a.ziel:        # aelterer Plan ohne Markierung: ueber die gemeinsame Zielzeile
            gleiche = {k.id: k for k in koepfe.values() if k.ziel == a.ziel}
            if len(gleiche) == 1:
                kopf = next(iter(gleiche.values()))
        if kopf is not None:
            gruppen.setdefault(kopf.id, []).append(a)
    return gruppen


def _gekoppelt(aenderungen: list[Aenderung], gruppen: dict[str, list[Aenderung]]) -> dict[str, set[str]]:
    """Je Aenderung alle Aenderungen, die nur gemeinsam gelten (gleiche `gruppe`, Schema 3, oder neue Section + Folgeabsaetze)."""
    nachbarn: dict[str, set[str]] = {a.id: {a.id} for a in aenderungen}
    je_gruppe: dict[str, list[str]] = {}
    for a in aenderungen:
        if a.gruppe:
            je_gruppe.setdefault(a.gruppe, []).append(a.id)
    kanten = list(je_gruppe.values()) + [[k] + [f.id for f in fl] for k, fl in gruppen.items()]
    for ids in kanten:
        for i in ids:
            if i in nachbarn:
                nachbarn[i].update(x for x in ids if x in nachbarn)
    out: dict[str, set[str]] = {}
    for start in nachbarn:                       # Huelle (Gruppen koennen sich ueber gemeinsame Mitglieder verketten)
        gesehen, offen = {start}, [start]
        while offen:
            for n in nachbarn[offen.pop()]:
                if n not in gesehen:
                    gesehen.add(n)
                    offen.append(n)
        out[start] = gesehen
    return out


_FESTE_STELLEN = re.compile(r"^0\.(0+)$")      # Zahlenformat "0.0", "0.00", ... (wie im Excel-Modell)


def _section_zahl(text: str, zahlenformat: str) -> int | float | None:
    """Neue Section-Nummer als Zahl, wenn sie mit diesem Zahlenformat buchstabengleich angezeigt wird (RT-212), sonst None.
    Nur einfache Formate ("General", "0", "0.0", "0.00", ...), deren Anzeige sicher bekannt ist."""
    fmt = zahlenformat or "General"
    if fmt != "General" and fmt != "0" and not _FESTE_STELLEN.match(fmt):
        return None
    wert = zahl_aus_text(text, fmt)
    if wert is None:
        return None
    if fmt == "0" and not float(wert).is_integer():           # "0" zeigt keine Nachkommastellen: 2.3 sähe aus wie "2"
        return None
    return wert


def _nur_loeschen(o: Operation) -> bool:
    return o.art == "ersetzen" and o.text == ""


def _schneiden(a: Operation, b: Operation, text: str | None = None) -> bool:
    """Ueberschneiden sich zwei Operationen? Zwei reine Loeschungen, die nur ein gemeinsames Leerzeichen teilen
    (zwei aufeinanderfolgende gestrichene Absaetze nehmen beide das Leerzeichen dazwischen mit), gelten als vertraeglich:
    `_wende_text_an` loescht den gemeinsamen Teil nur einmal."""
    if text is not None and _nur_loeschen(a) and _nur_loeschen(b):
        von, bis = max(a.start, b.start), min(a.ende, b.ende)
        if von < bis and not text[von:bis].strip():
            return False
    ins_a, ins_b = a.start == a.ende, b.start == b.ende
    if ins_a and ins_b:
        return False
    if ins_b:
        return a.start < b.start < a.ende
    if ins_a:
        return b.start < a.start < b.ende
    return a.start < b.ende and b.start < a.ende


def _ueberlappen(ops: list[Operation], text: str | None = None) -> set[str]:
    """IDs der Aenderungen, deren Operationen sich in einer Zelle ueberschneiden (die spaetere verliert).
    Mit `text` (Zelltext) zaehlt ein gemeinsames Leerzeichen zweier Loeschungen nicht als Ueberschneidung."""
    verlierer: set[str] = set()
    s = sorted(ops, key=lambda o: (o.start, o.ende, o.seq))
    for i, a in enumerate(s):
        for b in s[:i]:
            if b.aenderung == a.aenderung or a.aenderung in verlierer or b.aenderung in verlierer:
                continue
            if _schneiden(a, b, text):
                verlierer.add(a.aenderung)
    return verlierer


def _roh_anwenden(orig: str, ops: list[Operation]) -> tuple[str, list[Operation]]:
    """Operationen ohne Aufraeumen anwenden. Rueckgabe: Text und die WIRKLICH angewendeten Operationen (eine Loeschung,
    die ein gemeinsames Leerzeichen mit einer spaeteren teilt, endet frueher) -- fuer die Gegenprobe A3."""
    text = orig
    wirklich: list[Operation] = []
    grenze = len(orig)        # kleinster Start der schon angewendeten Operationen: dahinter hat sich der Text geaendert
    # Von hinten nach vorn. Bei gleichem Start wird erst ersetzt und dann eingefuegt: der eingefuegte Text steht vor dem
    # ersetzten (Reihenfolge unabhaengig von `seq`, sonst wird der Text zerstueckelt).
    for o in sorted(ops, key=lambda o: (o.start, 1 if o.art == "ersetzen" else 0, o.seq), reverse=True):
        if o.art == "ersetzen":
            ende = o.ende
            if _nur_loeschen(o) and ende > grenze:       # gemeinsames Leerzeichen mit einer spaeteren Loeschung: nur einmal weg
                ende = max(o.start, grenze)
            text = text[:o.start] + o.text + text[ende:]
            wirklich.append(o if ende == o.ende else replace(o, ende=ende))
        else:   # einfuegen
            text = text[:o.start] + o.text + text[o.start:]
            wirklich.append(o)
        grenze = min(grenze, o.start)
    return text, wirklich


def _wende_text_an(orig: str, ops: list[Operation]) -> str:
    text = _roh_anwenden(orig, ops)[0]
    if any(o.absatz_weg for o in ops) and "\n\n" not in orig and not orig.startswith("\n") and not orig.endswith("\n"):
        text = "\n".join(z for z in text.split("\n") if z.strip())
    if not orig.endswith("\n"):
        text = text.rstrip("\n")
    if not orig.startswith("\n"):
        text = text.lstrip("\n")
    if not orig.endswith(" "):
        text = text.rstrip(" ")
    return text


def _a3_zelle(orig: str, gespeichert: str, info: dict | None, blatt: str, zeile: int,
              spalte: int) -> tuple[str, str, list[Operation]]:
    """Daten einer Zelle fuer die Gegenprobe A3: (Originaltext, gespeicherter Text, wirklich angewendete Operationen).

    Steht genau das in der Zelle, was der Applier schreiben wollte, wird der Text VOR dem Aufraeumen (leere Zeilen,
    Rand-Umbrueche) verglichen: Das Aufraeumen entfernt nur Leerraum, verschiebt aber die Zeichenpositionen."""
    if info is None:
        return orig, gespeichert, []
    ops = [o for o in info["ops"] if not o.abgedeckt]
    leeren = next((o for o in ops if o.art == "zelle_leeren"), None)
    if leeren is not None:                     # geleerte Zelle: wie EIN Ersetzen des ganzen Texts
        return orig, gespeichert, [Operation("ersetzen", blatt, zeile, spalte, 0, len(orig), info["neu"],
                                             aenderung=leeren.aenderung)]
    roh, wirklich = _roh_anwenden(orig, [o for o in ops if o.art in ("ersetzen", "einfuegen")])
    return orig, (roh if gespeichert == info["neu"] else gespeichert), wirklich


def _neuer_absatz_in_zelle(ae: Aenderung) -> bool:
    """Neuer Absatz, der in eine bestehende Zelle eingefuegt wird (nur Einfuegungen mit leerer Spanne). Der Planer legt dafuer
    keinen PDF-Absatz an; geprueft wird er nach dem Speichern mit `_als_eigene_zeile`."""
    return (ae.typ == "absatz_neu" and bool(ae.ops)
            and all(o.art in ("ersetzen", "einfuegen") and o.start == o.ende and o.text.strip() for o in ae.ops))


def _als_eigene_zeile(gespeichert: str, ohne: str, soll: str, cfg: dict) -> bool:
    """Steht der neue Absatz `soll` als EIGENE Zeile in der Zelle (Normalform gleich), und ist der Rest der Zelle ohne diese
    Zeile (und ihren einen Umbruch) bytegleich mit dem Zelltext ohne diese Aenderung (`ohne`)?"""
    zeilen = gespeichert.split("\n")
    for i, z in enumerate(zeilen):
        if gleich(soll, z, cfg=cfg) and "\n".join(zeilen[:i] + zeilen[i + 1:]) == ohne:
            return True
    return False


def _zelltext(orig: str, ops: list[Operation]) -> str:
    """Endtext einer Zelle nach allen Operationen dieser Zelle."""
    if any(o.art == "zelle_leeren" for o in ops):
        ein = [o for o in ops if o.art == "einfuegen"]
        # Einfuegungen in geleerte Zelle: Reihenfolge nach seq, Position 0
        return "".join(o.text for o in sorted(ein, key=lambda o: o.seq)).strip("\n") if ein else ""
    return _wende_text_an(orig, [o for o in ops if not o.abgedeckt])


def _verschiebe_hoehen(ws, ab: int, delta: int) -> None:
    """Zeilenmerkmale (Hoehe, ausgeblendet, Gliederungsebene, Zeilenformat) wandern mit ihrer Zeile.
    Ab Zeile `ab` um `delta` Zeilen; bei delta < 0 verschwinden die geloeschten Zeilen davor."""
    unten = min(ab, ab + delta) if delta < 0 else ab
    merkmale = {}
    for r in [r for r in ws.row_dimensions.keys() if r >= unten]:
        d = ws.row_dimensions[r]
        merkmale[r] = (d.height, d.hidden, d.outlineLevel, d.collapsed, copy(d._style))
    for r in merkmale:
        d = ws.row_dimensions[r]
        d.height, d.hidden, d.outlineLevel, d.collapsed = None, False, 0, False
        d._style = None
    for r, (hoehe, versteckt, ebene, zu, stil) in merkmale.items():
        ziel = r + delta
        if r >= ab and ziel >= 1:
            d = ws.row_dimensions[ziel]
            d.height, d.hidden, d.outlineLevel, d.collapsed = hoehe, versteckt, ebene, zu
            d._style = stil


def _links_nachziehen(ws) -> None:
    """Hyperlinks stehen mit ihrer Zelladresse in der Datei: nach dem Verschieben von Zellen neu eintragen."""
    for zelle in ws._cells.values():
        link = getattr(zelle, "_hyperlink", None)           # verbundene Zellen (MergedCell) haben keinen Link
        if link is not None:
            link.ref = zelle.coordinate


def _verschiebe_verbunden(ws, ab: int, delta: int) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= ab:
            ws.merged_cells.remove(rng)
            rng.shift(row_shift=delta)
            ws.merged_cells.add(rng)


# ---------------------------------------------------------------- Schreiber und Markierung (ADR 0005/0007, Auftrag 003)
def _schreiber(anw: dict) -> str:
    """anwenden.schreiber: patch (Standard, verlustfrei) | openpyxl (frueheres Speichern)."""
    return "openpyxl" if str(anw.get("schreiber", "patch")) == "openpyxl" else "patch"


def _sicht(xl: ExcelModell, blatt: str, zeile: int, spalte: int, zelle) -> str:
    """Welche Sicht hatte das Excel-Modell (und damit der Plan) auf die Zelle? flach = ganzer Text, auch durchgestrichene
    Laeufe (heutiges Modell); neu = Neu-Ansicht (durchgestrichene Laeufe zaehlen nicht, ADR 0006). Erkannt am Text."""
    if zelle is None or not isinstance(zelle.value, CellRichText):
        return mk.SICHT_FLACH
    m = xl.text(blatt, zeile, spalte)
    if m == str(zelle.value):
        return mk.SICHT_FLACH
    if m == mk.sicht_text(mk.stuecke_aus_zelle(zelle), mk.SICHT_NEU):
        return mk.SICHT_NEU
    return mk.SICHT_FLACH


def _sichttext(zelle, sicht: str) -> str:
    if zelle is not None and sicht == mk.SICHT_NEU:
        return mk.sicht_text(mk.stuecke_aus_zelle(zelle), mk.SICHT_NEU)
    return _text(zelle) if zelle is not None else ""


def _zahl_zeichen(orig: str, neu: str | None, aid: str) -> list[tuple[str, str, int, int, str]]:
    """Zahl-Zelle mit Markierung (nur Schriftfarbe, ADR 0005): fuer die Gegenprobe so, als stuende alt (durchgestrichen)
    vor neu -- die Datei selbst hat nur den neuen Wert in roter Schrift (bzw. den alten Wert blau durchgestrichen)."""
    out = [(c, mk.ALT, i, i, aid) for i, c in enumerate(orig)]
    return out + [(c, mk.NEU, -1, len(orig), aid) for c in (neu or "")]


def _merkmal(schrift: Any) -> tuple[bool, str]:
    """(durchgestrichen, Farbe) einer openpyxl-Schrift -- das, woran man alt/neu in der Datei erkennt."""
    if schrift is None:
        return False, ""
    c = getattr(schrift, "color", None)
    farbe = ""
    if c is not None:
        wert = getattr(c, "value", "")
        farbe = f"{getattr(c, 'type', '')}:{str(wert).upper()}:{getattr(c, 'tint', 0) or 0}"
    return bool(getattr(schrift, "strike", False)), farbe


def _ist_merkmale(zelle) -> list[tuple[str, tuple[bool, str]]]:
    """Zeichen einer gespeicherten Zelle mit (durchgestrichen, Farbe)."""
    v = zelle.value
    if isinstance(v, CellRichText):
        out: list[tuple[str, tuple[bool, str]]] = []
        for b in v:
            m = _merkmal(b.font if isinstance(b, TextBlock) else zelle.font)
            out += [(c, m) for c in (b.text if isinstance(b, TextBlock) else str(b))]
        return out
    return [(c, _merkmal(zelle.font)) for c in _text(zelle)]


def _soll_merkmale(stuecke: list[mk.Stueck], zellschrift: Any, cfg: dict) -> list[tuple[str, tuple[bool, str]]]:
    k = mk.konfig(cfg)
    out: list[tuple[str, tuple[bool, str]]] = []
    for s in stuecke:
        if s.behalten or s.format == mk.NORMAL:
            m = _merkmal(s.basis if s.basis is not None else zellschrift)
        elif s.format == mk.ALT:
            m = (True, f"rgb:{str(k.get('alt_farbe', 'FF00B0F0')).upper()}:0")
        else:
            m = (False, f"rgb:{str(k.get('neu_farbe', 'FFFF0000')).upper()}:0")
        out += [(c, m) for c in s.text]
    return out


def _formel_text(zelle) -> str | None:
    v = zelle.value if zelle is not None else None
    return v if isinstance(v, str) and v.startswith("=") and zelle.data_type == "f" else None


def _konstanten(ws, vorlage: int, belegt: set[int]) -> tuple[dict[int, ZellInhalt], list[int]]:
    """Konstanten einer neuen Zeile (neue_zeile.konstanten_kopieren): Spalten, deren Wert in der Vorlagezeile und in
    mindestens zwei ihrer Nachbarzeilen (bis zwei darueber/darunter) gleich ist, werden kopiert; Formeln, die in den
    Nachbarzeilen gleich gebaut sind (gleich nach dem Verschieben), werden in die neue Zeile uebertragen. Alles andere mit
    Inhalt in der Vorlagezeile bleibt leer (zweite Rueckgabe)."""
    kopie: dict[int, ZellInhalt] = {}
    leer: list[int] = []
    nachbarn = [r for r in (vorlage - 2, vorlage - 1, vorlage + 1, vorlage + 2) if r >= 1]
    for c in range(1, ws.max_column + 1):
        if c in belegt:
            continue
        z = ws._cells.get((vorlage, c))
        if z is None or z.value is None or (isinstance(z.value, str) and not z.value.strip()):
            continue
        f = _formel_text(z)
        gleich_n = 0
        for r in nachbarn:
            n = ws._cells.get((r, c))
            if n is None or n.value is None:
                continue
            if f is not None:
                nf = _formel_text(n)
                try:
                    gleich_n += nf is not None and Translator(nf, origin=n.coordinate).translate_formula(z.coordinate) == f
                except Exception:                          # noqa: BLE001 -- nicht uebertragbar: zaehlt nicht als gleich
                    pass
            elif _formel_text(n) is None and n.value == z.value and type(n.value) is type(z.value):
                gleich_n += 1
        if gleich_n >= 2:
            kopie[c] = (ZellInhalt(ws.title, 0, c, formel=f, formel_von=vorlage) if f is not None
                        else ZellInhalt(ws.title, 0, c, kopie_von=vorlage))
        else:
            leer.append(c)
    return kopie, leer


def _gefaehrdet(quelle: Path) -> list[str]:
    """Was openpyxl beim Speichern dieser Datei verliert (ADR 0007), in einfachen Worten."""
    try:
        with zipfile.ZipFile(quelle) as z:
            namen = z.namelist()
    except (zipfile.BadZipFile, OSError):
        return []
    out = []
    if any(n.startswith(("xl/drawings/", "xl/media/")) for n in namen):
        out.append("Bilder und Formen")
    if any(n.startswith(("customXml/", "docMetadata/")) for n in namen):
        out.append("die Vertraulichkeits-Kennzeichnung")
    if any(n.startswith("xl/printerSettings/") for n in namen):
        out.append("Druckereinstellungen")
    return out


def uebernehmen(aenderungen: list[Aenderung], quelle: str | Path, ziel: str | Path,
                freigaben: set[str] | None = None, cfg: dict | None = None,
                xl: ExcelModell | None = None, abwahl: set[str] | None = None,
                absaetze: list[Absatz] | None = None) -> UebernahmeErgebnis:
    """Trägt alle Aenderungen mit Status 'uebernommen' sowie freigegebene 'pruefen' ein.

    Die Aenderung-Objekte werden aktualisiert (Konflikte -> nicht_uebernommen,
    freigegebene Vorschlaege -> uebernommen/bestaetigt, Nachkontrolle und Wirkung; Aenderungen ohne
    nachweisbare Wirkung werden zurueckgestuft).

    `xl`: schon gebautes Modell derselben Quell-Excel (spart das erneute Einlesen); ohne Angabe wird es gebaut.
    `abwahl`: IDs automatischer Aenderungen, die ein Mensch abgewaehlt hat (Schema 3): sie werden nicht eingetragen
    (nicht_uebernommen, "Vom Menschen abgewählt."); gekoppelte Aenderungen (gleiche `gruppe`, Section + Folgeabsaetze)
    fallen mit weg.
    `absaetze`: PDF-Absaetze aus dem Plan fuer die Gegenprobe A3 (Schema 3, ADR 0004); None = Plan ohne Absaetze (Schema <= 2).
    A3: Nach dem Speichern muss jeder Absatz mit angewendeter Aenderung so in der Excel stehen, wie das PDF es verlangt,
    sonst "pruefen" (Wirkung "abweichung", Pruefcode gegenprobe_*). Invariante (nur mit `anwenden.gegenprobe_pflicht`):
    eine automatische Text-Operation, deren Aenderung in keinem Absatz vorkommt, wird nicht angewendet ("ohne_absatz").
    Mit `markierung.aenderungen: farbe` wird A3 schon VOR dem Schreiben am fertigen Zellinhalt geprueft; was nicht besteht,
    kommt nicht in die Datei (keine ungepruefte Automatik), der Status ist derselbe ("pruefen", Wirkung "abweichung").
    Schreiber "patch" (ADR 0007): danach Pruefung der Datei (unberuehrte Teile bytegleich, openpyxl liest sie, jede Zelle
    wie geplant); scheitert das, gibt es keine Ausgabedatei, sondern einen PatchFehler mit klarer Meldung.
    Eine vorhandene Zieldatei wird nie ueberschrieben (Regel 29): dann gibt es einen DateiVorhanden-Fehler (auch FileExistsError)."""
    cfg = cfg or config.lade("zuordnung")
    anw = cfg.get("anwenden") or config.lade("zuordnung")["anwenden"]
    farbe, autor = anw["markierung_farbe"], anw["kommentar_autor"]
    patch = _schreiber(anw) == "patch"
    farbig = patch and mk.aktiv(cfg)          # Markierung nur beim Patch-Schreiber; openpyxl = fruehere Reinigungs-Logik
    trenner = str(mk.konfig(cfg).get("trenner", ""))
    freigaben = freigaben or set()
    quelle, ziel = Path(quelle), Path(ziel)
    if ziel.exists():                                           # Regel 29: Ergebnisse werden nie ueberschrieben
        raise DateiVorhanden("Die Ergebnis-Datei gibt es schon und wird nicht überschrieben. "
                             "Bitte einen neuen Dateinamen oder einen neuen Lauf nehmen.")
    ergebnis = UebernahmeErgebnis(ausgabe=str(ziel))
    if xl is None:
        xl = ExcelModell(quelle)                                # nur zum Lesen (Layout, Sections)
    wb = load_workbook(str(quelle), rich_text=True, keep_vba=(quelle.suffix.lower() == ".xlsm"))
    sichten: dict[tuple[str, int, int], str] = {}

    def sicht(blatt: str, zeile: int, spalte: int) -> str:
        key = (blatt, zeile, spalte)
        if key not in sichten:
            sichten[key] = _sicht(xl, blatt, zeile, spalte, wb[blatt]._cells.get((zeile, spalte)))
        return sichten[key]

    def sichttext(blatt: str, zeile: int, spalte: int) -> str:
        return _sichttext(wb[blatt]._cells.get((zeile, spalte)), sicht(blatt, zeile, spalte))

    # Originaltexte der Absatz-Zellen (fuer die Gegenprobe A3 nach dem Speichern); ohne neue Zellen anzulegen
    absatz_orig: dict[tuple[str, int, int], str] = {}
    zwischen: dict[str, str] = {}                     # Leerraum zwischen Alt und Neu einer Aenderung laut PDF
    for ab in absaetze or []:
        zwischen.update(mk.trenner_aus_laeufen(ab.laeufe))
        if ab.blatt in wb.sheetnames:
            for z, sp, _, _ in ab.zellen:
                absatz_orig[(ab.blatt, z, sp)] = sichttext(ab.blatt, z, sp)

    # --- Welche Aenderungen werden angewendet? (neue Section + Folgeabsaetze gelten als Gruppe) ---------------
    gruppen = _folge_gruppen(aenderungen)
    folge_zu_kopf = {f.id: k for k, fl in gruppen.items() for f in fl}
    alle_nach_id = {a.id: a for a in aenderungen}
    kopplung = _gekoppelt(aenderungen, gruppen)

    # --- Abwahl: automatische Aenderungen, die ein Mensch nicht will (gekoppelte fallen mit weg) -----------------
    abgewaehlt: set[str] = set()
    for aid in abwahl or ():
        a = alle_nach_id.get(aid)
        if a is not None and a.status == UEBERNOMMEN:
            abgewaehlt |= kopplung.get(aid, {aid})
    for aid in sorted(abgewaehlt):
        a = alle_nach_id[aid]
        if a.status == UEBERNOMMEN:
            a.status = NICHT
            a.grund = (a.grund + "; " if a.grund else "") + "Vom Menschen abgewählt."
            _code(a, "abgewaehlt")
            if aid not in (abwahl or ()):
                _hinweise(a, "Gehört zu einer Gruppe, von der ein Mensch eine Änderung abgewählt hat; "
                             "die Gruppe gilt nur gemeinsam.")
        elif a.status == PRUEFEN and aid in freigaben:
            _hinweise(a, "Nicht eingetragen: gehört zu einer Gruppe, von der ein Mensch eine Änderung abgewählt hat.")
    freigaben = freigaben - abgewaehlt

    def freigegeben(a: Aenderung) -> bool:
        return a.status == PRUEFEN and a.id in freigaben

    aktiv_ids: set[str] = set()
    for kopf_id, folgen in gruppen.items():
        kopf = alle_nach_id[kopf_id]
        mitglieder = [kopf] + folgen
        mensch = any(freigegeben(m) for m in mitglieder)          # nur ein Mensch (Freigabe-ID) darf die Gruppe mitnehmen
        if (not mensch and kopf.status in (UEBERNOMMEN, PRUEFEN)
                and {m.status for m in mitglieder} >= {UEBERNOMMEN, PRUEFEN}):
            # Gemischte Gruppe ohne Freigabe: der "uebernommene" Teil darf den "pruefen"-Teil nicht ohne Mensch mitziehen
            # (und umgekehrt). Der Applier sichert das selbst ab, auch wenn ein Plan von Hand gebaut wurde.
            for m in mitglieder:
                if m.status == UEBERNOMMEN:
                    _vorschlag(m, "Gehört zu einer neuen Section mit Absätzen, von denen mindestens einer noch geprüft werden muss.",
                               "Section und Absätze stehen in einer Excel-Zeile und werden nur gemeinsam eingetragen; "
                               "ein Mensch muss mindestens einen davon freigeben.")
        elif kopf.status in (UEBERNOMMEN, PRUEFEN) and (kopf.status == UEBERNOMMEN or mensch):
            for m in mitglieder:
                if m.status == PRUEFEN:
                    m.status, m.bestaetigt = UEBERNOMMEN, True
                    if m.id not in freigaben:
                        _hinweise(m, "Zusammen mit der zugehörigen neuen Section bzw. ihrem Absatz freigegeben, "
                                     "weil beide in einer Excel-Zeile stehen.")
                if m.status == UEBERNOMMEN:
                    aktiv_ids.add(m.id)
        else:
            for m in folgen:
                if m.status == UEBERNOMMEN:         # darf nicht allein "uebernommen" sein
                    if kopf.status == NICHT:
                        _konflikt(m, "Gehört zu einer neuen Section, die nicht eingetragen wird.")
                    else:
                        _vorschlag(m, "Gehört zu einer neuen Section, die noch nicht freigegeben ist.",
                                   "Dieser Absatz wird nur zusammen mit seiner neuen Section eingetragen.")
    for ae in aenderungen:
        if ae.id in folge_zu_kopf or ae.id in gruppen:
            continue
        if ae.status == UEBERNOMMEN:
            aktiv_ids.add(ae.id)
        elif freigegeben(ae):
            ae.status, ae.bestaetigt = UEBERNOMMEN, True
            aktiv_ids.add(ae.id)
    aktiv: list[Aenderung] = [a for a in aenderungen if a.id in aktiv_ids]
    nach_id = {a.id: a for a in aktiv}

    def entferne(ae: Aenderung, art: str, grund: str, *hinweise: str) -> None:
        """Nimmt eine Aenderung (und die Gruppe, zu der sie gehoert) aus der Anwendung."""
        if ae.id not in nach_id:
            return
        (_konflikt if art == NICHT else _vorschlag)(ae, grund, *hinweise)
        nach_id.pop(ae.id)
        kopf_id = folge_zu_kopf.get(ae.id)
        weitere = list(gruppen.get(ae.id, []))
        if kopf_id and kopf_id in nach_id:
            weitere.append(nach_id[kopf_id])
            weitere += [f for f in gruppen[kopf_id] if f.id != ae.id]
        for w in weitere:
            entferne(w, art, "Gehört zu einer neuen Section, deren Eintragung nicht möglich war.")
        if ae.gruppe:                                   # Schema 3: Aenderungen derselben Gruppe gelten nur gemeinsam
            for w in [x for x in list(nach_id.values()) if x.gruppe == ae.gruppe]:
                entferne(w, art, "Gehört zu einer Gruppe von Änderungen, von der eine nicht eingetragen werden kann.")

    # --- Invariante (ADR 0004): automatische Text-Operation ohne gespeicherten PDF-Absatz -> nicht automatisch -----
    # Ausgenommen: Freigaben, neue Zeilen (Zeilen-Nachkontrolle) und neue Absaetze in einer bestehenden Zelle (eigene Pruefung
    # nach dem Speichern: `_als_eigene_zeile`).
    mit_absatz = {i for ab in absaetze or [] for i in ab.aenderungen}
    if absaetze is not None and anw.get("gegenprobe_pflicht", False):
        for ae in list(aktiv):
            if (ae.id in nach_id and not ae.bestaetigt and ae.id not in mit_absatz and not _neuer_absatz_in_zelle(ae)
                    and any(o.art in ("ersetzen", "einfuegen") and not o.abgedeckt for o in ae.ops)):
                entferne(ae, PRUEFEN, "Zu dieser Änderung fehlt der PDF-Absatz für die Gegenprobe; sie wird nicht automatisch eingetragen.",
                         "Das Ergebnis könnte nach dem Eintragen nicht gegen das PDF geprüft werden. Bitte von Hand prüfen.")
                _code(ae, "ohne_absatz")

    # --- Stimmt der erwartete Text noch? Ist der Neutext schreibbar? ------------------------------------------
    for ae in list(aktiv):
        if ae.id not in nach_id:
            continue
        for op in ae.ops:
            if op.blatt not in wb.sheetnames:
                entferne(ae, NICHT, "Tabellenblatt nicht gefunden.")
                break
            if op.art in ("ersetzen", "einfuegen", "zelle_leeren"):
                zelle = wb[op.blatt].cell(op.zeile, op.spalte)
                if zelle.data_type == "f":
                    entferne(ae, NICHT, "In der Zelle steht eine Formel; sie wird nie verändert.")
                    break
                if not farbig and sicht(op.blatt, op.zeile, op.spalte) == mk.SICHT_NEU:
                    entferne(ae, PRUEFEN, "In der Zelle steht durchgestrichener Text früherer Runden; ohne farbige Markierung "
                                          "kann er beim Eintragen nicht sicher erhalten werden.",
                             "Bitte von Hand eintragen oder mit farbiger Markierung (markierung.aenderungen = farbe) arbeiten.")
                    break
                aktuell = sichttext(op.blatt, op.zeile, op.spalte)
                if op.art != "zelle_leeren" and (op.ende > len(aktuell) or
                                                 (op.art == "ersetzen" and aktuell[op.start:op.ende] != op.erwartet)):
                    entferne(ae, NICHT, "Die Excel hat sich seit der Analyse geändert (Text an der Stelle stimmt nicht mehr).")
                    break
                if op.art == "zelle_leeren" and op.erwartet:
                    if aktuell != op.erwartet:
                        entferne(ae, NICHT, "Die Excel hat sich seit der Analyse geändert (Text an der Stelle stimmt nicht mehr).")
                        break
                    ungestrichen = Counter(norm(w) or w for w in aktuell.split()) - Counter(norm(w) or w for w in ae.alt.split())
                    if ungestrichen:       # eine Zelle wird nur geleert, wenn ALLE ihre Woerter gestrichen sind
                        entferne(ae, PRUEFEN, "In der Zelle stehen Wörter, die im PDF nicht gestrichen sind. Sie wird nicht geleert.",
                                 "Die Zelle enthält mehr Text als im PDF gestrichen ist. Bitte von Hand entscheiden.")
                        break
        if ae.id not in nach_id:
            continue
        texte = [o.text for o in ae.ops if o.art in ("ersetzen", "einfuegen")] + [t for o in ae.ops for t in o.zellen.values()]
        if ae.id in folge_zu_kopf:
            texte.append(ae.neu)
        p = neutext_probleme(texte, cfg)
        if p["steuerzeichen"]:
            entferne(ae, NICHT, "Der neue Text enthält Steuerzeichen und kann nicht in die Excel geschrieben werden.")
        elif p["verdaechtig"] and not ae.bestaetigt:
            entferne(ae, PRUEFEN, "Der neue Text könnte aus dem PDF verfälscht gelesen worden sein; er wird nicht automatisch eingetragen.",
                     *p["verdaechtig"])
        elif p["formel"]:
            _hinweise(ae, "Der neue Text beginnt mit „=“. Er wird als Text eingetragen, nicht als Formel.")
    aktiv = [a for a in aktiv if a.id in nach_id]

    # --- Zeile einfuegen: Risiken der Excel (Formeln, Formatierung, Filter, Links ...). Nie automatisch. -------
    # Patch-Schreiber: nur, was xlsx_patch nicht sicher mitziehen kann (geteilte Formel ueber die Stelle, Tabellenobjekt ...);
    # das ist ein festes Hindernis, auch nach Freigabe (die Datei wuerde sonst beschaedigt).
    risiko_cache: dict[tuple[str, int], list[str]] = {}

    def zeilen_risiken(ae: Aenderung) -> list[str]:
        out: list[str] = []
        for op in ae.ops:
            if op.art == "zeile_einfuegen" and not op.abgedeckt:
                pos = op.zeile if op.vor else op.zeile + 1
                if (op.blatt, pos) not in risiko_cache:
                    risiko_cache[(op.blatt, pos)] = (patch_hindernisse(quelle, op.blatt, pos) if patch
                                                     else xl.einfuege_hindernisse(op.blatt, pos, schreiber="openpyxl"))
                out += [r for r in risiko_cache[(op.blatt, pos)] if r not in out]
        return out

    for ae in list(aktiv):
        if ae.id not in nach_id:
            continue
        risiken = zeilen_risiken(ae)
        if not risiken:
            continue
        if patch:
            entferne(ae, PRUEFEN, "Die neue Zeile kann nicht verlustfrei eingefügt werden; dabei würde in der Excel etwas "
                                  "verrutschen oder brechen. Die Änderung wird nicht eingetragen.", *risiken)
        elif ae.bestaetigt:               # ein Mensch hat es freigegeben: eintragen, aber laut warnen
            ergebnis.warnungen.append(f"Änderung {ae.id} ({ae.ziel}): Die Zeile wurde auf Freigabe eingefügt, obwohl "
                                      f"in der Excel etwas verrutschen oder brechen kann: {' '.join(risiken)} Bitte die Ergebnis-Excel prüfen.")
            _hinweise(ae, *risiken)
        else:
            entferne(ae, PRUEFEN, "Beim Einfügen der neuen Zeile könnten Dinge in der Excel verrutschen oder brechen. "
                                  "Die Änderung wird nicht automatisch eingetragen.", *risiken)
    aktiv = [a for a in aktiv if a.id in nach_id]

    # --- Zelltext-Operationen je Zelle ---------------------------------------------------
    je_zelle: dict[tuple[str, int, int], list[Operation]] = {}
    for ae in aktiv:
        for op in ae.ops:
            if op.art in ("ersetzen", "einfuegen", "zelle_leeren"):
                je_zelle.setdefault((op.blatt, op.zeile, op.spalte), []).append(op)

    for key, ops in je_zelle.items():
        text_ops = [o for o in ops if o.art != "zelle_leeren"]
        if any(o.art == "zelle_leeren" for o in ops):
            for o in text_ops:                                   # Ersetzen in geleerter Zelle: bereits abgedeckt
                if o.art == "ersetzen":
                    o.abgedeckt = True
        else:
            for aid in _ueberlappen(text_ops, sichttext(*key)):
                if aid in nach_id:
                    entferne(nach_id[aid], NICHT, "Konflikt: überschneidet sich mit einer anderen Änderung in derselben Zelle.")
    aktiv = [a for a in aktiv if a.id in nach_id]

    # --- Zahl-Zellen: der neue Text muss wieder eine Zahl sein, sonst geht das Zahlenformat verloren -------------
    for (blatt, zeile, spalte), ops in je_zelle.items():
        zelle = wb[blatt].cell(zeile, spalte)
        ops = [o for o in ops if o.aenderung in nach_id and o.art != "zelle_leeren"]
        if not ops or not _ist_zahlzelle(zelle) or any(o.abgedeckt for o in ops):
            continue
        neu_text = _zelltext(_text(zelle), ops)
        if not neu_text.strip() or zahl_aus_text(neu_text, zelle.number_format) is not None:
            continue
        for aid in dict.fromkeys(o.aenderung for o in ops):
            ae = nach_id.get(aid)
            if ae is None:
                continue
            if ae.bestaetigt:
                _hinweise(ae, "Die Zelle war eine Zahl, jetzt steht dort Text. Das Zahlenformat geht verloren.")
                ergebnis.warnungen.append(f"Änderung {ae.id} ({ae.ziel}): Die Zelle war eine Zahl und ist jetzt Text; "
                                          f"das Zahlenformat geht verloren.")
            else:
                entferne(ae, PRUEFEN, "Zahlenformat geht verloren: Die Zelle ist in der Excel eine Zahl, der neue Text ist keine "
                                      "(oder sieht anders aus als das Zahlenformat).",
                         "Die Zelle ist eine Zahl. Wird dort Text eingetragen, geht das Zahlenformat verloren. Bitte von Hand entscheiden.")
    aktiv = [a for a in aktiv if a.id in nach_id]

    # --- Patch-Schreiber: was es dort nicht gibt (Kommentar/Gelb, Zeilen entfernen) -> "pruefen" mit klarer Meldung --------
    modus = cfg["markierung"]["geaenderte_zellen"]
    zeilen_weg = (str(mk.konfig(cfg).get("geloeschte_section", "markieren")) == "zeile_entfernen" if farbig
                  else cfg["geloeschtes"]["inhalt"] == "zeile_entfernen")
    if not patch and mk.aktiv(cfg):
        ergebnis.warnungen.append("Die farbige Markierung (markierung.aenderungen = farbe) gibt es nur beim verlustfreien "
                                  "Schreiben (anwenden.schreiber = patch). Mit openpyxl wurde der alte Text wie früher ersetzt.")
    if patch and modus != "keine":
        for ae in list(aktiv):
            entferne(ae, PRUEFEN, f"Die Markierung „{modus}“ gibt es nur mit dem Schreiber „openpyxl“; die Änderung wird "
                                  "nicht eingetragen.",
                     "Für das verlustfreie Schreiben (ADR 0007) bitte markierung.geaenderte_zellen auf „keine“ stellen.")
        ergebnis.warnungen.append(f"Die Einstellung markierung.geaenderte_zellen = „{modus}“ gibt es beim verlustfreien "
                                  "Schreiben nicht. Es wurde nichts eingetragen; alle Änderungen stehen auf „prüfen“.")
        aktiv = [a for a in aktiv if a.id in nach_id]
    if patch and zeilen_weg:
        for (blatt, zeile, spalte), ops in je_zelle.items():
            ops = [o for o in ops if o.aenderung in nach_id]
            lay = xl.layouts.get(blatt)
            if not lay or not any(o.art == "zelle_leeren" for o in ops):
                continue
            cols = [c for c in (lay.section_spalte, lay.titel_spalte, lay.inhalt_spalte) if c]

            def leer_nachher(c: int) -> bool:
                eigene = [o for o in je_zelle.get((blatt, zeile, c), []) if o.aenderung in nach_id]
                return not (_zelltext(sichttext(blatt, zeile, c), eigene) if eigene else sichttext(blatt, zeile, c)).strip()
            if all(leer_nachher(c) for c in cols):
                for aid in dict.fromkeys(o.aenderung for o in ops):
                    if aid in nach_id:
                        entferne(nach_id[aid], PRUEFEN, "Eine Zeile zu entfernen gibt es nur mit dem Schreiber „openpyxl“; "
                                                        "die Änderung wird nicht eingetragen.",
                                 "Beim verlustfreien Schreiben (ADR 0007) bleibt die Zeile stehen. Bitte von Hand entscheiden "
                                 "oder die gestrichene Zeile markieren lassen (geloeschte_section = markieren).")
        aktiv = [a for a in aktiv if a.id in nach_id]

    # --- Inhalt jeder geaenderten Zelle planen (ohne zu schreiben) ------------------------------------
    def baue_zellen() -> dict[tuple[str, int, int], dict]:
        out: dict[tuple[str, int, int], dict] = {}
        for (blatt, zeile, spalte), ops in je_zelle.items():
            ops = [o for o in ops if o.aenderung in nach_id]
            if not ops:
                continue
            zelle = wb[blatt].cell(zeile, spalte)
            sv = sicht(blatt, zeile, spalte)
            orig = _sichttext(zelle, sv)
            neu = _zelltext(orig, ops)
            zahl = zahl_aus_text(neu, zelle.number_format) if _ist_zahlzelle(zelle) else None
            info: dict[str, Any] = {"orig": orig, "neu": neu, "ops": ops, "zahl": zahl, "sicht": sv,
                                    "leeren": any(o.art == "zelle_leeren" for o in ops),
                                    "aen": [nach_id[i] for i in dict.fromkeys(o.aenderung for o in ops)]}
            if farbig and _ist_zahlzelle(zelle) and (zahl is not None or not neu.strip()):
                # Zahl bleibt Zahl: nur die Schrift wird gefaerbt (neu = rot; geloescht = Wert bleibt, blau durchgestrichen)
                aid = ops[0].aenderung
                if zahl is None:
                    info.update(farbe=mk.ALT, wert=zelle.value, zeichen=_zahl_zeichen(orig, None, aid), neu=orig)
                else:
                    t = zahl_als_text(zahl, zelle.number_format)
                    info.update(farbe=mk.NEU, wert=zahl, zeichen=_zahl_zeichen(orig, t, aid), neu=t)
            elif farbig:
                vorher = [mk.Stueck(mk.NORMAL, orig)] if _ist_zahlzelle(zelle) else mk.stuecke_aus_zelle(zelle)
                m = mk.markiere_ops(vorher, ops, sv, zwischen, cfg)
                info.update(mark=m, zeichen=m.zeichen, neu=m.text, zellschrift=zelle.font)
            out[(blatt, zeile, spalte)] = info
        return out

    # --- Gegenprobe A3 VOR dem Schreiben (nur mit Markierung): was nicht zum PDF passt, kommt nicht in die Datei -------
    def vorpruefung(zell_info: dict) -> tuple[dict[str, str], set[str]]:
        fehler: dict[str, str] = {}
        gut: set[str] = set()
        angewendet = set(nach_id)
        for ab in absaetze or []:
            if not ab.aenderungen & angewendet:
                continue
            daten: dict[tuple[int, int], tuple[str, list]] = {}
            for z, sp, _, _ in ab.zellen:
                info = zell_info.get((ab.blatt, z, sp))
                if info is not None:
                    daten[(z, sp)] = (info["orig"], info["zeichen"])
                elif ab.blatt in wb.sheetnames:
                    t = absatz_orig.get((ab.blatt, z, sp), "")
                    daten[(z, sp)] = (t, [(c, mk.NORMAL, i, i, "") for i, c in enumerate(t)])
            code = pruefe_absatz_markiert(ab, daten, angewendet, cfg, trenner)
            for aid in sorted(ab.aenderungen & angewendet):
                if code is None:
                    gut.add(aid)
                else:
                    fehler.setdefault(aid, code)
        for ae in aktiv:                    # neuer Absatz in einer bestehenden Zelle: eigene Zeile, Rest unveraendert
            if ae.id not in nach_id or ae.id in mit_absatz or not _neuer_absatz_in_zelle(ae):
                continue
            soll = ae.soll_absatz or ae.neu
            ok_ = True
            for o in ae.ops:
                info = zell_info.get((o.blatt, o.zeile, o.spalte))
                if info is None or "zeichen" not in info:
                    ok_ = False
                    break
                gesp = "".join(c for c, kl, _, _, _ in info["zeichen"] if kl != mk.ALT)
                ohne = "".join(c for c, kl, p, _, aid in info["zeichen"] if kl != mk.ALT and not (p < 0 and aid == ae.id))
                if not _als_eigene_zeile(gesp, ohne, soll, cfg):
                    ok_ = False
                    break
            if ok_:
                gut.add(ae.id)
            else:
                fehler.setdefault(ae.id, "gegenprobe_neuer_absatz")
        return fehler, gut

    zurueck: dict[str, str] = {}                   # vor dem Schreiben durchgefallen (nicht in der Datei)
    a3_vorab: set[str] = set()
    zell_info = baue_zellen()
    if farbig:
        # Idempotenz: eine Aenderung, deren Loeschungen schon durchgestrichen dastehen und die nichts Neues markiert oder
        # einfuegt, steht schon so in der Excel (zweiter Lauf auf dem eigenen Ergebnis) -- sie wird nicht noch einmal eingetragen
        neu_n: Counter = Counter()
        alt_n: Counter = Counter()
        for info in zell_info.values():
            if "mark" in info:
                neu_n.update(info["mark"].neu_markiert)
                alt_n.update(info["mark"].schon_alt)
        schon = [a for a in aktiv if a.id in nach_id and alt_n[a.id] and not neu_n[a.id]]
        for ae in schon:
            entferne(ae, NICHT, "Die Änderung steht schon so in der Excel (der gestrichene Text ist dort bereits "
                                "durchgestrichen markiert).")
            _code(ae, "schon_eingetragen")
        if schon:
            aktiv = [a for a in aktiv if a.id in nach_id]
            zell_info = baue_zellen()
    if farbig and absaetze is not None:
        for _ in range(len(nach_id) + 1):
            fehler_vorab, a3_vorab = vorpruefung(zell_info)
            if not fehler_vorab:
                break
            for aid, code in fehler_vorab.items():
                for k in sorted(kopplung.get(aid, {aid})):
                    if k in nach_id and k not in zurueck:
                        zurueck[k] = code if k == aid else "gegenprobe_gruppe"
            for k in zurueck:
                nach_id.pop(k, None)
            zell_info = baue_zellen()
    geleerte: set[tuple[str, int, int]] = {k for k, i in zell_info.items() if i["leeren"]}
    for (blatt, zeile, spalte), info in zell_info.items():          # Schema 4: Ergebnis der Zielzelle fuer die Anzeige
        if "zeichen" in info:
            for ae in info["aen"]:
                if not ae.stuecke:
                    ae.stuecke = mk.klassen_stuecke(info["zeichen"])

    # --- Zellen schreiben (openpyxl) bzw. sammeln (Patch) -----------------------------------------
    zell_inhalte: list[ZellInhalt] = []
    for (blatt, zeile, spalte), info in zell_info.items():
        if patch:
            if "mark" in info:
                zell_inhalte.append(ZellInhalt(blatt, zeile, spalte, stuecke=info["mark"].stuecke,
                                               text_praefix=mk.sicht_text(info["mark"].stuecke, mk.SICHT_NEU).startswith("=")))
            elif "farbe" in info:
                zell_inhalte.append(ZellInhalt(blatt, zeile, spalte, wert=info["wert"], farbe=info["farbe"]))
            elif info["zahl"] is not None:
                zell_inhalte.append(ZellInhalt(blatt, zeile, spalte, wert=info["zahl"]))
            else:
                zell_inhalte.append(ZellInhalt(blatt, zeile, spalte, wert=info["neu"] or None,
                                               text_praefix=info["neu"].startswith("=")))
            continue
        zelle = wb[blatt].cell(zeile, spalte)
        if info["zahl"] is not None:
            zelle.value = info["zahl"]                    # bleibt eine Zahl, das Zahlenformat der Zelle bleibt
        else:
            _setze_wert(zelle, info["neu"])

    # --- Markierung der geaenderten Zellen (optional, nur openpyxl) -------------------------------------
    if modus != "keine" and not patch:
        for (blatt, zeile, spalte), info in zell_info.items():
            zelle = wb[blatt].cell(zeile, spalte)
            if modus == "gelb":
                zelle.fill = PatternFill("solid", fgColor=farbe)
            elif modus == "kommentar":
                _kommentar(zelle, "\n".join(f"PDF S.{a.seite}: {a.alt[:80]} → {a.neu[:80]}" for a in info["aen"]), autor)

    # --- Zeilen einfuegen / loeschen ------------------------------------------------------
    einfuegungen: dict[tuple[str, int], list[Operation]] = {}
    for ae in aktiv:
        if ae.id not in nach_id:
            continue
        for op in ae.ops:
            if op.art == "zeile_einfuegen" and not op.abgedeckt:
                pos = op.zeile if op.vor else op.zeile + 1
                einfuegungen.setdefault((op.blatt, pos), []).append(op)
    loeschen: set[tuple[str, int]] = set()
    if zeilen_weg and not patch:
        for (blatt, zeile, spalte) in geleerte:
            lay = xl.layouts.get(blatt)
            if not lay:
                continue
            cols = [c for c in (lay.section_spalte, lay.titel_spalte, lay.inhalt_spalte) if c]

            if all(wb[blatt].cell(zeile, c).value in (None, "") for c in cols) and not wb[blatt].merged_cells.ranges:
                risiken = xl.einfuege_risiken(blatt, zeile)
                if risiken:                    # die leere Zeile bleibt stehen (nichts geht kaputt), dafuer ein Hinweis
                    for ae in zell_info[(blatt, zeile, spalte)]["aen"]:
                        _hinweise(ae, "Die leere Zeile wurde nicht entfernt, weil dabei Dinge in der Excel verrutschen "
                                      "oder brechen könnten:", *risiken)
                    ergebnis.warnungen.append(f"{blatt}, Zeile {zeile}: Die geleerte Zeile wurde nicht entfernt (sie bleibt leer "
                                              f"stehen), weil sonst in der Excel etwas verrutschen oder brechen könnte.")
                else:
                    loeschen.add((blatt, zeile))
    abbildung: dict[str, list[tuple[str, int]]] = {}
    neue_zeilen_patch: list[NeueZeile] = []
    nz_index: dict[int, int] = {}                     # id(Operation) -> Index in neue_zeilen_patch
    konstanten_an = bool((cfg.get("neue_zeile") or {}).get("konstanten_kopieren", False))
    if patch:
        for (blatt, pos), ops in sorted(einfuegungen.items(), key=lambda e: (e[0][0], e[0][1])):
            ws = wb[blatt]
            lay = xl.layouts.get(blatt)
            sek = lay.section_spalte if lay else None
            for op in sorted(ops, key=lambda o: o.seq):
                stil = op.stil_von or op.zeile
                zellen: dict[int, Any] = {}
                for spalte, text in op.zellen.items():
                    if spalte == sek and text.strip():
                        # RT-212: stehen die Section-Nummern als Zahl in der Excel, wird die neue auch eine Zahl
                        vorlage_r = next((r for r in (stil, op.zeile) if _ist_zahlzelle(ws.cell(r, spalte))), None)
                        if vorlage_r is not None:
                            zahl = _section_zahl(text, ws.cell(vorlage_r, spalte).number_format)
                            if zahl is not None:
                                zellen[spalte] = ZellInhalt(blatt, 0, spalte, wert=zahl, farbe=mk.NEU if farbig else "",
                                                            stil_von=vorlage_r if vorlage_r != stil else 0)
                                continue
                            _hinweise(alle_nach_id[op.aenderung], "Die Section-Nummern stehen in der Excel als Zahl. Die neue "
                                      "Nummer wurde als Text eingetragen, weil sie mit dem Zahlenformat anders aussehen würde.")
                    if farbig and text:
                        zellen[spalte] = ZellInhalt(blatt, 0, spalte, stuecke=[mk.Stueck(mk.NEU, text)],
                                                    text_praefix=text.startswith("="))
                    else:
                        zellen[spalte] = ZellInhalt(blatt, 0, spalte, wert=text or None, text_praefix=text.startswith("="))
                if konstanten_an:
                    kopie, leer = _konstanten(ws, stil, set(op.zellen))
                    zellen.update(kopie)
                    op.leer_spalten = leer
                    if leer:
                        _hinweise(alle_nach_id[op.aenderung],
                                  "In der neuen Zeile bleiben die Spalten " + ", ".join(get_column_letter(c) for c in leer)
                                  + " leer (sie sind in den Nachbarzeilen verschieden). Bitte von Hand ergänzen.")
                    if any(z.formel for z in kopie.values()):
                        op.formeln = {c: z.formel for c, z in kopie.items() if z.formel}
                nz_index[id(op)] = len(neue_zeilen_patch)
                neue_zeilen_patch.append(NeueZeile(blatt, pos - 1, stil, zellen))
    else:
        ereignisse = [("del", b, r) for b, r in loeschen] + [("ins", b, p) for (b, p) in einfuegungen]
        ereignisse.sort(key=lambda e: (e[1], -e[2], 0 if e[0] == "del" else 1))
        for art, blatt, pos in ereignisse:
            ws = wb[blatt]
            if art == "del":
                ws.delete_rows(pos, 1)
                _verschiebe_hoehen(ws, pos + 1, -1)
                _verschiebe_verbunden(ws, pos + 1, -1)
                _links_nachziehen(ws)
                abbildung.setdefault(blatt, []).append(("del", pos))
            else:
                ops = sorted(einfuegungen[(blatt, pos)], key=lambda o: o.seq)
                n = len(ops)
                ws.insert_rows(pos, n)
                _verschiebe_hoehen(ws, pos, n)
                _verschiebe_verbunden(ws, pos, n)
                _links_nachziehen(ws)
                abbildung.setdefault(blatt, []).append(("ins", pos, n))        # type: ignore[arg-type]
                lay = xl.layouts.get(blatt)
                sek = lay.section_spalte if lay else None
                for i, op in enumerate(ops):
                    neu_zeile = pos + i
                    stil = op.stil_von + (n if op.stil_von >= pos else 0)
                    anker = op.zeile + (n if op.zeile >= pos else 0)
                    for c in range(1, ws.max_column + 1):
                        ws.cell(neu_zeile, c)._style = copy(ws.cell(stil, c)._style)
                    ws.row_dimensions[neu_zeile].height = ws.row_dimensions[stil].height
                    for spalte, text in op.zellen.items():
                        zelle = ws.cell(neu_zeile, spalte)
                        if spalte == sek and text.strip():
                            # RT-212: stehen die Section-Nummern als Zahl in der Excel, wird die neue auch eine Zahl
                            vorlage = next((ws.cell(r, spalte) for r in (stil, anker)
                                            if r != neu_zeile and _ist_zahlzelle(ws.cell(r, spalte))), None)
                            if vorlage is not None:
                                zahl = _section_zahl(text, vorlage.number_format)
                                if zahl is not None:
                                    zelle.value = zahl
                                    zelle.number_format = vorlage.number_format
                                    continue
                                _hinweise(alle_nach_id[op.aenderung], "Die Section-Nummern stehen in der Excel als Zahl. Die neue "
                                          "Nummer wurde als Text eingetragen, weil sie mit dem Zahlenformat anders aussehen würde.")
                        _setze_wert(zelle, text)
                    if modus == "gelb":
                        for spalte in op.zellen:
                            ws.cell(neu_zeile, spalte).fill = PatternFill("solid", fgColor=farbe)
                    elif modus == "kommentar":
                        a = nach_id[op.aenderung]
                        _kommentar(ws.cell(neu_zeile, min(op.zellen)), f"PDF S.{a.seite}: neu eingefügt", autor)

    # --- Schreiben ----------------------------------------------------------------------------------
    perg = None
    if patch:
        perg = patch_schreiben(quelle, ziel, zell_inhalte, neue_zeilen_patch, cfg)
        treue = pruefe_treue(quelle, ziel, perg.geaenderte_teile)
        if treue:
            ziel.unlink(missing_ok=True)
            raise PatchFehler("Die Ergebnis-Excel wäre nicht verlustfrei geworden; es wurde keine Datei geschrieben. "
                              "Gefunden: " + "; ".join(treue[:3]))
        ergebnis.warnungen.extend(perg.warnungen)
    else:
        ziel.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(ziel))
        verloren = _gefaehrdet(quelle)
        if verloren:
            ergebnis.warnungen.append("Die Excel wurde mit openpyxl gespeichert (anwenden.schreiber = openpyxl). Dabei können "
                                      + ", ".join(verloren) + " verloren gehen. Bitte das Ergebnis mit dem Original vergleichen.")

    # --- Nachkontrolle: gespeicherte Datei neu lesen -----------------------------------------
    def neue_zeile(blatt: str, orig: int) -> int | None:
        if perg is not None:
            return perg.zeile(blatt, orig)
        z = orig
        for e in sorted(abbildung.get(blatt, []), key=lambda t: t[1]):
            if e[0] == "del":
                if e[1] == orig:
                    return None
                if e[1] < orig:
                    z -= 1
            else:
                if e[1] <= orig:
                    z += e[2]       # type: ignore[misc]
        return z

    # Je Aenderung sammeln: in wie vielen Zielen die Wirkung stimmt (ok), fehlt (keine) oder abweicht (abw)
    ok: Counter = Counter()
    keine: Counter = Counter()
    abw: Counter = Counter()
    grund_abw: dict[str, str] = {}

    try:
        kontrolle = load_workbook(str(ziel), rich_text=True, keep_vba=(quelle.suffix.lower() == ".xlsm"))
    except Exception as e:                                   # noqa: BLE001 -- die Datei waere fuer den Nutzer unbrauchbar
        if patch:
            ziel.unlink(missing_ok=True)
            raise PatchFehler("Die Ergebnis-Excel ließ sich nicht wieder öffnen; es wurde keine Datei geschrieben.") from e
        raise

    # --- Gegenprobe A3 (ADR 0004): steht jeder betroffene Absatz so in der Excel, wie das PDF es verlangt? ----------
    durchgefallen: dict[str, str] = {}             # Aenderungs-ID -> Pruefcode
    a3_bestanden: set[str] = set(a3_vorab)         # gegen das PDF geprueft und in Ordnung
    angewendet = set(nach_id)
    if not farbig:
        for ab in absaetze or []:
            if not ab.aenderungen & angewendet:
                continue
            zellen: dict[tuple[int, int], tuple[str, str, list[Operation]]] = {}
            if ab.blatt in kontrolle.sheetnames:
                for z, sp, _, _ in ab.zellen:
                    nz = neue_zeile(ab.blatt, z)
                    gesp = "" if nz is None else _text(kontrolle[ab.blatt].cell(nz, sp))
                    zellen[(z, sp)] = _a3_zelle(absatz_orig.get((ab.blatt, z, sp), ""), gesp,
                                                zell_info.get((ab.blatt, z, sp)), ab.blatt, z, sp)
            code = pruefe_absatz(ab, zellen, angewendet, cfg)
            for aid in sorted(ab.aenderungen & angewendet):
                if code is None:
                    a3_bestanden.add(aid)
                else:
                    durchgefallen.setdefault(aid, code)
        # Neue Absaetze in einer bestehenden Zelle (ohne PDF-Absatz): als eigene Zeile genau der Text aus dem PDF, Rest unveraendert
        if absaetze is not None:
            for ae in aktiv:
                if ae.id not in nach_id or ae.id in mit_absatz or not _neuer_absatz_in_zelle(ae):
                    continue
                soll = ae.soll_absatz or ae.neu
                gut = True
                for o in ae.ops:
                    info = zell_info.get((o.blatt, o.zeile, o.spalte))
                    nz = neue_zeile(o.blatt, o.zeile)
                    if info is None or nz is None:
                        gut = False
                        break
                    gesp = _text(kontrolle[o.blatt].cell(nz, o.spalte))
                    ohne = _zelltext(info["orig"], [x for x in info["ops"] if x.aenderung != ae.id])
                    if not _als_eigene_zeile(gesp, ohne, soll, cfg):
                        gut = False
                        break
                if gut:
                    a3_bestanden.add(ae.id)
                else:
                    durchgefallen.setdefault(ae.id, "gegenprobe_neuer_absatz")
    for aid, code in zurueck.items():
        durchgefallen.setdefault(aid, code)
    for aid in list(durchgefallen):                     # gekoppelte Aenderungen gelten nur gemeinsam
        for k in kopplung.get(aid, ()):
            if k in nach_id and k not in durchgefallen:
                durchgefallen[k] = "gegenprobe_gruppe"
    a3_bestanden -= set(durchgefallen)

    for (blatt, zeile, spalte), info in zell_info.items():
        z = neue_zeile(blatt, zeile)
        gz = None if z is None else kontrolle[blatt].cell(z, spalte)
        tatsaechlich = "" if gz is None else _text(gz)
        if gz is not None and "mark" in info:         # markiert: Zeichen UND Schrift (Strich, Farbe) je Zeichen wie geplant
            zelle_stimmt = _ist_merkmale(gz) == _soll_merkmale(info["mark"].stuecke, info["zellschrift"], cfg)
        elif gz is not None and "farbe" in info:
            k = mk.konfig(cfg)
            soll_m = ((True, f"rgb:{str(k.get('alt_farbe', 'FF00B0F0')).upper()}:0") if info["farbe"] == mk.ALT
                      else (False, f"rgb:{str(k.get('neu_farbe', 'FFFF0000')).upper()}:0"))
            zelle_stimmt = tatsaechlich == info["neu"] and _merkmal(gz.font) == soll_m
        else:
            zelle_stimmt = _ws(tatsaechlich) == _ws(info["neu"])
        for ae in info["aen"]:
            eigene = [o for o in info["ops"] if o.aenderung == ae.id]
            if not zelle_stimmt:
                abw[ae.id] += 1
                grund_abw[ae.id] = "Der Inhalt der Zelle in der gespeicherten Datei weicht vom Plan ab."
            elif (ae.id not in a3_bestanden        # die Gegenprobe gegen das PDF ist genauer als die Wortgrenzen-Regel
                  and any(o.art == "einfuegen" and not _wortgrenze(info["orig"], o.start) for o in eigene)):
                abw[ae.id] += 1
                grund_abw[ae.id] = "Der Text wurde mitten in ein Wort eingefügt (der Plan passt nicht zum Zellinhalt)."
            else:
                if "mark" in info:                            # markiert: wirkt, wenn dieser Lauf sichtbar etwas markiert hat
                    nm = info["mark"].neu_markiert
                    wirkt = (sum(nm.values()) if all(o.abgedeckt for o in eigene) else nm.get(ae.id, 0)) > 0
                elif "zeichen" in info:                       # Zahl: nur die Schrift wird gefaerbt
                    wirkt = True
                elif all(o.abgedeckt for o in eigene):          # Ersetzen in einer Zelle, die ohnehin geleert wurde
                    wirkt = _ws(info["neu"]) != _ws(info["orig"])
                else:
                    ohne = _zelltext(info["orig"], [o for o in info["ops"] if o.aenderung != ae.id])
                    wirkt = _ws(ohne) != _ws(info["neu"])
                (ok if wirkt else keine)[ae.id] += 1
    neue_zeilen: dict[str, tuple[str, int]] = {}
    for (blatt, pos), ops in einfuegungen.items():
        ws = kontrolle[blatt]
        vorher_eingefuegt = sum(e[2] for e in abbildung.get(blatt, []) if e[0] == "ins" and e[1] < pos)   # type: ignore[misc]
        geloescht_davor = sum(1 for e in abbildung.get(blatt, []) if e[0] == "del" and e[1] < pos)
        basis = pos + vorher_eingefuegt - geloescht_davor
        for i, op in enumerate(sorted(ops, key=lambda o: o.seq)):
            r = perg.neue_zeilen[nz_index[id(op)]][1] if perg is not None else basis + i
            stimmt = bool(op.zellen) and all(_ws(_text(ws.cell(r, c))) == _ws(t) for c, t in op.zellen.items())
            if stimmt:
                ok[op.aenderung] += 1
                neue_zeilen[op.aenderung] = (blatt, r)
            else:
                abw[op.aenderung] += 1
                grund_abw[op.aenderung] = "Die neue Zeile in der gespeicherten Datei weicht vom Plan ab."
    # Folgeabsaetze einer neuen Section: Wirkung = die Section-Zeile stimmt UND ihr Absatz steht in der Zelle
    for kopf_id, folgen in gruppen.items():
        kopf_op = next((o for o in nach_id[kopf_id].ops if o.art == "zeile_einfuegen" and not o.abgedeckt), None) \
            if kopf_id in nach_id else None
        for f in folgen:
            if f.id not in nach_id or kopf_op is None:
                continue
            spalten = [c for c, t in kopf_op.zellen.items() if _ws(f.neu) in [_ws(x) for x in t.split("\n")]]
            if kopf_id in neue_zeilen and spalten:
                blatt, r = neue_zeilen[kopf_id]
                zeilen = [_ws(x) for x in _text(kontrolle[blatt].cell(r, spalten[0])).split("\n")]
                if _ws(f.neu) in zeilen:
                    ok[f.id] += 1
                else:
                    abw[f.id] += 1
                    grund_abw[f.id] = "Der Absatz steht nicht in der neuen Section-Zeile."
            elif abw[kopf_id]:
                abw[f.id] += 1
                grund_abw[f.id] = grund_abw.get(kopf_id, "Die zugehörige neue Section stimmt nicht.")
            else:
                keine[f.id] += 1

    for ae in aktiv:
        if ae.id in zurueck:                            # vor dem Schreiben durchgefallen: Status setzt die Gegenprobe unten
            ae.wirkung, ae.nachkontrolle = WIRKUNG_ABWEICHUNG, "abweichung"
            continue
        if abw[ae.id] or (keine[ae.id] and ok[ae.id]):
            w = WIRKUNG_ABWEICHUNG
        elif ok[ae.id]:
            w = WIRKUNG_OK
        else:
            w = WIRKUNG_KEINE
        ae.wirkung = w
        if w == WIRKUNG_OK:
            ae.nachkontrolle = "ok"
            continue
        ae.nachkontrolle = "abweichung"
        if w == WIRKUNG_KEINE:
            _konflikt(ae, "Die Änderung hatte keine Wirkung auf die Excel (nichts hat sich geändert).",
                      "Die Excel ist an dieser Stelle unverändert geblieben. Bitte von Hand prüfen.")
        else:
            _vorschlag(ae, "Die Excel wurde an dieser Stelle geändert, aber nicht so wie geplant.",
                       grund_abw.get(ae.id, "Das Ergebnis weicht vom Plan ab."),
                       "Bitte die Stelle in der Ergebnis-Excel von Hand prüfen.")
            ae.konfidenz = NIEDRIG
            ergebnis.warnungen.append(f"Änderung {ae.id} ({ae.ziel}): in der Ergebnis-Excel weicht die Stelle vom Plan ab. Bitte von Hand prüfen.")

    # --- Gegenprobe A3: Ergebnis eintragen (nie "uebernommen/ok" fuer etwas, das nicht zum PDF passt) ---------------
    if durchgefallen:
        for aid, code in durchgefallen.items():
            ae = alle_nach_id[aid]
            _code(ae, code)
            if ae.status == NICHT:                      # schon zurueckgestuft (z. B. ohne Wirkung): nie "uebernommen"
                continue
            if ae.status == UEBERNOMMEN:
                _vorschlag(ae, "Gegenprobe: Das Ergebnis in der Excel passt nicht zu dem, was das PDF verlangt.")
                ergebnis.warnungen.append(f"Änderung {ae.id} ({ae.ziel}): Gegenprobe – das Ergebnis in der Excel passt nicht "
                                          f"zu dem, was das PDF verlangt. Bitte von Hand prüfen.")
            _hinweise(ae, "Gegenprobe: Das Ergebnis in der Excel passt nicht zu dem, was das PDF verlangt. Bitte prüfen."
                      if code != "gegenprobe_gruppe" else
                      "Gegenprobe: Eine Änderung derselben Gruppe passt nicht zum PDF; die Gruppe gilt nur gemeinsam. Bitte prüfen.")
            if aid in zurueck:
                _hinweise(ae, "Die Änderung wurde deshalb nicht in die Ergebnis-Excel geschrieben.")
            ae.konfidenz, ae.wirkung, ae.nachkontrolle = NIEDRIG, WIRKUNG_ABWEICHUNG, "abweichung"

    # --- Hinweise: Inhalt in gestrichenen Sections ----------------------------------------
    aktive_ops = [o for a in aktiv if a.status == UEBERNOMMEN for o in a.ops]
    for ae in aktiv:
        if ae.typ == "section_geloescht" and ae.status == UEBERNOMMEN:
            for op in ae.ops:
                lay = xl.layouts.get(op.blatt)
                if lay and op.zeile in lay.kopf_zeilen and lay.inhalt_spalte:
                    for r in xl.section_zeilen(op.blatt, op.zeile):
                        if (op.blatt, r, lay.inhalt_spalte) not in geleerte and xl.text(op.blatt, r, lay.inhalt_spalte).strip():
                            if not any(o.zeile == r and o.blatt == op.blatt for o in aktive_ops):
                                ergebnis.warnungen.append(
                                    f"Section '{ae.alt}' entfällt, aber in {op.blatt}!{get_column_letter(lay.inhalt_spalte)}{r} "
                                    f"steht noch Inhalt (im PDF nicht gestrichen). Bitte prüfen.")
                                break
    return ergebnis
