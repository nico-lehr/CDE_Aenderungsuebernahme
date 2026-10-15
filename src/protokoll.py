"""Bilanz und Aenderungsprotokoll (damit ein Mensch alles pruefen kann).

Bilanz-Regel: Jede im PDF erkannte Aenderung steht entweder in der Excel
(uebernommen), wartet auf Freigabe (pruefen) oder steht mit Grund auf der Liste
"nicht uebernommen". Die Woerter-Summe muss mit der Zahl der markierten PDF-Woerter
uebereinstimmen -- sonst ist die Bilanz rot.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import NICHT, PRUEFEN, SCHEMA_VERSION, UEBERNOMMEN, Aenderung

AUTOR = "Nico Lehr"
TOOL = "CDE-Änderungsübernahme"

_WASH = {UEBERNOMMEN: "E4F4E8", PRUEFEN: "FBEFDA", NICHT: "FBE7E5"}
_NAMEN = {UEBERNOMMEN: "übernommen", PRUEFEN: "zu prüfen", NICHT: "nicht übernommen"}
_TYPEN = {
    "text_ersetzt": "Text ersetzt", "text_eingefuegt": "Text eingefügt", "text_geloescht": "Text gelöscht",
    "absatz_neu": "Absatz neu", "absatz_geloescht": "Absatz gelöscht",
    "section_neu": "Section neu", "section_geloescht": "Section gelöscht",
}


def bilanz(aenderungen: list[Aenderung], markierte_woerter: int, ignoriert: list[dict[str, Any]]) -> dict[str, Any]:
    zaehler = {UEBERNOMMEN: 0, PRUEFEN: 0, NICHT: 0}
    woerter = {UEBERNOMMEN: 0, PRUEFEN: 0, NICHT: 0}
    for a in aenderungen:
        zaehler[a.status] = zaehler.get(a.status, 0) + 1
        woerter[a.status] = woerter.get(a.status, 0) + a.woerter
    ign = sum(i["woerter"] for i in ignoriert)
    summe = sum(woerter.values()) + ign
    return {
        "aenderungen": len(aenderungen),
        "uebernommen": zaehler[UEBERNOMMEN], "pruefen": zaehler[PRUEFEN], "nicht_uebernommen": zaehler[NICHT],
        "woerter_uebernommen": woerter[UEBERNOMMEN], "woerter_pruefen": woerter[PRUEFEN],
        "woerter_nicht_uebernommen": woerter[NICHT], "woerter_ignoriert": ign,
        "woerter_im_pdf": markierte_woerter,
        "stimmt": summe == markierte_woerter,
        "nachkontrolle_abweichungen": sum(1 for a in aenderungen if a.nachkontrolle == "abweichung"),
    }


def aenderung_aus_dict(d: dict[str, Any]) -> Aenderung:
    from .models import Operation
    d = dict(d)
    ops = []
    for o in d.pop("ops", []):
        o = dict(o)
        o["zellen"] = {int(k): v for k, v in (o.get("zellen") or {}).items()}
        ops.append(Operation(**o))
    d["bbox"] = tuple(d["bbox"])
    return Aenderung(ops=ops, **d)


def schreibe_json(pfad: Path, daten: dict[str, Any]) -> None:
    daten = {"schema_version": SCHEMA_VERSION, **daten}
    pfad.write_text(json.dumps(daten, ensure_ascii=False, indent=1), encoding="utf-8")


def schreibe_xlsx(pfad: Path, info: dict[str, Any], aenderungen: list[Aenderung], bil: dict[str, Any],
                  layout: list[dict], warnungen: list[str], farben: list[dict] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bilanz"
    ws["A1"] = f"{TOOL} – Änderungsprotokoll"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Erstellt von {AUTOR} · {datetime.now():%d.%m.%Y %H:%M}"
    zeilen = [
        ("PDF mit Änderungen", info.get("pdf_name", "")), ("Excel (Original, unverändert)", info.get("excel_name", "")),
        ("Ergebnis-Excel", info.get("ergebnis_name", "(noch nicht erstellt)")),
        ("", ""),
        ("Erkannte Änderungen", bil["aenderungen"]),
        ("übernommen", bil["uebernommen"]), ("zu prüfen (Vorschlag, wartet auf Freigabe)", bil["pruefen"]),
        ("nicht übernommen (mit Grund)", bil["nicht_uebernommen"]),
        ("", ""),
        ("Markierte Wörter im PDF", bil["woerter_im_pdf"]),
        ("… davon übernommen", bil["woerter_uebernommen"]), ("… davon zu prüfen", bil["woerter_pruefen"]),
        ("… davon nicht übernommen", bil["woerter_nicht_uebernommen"]),
        ("… davon ignoriert (Kopf-/Fußzeile)", bil["woerter_ignoriert"]),
        ("Bilanz stimmt (nichts ging still verloren)", "JA" if bil["stimmt"] else "NEIN – bitte prüfen"),
        ("Nachkontrolle-Abweichungen", bil["nachkontrolle_abweichungen"]),
    ]
    for i, (a, b) in enumerate(zeilen, 4):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
    r = 4 + len(zeilen) + 1
    if warnungen:
        ws.cell(r, 1, "Hinweise").font = Font(bold=True)
        for w in warnungen:
            r += 1
            ws.cell(r, 1, w)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60

    ws2 = wb.create_sheet("Änderungen")
    kopf = ["Nr.", "Seite", "Art", "Section", "Status", "Sicherheit", "Bestätigt", "Alt (PDF)", "Neu (PDF)",
            "Stelle in Excel", "Hinweis / Grund", "Nachkontrolle", "Wörter"]
    ws2.append(kopf)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8ED")
    for a in aenderungen:
        ws2.append([a.id, a.seite, _TYPEN.get(a.typ, a.typ), a.section or "", _NAMEN.get(a.status, a.status), a.konfidenz,
                    "ja" if a.bestaetigt else "", a.alt, a.neu, a.ziel, a.grund, a.nachkontrolle, a.woerter])
        ws2.cell(ws2.max_row, 5).fill = PatternFill("solid", fgColor=_WASH.get(a.status, "FFFFFF"))
        for c in ws2[ws2.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, b in enumerate((8, 7, 16, 9, 16, 11, 10, 48, 48, 22, 44, 14, 8), 1):
        ws2.column_dimensions[get_column_letter(i)].width = b
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("Nicht übernommen")
    ws3.append(kopf[:2] + ["Art", "Alt (PDF)", "Neu (PDF)", "Stelle in Excel (Vorschlag)", "Grund"])
    for c in ws3[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8ED")
    for a in aenderungen:
        if a.status != UEBERNOMMEN:
            ws3.append([a.id, a.seite, _TYPEN.get(a.typ, a.typ), a.alt, a.neu, a.ziel,
                        a.grund or ("Wartet auf Freigabe" if a.status == PRUEFEN else "")])
            for c in ws3[ws3.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, b in enumerate((8, 7, 16, 48, 48, 26, 56), 1):
        ws3.column_dimensions[get_column_letter(i)].width = b

    ws4 = wb.create_sheet("Erkannter Aufbau")
    ws4.append(["Blatt", "Section-Spalte", "Titel-Spalte", "Inhalt-Spalte", "Modell", "Sections", "Zeilen"])
    for l in layout:
        ws4.append([l["blatt"], l["section_spalte"], l["titel_spalte"], l["inhalt_spalte"], l["modell"], l["sections"], l["zeilen"]])
    for c in ws4[1]:
        c.font = Font(bold=True)
    if farben:
        ws5 = wb.create_sheet("Farben im PDF")
        ws5.append(["Farbe", "Anzahl", "Klasse", "Art", "Bedeutung"])
        for f in farben:
            ws5.append([f["farbe"], f["zeichen"], f["klasse"], f["art"], f["bedeutung"]])
        for c in ws5[1]:
            c.font = Font(bold=True)
    wb.save(str(pfad))
