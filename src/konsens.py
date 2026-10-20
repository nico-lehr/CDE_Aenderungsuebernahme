"""KONSENS: ein unabhaengiger zweiter Pruefweg gegen stille Fehler (Auftrag 003, Welle 4).

Selbst wenn PDF-Leser und Planer sich irren, soll eine automatisch uebernommene Aenderung nicht still falsch in der Excel
landen. Zwei Pruefungen, die voneinander und vom Planer unabhaengig sind (sie benutzen weder dessen Operationen noch dessen
Hilfsfunktionen, nur die Rohdaten: PDF-Woerter mit Lage, PDF-Absaetze als Laeufe, die beiden Excel-Dateien):

  (a) Wort-Bildprobe (nur PDF): Die Seite wird gerendert (pypdfium2, unter PDFIUM_SPERRE). An der Wortbox jedes Wortes eines
      Blocks mit automatisch uebernommener Aenderung muss sichtbar sein, was der Leser behauptet: neu = farbige Unterstreichung
      (Legende "linie") bzw. Farbe ("farbe"/"alt"), entfaellt = Durchstreichung in Zeilenmitte (farbig oder dunkel), bleibt =
      weder Linie noch Aenderungsfarbe. Ist ein Wort nicht sicher beurteilbar (zu klein, farbiger Grund, nur Satzzeichen ...),
      zaehlt es als "nicht pruefbar" -- nie als Abweichung. Jede Abweichung stuft die Aenderungen des Blocks auf "pruefen"
      (Code `konsens_bild`).
  (b) Konsens-Text NACH dem Schreiben: Aus den PDF-Absaetzen (plan.absaetze) wird ein eigenes Soll fuer alten und neuen Text
      je Zeichen (Zeichen + Klasse alt / neu / normal) mit einer eigenen, einfachen Normalisierung berechnet und mit der
      geschriebenen Zelle (gelesen mit openpyxl aus der Ergebnisdatei) verglichen. Dazu muss jede Zelle, die nicht zu einer
      uebernommenen Aenderung gehoert, unveraendert sein (Zellen anderer Zeilen oder Specs). Jede Abweichung stuft die
      betroffenen Aenderungen auf "pruefen" (Code `konsens_text`); der Applier schreibt die Datei dann neu, ohne sie.

Schwellen und Schalter stehen in config/zuordnung.yaml (Abschnitt `konsens:`, Regel 28); `konsens.aktiv: false` schaltet beides
ab. Rueckgaben enthalten nie Dokumentinhalte (Regel 16), nur Zahlen und Kategorien.
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

from . import config
from .models import ADD, DEL, KEEP, NICHT, PRUEFEN, UEBERNOMMEN, WIRKUNG_ABWEICHUNG, Absatz, Aenderung

CODE_TEXT = "konsens_text"
CODE_BILD = "konsens_bild"
ALT, NEU, NORMAL, ALT2 = "alt", "neu", "normal", "alt_anders"     # ALT2: durchgestrichen, aber nicht in der Alt-Farbe (frueher)


def _strich(klasse: str) -> bool:
    return klasse in (ALT, ALT2)


def konfig(cfg: dict | None = None) -> dict[str, Any]:
    """Abschnitt `konsens:` aus config/zuordnung.yaml."""
    if cfg and "konsens" in cfg:
        return cfg["konsens"] or {}
    return config.lade("zuordnung").get("konsens", {}) or {}


def aktiv(cfg: dict | None = None) -> bool:
    return bool(konfig(cfg).get("aktiv", False))


def _auto(a: Aenderung) -> bool:
    """Automatisch uebernommen (nicht von einem Menschen freigegeben)."""
    return a.status == UEBERNOMMEN and not a.bestaetigt


# =============================================================================================== (a) Wort-Bildprobe
@dataclass
class BildErgebnis:
    """Ergebnis der Wort-Bildprobe. `abweichungen`: Aenderungs-ID -> kurze Beschreibungen (ohne Dokumentinhalt)."""
    gelaufen: bool = False
    grund: str = ""                       # warum die Probe nicht (ganz) lief, einfaches Deutsch
    geprueft: int = 0                     # Woerter mit klarem Befund (in Ordnung oder abweichend)
    unpruefbar: int = 0                   # Woerter, bei denen die Probe unsicher war (keine Abweichung)
    abweichend: int = 0                   # Woerter mit Abweichung
    abweichungen: dict[str, list[str]] = field(default_factory=dict)
    gruende: dict[str, int] = field(default_factory=dict)          # warum Woerter nicht pruefbar waren (Zahlen, kein Inhalt)

    def nicht_pruefbar(self, grund: str, n: int = 1) -> None:
        self.unpruefbar += n
        self.gruende[grund] = self.gruende.get(grund, 0) + n

    def zusammenfassung(self) -> str:
        return (f"Wort-Bildprobe: {self.geprueft} Wörter geprüft, {self.abweichend} abweichend, "
                f"{self.unpruefbar} nicht prüfbar.")


def _hsv(rgb: tuple[int, int, int]) -> tuple[float, float, float]:
    import colorsys
    return colorsys.rgb_to_hsv(*(v / 255.0 for v in rgb))


def _bloecke_der_aenderungen(analyse: Any, aenderungen: Iterable[Aenderung]) -> dict[int, list[Aenderung]]:
    """Block-Nummer -> automatisch uebernommene Aenderungen in diesem Block."""
    vorhanden = {b.idx for b in analyse.blocks if not b.ignoriert}
    out: dict[int, list[Aenderung]] = {}
    for a in aenderungen:
        if _auto(a) and a.block_idx in vorhanden:
            out.setdefault(a.block_idx, []).append(a)
    return out


def _tinte(crop: Any, bg: tuple[int, int, int], abstand: int) -> Any:
    from PIL import Image, ImageChops
    d = ImageChops.difference(crop, Image.new("RGB", crop.size, bg))
    r, g, b = d.split()
    return ImageChops.lighter(ImageChops.lighter(r, g), b).point(lambda v: 255 if v >= abstand else 0)


def _laengste_strecke(zeile: bytes) -> int:
    return max((len(s) for s in zeile.split(b"\x00")), default=0)


def _linie_im_band(maske: Any, crop: Any, y0: int, y1: int, min_anteil: float, max_dicke: int) -> tuple[bool, tuple[int, int, int] | None]:
    """Gibt es im Hoehenband [y0, y1) eine duenne Linie: Pixelzeilen, in denen zusammenhaengende Tinte fast die ganze Breite
    deckt, hoechstens `max_dicke` Zeilen dick (dickere Gruppen sind Flaechen, z. B. eine Hinterlegung der Nachbarzeile)?
    Rueckgabe (ja/nein, mittlere Farbe der Linie)."""
    from PIL import Image
    b, h = maske.size
    daten = maske.tobytes()
    gruppe: list[int] = []
    gruppen: list[list[int]] = []
    for y in range(h):
        if _laengste_strecke(daten[y * b:(y + 1) * b]) >= min_anteil * b:
            gruppe.append(y)
            continue
        if gruppe:
            gruppen.append(gruppe)
            gruppe = []
    if gruppe:
        gruppen.append(gruppe)
    for g in gruppen:
        if len(g) <= max_dicke and g[-1] >= y0 and g[0] < y1:
            mitte = g[len(g) // 2]
            return True, crop.crop((0, mitte, b, mitte + 1)).resize((1, 1), Image.Resampling.BOX).getpixel((0, 0))
    return False, None


def _wort_befund(bild: Any, tok: Any, modus: str, bk: dict, seiten_breite: float, hoehe_ref: float) -> tuple[str, str]:
    """Befund fuer EIN Wort: ("ok" | "abweichung" | "unpruefbar", Kurzbeschreibung ohne Inhalt)."""
    from PIL import Image
    s = float(bk.get("aufloesung", 3.0))
    x0, top, x1, bot = tok.bbox
    w, h = x1 - x0, bot - top
    if w < float(bk.get("min_breite_pt", 9.0)) or h < float(bk.get("min_hoehe_pt", 5.0)):
        return "unpruefbar", "zu klein"
    if h < float(bk.get("hochgestellt_anteil", 0.75)) * hoehe_ref:
        return "unpruefbar", "hochgestellt"
    if not any(ch.isalnum() for ch in (tok.new or tok.old or "")):
        return "unpruefbar", "nur Satzzeichen"
    if tok.teile and len({st for st, _ in tok.teile}) > 1:
        return "unpruefbar", "Wort nur zum Teil markiert"
    if x0 < 0 or top < 0 or x1 > bild.width / s or bot > bild.height / s:
        return "unpruefbar", "ausserhalb der Seite"
    rand = float(bk.get("rand_anteil", 0.1)) * w
    unten = float(bk.get("unterstrichen_bis", 1.3))
    extra = float(bk.get("linie_max_dicke_pt", 1.6)) + 1.0          # etwas tiefer ausschneiden, um die Dicke einer Linie zu messen
    box = (int((x0 + rand) * s), int(top * s), int((x1 - rand) * s) + 1,
           min(bild.height, int((top + unten * h + extra) * s) + 1))
    crop = bild.crop(box).convert("RGB")
    cb, ch_ = crop.size
    if cb < 6 or ch_ < 6:
        return "unpruefbar", "zu klein"
    farben = crop.getcolors(maxcolors=cb * ch_) or []
    if not farben:
        return "unpruefbar", "kein Bild"
    bg = max(farben)[1]
    if min(bg) < int(bk.get("hintergrund_min_helligkeit", 225)):
        return "unpruefbar", "farbiger Hintergrund"
    maske = _tinte(crop, bg, int(bk.get("tinte_abstand", 70)))
    tinte_n = maske.histogram()[255]
    if tinte_n < int(bk.get("min_tinte_pixel", 8)):
        return "unpruefbar", "kaum Tinte"
    hoehe_px = h * s
    d0, d1 = bk.get("band_durchgestrichen", [0.3, 0.72])
    u0, u1 = bk.get("band_unterstrichen", [0.8, 1.3])
    min_a = float(bk.get("linie_min_anteil", 0.85))
    dicke = max(1, int(round(float(bk.get("linie_max_dicke_pt", 1.6)) * s)))
    strich, _ = _linie_im_band(maske, crop, int(d0 * hoehe_px), int(d1 * hoehe_px) + 1, min_a, dicke)
    unter, ucol = _linie_im_band(maske, crop, int(u0 * hoehe_px), int(u1 * hoehe_px) + 1, min_a, dicke)
    farbig_u = bool(unter and ucol and _hsv(ucol)[1] >= float(bk.get("linie_min_saettigung", 0.12)))
    # Farbe des Textes (fuer die Legenden farbe / alt): Anteil farbiger Tinte
    hsv = crop.convert("HSV")
    hh, ss, vv = hsv.split()
    smin, vmin = float(bk.get("text_min_saettigung", 0.35)) * 255, float(bk.get("text_min_helligkeit", 0.25)) * 255
    from PIL import ImageChops
    farb = ImageChops.multiply(ss.point(lambda v: 255 if v >= smin else 0), vv.point(lambda v: 255 if v >= vmin else 0))
    farb = ImageChops.multiply(farb, maske)
    farb_n = farb.histogram()[255]
    farbanteil = farb_n / max(1, tinte_n)
    blau = ImageChops.multiply(farb, hh.point(lambda v: 255 if 131 <= v <= 188 else 0)).histogram()[255]
    rot = ImageChops.multiply(farb, hh.point(lambda v: 255 if (v <= 13 or v >= 237) else 0)).histogram()[255]
    mind = float(bk.get("farbtext_min_anteil", 0.5))
    keep_max = float(bk.get("keep_max_farbanteil", 0.15))

    if modus == "linie":
        if tok.status == ADD:
            ok = farbig_u and not strich
            return ("ok", "") if ok else ("abweichung", "als neu gelesen, aber keine farbige Unterstreichung" if not strich
                                          else "als neu gelesen, aber durchgestrichen")
        if tok.status == DEL:
            return ("ok", "") if strich else ("abweichung", "als entfallen gelesen, aber keine Durchstreichung")
        if strich:
            return "abweichung", "als unverändert gelesen, aber durchgestrichen"
        if farbig_u:
            return "abweichung", "als unverändert gelesen, aber farbig unterstrichen"
        return "ok", ""
    # Legende farbe / alt: Farbe zaehlt
    if modus == "farbe":
        if tok.status == ADD:
            ok = farbanteil >= mind and not strich
            return ("ok", "") if ok else ("abweichung", "als neu gelesen, aber nicht farbig" if not strich
                                          else "als neu gelesen, aber durchgestrichen")
        if tok.status == DEL:
            return ("ok", "") if strich else ("abweichung", "als entfallen gelesen, aber keine Durchstreichung")
        if strich:
            return "abweichung", "als unverändert gelesen, aber durchgestrichen"
        if farbanteil > keep_max:
            return "abweichung", "als unverändert gelesen, aber farbig"
        return "ok", ""
    # alt: blau = neu, rot = entfaellt
    if tok.status == ADD:
        return ("ok", "") if (farbanteil >= mind and blau > rot) else ("abweichung", "als neu gelesen, aber nicht blau")
    if tok.status == DEL:
        return ("ok", "") if ((farbanteil >= mind and rot >= blau) or strich) else ("abweichung",
                                                                                   "als entfallen gelesen, aber nicht rot")
    if farbanteil > keep_max:
        return "abweichung", "als unverändert gelesen, aber farbig"
    return "ok", ""


def bildprobe(analyse: Any, aenderungen: Iterable[Aenderung], cfg: dict | None = None, pdf: str | Path | None = None) -> BildErgebnis:
    """Wort-Bildprobe fuer alle Bloecke mit automatisch uebernommenen Aenderungen (siehe Modul-Text). Aendert nichts an den
    Aenderungen; `bildprobe_anwenden` stuft ab. Ohne PDF (Word-Eingabe, `legende_modus` "docx") oder bei ausgeschaltetem
    `konsens` laeuft sie nicht (`gelaufen` False, `grund` sagt warum)."""
    aenderungen = list(aenderungen)
    erg = BildErgebnis()
    k = konfig(cfg)
    if not k.get("aktiv", False):
        erg.grund = "ausgeschaltet"
        return erg
    modus = str(getattr(analyse, "legende_modus", "") or "")
    if modus not in ("linie", "farbe", "alt"):
        erg.grund = "Die Wort-Bildprobe gibt es nur für PDF-Eingaben."
        return erg
    pfad = Path(pdf if pdf is not None else getattr(analyse, "pfad", ""))
    if pfad.suffix.lower() != ".pdf" or not pfad.is_file():
        erg.grund = "Das PDF liegt nicht (mehr) vor; die Wort-Bildprobe wurde übersprungen."
        return erg
    bloecke = _bloecke_der_aenderungen(analyse, aenderungen)
    if not bloecke:
        erg.gelaufen = True
        return erg
    bk = k.get("bild", {}) or {}
    s = float(bk.get("aufloesung", 3.0))
    try:
        import pypdfium2 as pdfium

        from .vorschau import PDFIUM_SPERRE
    except Exception as e:                               # noqa: BLE001 -- die Pruefung fehlt dann, mehr nicht
        erg.grund = f"Die Wort-Bildprobe konnte nicht starten ({type(e).__name__})."
        return erg
    nach_idx = {b.idx: b for b in analyse.blocks}
    je_seite: dict[int, list[int]] = {}
    for idx in bloecke:
        for seite in sorted({t.page for t in nach_idx[idx].tokens}):
            je_seite.setdefault(seite, []).append(idx)
    breiten = getattr(analyse, "seiten_masse", {}) or {}
    with PDFIUM_SPERRE:
        try:
            doc = pdfium.PdfDocument(str(pfad))
        except Exception as e:                           # noqa: BLE001
            erg.grund = f"Das PDF ließ sich für die Wort-Bildprobe nicht öffnen ({type(e).__name__})."
            return erg
        try:
            for seite, idxs in sorted(je_seite.items()):
                woerter = [(idx, t) for idx in idxs for t in nach_idx[idx].tokens if t.page == seite]
                try:
                    pg = doc[seite - 1]
                    breite_pt, hoehe_pt = pg.get_size()
                    gedreht = bool(pg.get_rotation())
                    cb, mb = pg.get_cropbox(), pg.get_mediabox()
                    masse = breiten.get(seite)
                    passt = (not gedreht and tuple(round(v, 1) for v in cb) == tuple(round(v, 1) for v in mb)
                             and (masse is None or (abs(masse[0] - breite_pt) < 1.5 and abs(masse[1] - hoehe_pt) < 1.5)))
                    bild = pg.render(scale=s).to_pil().convert("RGB") if passt else None
                except Exception:                        # noqa: BLE001
                    bild = None
                if bild is None:
                    erg.nicht_pruefbar("Seite nicht rendert", len(woerter))
                    continue
                hoehen: dict[int, float] = {}
                for idx in idxs:
                    hs = sorted(t.bbox[3] - t.bbox[1] for t in nach_idx[idx].tokens)
                    hoehen[idx] = hs[len(hs) // 2] if hs else 0.0
                for idx, t in woerter:
                    if t.status not in (KEEP, ADD, DEL):
                        erg.nicht_pruefbar("Status")
                        continue
                    if t.status == KEEP and t.hinweis:
                        erg.nicht_pruefbar("Sonderfall")      # Link, Hinterlegung, farbiges Wort ohne Linie ...
                        continue
                    if getattr(analyse, "markup_leiste", False) and t.bbox[0] >= 0.68 * bild.width / s:
                        erg.nicht_pruefbar("Markup-Leiste")
                        continue
                    try:
                        befund, was = _wort_befund(bild, t, modus, bk, bild.width / s, hoehen[idx])
                    except Exception:                    # noqa: BLE001
                        befund, was = "unpruefbar", "Fehler"
                    if befund == "unpruefbar":
                        erg.nicht_pruefbar(was)
                        continue
                    erg.geprueft += 1
                    if befund == "abweichung":
                        erg.abweichend += 1
                        for a in bloecke[idx]:
                            lst = erg.abweichungen.setdefault(a.id, [])
                            if len(lst) < 5:
                                lst.append(f"Seite {seite}: Ein Wort ist {was}.")
        finally:
            try:
                doc.close()
            except Exception:                            # pragma: no cover
                pass
    erg.gelaufen = True
    return erg


def abgestuft(a: Aenderung, code: str, grund: str, hinweise: Iterable[str] = ()) -> None:
    """Eine Aenderung auf "pruefen" zurueck (wie eine Abweichung der Gegenprobe A3): nie still, nie automatisch geschrieben."""
    from . import codes
    a.status, a.konfidenz = PRUEFEN, "niedrig"
    a.grund = (a.grund + "; " if a.grund else "") + grund
    if code not in a.pruefcodes:
        a.pruefcodes.append(code)
    for h in hinweise:
        if h and h not in a.hinweise:
            a.hinweise.append(h)
    a.grund_einfach = codes.grund_einfach(a)


def bildprobe_anwenden(aenderungen: list[Aenderung], analyse: Any, cfg: dict | None = None,
                       pdf: str | Path | None = None) -> BildErgebnis:
    """Wort-Bildprobe im PLAN (vor dem Schreiben): jede Aenderung mit Abweichung -- samt den Aenderungen, die nur gemeinsam
    gelten (Section + Folgeabsaetze, gleiche `gruppe`) -- wird "pruefen" mit Code `konsens_bild`. Dafuer reichen Plan und PDF."""
    erg = bildprobe(analyse, aenderungen, cfg, pdf)
    if not erg.abweichungen:
        return erg
    from .applier import _folge_gruppen, _gekoppelt
    gruppen = _folge_gruppen(aenderungen)
    kopplung = _gekoppelt(aenderungen, gruppen)
    nach_id = {a.id: a for a in aenderungen}
    for aid, was in erg.abweichungen.items():
        for k in sorted(kopplung.get(aid, {aid})):
            a = nach_id.get(k)
            if a is None or a.status != UEBERNOMMEN:
                continue
            if k == aid:
                abgestuft(a, CODE_BILD, "Zweitprüfung: Im Seitenbild ist die Markierung anders als gelesen.", was)
            else:
                abgestuft(a, "gegenprobe_gruppe", "Gehört zu einer Änderung, deren Markierung im Seitenbild nicht stimmt.")
    return erg


# =============================================================================================== (b) Konsens-Text
_Z = tuple[str, str]                  # (Zeichen, Klasse)


@dataclass
class TextErgebnis:
    """Ergebnis des Konsens-Texts. `abweichungen`: Aenderungs-ID -> Art ("text" | "fremd"), nie Dokumentinhalt."""
    gelaufen: bool = False
    geprueft: int = 0                     # Absaetze, deren Zellen verglichen wurden
    unpruefbar: int = 0                   # Absaetze, die nicht verglichen werden konnten (Zahl-Zelle, Mehrfachbelegung ...)
    fremd: int = 0                        # Zellen ausserhalb der Aenderungen, die sich geaendert haben
    abweichungen: dict[str, str] = field(default_factory=dict)

    def zusammenfassung(self) -> str:
        return (f"Konsens-Text: {self.geprueft} Absätze verglichen, {len(self.abweichungen)} Änderungen mit Abweichung, "
                f"{self.fremd} unerwartet geänderte Zellen, {self.unpruefbar} Absätze nicht prüfbar.")


class _Norm:
    """Eigene, einfache Normalisierung: Zeichen einzeln, NFKC, ohne Leerraum und unsichtbare Zeichen, Strich- und
    Anfuehrungszeichen-Varianten gleich (Tabellen aus der Konfiguration)."""

    def __init__(self, tk: dict) -> None:
        self.entfernen = set(tk.get("entfernen", []))
        self.ersetzen = {str(k): str(v) for k, v in (tk.get("ersetzen", {}) or {}).items()}
        self.listen = set(tk.get("listenzeichen", []))
        self._cache: dict[str, str] = {}

    def zeichen(self, c: str) -> str:
        r = self._cache.get(c)
        if r is None:
            if c in self.entfernen or c.isspace() or unicodedata.category(c) in ("Cf", "Cc"):
                r = ""
            else:
                r = "".join(self.ersetzen.get(t, t) for t in unicodedata.normalize("NFKC", self.ersetzen.get(c, c))
                            if not (t in self.entfernen or t.isspace() or unicodedata.category(t) in ("Cf", "Cc")))
            self._cache[c] = r
        return r

    def strom(self, zeichen: Iterable[tuple]) -> list[tuple]:
        """[(zeichen, klasse, ...)] -> [(einzelnes Zeichen, klasse, ...)] ohne Leerraum."""
        out = []
        for z in zeichen:
            for t in self.zeichen(z[0]):
                out.append((t,) + tuple(z[1:]))
        return out


def _rgb(font: Any) -> str:
    c = getattr(font, "color", None)
    if c is None or getattr(c, "type", "rgb") != "rgb":
        return ""
    v = getattr(c, "rgb", None)
    return v[-6:].upper() if isinstance(v, str) and len(v) >= 6 else ""


def _merkmal(font: Any) -> tuple[bool, str]:
    return (bool(getattr(font, "strike", False)), _rgb(font)) if font is not None else (False, "")


def _klasse(merkmal: tuple[bool, str], neu_rgb: str, alt_rgb: str) -> str:
    if merkmal[0]:
        return ALT if merkmal[1] == alt_rgb else ALT2
    return NEU if merkmal[1] == neu_rgb else NORMAL


def _zeichen_der_zelle(zelle: Any, neu_rgb: str, alt_rgb: str) -> list[_Z] | None:
    """(Zeichen, Klasse) je Zeichen; None fuer Zellen, die kein Text sind (Zahl, Formel)."""
    if zelle is None or zelle.value is None:
        return []
    v = zelle.value
    if isinstance(v, CellRichText):
        out: list[_Z] = []
        for b in v:
            if isinstance(b, TextBlock):
                k = _klasse(_merkmal(b.font), neu_rgb, alt_rgb)
                out += [(c, k) for c in b.text]
            else:
                k = _klasse(_merkmal(zelle.font), neu_rgb, alt_rgb)
                out += [(c, k) for c in str(b)]
        return out
    if isinstance(v, str) and zelle.data_type != "f":
        k = _klasse(_merkmal(zelle.font), neu_rgb, alt_rgb)
        return [(c, k) for c in v]
    return None


def _kanon(zelle: Any) -> Any:
    """Vergleichsform einer Zelle fuer "unveraendert?": Text je Zeichen mit (durchgestrichen, Farbe); Zahlen als Wert."""
    if zelle is None or zelle.value is None:
        return None
    v = zelle.value
    if isinstance(v, CellRichText) or isinstance(v, str):
        if isinstance(v, str) and zelle.data_type == "f":
            return "formel"
        m0 = _merkmal(zelle.font)
        if isinstance(v, str):
            return (v, m0)
        laeufe: list[list] = []
        for b in v:
            t, m = (b.text, _merkmal(b.font)) if isinstance(b, TextBlock) else (str(b), m0)
            if laeufe and laeufe[-1][1] == m:
                laeufe[-1][0] += t
            else:
                laeufe.append([t, m])
        if len(laeufe) == 1:
            return (laeufe[0][0], laeufe[0][1])
        return tuple((t, m) for t, m in laeufe)
    if zelle.data_type == "f":
        return "formel"
    return (v, _merkmal(zelle.font))


def _soll_strom(laeufe: Iterable, aktiv_ids: set[str], markiert: bool, norm: _Norm) -> list[tuple[str, str, bool]]:
    """Eigenes Soll eines Absatzes je Zeichen: (Zeichen, Klasse, optional). Entfaellt-Laeufe angewendeter Aenderungen
    bleiben stehen (alt) bzw. fehlen (ohne Markierung), neue Laeufe angewendeter Aenderungen sind neu bzw. normal, nicht
    angewendete neue fehlen, nicht angewendete Entfaellt-Laeufe bleiben normal."""
    out: list[tuple[str, str, bool]] = []
    for l in laeufe:
        if l.art == "s":
            continue
        if l.art == "d" and l.aenderung in aktiv_ids:
            if not markiert:
                continue
            k = ALT
        elif l.art == "a":
            if l.aenderung not in aktiv_ids:
                continue
            k = NEU if markiert else NORMAL
        else:
            k = NORMAL
        out += [(t[0], k, bool(l.optional)) for t in norm.strom([(c,) for c in l.text])]
    return out


def _stuecke_gleich(soll: list[tuple[str, str, bool]], ist: list[_Z], markiert: bool, neu_toleranz: bool,
                    listen: set[str]) -> bool:
    """Zeichen und (mit Markierung) Klasse stimmen ueberein; optionale Soll-Zeichen duerfen fehlen."""
    if markiert and not neu_toleranz:
        gleich: Callable[[str, str], bool] = lambda a, b: a == b            # noqa: E731
    elif markiert:
        gleich = lambda a, b: a == b or (a == NORMAL and b == NEU)           # noqa: E731
    else:
        gleich = lambda a, b: True                                            # noqa: E731
    if soll and ist and soll[0][0] != ist[0][0]:        # optionales Aufzaehlungszeichen am Anfang
        if soll[0][0] in listen:
            soll = soll[1:]
        elif ist[0][0] in listen:
            ist = ist[1:]
    i = j = 0
    while i < len(soll) and j < len(ist):
        if soll[i][0] == ist[j][0] and gleich(soll[i][1], ist[j][1]):
            i += 1
            j += 1
        elif soll[i][2]:
            i += 1
        else:
            return False
    while i < len(soll) and soll[i][2]:
        i += 1
    return i == len(soll) and j == len(ist)


def _bereich(o: list[_Z], neu_sicht: bool, start: int, ende: int) -> tuple[int, int] | None:
    """Zeichenbereich [start, ende) der Sicht -> Indizes in der Zeichenliste der Zelle (None = passt nicht)."""
    zaehler: list[int] = [i for i, (_, k) in enumerate(o) if not (neu_sicht and _strich(k))]
    n = len(zaehler)
    if not 0 <= start <= ende <= n:
        return None
    i0 = zaehler[start] if start < n else len(o)
    i1 = (zaehler[ende - 1] + 1) if ende > start else i0
    return i0, i1


def neue_zeilen_ids(neue_zeilen: dict[str, tuple[str, int]] | None) -> set[str]:
    return set(neue_zeilen or ())


def pruefe_text(quelle: str | Path, ziel: str | Path, absaetze: list[Absatz] | None, angewendet: set[str],
                aenderungen: Iterable[Aenderung], zeile_von: Callable[[str, int], int | None],
                ziele: set[tuple[str, int, int]] | None = None, cfg: dict | None = None,
                auto_ids: set[str] | None = None, neue_zeilen: dict[str, tuple[str, int]] | None = None,
                _wb_quelle: Any = None, _wb_ziel: Any = None) -> TextErgebnis:
    """Konsens-Text: Stimmt die geschriebene Datei `ziel` mit dem eigenen Soll aus den PDF-Absaetzen ueberein?

    angewendet: IDs der Aenderungen, die in `ziel` geschrieben wurden. aenderungen: alle Aenderungen (Status, `bestaetigt`).
    zeile_von(blatt, zeile): Zeile in `ziel` zu einer Zeile der `quelle` (None = Zeile entfernt); ziele: Zellen der Quelle
    (blatt, zeile, spalte), die Operationen einer angewendeten Aenderung treffen (sie duerfen sich aendern).
    auto_ids: die automatisch (ohne Freigabe) angewendeten Aenderungen; ohne Angabe aus `aenderungen` (Status uebernommen und
    nicht bestaetigt). Nur sie werden bei einer Abweichung beanstandet.
    neue_zeilen: Aenderungs-ID -> (Blatt, Zeile in `ziel`) der neu eingefuegten Zeile. Fuer neue Absaetze (Typ absatz_neu) in einer
    neuen Zeile bzw. in einer bestehenden Zelle gilt: der neue Text laut PDF (`Aenderung.neu`) muss rot (neu) dastehen."""
    erg = TextErgebnis()
    k = konfig(cfg)
    if not k.get("aktiv", False) or absaetze is None:
        return erg
    tk = k.get("text", {}) or {}
    norm = _Norm(tk)
    mk_cfg = (cfg or config.lade("zuordnung")).get("markierung", {}) or {}
    markiert = str(mk_cfg.get("aenderungen", "keine")) == "farbe"
    neu_rgb = str(mk_cfg.get("neu_farbe", "FFFF0000"))[-6:].upper()
    alt_rgb = str(mk_cfg.get("alt_farbe", "FF00B0F0"))[-6:].upper()
    nach_id = {a.id: a for a in aenderungen}
    if auto_ids is None:
        auto_ids = {i for i in angewendet if i in nach_id and _auto(nach_id[i])}
    erlaubt: set[tuple[str, int, int]] = set(ziele or ())

    # (_wb_quelle / _wb_ziel: schon geladene Mappen, nur fuer Tests, die viele Abwandlungen derselben Datei pruefen)
    wb_q = _wb_quelle if _wb_quelle is not None else load_workbook(str(quelle), rich_text=True)
    wb_z = _wb_ziel if _wb_ziel is not None else load_workbook(str(ziel), rich_text=True)
    try:
        # --- Absaetze mit angewendeten automatischen Aenderungen --------------------------------------------------
        dran = [ab for ab in absaetze if ab.aenderungen & angewendet and ab.blatt in wb_q.sheetnames and ab.blatt in wb_z.sheetnames]
        for ab in dran:
            for z, sp, _, _ in ab.zellen:
                erlaubt.add((ab.blatt, z, sp))
        je_zelle: dict[tuple[str, int, int], list[tuple[int, Absatz, int, int, int]]] = {}
        for n, ab in enumerate(dran):
            for j, (z, sp, a, e) in enumerate(ab.zellen):
                je_zelle.setdefault((ab.blatt, z, sp), []).append((a, ab, j, e, n))
        bereiche: dict[int, list[_Z] | None] = {}       # Absatz-Nr -> Ist-Strom seiner Zellen (Teilstuecke je Zelle)
        schlecht: dict[int, bool] = {}
        unpruefbar: set[int] = set()
        teile: dict[int, dict[int, list[tuple[str, str]]]] = {}
        for (blatt, zeile, spalte), liste in je_zelle.items():
            ab_nrs = {n for _, _, _, _, n in liste}
            mensch = any(dran[n].aenderungen & angewendet - auto_ids for n in ab_nrs)
            if mensch:
                unpruefbar.update(ab_nrs)                 # von einem Menschen freigegeben: kein Soll aus dem PDF
                continue
            zq = wb_q[blatt]._cells.get((zeile, spalte))
            nz = zeile_von(blatt, zeile)
            zz = None if nz is None else wb_z[blatt]._cells.get((nz, spalte))
            o = _zeichen_der_zelle(zq, neu_rgb, alt_rgb)
            r = _zeichen_der_zelle(zz, neu_rgb, alt_rgb) if nz is not None else None
            if o is None or r is None:
                unpruefbar.update(ab_nrs)
                continue
            neu_sicht = any(_strich(kl) for _, kl in o)
            grenzen = []
            for a, ab, j, e, n in sorted(liste, key=lambda t: (t[0], t[3])):
                g = _bereich(o, neu_sicht, a, e)
                if g is None:
                    schlecht[n] = True
                    break
                grenzen.append((g, ab, j, n))
            else:
                if any(g1[0][1] > g2[0][0] for g1, g2 in zip(grenzen, grenzen[1:])) or \
                        any(_strich(kl) for (i0, i1), *_ in grenzen for _, kl in o[i0:i1]):
                    unpruefbar.update(ab_nrs)             # ueberlappend oder alter Strich mitten im Absatz
                    continue
                rn = norm.strom(r)
                tol = any(kl == NEU for (i0, i1), *_ in grenzen for _, kl in o[i0:i1])
                if all(len(ab.zellen) == 1 for _, ab, _, _ in grenzen):
                    erwartet: list[tuple[str, str, bool]] = []
                    pos = 0
                    for (i0, i1), ab, j, n in grenzen:
                        erwartet += [(c, kl, False) for c, kl in norm.strom(o[pos:i0])]
                        erwartet += _soll_strom(ab.laeufe, ab.aenderungen & angewendet, markiert, norm)
                        pos = i1
                    erwartet += [(c, kl, False) for c, kl in norm.strom(o[pos:])]
                    if not _stuecke_gleich(erwartet, rn, markiert, tol, norm.listen):
                        for _, ab, _, n in grenzen:
                            schlecht[n] = True
                elif len(grenzen) == 1:
                    (i0, i1), ab, j, n = grenzen[0]
                    vor, nach = norm.strom(o[:i0]), norm.strom(o[i1:])
                    if (len(rn) < len(vor) + len(nach) or rn[:len(vor)] != vor or (nach and rn[len(rn) - len(nach):] != nach)):
                        schlecht[n] = True
                    else:
                        teile.setdefault(n, {})[j] = rn[len(vor):len(rn) - len(nach)]
                else:
                    unpruefbar.update(ab_nrs)             # mehrzellige Absaetze in einer Zelle mit weiteren Absaetzen
        for n, ab in enumerate(dran):
            if n in unpruefbar or n in schlecht:
                continue
            if n in teile:                                 # mehrzellig: die Teilstuecke zusammensetzen und einmal vergleichen
                strom = [z for j in sorted(teile[n]) for z in teile[n][j]]
                tol = False
                if not _stuecke_gleich(_soll_strom(ab.laeufe, ab.aenderungen & angewendet, markiert, norm), strom, markiert,
                                       tol, norm.listen):
                    schlecht[n] = True
        for n, ab in enumerate(dran):
            if n in unpruefbar and n not in schlecht:
                erg.unpruefbar += 1
            elif n not in unpruefbar:
                erg.geprueft += 1
            if schlecht.get(n):
                for aid in ab.aenderungen & auto_ids:
                    erg.abweichungen.setdefault(aid, "text")
        # --- neue Absaetze (neue Zeile oder neuer Absatz in einer bestehenden Zelle): der Text laut PDF steht rot da ------------
        if markiert:
            zeilen_von: dict[tuple[str, int], list[Any]] = {}
            for blatt, zeile in {v for v in (neue_zeilen or {}).values()}:
                if blatt in wb_z.sheetnames:
                    zeilen_von[(blatt, zeile)] = [z for (r, _), z in wb_z[blatt]._cells.items() if r == zeile]
            for aid in sorted(auto_ids):
                a = nach_id.get(aid)
                if a is None or a.typ != "absatz_neu" or not a.neu.strip():
                    continue
                soll = "".join(t[0] for t in norm.strom([(c,) for c in a.neu]))
                zellen_ = None
                if neue_zeilen and aid in neue_zeilen:
                    zellen_ = zeilen_von.get(neue_zeilen[aid], [])
                    zellen_ = sorted(zellen_, key=lambda z: z.column)
                elif not any(aid in ab.aenderungen for ab in absaetze):
                    cs = []
                    for o in a.ops:
                        nz = zeile_von(o.blatt, o.zeile) if o.art == "einfuegen" and o.blatt in wb_z.sheetnames else None
                        if nz is not None:
                            cs.append(wb_z[o.blatt]._cells.get((nz, o.spalte)))
                    zellen_ = [z for z in cs if z is not None] or None
                if zellen_ is None:
                    continue
                strom = ""
                for z in zellen_:
                    zs = _zeichen_der_zelle(z, neu_rgb, alt_rgb)
                    strom += "".join(t[0] for t in norm.strom([c for c in (zs or []) if c[1] == NEU]))
                if not soll:
                    continue
                erg.geprueft += 1
                if aid in neue_zeilen_ids(neue_zeilen):
                    gut = strom == soll or any(soll[k:] == strom for k in range(1, min(7, len(soll))))   # Listenzeichen: eigene Spalte
                else:
                    gut = soll in strom or any(soll[k:] in strom for k in range(1, min(7, len(soll))))
                if not gut:
                    erg.abweichungen.setdefault(aid, "text")
        # --- Zellen, die sich nicht aendern sollten --------------------------------------------------------------
        fremd = 0
        fremd_blaetter: set[str] = set()
        for blatt in wb_q.sheetnames:
            fremd_vorher = fremd
            if blatt not in wb_z.sheetnames:
                fremd += 1
                fremd_blaetter.add(blatt)
                continue
            wq, wz = wb_q[blatt], wb_z[blatt]
            abbild: dict[int, int] = {}
            for (zeile, _) in {key for key in wq._cells}:
                if zeile not in abbild:
                    nz = zeile_von(blatt, zeile)
                    if nz is not None:
                        abbild[zeile] = nz
            zurueck = {v: kk for kk, v in abbild.items()}
            gesehen: set[tuple[int, int]] = set()
            for (zeile, spalte), zq in wq._cells.items():
                if (blatt, zeile, spalte) in erlaubt:
                    continue
                gesehen.add((zeile, spalte))
                nz = abbild.get(zeile)
                if nz is None:
                    if zq.value is not None:
                        fremd += 1                         # Zeile mit Inhalt ist weg
                    continue
                zz = wz._cells.get((nz, spalte))
                a, b = _kanon(zq), _kanon(zz)
                if a != b and "formel" not in (a, b):
                    fremd += 1
            for (zeile, spalte), zz in wz._cells.items():
                orig = zurueck.get(zeile)
                if orig is None or (orig, spalte) in gesehen or (blatt, orig, spalte) in erlaubt:
                    continue                               # neue Zeile oder schon verglichen
                if zz.value is not None:
                    fremd += 1                             # Zelle hat jetzt Inhalt, hatte keinen
            if fremd > fremd_vorher:
                fremd_blaetter.add(blatt)
        erg.fremd = fremd
        if fremd:                                          # schuld sind die Aenderungen in den betroffenen Blaettern
            for aid in auto_ids:
                a = nach_id.get(aid)
                blaetter = {o.blatt for o in (a.ops if a is not None else [])} | {ab.blatt for ab in absaetze if aid in ab.aenderungen}
                if not blaetter or blaetter & fremd_blaetter:
                    erg.abweichungen.setdefault(aid, "fremd")
        erg.gelaufen = True
    finally:
        for w, fremd_geladen in ((wb_q, _wb_quelle is not None), (wb_z, _wb_ziel is not None)):
            if fremd_geladen:
                continue
            try:
                w.close()
            except Exception:                              # pragma: no cover
                pass
    return erg
