"""Steckbrief eines Beispiel-Paars (PDF mit Aenderungen + Excel vorher + Excel nachher): NUR Zahlen, Farbcodes und Formen.

Wozu: Echte Beispiele darf nur eine lokale KI sehen. Dieses Werkzeug laeuft lokal und beschreibt, WIE ein Paar aussieht
(Farben, Striche, Aufbau, Formatierung der Aenderungen in der Excel), ohne je ein Wort aus den Dokumenten auszugeben.
Aus der Datei lassen sich spaeter erfundene Nachbau-Paare bauen (ADR 0005: der Steckbrief bestaetigt, welche Farben und
welchen Trenner die echte CDE benutzt).

Datenschutz (Regel 16/17) -- Positivliste statt Verbotsliste:
  * Alles, was in `steckbrief.json` steht, muss zu `SCHEMA` passen: Zahlen, Wahrheitswerte, Hex-Farbcodes, feste Aufzaehlungen,
    Schluessel aus engen Mustern und FORMEN. Eine Form (`schablone`) kennt nur die Zeichen aus `FORM_ZEICHEN`: Grossbuchstabe = A,
    Kleinbuchstabe = a, Ziffer = 9, Leerraum bleibt sichtbar, Satzzeichen bleiben, alles andere ist ?.
  * `pruefe(daten)` prueft jedes Feld und jeden Schluessel; der Schreiber ruft sie VOR dem Schreiben (bei Verstoss wird nichts geschrieben),
    und `steckbrief --pruefe DATEI.json` prueft eine vorhandene Datei erneut.
  * Es gibt keine Datei-, Blatt-, Spaltenkopf-, Zellinhalte, keine Autoren und keine Metadaten.

Aufruf (siehe docs/anleitung_steckbrief.md):
  python -m src.cli steckbrief --pdf A.pdf --vorher V.xlsx --nachher N.xlsx [--name NAME] [--ausgabe ORDNER]
  python -m src.cli steckbrief --ordner PFAD_MIT_PAAREN [--name NAME]      (je Unterordner ein Steckbrief + zusammenfassung.json)
  python -m src.cli steckbrief --pruefe DATEI.json
"""
from __future__ import annotations

import argparse
import colorsys
import json
import math
import re
import sys
import zipfile
from bisect import bisect_left, bisect_right
from collections import Counter
from difflib import SequenceMatcher
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from . import config, markierung
from .markierung import ALT, NEU, NORMAL

STECKBRIEF_FORMAT = 1
MAX_BEISPIELE = 40               # Schablonen-Beispiele je Steckbrief
FENSTER_WOERTER = 12             # hoechstens so viele Woerter Umgebung je Form
FENSTER_DAVOR = 3                # davon vor der Aenderung
KURZ_AB = 30                     # Woerter mit so vielen Zeichen werden zu W… gekuerzt
FORM_MAX = 400                   # laengste Form
NAME_RE = re.compile(r"^[a-z0-9_-]{1,40}$")


# ======================================================================================
# 1. Formen (Schablonen)
# ======================================================================================
_LEERRAUM = {" ": " ", "\u00a0": "·", "\u202f": "·", "\t": "⇥", "\n": "¶"}
_PUNKT = frozenset(".,;:!?()[]{}<>\"'-–—/\\*+=%&#@_|~^$…«»„“”‚‘’")
# Der Zeichensatz einer Form (alles andere in einer Form ist ein Fehler): A a W 9, Leerzeichen, · (geschuetztes Leerzeichen),
# ⇥ (Tab), ¶ (Zeilenumbruch), … (gekuerzt) und die Satzzeichen oben. ? kommt als Satzzeichen UND als "anderes Zeichen" vor.
FORM_ZEICHEN = frozenset("AaW9 ·⇥¶…") | _PUNKT
_WS = " \t\n\r\u00a0\u202f"
_TOKEN = re.compile(r"[ \t\n\r\u00a0\u202f]+|[^ \t\n\r\u00a0\u202f]+")
_KUERZEL = {NORMAL: "-", ALT: "a", NEU: "n"}


def _nl(text: Any) -> str:
    return "" if text is None else str(text).replace("\r\n", "\n").replace("\r", "\n")


def _zeichen_form(ch: str) -> str:
    if ch in _LEERRAUM:
        return _LEERRAUM[ch]
    if "A" <= ch <= "Z" or ch in "ÄÖÜ":
        return "A"
    if "a" <= ch <= "z" or ch in "äöüß":
        return "a"
    if "0" <= ch <= "9":
        return "9"
    if ch in _PUNKT:
        return ch
    return "?"


def schablone(text: Any, max_woerter: int = 0) -> str:
    """Die FORM eines Textes: Grossbuchstabe -> A, Kleinbuchstabe -> a, Ziffer -> 9, Umlaute und ß wie Buchstaben,
    Leerzeichen bleibt ' ', geschuetztes/schmales Leerzeichen -> ·, Tab -> ⇥, Zeilenumbruch -> ¶, Satzzeichen bleiben, jedes andere
    Zeichen -> ?. Woerter ab 30 Zeichen werden zu W… gekuerzt. max_woerter > 0: hoechstens so viele Woerter (dann …)."""
    n = 0
    teile: list[str] = []
    for tok in _TOKEN.findall(_nl(text)):
        if tok[0] in _WS:
            teile.append("".join(_zeichen_form(c) for c in tok))
            continue
        n += 1
        if max_woerter and n > max_woerter:
            teile.append("…")
            break
        teile.append("W…" if len(tok) >= KURZ_AB else "".join(_zeichen_form(c) for c in tok))
    s = "".join(teile)
    return s if len(s) <= FORM_MAX else s[:FORM_MAX - 1] + "…"


def fenster(stuecke: Iterable[tuple[str, str]], start: int = 0, max_woerter: int = FENSTER_WOERTER,
            davor: int = FENSTER_DAVOR) -> list[tuple[str, str]]:
    """Die Umgebung einer Stelle als Formen. `stuecke` = [(Format 'normal'|'alt'|'neu', Text)] (Zeichenposition `start` zaehlt im
    verbundenen Text, Zeilenenden vereinheitlicht). Ergebnis: [(Kuerzel '-'|'a'|'n', Form)] mit hoechstens `max_woerter` Woertern,
    beginnend `davor` Woerter vor der Stelle. Ist der Text an einer Seite gekuerzt, steht dort ein …."""
    teile = [(f, _nl(t)) for f, t in stuecke if t]
    text = "".join(t for _, t in teile)
    fmts: list[str] = []
    for f, t in teile:
        fmts.extend([f] * len(t))
    toks = [(m.start(), m.end()) for m in _TOKEN.finditer(text)]
    woerter = [i for i, (a, b) in enumerate(toks) if text[a] not in _WS]
    if not woerter:
        return [(_KUERZEL.get(fmts[0], "-"), schablone(text))] if text else []
    start = max(0, min(start, len(text)))
    k = next((j for j, i in enumerate(woerter) if toks[i][1] > start), len(woerter) - 1)
    erst = max(0, k - davor)
    letzt = min(len(woerter), erst + max_woerter)
    erst = max(0, letzt - max_woerter) if letzt == len(woerter) else erst
    a, b = toks[woerter[erst]][0], toks[woerter[letzt - 1]][1]
    laeufe: list[list[str]] = []                     # [Format, Form]
    for ta, tb in toks:
        if tb <= a or ta >= b:
            continue
        wort = text[ta:tb]
        if wort[0] not in _WS and len(wort) >= KURZ_AB:
            zeichen = [(fmts[ta], "W…")]                     # ein langes Wort: nur sein Anfangsformat zaehlt
        else:
            zeichen = [(fmts[p], _zeichen_form(text[p])) for p in range(ta, tb)]
        for f, c in zeichen:
            if laeufe and laeufe[-1][0] == f:
                laeufe[-1][1] += c
            else:
                laeufe.append([f, c])
    if laeufe and a > 0:
        laeufe[0][1] = "…" + laeufe[0][1]
    if laeufe and b < len(text):
        laeufe[-1][1] += "…"
    return [(_KUERZEL.get(f, "-"), form[:FORM_MAX]) for f, form in laeufe]


# ======================================================================================
# 2. Positivliste: das Schema
# ======================================================================================
class _Spec:
    def pruefe(self, w: Any, pfad: str, fehler: list[str]) -> None:       # pragma: no cover - Basis
        raise NotImplementedError


def _f(fehler: list[str], pfad: str, grund: str) -> None:
    fehler.append(f"{pfad or '(oben)'}: {grund}")


class Ganz(_Spec):
    def __init__(self, lo: int = 0, hi: int = 10 ** 9) -> None:
        self.lo, self.hi = lo, hi

    def pruefe(self, w, pfad, fehler):
        if isinstance(w, bool) or not isinstance(w, int):
            _f(fehler, pfad, "ganze Zahl erwartet")
        elif not self.lo <= w <= self.hi:
            _f(fehler, pfad, "Zahl außerhalb des erlaubten Bereichs")


class Dezimal(_Spec):
    def __init__(self, lo: float = -1e9, hi: float = 1e9) -> None:
        self.lo, self.hi = lo, hi

    def pruefe(self, w, pfad, fehler):
        if isinstance(w, bool) or not isinstance(w, (int, float)) or (isinstance(w, float) and not math.isfinite(w)):
            _f(fehler, pfad, "Zahl erwartet")
        elif not self.lo <= w <= self.hi:
            _f(fehler, pfad, "Zahl außerhalb des erlaubten Bereichs")


class Wahr(_Spec):
    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, bool):
            _f(fehler, pfad, "Ja/Nein-Wert erwartet")


class Muster(_Spec):
    """Text, der als GANZES zu einem engen Muster passen muss (und höchstens `maxlen` Zeichen lang ist)."""

    def __init__(self, regex: str, maxlen: int = 60) -> None:
        self.re, self.maxlen = re.compile(regex), maxlen

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, str):
            _f(fehler, pfad, "Text erwartet")
        elif len(w) > self.maxlen:
            _f(fehler, pfad, "Text zu lang")
        elif not self.re.fullmatch(w):
            _f(fehler, pfad, "Text passt nicht zum erlaubten Muster")


class Auswahl(_Spec):
    def __init__(self, *werte: str) -> None:
        self.werte = frozenset(werte)

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, str) or w not in self.werte:
            _f(fehler, pfad, "Wert nicht in der erlaubten Aufzählung")


class Gestalt(_Spec):
    """Eine Form (siehe `schablone`): nur Zeichen aus FORM_ZEICHEN, höchstens `maxlen` Zeichen."""

    def __init__(self, maxlen: int = FORM_MAX) -> None:
        self.maxlen = maxlen

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, str):
            _f(fehler, pfad, "Form (Text) erwartet")
        elif len(w) > self.maxlen:
            _f(fehler, pfad, "Form zu lang")
        elif any(c not in FORM_ZEICHEN for c in w):
            _f(fehler, pfad, "Zeichen außerhalb des Formen-Zeichensatzes")


_UPPER = re.compile(r"[A-Z]+")
_REL = re.compile(r"\{r(?:[+-]\d{1,6})?\}")
FUNKTIONEN = frozenset((
    "SUM SUMME COUNT ANZAHL COUNTA ANZAHL2 COUNTIF ZAEHLENWENN COUNTIFS ZAEHLENWENNS SUMIF SUMMEWENN SUMIFS SUMMEWENNS AVERAGE "
    "MITTELWERT MIN MAX IF WENN IFERROR WENNFEHLER IFS AND UND OR ODER NOT NICHT LEN LAENGE LEFT LINKS RIGHT RECHTS MID TEIL "
    "CONCATENATE VERKETTEN CONCAT TEXTJOIN TEXTKETTE TRIM GLAETTEN UPPER GROSS LOWER KLEIN PROPER GROSS2 ROUND RUNDEN ROUNDUP "
    "AUFRUNDEN ROUNDDOWN ABRUNDEN ROW ZEILE COLUMN SPALTE ROWS ZEILEN COLUMNS SPALTEN VLOOKUP SVERWEIS HLOOKUP WVERWEIS XLOOKUP "
    "XVERWEIS LOOKUP VERWEIS INDEX MATCH VERGLEICH INDIRECT INDIREKT OFFSET BEREICH.VERSCHIEBEN SUBTOTAL TEILERGEBNIS AGGREGATE "
    "AGGREGAT TEXT TODAY HEUTE NOW JETZT DATE DATUM YEAR JAHR MONTH MONAT DAY TAG ISBLANK ISTLEER ISNUMBER ISTZAHL ISTEXT ISTTEXT "
    "ISERROR ISTFEHLER VALUE WERT SUBSTITUTE WECHSELN FIND FINDEN SEARCH SUCHEN REPT WIEDERHOLEN COUNTBLANK ANZAHLLEEREZELLEN ABS "
    "INT GANZZAHL MOD REST TRUNC KUERZEN CEILING OBERGRENZE FLOOR UNTERGRENZE CHAR ZEICHEN CODE UNICODE ROMAN ROEMISCH ARABIC "
    "ARABISCH SUMPRODUCT SUMMENPRODUKT LET LAMBDA FILTER SORT UNIQUE SEQUENCE CHOOSE WAHL TRUE FALSE WAHR FALSCH").split())
_FORMEL_ERLAUBT = ("NAME", "FUNC", "TEXT")


class Formel(_Spec):
    """Formel-MUSTER (`=SUM(D{r+1}:D{r+3})`): nur Großbuchstaben-Folgen, die Spaltenbuchstaben (höchstens 3) oder bekannte
    Funktionsnamen sind. Ein Wort aus einem Dokument kann so nicht hineinrutschen."""

    def __init__(self, maxlen: int = 120) -> None:
        self.maxlen = maxlen

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, str):
            _f(fehler, pfad, "Formelmuster (Text) erwartet")
            return
        if len(w) > self.maxlen:
            _f(fehler, pfad, "Formelmuster zu lang")
            return
        if not w.startswith("="):
            _f(fehler, pfad, "Formelmuster muss mit = beginnen")
            return
        rest = _REL.sub("", w[1:])
        if re.search(r"[^A-Z0-9(),:;+\-*/<>&^%$=\" .!_#@|~\[\]{}?]", rest):
            _f(fehler, pfad, "Zeichen außerhalb des Formelmuster-Zeichensatzes")
            return
        for lauf in _UPPER.findall(rest):
            if len(lauf) > 3 and lauf not in FUNKTIONEN and lauf not in _FORMEL_ERLAUBT:
                _f(fehler, pfad, "unbekanntes Wort im Formelmuster")
                return


class Objekt(_Spec):
    def __init__(self, felder: dict[str, _Spec], pflicht: tuple[str, ...] = ()) -> None:
        self.felder, self.pflicht = felder, pflicht

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, dict):
            _f(fehler, pfad, "Objekt erwartet")
            return
        for name in self.pflicht:
            if name not in w:
                _f(fehler, pfad, f"Pflichtfeld '{name}' fehlt")
        for k, v in w.items():
            spec = self.felder.get(k) if isinstance(k, str) else None
            if spec is None:                                    # der Name selbst wird nicht wiederholt (koennte Text sein)
                _f(fehler, pfad, "unbekanntes Feld")
                continue
            spec.pruefe(v, f"{pfad}.{k}" if pfad else k, fehler)


class Zaehl(_Spec):
    """Objekt mit beliebig vielen Schlüsseln (je ein enges Muster) und gleichartigen Werten (Häufigkeitstabelle)."""

    def __init__(self, schluessel: _Spec, wert: _Spec, maxn: int = 300) -> None:
        self.schluessel, self.wert, self.maxn = schluessel, wert, maxn

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, dict):
            _f(fehler, pfad, "Objekt erwartet")
            return
        if len(w) > self.maxn:
            _f(fehler, pfad, "zu viele Einträge")
        for k, v in w.items():
            self.schluessel.pruefe(k, f"{pfad}{{*}}", fehler)
            self.wert.pruefe(v, f"{pfad}{{*}}", fehler)


class Folge(_Spec):
    def __init__(self, elem: _Spec, maxn: int = 500) -> None:
        self.elem, self.maxn = elem, maxn

    def pruefe(self, w, pfad, fehler):
        if not isinstance(w, list):
            _f(fehler, pfad, "Liste erwartet")
            return
        if len(w) > self.maxn:
            _f(fehler, pfad, "zu viele Einträge")
        for x in w:
            self.elem.pruefe(x, f"{pfad}[*]", fehler)


# ---- feste Aufzaehlungen und enge Muster -------------------------------------------------------------
Z = Ganz()
KLASSEN = ("schwarz", "weiss", "grau", "rot", "orange", "gelb", "gruen", "tuerkis", "blau", "violett", "standard", "andere",
           "unklar")
FARBRAEUME = ("DeviceRGB", "DeviceGray", "DeviceCMYK", "ICCBased", "CalRGB", "CalGray", "Lab", "Indexed", "Separation",
              "DeviceN", "Pattern", "anderer")
ANNOTATIONEN = ("Link", "Highlight", "StrikeOut", "Underline", "Squiggly", "Text", "FreeText", "Stamp", "Ink", "Popup",
                "Widget", "Square", "Circle", "Line", "Caret", "FileAttachment", "Polygon", "PolyLine", "Redact", "andere")
RICHTUNGEN = ("h", "um", "ab", "auf", "schraeg")
LUECKEN = ("scan_seite", "mehrspaltig", "tabelle", "bild", "unbekannte_farbe", "nicht_lesbarer_text", "unerklaerte_markierung",
           "sonstiges")
TYPEN = ("text_ersetzt", "text_eingefuegt", "text_geloescht", "absatz_geloescht", "section_geloescht", "section_neu",
         "absatz_neu", "andere")
STATUS = ("uebernommen", "pruefen", "nicht_uebernommen")
AUFZAEHLUNG = ("punkt", "strich", "stern", "zahl", "buchstabe", "roemisch")
_ALLE_FORMZEICHEN = "".join(sorted(re.escape(c) for c in FORM_ZEICHEN))

CODE = Muster(r"^[a-z0-9_:.\-]{1,60}$", 60)
ZAHLKEY = Muster(r"^-?\d{1,7}(?:\.\d{1,2})?$", 12)
INTKEY = Muster(r"^\d{1,7}$", 7)
SPALTE = Muster(r"^[A-Z]{1,3}$", 3)
MASS = Muster(r"^\d{1,5}x\d{1,5}$", 11)
SCHRIFT = Muster(r"^[A-Za-z0-9\-+ ]{1,40}$", 40)
FARBCODE = Muster(r"^(?:#?[0-9A-Fa-f]{6,8}|theme_\d{1,2}|index_\d{1,3}|auto|keine|unklar)$", 12)
FARBEN_KEY = Muster(r"^(?:" + "|".join(FARBRAEUME) + r")\|(?:#[0-9a-f]{6}|unklar)$", 40)
LEGENDE_KEY = Muster(r"^(?:" + "|".join(KLASSEN) + r")\|(?:strich|-)\|(?:unterstrichen|-)\|(?:hinterlegt|-)$", 60)
STUECK_MUSTER = Muster(r"^(?:weitere|(?:normal|alt|neu)(?:,(?:normal|alt|neu)){0,11}(?:,mehr)?)$", 80)
TRENNER = Muster(r"^(?:kein|[" + _ALLE_FORMZEICHEN + r"]{1,10})$", 12)
ZAHLFORMAT = Muster(r"^(?:General|Standard|@|anderes|[0-9#.,%/;\- +():_*?]{1,24})$", 24)
FARBE_STRICH = Muster(r"^(?:" + "|".join(KLASSEN) + r")\|(?:strich|-)$", 30)
NAME = Muster(NAME_RE.pattern, 40)
PAARNAME = NAME
FORM = Gestalt()

_STRICHE = Objekt({
    "anzahl": Z, "laenge_zur_seitenbreite": Zaehl(INTKEY, Z), "laenge_zum_wort": Zaehl(INTKEY, Z),
    "hoehenband": Zaehl(ZAHLKEY, Z), "staerke": Zaehl(ZAHLKEY, Z), "farben": Zaehl(FARBCODE, Z),
    "farbe_wie_text": Z, "farbe_anders": Z})

PDF_SCHEMA = Objekt({
    "seiten": Z, "seiten_masse": Zaehl(MASS, Z), "gedrehte_seiten": Z, "seiten_ohne_text": Z,
    "seite": Folge(Objekt({"zeichen": Z, "woerter": Z, "bilder": Z}), 3000),
    "schriften": Zaehl(SCHRIFT, Z), "schriftgroessen": Zaehl(ZAHLKEY, Z),
    "textrichtungen": Zaehl(Auswahl(*RICHTUNGEN), Z),
    "farben": Zaehl(FARBEN_KEY, Objekt({"woerter": Z, "durchgestrichen": Z, "unterstrichen": Z, "hinterlegt": Z})),
    "hinterlegungen": Zaehl(FARBCODE, Z), "legende": Zaehl(LEGENDE_KEY, Z),
    "striche": _STRICHE, "unterstreichungen": _STRICHE,
    "annotationen": Zaehl(Auswahl(*ANNOTATIONEN), Z), "hyperlinks": Z,
    "luecken": Zaehl(Auswahl(*LUECKEN), Objekt({"anzahl": Z, "woerter": Z})),
    "kopf_fuss": Objekt({"eintraege": Z, "woerter": Z}),
    "silbentrennung": Z, "aufzaehlungszeichen": Zaehl(Auswahl(*AUFZAEHLUNG), Z), "hochzahlen": Z, "bilder": Z,
    "geometrie_nicht_gemessen_seiten": Z,
})

TOOL_SCHEMA = Objekt({
    "lauf_ok": Wahr(), "aenderungen": Z, "je_typ": Zaehl(Auswahl(*TYPEN), Z), "je_status": Zaehl(Auswahl(*STATUS), Z),
    "hinweise": Zaehl(CODE, Z), "pruefcodes": Zaehl(CODE, Z),
    "woerter_status": Zaehl(Auswahl("keep", "del", "add", "mod"), Z), "warnungen": Z, "markierte_woerter": Z})

_BLATT = Objekt({
    "nr": Z, "sichtbarkeit": Auswahl("sichtbar", "versteckt", "sehr_versteckt"), "zeilen": Z, "spalten": Z, "zellen": Z,
    "rollen": Zaehl(Auswahl("section", "titel", "inhalt"), SPALTE),
    "modell": Auswahl("zeile_pro_section", "zeile_pro_absatz"), "sicherheit": Auswahl("hoch", "mittel", "niedrig"),
    "kombiniert": Wahr(), "sections": Z,
    "zeilentypen": Zaehl(Auswahl("kopf", "inhalt", "sonstige", "leer"), Z),
    "versteckte_zeilen": Z, "versteckte_spalten": Z, "verbunden": Z, "formeln": Z})

EXCEL_VORHER_SCHEMA = Objekt({
    "blaetter": Z, "blaetter_versteckt": Z, "blatt": Folge(_BLATT, 300),
    "absaetze_je_zelle": Zaehl(INTKEY, Z), "woerter_je_absatz": Zaehl(INTKEY, Z), "zellen_mit_mehreren_absaetzen": Z,
    "formeln": Objekt({"anzahl": Z, "muster": Zaehl(Formel(), Z), "funktionen": Zaehl(Auswahl(*sorted(FUNKTIONEN), "FUNC"), Z),
                       "zu_anderem_blatt": Z}),
    "konstanten_spalten": Folge(Objekt({"blatt": Z, "spalte": SPALTE, "anteil": Dezimal(0, 1), "zellen": Z, "form": FORM}), 200),
    "verbundene_zellen": Zaehl(MASS, Z), "versteckte_zeilen": Z, "versteckte_spalten": Z,
    "datenpruefung": Z, "datenpruefung_typen": Zaehl(Auswahl("list", "whole", "decimal", "date", "time", "textLength", "custom",
                                                           "any", "andere"), Z),
    "bedingte_formatierung": Z, "filter": Z, "tabellen": Z, "namen": Z,
    "section_zahlenformat": Zaehl(ZAHLFORMAT, Z), "section_nummern": Zaehl(Auswahl("zahl", "text"), Z),
    "rich_text_zellen": Z, "inhalt_schriften": Zaehl(SCHRIFT, Z), "inhalt_groessen": Zaehl(ZAHLKEY, Z),
    "hyperlinks": Z, "kommentare": Z, "bilder": Z, "diagramme": Z})

_FORMAT3 = Auswahl("normal", "alt", "neu")
EXCEL_NACHHER_SCHEMA = Objekt({
    "blaetter_vorher": Z, "blaetter_nachher": Z, "blaetter_gleich": Wahr(),
    "zeilen": Objekt({"vorher": Z, "nachher": Z, "gepaart": Z, "neu": Z, "entfernt": Z, "gestrichen_markiert": Z,
                      "inhalt_geleert": Z}),
    "zellen": Objekt({"gleich": Z, "mit_markierung": Z, "markiert_text_geaendert": Z, "markiert_nur_format": Z,
                      "sauber_ersetzt": Z, "sauber_ergaenzt": Z, "sauber_geleert": Z, "ganze_zelle_markiert": Z,
                      "alt_ansicht_stimmt": Zaehl(Auswahl("ja", "nur_leerraum", "nein"), Z), "zahl_formel_geaendert": Z}),
    "zahl_formel_markiert": Zaehl(_FORMAT3, Z),
    "muster": Zaehl(STUECK_MUSTER, Z), "stueck_farben": Zaehl(_FORMAT3, Zaehl(FARBCODE, Z)),
    "farbe_strich": Zaehl(FARBE_STRICH, Z),
    "strich_anteil": Zaehl(Auswahl("rot", "blau", "andere_farbe", "schwarz"), Dezimal(0, 1)),
    "trenner": Zaehl(TRENNER, Z), "trenner_ort": Zaehl(Auswahl("eigenes_stueck", "im_alt", "im_neu", "kein", "beides"), Z),
    "reihenfolge": Zaehl(Auswahl("alt_vor_neu", "neu_vor_alt"), Z), "nur_alt": Z, "nur_neu": Z,
    "geaenderte_rollen": Zaehl(Auswahl("section", "titel", "inhalt", "andere"), Z), "geaenderte_spalten": Zaehl(SPALTE, Z),
    "formatierung": Objekt({"zellen_andere_schriftfarbe": Z, "schriftfarben": Zaehl(FARBCODE, Z),
                            "zellen_neu_gefuellt": Z, "fuellfarben": Zaehl(FARBCODE, Z),
                            "stuecke_andere_farbe": Zaehl(FARBCODE, Z), "stuecke_fett": Z, "stuecke_kursiv": Z,
                            "stuecke_unterstrichen": Z}),
    "neue_zeilen": Objekt({"anzahl": Z, "zellfarben": Zaehl(FARBCODE, Z), "spalten_befuellt": Zaehl(SPALTE, Z),
                           "muster": Zaehl(STUECK_MUSTER, Z),
                           "alle_stuecke_neu": Z, "konstanten_uebernommen": Z, "konstanten_fehlen": Z, "formeln": Z}),
    "gestrichene_zeilen": Objekt({"anzahl": Z, "zellfarben": Zaehl(FARBCODE, Z), "alle_durchgestrichen": Z,
                                  "inhalt_noch_da": Z}),
    "masse": Objekt({"zeilenhoehe_geaendert": Z, "spaltenbreite_geaendert": Z, "zeilenhoehe_geaendert_ja": Wahr(),
                     "spaltenbreite_geaendert_ja": Wahr()}),
})

_LAUF = Objekt({"f": Auswahl("-", "a", "n"), "form": FORM}, pflicht=("f", "form"))
SCHABLONEN_SCHEMA = Folge(Objekt({
    "typ": Auswahl(*TYPEN), "status": Auswahl(*STATUS), "pdf_alt": Folge(_LAUF, 60), "pdf_neu": Folge(_LAUF, 60),
    "excel_vorher": Folge(_LAUF, 60), "excel_nachher": Folge(_LAUF, 60)}), MAX_BEISPIELE)

SCHEMA = Objekt({
    "art": Auswahl("steckbrief"), "format": Ganz(1, 99), "name": NAME,
    "pdf": PDF_SCHEMA, "excel_vorher": EXCEL_VORHER_SCHEMA, "excel_nachher": EXCEL_NACHHER_SCHEMA,
    "tool": TOOL_SCHEMA, "schablonen": SCHABLONEN_SCHEMA}, pflicht=("art", "format"))

# Fehler-Schablonen von `messen --schablonen`: je Eintrag Fehlerart, Aenderungstyp, Erzeuger und die Formen Soll/Ist
FEHLER_SCHABLONE = Objekt({
    "quelle": Auswahl("stiller_fehler", "falscher_vorschlag"), "fehlerart": Auswahl(
        "pdf_falsch_gelesen", "layout_falsch", "zuordnung_falsch", "plan_falsch", "bilanz_kaputt", "unklar", "ohne_art"),
    "aenderungstyp": Auswahl(*TYPEN, "ohne_typ"), "erzeuger": Auswahl("writer", "libreoffice", "sonstiger", "unbekannt"),
    "soll": Folge(_LAUF, 60), "ist": Folge(_LAUF, 60)}, pflicht=("quelle", "soll", "ist"))
FEHLER_SCHABLONEN_SCHEMA = Folge(FEHLER_SCHABLONE, MAX_BEISPIELE)


# ---- Zusammenfassung ueber mehrere Paare (aus dem Schema abgeleitet) ------------------------------------------------
def _skalar(spec: _Spec) -> bool:
    return isinstance(spec, (Muster, Auswahl, Gestalt, Formel))


_GROSS = 10 ** 12


def _summen_spec(spec: _Spec) -> _Spec | None:
    """Das Schema der Zusammenfassung: Zahlen werden addiert (Wahr -> Anzahl der Paare, Dezimal -> Mittelwert), Text-Werte
    (Aufzaehlung, Form ...) werden zu Haeufigkeitstabellen; Listen von Objekten (z. B. Schablonen, Blaetter) entfallen."""
    if isinstance(spec, (Ganz, Wahr)):
        return Ganz(0, _GROSS)
    if isinstance(spec, Dezimal):
        return Dezimal(-_GROSS, _GROSS)
    if _skalar(spec):
        return Zaehl(spec, Ganz(0, _GROSS), 5000)
    if isinstance(spec, Objekt):
        felder = {k: s for k, s in ((k, _summen_spec(v)) for k, v in spec.felder.items()) if s is not None}
        return Objekt(felder)
    if isinstance(spec, Zaehl):
        w = _summen_spec(spec.wert)
        return Zaehl(spec.schluessel, w, spec.maxn * 50) if w is not None else None
    if isinstance(spec, Folge) and _skalar(spec.elem):
        return Zaehl(spec.elem, Ganz(0, _GROSS), 5000)
    return None


def _summe(spec: _Spec, werte: list[Any]) -> Any:
    werte = [w for w in werte if w is not None]
    if not werte:
        return None
    if isinstance(spec, Ganz):
        return sum(werte)
    if isinstance(spec, Wahr):
        return sum(1 for w in werte if w)
    if isinstance(spec, Dezimal):
        return round(sum(werte) / len(werte), 4)
    if _skalar(spec):
        return dict(sorted(Counter(werte).items()))
    if isinstance(spec, Objekt):
        out: dict[str, Any] = {}
        for k, s in spec.felder.items():
            if _summen_spec(s) is None:
                continue
            r = _summe(s, [w[k] for w in werte if isinstance(w, dict) and k in w])
            if r is not None:
                out[k] = r
        return out
    if isinstance(spec, Zaehl):
        schluessel = sorted({k for w in werte if isinstance(w, dict) for k in w})
        return {k: _summe(spec.wert, [w[k] for w in werte if isinstance(w, dict) and k in w]) for k in schluessel}
    if isinstance(spec, Folge) and _skalar(spec.elem):
        return dict(sorted(Counter(x for w in werte if isinstance(w, list) for x in w).items()))
    return None


_TEILE = ("pdf", "excel_vorher", "excel_nachher", "tool")
SCHEMA_ZUSAMMENFASSUNG = Objekt({
    "art": Auswahl("zusammenfassung"), "format": Ganz(1, 99), "paare": Z,
    **{k: _summen_spec(SCHEMA.felder[k]) for k in _TEILE}}, pflicht=("art", "format", "paare"))                # type: ignore[arg-type]


def fasse_zusammen(steckbriefe: list[dict[str, Any]]) -> dict[str, Any]:
    """Summen und Haeufigkeiten ueber mehrere Steckbriefe (nur Zahlen, Farbcodes, Formen)."""
    out: dict[str, Any] = {"art": "zusammenfassung", "format": STECKBRIEF_FORMAT, "paare": len(steckbriefe)}
    for k in _TEILE:
        r = _summe(SCHEMA.felder[k], [s.get(k) for s in steckbriefe])
        if r is not None:
            out[k] = r
    return out


def pruefe(daten: Any) -> list[str]:
    """Prueft ein Ergebnis (Steckbrief oder Zusammenfassung) gegen die Positivliste. -> Liste der Verstoesse (leer = alles in Ordnung).
    Die Meldungen nennen nur den Ort im Schema (Pfad; {*} = ein Schluessel), nie den Inhalt."""
    fehler: list[str] = []
    if not isinstance(daten, dict):
        return ["(oben): Objekt erwartet"]
    schema = SCHEMA_ZUSAMMENFASSUNG if daten.get("art") == "zusammenfassung" else SCHEMA
    schema.pruefe(daten, "", fehler)
    gezaehlt = Counter(fehler)
    return [m if n == 1 else f"{m} ({n}x)" for m, n in gezaehlt.items()][:200]


def pruefe_fehler_schablonen(eintraege: Any) -> list[str]:
    """Wie `pruefe`, fuer die Fehler-Schablonen von `messen --schablonen`."""
    fehler: list[str] = []
    FEHLER_SCHABLONEN_SCHEMA.pruefe(eintraege, "fehler_schablonen", fehler)
    return [m if n == 1 else f"{m} ({n}x)" for m, n in Counter(fehler).items()][:200]


# ======================================================================================
# 3. Farben
# ======================================================================================
def _klasse_rgb(r: float, g: float, b: float) -> str:
    h, s, v = colorsys.rgb_to_hsv(max(0.0, min(1.0, r)), max(0.0, min(1.0, g)), max(0.0, min(1.0, b)))
    if s < 0.25:
        return "schwarz" if v < 0.35 else "weiss" if v > 0.92 else "grau"
    grad = h * 360.0
    if grad < 15 or grad >= 345:
        return "rot"
    if grad < 40:
        return "orange"
    if grad < 70:
        return "gelb"
    if grad < 165:
        return "gruen"
    if grad < 200:
        return "tuerkis"
    if grad < 255:
        return "blau"
    return "violett"


def _klasse_hex(code: str) -> str:
    if code.startswith("#") and len(code) in (7, 9):
        try:
            r, g, b = (int(code[-6:][i:i + 2], 16) / 255.0 for i in (0, 2, 4))
        except ValueError:
            return "unklar"
        return _klasse_rgb(r, g, b)
    if code in ("keine", "auto", "theme_0", "theme_1", "index_8", "index_64"):
        return "standard"
    if code == "unklar":
        return "unklar"
    return "andere"


def _hex_pdf(farbe: Any, raum: str) -> str:
    """Hex-Farbcode ('#rrggbb') einer PDF-Farbe, sonst 'unklar' (Sonderfarbraum, Muster)."""
    if farbe is None:
        return "#000000"
    if isinstance(farbe, str):
        return "unklar"
    try:
        w = tuple(float(x) for x in farbe)
    except (TypeError, ValueError):
        return "unklar"
    if len(w) == 1 and raum not in ("Separation", "DeviceN", "Indexed", "Pattern", "Lab"):
        w = (w[0], w[0], w[0])
    if len(w) == 3 and raum != "Lab":
        rgb = w
    elif len(w) == 4:
        c, m, y, k = w
        rgb = ((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    else:
        return "unklar"
    return "#%02x%02x%02x" % tuple(max(0, min(255, round(v * 255))) for v in rgb)


def _raum(roh: Any) -> str:
    s = str(roh or "")
    return s if s in FARBRAEUME else "anderer"


def _farbcode(farbe: Any) -> str:
    """Farbcode einer Excel-Farbe: '#rrggbb' | theme_N | index_N | auto | keine."""
    if farbe is None:
        return "keine"
    typ = getattr(farbe, "type", None)
    try:
        if typ == "rgb":
            rgb = farbe.rgb
            if isinstance(rgb, str) and len(rgb) in (6, 8):
                return "#" + rgb[-6:].lower()
        elif typ == "theme":
            return f"theme_{int(farbe.theme)}"
        elif typ == "indexed":
            return f"index_{int(farbe.indexed)}"
        elif typ == "auto" or getattr(farbe, "auto", None):
            return "auto"
    except (TypeError, ValueError, AttributeError):
        pass
    return "keine"


# ======================================================================================
# 4. Formeln als Muster
# ======================================================================================
_F_TOKEN = re.compile(
    r"""(?P<text>"(?:[^"]|"")*")"""
    r"""|(?P<blatt>(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)"""
    r"""|(?P<bezug>\$?[A-Za-z]{1,3}\$?\d+)(?![\w(])"""
    r"""|(?P<zahl>\d+(?:[.,]\d+)?(?:[eE][+-]?\d+)?%?)"""
    r"""|(?P<name>[A-Za-z_À-ɏ][\w.À-ɏ]*)"""
    r"""|(?P<klammer>[\[{][^\]}]*[\]}])"""
    r"""|(?P<op>[()+\-*/^&=<>,;:%])"""
    r"""|(?P<leer>\s+)"""
    r"""|(?P<rest>.)""", re.S)
_UMLAUT = str.maketrans({"Ä": "AE", "Ö": "OE", "Ü": "UE", "ß": "SS"})


def formel_muster(formel: str, zeile: int) -> tuple[str, list[str]]:
    """Formel -> (Muster, Funktionsnamen). Bezuege werden relativ zur Zeile der Formel geschrieben ({r}, {r+1}, {r-2}), Textkonstanten
    zu "T", Zahlen zu 9, Namen zu NAME, Funktionen nur aus der Whitelist (sonst FUNC). Beispiel: =LEN(C5) in Zeile 5 -> =LEN(C{r})."""
    out: list[str] = []
    funktionen: list[str] = []
    text = str(formel)
    if text.startswith("="):
        text = text[1:]
    pos = 0
    while pos < len(text):
        m = _F_TOKEN.match(text, pos)
        if m is None:                                                       # pragma: no cover - rest faengt alles
            break
        pos = m.end()
        art = m.lastgroup
        s = m.group()
        if art == "text":
            out.append('"T"')
        elif art == "blatt":
            out.append("@!")
        elif art == "bezug":
            sp = re.match(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)", s)
            assert sp is not None
            d1, col, d2, row = sp.groups()
            if d2:
                out.append(f"{d1}{col.upper()}$9")
            else:
                diff = int(row) - zeile
                out.append(f"{d1}{col.upper()}" + ("{r}" if diff == 0 else f"{{r{diff:+d}}}"))
        elif art == "zahl":
            out.append("9")
        elif art == "name":
            nxt = text[pos:pos + 1]
            roh = s.upper().translate(_UMLAUT)
            if nxt == "(":
                fn = roh if roh in FUNKTIONEN else "FUNC"
                funktionen.append(fn)
                out.append(fn)
            else:
                out.append(roh if roh in ("TRUE", "FALSE", "WAHR", "FALSCH") else "NAME")
        elif art == "klammer":
            out.append("[T]" if s[0] == "[" else "{T}")
        elif art == "op":
            out.append(s)
        elif art == "rest":
            out.append("?")
    muster = "=" + "".join(out)
    if len(muster) > 120:
        muster = muster[:117] + "..."
    return muster, funktionen


# ======================================================================================
# 5. PDF messen
# ======================================================================================
_FAMILIE = re.compile(r"^[A-Za-z0-9\-+ ]{1,40}$")
_SUBSET = re.compile(r"^[A-Z]{6}\+")
_PUNKT_ZEICHEN = "•·▪◦‣●○■□\uf0b7\uf0a7\uf0d8"
_HOCH_UNICODE = re.compile("[¹²³⁰⁴-⁹]")


def _richtung(c: dict) -> str:
    m = c.get("matrix") or (1, 0, 0, 1, 0, 0)
    a, b = float(m[0]), float(m[1])
    if abs(b) <= 0.2 * abs(a):
        return "h" if a >= 0 else "um"
    if abs(a) <= 0.2 * abs(b):
        return "ab" if b < 0 else "auf"
    return "schraeg"


BEKANNTE_SCHRIFTEN = (
    "Arial", "Arial Narrow", "Aptos", "Book Antiqua", "Bookman", "Calibri", "Cambria", "Candara", "Century Gothic", "Comic Sans",
    "Consolas", "Courier New", "Courier", "DejaVu Sans", "DejaVu Serif", "Franklin Gothic", "Garamond", "Georgia", "Helvetica",
    "Impact", "Lato", "Liberation Sans", "Liberation Serif", "Liberation Mono", "Lucida Console", "Lucida Sans", "Montserrat",
    "MS Gothic", "Noto Sans", "Noto Serif", "Open Sans", "Palatino", "Roboto", "Segoe UI", "Symbol", "Tahoma", "Times New Roman",
    "Times", "Trebuchet MS", "Verdana", "Wingdings", "Zapf Dingbats", "Source Sans", "Carlito", "Caladea", "Univers", "Frutiger")
_SCHRIFT_KEY = {re.sub(r"[^a-z0-9]", "", b.lower()): b for b in BEKANNTE_SCHRIFTEN}


def _familie(name: Any) -> str:
    """Schriftfamilie nur als Muster: ein Name aus der Liste der gaengigen Familien (BEKANNTE_SCHRIFTEN), sonst 'anderer'.
    Die Familie wird als Anfang des Namens erkannt (Arial-BoldMT, ABCDEF+Calibri-Italic); ausgegeben wird immer der Name aus der Liste."""
    roh = re.sub(r"[^a-z0-9]", "", _SUBSET.sub("", str(name or "")).lower())
    for k in sorted(_SCHRIFT_KEY, key=len, reverse=True):
        if roh.startswith(k):
            return _SCHRIFT_KEY[k]
    return "anderer"


def _gerundet(x: float, schritt: float = 0.5) -> str:
    """Auf `schritt` gerundete Zahl als Schluessel ('10.5'); unsinnige Werte werden auf 0..9999 begrenzt."""
    x = x if math.isfinite(x) else 0.0
    return f"{round(max(0.0, min(9999.0, x)) / schritt) * schritt:g}"


def _annot_art(a: dict) -> str:
    try:
        st = (a.get("data") or {}).get("Subtype")
        name = getattr(st, "name", None) or (st if isinstance(st, str) else "")
    except Exception:
        name = ""
    return name if name in ANNOTATIONEN else "andere"


def _aufzaehlung(wort: str) -> str | None:
    if wort and all(c in _PUNKT_ZEICHEN for c in wort):
        return "punkt"
    if wort in ("-", "–", "—"):
        return "strich"
    if wort == "*":
        return "stern"
    if re.fullmatch(r"\d{1,3}[.)]", wort):
        return "zahl"
    if re.fullmatch(r"\(?[A-Za-z][.)]", wort):
        return "buchstabe"
    if re.fullmatch(r"\(?(?:[ivxIVX]{2,6})[.)]", wort):
        return "roemisch"
    return None


@dataclass
class _Strich:
    x0: float
    x1: float
    y: float
    staerke: float
    farbe: str


@dataclass
class _PdfSumme:
    seiten: int = 0
    masse: Counter = field(default_factory=Counter)
    gedreht: int = 0
    ohne_text: int = 0
    pro_seite: list[dict[str, int]] = field(default_factory=list)
    schriften: Counter = field(default_factory=Counter)
    groessen: Counter = field(default_factory=Counter)
    richtungen: Counter = field(default_factory=Counter)
    farben: dict[str, dict[str, int]] = field(default_factory=dict)
    hinterlegungen: Counter = field(default_factory=Counter)
    legende: Counter = field(default_factory=Counter)
    striche: dict[str, Any] = field(default_factory=lambda: {
        "anzahl": 0, "laenge_zur_seitenbreite": Counter(), "laenge_zum_wort": Counter(), "hoehenband": Counter(),
        "staerke": Counter(), "farben": Counter(), "farbe_wie_text": 0, "farbe_anders": 0})
    unter: dict[str, Any] = field(default_factory=lambda: {
        "anzahl": 0, "laenge_zur_seitenbreite": Counter(), "laenge_zum_wort": Counter(), "hoehenband": Counter(),
        "staerke": Counter(), "farben": Counter(), "farbe_wie_text": 0, "farbe_anders": 0})
    annotationen: Counter = field(default_factory=Counter)
    hyperlinks: int = 0
    silben: int = 0
    aufzaehlung: Counter = field(default_factory=Counter)
    hoch: int = 0
    bilder: int = 0
    ohne_geometrie: int = 0


def _striche_der_seite(page: Any, breite: float, hoehe: float) -> tuple[list[_Strich], list[tuple[float, float, float, float, str]]]:
    """Dünne waagerechte Linien (Striche, Unterstreichungen) und gefüllte Flächen (Hinterlegungen) einer Seite."""
    linien: list[_Strich] = []
    flaechen: list[tuple[float, float, float, float, str]] = []
    for l in page.lines:
        if abs(l["bottom"] - l["top"]) <= 1.0 and l["x1"] - l["x0"] >= 1.0:
            linien.append(_Strich(l["x0"], l["x1"], (l["top"] + l["bottom"]) / 2, float(l.get("linewidth") or 0.0),
                                  _hex_pdf(l.get("stroking_color"), _raum(l.get("scs")))))
    for r in page.rects:
        hh, ww = r["bottom"] - r["top"], r["x1"] - r["x0"]
        if hh <= 1.6 and ww >= 1.0:
            farbe = r.get("non_stroking_color") if r.get("fill") else r.get("stroking_color")
            raum = _raum(r.get("ncs") if r.get("fill") else r.get("scs"))
            linien.append(_Strich(r["x0"], r["x1"], (r["top"] + r["bottom"]) / 2, float(hh), _hex_pdf(farbe, raum)))
        elif r.get("fill") and hh >= 3.0 and ww >= 3.0 and ww * hh <= 0.5 * breite * hoehe:
            code = _hex_pdf(r.get("non_stroking_color"), _raum(r.get("ncs")))
            if code != "unklar" and min(int(code[i:i + 2], 16) for i in (1, 3, 5)) < 0xF8:           # weisser Hintergrund ist keine Markierung
                flaechen.append((r["x0"], r["top"], r["x1"], r["bottom"], code))
    return sorted(linien, key=lambda s: s.y), flaechen


def _zeilen_gruppen(woerter: list[dict]) -> list[list[dict]]:
    zeilen: list[list[dict]] = []
    for w in sorted(woerter, key=lambda w: (round(w["top"], 1), w["x0"])):
        if zeilen and abs(zeilen[-1][0]["top"] - w["top"]) <= 2.5:
            zeilen[-1].append(w)
        else:
            zeilen.append([w])
    for z in zeilen:
        z.sort(key=lambda w: w["x0"])
    return zeilen


def _seite_messen(page: Any, a: _PdfSumme) -> None:
    breite, hoehe = float(page.width), float(page.height)
    a.masse[f"{round(breite)}x{round(hoehe)}"] += 1
    if int(getattr(page, "rotation", 0) or 0) % 360:
        a.gedreht += 1
    zeichen = [c for c in page.chars if (c.get("text") or "").strip()]
    richtung = Counter(_richtung(c) for c in zeichen)
    a.richtungen.update(richtung)
    for c in zeichen:
        a.schriften[_familie(c.get("fontname"))] += 1
        a.groessen[_gerundet(float(c.get("size") or 0.0))] += 1
    a.hoch += sum(len(_HOCH_UNICODE.findall(c.get("text") or "")) for c in zeichen)
    n_bilder = len(page.images)
    a.bilder += n_bilder
    a.annotationen.update(_annot_art(x) for x in (page.annots or []))
    a.hyperlinks += len(page.hyperlinks or [])
    if not zeichen:
        a.ohne_text += 1
        a.pro_seite.append({"zeichen": 0, "woerter": 0, "bilder": n_bilder})
        return
    a.pro_seite.append({"zeichen": len(zeichen), "woerter": len(page.extract_words()), "bilder": n_bilder})
    woerter = page.extract_words(extra_attrs=["non_stroking_color", "ncs", "size", "fontname"])
    gerade = max(richtung, key=richtung.get) == "h" and not int(getattr(page, "rotation", 0) or 0) % 360   # type: ignore[arg-type]
    if gerade:
        linien, flaechen = _striche_der_seite(page, breite, hoehe)
    else:
        a.ohne_geometrie += 1
        linien, flaechen = [], []
    ys = [s.y for s in linien]
    gesehen_strich: set[int] = set()
    gesehen_unter: set[int] = set()
    for w in woerter:
        x0, x1, top, bot = w["x0"], w["x1"], w["top"], w["bottom"]
        wh, ww = max(bot - top, 0.1), max(x1 - x0, 0.1)
        raum = _raum(w.get("ncs"))
        code = _hex_pdf(w.get("non_stroking_color"), raum)
        klasse = _klasse_hex(code) if code != "unklar" else "unklar"
        strich = unter = False
        lo, hi = bisect_left(ys, top - 0.2 * wh), bisect_right(ys, bot + 0.45 * wh)
        for i in range(lo, hi):
            s = linien[i]
            if min(x1, s.x1) - max(x0, s.x0) < 0.5 * ww:
                continue
            rel = (s.y - top) / wh
            if 0.2 <= rel <= 0.8:
                strich = True
                art, gesehen = a.striche, gesehen_strich
            elif 0.8 < rel <= 1.45:
                unter = True
                art, gesehen = a.unter, gesehen_unter
            else:
                continue
            if id(s) in gesehen:
                continue
            gesehen.add(id(s))
            art["anzahl"] += 1
            art["laenge_zur_seitenbreite"][str(min(10, int((s.x1 - s.x0) / breite * 10)))] += 1
            art["laenge_zum_wort"][str(min(30, round((s.x1 - s.x0) / ww * 10)))] += 1
            art["hoehenband"][f"{round(rel, 1):g}"] += 1
            art["staerke"][_gerundet(s.staerke, 0.25)] += 1
            art["farben"][s.farbe] += 1
            art["farbe_wie_text" if s.farbe == code else "farbe_anders"] += 1
        hint = False
        flaeche = ww * wh
        for fx0, ftop, fx1, fbot, fcode in flaechen:
            ueber = max(0.0, min(x1, fx1) - max(x0, fx0)) * max(0.0, min(bot, fbot) - max(top, ftop))
            if ueber >= 0.5 * flaeche:
                hint = True
                a.hinterlegungen[fcode] += 1
                break
        e = a.farben.setdefault(f"{raum}|{code}", {"woerter": 0, "durchgestrichen": 0, "unterstrichen": 0, "hinterlegt": 0})
        e["woerter"] += 1
        e["durchgestrichen"] += int(strich)
        e["unterstrichen"] += int(unter)
        e["hinterlegt"] += int(hint)
        if klasse != "schwarz" or strich or unter or hint:
            a.legende[f"{klasse}|{'strich' if strich else '-'}|{'unterstrichen' if unter else '-'}|{'hinterlegt' if hint else '-'}"] += 1
    for zeile in _zeilen_gruppen(woerter):
        letztes = zeile[-1]["text"]
        if len(letztes) >= 3 and letztes[-1] in "-\u00ad" and letztes[-2].isalpha():
            a.silben += 1
        art = _aufzaehlung(zeile[0]["text"])
        if art:
            a.aufzaehlung[art] += 1
        ref = max(w["size"] for w in zeile)
        boden = sorted(w["bottom"] for w in zeile)[len(zeile) // 2]
        a.hoch += sum(1 for w in zeile if w["size"] <= 0.75 * ref and w["bottom"] <= boden - 0.1 * ref)


def _pdf_zahlen(pfad: Path) -> dict[str, Any]:
    import pdfplumber
    a = _PdfSumme()
    with pdfplumber.open(str(pfad)) as pdf:
        a.seiten = len(pdf.pages)
        for page in pdf.pages:
            try:
                _seite_messen(page, a)
            except Exception:                           # eine ungewoehnliche Seite verhindert den Steckbrief nicht
                a.ohne_geometrie += 1
                a.pro_seite.append({"zeichen": 0, "woerter": 0, "bilder": 0})

    def striche(d: dict[str, Any]) -> dict[str, Any]:
        return {k: (dict(sorted(v.items(), key=lambda kv: -kv[1])[:100]) if isinstance(v, Counter) else v) for k, v in d.items()}

    return {
        "seiten": a.seiten, "seiten_masse": dict(a.masse.most_common(100)), "gedrehte_seiten": a.gedreht, "seiten_ohne_text": a.ohne_text,
        "seite": a.pro_seite[:3000], "schriften": dict(a.schriften.most_common(200)),
        "schriftgroessen": dict(sorted(a.groessen.items(), key=lambda kv: float(kv[0]))[:200]),
        "textrichtungen": dict(a.richtungen), "farben": dict(sorted(a.farben.items(), key=lambda kv: -kv[1]["woerter"])[:300]),
        "hinterlegungen": dict(a.hinterlegungen.most_common(100)), "legende": dict(a.legende.most_common(200)),
        "striche": striche(a.striche), "unterstreichungen": striche(a.unter),
        "annotationen": dict(a.annotationen), "hyperlinks": a.hyperlinks, "silbentrennung": a.silben,
        "aufzaehlungszeichen": dict(a.aufzaehlung), "hochzahlen": a.hoch, "bilder": a.bilder,
        "geometrie_nicht_gemessen_seiten": a.ohne_geometrie}


# ======================================================================================
# 6. Excel lesen (vorher)
# ======================================================================================
_ZAHLFORMAT = re.compile(r"^(?:General|Standard|@|[0-9#.,%/;\- +():_*?]{1,24})$")
_WORT_BUCKETS = (1, 6, 11, 21, 41, 81, 161)


def _sichtbarkeit(ws: Any) -> str:
    return {"visible": "sichtbar", "hidden": "versteckt", "veryHidden": "sehr_versteckt"}.get(ws.sheet_state, "sichtbar")


def _spaltenbuchstabe(n: int | None) -> str | None:
    from openpyxl.utils import get_column_letter
    return get_column_letter(n) if n else None


def _belegte(ws: Any) -> list[Any]:
    roh = getattr(ws, "_cells", None)
    if roh is not None:
        return [roh[k] for k in sorted(roh)]
    return [c for row in ws.iter_rows(max_row=min(ws.max_row, 20000), max_col=min(ws.max_column, 200)) for c in row]


def _bucket(n: int) -> str:
    return str(max((b for b in _WORT_BUCKETS if b <= max(n, 1)), default=1))


def _schrift_name(font: Any) -> str:
    return _familie(getattr(font, "name", None)) if getattr(font, "name", None) else "keine Angabe"


@dataclass
class _VorherKontext:
    """Merkt sich, was der Vergleich mit 'nachher' braucht (bleibt im Speicher)."""
    rollen: dict[int, dict[str, int]] = field(default_factory=dict)        # Blatt-Index -> {section|titel|inhalt: Spalte}
    konstanten: dict[tuple[int, int], str] = field(default_factory=dict)   # (Blatt-Index, Spalte) -> haeufigster Wert (streng)


def _excel_vorher(pfad: Path, xl: Any) -> tuple[dict[str, Any], _VorherKontext]:
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText
    from openpyxl.worksheet.formula import ArrayFormula
    from . import messen as _m
    ctx = _VorherKontext()
    wb = load_workbook(str(pfad), rich_text=True)
    try:
        absaetze, woerter = Counter(), Counter()
        formeln, f_muster, f_funk, f_fremd = 0, Counter(), Counter(), 0
        verbunden = Counter()
        pruefung, pruefung_typen = 0, Counter()
        bedingt = filter_n = tabellen = rich = links = kommentare = versteckt_z = versteckt_s = mehrabs = 0
        sec_format, sec_nummern = Counter(), Counter()
        inh_schrift, inh_gr = Counter(), Counter()
        konst_liste: list[dict[str, Any]] = []
        blaetter: list[dict[str, Any]] = []
        namen = len(wb.defined_names)
        for i, ws in enumerate(wb.worksheets):
            lay = xl.layouts.get(ws.title) if xl is not None else None
            rollen: dict[str, int] = {}
            if lay is not None:
                for k, v in (("section", lay.section_spalte), ("titel", lay.titel_spalte), ("inhalt", lay.inhalt_spalte)):
                    if v:
                        rollen[k] = v
            ctx.rollen[i] = rollen
            sec, inh = rollen.get("section"), rollen.get("inhalt")
            kopf = set(lay.kopf_zeilen) if lay is not None else set()
            erste = min(kopf, default=0)                       # Zeilen davor sind Spaltenueberschriften
            zeilen_mit_wert: set[int] = set()
            zeilen_inhalt: set[int] = set()
            zellen_n, f_blatt = 0, 0
            spalten_werte: dict[int, Counter] = {}
            max_z = max_s = 0
            for c in _belegte(ws):
                if getattr(c, "_hyperlink", None) is not None:
                    links += 1
                if getattr(c, "comment", None) is not None:
                    kommentare += 1
                v = c.value
                if v is None:
                    continue
                zellen_n += 1
                zeilen_mit_wert.add(c.row)
                max_z, max_s = max(max_z, c.row), max(max_s, c.column)
                if isinstance(v, ArrayFormula):
                    v = str(v.text or "=")
                if isinstance(v, CellRichText):
                    rich += 1
                    v = str(v)
                if isinstance(v, str) and v.startswith("="):
                    formeln += 1
                    f_blatt += 1
                    muster, fn = formel_muster(v, c.row)
                    f_muster[muster] += 1
                    f_funk.update(fn)
                    f_fremd += int("@!" in muster)
                    continue
                if sec and c.column == sec and c.row >= erste:
                    sec_nummern["text" if isinstance(v, str) else "zahl"] += 1
                    fmt = str(c.number_format or "General")
                    sec_format[fmt if _ZAHLFORMAT.fullmatch(fmt) else "anderes"] += 1
                if c.column not in (sec, rollen.get("titel"), inh) and c.row >= erste:
                    spalten_werte.setdefault(c.column, Counter())[_m._s(_m._zelltext(v)) if isinstance(v, str) else repr(v)] += 1
                if inh and c.column == inh and isinstance(v, str) and v.strip():
                    zeilen_inhalt.add(c.row)
                    teile = [t for t in _nl(v).split("\n") if t.strip()]
                    absaetze[str(min(len(teile), 10))] += 1
                    mehrabs += int(len(teile) > 1)
                    woerter.update(_bucket(len(t.split())) for t in teile)
                    inh_schrift[_schrift_name(c.font)] += 1
                    if c.font.sz:
                        inh_gr[_gerundet(float(c.font.sz))] += 1
            for spalte, zaehler in sorted(spalten_werte.items()):
                n = sum(zaehler.values())
                wert, haeufig = zaehler.most_common(1)[0]
                if n >= 5 and haeufig / n >= 0.8:
                    ctx.konstanten[(i, spalte)] = wert
                    konst_liste.append({"blatt": i + 1, "spalte": _spaltenbuchstabe(spalte), "anteil": round(haeufig / n, 3),
                                        "zellen": n, "form": schablone(wert, max_woerter=6)})
            vz = sum(1 for d in ws.row_dimensions.values() if d.hidden)
            vs = sum(1 for d in ws.column_dimensions.values() if d.hidden)
            versteckt_z += vz
            versteckt_s += vs
            v_n = 0
            for rng in ws.merged_cells.ranges:
                verbunden[f"{rng.max_row - rng.min_row + 1}x{rng.max_col - rng.min_col + 1}"] += 1
                v_n += 1
            bedingt += len(list(ws.conditional_formatting))
            for dv in ws.data_validations.dataValidation:
                pruefung += 1
                t = dv.type or "any"
                pruefung_typen[t if t in ("list", "whole", "decimal", "date", "time", "textLength", "custom", "any") else "andere"] += 1
            filter_n += int(bool(ws.auto_filter and ws.auto_filter.ref))
            tabellen += len(ws.tables)
            namen += len(ws.defined_names)
            typen = Counter()
            for r in zeilen_mit_wert:
                typen["kopf" if r in kopf else "inhalt" if r in zeilen_inhalt else "sonstige"] += 1
            typen["leer"] = max(0, max_z - len(zeilen_mit_wert))
            blatt: dict[str, Any] = {
                "nr": i + 1, "sichtbarkeit": _sichtbarkeit(ws), "zeilen": max_z, "spalten": max_s, "zellen": zellen_n,
                "rollen": {k: _spaltenbuchstabe(v) for k, v in rollen.items()}, "zeilentypen": dict(typen),
                "versteckte_zeilen": vz, "versteckte_spalten": vs, "verbunden": v_n, "formeln": f_blatt}
            if lay is not None:
                blatt.update({"modell": lay.modell, "sicherheit": lay.sicherheit, "kombiniert": bool(lay.kombiniert),
                              "sections": len(lay.kopf_zeilen)})
            blaetter.append(blatt)
    finally:
        wb.close()
    bilder = diagramme = 0
    try:
        with zipfile.ZipFile(pfad) as z:
            for n in z.namelist():
                bilder += int(n.startswith("xl/media/"))
                diagramme += int(n.startswith("xl/charts/chart"))
    except (zipfile.BadZipFile, OSError):
        pass
    return {
        "blaetter": len(blaetter), "blaetter_versteckt": sum(1 for b in blaetter if b["sichtbarkeit"] != "sichtbar"),
        "blatt": blaetter[:300], "absaetze_je_zelle": dict(sorted(absaetze.items(), key=lambda kv: int(kv[0]))),
        "woerter_je_absatz": dict(sorted(woerter.items(), key=lambda kv: int(kv[0]))), "zellen_mit_mehreren_absaetzen": mehrabs,
        "formeln": {"anzahl": formeln, "muster": dict(f_muster.most_common(60)), "funktionen": dict(f_funk.most_common(60)),
                    "zu_anderem_blatt": f_fremd},
        "konstanten_spalten": konst_liste[:200], "verbundene_zellen": dict(verbunden.most_common(100)), "versteckte_zeilen": versteckt_z,
        "versteckte_spalten": versteckt_s, "datenpruefung": pruefung, "datenpruefung_typen": dict(pruefung_typen),
        "bedingte_formatierung": bedingt, "filter": filter_n, "tabellen": tabellen, "namen": namen,
        "section_zahlenformat": dict(sec_format.most_common(60)), "section_nummern": dict(sec_nummern), "rich_text_zellen": rich,
        "inhalt_schriften": dict(inh_schrift.most_common(60)), "inhalt_groessen": dict(inh_gr.most_common(60)),
        "hyperlinks": links, "kommentare": kommentare, "bilder": bilder, "diagramme": diagramme}, ctx


# ======================================================================================
# 7. Excel nachher: wie sind Aenderungen formatiert?
# ======================================================================================
def _stueck_font(stueck: Any, zelle: Any) -> Any:
    return stueck.basis if stueck.basis is not None else zelle.font


def _muster_text(formate: list[str]) -> str:
    return ",".join(formate[:12]) + (",mehr" if len(formate) > 12 else "")


def _trenner_paare(st: list[Any], z: dict[str, Any]) -> None:
    """Zaehlt Reihenfolge und Trenner zwischen Alt- und Neu-Stuecken einer Zelle."""
    n = len(st)
    benutzt: set[int] = set()
    for i, s in enumerate(st):
        if s.format not in (ALT, NEU):
            continue
        gegen = NEU if s.format == ALT else ALT
        # direkt benachbart oder durch ein reines Leerraum-Stueck getrennt
        j, mitte = i + 1, ""
        if j < n and st[j].format == NORMAL and st[j].text.strip() == "" and j + 1 < n:
            mitte, j = st[j].text, j + 1
        if j < n and st[j].format == gegen and i not in benutzt and j not in benutzt:
            benutzt.update((i, j))
            erst, zweit = s.text, st[j].text
            rest_erst = erst[len(erst.rstrip(_WS)):]
            rest_zweit = zweit[:len(zweit) - len(zweit.lstrip(_WS))]
            trenner = rest_erst + mitte + rest_zweit
            z["reihenfolge"]["alt_vor_neu" if s.format == ALT else "neu_vor_alt"] += 1
            z["trenner"][schablone(trenner)[:10] if trenner else "kein"] += 1
            ort = ("eigenes_stueck" if mitte else "beides" if rest_erst and rest_zweit else "im_alt" if rest_erst
                   else "im_neu" if rest_zweit else "kein")
            z["trenner_ort"][ort] += 1
    for i, s in enumerate(st):
        if s.format in (ALT, NEU) and i not in benutzt:
            z["nur_alt" if s.format == ALT else "nur_neu"] += 1


def _fuellcode(zelle: Any) -> str | None:
    f = zelle.fill
    if f is None or getattr(f, "fill_type", None) in (None, "none"):
        return None
    return _farbcode(getattr(f, "fgColor", None))


@dataclass
class _NachherKontext:
    zuordnung: dict[int, dict[int, int]] = field(default_factory=dict)      # Blatt-Index -> vorher-Zeile -> nachher-Zeile
    wv: Any = None
    wn: Any = None
    v_texte: dict[int, list[list[str]]] = field(default_factory=dict)


ZEILEN_BLOCK_MAX = 800      # Austausch-Bloecke bis so viele Zeilenpaare werden genau gepaart, groessere nach Position
LANGE_ZELLE = 200           # ab so vielen Zeichen wird die Zell-Aehnlichkeit ueber Woerter statt Zeichen gerechnet (Laufzeit)


def _zell_aehnlich_schnell(_m: Any, original: Any):
    """Wie `messen._zell_aehnlich`, aber lange Zellen werden wortweise verglichen (sonst braucht ein Block veraenderter Zeilen mit
    langen Zellen im Quadrat Zeit: ueber 30 s statt unter 1 s). Kurze Zellen: unveraendert dieselbe Rechnung wie `messen`."""
    def f(x: str, y: str) -> float:
        if max(len(x), len(y)) <= LANGE_ZELLE or x == y or not x or not y:
            return original(x, y)
        kurz, lang = (x, y) if len(x) <= len(y) else (y, x)
        best = 0.8 if _m._enthalten(kurz, lang) else 0.0
        sm = SequenceMatcher(None, x.split(), y.split(), autojunk=False)
        if sm.real_quick_ratio() >= 0.5 and sm.quick_ratio() >= 0.5:
            r = sm.ratio()
            if r >= 0.5:
                best = max(best, r)
        return best
    return f


def _zeilen_vergleich(_m: Any, v: list, n: list, spalten: Any) -> Any:
    """Zeilen-Zuordnung vorher <-> nachher wie `messen` (`messen._diff`), mit zwei Laufzeit-Schutzen (nur waehrend dieses Aufrufs):
    kleinere Blockgrenze und wortweiser Vergleich langer Zellen."""
    alt_max, alt_fn = _m.BLOCK_MAX, _m._zell_aehnlich
    _m.BLOCK_MAX = min(alt_max, ZEILEN_BLOCK_MAX)
    _m._zell_aehnlich = _zell_aehnlich_schnell(_m, alt_fn)
    try:
        return _m._diff(v, n, None, spalten)
    finally:
        _m.BLOCK_MAX, _m._zell_aehnlich = alt_max, alt_fn


def _stueck_zahlen(st: list[Any], zelle: Any, hat: bool, s: dict[str, Any]) -> None:
    """Farben, Strich und Schriftmerkmale der Stuecke einer Zelle in die Zaehler `s` (nur wo es Markierung oder Farbe gibt)."""
    for x in st:
        font = _stueck_font(x, zelle)
        code = _farbcode(getattr(font, "color", None))
        klasse = _klasse_hex(code)
        strich = bool(getattr(font, "strike", None))
        if hat or klasse not in ("standard", "schwarz"):
            s["stueck_farben"][x.format][code] += 1
            key = "rot" if klasse == "rot" else "blau" if klasse == "blau" else "schwarz" if klasse in ("schwarz", "standard") \
                else "andere_farbe"
            s["strich_summe"][key][0] += 1
            s["strich_summe"][key][1] += int(strich)
            s["farbe_strich"][f"{klasse}|{'strich' if strich else '-'}"] += 1
        if hat and x.format in (ALT, NEU):
            s["form"]["stuecke_fett"] += int(bool(getattr(font, "b", None)))
            s["form"]["stuecke_kursiv"] += int(bool(getattr(font, "i", None)))
            s["form"]["stuecke_unterstrichen"] += int(bool(getattr(font, "u", None)))
        if hat and x.format == NORMAL and klasse not in ("standard", "schwarz"):
            s["andere_farbe"][code] += 1


def _excel_nachher(vorher: Path, nachher: Path, vk: _VorherKontext) -> tuple[dict[str, Any], _NachherKontext]:
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText
    from . import messen as _m
    wv = load_workbook(str(vorher), rich_text=True)
    wn = load_workbook(str(nachher), rich_text=True)
    nk = _NachherKontext(wv=wv, wn=wn)
    try:
        return _excel_nachher_messen(wv, wn, vk, nk, _m, CellRichText), nk
    finally:
        wn.close()


def _excel_nachher_messen(wv: Any, wn: Any, vk: _VorherKontext, nk: _NachherKontext, _m: Any, CellRichText: Any) -> dict[str, Any]:
    paar = {"reihenfolge": Counter(), "trenner": Counter(), "trenner_ort": Counter(), "nur_alt": 0, "nur_neu": 0}
    zeilen, zellen, alt_ansicht = Counter(), Counter(), Counter()
    muster, rollen_z, spalten_z, zf_markiert = Counter(), Counter(), Counter(), Counter()
    s = {"stueck_farben": {f: Counter() for f in (NORMAL, ALT, NEU)}, "farbe_strich": Counter(), "andere_farbe": Counter(),
         "strich_summe": {k: [0, 0] for k in ("rot", "blau", "andere_farbe", "schwarz")}, "form": Counter()}
    schriftfarben, fuellfarben = Counter(), Counter()
    neu = {"anzahl": 0, "zellfarben": Counter(), "spalten_befuellt": Counter(), "alle_stuecke_neu": 0, "muster": Counter(),
           "konstanten_uebernommen": 0, "konstanten_fehlen": 0, "formeln": 0}
    gestr = {"anzahl": 0, "zellfarben": Counter(), "alle_durchgestrichen": 0, "inhalt_noch_da": 0}
    hoehe_geaendert = breite_geaendert = 0
    n_blaetter = min(len(wv.worksheets), len(wn.worksheets))
    for bi in range(n_blaetter):
        sv, sn = wv.worksheets[bi], wn.worksheets[bi]
        v_rows = [[_m._zelltext(c.value) for c in row] for row in sv.iter_rows()]
        n_rows = [[_m._zelltext(c.value) for c in row] for row in sn.iter_rows()]
        nk.v_texte[bi] = v_rows
        rollen = vk.rollen.get(bi, {})
        rolle_von = {col: k for k, col in rollen.items()}
        zv_, zn_ = _m._zeilen(v_rows), _m._zeilen(n_rows)
        d = _zeilen_vergleich(_m, zv_, zn_, set(rollen.values()) or None)
        nk.zuordnung[bi] = dict(d.zuordnung)
        zeilen["vorher"] += len(zv_)
        zeilen["nachher"] += len(zn_)
        zeilen["gepaart"] += len(d.zuordnung)
        zeilen["neu"] += len(d.neue_nr)
        zeilen["entfernt"] += len(d.weg)
        zellen["sauber_geleert"] += sum(1 for (vr, _c), t in d.zellen.items() if t == "" and vr in d.zuordnung)
        n_zu_v = {n: v for v, n in d.zuordnung.items()}
        neue_nr = set(d.neue_nr)
        for r_n, row in enumerate(sn.iter_rows(), 1):
            if r_n not in n_zu_v and r_n not in neue_nr:
                continue
            r_v = n_zu_v.get(r_n)
            ist_neu = r_v is None
            zeile_zellen: list[tuple[Any, list[Any], str, str]] = []           # (Zelle, Stuecke, nachher-Text, vorher-Text)
            neu_stuecke: list[Any] = []
            for c in row:
                wert = c.value
                if wert is None:
                    continue
                col = c.column
                vt = ""
                if not ist_neu and r_v - 1 < len(v_rows) and col - 1 < len(v_rows[r_v - 1]):
                    vt = v_rows[r_v - 1][col - 1]
                nt = _m._zelltext(wert)
                ist_text = isinstance(wert, CellRichText) or (isinstance(wert, str) and not wert.startswith("="))
                st = markierung.stuecke_aus_zelle(c)
                formate = [x.format for x in st]
                hat = markierung.hat_markierung(st)
                if not ist_text:                                   # Zahl / Formel / Datum
                    if hat:
                        zf_markiert.update({f: formate.count(f) for f in set(formate) if f != NORMAL})
                    if ist_neu:
                        neu["formeln"] += int(isinstance(wert, str))
                    elif _m._s(nt) != _m._s(vt):
                        zellen["zahl_formel_geaendert"] += 1
                elif ist_neu:
                    if hat:
                        muster_n = _muster_text(formate)
                        neu["muster"][muster_n] += 1
                elif hat:
                    zellen["mit_markierung"] += 1
                    zellen["markiert_text_geaendert" if _m._s(nt) != _m._s(vt) else "markiert_nur_format"] += 1
                    alt = markierung.ansicht(st, ALT)
                    alt_ansicht["ja" if _m._s(alt) == _m._s(vt) else
                                "nur_leerraum" if " ".join(alt.split()) == " ".join(vt.split()) else "nein"] += 1
                    muster[_muster_text(formate)] += 1
                    zellen["ganze_zelle_markiert"] += int(len(st) == 1)
                    _trenner_paare(st, paar)
                    rollen_z[rolle_von.get(col, "andere")] += 1
                    spalten_z[_spaltenbuchstabe(col)] += 1
                elif _m._s(nt) == _m._s(vt):
                    zellen["gleich"] += 1
                else:
                    zellen["sauber_ergaenzt" if not _m._s(vt) else "sauber_ersetzt"] += 1
                    rollen_z[rolle_von.get(col, "andere")] += 1
                    spalten_z[_spaltenbuchstabe(col)] += 1
                _stueck_zahlen(st, c, hat, s)
                if ist_neu:
                    neu["spalten_befuellt"][_spaltenbuchstabe(col)] += 1
                    for x in st:
                        neu["zellfarben"][_farbcode(getattr(_stueck_font(x, c), "color", None))] += 1
                    neu_stuecke += [x for x in st if x.text.strip()]
                else:
                    if col in rolle_von:
                        zeile_zellen.append((c, st, nt, vt))
                    vz = sv.cell(r_v, col)
                    if not hat:                                    # andere Schriftfarbe / Fuellung als vorher, ohne Markierung
                        code_n, code_v = _farbcode(getattr(c.font, "color", None)), _farbcode(getattr(vz.font, "color", None))
                        if code_n != code_v and _klasse_hex(code_n) not in ("standard", "schwarz"):
                            s["form"]["zellen_andere_schriftfarbe"] += 1
                            schriftfarben[code_n] += 1
                    fn, fv = _fuellcode(c), _fuellcode(vz)
                    if fn and fn != fv:
                        s["form"]["zellen_neu_gefuellt"] += 1
                        fuellfarben[fn] += 1
            if ist_neu:
                if neu_stuecke or any(c.value is not None for c in row):
                    neu["anzahl"] += 1
                    neu["alle_stuecke_neu"] += int(bool(neu_stuecke) and all(x.format == NEU for x in neu_stuecke))
                    for (b2, col), konst in vk.konstanten.items():
                        if b2 == bi:
                            zelle_n = n_rows[r_n - 1][col - 1] if col - 1 < len(n_rows[r_n - 1]) else ""
                            neu["konstanten_uebernommen" if _m._s(zelle_n) == konst else "konstanten_fehlen"] += 1
                continue
            text_zellen = [z for z in zeile_zellen if z[2].strip()]
            if text_zellen and all(z[1] and all(x.format == ALT for x in z[1]) for z in text_zellen):
                gestr["anzahl"] += 1
                gestr["alle_durchgestrichen"] += 1
                for c, st, _nt, _vt in text_zellen:
                    for x in st:
                        gestr["zellfarben"][_farbcode(getattr(_stueck_font(x, c), "color", None))] += 1
                gestr["inhalt_noch_da"] += int(all(_m._s(z[2]) == _m._s(z[3]) for z in text_zellen))
            inh = rollen.get("inhalt")
            if inh and r_v - 1 < len(v_rows) and inh - 1 < len(v_rows[r_v - 1]) and _m._s(v_rows[r_v - 1][inh - 1]):
                nz = n_rows[r_n - 1][inh - 1] if inh - 1 < len(n_rows[r_n - 1]) else ""
                zeilen["inhalt_geleert"] += int(not _m._s(nz))
        for r_v2, r_n2 in d.zuordnung.items():                     # Zeilenhoehen / Spaltenbreiten
            hv = sv.row_dimensions[r_v2].height if r_v2 in sv.row_dimensions else None
            hn = sn.row_dimensions[r_n2].height if r_n2 in sn.row_dimensions else None
            hoehe_geaendert += int(_runde(hv) != _runde(hn))
        for buchstabe in set(sv.column_dimensions.keys()) | set(sn.column_dimensions.keys()):
            bv = sv.column_dimensions[buchstabe].width if buchstabe in sv.column_dimensions else None
            bn = sn.column_dimensions[buchstabe].width if buchstabe in sn.column_dimensions else None
            breite_geaendert += int(_runde(bv) != _runde(bn))
    zeilen["gestrichen_markiert"] = gestr["anzahl"]
    n_v, n_n = len(wv.worksheets), len(wn.worksheets)
    zellen_out = {k: int(zellen.get(k, 0)) for k in ("gleich", "mit_markierung", "markiert_text_geaendert", "markiert_nur_format",
                                                    "sauber_ersetzt", "sauber_ergaenzt", "sauber_geleert", "ganze_zelle_markiert",
                                                    "zahl_formel_geaendert")}
    zellen_out["alt_ansicht_stimmt"] = dict(alt_ansicht)
    form = s["form"]

    def tabelle(c: Counter, n: int = 30) -> dict[str, int]:
        return dict(c.most_common(n))

    return {
        "blaetter_vorher": n_v, "blaetter_nachher": n_n, "blaetter_gleich": n_v == n_n,
        "zeilen": {k: int(zeilen.get(k, 0)) for k in ("vorher", "nachher", "gepaart", "neu", "entfernt", "gestrichen_markiert",
                                                      "inhalt_geleert")},
        "zellen": zellen_out, "zahl_formel_markiert": dict(zf_markiert),
        "muster": tabelle(muster, 100),
        "stueck_farben": {f: tabelle(c, 60) for f, c in s["stueck_farben"].items() if c},
        "farbe_strich": tabelle(s["farbe_strich"], 60),
        "strich_anteil": {k: round(st / n, 3) for k, (n, st) in s["strich_summe"].items() if n},
        "trenner": tabelle(paar["trenner"]), "trenner_ort": dict(paar["trenner_ort"]), "reihenfolge": dict(paar["reihenfolge"]),
        "nur_alt": paar["nur_alt"], "nur_neu": paar["nur_neu"], "geaenderte_rollen": dict(rollen_z),
        "geaenderte_spalten": tabelle(spalten_z, 60),
        "formatierung": {"zellen_andere_schriftfarbe": form["zellen_andere_schriftfarbe"], "schriftfarben": tabelle(schriftfarben),
                         "zellen_neu_gefuellt": form["zellen_neu_gefuellt"], "fuellfarben": tabelle(fuellfarben),
                         "stuecke_andere_farbe": tabelle(s["andere_farbe"]), "stuecke_fett": form["stuecke_fett"],
                         "stuecke_kursiv": form["stuecke_kursiv"], "stuecke_unterstrichen": form["stuecke_unterstrichen"]},
        "neue_zeilen": {k: (tabelle(v) if isinstance(v, Counter) else v) for k, v in neu.items()},
        "gestrichene_zeilen": {k: (tabelle(v) if isinstance(v, Counter) else v) for k, v in gestr.items()},
        "masse": {"zeilenhoehe_geaendert": hoehe_geaendert, "spaltenbreite_geaendert": breite_geaendert,
                  "zeilenhoehe_geaendert_ja": hoehe_geaendert > 0, "spaltenbreite_geaendert_ja": breite_geaendert > 0}}


def _runde(x: Any) -> Any:
    return None if x is None else round(float(x), 1)


# ======================================================================================
# 8. Schablonen-Beispiele (Umgebung jeder Aenderung als Formen)
# ======================================================================================
def _typ(t: Any) -> str:
    return t if t in TYPEN else "andere"


def _gleichmaessig(aenderungen: list[Any], maximal: int = MAX_BEISPIELE) -> list[Any]:
    """Hoechstens `maximal` Aenderungen, gleichmaessig ueber die Typen verteilt (reihum), je Typ gleichmaessig ueber die Liste."""
    gruppen: dict[str, list[Any]] = {}
    for a in sorted(aenderungen, key=lambda x: x.id):
        gruppen.setdefault(_typ(a.typ), []).append(a)
    kontingent = {t: 0 for t in gruppen}
    n = 0
    while n < maximal and any(kontingent[t] < len(g) for t, g in gruppen.items()):
        for t in sorted(gruppen):
            if n < maximal and kontingent[t] < len(gruppen[t]):
                kontingent[t] += 1
                n += 1
    out: list[Any] = []
    for t in sorted(gruppen):
        g, k = gruppen[t], kontingent[t]
        out += [g[int(i * len(g) / k)] for i in range(k)] if k else []
    return out


def _lauf(paare: list[tuple[str, str]]) -> list[dict[str, str]]:
    return [{"f": k, "form": f} for k, f in paare]


def fehler_formen(soll: Any, ist: Any) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Formen von Soll und Ist rund um die erste Abweichung (fuer `messen --schablonen`): zwei Laufverzeichnisse, kein Klartext."""
    a, b = _nl(soll), _nl(ist)
    pos = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), min(len(a), len(b)))
    return _lauf(fenster([(NORMAL, a)], pos)), _lauf(fenster([(NORMAL, b)], pos))


def fehler_eintrag(quelle: str, fehlerart: str, typ: str, erzeuger: str, soll: Any, ist: Any) -> dict[str, Any]:
    """Ein Eintrag der Fehler-Schablonen (siehe FEHLER_SCHABLONE): nur feste Begriffe und Formen."""
    f_soll, f_ist = fehler_formen(soll, ist)
    return {"quelle": quelle,
            "fehlerart": fehlerart if fehlerart in FEHLER_SCHABLONE.felder["fehlerart"].werte else "ohne_art",
            "aenderungstyp": typ if typ in TYPEN else "ohne_typ",
            "erzeuger": erzeuger if erzeuger in ("writer", "libreoffice", "sonstiger") else "unbekannt",
            "soll": f_soll, "ist": f_ist}


def pruefe_eintrag(eintrag: Any) -> list[str]:
    """Prueft EINEN Eintrag der Fehler-Schablonen gegen die Positivliste."""
    fehler: list[str] = []
    FEHLER_SCHABLONE.pruefe(eintrag, "fehler_schablonen[*]", fehler)
    return fehler


def _pdf_ansicht(absatz: Any, ae_id: str, welche: str) -> list[tuple[str, str]]:
    """Die PDF-Ansicht eines Absatzes (alt: bleibt + entfaellt, neu: bleibt + neu) als Formen rund um die Aenderung."""
    stuecke: list[tuple[str, str]] = []
    start, pos, letzter_trenner = None, 0, False
    for l in absatz.laeufe:
        if l.aenderung == ae_id and l.art in ("d", "a") and start is None:
            start = pos
        if l.art == "s":
            if letzter_trenner:                          # ein ausgelassener Lauf haette zwei Trenner nebeneinander gelassen
                continue
            f = NORMAL
        elif l.art == "k":
            f = NORMAL
        elif l.art == "d" and welche == "alt":
            f = ALT
        elif l.art == "a" and welche == "neu":
            f = NEU
        else:
            continue
        letzter_trenner = l.art == "s"
        stuecke.append((f, l.text))
        pos += len(_nl(l.text))
    return fenster(stuecke, start or 0)


_ZIEL = re.compile(r"^(.+)!([A-Z]+)(\d+)$")


def _excel_ansichten(ae: Any, xl: Any, vk: _VorherKontext, nk: _NachherKontext) -> tuple[list, list]:
    """(Excel vorher, Excel nachher) als Formen rund um die Aenderung; leer, wenn die Aenderung keine einzelne Zelle trifft."""
    from openpyxl.utils import column_index_from_string
    op = next((o for o in ae.ops if o.art in ("ersetzen", "einfuegen", "zelle_leeren")), None)
    blatt, zeile, spalte, start = (op.blatt, op.zeile, op.spalte, op.start) if op is not None else (None, 0, 0, 0)
    if op is None:
        m = _ZIEL.match(ae.ziel or "")
        if not m:
            return [], []
        blatt, zeile, spalte, start = m.group(1), int(m.group(3)), column_index_from_string(m.group(2)), 0
    wv, wn = nk.wv, nk.wn
    if blatt not in wv.sheetnames:
        return [], []
    bi = wv.sheetnames.index(blatt)
    if bi >= len(wn.worksheets):
        return [], []
    zv = wv[blatt].cell(zeile, spalte)
    st_v = markierung.stuecke_aus_zelle(zv)
    vor = fenster([(s.format, s.text) for s in st_v], start)
    r_n = nk.zuordnung.get(bi, {}).get(zeile)
    if r_n is None:
        return vor, []
    zn = wn.worksheets[bi].cell(r_n, spalte)
    st_n = markierung.stuecke_aus_zelle(zn)
    text_n = "".join(_nl(s.text) for s in st_n)
    pos, gefunden = 0, False
    for s in st_n:
        if s.format != NORMAL:
            gefunden = True
            break
        pos += len(_nl(s.text))
    if not gefunden:
        text_v = "".join(_nl(s.text) for s in st_v)
        pos = next((i for i, (x, y) in enumerate(zip(text_v, text_n)) if x != y), min(len(text_v), len(text_n)))
    return vor, fenster([(s.format, s.text) for s in st_n], pos)


def _beispiele(plan: Any, xl: Any, vk: _VorherKontext, nk: _NachherKontext) -> list[dict[str, Any]]:
    absatz_von: dict[str, Any] = {}
    for ab in getattr(plan, "absaetze", []):
        for i in ab.aenderungen:
            absatz_von.setdefault(i, ab)
    out: list[dict[str, Any]] = []
    for ae in _gleichmaessig(plan.aenderungen):
        ab = absatz_von.get(ae.id)
        if ab is not None:
            pa, pn = _pdf_ansicht(ab, ae.id, "alt"), _pdf_ansicht(ab, ae.id, "neu")
        else:
            pa = fenster([(ALT, ae.alt)]) if ae.alt else []
            pn = fenster([(NEU, ae.neu)]) if ae.neu else []
        try:
            ev, en = _excel_ansichten(ae, xl, vk, nk)
        except Exception:                                    # eine Stelle ohne Zelle darf das Beispiel nicht verhindern
            ev, en = [], []
        out.append({"typ": _typ(ae.typ), "status": ae.status if ae.status in STATUS else "nicht_uebernommen",
                    "pdf_alt": _lauf(pa), "pdf_neu": _lauf(pn), "excel_vorher": _lauf(ev), "excel_nachher": _lauf(en)})
    return out


# ======================================================================================
# 9. Das Tool selbst (nur Zaehler)
# ======================================================================================
def _bekannte_codes() -> frozenset[str]:
    try:
        from .pdf_reader import _HINWEIS_TEXT
        return frozenset(_HINWEIS_TEXT)
    except Exception:                                       # pragma: no cover
        return frozenset()


def _hinweis_code(roh: str) -> str:
    c = roh.strip().lower()
    if c.startswith("kommentar_"):
        c = "kommentar"
    if c in _bekannte_codes() or c in ("mehrspaltig", "tabelle"):
        return c
    return "anderer_code"


def _tool_zahlen(pa: Any, plan: Any | None) -> dict[str, Any]:
    woerter = Counter(t.status for b in pa.blocks for t in b.tokens)
    hinweise = Counter()
    for b in pa.blocks:
        for t in b.tokens:
            for c in filter(None, (t.hinweis or "").split(";")):
                hinweise[_hinweis_code(c)] += 1
    out: dict[str, Any] = {"lauf_ok": plan is not None, "aenderungen": 0, "je_typ": {}, "je_status": {},
                           "hinweise": dict(hinweise), "pruefcodes": {}, "woerter_status": dict(woerter),
                           "warnungen": len(pa.warnungen), "markierte_woerter": int(pa.markierte_woerter)}
    if plan is not None:
        out["aenderungen"] = len(plan.aenderungen)
        out["je_typ"] = dict(Counter(_typ(a.typ) for a in plan.aenderungen))
        out["je_status"] = dict(Counter(a.status for a in plan.aenderungen if a.status in STATUS))
        codes = Counter()
        for a in plan.aenderungen:
            for c in set(a.pruefcodes or []):
                c = str(c).strip().lower()
                codes[c if CODE.re.fullmatch(c) else "anderer_code"] += 1
        out["pruefcodes"] = dict(codes)
    return out


def _luecken_zahlen(pa: Any) -> dict[str, Any]:
    out: dict[str, dict[str, int]] = {}
    for l in pa.luecken:
        art = l.get("art") if l.get("art") in LUECKEN else "sonstiges"
        e = out.setdefault(art, {"anzahl": 0, "woerter": 0})
        e["anzahl"] += 1
        e["woerter"] += int(l.get("woerter") or 0)
    return out


# ======================================================================================
# 10. Ein Steckbrief
# ======================================================================================
def erzeuge(pdf: Path, vorher: Path, nachher: Path, name: str) -> dict[str, Any]:
    """Baut den Steckbrief eines Paars (nur Zahlen, Farbcodes, Formen). Wirft bei kaputten Dateien (ohne Inhalte in der Meldung)."""
    from .excel_model import ExcelModell
    from .pdf_reader import analysiere_pdf
    from .planner import plane
    pa = analysiere_pdf(pdf)
    xl = ExcelModell(vorher)
    try:
        plan = plane(pa, xl)
    except Exception:                                         # das Tool darf den Steckbrief nicht verhindern
        plan = None
    pdf_teil = _pdf_zahlen(pdf)
    pdf_teil["luecken"] = _luecken_zahlen(pa)
    pdf_teil["kopf_fuss"] = {"eintraege": len(pa.ignoriert), "woerter": sum(int(i.get("woerter") or 0) for i in pa.ignoriert)}
    ev, vk = _excel_vorher(vorher, xl)
    en, nk = _excel_nachher(vorher, nachher, vk)
    try:
        beispiele = _beispiele(plan, xl, vk, nk) if plan is not None else []
    finally:
        nk.wv.close()
    return {"art": "steckbrief", "format": STECKBRIEF_FORMAT, "name": name, "pdf": pdf_teil, "excel_vorher": ev,
            "excel_nachher": en, "tool": _tool_zahlen(pa, plan), "schablonen": beispiele}


# ======================================================================================
# 11. Markdown (nur aus den geprueften Daten)
# ======================================================================================
def _wert_md(v: Any) -> str:
    if isinstance(v, bool):
        return "ja" if v else "nein"
    if isinstance(v, float):
        return f"{v:.3g}"
    if isinstance(v, dict):
        return ", ".join(f"{k}={_wert_md(x)}" for k, x in v.items()) or "-"
    if isinstance(v, list):
        return ", ".join(_wert_md(x) for x in v) or "-"
    return str(v)


def _code(s: Any) -> str:
    return "`" + str(s).replace("`", "'") + "`"


def _zelle(s: Any, code: bool = False) -> str:
    t = _code(s) if code else _wert_md(s)
    return t.replace("|", "\\|")


def _lauf_md(laeufe: list[dict[str, str]]) -> str:
    return " ‖ ".join(f"{l['f']}:{l['form']}" for l in laeufe) if laeufe else "(leer)"


def _ist_skalar(v: Any) -> bool:
    return not isinstance(v, (dict, list))


def _skalar_dict(v: Any) -> bool:
    return isinstance(v, dict) and all(_ist_skalar(x) for x in v.values())


def _md_teil(wert: Any, titel: str, ebene: int, z: list[str]) -> None:
    kopf = "#" * min(ebene, 6)
    if isinstance(wert, dict):
        if not wert:
            return
        if all(_skalar_dict(x) and x for x in wert.values()):             # Tabelle: Zeile = Schluessel, Spalte = Unterschluessel
            spalten = list(dict.fromkeys(k for x in wert.values() for k in x))
            z += ["", f"{kopf} {titel}", "", "| | " + " | ".join(_zelle(s, True) for s in spalten) + " |", "|---|" + "---|" * len(spalten)]
            z += ["| " + _zelle(k, True) + " | " + " | ".join(_zelle(x.get(s, "")) for s in spalten) + " |" for k, x in wert.items()]
            return
        z += ["", f"{kopf} {titel}", ""]
        if _skalar_dict(wert) and (len(wert) > 8 or any(not re.fullmatch(r"[a-z_]+", k) for k in wert)):
            z += ["| Eintrag | Wert |", "|---|---|"] + [f"| {_zelle(k, True)} | {_zelle(v)} |" for k, v in wert.items()]
            return
        z += [f"- {k.replace('_', ' ')}: {_wert_md(v)}" for k, v in wert.items() if _ist_skalar(v)]
        for k, v in wert.items():
            if not _ist_skalar(v):
                _md_teil(v, k.replace("_", " "), ebene + 1, z)
    elif isinstance(wert, list):
        if not wert:
            return
        z += ["", f"{kopf} {titel}", ""]
        if wert and all(isinstance(x, dict) and any(isinstance(y, list) for y in x.values()) for x in wert):     # Schablonen
            for i, x in enumerate(wert, 1):
                z.append(f"{i}. " + "; ".join(f"{k.replace('_', ' ')} {_code(_lauf_md(y))}" if isinstance(y, list) else f"{k}: {y}"
                                              for k, y in x.items()))
        elif wert and all(isinstance(x, dict) for x in wert):
            spalten = list(dict.fromkeys(k for x in wert for k in x))
            z += ["| " + " | ".join(k.replace("_", " ") for k in spalten) + " |", "|" + "---|" * len(spalten)]
            for x in wert:
                z.append("| " + " | ".join(_zelle(x.get(k, ""), isinstance(x.get(k), str)) for k in spalten) + " |")
        else:
            z += [f"- {_wert_md(x)}" for x in wert]


def markdown(daten: dict[str, Any]) -> str:
    titel = "Zusammenfassung" if daten.get("art") == "zusammenfassung" else "Steckbrief"
    z = [f"# {titel}", "",
         "Diese Datei enthält nur Zahlen, Farbcodes und Formen (A = Großbuchstabe, a = Kleinbuchstabe, 9 = Ziffer, · = geschütztes "
         "Leerzeichen, ¶ = Zeilenumbruch, ⇥ = Tab, W… = langes Wort, ? = anderes Zeichen). Kein Text aus den Dokumenten.",
         "Bei Läufen steht vor jeder Form das Kürzel: `a:` = alt (durchgestrichen), `n:` = neu, `-:` = normal; die Läufe sind durch ‖ getrennt.", ""]
    for k in ("name", "paare", "format"):
        if k in daten:
            z.append(f"- {k}: {_wert_md(daten[k])}")
    for k, v in daten.items():
        if k not in ("art", "format", "name", "paare"):
            _md_teil(v, k.replace("_", " "), 2, z)
    return "\n".join(z) + "\n"


# ======================================================================================
# 12. Schreiben (append-only) und Kommandozeile
# ======================================================================================
def _neuer_ordner(basis: Path, name: str) -> Path:
    basis.mkdir(parents=True, exist_ok=True)
    nummern = [int(m.group(1)) for d in basis.iterdir() if (m := re.match(r"^(\d{3,})_", d.name))]
    nr = max(nummern, default=0) + 1
    while True:
        ziel = basis / f"{nr:03d}_{name}"
        try:
            ziel.mkdir(exist_ok=False)
            return ziel
        except FileExistsError:
            nr += 1


def _schreibe(ordner: Path, daten: dict[str, Any], dateiname: str = "steckbrief") -> None:
    ordner.mkdir(parents=True, exist_ok=True)
    with open(ordner / f"{dateiname}.json", "x", encoding="utf-8") as f:
        json.dump(daten, f, ensure_ascii=False, indent=1)
    with open(ordner / f"{dateiname}.md", "x", encoding="utf-8") as f:
        f.write(markdown(daten))


class _Parser(argparse.ArgumentParser):
    def error(self, message: str) -> None:                  # argparse wuerde Eingaben (z. B. Dateinamen) wiederholen
        raise ValueError("Ungültige Angaben. Aufruf siehe docs/anleitung_steckbrief.md.")


def _kurz(fehler: list[str], n: int = 5) -> list[str]:
    return [f"  - {m}" for m in fehler[:n]] + ([f"  - ... und {len(fehler) - n} weitere"] if len(fehler) > n else [])


def pruefe_datei(pfad: Path) -> int:
    try:
        daten = json.loads(Path(pfad).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        print("Die Datei konnte nicht gelesen werden oder ist keine gültige JSON-Datei.")
        return 2
    fehler = pruefe(daten)
    if not fehler:
        print("OK: enthält nur Zahlen, Farbcodes und Formen")
        return 0
    print(f"NICHT in Ordnung: {len(fehler)} Verstoß/Verstöße gegen die Positivliste (die Inhalte werden hier nicht wiederholt):")
    print("\n".join(_kurz(fehler, 20)))
    print("Bitte diese Datei NICHT weitergeben.")
    return 1


def _paar_ordner(eltern: Path) -> list[Path]:
    dateien = ("aenderungen.pdf", "vorher.xlsx", "nachher.xlsx")
    return sorted(d for d in eltern.iterdir() if d.is_dir() and not d.name.startswith(".") and all((d / n).is_file() for n in dateien))


def main(argv: list[str] | None = None) -> int:
    ap = _Parser(prog="cde steckbrief", description="Steckbrief eines Paars: nur Zahlen, Farbcodes und Formen, kein Text.")
    ap.add_argument("--pdf")
    ap.add_argument("--vorher")
    ap.add_argument("--nachher")
    ap.add_argument("--name")
    ap.add_argument("--ausgabe")
    ap.add_argument("--ordner")
    ap.add_argument("--pruefe")
    try:
        a = ap.parse_args(sys.argv[1:] if argv is None else argv)
    except ValueError as e:
        print(e)
        return 2
    except SystemExit:
        return 2
    if a.pruefe:
        return pruefe_datei(Path(a.pruefe))
    if a.name is not None and not NAME_RE.fullmatch(a.name):
        print("Der Name darf nur aus kleinen Buchstaben (a-z), Ziffern, Unterstrich und Bindestrich bestehen (höchstens 40 Zeichen).")
        return 2
    basis = Path(a.ausgabe) if a.ausgabe else config.pfad("steckbrief_dir")
    if a.ordner:
        eltern = Path(a.ordner)
        if not eltern.is_dir():
            print("Der Ordner mit den Paaren wurde nicht gefunden.")
            return 2
        paare = _paar_ordner(eltern)
        if not paare:
            print("Es wurden keine Paare gefunden. Jedes Paar ist ein Ordner mit aenderungen.pdf, vorher.xlsx und nachher.xlsx.")
            return 2
        quellen = [(d / "aenderungen.pdf", d / "vorher.xlsx", d / "nachher.xlsx") for d in paare]
        namen = [f"paar_{i:02d}" for i in range(1, len(quellen) + 1)]
        gruppe = a.name or "sammlung"
    else:
        if not (a.pdf and a.vorher and a.nachher):
            print("Bitte --pdf, --vorher und --nachher angeben (oder --ordner, oder --pruefe).")
            return 2
        quellen = [(Path(a.pdf), Path(a.vorher), Path(a.nachher))]
        if not all(p.is_file() for p in quellen[0]):
            print("Eine der drei Dateien wurde nicht gefunden.")
            return 2
        namen = [a.name or "paar"]
        gruppe = namen[0]
    ergebnisse: list[dict[str, Any]] = []
    for nr, (pdf, vorher, nachher) in enumerate(quellen):
        try:
            sb = erzeuge(pdf, vorher, nachher, namen[nr])
        except Exception:                                    # kaputte Datei o. ae.: einfache Meldung, nie Rohtext oder Pfade
            print(f"{namen[nr]}: Dieses Paar konnte nicht gelesen werden (beschädigte oder nicht unterstützte Datei).")
            continue
        fehler = pruefe(sb)
        if fehler:
            print(f"{namen[nr]}: Der Steckbrief verstößt gegen die Positivliste und wird NICHT geschrieben "
                  f"({len(fehler)} Verstoß/Verstöße). Bitte melden:")
            print("\n".join(_kurz(fehler)))
            return 3
        ergebnisse.append(sb)
        print(f"{namen[nr]}: gelesen", flush=True)
    if not ergebnisse:
        print("Es wurde kein Steckbrief erzeugt.")
        return 2
    zusammen = fasse_zusammen(ergebnisse) if a.ordner else None
    if zusammen is not None:
        fz = pruefe(zusammen)
        if fz:
            print("Die Zusammenfassung verstößt gegen die Positivliste und wird NICHT geschrieben. Bitte melden:")
            print("\n".join(_kurz(fz)))
            return 3
    ziel = _neuer_ordner(basis, gruppe)
    if a.ordner:
        for sb in ergebnisse:
            _schreibe(ziel / sb["name"], sb)
        _schreibe(ziel, zusammen, "zusammenfassung")            # type: ignore[arg-type]
        print(f"Zusammenfassung: {ziel / 'zusammenfassung.json'}")
    else:
        _schreibe(ziel, ergebnisse[0])
        print(f"Steckbrief: {ziel / 'steckbrief.json'}")
    print(f"Vor dem Weitergeben bitte prüfen:  python -m src.cli steckbrief --pruefe \"{(ziel / ('zusammenfassung.json' if a.ordner else 'steckbrief.json'))}\"")
    return 0


if __name__ == "__main__":
    sys.exit(main())
