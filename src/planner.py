"""Plant, was in die Excel eingetragen wird -- ohne die Excel zu veraendern.

Eingabe: PDF-Analyse (Woerter mit Status) + Excel-Modell.
Ausgabe: Liste von Aenderungen. Jede Aenderung ist entweder
  * uebernommen  (Konfidenz hoch: wird automatisch eingetragen),
  * pruefen      (Vorschlag steht fest, ein Mensch bestaetigt),
  * nicht_uebernommen (mit Grund). Nichts geht still verloren.

Alles Regeln, keine KI. Alle Operationen beziehen sich auf den Originalzustand
der Excel; die Reihenfolge beim Eintragen regelt der Applier.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any, Iterable, Iterator

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from . import codes as _codes
from . import config
from . import markierung as mk
from .excel_model import REGISTER_UEBERSCHRIFT, Kandidat, ExcelModell, zahl_als_text
from .abgleich import (ExaktIndex, Label, _aufbau, _positionen, _unscharfe_abbildung, Fundstelle, Spur, Zuordnung, _woerter_norm,
                       ganz_loeschen, label_der_spur, ops_fuer_spur, reissverschluss, spuren, umbenennen)
from .gegenprobe import normalform, vergleiche, zustand, zustand_mit_optional
from .matcher import Matcher, Treffer
from .models import (ADD, DEL, HOCH, KEEP, MITTEL, MOD, NICHT, NIEDRIG, PRUEFEN, UEBERNOMMEN,
                     Absatz, Aenderung, Lauf, Operation, PdfBlock, PdfToken)
from .pdf_reader import PdfAnalyse, hinweis_text
from .textutil import kopf_zerlegen, norm, section_nummer

STATUS_VON_STUFE = {HOCH: UEBERNOMMEN, MITTEL: PRUEFEN}


def _planer_cfg(cfg: dict | None) -> dict:
    """Abschnitt `planer:` der Konfiguration (Schwellen und Zeichenlisten, Regel 28)."""
    if cfg and "planer" in cfg:
        return cfg["planer"]
    return config.lade("zuordnung").get("planer", {})


def _trennung_ohne_ausnahme(muster: "re.Pattern[str]", text: str, ausnahmen: set[str]) -> bool:
    """Gibt es im Text einen Treffer des Trennmusters, dessen Rest-Wort keine erlaubte Ausnahme ist ("Ein- und Ausgang")?"""
    for m in muster.finditer(text):
        rest = (m.groupdict().get("rest") or "").casefold()
        if rest not in ausnahmen:
            return True
    return False


def neutext_probleme(texte: Iterable[str], cfg: dict | None = None) -> dict[str, Any]:
    """Prueft Neutext auf Dinge, die in der Excel Schaden anrichten oder auf kaputtes PDF-Lesen hindeuten
    (Ligaturen, Zeichencode-Reste, nicht zusammengefuegte Silbentrennung, Tabellen-Spaltentrenner; Muster in config planer:).

    Rueckgabe: steuerzeichen (nicht schreibbar), verdaechtig (Liste einfacher Hinweise -> nie automatisch),
    formel (Text beginnt mit '=' und wuerde von Excel als Formel gelesen)."""
    pl = _planer_cfg(cfg)
    zeichen = "".join(pl.get("verdaechtige_zeichen", []))
    muster = [re.compile(m) for m in pl.get("verdaechtige_muster", [])]
    trennung = [re.compile(m) for m in pl.get("trennmuster", [])]
    ausnahmen = {str(w).casefold() for w in pl.get("trennung_ausnahmen", [])}
    tabelle = [re.compile(m, re.MULTILINE) for m in pl.get("tabellen_muster", [])]
    tab_min = max(int(pl.get("tabellen_min_treffer", 1)), 1)
    out: dict[str, Any] = {"steuerzeichen": False, "verdaechtig": [], "formel": False}
    for t in texte:
        if not isinstance(t, str) or not t:
            continue
        if ILLEGAL_CHARACTERS_RE.search(t):
            out["steuerzeichen"] = True
        if zeichen and any(z in t for z in zeichen):
            h = "Der neue Text enthält Zeichen (z. B. Ligaturen), die beim Lesen des PDF verfälscht sein können. Bitte mit dem PDF vergleichen."
            if h not in out["verdaechtig"]:
                out["verdaechtig"].append(h)
        if any(m.search(t) for m in muster):
            h = "Der neue Text enthält Zeichencode-Reste aus dem PDF (z. B. „(cid:12)“). Die PDF-Schrift wurde nicht sauber gelesen."
            if h not in out["verdaechtig"]:
                out["verdaechtig"].append(h)
        if any(_trennung_ohne_ausnahme(m, t, ausnahmen) for m in trennung):
            h = ("Der neue Text enthält ein Trennmuster wie „Wort- Rest“ (Silbentrennung am Zeilenende). Das Wort könnte "
                 "zerrissen sein. Bitte mit dem PDF vergleichen.")
            if h not in out["verdaechtig"]:
                out["verdaechtig"].append(h)
        if tabelle and sum(len(m.findall(t)) for m in tabelle) >= tab_min:
            h = ("Der neue Text enthält senkrechte Striche als Spaltentrenner (Text aus einer Tabelle?). Tabellen werden nicht "
                 "gelesen. Bitte mit dem PDF vergleichen.")
            if h not in out["verdaechtig"]:
                out["verdaechtig"].append(h)
        if t.startswith("="):
            out["formel"] = True
    return out


_ZAHL_TEXT = re.compile(r"-?\d+(\.\d+)?")


def zahl_aus_text(text: str, zahlenformat: str | None = None) -> int | float | None:
    """Zahl, die genau diesen Text als Anzeige hat -- sonst None (dann bliebe nur Text, das Zahlenformat ginge verloren).

    Ohne `zahlenformat` wird nur geprueft, ob der Text eine einfache Zahl ist (Planer, kennt das Format nicht);
    mit Format muss die Anzeige der Zahl (siehe `zahl_als_text`) buchstabengleich sein (Applier). Hoechstens 15 Ziffern."""
    t = text.strip()
    if not _ZAHL_TEXT.fullmatch(t) or sum(c.isdigit() for c in t) > 15:
        return None
    wert: int | float = float(t) if "." in t else int(t)
    if zahlenformat is not None and zahl_als_text(wert, zahlenformat) != t:
        return None
    return wert


def _wort_gleich(a: str, b: str) -> bool:
    """Zwei Woerter gelten als gleich: gleiche Normform; Woerter ohne Normform ('•') muessen buchstabengleich sein."""
    na, nb = norm(a), norm(b)
    return na == nb if (na or nb) else a == b


def _woerter_gleich(excel: list[str], pdf: list[str]) -> bool:
    return len(excel) == len(pdf) and bool(pdf) and all(_wort_gleich(x, y) for x, y in zip(excel, pdf))


@dataclass
class Plan:
    aenderungen: list[Aenderung]
    layout: list[dict]
    warnungen: list[str] = field(default_factory=list)
    ignoriert: list[dict] = field(default_factory=list)
    markierte_woerter: int = 0
    luecken: list[dict] = field(default_factory=list)   # Schema 2: Dinge, die nicht gelesen werden konnten
    absaetze: list[Absatz] = field(default_factory=list)   # Schema 3: PDF-Absaetze fuer die Gegenprobe (ADR 0004)
    # Spec-Filter (Welle 3): {"erkannt": "11 22 33" | None, "quelle": vorgabe | pdf | einzige | "", "zeilen": [von, bis] | None,
    # "kandidaten": [Specs der CDE], "pdf": [im PDF genannte Specs der CDE]}; None = die CDE hat keine Spec-Spalte
    spec: dict[str, Any] | None = None
    # Welle 6: Luecken auf Seiten, auf denen nichts zur CDE gehoert (Deckblatt, Pruef- und Revisionstabelle, Anhang). Nur Info:
    # sie zaehlen nicht als "nicht gelesen" (`luecken`), bleiben aber im Protokoll sichtbar (Bilanz-Prinzip: nichts still).
    luecken_ausserhalb: list[dict] = field(default_factory=list)

    @property
    def alle_luecken(self) -> list[dict]:
        """Alle gemeldeten Luecken, auch die ausserhalb der CDE (fuer das Messwerkzeug und die Gesamtzahl)."""
        return list(self.luecken) + list(self.luecken_ausserhalb)


def _bbox(tokens: list[PdfToken]) -> tuple[float, float, float, float]:
    return (min(t.bbox[0] for t in tokens), min(t.bbox[1] for t in tokens),
            max(t.bbox[2] for t in tokens), max(t.bbox[3] for t in tokens))


@dataclass
class _Ort:
    """Ankerstelle fuer einen NEUEN Absatz bzw. eine neue Section (aus der Kette oder, unsicher, aus dem alten Matcher)."""
    blatt: str
    zeile: int
    spalte: int
    art: str                 # kopf | absatz | zelle | teil
    start: int
    end: int
    text: str
    score: float
    stufe: str
    davor: bool              # der Anker steht VOR dem neuen Block (sonst danach)
    gruende: list[str] = field(default_factory=list)
    codes: list[str] = field(default_factory=list)

    @property
    def ziel(self) -> str:
        return f"{self.blatt}!{get_column_letter(self.spalte)}{self.zeile}"


# Saetze zu Pruefcodes in einfachem Deutsch: EINE Tabelle in src/codes.py (frueher hier _GRUND, _CODE_SATZ, _A1_SATZ)
_GRUND = _codes.SAETZE
_A1_SATZ = _codes.A1_ART


def _in(a: Fundstelle, b: Fundstelle) -> bool:
    """Liegt Fundstelle a ganz in b (gleiche Zelle)?"""
    if a.blatt != b.blatt or len(a.zellen) != 1:
        return False
    z, s, x, y = a.zellen[0]
    return any((z2, s2) == (z, s) and x2 <= x and y <= y2 for z2, s2, x2, y2 in b.zellen)


class _Planer:
    def __init__(self, pdf: PdfAnalyse, xl: ExcelModell, cfg: dict, spec: str | None = None) -> None:
        self.pdf, self.xl, self.cfg = pdf, xl, cfg
        self.spec_info: dict[str, Any] | None = None
        self.spec_warnungen: list[str] = []
        self._spec_zeilen: dict[str, set[int]] = {}          # Blatt -> Zeilen der Spec des PDFs (leer = kein Filter)
        self._spec_unklar = False
        self._spec_waehlen(spec)
        self.matcher = Matcher(xl, cfg, self._erlaubt)
        self.out: list[Aenderung] = []
        self.absaetze: list[Absatz] = []
        self._tr: dict[tuple, Treffer | None] = {}
        self.verbraucht: set[int] = set()   # Spuren, die schon als Teil einer neuen Section behandelt wurden
        self._ag: list[str] = []                    # Gruende, warum die Ankerstelle des aktuellen neuen Blocks unsicher ist
        self._risiko_cache: dict[tuple[str, int], list[str]] = {}               # (Blatt, Einfuegezeile) -> Risiken
        self._sperre: dict[str, list[str]] = {}                                 # Aenderung -> Sperr-Codes des Absatzes
        self._fenster: tuple | None = None          # Reissverschluss-Fenster fuer den Ersatz-Matcher
        self._deckel: list[str] | None = None       # Ersatz-Matcher aktiv: hoechstens "mittel", mit diesen Codes
        self._spur_von: dict[str, int] = {}         # Aenderung -> Spur-Nummer (Index)
        self._geloescht_bis: dict[tuple[str, int, int], list[tuple[int, int]]] = {}   # Zelle -> geloeschte Bereiche ganzer Absaetze
        self._ist: str = ""                         # Excel-Text der Ersatz-Stelle (fuer ist_absatz)
        self._soll: str = ""

    # ------------------------------------------------------------ Spec-Filter (Welle 3, Nico: eine CDE enthaelt mehrere Specs)
    def _erlaubt(self, blatt: str, zeile: int) -> bool:
        """Darf diese Excel-Zeile benutzt werden? Im Register mit mehreren Specs nur die Zeilen der Spec des PDFs."""
        zeilen = self._spec_zeilen.get(blatt)
        return zeilen is None or zeile in zeilen

    def _zeilen_der_spec(self, blatt: str) -> set[int] | None:
        return self._spec_zeilen.get(blatt)

    def _pdf_specs(self, bekannt: set[str]) -> tuple[Counter, Counter]:
        """(sicher, alle): Specs der CDE, die das PDF in der Titel-/Fusszeile (planer.spec_pdf_muster) bzw. irgendwo nennt."""
        from .excel_model import _register_cfg, spec_norm
        rc = _register_cfg(config.lade("zuordnung")["excel"])
        kern = rc.get("spec_muster", r"(\d{2})[ _.\-](\d{2})[ _.\-](\d{2})")
        stark_rx = [re.compile(str(m).replace("{spec}", "(?P<spec>" + kern + ")")) for m in _planer_cfg(self.cfg).get(
            "spec_pdf_muster", [])]
        texte = [x for b in self.pdf.blocks for x in {b.alt_text, b.neu_text} if x]
        texte += [str(i.get("text", "")) for i in self.pdf.ignoriert if i.get("text")]
        sicher: Counter = Counter()
        alle: Counter = Counter()
        for t in texte:
            for rx in stark_rx:
                m = rx.search(t)
                if m and spec_norm(m.group("spec"), rc) in bekannt:
                    sicher[spec_norm(m.group("spec"), rc)] += 1
            for m in re.finditer(kern, t):
                sp = spec_norm(m.group(0), rc)
                if sp in bekannt:
                    alle[sp] += 1
        return sicher, alle

    def _spec_waehlen(self, vorgabe: str | None) -> None:
        """Spec des PDFs bestimmen. Mehrere Specs in der CDE: nur deren Zeilen werden benutzt; keine oder mehrere passende:
        Warnung und nichts automatisch. Eine Spec in der CDE: wie bisher (kein Filter)."""
        from .excel_model import _register_cfg, spec_norm
        je_spec: dict[str, dict[str, list[int]]] = {}
        for blatt, lay in self.xl.layouts.items():
            for sp, zeilen in self.xl.specs(blatt).items():
                je_spec.setdefault(sp, {})[blatt] = zeilen
        if not je_spec:
            if vorgabe:
                self.spec_warnungen.append("Die CDE hat keine Spalte mit Spezifikationsnummern; die vorgegebene Spec wird "
                                           "nicht benutzt.")
            return
        sicher, alle = self._pdf_specs(set(je_spec))
        gewaehlt, quelle = None, ""
        if vorgabe:
            v = spec_norm(vorgabe, _register_cfg(config.lade("zuordnung")["excel"]))
            if v in je_spec:
                gewaehlt, quelle = v, "vorgabe"
            else:
                self.spec_warnungen.append("Die vorgegebene Spec steht nicht in der CDE. Es wird nichts automatisch eingetragen.")
        elif len(je_spec) == 1:
            gewaehlt, quelle = next(iter(je_spec)), "einzige"
        elif len(sicher) == 1:
            gewaehlt, quelle = next(iter(sicher)), "pdf"
        elif not sicher and len(alle) == 1:
            gewaehlt, quelle = next(iter(alle)), "pdf"
        self.spec_info = {"erkannt": gewaehlt, "quelle": quelle, "zeilen": None, "kandidaten": sorted(je_spec),
                          "pdf": sorted(set(sicher) | set(alle))}
        if gewaehlt is None:
            self._spec_unklar = len(je_spec) > 1 or bool(vorgabe)
            if not vorgabe:
                self.spec_warnungen.append(
                    "Spec des PDFs nicht in der CDE gefunden" + (" (mehrere passen)" if len(sicher) > 1 or len(alle) > 1 else "")
                    + f": die CDE enthält {len(je_spec)} Specs. Es wird nichts automatisch eingetragen; bitte die Spec wählen.")
            return
        bereiche = je_spec[gewaehlt]
        zeilen = sorted(r for rs in bereiche.values() for r in rs)
        self.spec_info["zeilen"] = [zeilen[0], zeilen[-1]]
        if len(je_spec) > 1:
            self._spec_zeilen = {blatt: set(rs) for blatt, rs in bereiche.items()}
            for blatt in {b for d in je_spec.values() for b in d}:
                self._spec_zeilen.setdefault(blatt, set())
            self.spec_warnungen.append(f"Die CDE enthält {len(je_spec)} Specs; benutzt werden nur die {len(zeilen)} Zeilen der "
                                       f"Spec {gewaehlt}. Die Zeilen der anderen Specs bleiben unverändert.")

    # ------------------------------------------------------------ Helfer
    def treffer(self, b: PdfBlock) -> Treffer | None:
        key = (b.idx, self._fenster)
        if key not in self._tr:
            a = [(i, norm(t.old)) for i, t in enumerate(b.tokens) if t.old and norm(t.old)]
            self._tr[key] = self.matcher.finde(a, b.section_alt, ist_kopf=(b.kind == "ueberschrift"), fenster=self._fenster)
        return self._tr[key]

    def _basis(self, b: PdfBlock, typ: str, toks: list[PdfToken], alt: str, neu: str) -> Aenderung:
        ae = Aenderung(id=f"A{len(self.out) + 1:04d}", typ=typ, seite=toks[0].page, bbox=_bbox(toks),
                       section=b.section_neu or b.section_alt, alt=alt, neu=neu,
                       woerter=sum(1 for t in toks if t.status != KEEP), block_idx=b.idx,
                       kontext=(b.neu_text or b.alt_text)[:400])
        self.out.append(ae)
        if self._aktuelle_spur is not None:
            self._spur_von[ae.id] = self._aktuelle_spur
            self._sperre[ae.id] = list(self._spur_sperre)
            self._hinweise_nur(ae)
        ae.ist_absatz = self._ist
        ae.soll_absatz = self._soll
        return ae

    _a1_satz: str = ""
    _aktuelle_spur: int | None = None
    _spur_sperre: list[str] = []
    _spur_hinweise: list[str] = []

    def _hinweise_nur(self, ae: Aenderung) -> None:
        for c in self._spur_hinweise:
            satz = hinweis_text(c)
            if satz:
                self._hinweis(ae, satz)

    def _codes_des_absatzes(self, sp: Spur, mit: Iterable[Spur] = ()) -> tuple[list[str], list[str]]:
        """(Sperr-Codes, reine Hinweis-Codes) ueber ALLE Woerter des Absatzes (nicht nur die geaenderten).
        Sperr-Codes stehen in config planer.unsichere_hinweise (Praefix "kommentar" gilt fuer kommentar_*)."""
        pl = _planer_cfg(self.cfg)
        liste = [str(c) for c in pl.get("unsichere_hinweise", [])]
        nur = {str(c) for c in pl.get("hinweis_codes", [])}
        sperre: list[str] = []
        hinweise: list[str] = []
        for s in [sp, *mit]:
            codes = list(s.codes)
            for t in s.tokens:
                codes += [c for c in t.hinweis.split(";") if c]
            for c in codes:
                if any(c == x or c.startswith(x + "_") for x in liste):
                    if c not in sperre:
                        sperre.append(c)
                elif (c in nur or c.split("_")[0] in nur) and c not in hinweise:
                    hinweise.append(c)
        return sperre, hinweise

    def _setze(self, ae: Aenderung, stufe: str, grund: str = "", codes: Iterable[str] = ()) -> None:
        codes = list(codes)
        # Neutext pruefen: Steuerzeichen koennen nicht geschrieben werden, Ligaturen/(cid:N) sind nie "sicher"
        texte = [ae.neu] + [o.text for o in ae.ops] + [t for o in ae.ops for t in o.zellen.values()]
        p = neutext_probleme(texte, self.cfg)
        if p["steuerzeichen"]:
            self._nicht(ae, "Der neue Text enthält Steuerzeichen und kann nicht in die Excel geschrieben werden.")
            self._code(ae, "steuerzeichen")
            return
        if self._deckel is not None:
            if stufe == HOCH:
                stufe = MITTEL
            codes += self._deckel
            if self._deckel and "excel_weicht_ab" in self._deckel:
                grund = (grund + "; " if grund else "") + _GRUND["excel_weicht_ab"]
                if self._a1_satz:
                    self._hinweis(ae, self._a1_satz)
        sperre = self._sperre.get(ae.id, [])
        if sperre:
            if stufe == HOCH:
                stufe = MITTEL
            for c in sperre:
                satz = hinweis_text(c) or _codes.satz(c)
                if satz:
                    self._hinweis(ae, satz)
                codes.append("sperre:" + c)
            grund = (grund + "; " if grund else "") + ("Im Absatz gibt es eine Stelle, die nicht sicher gelesen werden konnte "
                                                       "(siehe Hinweise). Der ganze Absatz wird nicht automatisch eingetragen.")
            # Auftrag 004: Bei diesen Sperren ist schon die Stelle des Textes unsicher (z. B. ob ein Absatz auf der naechsten
            # Seite weitergeht); ein fertiger Vorschlag waere geraten. Besser kein Vorschlag als ein falscher.
            ohne = [c for c in sperre if c in [str(x) for x in _planer_cfg(self.cfg).get("sperren_ohne_vorschlag", [])]]
            if ohne and ae.ops:
                ae.ops = []
                codes.append("kein_vorschlag")
                self._hinweis(ae, _codes.satz("kein_vorschlag"))
        if p["verdaechtig"]:
            if stufe == HOCH:
                stufe = MITTEL
            for h in p["verdaechtig"]:
                self._hinweis(ae, h)
            codes.append("neutext_verdaechtig")
            grund = (grund + "; " if grund else "") + "Der neue Text könnte aus dem PDF verfälscht gelesen worden sein."
        if p["formel"] and any(o.art in ("ersetzen", "einfuegen", "zeile_einfuegen") for o in ae.ops):
            self._hinweis(ae, "Der neue Text beginnt mit „=“. Er wird als Text eingetragen, nicht als Formel.")
        for op in ae.ops:                # Zahl-Zelle: der neue Text muss wieder eine Zahl sein, sonst geht das Zahlenformat verloren
            if op.art in ("ersetzen", "einfuegen") and self.xl.ist_zahl(op.blatt, op.zeile, op.spalte):
                alt_z = self.xl.text(op.blatt, op.zeile, op.spalte)
                ergebnis = alt_z[:op.start] + op.text + alt_z[op.ende:]
                if ergebnis.strip() and zahl_aus_text(ergebnis) is None:
                    if stufe == HOCH:
                        stufe = MITTEL
                    self._hinweis(ae, "Die Zelle in der Excel ist eine Zahl, der neue Text ist keine. Das Zahlenformat geht verloren.")
                    grund = (grund + "; " if grund else "") + "Zahlenformat geht verloren (die Zelle ist eine Zahl, der neue Text nicht)."
                    codes.append("zahlformat")
                    break
        risiken = self._zeilen_risiken(ae)
        if risiken:                      # Zeile einfuegen verschiebt Dinge, die openpyxl nicht mitzieht: nie automatisch
            if stufe == HOCH:
                stufe = MITTEL
            for r in risiken:
                self._hinweis(ae, r)
            codes.append("zeilen_risiko")
            grund = (grund + "; " if grund else "") + ("Beim Einfügen der neuen Zeile könnten Dinge in der Excel verrutschen "
                                                       "oder brechen (siehe Hinweise). Bitte vor dem Eintragen prüfen.")
        ae.konfidenz = stufe
        ae.status = STATUS_VON_STUFE.get(stufe, NICHT)
        for c in codes:
            self._code(ae, c)
        if grund:
            ae.grund = (ae.grund + "; " if ae.grund else "") + grund

    @staticmethod
    def _code(ae: Aenderung, code: str) -> None:
        if code and code not in ae.pruefcodes:
            ae.pruefcodes.append(code)

    def _zeilen_risiken(self, ae: Aenderung) -> list[str]:
        """Hindernisse (einfache Saetze) aller Zeilen-Einfuegungen einer Aenderung. Leer = unbedenklich.
        Was der Applier selbst mitzieht (verbundene Zellen unterhalb, ausgeblendete Zeilen darunter), zaehlt nicht."""
        out: list[str] = []
        for op in ae.ops:
            if op.art != "zeile_einfuegen" or op.abgedeckt:
                continue
            pos = op.zeile if op.vor else op.zeile + 1
            if (op.blatt, pos) not in self._risiko_cache:
                self._risiko_cache[(op.blatt, pos)] = self.xl.einfuege_hindernisse(op.blatt, pos)
            out += [r for r in self._risiko_cache[(op.blatt, pos)] if r not in out]
        return out

    @staticmethod
    def _hinweis(ae: Aenderung, text: str) -> None:
        if text not in ae.hinweise:
            ae.hinweise.append(text)

    def _nicht(self, ae: Aenderung, grund: str, ziel: str = "") -> None:
        ae.status, ae.konfidenz, ae.ops = NICHT, NIEDRIG, []
        ae.grund = (ae.grund + "; " if ae.grund else "") + grund
        if ziel:
            ae.ziel = ziel

    def _kein_treffer(self, ae: Aenderung, b: PdfBlock) -> None:
        tr = self.treffer(b)
        if tr is None:
            self._nicht(ae, "Der Absatz wurde in der Excel nicht gefunden.")
            self._code(ae, "nicht_gefunden")
        else:
            self._nicht(ae, f"Zu unsichere Stelle in der Excel (Übereinstimmung {tr.score:.0%}).", tr.kandidat.ziel)
            self._code(ae, "zu_unsicher")
            ae.excel_text = tr.kandidat.text[:300]
        ae.score = tr.score if tr else 0.0

    def _ziel(self, xtok) -> str:
        return f"{xtok.blatt}!{get_column_letter(xtok.spalte)}{xtok.zeile}"

    @staticmethod
    def _gruppen(b: PdfBlock) -> Iterator[tuple[int, int]]:
        i, n = 0, len(b.tokens)
        while i < n:
            if b.tokens[i].status == KEEP:
                i += 1
                continue
            j = i
            while j < n and b.tokens[j].status != KEEP:
                j += 1
            yield i, j
            i = j

    # ------------------------------------------------------------ Hauptschleife
    def lauf(self) -> None:
        self.idx = ExaktIndex(self.xl, self.cfg, self.matcher.index, self._erlaubt)
        self.spuren = spuren(self.pdf, self.idx, self.cfg)
        # ADR 0006: CDE-Register in der Mappe -> Labels der PDF-Absaetze (Sub-Section, Section, Part stehen in eigenen Spalten)
        self._register = any(l.register for l in self.xl.layouts.values())
        self._labels = [label_der_spur(sp) if self._register else Label() for sp in self.spuren]
        self.zuo = self._zuordnen()
        self._zugeordnet = self._zugeordnete_zeilen()
        self._anker_je_blatt = Counter(z.fund.blatt for z in self.zuo if z.automatisch and z.fund is not None)
        self._ausserhalb = self._bereich()
        self._sub_neu = self._subs_im_neuen_stand() if self._register else []
        for i, sp in enumerate(self.spuren):
            if i in self.verbraucht or not sp.hat_aenderung:
                continue
            if i in self._ausserhalb:
                self._ausserhalb_spur(i, self._ausserhalb[i])
                continue
            self._spur_start(i)
            if sp.ganz_neu:
                self._neu(i)
            elif sp.ganz_geloescht:
                self._geloescht_spur(i)
            else:
                self._text_spur(i)
            self._spur_ende()
        self._loeschungen_koppeln()
        self._leere_zellen_naht()

    def _leere_zellen_naht(self) -> None:
        """Wird der ganze Inhalt einer Zelle geloescht und an ihrem Rand ein neuer Absatz eingesetzt, braucht der neue Absatz
        keinen Zeilenumbruch zum (verschwundenen) Nachbarn."""
        leer = {(o.blatt, o.zeile, o.spalte) for a in self.out for o in a.ops
                if o.art == "ersetzen" and o.text == "" and o.start == 0 and o.ende == len(self.xl.text(o.blatt, o.zeile, o.spalte))
                and o.ende > 0}
        for a in self.out:
            if a.typ not in ("absatz_neu",):
                continue
            for o in a.ops:
                if o.art == "ersetzen" and o.start == o.ende and (o.blatt, o.zeile, o.spalte) in leer:
                    o.text = o.text.strip("\n")

    def _spur_start(self, i: int, mit: Iterable[Spur] = ()) -> None:
        self._aktuelle_spur = i
        self._spur_sperre, self._spur_hinweise = self._codes_des_absatzes(self.spuren[i], mit)
        self._fenster = self.zuo[i].fenster
        self._deckel = None
        self._a1_satz = ""
        self._ist, self._soll = "", ""

    def _spur_ende(self) -> None:
        self._aktuelle_spur = None
        self._spur_sperre, self._spur_hinweise = [], []
        self._fenster, self._deckel = None, None
        self._ist, self._soll = "", ""

    # ------------------------------------------------------------ Reissverschluss
    def _zuordnen(self) -> list[Zuordnung]:
        nf = self.cfg.get("normalform")
        alle: list[list[tuple[Fundstelle, str]]] = []
        for i, sp in enumerate(self.spuren):
            cands: list[tuple[Fundstelle, str]] = []
            ids = {l.aenderung for l in sp.laeufe if l.aenderung}
            alt, opt_alt = zustand_mit_optional(sp.laeufe, set())
            neu, opt_neu = zustand_mit_optional(sp.laeufe, ids)
            orte: set[tuple] = set()
            if alt.strip():
                for f in self.idx.finde(alt, opt_alt):
                    cands.append((f, "alt"))
                    orte.add((f.blatt, f.zellen))
            if ids and neu.strip() and normalform(neu, nf)[0] != normalform(alt, nf)[0]:
                for f in self.idx.finde(neu, opt_neu):
                    if (f.blatt, f.zellen) in orte:
                        continue
                    # Ein Stueck des alten Absatzes ist kein "schon eingetragen" (am Ende gestrichenes Wort); umgekehrt ist der
                    # alte Absatz als Stueck des schon eingetragenen neuen (am Ende angehaengter Satz) keine Stelle zum Aendern.
                    if any(_in(f, g) for g, art in cands if art == "alt"):
                        continue
                    cands = [(g, art) for g, art in cands if not (art == "alt" and _in(g, f))]
                    cands.append((f, "neu"))
            if self._labels[i].art in ("sub", "section", "part"):
                schon = {(f.blatt, f.zellen) for f, _ in cands}
                cands += [c for c in self._register_kandidaten(i) if (c[0].blatt, c[0].zellen) not in schon]
            alle.append(self._filtern(sp, cands))
        return reissverschluss(alle)

    # ------------------------------------------------------------ CDE-Register (ADR 0006 Punkt 5)
    def _register_kandidaten(self, i: int) -> list[tuple[Fundstelle, str]]:
        """Absatz mit Sub-Section-Buchstabe ("A. Text") bzw. Ueberschrift ("1.1 TITEL", "PART 1 TITEL"): im Register steht das
        Label in der Sub-Section- bzw. Section-Spalte, in der Textspalte nur der Text. Gesucht wird der Text ohne Label; die
        Zeile zaehlt nur, wenn ihr Label passt (gleicher Buchstabe / gleiche Section, keine Listennummer im Text)."""
        sp, lab = self.spuren[i], self._labels[i]
        koerper = sp.laeufe[lab.n:]
        ids = {l.aenderung for l in sp.laeufe if l.aenderung}
        nf = self.cfg.get("normalform")
        alt, opt_alt = zustand_mit_optional(koerper, set())
        neu, opt_neu = zustand_mit_optional(koerper, ids)
        out: list[tuple[Fundstelle, str]] = []
        if lab.alt and alt.strip():
            out += [(f, "alt") for f in self.idx.finde(alt, opt_alt) if self._label_passt(f, lab.art, lab.alt)]
        if lab.neu and neu.strip() and ids and (lab.neu != lab.alt or normalform(neu, nf)[0] != normalform(alt, nf)[0]):
            orte = {(f.blatt, f.zellen) for f, _ in out}
            out += [(f, "neu") for f in self.idx.finde(neu, opt_neu) if self._label_passt(f, lab.art, lab.neu)
                    and (f.blatt, f.zellen) not in orte]
        return out

    def _label_passt(self, f: Fundstelle, art: str, label: str) -> bool:
        """Steht der Text an einer ganzen Registerzeile, deren Sub-Section-Buchstabe bzw. Section zum Label passt?"""
        if len(f.zellen) != 1 or f.art not in ("absatz", "zelle"):
            return False
        lay = self.xl.layouts.get(f.blatt)
        if lay is None or not lay.register or f.spalte != lay.inhalt_spalte:
            return False
        rz = self.xl.register_zeile(f.blatt, f.zeile)
        if rz is None or rz.label:
            return False
        if art == "sub":
            return rz.art in ("sub", "sub_titel") and rz.sub == label.rstrip(".")
        return rz.art in ("section", "part") and (rz.section or "").casefold() == label.casefold()

    def _koerper_modus(self, i: int, fund: Fundstelle) -> bool:
        """Die Fundstelle ist die Textzelle einer Registerzeile, das Label steht in einer eigenen Spalte (siehe oben)."""
        lab = self._labels[i] if i < len(self._labels) else Label()
        return lab.art in ("sub", "section", "part") and self._label_passt(fund, lab.art, lab.alt or lab.neu)

    def _label_spalte(self, blatt: str, art: str) -> int | None:
        lay = self.xl.layouts[blatt]
        return lay.untergliederung_spalte if art == "sub" else lay.section_spalte

    def _label_ops(self, i: int, label_laeufe: list[Lauf], fund: Fundstelle, seq: int) -> dict[str, list[Operation]]:
        """Operation fuer ein geaendertes Label in der Sub-Section- oder Section-Spalte der Zeile (dazu ein eigener Absatz fuer
        die Gegenprobe). Der Wert wird in der Schreibweise der Excel ersetzt ("B." -> "B")."""
        lab = self._labels[i]
        gid = next((l.aenderung for l in label_laeufe if l.aenderung), "")
        if not gid:
            return {}
        spalte = self._label_spalte(fund.blatt, lab.art)
        if not spalte:
            return {gid: []}
        roh = self.xl.text(fund.blatt, fund.zeile, spalte)
        s0 = len(roh) - len(roh.lstrip())
        e0 = len(roh.rstrip())
        alt_form = roh[s0:e0]
        neu_form = "" if not lab.neu else (lab.neu.rstrip(".") if lab.art == "sub" else lab.neu)
        if neu_form == alt_form:
            return {}
        op = Operation("ersetzen", fund.blatt, fund.zeile, spalte, s0, e0, neu_form, seq, aenderung=gid)
        laeufe = [Lauf("d", alt_form, gid)] + ([Lauf("a", neu_form, gid)] if neu_form else [])
        if alt_form:
            self._absatz(fund.blatt, [(fund.zeile, spalte, s0, e0)], laeufe)
        return {gid: [op]}

    def _subs_im_neuen_stand(self) -> list[str]:
        """Je PDF-Absatz der Sub-Section-Buchstabe, der im NEUEN Stand gilt (fuer die Sub-Section-Spalte neuer Zeilen)."""
        out: list[str] = []
        aktuell = ""
        for sp, lab in zip(self.spuren, self._labels):
            if not sp.ganz_geloescht:
                if lab.art in ("section", "part") and lab.neu:
                    aktuell = ""
                elif lab.art == "sub" and lab.neu:
                    aktuell = lab.neu.rstrip(".")
            out.append(aktuell)
        return out

    # ------------------------------------------------------------ Bereich der CDE (ADR 0006 Punkt 6)
    def _bereich(self) -> dict[int, str]:
        """PDF-Absaetze vor dem ersten bzw. nach dem letzten sicheren Anker im Hauptblatt (das Blatt mit den meisten Ankern)
        gehoeren nicht zur CDE (Deckblatt, Vorspann, Revisionstabelle, Anhang) -> {Spur: "vorspann" | "anhang"}.
        Ausnahme: der Absatz steht im Hauptblatt VOR (bzw. HINTER) dem Anker -- dann passt er in die Reihenfolge und gehoert
        dazu; hinter dem letzten Anker ausserdem alles ab einem Gliederungselement und alles auf derselben Seite. Die Regel gilt
        erst ab planer.umfang_min_anker Ankern; planer.umfang: register (nur wenn das Hauptblatt ein CDE-Register ist,
        Standard) | immer | aus."""
        pl = _planer_cfg(self.cfg)
        umfang = str(pl.get("umfang", "register"))
        if umfang == "aus" or not self._anker_je_blatt:
            return {}
        haupt = max(self._anker_je_blatt, key=lambda b: (self._anker_je_blatt[b], -list(self.xl.layouts).index(b)))
        self._hauptblatt = haupt
        if umfang == "register" and not self.xl.layouts[haupt].register:
            return {}
        anker = [i for i, z in enumerate(self.zuo) if z.automatisch and z.fund is not None and z.fund.blatt == haupt]
        if len(anker) < int(pl.get("umfang_min_anker", 5)):
            return {}
        lo = self.zuo[anker[0]].fund.schluessel                                   # type: ignore[union-attr]
        hi = self.zuo[anker[-1]].fund.end_schluessel                              # type: ignore[union-attr]
        out: dict[int, str] = {}
        for i in range(anker[0]):
            if not any(f.blatt == haupt and f.schluessel < lo for f, _ in self.zuo[i].kandidaten):
                out[i] = "vorspann"
        # Hinter dem letzten Anker vorsichtiger (ein neuer Absatz am Ende der Spezifikation darf nie still wegfallen): Anhang ist
        # nur, was auf einer spaeteren Seite steht und vor dem kein Gliederungselement (Ueberschrift, Listennummer, Sub-Section)
        # mehr kommt -- sonst geht die Gliederung weiter.
        seite = self.spuren[anker[-1]].bloecke[-1].page
        for i in range(anker[-1] + 1, len(self.spuren)):
            if self._labels[i].art or (self.spuren[i].kind == "ueberschrift" and
                                        (self.spuren[i].erster.section_alt or self.spuren[i].erster.section_neu)):
                break
            if self.spuren[i].erster.page > seite and \
                    not any(f.blatt == haupt and f.schluessel > hi for f, _ in self.zuo[i].kandidaten):
                out[i] = "anhang"
        return out

    _hauptblatt: str = ""

    def seiten_der_cde(self) -> tuple[int, int] | None:
        """Erste und letzte PDF-Seite, auf der etwas zur CDE gehoert (alle Absaetze ausser Vorspann/Anhang), sonst None
        (kein sicherer Bereich: dann gilt jede Luecke als echte Luecke). Seite 0 = Word-Datei ohne Seiten: immer None."""
        if not self._ausserhalb:
            return None
        seiten = [b.page for i, sp in enumerate(self.spuren) if i not in self._ausserhalb for b in sp.bloecke]
        if not seiten or min(seiten) <= 0:
            return None
        return min(seiten), max(seiten)

    def _ausserhalb_spur(self, i: int, wo: str) -> None:
        """Ein PDF-Absatz ausserhalb der CDE: EINE Zeile im Protokoll (zaehlt in der Bilanz), nicht uebernommen, kein offener
        Punkt (Aenderung.bereich = "ausserhalb")."""
        sp = self.spuren[i]
        toks = sp.tokens
        seite = toks[0].page
        ist_kopf = sp.kind == "ueberschrift"
        if sp.ganz_neu:
            typ = "section_neu" if ist_kopf else "absatz_neu"
        elif sp.ganz_geloescht:
            typ = "section_geloescht" if ist_kopf else "absatz_geloescht"
        else:
            arten = {t.status for t in toks if t.status != KEEP}
            typ = "text_geloescht" if arten == {DEL} else "text_eingefuegt" if arten == {ADD} else "text_ersetzt"
        ae = Aenderung(id=f"A{len(self.out) + 1:04d}", typ=typ, seite=seite, bbox=_bbox([t for t in toks if t.page == seite]),
                       section=sp.erster.section_neu or sp.erster.section_alt, alt=sp.alt, neu=sp.neu,
                       woerter=sum(1 for t in toks if t.status != KEEP), block_idx=sp.erster.idx, kontext=(sp.neu or sp.alt)[:400])
        self.out.append(ae)
        self._spur_von[ae.id] = i
        ae.bereich = "ausserhalb"
        self._nicht(ae, "Gehört nicht zur CDE: " + ("Deckblatt/Vorspann vor dem ersten Absatz der CDE." if wo == "vorspann"
                                                    else "steht im PDF hinter dem letzten Absatz der CDE (Anhang)."))
        self._code(ae, "ausserhalb_cde")
        ae.soll_absatz = sp.neu

    @staticmethod
    def _filtern(sp: Spur, cands: list[tuple[Fundstelle, str]]) -> list[tuple[Fundstelle, str]]:
        """Fundstellen, die zaehlen: Ueberschriften nur fuer Ueberschriften; Blaetter ohne Gliederung (z. B. "Historie") nur, wenn
        es im gegliederten Blatt keine gibt; Nummern-/Titelspalte fuer Fliesstext nur, wenn es sonst keine gibt; sichtbare
        Stellen schlagen versteckte; ganze Absaetze schlagen Teilstuecke."""
        if sp.kind != "ueberschrift":
            cands = [c for c in cands if c[0].art != "kopf"]
            if any(not c[0].kopf_spalte for c in cands):
                cands = [c for c in cands if not c[0].kopf_spalte]
        # ein ganzer Excel-Absatz schlaegt Stuecke laengerer Absaetze (gleicher Satz zufaellig mitten in einem anderen Absatz)
        if any(c[0].art != "teil" for c in cands):
            cands = [c for c in cands if c[0].art != "teil"]
        if any(c[0].gliederung for c in cands):
            cands = [c for c in cands if c[0].gliederung]
        if any(not c[0].versteckt for c in cands):
            cands = [c for c in cands if not c[0].versteckt]
        return cands

    def _zugeordnete_zeilen(self) -> set[tuple[str, int, int, int]]:
        """(Blatt, Zeile, Spalte, Zeilenanfang) aller Excel-Textzeilen, die ein PDF-Absatz sicher (auf der Kette) belegt."""
        out: set[tuple[str, int, int, int]] = set()
        for z in self.zuo:
            if not z.automatisch or z.fund is None:
                continue
            for zeile, spalte, a, e in z.fund.zellen:
                t = self.xl.text(z.fund.blatt, zeile, spalte)
                for m in re.finditer(r"[^\n]+", t):
                    if m.start() < max(e, a + 1) and m.end() > a:
                        out.add((z.fund.blatt, zeile, spalte, m.start()))
        return out

    def _unzugeordnet_dazwischen(self, blatt: str, lo: tuple, hi: tuple | None) -> bool:
        for key, (z, s, a, _) in self.idx.verwaltete_zeilen(blatt):
            if key <= lo:
                continue
            if hi is not None and key >= hi:
                break
            if (blatt, z, s, a) not in self._zugeordnet:
                return True
        return False

    def _stufe_an_stelle(self, sp: Spur, fund: Fundstelle, z: Zuordnung, gruende: list[str], codes: list[str],
                         text_aenderung: bool = True) -> str:
        """Sicherheit einer In-Zellen-Aenderung an einer exakten Fundstelle (A1 ist erfuellt)."""
        stufe = HOCH
        if not z.automatisch:
            stufe = MITTEL
            codes.append(z.status)
            gruende.append(_GRUND.get(z.status, ""))
        if fund.versteckt:
            stufe = MITTEL
            codes.append("versteckt")
            gruende.append({"blatt": f"das Blatt '{fund.blatt}' ist ausgeblendet", "zeile": "die Zeile ist ausgeblendet",
                            "spalte": "die Spalte ist ausgeblendet"}.get(fund.versteckt, _GRUND["versteckt"]))
        if sp.kind != "ueberschrift" and fund.kopf_spalte:
            stufe = MITTEL
            codes.append("kopfspalte")
            gruende.append(_GRUND["kopfspalte"])
        lay = self.xl.layouts.get(fund.blatt)
        # Blatt-Aufbau-Sperre (ADR 0004, 7): gilt nicht fuer sichtbare In-Zellen-Aenderungen im Fliesstext an einer eindeutigen
        # (auch ueber die Reihenfolge eindeutigen) exakten Stelle; fuer Nummer/Titel und neue Zeilen bleibt sie.
        if lay is not None and lay.sicherheit != "hoch" and (fund.kopf_spalte or not text_aenderung or not z.automatisch):
            stufe = MITTEL
            codes.append("layout_unsicher")
            gruende.append(f"der Aufbau des Blatts '{fund.blatt}' wurde nur unsicher erkannt")
        sec = sp.erster.section_alt
        if sec and fund.section and not fund.section_unsicher and fund.section != sec and \
                self.cfg["matcher"].get("section_unpassend_stufe", MITTEL) != HOCH:
            stufe = MITTEL
            codes.append("section_unpassend")
            gruende.append("Section passt nicht zur gefundenen Stelle")
        if fund.section_unsicher and fund.kopf_spalte:
            stufe = MITTEL
            codes.append("section_unsicher")
            gruende.append("die Section-Nummer steht als Zahl in der Excel und ist nicht eindeutig (2.1 oder 2.10?)")
        if len(_woerter_norm(sp.alt)) <= self.cfg["schwellen"]["kurz_tokens"] and z.status != "eindeutig" and \
                not (sec and fund.section == sec):
            stufe = MITTEL
            codes.append("kurz")
            gruende.append(_GRUND["kurz"])
        return stufe

    # ------------------------------------------------------------ Aenderungen je Absatz (exakt)
    def _spur_aenderungen(self, sp: Spur) -> tuple[list[Aenderung], list[Lauf]]:
        """Je Gruppe eine Aenderung (IDs in Reihenfolge); dazu die Laeufe mit den echten IDs."""
        a = sp.aufbau
        aes: list[Aenderung] = []
        ids: dict[str, str] = {}
        for gi, runs in enumerate(a.gruppen):
            gid = sp.laeufe[runs[0]].aenderung
            seg = sp.laeufe[runs[0]:runs[-1] + 1]
            alt = zustand(seg, set())
            neu = zustand(seg, {gid})
            toks = [sp.tokens[i] for i in a.gruppen_token[gi]]
            seite = toks[0].page
            b = sp.block_von_token(a.gruppen_token[gi][0])
            typ = "text_ersetzt" if alt and neu else "text_eingefuegt" if neu else "text_geloescht"
            ae = Aenderung(id=f"A{len(self.out) + 1:04d}", typ=typ, seite=seite,
                           bbox=_bbox([t for t in toks if t.page == seite]), section=b.section_neu or b.section_alt,
                           alt=alt, neu=neu, woerter=len(a.zaehl_token[gi]), block_idx=b.idx,
                           kontext=(sp.neu or sp.alt)[:400])
            self.out.append(ae)
            self._spur_von[ae.id] = self._aktuelle_spur if self._aktuelle_spur is not None else -1
            self._sperre[ae.id] = list(self._spur_sperre)
            self._hinweise_nur(ae)
            ids[gid] = ae.id
            aes.append(ae)
        laeufe = umbenennen(sp.laeufe, ids)
        soll = zustand(laeufe, set(ids.values()))
        for ae in aes:
            ae.soll_absatz = soll
        return aes, laeufe

    def _absatz(self, blatt: str, zellen: Iterable[tuple[int, int, int, int]], laeufe: list[Lauf]) -> None:
        self.absaetze.append(Absatz(id=f"P{len(self.absaetze) + 1:04d}", blatt=blatt, zellen=[tuple(z) for z in zellen],
                                    laeufe=laeufe))

    def _text_spur(self, i: int) -> None:
        sp, z = self.spuren[i], self.zuo[i]
        if z.fund is not None and z.art == "neu" and z.automatisch:
            self._schon_eingetragen(sp, z.fund)
            return
        if z.fund is not None and z.art == "alt" and z.status in ("eindeutig", "durch_reihenfolge", "mehrdeutig"):
            self._exakt(i, z.fund)
            return
        # keiner / widerspruch: alter Matcher im Fenster zwischen den Ankern, hoechstens "mittel"
        if self._ersatz_text(i):
            return
        if z.fund is not None and z.art == "alt":            # Widerspruch: die exakte Stelle als Vorschlag
            self._exakt(i, z.fund)
            return
        self._deckel = ["excel_weicht_ab"] if z.status == "keiner" else [z.status]
        for b in sp.bloecke:
            for i0, i1 in self._gruppen(b):
                self._text_gruppe(b, i0, i1)

    def _schon_eingetragen(self, sp: Spur, fund: Fundstelle) -> None:
        aes, laeufe = self._spur_aenderungen(sp)
        for ae in aes:
            ae.ist_absatz, ae.excel_text, ae.score = fund.text, fund.text[:300], 1.0
            self._nicht(ae, "Der neue Text steht schon so in der Excel (die Änderung ist offenbar schon eingetragen).", fund.ziel)
            self._code(ae, "schon_eingetragen")

    def _exakt(self, i: int, fund: Fundstelle) -> None:
        sp, z = self.spuren[i], self.zuo[i]
        aes, laeufe = self._spur_aenderungen(sp)
        seq = sp.erster.idx * 1000
        koerper = self._koerper_modus(i, fund)
        n = self._labels[i].n if koerper else 0            # Register: das Label steht in einer eigenen Spalte
        so = ops_fuer_spur(laeufe[n:], fund, self.xl.text, self.cfg, seq)
        if not so.vergleich.gleich:                       # sollte nicht vorkommen (Fundstelle ist exakt)
            for ae in aes:
                self._nicht(ae, "Der alte Absatz steht so nicht in der Excel. Bitte von Hand prüfen.", fund.ziel)
                self._code(ae, "a1_" + (so.vergleich.art or "abweichung"))
            return
        self._absatz(fund.blatt, fund.zellen, laeufe[n:])
        label_ops = self._label_ops(i, laeufe[:n], fund, seq) if koerper else {}
        for ae in aes:
            ae.ist_absatz, ae.excel_text, ae.score, ae.ziel = fund.text, fund.text[:300], 1.0, fund.ziel
            gruende: list[str] = []
            codes: list[str] = []
            stufe = self._stufe_an_stelle(sp, fund, z, gruende, codes)
            ops = so.ops.get(ae.id, []) + label_ops.get(ae.id, [])
            if ae.id in label_ops:
                # Sub-Section-Buchstabe / Section-Nummer stehen auch in den Zeilen darunter (dort bleiben sie): ein Mensch prueft
                stufe = MITTEL
                codes.append("label_spalte")
                gruende.append(_codes.satz("label_spalte"))
                if not label_ops[ae.id]:
                    codes.append("naht_unklar")
            for c in so.probleme.get(ae.id, []):
                stufe = MITTEL
                codes.append(c)
                gruende.append(_GRUND.get(c, ""))
            if not ops:
                self._nicht(ae, "Die Stelle der Änderung im Excel-Absatz konnte nicht sicher bestimmt werden. Bitte von Hand "
                                "prüfen.", fund.ziel)
                for c in codes or ["naht_unklar"]:
                    self._code(ae, c)
                continue
            for o in ops:
                o.seq = seq + sp.aufbau.gruppen_token[aes.index(ae)][0]
            ae.ops = ops
            ae.ziel = f"{fund.blatt}!{get_column_letter(ops[0].spalte)}{ops[0].zeile}"
            self._setze(ae, stufe, "; ".join(g for g in dict.fromkeys(gruende) if g), codes)

    def _a1_codes(self, alt: str, excel: str) -> list[str]:
        v = vergleiche(alt, excel, frozenset(), self.cfg.get("normalform"))
        return ["a1_" + v.art] if v.art else []

    def _ersatz_text(self, i: int) -> bool:
        """keiner/widerspruch: ein Vorschlag ueber den alten Matcher im Fenster. True, wenn er etwas geplant hat."""
        sp, z = self.spuren[i], self.zuo[i]
        schw = self.cfg["schwellen"]
        trs = [self.treffer(b) for b in sp.bloecke]
        if not any(tr is not None and tr.score >= schw["mittel"] for tr in trs):
            return False
        for b, tr in zip(sp.bloecke, trs):
            basis = ["excel_weicht_ab"] if z.status == "keiner" else [z.status, "excel_weicht_ab"]
            if tr is not None:
                a1 = self._a1_codes(b.alt_text if len(sp.bloecke) > 1 else sp.alt, tr.kandidat.text)
                art = a1[0][3:] if a1 else ""
                self._a1_satz = ""
                self._ist = tr.kandidat.text
                self._soll = sp.neu
                self._deckel = basis + a1
                if art:
                    self._a1_satz = f"Der alte Absatz aus dem PDF steht so nicht in der Excel (Abweichung: {_A1_SATZ.get(art, art)})."
            else:
                self._deckel = basis
            if b.hat_aenderung:
                n0 = len(self.out)
                if b.ganz_geloescht:
                    self._geloescht(b)
                else:
                    for i0, i1 in self._gruppen(b):
                        self._text_gruppe(b, i0, i1)
                    self._vorschlag_unscharf(b, tr, self.out[n0:])
                self._ersatz_absatz(b, tr, self.out[n0:])
        return True

    def _vorschlag_unscharf(self, b: PdfBlock, tr: Treffer | None, aes: list[Aenderung]) -> None:
        """Nicos Entscheidung (ADR 0006 Punkt 7): ist der Absatz in der Excel verortet, die Excel weicht aber anderswo im Absatz
        vom PDF ab, wird eine Aenderung, die Wort fuer Wort nicht passte (sonst "nicht uebernommen"), ein fertiger Vorschlag:
        nur die markierte Stelle aendert sich (zeichengenau ueber die Normalform ausgerichtet), Excel-eigene Zeichen bleiben.
        Nie automatisch ("pruefen", Codes excel_weicht_ab + a1_*).

        Auftrag 004 (Vorschlaege): Auch ein Vorschlag des alten Matchers (Wort fuer Wort) wird durch den zeichengenauen
        ersetzt, wenn es ihn gibt -- der alte Matcher ersetzt ganze Woerter samt Satzzeichen ("Pumpen." -> "Pumpen nachweislich."),
        liest Leerraum an der Naht anders als Word und kennt keine Klammern, deren Zeichen allein gestrichen sind. Gibt es keinen
        zeichengenauen Vorschlag, bleibt der alte nur, wenn `planer.vorschlag_wortgenau` es erlaubt (Standard: nein, besser kein
        Vorschlag als ein falscher)."""
        offen = [ae for ae in aes if ae.status == NICHT and not ae.ops and not ae.pruefcodes]
        alte = [ae for ae in aes if ae.status == PRUEFEN and ae.ops and "excel_weicht_ab" in ae.pruefcodes
                and all(o.art in ("ersetzen", "einfuegen") for o in ae.ops)]
        if tr is None or not (offen or alte) or tr.kandidat.art == "kopf":
            return
        k = tr.kandidat
        wortgenau = bool(_planer_cfg(self.cfg).get("vorschlag_wortgenau", False))

        def verwerfen() -> None:
            if not wortgenau:
                for ae in alte:
                    self._vorschlag_verwerfen(ae)
        aufbau = _aufbau(b.tokens, "X")
        # Eine Aenderung des alten Matchers = zusammenhaengende markierte Woerter; der Aufbau kann daraus mehrere Gruppen machen
        # (gestrichene Klammern um einen unveraenderten Wert: "[" und "]" sind zwei Gruppen). Jede Gruppe gehoert zu genau einer.
        bereiche = list(self._gruppen(b))
        if len(bereiche) != len(aes):
            return verwerfen()
        je_aenderung: dict[int, list[int]] = {n: [] for n in range(len(aes))}
        for gi, toks in enumerate(aufbau.gruppen_token):
            ziele = {n for n, (i0, i1) in enumerate(bereiche) for t in toks if i0 <= t < i1}
            if len(ziele) != 1:
                return verwerfen()
            je_aenderung[next(iter(ziele))].append(gi)
        ids = {f"X{gi + 1}": f"{aes[n].id}~{gi}" for n, gis in je_aenderung.items() for gi in gis}
        laeufe = umbenennen(aufbau.laeufe, ids)
        fund = Fundstelle(blatt=k.blatt, zellen=((k.zeile, k.spalte, k.start, k.end),), art="absatz",
                          schluessel=self.idx.schluessel(k.blatt, k.zeile, k.spalte, k.start), text=k.text,
                          end_schluessel=self.idx.schluessel(k.blatt, k.zeile, k.spalte, k.end))
        so = ops_fuer_spur(laeufe, fund, self.xl.text, self.cfg, b.idx * 1000, unscharf=True)
        a1 = self._a1_codes(b.alt_text, k.text)

        def ops_von(n: int) -> list[Operation] | None:
            """Alle Operationen der Gruppen dieser Aenderung (None = eine Gruppe ohne sichere Stelle)."""
            out: list[Operation] = []
            for gi in je_aenderung[n]:
                gid = f"{aes[n].id}~{gi}"
                o = so.ops.get(gid, [])
                if not o or so.probleme.get(gid):
                    return None
                out += o
            for o in out:
                o.aenderung = aes[n].id
            return sorted(out, key=lambda o: (o.zeile, o.spalte, o.start)) or None

        for ae in alte:
            ops = ops_von(aes.index(ae))
            if ops and self._ops_unbedenklich(ae, ops):
                ae.ops = ops
                ae.ziel = f"{k.blatt}!{get_column_letter(ops[0].spalte)}{ops[0].zeile}"
                ae.ist_absatz = k.text
            elif not wortgenau:
                self._vorschlag_verwerfen(ae)
        for ae in offen:
            ops = ops_von(aes.index(ae))
            if not ops:
                continue
            ae.status, ae.konfidenz, ae.grund = PRUEFEN, MITTEL, ""
            ae.ops = ops
            ae.ziel = f"{k.blatt}!{get_column_letter(ops[0].spalte)}{ops[0].zeile}"
            ae.ist_absatz = k.text
            self._hinweis(ae, "Die Excel weicht im Absatz an einer anderen Stelle vom PDF ab; der Vorschlag ändert nur die "
                              "markierte Stelle. Bitte prüfen und freigeben.")
            doppelt = bool(self._deckel and "excel_weicht_ab" in self._deckel)       # der Deckel nennt den Grund schon
            self._setze(ae, MITTEL, "" if doppelt else _GRUND["excel_weicht_ab"], ["excel_weicht_ab"] + a1)
        if all(ae.ops and ae.status == PRUEFEN for ae in aes):
            self._absatz_an_excel_ausrichten(k, so.excel_text, laeufe, {f"{aes[n].id}~{gi}": aes[n].id
                                                                       for n, gis in je_aenderung.items() for gi in gis}, aes)

    def _absatz_an_excel_ausrichten(self, k: Any, excel: str, laeufe: list[Lauf], ids: dict[str, str],
                                    aes: list[Aenderung]) -> None:
        """Der PDF-Absatz fuer die Markierung der Vorschlaege bei abweichender Excel (Auftrag 004).

        Ohne PDF-Absatz markiert der Applier nur nach den Operationen: Leerraum am Rand eines neuen Stuecks bleibt unmarkiert,
        waehrend Word ihn mit unterstreicht (" Satz" als Ganzes neu). Die Projektion der PDF-Laeufe markiert genau wie das PDF --
        sie braucht aber Laeufe, deren alter Text Zeichen fuer Zeichen die Excel ist. Hier werden die PDF-Laeufe an die Excel
        angelegt: bleibende und gestrichene Stuecke tragen den Text der Excel (auch ihre Abweichung), neue Stuecke bleiben
        wie im PDF. Der Absatz wird nur registriert, wenn (1) die alte Fassung gleich der Excel ist und (2) die Projektion
        denselben neuen Text ergibt wie die Operationen -- sonst bleibt es bei den Operationen."""
        nf = self.cfg.get("normalform")
        alt, lage = _positionen(laeufe, set())
        abb = _unscharfe_abbildung(alt, excel, nf)
        out: list[Lauf] = []
        ende = 0                                              # bis hierher ist die Excel in `out` untergebracht
        offene_trenner: list[int] = []                        # Indizes (in out) von Trennern seit dem letzten Stueck mit Zeichen

        def anfuegen(art: str, text: str, aenderung: str = "", optional: bool = False) -> None:
            if art in ("k", "s") and out and out[-1].art == art and out[-1].aenderung == aenderung and not optional:
                out[-1] = Lauf(art, out[-1].text + text, aenderung, False)
            else:
                out.append(Lauf(art, text, aenderung, optional))

        for i, l in enumerate(laeufe):
            if l.art == "a":
                out.append(Lauf("a", l.text, ids.get(l.aenderung, l.aenderung), l.optional))
                continue
            if i not in lage:
                continue                                      # nicht sichtbar in der alten Fassung (zusammengelegter Trenner)
            a0, e0 = lage[i]
            ziel = [abb[p] for p in range(a0, e0) if p in abb]
            if l.art == "d" and not l.text.strip():
                # gestrichenes Leerzeichen als eigener Lauf: es ist das Leerzeichen der Excel an dieser Stelle (sonst entfaellt es)
                rest = len(l.text)
                nach = ende
                while rest > 0 and nach < len(excel) and excel[nach].isspace():
                    nach, rest = nach + 1, rest - 1
                if nach > ende:
                    out.append(Lauf("d", excel[ende:nach], ids.get(l.aenderung, l.aenderung), l.optional))
                    ende = nach
                continue
            if l.art == "s" or not ziel:
                if l.art == "d":
                    return                                    # gestrichener Text ohne Gegenstueck in der Excel: nicht sicher
                offene_trenner.append(len(out))
                out.append(Lauf("s", "", "", l.optional))      # Text folgt, sobald klar ist, was die Excel dort hat
                continue
            xs, xe = min(ziel), max(ziel) + 1
            if l.art == "d":
                # Word streicht das Leerzeichen am Rand eines gestrichenen Stuecks mit; es steht in der Excel zwischen den Woertern
                # und gehoert dann zum gestrichenen Stueck (nicht zur Luecke), sonst entstuende ein Leerzeichen zu viel bzw. zu wenig
                vorn = len(l.text) - len(l.text.lstrip())
                hinten = len(l.text) - len(l.text.rstrip())
                while vorn > 0 and xs > ende and excel[xs - 1].isspace():
                    xs, vorn = xs - 1, vorn - 1
                while hinten > 0 and xe < len(excel) and excel[xe].isspace():
                    xe, hinten = xe + 1, hinten - 1
            if xs < ende:
                return                                        # Reihenfolge nicht eindeutig
            luecke = excel[ende:xs]
            if luecke:
                if offene_trenner and not luecke.strip():
                    j = offene_trenner[0]
                    out[j] = Lauf("s", luecke, "", out[j].optional)
                else:
                    out.insert(offene_trenner[0] if offene_trenner else len(out), Lauf("k", luecke))
                    if offene_trenner:
                        offene_trenner = [j + 1 for j in offene_trenner]
            offene_trenner = []
            if l.art == "d":
                kern = excel[xs:xe]
                anfuegen("d", kern, ids.get(l.aenderung, l.aenderung), l.optional)
            else:
                anfuegen("k", excel[xs:xe], "", l.optional)
            ende = xe
        rest = excel[ende:]
        if rest:
            if offene_trenner and not rest.strip():
                j = offene_trenner[0]
                out[j] = Lauf("s", rest, "", out[j].optional)
            else:
                anfuegen("k", rest)
        out = [l for l in out if l.text]
        aktiv = set(ids.values())
        alt_neu, opt_alt = zustand_mit_optional(out, set())
        if not vergleiche(alt_neu, excel, opt_alt, nf).gleich:
            return
        res = mk.projiziere([excel], out, aktiv, self.cfg)
        if res is None:
            return
        zelle = self.xl.text(k.blatt, k.zeile, k.spalte)
        if zelle[k.start:k.end] != excel:
            return
        projiziert = zelle[:k.start] + "".join(z[0] for z in res[0] if z[1] != mk.ALT) + zelle[k.end:]
        ops = [o for ae in aes for o in ae.ops]
        if any((o.blatt, o.zeile, o.spalte) != (k.blatt, k.zeile, k.spalte) or o.art not in ("ersetzen", "einfuegen") for o in ops):
            return
        try:
            geplant = "".join(z[0] for z in mk.markiere_ops([mk.Stueck(mk.NORMAL, zelle)], ops, mk.SICHT_FLACH, {}, self.cfg).zeichen
                              if z[1] != mk.ALT)
        except ValueError:
            return
        if " ".join(projiziert.split()) != " ".join(geplant.split()):
            return
        self._absatz(k.blatt, [(k.zeile, k.spalte, k.start, k.end)], out)

    def _ops_unbedenklich(self, ae: Aenderung, ops: list[Operation]) -> bool:
        """Neuer Text der Operationen: keine Steuerzeichen, nichts Verdaechtiges, Zahl-Zellen bleiben Zahlen (wie `_setze`)."""
        p = neutext_probleme([o.text for o in ops], self.cfg)
        if p["steuerzeichen"] or p["verdaechtig"]:
            return False
        for op in ops:
            if self.xl.ist_zahl(op.blatt, op.zeile, op.spalte):
                return False
        return True

    def _vorschlag_verwerfen(self, ae: Aenderung) -> None:
        """Kein sicherer Vorschlag: die Aenderung bleibt auf "pruefen", aber ohne fertige Operationen (Regel 27, besser kein
        Vorschlag als ein falscher). Freigeben schreibt dann nichts; der Mensch traegt die Stelle von Hand ein."""
        ae.ops = []
        ae.status, ae.konfidenz = PRUEFEN, MITTEL
        self._hinweis(ae, _codes.satz("kein_vorschlag"))
        self._code(ae, "kein_vorschlag")

    def _ersatz_absatz(self, b: PdfBlock, tr: Treffer | None, aes: list[Aenderung]) -> None:
        """PDF-Absatz fuer die Gegenprobe A3 auch auf dem Weg ueber den alten Matcher -- aber nur, wenn die Laeufe des Blocks
        genau zur Excel-Stelle passen (A1 gleich) und alle Operationen in dieser Spanne liegen. Sonst keinen Absatz (dann
        greift beim automatischen Eintragen die Invariante "ohne_absatz")."""
        if tr is None or not aes:
            return
        k = tr.kandidat
        aufbau = _aufbau(b.tokens, "X")
        if len(aufbau.gruppen) != len(aes):
            return
        ids = {f"X{n + 1}": ae.id for n, ae in enumerate(aes)}
        laeufe = umbenennen(aufbau.laeufe, ids)
        if k.art == "kopf":
            zellen = list(dict.fromkeys((t.zeile, t.spalte) for t in k.tokens))
            bereiche = [(z, s, 0, len(self.xl.text(k.blatt, z, s))) for z, s in zellen]
        else:
            bereiche = [(k.zeile, k.spalte, k.start, k.end)]
        excel = " ".join(self.xl.text(k.blatt, z, s)[a:e] for z, s, a, e in bereiche)
        alt, optional = zustand_mit_optional(laeufe, set())
        if not vergleiche(alt, excel, optional, self.cfg.get("normalform")).gleich:
            return
        for ae in aes:
            for o in ae.ops:
                if o.art not in ("ersetzen", "einfuegen") or not any(
                        (o.blatt, o.zeile, o.spalte) == (k.blatt, z, s) and a <= o.start and o.ende <= e for z, s, a, e in bereiche):
                    return
        self._absatz(k.blatt, bereiche, laeufe)

    # ------------------------------------------------------------ Geloeschte Absaetze
    def _geloescht_spur(self, i: int) -> None:
        sp, z = self.spuren[i], self.zuo[i]
        if z.fund is not None and z.art == "alt" and (z.status in ("eindeutig", "durch_reihenfolge", "mehrdeutig")
                                                       or (z.status == "widerspruch" and not self._ersatz_moeglich(i))):
            fund = z.fund
            aes, laeufe = self._spur_aenderungen(sp)
            ist_kopf = sp.kind == "ueberschrift"
            for ae in aes:
                ae.typ = "section_geloescht" if ist_kopf else "absatz_geloescht"
            if len(aes) != 1:                                  # ein ganz gestrichener Absatz ist genau eine Gruppe
                for ae in aes:
                    self._nicht(ae, "Der gestrichene Absatz konnte nicht als Ganzes zugeordnet werden. Bitte von Hand prüfen.",
                                fund.ziel)
                return
            ae = aes[0]
            ae.alt = sp.alt
            belegt = {(z0, s0): r for (b0, z0, s0), r in self._geloescht_bis.items() if b0 == fund.blatt}
            ops, bereiche = ganz_loeschen(fund, self.xl.text, ae.id, sp.erster.idx * 1000, belegt)
            # Die Operationen nehmen den Zeilenumbruch selbst mit (kein Aufraeumen leerer Zeilen durch den Applier noetig,
            # sonst stimmen die Stellen der anderen Absaetze derselben Zelle fuer die Gegenprobe A3 nicht mehr).
            for z0, s0, a0, e0 in bereiche:
                self._geloescht_bis[(fund.blatt, z0, s0)] = self._geloescht_bis.get((fund.blatt, z0, s0), []) + [(a0, e0)]
            # Register: nur die Textzelle wird markiert; Sub-Section-Buchstabe und Section-Nummer der Zeile bleiben (wie bei Nico)
            n = self._labels[i].n if self._koerper_modus(i, fund) else 0
            self._absatz(fund.blatt, bereiche, laeufe[n:])
            ae.ist_absatz, ae.excel_text, ae.score, ae.ziel, ae.ops = fund.text, fund.text[:300], 1.0, fund.ziel, ops
            gruende: list[str] = []
            codes: list[str] = []
            stufe = self._stufe_an_stelle(sp, fund, z, gruende, codes, text_aenderung=not ist_kopf)
            if fund.art == "teil":
                # der gestrichene Absatz ist nur ein Stueck eines laengeren Excel-Absatzes: der Rest bleibt stehen -- ob das so
                # gemeint ist, entscheidet ein Mensch
                stufe = MITTEL
                codes.append("teil_absatz")
                gruende.append("In der Excel steht im selben Absatz mehr Text, der im PDF nicht gestrichen ist. Nur der gestrichene "
                               "Text wird entfernt, der Rest bleibt stehen.")
                self._hinweis(ae, gruende[-1])
            self._setze(ae, stufe, "; ".join(g for g in dict.fromkeys(gruende) if g), codes)
            return
        schon = self._schon_gestrichen(i) if z.fund is None else None
        if schon is not None:              # zweiter Lauf auf dem markierten Ergebnis: der Absatz steht schon durchgestrichen da
            aes, _ = self._spur_aenderungen(sp)
            ziel = f"{schon[0]}!{get_column_letter(schon[2])}{schon[1]}"
            for ae in aes:
                ae.typ = "section_geloescht" if sp.kind == "ueberschrift" else "absatz_geloescht"
                self._nicht(ae, "Der gestrichene Absatz steht schon durchgestrichen in der Excel (die Änderung ist offenbar "
                                "schon eingetragen).", ziel)
                self._code(ae, "schon_eingetragen")
            return
        self._deckel = ["excel_weicht_ab"] if z.status == "keiner" else [z.status, "excel_weicht_ab"]
        for b in sp.bloecke:
            tr = self.treffer(b)
            self._ist = tr.kandidat.text if tr is not None else ""
            self._deckel = (["excel_weicht_ab"] if z.status == "keiner" else [z.status, "excel_weicht_ab"]) + \
                (self._a1_codes(b.alt_text, tr.kandidat.text) if tr is not None else [])
            n0 = len(self.out)
            self._geloescht(b)
            self._ersatz_absatz(b, tr, self.out[n0:])

    def _schon_gestrichen(self, i: int) -> tuple[str, int, int] | None:
        """Zelle, deren durchgestrichener Text (Markierung einer Runde) genau der ganz gestrichene PDF-Absatz ist; im Register
        auch nur der Text ohne Label (das Label steht durchgestrichen in seiner eigenen Spalte)."""
        if not self.xl.gestrichen:
            return None
        sp, lab = self.spuren[i], self._labels[i]
        nf = self.cfg.get("normalform")
        texte = {normalform(sp.alt, nf)[0]}
        if lab.art in ("sub", "section", "part"):
            texte.add(normalform(zustand(sp.laeufe[lab.n:], set()), nf)[0])
        texte.discard("")
        for key, t in sorted(self.xl.gestrichen.items()):
            if not self._erlaubt(key[0], key[1]):
                continue
            if normalform(t, nf)[0] in texte and not self.xl.text(*key).strip():
                return key
        return None

    def _ersatz_moeglich(self, i: int) -> bool:
        schw = self.cfg["schwellen"]
        return any((tr := self.treffer(b)) is not None and tr.score >= schw["mittel"] for b in self.spuren[i].bloecke)

    def _loeschungen_koppeln(self) -> None:
        """Eine gestrichene Section und der gestrichene Inhalt darunter sind eine Gruppe (`gruppe`): der Status gilt fuer alle
        gemeinsam.

        Sonst bliebe eine verwaiste Zeile zurueck (Nummer und Titel weg, alter Inhalt steht noch da) oder ein Teil des
        Inhalts verschwindet, der Rest nicht. Ist ein Mitglied der Gruppe nicht sicher (pruefen) oder nicht eintragbar,
        bleiben alle anderen auf pruefen. Die Vorschlaege (Operationen) bleiben stehen."""
        je_spur: dict[int, list[Aenderung]] = {}
        for ae in self.out:
            if ae.typ in ("absatz_geloescht", "section_geloescht"):
                je_spur.setdefault(self._spur_von.get(ae.id, -1), []).append(ae)
        for kopf in [ae for ae in self.out if ae.typ == "section_geloescht" and not ae.bereich]:
            i = self._spur_von.get(kopf.id, -1)
            if i < 0:
                continue
            folge: list[Aenderung] = []
            k = i + 1
            while k < len(self.spuren) and self.spuren[k].ganz_geloescht and self.spuren[k].kind == "absatz":
                folge += [a for a in je_spur.get(k, []) if a.typ == "absatz_geloescht" and not a.bereich]
                k += 1
            if not folge:
                continue
            gruppe = [kopf] + folge
            alte = {m.gruppe for m in gruppe if m.gruppe}
            gruppe += [a for a in self.out if a.gruppe in alte and a not in gruppe]     # z. B. neuer Absatz in geleerter Zelle
            for m in gruppe:
                m.gruppe = kopf.id
            if all(m.status == UEBERNOMMEN for m in gruppe):
                continue
            for m in gruppe:
                if m.status == UEBERNOMMEN:
                    self._herabstufen(m, "Gehört zu einer gestrichenen Section, deren Löschung nicht ganz sicher ist: Section und "
                                         "gestrichener Inhalt werden nur zusammen gelöscht, sonst bliebe ein Rest stehen "
                                         "(z. B. eine Zeile mit altem Inhalt ohne Nummer und Titel).",
                                      "Section und gestrichener Inhalt gehören zusammen. Bitte alles zusammen prüfen.")
                    self._code(m, "gruppe")


    def _herabstufen(self, ae: Aenderung, grund: str, hinweis: str) -> None:
        """Aenderung von 'uebernommen' auf 'pruefen' (der Vorschlag mit seinen Operationen bleibt stehen)."""
        ae.status, ae.konfidenz = PRUEFEN, MITTEL
        ae.grund = (ae.grund + "; " if ae.grund else "") + grund
        self._hinweis(ae, hinweis)

    # ------------------------------------------------------------ Text in Absatz/Ueberschrift
    def _text_gruppe(self, b: PdfBlock, i0: int, i1: int) -> None:
        toks = b.tokens[i0:i1]
        alt = " ".join(t.old for t in toks if t.old)
        neu = " ".join(t.new for t in toks if t.new)
        typ = "text_ersetzt" if alt and neu else "text_eingefuegt" if neu else "text_geloescht"
        ae = self._basis(b, typ, toks, alt, neu)
        tr = self.treffer(b)
        schw = self.cfg["schwellen"]
        if tr is None or tr.score < schw["mittel"]:
            self._kein_treffer(ae, b)
            return
        cand = tr.kandidat
        ae.score = tr.score
        ae.ziel = cand.ziel
        ae.excel_text = cand.text[:300]
        stufe = tr.stufe
        gruende = list(tr.gruende)
        alt_idx = [i for i in range(i0, i1) if b.tokens[i].old]
        d_idx = [i for i in alt_idx if norm(b.tokens[i].old)]
        leer_idx = [i for i in alt_idx if not norm(b.tokens[i].old)]     # alte Woerter ohne Normform ("•", ".")
        seq = b.idx * 1000 + i0
        neu_w = [t.new for t in toks if t.new]

        if d_idx:
            fehlt = [i for i in d_idx if i not in tr.zuordnung]
            if fehlt:
                if self._schon_da(b, i0, i1, tr, neu_w):
                    self._nicht(ae, "Der neue Text steht schon in der Excel (die Änderung ist offenbar schon eingetragen).", cand.ziel)
                    self._code(ae, "schon_eingetragen")
                elif not neu_w and len(fehlt) == len(d_idx):
                    self._nicht(ae, "Der gestrichene Text steht nicht mehr in der Excel (vielleicht schon entfernt, vielleicht "
                                    "anders formuliert). Bitte von Hand prüfen.", cand.ziel)
                else:
                    self._nicht(ae, "Der alte Text steht in der Excel anders (abweichende Formulierung). "
                                    "Bitte von Hand prüfen.", cand.ziel)
                return
            js = sorted(tr.zuordnung[i] for i in d_idx)
            xt = [cand.tokens[j] for j in js]
            zellen = {(t.blatt, t.zeile, t.spalte) for t in xt}
            if len(zellen) > 1:
                if not self._kopf_mehrzellig(ae, b, i0, i1, cand, xt, neu, seq, stufe):
                    self._nicht(ae, "Die Änderung betrifft mehrere Excel-Zellen.", cand.ziel)
                return
            if js[-1] - js[0] + 1 != len(js):
                stufe = MITTEL
                gruende.append("in der Excel stehen zusätzliche Wörter zwischen dem alten Text")
            blatt, zeile, spalte = xt[0].blatt, xt[0].zeile, xt[0].spalte
            zt = self.xl.text(blatt, zeile, spalte)
            start, ende = xt[0].start, xt[-1].end
            vorn = [b.tokens[i].old for i in leer_idx if i < d_idx[0]]
            hinten = [b.tokens[i].old for i in leer_idx if i > d_idx[-1]]
            if vorn or hinten:       # gestrichene Aufzaehlungszeichen/Satzzeichen am Rand gehoeren zum Bereich
                rand = self._rand_leer(zt, start, ende, vorn, hinten, cand)
                if rand is None:
                    self._nicht(ae, "Ein gestrichenes Zeichen ohne Buchstaben (z. B. Aufzählungszeichen) wurde in der Excel "
                                    "nicht sicher gefunden. Bitte von Hand prüfen.", cand.ziel)
                    return
                start, ende = rand
            if neu:
                op = Operation("ersetzen", blatt, zeile, spalte, start, ende, neu, seq, aenderung=ae.id)
            else:
                s, e = self._luecke(zt, start, ende)
                op = Operation("ersetzen", blatt, zeile, spalte, s, e, "", seq, aenderung=ae.id)
            ae.ops = [op]
            ae.ziel = self._ziel(xt[0])
        elif leer_idx:
            # Nur Woerter ohne Normform ("•", "."): ihre Stelle aus dem Rohtext zwischen zwei sicheren Nachbarwoertern bestimmen
            span = self._leer_span(b, i0, i1, tr, cand)
            if span is None:
                self._nicht(ae, "Ein gestrichenes Zeichen ohne Buchstaben (z. B. Aufzählungszeichen) wurde in der Excel "
                                "nicht sicher gefunden. Bitte von Hand prüfen.", cand.ziel)
                return
            blatt, zeile, spalte, start, ende = span
            zt = self.xl.text(blatt, zeile, spalte)
            if neu:
                op = Operation("ersetzen", blatt, zeile, spalte, start, ende, neu, seq, aenderung=ae.id)
            else:
                s, e = self._luecke(zt, start, ende)
                op = Operation("ersetzen", blatt, zeile, spalte, s, e, "", seq, aenderung=ae.id)
            ae.ops = [op]
            ae.ziel = f"{blatt}!{get_column_letter(spalte)}{zeile}"
        else:
            # reine Einfuegung: nach dem naechsten vorangehenden, zugeordneten Wort
            vor = next((i for i in range(i0 - 1, -1, -1) if i in tr.zuordnung), None)
            if vor is not None:
                xt = cand.tokens[tr.zuordnung[vor]]
                if vor != i0 - 1:
                    stufe = MITTEL
                    gruende.append("Einfügestelle nur ungefähr bestimmbar")
                op = Operation("einfuegen", xt.blatt, xt.zeile, xt.spalte, xt.end, xt.end, " " + neu, seq, aenderung=ae.id)
            else:
                nach = next((i for i in range(i1, len(b.tokens)) if i in tr.zuordnung), None)
                if nach is None:
                    self._nicht(ae, "Keine feste Stelle für die Einfügung gefunden.", cand.ziel)
                    return
                xt = cand.tokens[tr.zuordnung[nach]]
                op = Operation("einfuegen", xt.blatt, xt.zeile, xt.spalte, xt.start, xt.start, neu + " ", seq, aenderung=ae.id)
            if self._schon_da(b, i0, i1, tr, neu_w):
                self._nicht(ae, "Der neue Text steht schon in der Excel (die Änderung ist offenbar schon eingetragen).",
                            self._ziel(xt))
                self._code(ae, "schon_eingetragen")
                return
            ae.ops = [op]
            ae.ziel = self._ziel(xt)
        self._setze(ae, stufe, "; ".join(gruende))

    # ------------------------------------------------------------ Hilfen fuer Woerter ohne Normform und Idempotenz
    def _schon_da(self, b: PdfBlock, i0: int, i1: int, tr: Treffer, neu_w: list[str]) -> bool:
        """Steht der neue Text schon direkt an der Einfuegestelle? (zweiter Lauf auf dem eigenen Ergebnis)"""
        if not neu_w:
            return False
        cand = tr.kandidat
        zuo = tr.zuordnung
        vor = next((i for i in range(i0 - 1, -1, -1) if i in zuo), None)
        if vor is not None and cand.art != "kopf":
            xt = cand.tokens[zuo[vor]]
            rest = re.findall(r"\S+", self.xl.text(xt.blatt, xt.zeile, xt.spalte)[xt.end:cand.end])
            if _woerter_gleich(rest[:len(neu_w)], neu_w):
                return True
        nach = next((i for i in range(i1, len(b.tokens)) if i in zuo), None)
        if nach is not None and cand.art != "kopf":
            xt = cand.tokens[zuo[nach]]
            davor = re.findall(r"\S+", self.xl.text(xt.blatt, xt.zeile, xt.spalte)[cand.start:xt.start])
            if _woerter_gleich(davor[-len(neu_w):], neu_w):
                return True
        return False

    def _rand_leer(self, zt: str, s: int, e: int, vorn: list[str], hinten: list[str], cand: Kandidat) -> tuple[int, int] | None:
        """Erweitert den Bereich [s, e) um direkt angrenzende Woerter ohne Normform, falls sie wirklich dort stehen."""
        lo = cand.start if cand.art != "kopf" else 0
        hi = cand.end if cand.art != "kopf" else len(zt)
        for w in reversed(vorn):
            m = re.search(r"(\S+)\s*$", zt[:s])
            if not m or m.group(1) != w or m.start(1) < lo:
                return None
            s = m.start(1)
        for w in hinten:
            m = re.match(r"\s*(\S+)", zt[e:])
            if not m or m.group(1) != w or e + m.end() > hi:
                return None
            e += m.end()
        return s, e

    def _leer_span(self, b: PdfBlock, i0: int, i1: int, tr: Treffer, cand: Kandidat) -> tuple[str, int, int, int, int] | None:
        """Stelle (Blatt, Zeile, Spalte, Start, Ende) der alten Woerter ohne Normform einer Gruppe.

        Der Rohtext zwischen den beiden naechsten sicher zugeordneten Nachbarwoertern muss genau die alten Woerter des
        PDF enthalten -- sonst gibt es keine Stelle (None)."""
        if cand.art == "kopf":
            return None
        zuo = tr.zuordnung
        vor = next((i for i in range(i0 - 1, -1, -1) if i in zuo), None)
        nach = next((i for i in range(i1, len(b.tokens)) if i in zuo), None)
        zt = cand.zelle_text
        s = cand.tokens[zuo[vor]].end if vor is not None else cand.start
        e = cand.tokens[zuo[nach]].start if nach is not None else cand.end
        if e < s:
            return None
        von = vor + 1 if vor is not None else 0
        bis = nach if nach is not None else len(b.tokens)
        idx = [i for i in range(von, bis) if b.tokens[i].old]
        excel = [(m.group(), s + m.start(), s + m.end()) for m in re.finditer(r"\S+", zt[s:e])]
        if [w for w, _, _ in excel] != [b.tokens[i].old for i in idx]:
            return None
        ks = [k for k, i in enumerate(idx) if i0 <= i < i1]
        if not ks:
            return None
        return cand.blatt, cand.zeile, cand.spalte, excel[ks[0]][1], excel[ks[-1]][2]

    @staticmethod
    def _luecke(text: str, s: int, e: int) -> tuple[int, int]:
        """Beim Loeschen eines Wortbereichs ein angrenzendes Leerzeichen mitnehmen."""
        if s > 0 and text[s - 1] == " ":
            while s > 0 and text[s - 1] == " ":
                s -= 1
        else:
            while e < len(text) and text[e] == " ":
                e += 1
        return s, e

    def _kopf_mehrzellig(self, ae, b, i0, i1, cand, xt, neu, seq, stufe) -> bool:
        """Ueberschrift, bei der Nummer UND Titel (zwei Zellen) geaendert werden."""
        if cand.art != "kopf":
            return False
        zellen = []
        for t in xt:
            if (t.blatt, t.zeile, t.spalte) not in zellen:
                zellen.append((t.blatt, t.zeile, t.spalte))
        if len(zellen) != 2:
            return False
        teile = neu.split()
        nr = teile[0] if teile and section_nummer(teile[0]) else ""
        titel = " ".join(teile[1:] if nr else teile)
        ops = []
        for (blatt, zeile, spalte), text in zip(zellen, (nr, titel)):
            toks = [t for t in xt if (t.blatt, t.zeile, t.spalte) == (blatt, zeile, spalte)]
            ops.append(Operation("ersetzen", blatt, zeile, spalte, toks[0].start, toks[-1].end, text, seq, aenderung=ae.id))
        ae.ops = ops
        ae.ziel = self._ziel(xt[0])
        self._setze(ae, MITTEL if stufe != HOCH else HOCH)
        return True

    # ------------------------------------------------------------ Geloeschte Bloecke
    def _geloescht(self, b: PdfBlock) -> None:
        ist_kopf = b.kind == "ueberschrift"
        toks = b.tokens
        ae = self._basis(b, "section_geloescht" if ist_kopf else "absatz_geloescht", toks,
                         b.alt_text, "")
        tr = self.treffer(b)
        schw = self.cfg["schwellen"]
        if tr is None or tr.score < schw["mittel"]:
            self._kein_treffer(ae, b)
            return
        cand = tr.kandidat
        ae.score, ae.ziel, ae.excel_text = tr.score, cand.ziel, cand.text[:300]
        stufe = tr.stufe
        n_alt = sum(1 for t in toks if t.old and norm(t.old))
        js = sorted(set(tr.zuordnung.values()))
        if len(js) / max(n_alt, 1) < schw["hoch"]:
            stufe = MITTEL
            ae.grund = "Der Text in der Excel stimmt nur teilweise mit dem gestrichenen Text überein."
        seq = b.idx * 1000
        alle = self._alle_gestrichen(cand, js, toks)
        if cand.art == "kopf" and alle:
            zellen = []
            for t in cand.tokens:
                z = (t.blatt, t.zeile, t.spalte)
                if z not in zellen:
                    zellen.append(z)
            ae.ops = [Operation("zelle_leeren", bl, ze, sp, seq=seq, aenderung=ae.id) for bl, ze, sp in zellen]
        else:
            ae.ops = self._loesch_ops(ae, cand, js, seq, alle)
            # Nicht alles gestrichen: nur wenn genau ein Stueck aus ganzen Zeilen entfernt wird, ist es sicher
            if not alle and not (len(ae.ops) == 1 and ae.ops[0].absatz_weg):
                stufe = MITTEL
                self._hinweis(ae, "In der Excel steht in derselben Zelle mehr Text, der im PDF nicht gestrichen ist. "
                                  "Nur die gestrichenen Wörter werden entfernt, der Rest bleibt stehen.")
                ae.grund = (ae.grund + "; " if ae.grund else "") + "Nur ein Teil des Excel-Textes ist im PDF gestrichen."
        self._setze(ae, stufe)

    def _alle_gestrichen(self, cand: Kandidat, js: list[int], toks: list[PdfToken]) -> bool:
        """Sind ALLE Woerter der Excel-Stelle im PDF gestrichen? (nur dann darf eine Zelle/ein Absatz ganz entfernt werden)"""
        if len(set(js)) != len(cand.tokens):
            return False
        # Woerter ohne Normform ("•") stehen nicht in cand.tokens, muessen aber auch mitgestrichen sein
        if cand.art == "kopf":
            texte = [self.xl.text(*z) for z in dict.fromkeys((t.blatt, t.zeile, t.spalte) for t in cand.tokens)]
        else:
            texte = [cand.zelle_text[cand.start:cand.end]]
        leer = Counter(w for t in texte for w in t.split() if not norm(w))
        gestrichen = Counter(t.old for t in toks if t.old and not norm(t.old))
        return not (leer - gestrichen)

    def _loesch_ops(self, ae: Aenderung, cand: Kandidat, js: list[int], seq: int, alle: bool) -> list[Operation]:
        zt = cand.zelle_text
        if alle and cand.art in ("absatz", "zelle"):
            rest = (zt[:cand.start] + zt[cand.end:]).strip()
            if not rest:
                return [Operation("zelle_leeren", cand.blatt, cand.zeile, cand.spalte, seq=seq, aenderung=ae.id)]
            return [Operation("ersetzen", cand.blatt, cand.zeile, cand.spalte, cand.start, cand.end, "", seq,
                              aenderung=ae.id, absatz_weg=True)]
        # gezielt nur die gestrichenen Woerter entfernen: je zusammenhaengendem Stueck (gleiche Zelle, nur Leerraum dazwischen)
        stuecke: list[list[int]] = []
        for j in sorted(set(js)):
            if stuecke:
                p, t = cand.tokens[stuecke[-1][-1]], cand.tokens[j]
                text = self.xl.text(t.blatt, t.zeile, t.spalte)
                if (p.blatt, p.zeile, p.spalte) == (t.blatt, t.zeile, t.spalte) and j == stuecke[-1][-1] + 1 \
                        and not text[p.end:t.start].strip():
                    stuecke[-1].append(j)
                    continue
            stuecke.append([j])
        ops: list[Operation] = []
        for k, st in enumerate(stuecke):
            t0, t1 = cand.tokens[st[0]], cand.tokens[st[-1]]
            text = self.xl.text(t0.blatt, t0.zeile, t0.spalte)
            s0, e0 = t0.start, t1.end
            davor, danach = text[:s0].rstrip(" "), text[e0:].lstrip(" ")
            ganze_zeilen = (davor == "" or davor.endswith("\n")) and (danach == "" or danach.startswith("\n"))
            if ganze_zeilen:
                ops.append(Operation("ersetzen", t0.blatt, t0.zeile, t0.spalte, s0, e0, "", seq + k, aenderung=ae.id, absatz_weg=True))
            else:
                s, e = self._luecke(text, s0, e0)
                ops.append(Operation("ersetzen", t0.blatt, t0.zeile, t0.spalte, s, e, "", seq + k, aenderung=ae.id))
        return ops

    # ------------------------------------------------------------ Neue Bloecke
    def _ort_aus_fund(self, i: int, fund: Fundstelle, z: Zuordnung, direkt: bool, davor: bool) -> _Ort:
        """Ankerstelle aus einer exakten Fundstelle. "hoch" nur, wenn der Anker direkt davor steht, sicher auf der Kette liegt,
        sein Ende das Ende einer Excel-Zeile ist und dazwischen in den verwalteten Spalten nichts Unzugeordnetes liegt."""
        stufe = HOCH
        gruende: list[str] = []
        codes: list[str] = []
        if not z.automatisch:
            stufe = MITTEL
            codes.append("anker_unsicher")
            gruende.append(_GRUND.get(z.status, _GRUND["anker_unsicher"]))
        if not direkt:
            stufe = MITTEL
            codes.append("anker_unsicher")
            gruende.append("der Absatz davor konnte in der Excel nicht zugeordnet werden, die Einfügestelle ist nur ungefähr bestimmbar")
        if davor and not fund.zeilen_ende:
            stufe = MITTEL
            codes.append("anker_teil")
            gruende.append("der PDF-Text deckt nur den Anfang oder ein Stück des längeren Excel-Absatzes ab (Seitenumbruch?); "
                           "neuer Absatz könnte mitten im Satz landen")
        if not davor and not fund.zeilen_anfang:
            stufe = MITTEL
            codes.append("anker_teil")
            gruende.append(_GRUND["anker_teil"])
        if davor:
            hi = self.zuo[i].fenster[1]
            if hi is not None and hi[0] != fund.end_schluessel[0]:
                hi = None
            if self._unzugeordnet_dazwischen(fund.blatt, fund.end_schluessel, hi):
                stufe = MITTEL
                codes.append("anker_luecke")
                gruende.append("zwischen der Ankerstelle und dem nächsten zugeordneten Absatz steht in der Excel Text, der keinem "
                               "PDF-Absatz zugeordnet ist")
        if fund.versteckt:
            stufe = MITTEL
            codes.append("versteckt")
            gruende.append(_GRUND["versteckt"][:-1].lower())
        lay = self.xl.layouts.get(fund.blatt)
        if lay is not None and lay.sicherheit != "hoch":
            stufe = MITTEL
            codes.append("layout_unsicher")
            gruende.append(f"der Aufbau des Blatts '{fund.blatt}' wurde nur unsicher erkannt")
        z0, s0, a0, _ = fund.zellen[0]
        zl, sl, _, el = fund.zellen[-1]
        return _Ort(fund.blatt, z0, s0, fund.art, a0, el, fund.text, 1.0, stufe, davor, gruende, list(dict.fromkeys(codes)))

    def _reihenfolge_sicher(self, k: int, tr: Treffer) -> Fundstelle | None:
        """Der Vorgaenger-Absatz k steht in der Excel nicht genau so (Excel weicht ab), ist aber ueber die Reihenfolge eindeutig:
        zwischen den beiden sicheren Nachbarankern gibt es im PDF genau diesen einen (nicht neuen) Absatz ohne Anker und in der
        Excel genau eine Textzeile -- die, die der Matcher gefunden hat. Dann ist sie ein sicherer Anker fuer den neuen Absatz
        dahinter (Welle 3). Sonst None."""
        z = self.zuo[k]
        lo, hi = z.fenster
        kd = tr.kandidat
        if lo is None or hi is None or kd.art == "kopf":
            return None
        vor = max((j for j in range(k) if self.zuo[j].automatisch), default=None)
        nach = min((j for j in range(k + 1, len(self.spuren)) if self.zuo[j].automatisch), default=None)
        if vor is None or nach is None:
            return None
        offen = [j for j in range(vor + 1, nach) if not self.spuren[j].ganz_neu and not self.zuo[j].automatisch]
        if offen != [k]:
            return None
        zeilen = [ort for key, ort in self.idx.verwaltete_zeilen(kd.blatt) if lo < key < hi]
        if len(zeilen) != 1 or zeilen[0] != (kd.zeile, kd.spalte, kd.start, kd.end):
            return None
        return self.idx._fund(kd.blatt, [zeilen[0]], "absatz", kd)

    def _ort_aus_treffer(self, tr: Treffer, davor: bool) -> _Ort:
        k = tr.kandidat
        if k.art != "kopf" and tr.zuordnung:
            ende = k.tokens[max(tr.zuordnung.values())].end
        else:
            ende = k.end
        gruende = list(tr.gruende) + [g for g in tr.anker_gruende if g not in tr.gruende]
        gruende.append("der Absatz davor steht so nicht in der Excel, die Einfügestelle ist nur ungefähr bestimmbar")
        return _Ort(k.blatt, k.zeile, k.spalte, k.art, k.start, ende, k.text, tr.score, MITTEL, davor, gruende,
                    ["anker_unsicher"])

    def _anker(self, i: int) -> _Ort | None:
        """Naechster Kettenanker vor dem neuen Block (sonst danach). Ohne Kette: alter Matcher (hoechstens "mittel")."""
        schw = self.cfg["schwellen"]
        uebersprungen = 0
        for k in range(i - 1, -1, -1):
            sp, z = self.spuren[k], self.zuo[k]
            if z.fund is not None and (z.automatisch or not sp.ganz_neu):
                return self._ort_aus_fund(i, z.fund, z, uebersprungen == 0, True)
            if sp.ganz_neu:
                continue
            self._fenster = z.fenster
            tr = self.treffer(sp.bloecke[-1])
            if tr is not None and tr.score >= schw["mittel"]:
                sicher = self._reihenfolge_sicher(k, tr)
                if sicher is not None:
                    return self._ort_aus_fund(i, sicher, Zuordnung("durch_reihenfolge", sicher, fenster=z.fenster),
                                              uebersprungen == 0, True)
                return self._ort_aus_treffer(tr, True)
            uebersprungen += 1
        for k in range(i + 1, len(self.spuren)):
            sp, z = self.spuren[k], self.zuo[k]
            if z.fund is not None and (z.automatisch or not sp.ganz_neu):
                o = self._ort_aus_fund(i, z.fund, z, False, False)
                return o
            if sp.ganz_neu:
                continue
            self._fenster = z.fenster
            tr = self.treffer(sp.bloecke[0])
            if tr is not None and tr.score >= schw["mittel"]:
                return self._ort_aus_treffer(tr, False)
        return None

    def _basis_neu(self, sp: Spur, typ: str) -> Aenderung:
        toks = sp.tokens
        seite = toks[0].page
        ae = Aenderung(id=f"A{len(self.out) + 1:04d}", typ=typ, seite=seite, bbox=_bbox([t for t in toks if t.page == seite]),
                       section=sp.erster.section_neu or sp.erster.section_alt, alt="", neu=sp.neu,
                       woerter=sum(1 for t in toks if t.status != KEEP), block_idx=sp.erster.idx, kontext=sp.neu[:400])
        self.out.append(ae)
        self._spur_von[ae.id] = self._aktuelle_spur if self._aktuelle_spur is not None else -1
        self._sperre[ae.id] = list(self._spur_sperre)
        self._hinweise_nur(ae)
        ae.soll_absatz = sp.neu
        return ae

    def _neu(self, i: int) -> None:
        sp = self.spuren[i]
        ist_kopf = sp.kind == "ueberschrift"
        koerper: list[int] = []
        if ist_kopf:
            k = i + 1
            while k < len(self.spuren) and self.spuren[k].ganz_neu and self.spuren[k].kind == "absatz":
                koerper.append(k)
                k += 1
        # die Folgeabsaetze stehen in derselben Excel-Zeile: ihr Zweifel gilt auch fuer die Section (und umgekehrt)
        self._spur_start(i, [self.spuren[k] for k in koerper])
        if any(self._seitenende_offen(k) for k in [i, *koerper]) and "fortsetzung_unklar" not in self._spur_sperre:
            # neuer Text endet ohne Satzende unten auf der Seite (z. B. vor einer Fussnote) und ist nicht sicher mit der
            # naechsten Seite verbunden: der Absatz geht vielleicht weiter -> nie automatisch
            if "fortsetzung_unklar" in [str(c) for c in _planer_cfg(self.cfg).get("unsichere_hinweise", [])]:
                self._spur_sperre = self._spur_sperre + ["fortsetzung_unklar"]
        ae = self._basis_neu(sp, "section_neu" if ist_kopf else "absatz_neu")
        if koerper:
            ae.gruppe = ae.id
        z = self.zuo[i]
        if z.fund is not None and z.art == "neu" and z.automatisch:          # steht schon an genau dieser Stelle (Idempotenz)
            schon = z.fund
            if ist_kopf:
                grund = f"Die Section steht schon in der Excel ({schon.ziel}); die Änderung ist offenbar schon eingetragen."
                self._nicht(ae, grund, schon.ziel)
                self._code(ae, "schon_eingetragen")
                for k in koerper:
                    self.verbraucht.add(k)
                    kae = self._basis_neu(self.spuren[k], "absatz_neu")
                    kae.gruppe = ae.id
                    self._nicht(kae, "Die zugehörige neue Section steht schon in der Excel.", schon.ziel)
                    self._code(kae, "schon_eingetragen")
                return
            if len([w for w in sp.neu.split() if norm(w)]) >= _planer_cfg(self.cfg).get("doppelt_min_woerter", 3):
                self._nicht(ae, f"Der Absatz steht schon in der Excel ({schon.ziel}); die Änderung ist offenbar schon eingetragen.",
                            schon.ziel)
                self._code(ae, "schon_eingetragen")
                ae.excel_text = schon.text[:300]
                return
            self._hinweis(ae, f"Ein sehr kurzer Absatz mit gleichem Text steht schon in der Excel ({schon.ziel}). "
                              "Bitte prüfen, ob er doppelt wäre.")
            self._code(ae, "vielleicht_doppelt")
            doppelt = True
        else:
            doppelt = False
        if not ist_kopf and not doppelt and self._in_geleerte_zelle(i, ae):
            return
        ort = self._anker(i)
        if ort is not None and self._register and not self._anker_je_blatt.get(ort.blatt):
            # ADR 0006: Nebenblaetter (Hinweise, Deckblatt, Unterschriften) bekommen keine neuen Zeilen, solange dort kein
            # Absatz sicher zugeordnet ist
            self._nicht(ae, _codes.satz("blatt_ohne_anker"), ort.ziel)
            self._code(ae, "blatt_ohne_anker")
            for k in koerper:
                self.verbraucht.add(k)
                kae = self._basis_neu(self.spuren[k], "absatz_neu")
                kae.gruppe = ae.id
                self._nicht(kae, _codes.satz("blatt_ohne_anker"), ort.ziel)
                self._code(kae, "blatt_ohne_anker")
            return
        if ort is None:
            self._nicht(ae, "Keine Ankerstelle in der Excel gefunden (kein benachbarter Text zuordenbar).")
            self._code(ae, "kein_anker")
            for k in koerper:
                self.verbraucht.add(k)
                kae = self._basis_neu(self.spuren[k], "absatz_neu")
                kae.gruppe = ae.id
                self._nicht(kae, "Gehört zu einer neuen Section ohne Ankerstelle.")
                self._code(kae, "kein_anker")
            return
        stufe = ort.stufe if not doppelt else MITTEL
        self._ag = list(dict.fromkeys(g for g in ort.gruende if g)) if stufe != HOCH else []
        self._ag_codes = list(ort.codes) if stufe != HOCH else []
        if stufe != HOCH:
            self._anker_grund(ae, self._ag)
        ae.score = ort.score
        ae.excel_text = ort.text[:200]
        lay = self.xl.layouts[ort.blatt]
        if lay.inhalt_spalte is None and not (ist_kopf and lay.section_spalte):
            self._nicht(ae, "Die Inhaltsspalte der Excel konnte nicht erkannt werden.", ort.ziel)
            self._code(ae, "keine_inhaltsspalte")
            return
        z_cfg = self.cfg["einfuegen"]
        seq = sp.erster.idx * 1000
        self._nachbar_spalte = self._naechste_textspalte(i, ort.blatt) if ort.art == "kopf" else None
        if ist_kopf:
            self._neue_section(ae, sp, [self.spuren[k] for k in koerper], koerper, ort, lay, stufe, ort.davor, seq)
        else:
            modus = z_cfg["neuer_absatz"]
            if modus == "auto":
                modus = "zelle" if lay.modell == "zeile_pro_section" else "zeile"
            if modus == "zelle":
                self._absatz_in_zelle(ae, sp, ort, lay, stufe, seq)
            else:
                self._absatz_als_zeile(ae, sp, ort, lay, stufe, ort.davor, seq)

    _ag_codes: list[str] = []

    def _seitenende_offen(self, i: int) -> bool:
        """Endet der (neue) Absatz ohne Satzende unten auf seiner Seite, und folgt auf dieser Seite nur noch wenig unveraenderter
        Text (z. B. eine Fussnote), bevor die naechste Seite beginnt? Dann koennte er dort weitergehen."""
        sp = self.spuren[i]
        if sp.kind != "absatz" or any(b.fortsetzung == "sicher" for b in sp.bloecke):
            return False
        pl = _planer_cfg(self.cfg)
        ende = sp.neu.rstrip().rstrip("\"'“”»«)]")
        if not ende or ende[-1:] in str(pl.get("satzende_zeichen", ".!?:;")):
            return False
        seite = sp.bloecke[-1].page
        rest = 0
        for k in range(i + 1, len(self.spuren)):
            nx = self.spuren[k]
            if nx.bloecke[0].page > seite:
                return True
            if nx.hat_aenderung:
                return False
            rest += len(nx.alt.split())
            if rest > int(pl.get("seitenende_rest_woerter", 25)):
                return False
        return False

    def _in_geleerte_zelle(self, i: int, ae: Aenderung) -> bool:
        """Ein ganz gestrichener Absatz, direkt gefolgt von einem ganz neuen: wie ein Mensch den neuen Text in die Zelle des
        gestrichenen schreiben (Ersetzung), statt die Zelle zu leeren und eine neue Zeile anzulegen. Nur wenn die Loeschung die
        ganze Zelle leert; beide gelten nur gemeinsam (`gruppe`)."""
        if i == 0 or self.spuren[i].kind != "absatz":
            return False
        vor = self.spuren[i - 1]
        if not (vor.ganz_geloescht and vor.kind == "absatz"):
            return False
        loesch = [a for a in self.out if self._spur_von.get(a.id) == i - 1 and a.typ == "absatz_geloescht"]
        if len(loesch) != 1 or loesch[0].status == NICHT or len(loesch[0].ops) != 1:
            return False
        lo = loesch[0]
        op = lo.ops[0]
        if self.xl.layouts[op.blatt].register:
            return False                                  # Register: der gestrichene bleibt stehen, der neue bekommt eine eigene Zeile
        voll = self.xl.text(op.blatt, op.zeile, op.spalte)
        if not (op.art == "ersetzen" and op.text == "" and op.start == 0 and op.ende == len(voll) and voll):
            return False
        if any(o.art == "ersetzen" and o.start == o.ende == len(voll) and (o.blatt, o.zeile, o.spalte) == (op.blatt, op.zeile, op.spalte)
               for a in self.out if a is not ae for o in a.ops):
            return False                                  # dort steht schon ein neuer Absatz
        sp = self.spuren[i]
        ae.ops = [Operation("ersetzen", op.blatt, op.zeile, op.spalte, len(voll), len(voll), sp.neu, sp.erster.idx * 1000,
                            aenderung=ae.id)]
        ae.ziel = lo.ziel
        ae.score = 1.0
        ae.excel_text = voll[:200]
        gruppe = lo.gruppe or lo.id
        for a in self.out:
            if a.gruppe and a.gruppe == lo.gruppe:
                a.gruppe = gruppe
        lo.gruppe = ae.gruppe = gruppe
        self._hinweis(ae, "Der neue Absatz ersetzt den gestrichenen Absatz direkt davor und steht in dessen Zelle; beide werden nur "
                          "gemeinsam eingetragen.")
        stufe = HOCH if lo.status == UEBERNOMMEN else MITTEL
        self._setze(ae, stufe, "" if stufe == HOCH else "Gehört zur Streichung davor, die noch geprüft werden muss.",
                    [] if stufe == HOCH else ["gruppe"])
        return True

    def _anker_grund(self, ae: Aenderung, gruende: list[str]) -> None:
        """Gruende, warum die Ankerstelle fuer einen neuen Absatz unsicher ist: in `grund` und als Hinweis."""
        for g in gruende:
            self._hinweis(ae, g[:1].upper() + g[1:] + ".")
        if gruende:
            ae.grund = (ae.grund + "; " if ae.grund else "") + "Die Stelle für den neuen Absatz ist unsicher: " + "; ".join(gruende)

    def _absatz_stil(self, cand: Any, lay) -> int:
        """Zeile, deren Format eine neue Absatz-Zeile bekommt: eine Absatzzeile, nie die Ueberschriftszeile."""
        kopf = set(lay.kopf_zeilen)
        if cand.zeile not in kopf or not lay.inhalt_spalte:
            return cand.zeile
        for r in self.xl.section_zeilen(cand.blatt, cand.zeile):
            if r not in kopf and self.xl.text(cand.blatt, r, lay.inhalt_spalte).strip():
                return r
        for r in range(cand.zeile + 1, lay.max_zeile + 1):
            if r not in kopf and self.xl.text(cand.blatt, r, lay.inhalt_spalte).strip():
                return r
        return cand.zeile

    def _register_neu(self, i: int, ort: _Ort) -> tuple[dict[int, str], int]:
        """Werte einer neuen Registerzeile fuer den PDF-Absatz i (ADR 0006) und ihre Vorlagezeile (naechste Zeile gleichen Typs):
        Section aus der PDF-Gliederung (Schreibweise wie die Ankerzeile, wenn gleich), Sub-Section = Buchstabe der geltenden
        Sub-Section (bzw. der eigene bei einer neuen Sub-Section), Text mit Listennummer + Polster wie die Nachbarzeilen
        gleicher Art (ohne Label bei Sub-Section, Section, Part), Kennwort nur bei Ueberschrift-Typ. Konstanten (Spezifikation,
        Revision, Datum ...) kopiert der Schreiber aus der Vorlagezeile."""
        sp, lab = self.spuren[i], self._labels[i]
        lay = self.xl.layouts[ort.blatt]
        ids = {l.aenderung for l in sp.laeufe if l.aenderung}
        text = zustand(sp.laeufe[lab.n:], ids).strip() if lab.art and lab.neu else sp.neu
        sec = sp.erster.section_neu or ""
        anker = self.xl.register_zeile(ort.blatt, ort.zeile)
        if anker is not None and sec and (anker.section or "").casefold() == sec.casefold():
            e = anker.section_roh
        else:
            e = sec or (anker.section_roh if anker is not None else "")
        sub = self._sub_neu[i] if i < len(self._sub_neu) else ""
        if lab.art in ("section", "part") and lab.neu:
            art = lab.art
        elif lab.art == "sub" and lab.neu:
            sub = lab.neu.rstrip(".")
            art = "sub_titel" if self._titel_zeile(ort.blatt, "sub_titel", text) else "sub"
        elif lab.art == "liste" and lab.neu:
            art = "liste_titel" if self._titel_zeile(ort.blatt, "liste_titel", text) else "liste"
            text = lab.neu + self.xl.register_polster(ort.blatt, ort.zeile, lab.neu, self._zeilen_der_spec(ort.blatt)) + text
        else:
            art = "absatz"
        zellen: dict[int, str] = {}
        if lay.section_spalte and e:
            zellen[lay.section_spalte] = e
        if sub and lay.untergliederung_spalte and art not in ("section", "part"):
            zellen[lay.untergliederung_spalte] = sub
        zellen[lay.inhalt_spalte] = text                                    # type: ignore[index]
        if art in REGISTER_UEBERSCHRIFT and lay.kennwort_spalte and lay.kennwort:
            zellen[lay.kennwort_spalte] = lay.kennwort
        return zellen, self.xl.register_vorlage(ort.blatt, ort.zeile, art, self._zeilen_der_spec(ort.blatt))

    def _titel_zeile(self, blatt: str, art: str, text: str) -> bool:
        """Neuer Listenpunkt / Sub-Section-Absatz mit Doppelpunkt am Ende: Ueberschriftzeile, wenn das Register solche hat."""
        if not text.rstrip().endswith(":"):
            return False
        lay = self.xl.layouts[blatt]
        return any((rz := self.xl.register_zeile(blatt, r)) is not None and rz.art == art
                   for r in range(lay.daten_ab, lay.max_zeile + 1) if self._erlaubt(blatt, r))

    def _neue_section(self, ae, sp: Spur, koerper: list[Spur], koerper_nr: list[int], cand: _Ort, lay, stufe, davor, seq) -> None:
        if lay.register:
            self._neue_section_register(ae, koerper_nr, cand, stufe, davor, seq)
            return
        text = sp.neu
        nr_titel = kopf_zerlegen(text)
        nr, titel = nr_titel if nr_titel else ("", text)
        zellen: dict[int, str] = {}
        if lay.section_spalte is None:
            zellen[lay.inhalt_spalte] = text            # type: ignore[index]
        elif lay.kombiniert:
            zellen[lay.section_spalte] = text
        elif lay.titel_spalte:
            zellen[lay.section_spalte] = nr
            zellen[lay.titel_spalte] = titel
        else:
            zellen[lay.section_spalte] = text
        modus_zeile = lay.modell == "zeile_pro_section" and lay.inhalt_spalte is not None
        if modus_zeile and koerper:
            zellen[lay.inhalt_spalte] = "\n".join(kb.neu for kb in koerper)       # type: ignore[index]
        kopf_zeilen = [r for r in lay.kopf_zeilen if r <= cand.zeile]
        stil = kopf_zeilen[-1] if kopf_zeilen else cand.zeile
        op = Operation("zeile_einfuegen", cand.blatt, cand.zeile, 0, text="", seq=seq, zellen=zellen,
                       stil_von=stil, vor=not davor, aenderung=ae.id)
        ae.ops = [op]
        ae.ziel = f"{cand.blatt}!Zeile {cand.zeile + (0 if not davor else 1)} (neu)"
        self._setze(ae, stufe, codes=self._neu_codes(stufe))
        folge = seq
        for k, kb in zip(koerper_nr, koerper):
            self.verbraucht.add(k)
            kae = self._basis_neu(kb, "absatz_neu")
            kae.gruppe = ae.id
            kae.score = ae.score
            kae.excel_text = ""
            if ae.status == NICHT:            # die Section selbst ging nicht (z. B. Steuerzeichen) -> Folgeabsatz auch nicht
                self._nicht(kae, "Gehört zu einer neuen Section, die nicht eingetragen werden kann.")
            elif modus_zeile:
                for h in ae.hinweise:                 # Risiken und Ankerzweifel der Section gelten auch fuer den Absatz in ihrer Zeile
                    self._hinweis(kae, h)
                # Der Folgeabsatz steht in der Zelle der neuen Section-Zeile. Eine "abgedeckte" Markierungs-Operation
                # (ohne Zellen) bindet ihn an die Section: Status und Freigabe gelten fuer beide gemeinsam (applier).
                kae.ops = [Operation("zeile_einfuegen", cand.blatt, cand.zeile, lay.inhalt_spalte or 0, seq=seq,
                                     abgedeckt=True, vor=not davor, aenderung=kae.id)]
                kae.ziel = ae.ziel
                self._setze(kae, stufe, "als Teil der neuen Section eingetragen", codes=self._neu_codes(stufe))
                kae.konfidenz, kae.status = ae.konfidenz, ae.status
                for c in ae.pruefcodes:
                    self._code(kae, c)
                if ae.status == PRUEFEN and ae.grund:
                    kae.grund = (kae.grund + "; " if kae.grund else "") + ae.grund
                self._hinweis(kae, "Dieser Absatz gehört zur neuen Section und wird nur zusammen mit ihr eingetragen.")
            else:
                folge += 1
                kop = Operation("zeile_einfuegen", cand.blatt, cand.zeile, 0, text="", seq=folge,
                                zellen={lay.inhalt_spalte: kb.neu}, stil_von=self._absatz_stil(cand, lay),   # type: ignore[dict-item]
                                vor=not davor, aenderung=kae.id)
                kae.ops = [kop]
                kae.ziel = ae.ziel
                self._anker_grund(kae, self._ag)
                self._setze(kae, stufe, codes=self._neu_codes(stufe))

    def _neue_section_register(self, ae, koerper_nr: list[int], cand: _Ort, stufe, davor, seq) -> None:
        """Neue Section im Register: eine Zeile fuer die Ueberschrift (Section-Spalte, Titel, Kennwort) und je neuem Absatz
        darunter eine eigene Zeile, alle direkt hinter der Ankerzeile in PDF-Reihenfolge; sie gelten nur gemeinsam."""
        i = self._aktuelle_spur if self._aktuelle_spur is not None else -1
        zellen, stil = self._register_neu(i, cand)
        ae.ops = [Operation("zeile_einfuegen", cand.blatt, cand.zeile, 0, text="", seq=seq, zellen=zellen, stil_von=stil,
                            vor=not davor, aenderung=ae.id)]
        ae.ziel = f"{cand.blatt}!Zeile {cand.zeile + (1 if davor else 0)} (neu)"
        self._setze(ae, stufe, codes=self._neu_codes(stufe))
        folge = seq
        for k in koerper_nr:
            self.verbraucht.add(k)
            kae = self._basis_neu(self.spuren[k], "absatz_neu")
            kae.gruppe = ae.id
            kae.score = ae.score
            if ae.status == NICHT:
                self._nicht(kae, "Gehört zu einer neuen Section, die nicht eingetragen werden kann.")
                continue
            folge += 1
            z2, st2 = self._register_neu(k, cand)
            kae.ops = [Operation("zeile_einfuegen", cand.blatt, cand.zeile, 0, text="", seq=folge, zellen=z2, stil_von=st2,
                                 vor=not davor, aenderung=kae.id)]
            kae.ziel = ae.ziel
            self._anker_grund(kae, self._ag)
            self._setze(kae, stufe, codes=self._neu_codes(stufe))

    def _neu_codes(self, stufe: str) -> list[str]:
        return list(self._ag_codes) if stufe != HOCH else []

    _nachbar_spalte: int | None = None

    def _text_spalte(self, cand: _Ort, lay) -> int | None:
        """Spalte fuer einen neuen Absatz: wie ein Mensch neben den Nachbarabsatz (Ankerstelle im Fliesstext; ist der Anker eine
        Ueberschrift, die Spalte des naechsten zugeordneten Absatzes), sonst die erkannte Inhaltsspalte."""
        if cand.art in ("absatz", "zeilen", "zelle", "teil") and cand.spalte not in (lay.section_spalte, lay.titel_spalte):
            return cand.spalte
        if self._nachbar_spalte is not None:
            return self._nachbar_spalte
        return lay.inhalt_spalte

    def _naechste_textspalte(self, i: int, blatt: str) -> int | None:
        for k in range(i + 1, len(self.spuren)):
            z = self.zuo[k]
            if self.spuren[k].kind == "ueberschrift":
                return None
            if z.automatisch and z.fund is not None and z.fund.blatt == blatt and not z.fund.kopf_spalte:
                return z.fund.spalte
        return None

    def _absatz_in_zelle(self, ae, sp: Spur, cand: _Ort, lay, stufe, seq) -> None:
        col = self._text_spalte(cand, lay)
        zeile = cand.zeile
        ist_inhalt = cand.spalte == col and cand.art in ("absatz", "zeilen", "zelle", "teil")
        zt = self.xl.text(cand.blatt, zeile, col)            # type: ignore[arg-type]
        if ist_inhalt and cand.davor:
            pos, text = cand.end, "\n" + sp.neu
            for a0, e0 in self._geloescht_bis.get((cand.blatt, zeile, col), []):    # type: ignore[arg-type]
                if a0 < pos < e0:            # der Anker (erste Zeile) wird samt folgendem Umbruch geloescht: dahinter einsetzen
                    pos, text = e0, sp.neu + "\n"
        elif ist_inhalt:
            pos, text = cand.start, sp.neu + "\n"
        elif cand.art == "kopf" or cand.spalte in (lay.section_spalte, lay.titel_spalte):
            pos, text = 0, sp.neu + ("\n" if zt.strip() else "")
        else:
            pos, text = len(zt), ("\n" if zt.strip() else "") + sp.neu
        # "ersetzen" mit leerer Spanne (wie die Text-Aenderungen): bei gleicher Stelle entscheidet `seq` die Reihenfolge, so steht
        # ein neuer Absatz HINTER einem Satz, der im PDF davor ans Absatzende angehaengt wird.
        ae.ops = [Operation("ersetzen", cand.blatt, zeile, col, pos, pos, text, seq, aenderung=ae.id)]    # type: ignore[arg-type]
        ae.ziel = f"{cand.blatt}!{get_column_letter(col)}{zeile}"                                         # type: ignore[arg-type]
        self._setze(ae, stufe, codes=self._neu_codes(stufe))

    def _absatz_als_zeile(self, ae, sp: Spur, cand: _Ort, lay, stufe, davor, seq) -> None:
        if lay.inhalt_spalte is None:
            self._nicht(ae, "Die Inhaltsspalte der Excel konnte nicht erkannt werden.", cand.ziel)
            return
        if lay.register:
            zellen, stil = self._register_neu(self._aktuelle_spur if self._aktuelle_spur is not None else -1, cand)
        else:
            zellen, stil = {self._text_spalte(cand, lay): sp.neu}, self._absatz_stil(cand, lay)      # type: ignore[dict-item]
        op = Operation("zeile_einfuegen", cand.blatt, cand.zeile, 0, text="", seq=seq,
                       zellen=zellen, stil_von=stil, vor=not davor, aenderung=ae.id)
        ae.ops = [op]
        ae.ziel = f"{cand.blatt}!Zeile {cand.zeile + (1 if davor else 0)} (neu)"
        self._setze(ae, stufe, codes=self._neu_codes(stufe))


_TEXT_OPS = ("ersetzen", "einfuegen", "zelle_leeren")


def end_stuecke(text: str, ops: list[Operation], zwischen: dict[str, str] | None = None, cfg: dict | None = None,
                absaetze: list[tuple[int, int, list[Lauf], set[str]]] | None = None) -> list[list[str]]:
    """Voraussichtliche End-Laeufe einer Zelle als [[format, text], ...] (normal | alt | neu) fuer die Anzeige (Schema 4,
    `Aenderung.stuecke`): geloeschter Text bleibt als alt stehen, neuer steht rot an seiner Stelle (ADR 0005). EINE kleine
    Funktion, damit sie austauschbar bleibt. Sie rechnet wie der Applier: je PDF-Absatz der Zelle (`absaetze`: Bereich,
    Laeufe, angewendete Aenderungen) die Projektion der PDF-Laeufe auf die Excel-Zeichen (markierung.projiziere), sonst die
    Operationen (markierung.markiere_ops) -- auf dem Text des Excel-Modells, ohne Schriften und frueher markierte Laeufe."""
    vorher = [mk.Stueck(mk.NORMAL, text)] if text else []
    bereiche: list[tuple[int, int, list]] = []
    pids: set[str] = set()
    for a, e, laeufe, aktiv in absaetze or []:
        res = mk.projiziere([text[a:e]], laeufe, aktiv, cfg) if 0 <= a <= e <= len(text) else None
        if res is not None:
            bereiche.append((a, e, res[0]))
            pids |= aktiv
    try:
        if bereiche:
            try:
                m = mk.markiere_ops(vorher, [o for o in ops if o.aenderung not in pids], mk.SICHT_FLACH, zwischen or {}, cfg,
                                    bereiche=bereiche)
                return mk.klassen_stuecke(m.zeichen)
            except ValueError:
                pass
        m = mk.markiere_ops(vorher, ops, mk.SICHT_FLACH, zwischen or {}, cfg)
    except ValueError:
        return []
    return mk.klassen_stuecke(m.zeichen)


def _stuecke_der_aenderung(ae: Aenderung, je_zelle: dict[tuple[str, int, int], list[Operation]], xl: ExcelModell,
                           zwischen: dict[str, str], cfg: dict,
                           absatz_von: dict[tuple[str, int, int], list[Absatz]] | None = None) -> list[list[str]]:
    """End-Laeufe der Zielzelle einer Aenderung: ihre eigenen Operationen plus die der automatisch eingetragenen Aenderungen in
    derselben Zelle. Neue Zeile: der neue Text (Textspalte) ganz neu."""
    neu_zeile = next((o for o in ae.ops if o.art == "zeile_einfuegen" and not o.abgedeckt and o.zellen), None)
    if neu_zeile is not None:
        lay = xl.layouts.get(neu_zeile.blatt)
        text = neu_zeile.zellen.get(lay.inhalt_spalte) if lay is not None and lay.inhalt_spalte else None
        text = text if text is not None else max(neu_zeile.zellen.values(), key=len)
        return [[mk.NEU, text]] if text else []
    eigene = [o for o in ae.ops if o.art in _TEXT_OPS and not o.abgedeckt]
    if not eigene:
        return []
    ziel = (eigene[0].blatt, eigene[0].zeile, eigene[0].spalte)
    ops = [o for o in eigene if (o.blatt, o.zeile, o.spalte) == ziel]
    ops += [o for o in je_zelle.get(ziel, []) if o.aenderung != ae.id]
    angewendet = {o.aenderung for o in ops}
    bereiche = []
    for ab in (absatz_von or {}).get(ziel, []):
        aktiv = ab.aenderungen & angewendet
        if aktiv and len(ab.zellen) == 1:
            _, _, a, e = ab.zellen[0]
            bereiche.append((a, e, ab.laeufe, aktiv))
    return end_stuecke(xl.text(*ziel), ops, zwischen, cfg, bereiche)


def plane(pdf: PdfAnalyse, xl: ExcelModell, cfg: dict | None = None, spec: str | None = None) -> Plan:
    """Plan fuer ein PDF und eine Excel. `spec`: Spezifikationsnummer, die der Nutzer vorgibt (sonst aus dem PDF erkannt;
    nur wichtig, wenn die CDE mehrere Specs enthaelt, Welle 3)."""
    cfg = cfg or config.lade("zuordnung")
    p = _Planer(pdf, xl, cfg, spec)
    p.lauf()
    if p._spec_unklar:                   # Spec unklar: nichts automatisch, alles bleibt ein Vorschlag
        for ae in p.out:
            if ae.status == UEBERNOMMEN:
                ae.status, ae.konfidenz = PRUEFEN, MITTEL
                ae.grund = (ae.grund + "; " if ae.grund else "") + _codes.satz("spec_unklar")
                p._code(ae, "spec_unklar")
    for ae in p.out:                     # Absicherung: an der Stelle muss beim Eintragen derselbe Text stehen
        for op in ae.ops:
            if op.art == "ersetzen":
                op.erwartet = xl.text(op.blatt, op.zeile, op.spalte)[op.start:op.ende]
            elif op.art == "zelle_leeren":       # der Applier leert nur, wenn die Zelle noch genau so aussieht
                op.erwartet = xl.text(op.blatt, op.zeile, op.spalte)
    # Schema 4: Grund in einem Satz und voraussichtliches Ergebnis der Zielzelle (Anzeige)
    zwischen: dict[str, str] = {}
    absatz_von: dict[tuple[str, int, int], list[Absatz]] = {}
    for ab in p.absaetze:
        zwischen.update(mk.trenner_aus_laeufen(ab.laeufe))
        for z, s, _, _ in ab.zellen:
            absatz_von.setdefault((ab.blatt, z, s), []).append(ab)
    je_zelle: dict[tuple[str, int, int], list[Operation]] = {}
    for ae in p.out:
        if ae.status == UEBERNOMMEN:
            for o in ae.ops:
                if o.art in _TEXT_OPS and not o.abgedeckt:
                    je_zelle.setdefault((o.blatt, o.zeile, o.spalte), []).append(o)
    for ae in p.out:
        ae.grund_einfach = _codes.grund_einfach(ae)
        ae.stuecke = _stuecke_der_aenderung(ae, je_zelle, xl, zwischen, cfg, absatz_von)
    warn = list(pdf.warnungen) + list(xl.warnungen) + p.spec_warnungen
    luecken, ausserhalb = list(pdf.luecken), []
    bereich = p.seiten_der_cde()
    if bereich is not None:             # Welle 6: eine Luecke strikt vor der ersten / hinter der letzten CDE-Seite ist nur Info
        ausserhalb = [l for l in luecken if isinstance(l.get("seite"), int) and l["seite"] > 0
                      and (l["seite"] < bereich[0] or l["seite"] > bereich[1])]
        luecken = [l for l in luecken if l not in ausserhalb]
    plan = Plan(aenderungen=p.out, layout=xl.beschreibung(), warnungen=warn,
                ignoriert=pdf.ignoriert, markierte_woerter=pdf.markierte_woerter,
                luecken=luecken, absaetze=p.absaetze, spec=p.spec_info, luecken_ausserhalb=ausserhalb)
    return plan
