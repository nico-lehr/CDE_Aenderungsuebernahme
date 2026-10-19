"""Verlustfreies ("chirurgisches") Schreiben der Excel (ADR 0007).

Das Ergebnis ist eine KOPIE des Original-ZIP. Neu geschrieben werden nur die Teile, die sich aendern muessen:
  * das betroffene Blatt-XML (Zellen; bei neuen Zeilen auch alle Bezuege darunter),
  * xl/sharedStrings.xml (neue Eintraege hinten anfuegen, Zaehler anpassen; fehlt die Datei, wird sie samt Verweis und
    Inhaltstyp angelegt),
  * xl/styles.xml nur, wenn eine GANZE Zelle eine Schriftfarbe braucht (Zahl, Formel) oder ein Text mit "=" beginnt:
    neue <font>/<xf> werden hinten angefuegt, vorhandene nie geaendert,
  * bei neuen Zeilen zusaetzlich xl/workbook.xml (Druckbereich, Filter, Namen), andere Blaetter mit Formeln auf das Blatt
    und die Zeichnungs-Anker des Blatts,
  * xl/calcChain.xml wird entfernt (samt Verweis und Inhaltstyp); Excel baut sie beim Oeffnen neu.
Alle anderen Teile bleiben entpackt bytegleich und in derselben Reihenfolge (customXml, docMetadata/LabelInfo.xml,
Druckereinstellungen, Bilder, Formen, vbaProject.bin ...).

Gearbeitet wird auf der Textebene des XML mit eng begrenzten regulaeren Ausdruecken (die Dateien sind maschinell
erzeugt): Namensraeume, mc:Ignorable und unbekannte Erweiterungen bleiben unberuehrt. ElementTree dient nur der Pruefung
(wohlgeformt), nie dem Schreiben -- es wuerde Namensraum-Praefixe umbenennen.

Texte in der openpyxl-Lesart: openpyxl liest Excel-Escapes wie "_x000D_" woertlich. Text, der aus der Excel stammt, wird
deshalb genau so zurueckgeschrieben (kein zweites Escapen); nur neu eingefuegter Text (Format "neu") wird escapt.

Was beim Einfuegen einer Zeile nicht sicher mitgezogen werden kann, nennt `hindernisse` (einfache deutsche Saetze);
`schreibe` verweigert dann das Einfuegen (PatchFehler) -- es entsteht nie eine halbe oder kaputte Datei.
"""
from __future__ import annotations

import bisect
import os
import posixpath
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape, unescape

from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import column_index_from_string, get_column_letter

from . import config
from .fehler import DateiVorhanden, NutzerFehler
from .markierung import ALT, NEU, Stueck
from .markierung import konfig as markierung_konfig

MAX_ZEILE = 1048576
_ENT = {"&quot;": '"', "&apos;": "'"}
_ATTR_ESC = {'"': "&quot;"}

# Beziehungstypen (Ende der Type-URI)
_T_DOKUMENT = "/officeDocument"
_T_BLATT = "/worksheet"
_T_TEXTE = "/sharedStrings"
_T_STILE = "/styles"
_T_KETTE = "/calcChain"
_T_ZEICHNUNG = "/drawing"
_T_KOMMENTARE = "/comments"
_T_VML = "/vmlDrawing"
_T_PIVOT = "/pivotTable"
_T_PIVOT_CACHE = "/pivotCacheDefinition"
_T_EXTERN = "/externalLink"
_T_DIAGRAMM = "/chart"
_T_THREAD = "/threadedComment"
_TEXTE_TYP_URI = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"
_TEXTE_INHALT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"

# Reihenfolge der Schrift-Eigenschaften, wie Excel sie schreibt (Lauf: rFont statt name)
_RPR_REIHE = ("b", "i", "strike", "condense", "extend", "outline", "shadow", "u", "vertAlign", "sz", "color", "rFont",
              "family", "charset", "scheme")
_FONT_REIHE = ("b", "i", "strike", "condense", "extend", "outline", "shadow", "u", "vertAlign", "sz", "color", "name",
               "family", "charset", "scheme")
_SCHALTER = ("b", "i", "strike", "condense", "extend", "outline", "shadow")

_ATTR = re.compile(r'([\w:.-]+)\s*=\s*(?:"([^"]*)"|\'([^\']*)\')')
_TAG = re.compile(r"<[^>]*>")
_ZEILE = re.compile(r"<row\b[^>]*?(?:/>|>.*?</row>)", re.S)
_ZEILE_TAG = re.compile(r"<row\b[^>]*?/?>")
_ZELLE = re.compile(r"<c\b[^>]*?(?:/>|>.*?</c>)", re.S)
_ZELL_TAG = re.compile(r"<c\b[^>]*?/?>")
_ZELL_R = re.compile(r'(<c\b[^>]*?\sr=")([A-Za-z]{1,3})(\d+)(")')
_FORMEL = re.compile(r"(<f\b[^>]*?)(/>|>(.*?)</f>)", re.S)
_KIND = re.compile(r"<(\w+)\b([^>]*?)\s*(?:/>|>\s*</\1>)")
_ZELL_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_SPALTEN_REF = re.compile(r"^\$?[A-Za-z]{1,3}$")
_ZEILEN_REF = re.compile(r"^(\$?)(\d+)$")


class PatchFehler(NutzerFehler):
    """Die Excel kann nicht verlustfrei geschrieben werden. `meldung` ist ein einfacher deutscher Satz."""

    def __init__(self, meldung: str, status: int = 422, art: str = "schreiben"):
        super().__init__(meldung, status, art)


# ===================================================================================== Schnittstelle
@dataclass
class ZellInhalt:
    """Neuer Inhalt einer Zelle. Koordinaten in der ORIGINAL-Nummerierung (vor neuen Zeilen).

    Genau eins davon: `stuecke` (Rich Text, je Stueck Format normal|alt|neu), `formel` (mit oder ohne "=") oder `wert`
    (str = einfacher Text, int/float = Zahl, bool, None = Zelle leeren; das Format der Zelle bleibt).
    `farbe` ("" | "neu" | "alt") faerbt die GANZE Zelle ueber ihren Zellstil (fuer Zahlen und Formeln).
    `formel_von`: Original-Zeile, auf die sich die Formel bezieht (neue Zeilen: die Vorlagezeile); 0 = `zeile`.
    `stil_von`: Original-Zeile derselben Spalte, deren Zellstil die Zelle bekommt (0 = eigener Stil bzw. Vorlagezeile).
    `text_praefix`: Text, der mit "=" beginnt, als Text kennzeichnen (Zellstil mit quotePrefix).
    `kopie_von` (nur neue Zeilen): Wert der Original-Zeile derselben Spalte genau so uebernehmen (Konstante: Text, Zahl,
    Datum, Rich Text -- ohne Formel)."""
    blatt: str
    zeile: int
    spalte: int
    stuecke: list[Stueck] | None = None
    wert: Any = None
    farbe: str = ""
    formel: str = ""
    formel_von: int = 0
    stil_von: int = 0
    text_praefix: bool = False
    kopie_von: int = 0


@dataclass
class NeueZeile:
    """Eine neue Zeile NACH der Original-Zeile `nach` (0 = ganz oben); mehrere an derselben Stelle in Listenreihenfolge.
    Zellstil je Spalte und Zeilenhoehe kommen aus der Original-Zeile `vorlage`. `zellen`: Spalte -> ZellInhalt oder Wert."""
    blatt: str
    nach: int
    vorlage: int
    zellen: dict[int, Any] = field(default_factory=dict)
    hoehe: float | None = None


@dataclass
class PatchErgebnis:
    ausgabe: str
    geaenderte_teile: list[str] = field(default_factory=list)       # neu geschriebene, angelegte und entfernte Teile
    zeilen_abbildung: dict[str, list[int]] = field(default_factory=dict)   # Blatt -> sortiert: `nach` je neuer Zeile
    warnungen: list[str] = field(default_factory=list)
    neue_zeilen: list[tuple[str, int]] = field(default_factory=list)       # je NeueZeile (Eingabereihenfolge): Blatt, Zeile

    def zeile(self, blatt: str, zeile: int) -> int:
        """Nummer einer Original-Zeile im Ergebnis."""
        return zeile + bisect.bisect_left(self.zeilen_abbildung.get(blatt, []), zeile)


# ===================================================================================== kleine XML-Helfer
def _attrs(tag: str) -> dict[str, str]:
    return {m.group(1): unescape(m.group(2) if m.group(2) is not None else m.group(3), _ENT)
            for m in _ATTR.finditer(tag)}


def _attr_name(attrs: dict[str, str], ende: str) -> str | None:
    """Wert eines Attributs mit beliebigem Namensraum-Praefix (z. B. r:id)."""
    for k, v in attrs.items():
        if k == ende or k.endswith(":" + ende):
            return v
    return None


def _setze_attr(tag: str, name: str, wert: str | None) -> str:
    """Attribut in einem Start-Tag setzen (None = entfernen); alle anderen Zeichen bleiben."""
    muster = re.compile(r"(\s)" + re.escape(name) + r"""\s*=\s*(?:"[^"]*"|'[^']*')""")
    if wert is None:
        return muster.sub("", tag, count=1)
    neu = f'{name}="{escape(wert, _ATTR_ESC)}"'
    if muster.search(tag):
        return muster.sub(lambda m: m.group(1) + neu, tag, count=1)
    m = re.search(r"\s*(/?>)$", tag)
    return tag[:m.start()] + " " + neu + tag[m.start():]


def _zahl(x: float) -> str:
    f = float(x)
    return str(int(f)) if f.is_integer() else repr(f)


def _text_xml(text: str, neu: bool = False) -> str:
    """Text fuer <t>: XML-Escapes; Wagenruecklauf als Zeichenreferenz; verbotene Steuerzeichen als Excel-Escape.
    Neu eingefuegter Text bekommt zusaetzlich das Excel-Escape fuer woertliches "_xHHHH_" (sonst zeigt Excel ein Zeichen)."""
    s = escape(text)
    if neu:
        s = re.sub(r"_(x[0-9A-Fa-f]{4}_)", r"_x005F_\1", s)
    s = s.replace("\r", "&#13;")
    return re.sub("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]", lambda m: "_x%04X_" % ord(m.group()), s)


def _rels_name(teil: str) -> str:
    d, n = posixpath.split(teil)
    return posixpath.join(d, "_rels", n + ".rels")


def _ziel_pfad(basis: str, target: str) -> str:
    if target.startswith("/"):
        return target[1:]
    return posixpath.normpath(posixpath.join(posixpath.dirname(basis), target))


# ===================================================================================== Bezuege verschieben
class _Verschiebung:
    """Zeilen-Abbildung fuer neue Zeilen in einem Blatt: Original-Zeile r -> r + (Anzahl neuer Zeilen mit nach < r)."""

    def __init__(self, paare: list[tuple[int, int]], formeln_wachsen: bool):
        self.nachs = sorted(n for n, _ in paare)
        self.paare = paare                                  # (nach, vorlage) je neuer Zeile
        self.formeln_wachsen = formeln_wachsen              # zeilen.listen_summen_erweitern
        self.gewachsen = 0                                  # Zaehler: Bereiche, die am Ende um die neue Zeile wuchsen

    def zeile(self, r: int) -> int:
        return min(r + bisect.bisect_left(self.nachs, r), MAX_ZEILE)

    def bereich(self, r0: int, r1: int, wachsen: bool) -> tuple[int, int]:
        """Excel-Regel: ein Bereich waechst, wenn die neue Zeile innerhalb liegt (r0 <= nach < r1). Mit `wachsen` waechst er
        auch, wenn er genau mit der Vorlagezeile endet und die neue Zeile direkt dahinter kommt (Liste, die weiterlaeuft)."""
        a, b = self.zeile(r0), self.zeile(r1)
        if wachsen and r0 < r1:
            # Alles oder nichts: kommen mehrere neue Zeilen nacheinander direkt hinter das Ende, waechst der Bereich nur, wenn
            # ALLE ihre Vorlage im Bereich haben (sonst haette er danach eine Luecke bzw. endete mitten in der Gruppe)
            hinter = [v for n, v in self.paare if n == r1]
            extra = len(hinter) if hinter and all(r0 <= v <= r1 for v in hinter) else 0
            if extra:
                b += extra
                self.gewachsen += 1
        return a, min(b, MAX_ZEILE)


def _teile_bezug(wert: str) -> tuple[str | None, str]:
    """"'Mein Blatt'!A1:B2" -> ("Mein Blatt", "A1:B2"); "A1" -> (None, "A1")."""
    if wert.startswith("'"):
        i = 1
        while i < len(wert):
            if wert[i] == "'":
                if i + 1 < len(wert) and wert[i + 1] == "'":
                    i += 2
                    continue
                break
            i += 1
        if i + 1 < len(wert) and wert[i + 1] == "!":
            return wert[1:i].replace("''", "'"), wert[i + 2:]
        return None, wert
    if "!" in wert:
        b, _, r = wert.rpartition("!")
        return b, r
    return None, wert


def _bezug_zeilen(ref: str) -> list[tuple[int, bool]] | None:
    """Zeilennummern eines Bezugs (A1, A1:B2, 3:5) mit 'absolut'-Merker; Spaltenbezuege (A:B) -> []; kein Bezug -> None."""
    teile = ref.split(":")
    if len(teile) > 2:
        return None
    out: list[tuple[int, bool]] = []
    spalten = 0
    for t in teile:
        m = _ZELL_REF.match(t)
        if m:
            out.append((int(m.group(4)), bool(m.group(3))))
            continue
        m = _ZEILEN_REF.match(t)
        if m and len(teile) == 2:
            out.append((int(m.group(2)), bool(m.group(1))))
            continue
        if _SPALTEN_REF.match(t) and len(teile) == 2:
            spalten += 1
            continue
        return None
    if spalten and out:
        return None
    return out


def _ref_verschieben(ref: str, v: _Verschiebung, wachsen: bool, formel_zeile: int | None = None) -> str | None:
    """Ein Bezug ohne Blattnamen (A1, $A$1:B9, 5:6, H:H) nach dem Einfuegen; None = kein Bezug (Name, Tabelle ...).
    `formel_zeile`: Zeile der Formel im selben Blatt -- ein Bereich, in dem die Formel selbst steht (laufende Summe), waechst
    nie an der Vorlage; nur Summen ausserhalb der Liste wachsen mit."""
    teile = ref.split(":")
    if _bezug_zeilen(ref) is None:
        return None
    if all(_SPALTEN_REF.match(t) for t in teile) and len(teile) == 2:
        return ref
    zeilen = []
    for t in teile:
        m = _ZELL_REF.match(t) or _ZEILEN_REF.match(t)
        zeilen.append(int(m.group(m.lastindex)))
    if len(zeilen) == 1:
        neu = [v.zeile(zeilen[0])]
    else:
        lo, hi = min(zeilen), max(zeilen)
        a, b = v.bereich(lo, hi, wachsen and not (formel_zeile is not None and lo <= formel_zeile <= hi))
        neu = [a if z == lo and (z != hi or i == 0) else b for i, z in enumerate(zeilen)]
    out = []
    for t, z in zip(teile, neu):
        m = _ZELL_REF.match(t)
        if m:
            out.append(f"{m.group(1)}{m.group(2)}{m.group(3)}{z}")
        else:
            m = _ZEILEN_REF.match(t)
            out.append(f"{m.group(1)}{z}")
    return ":".join(out)


def _sqref_verschieben(sqref: str, v: _Verschiebung, wachsen: bool = True) -> str:
    teile = []
    for t in sqref.split():
        neu = _ref_verschieben(t, v, wachsen)
        teile.append(t if neu is None else neu)
    return " ".join(teile)


def _formel_tokens(formel: str) -> list[Token]:
    return Tokenizer(formel if formel.startswith("=") else "=" + formel).items


def _formel_verschieben(formel: str, eigenes: str | None, ziel: str, v: _Verschiebung, wachsen: bool = False,
                        formel_zeile: int | None = None) -> str:
    """Formel (mit oder ohne "=") nach dem Einfuegen von Zeilen im Blatt `ziel`. `eigenes`: Blatt, in dem die Formel steht
    (Bezuege ohne Blattnamen zeigen dorthin), `formel_zeile` ihre Original-Zeile. Bezuege auf andere Blaetter, Namen und
    ganze Spalten bleiben."""
    mit_gleich = formel.startswith("=")
    tok = Tokenizer(formel if mit_gleich else "=" + formel)
    geaendert = False
    z = ziel.casefold()
    for t in tok.items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        blatt, ref = _teile_bezug(t.value)
        if blatt is None:
            if eigenes is None or eigenes.casefold() != z:
                continue
        elif blatt.casefold() != z:
            continue
        eigen = eigenes is not None and eigenes.casefold() == z
        neu = _ref_verschieben(ref, v, wachsen, formel_zeile if eigen else None)
        if neu is not None and neu != ref:
            t.value = t.value[:len(t.value) - len(ref)] + neu
            geaendert = True
    if not geaendert:
        return formel
    s = tok.render()
    return s if mit_gleich else s[1:]


def _formel_info(formel: str) -> dict[str, Any]:
    """Was eine Formel enthaelt (fuer die Hindernisse): Bezuege [(Blatt|None, [(Zeile, absolut)])], INDIREKT/BEREICH.VERSCHIEBEN,
    3D- und externe Bezuege. Nicht lesbar -> {"fehler": True}."""
    try:
        items = _formel_tokens(formel)
    except (TokenizerError, ValueError, IndexError):
        return {"fehler": True, "bezuege": [], "dynamisch": False, "drei_d": False, "extern": False}
    info: dict[str, Any] = {"fehler": False, "bezuege": [], "dynamisch": False, "drei_d": False, "extern": False}
    for t in items:
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            name = t.value.rstrip("(").upper().split(".")[-1]
            if name in ("INDIRECT", "OFFSET"):
                info["dynamisch"] = True
        elif t.type == Token.OPERAND and t.subtype == Token.RANGE:
            blatt, ref = _teile_bezug(t.value)
            if blatt is not None and (blatt.startswith("[") or "]" in blatt):
                info["extern"] = True
                continue
            if blatt is not None and ":" in blatt:
                info["drei_d"] = True
                continue
            z = _bezug_zeilen(ref)
            if z:
                info["bezuege"].append((blatt, z))
    return info


# ===================================================================================== Schrift (rPr)
def _schrift_aus_xml(xml: str) -> dict[str, str]:
    """Kinder eines <font>/<rPr> -> {Eigenschaft: Attribut-Text}; name heisst im Lauf rFont. Aus-Schalter fallen weg."""
    d: dict[str, str] = {}
    for m in _KIND.finditer(xml):
        tag, attrs = m.group(1), m.group(2)
        if tag in ("font", "rPr"):
            continue
        if tag == "name":
            tag = "rFont"
        if tag in _SCHALTER or tag == "u":
            val = _attrs(attrs).get("val")
            if val is not None and val.lower() in ("0", "false", "none"):
                continue
        d[tag] = attrs.strip()
    return d


def _farbe_attr(farbe: Any) -> str | None:
    if farbe is None:
        return None
    if isinstance(farbe, str):
        return f'rgb="{farbe}"'
    typ, wert = getattr(farbe, "type", None), getattr(farbe, "value", None)
    teile = []
    if typ == "rgb" and isinstance(wert, str):
        teile.append(f'rgb="{wert}"')
    elif typ == "theme":
        teile.append(f'theme="{int(wert)}"')
    elif typ == "indexed":
        teile.append(f'indexed="{int(wert)}"')
    elif typ == "auto":
        teile.append('auto="1"')
    else:
        return None
    tint = getattr(farbe, "tint", 0) or 0
    if tint:
        teile.append(f'tint="{tint!r}"')
    return " ".join(teile)


def _schrift_aus_openpyxl(f: Any) -> dict[str, str]:
    """openpyxl-Schrift (InlineFont eines Laufs oder Font einer Zelle) -> {Eigenschaft: Attribut-Text}."""
    d: dict[str, str] = {}
    for tag in _SCHALTER:
        if getattr(f, tag, None):
            d[tag] = ""
    u = getattr(f, "u", None)
    if u:
        d["u"] = "" if u == "single" else f'val="{u}"'
    for tag in ("vertAlign", "scheme"):
        w = getattr(f, tag, None)
        if w:
            d[tag] = f'val="{w}"'
    sz = getattr(f, "sz", None) or getattr(f, "size", None)
    if sz:
        d["sz"] = f'val="{_zahl(sz)}"'
    fa = _farbe_attr(getattr(f, "color", None))
    if fa:
        d["color"] = fa
    name = getattr(f, "rFont", None) or getattr(f, "name", None)
    if name:
        d["rFont"] = f'val="{escape(str(name), _ATTR_ESC)}"'
    for tag in ("family", "charset"):
        w = getattr(f, tag, None)
        if w is not None:
            d[tag] = f'val="{int(w)}"'
    return d


def _schrift_aus_basis(basis: Any) -> dict[str, str]:
    if isinstance(basis, dict):
        return dict(basis)
    if isinstance(basis, str):
        return _schrift_aus_xml(basis)
    return _schrift_aus_openpyxl(basis)


def _mit_format(schrift: dict[str, str], format_: str, cfg: dict | None) -> dict[str, str]:
    """Schrift eines markierten Stuecks: Grundschrift + Farbe (+ Durchstreichung bei alt, ADR 0005)."""
    k = markierung_konfig(cfg)
    d = dict(schrift)
    if format_ == ALT:
        d["color"] = f'rgb="{k.get("alt_farbe", "FF00B0F0")}"'
        d["strike"] = ""
    elif format_ == NEU:
        d["color"] = f'rgb="{k.get("neu_farbe", "FFFF0000")}"'
        d.pop("strike", None)
    return d


def _render_schrift(d: dict[str, str], reihe: tuple[str, ...], aussen: str) -> str:
    teile = []
    for tag in reihe:
        if tag not in d:
            continue
        name = "name" if (tag == "rFont" and aussen == "font") else tag
        a = d[tag]
        teile.append(f"<{name} {a}/>" if a else f"<{name}/>")
    for tag in d:                                       # unbekannte Eigenschaften (Erweiterungen) hinten anhaengen
        if tag not in reihe and not (aussen == "font" and tag == "rFont"):
            a = d[tag]
            teile.append(f"<{tag} {a}/>" if a else f"<{tag}/>")
    return f"<{aussen}>" + "".join(teile) + f"</{aussen}>" if teile else (f"<{aussen}/>" if aussen == "font" else "")


def rpr(schrift: dict[str, str]) -> str:
    """<rPr> eines Laufs in der Reihenfolge, die Excel schreibt."""
    if "name" in schrift and "rFont" not in schrift:
        schrift = {("rFont" if k == "name" else k): v for k, v in schrift.items()}
    return _render_schrift(schrift, _RPR_REIHE, "rPr")


# ===================================================================================== Paket
class _Paket:
    def __init__(self, pfad: Path):
        try:
            with zipfile.ZipFile(pfad) as z:
                self.infos = z.infolist()
                self.roh = {i.filename: z.read(i) for i in self.infos}
        except (zipfile.BadZipFile, OSError, KeyError, RuntimeError) as e:
            raise PatchFehler("Die Excel-Datei ist beschädigt oder keine .xlsx-Datei.") from e
        self.neu: dict[str, str] = {}
        self.weg: set[str] = set()
        self.anhang: list[str] = []

    def gibt(self, name: str) -> bool:
        return (name in self.roh or name in self.neu) and name not in self.weg

    def text(self, name: str) -> str:
        if name in self.neu:
            return self.neu[name]
        try:
            return self.roh[name].decode("utf-8")
        except UnicodeDecodeError as e:
            raise PatchFehler(f"Ein Teil der Excel ({name}) ist nicht als UTF-8 gespeichert und kann nicht sicher "
                              "bearbeitet werden.") from e

    def setze(self, name: str, text: str) -> None:
        if name not in self.roh and name not in self.anhang:
            self.anhang.append(name)
        self.neu[name] = text
        self.weg.discard(name)

    def entferne(self, name: str) -> None:
        self.weg.add(name)
        self.neu.pop(name, None)

    def geaendert(self) -> list[str]:
        out = [n for n in self.neu if n not in self.roh or self.neu[n].encode("utf-8") != self.roh[n]]
        return out + sorted(n for n in self.weg if n in self.roh)

    def rels(self, teil: str) -> list[dict[str, str]]:
        name = _rels_name(teil)
        if not self.gibt(name):
            return []
        out = []
        for m in re.finditer(r"<Relationship\b[^>]*?/?>", self.text(name)):
            a = _attrs(m.group(0))
            if a.get("TargetMode", "") != "External" and "Target" in a:
                a["_pfad"] = _ziel_pfad(teil, a["Target"])
            a["_xml"] = m.group(0)
            out.append(a)
        return out

    def schreibe(self, ziel: Path) -> None:
        tmp = ziel.with_name(ziel.name + ".teil")
        try:
            with zipfile.ZipFile(tmp, "w") as z:
                for info in self.infos:
                    n = info.filename
                    if n in self.weg:
                        continue
                    daten = self.neu[n].encode("utf-8") if n in self.neu else self.roh[n]
                    zi = zipfile.ZipInfo(n, info.date_time)
                    zi.compress_type = info.compress_type
                    zi.external_attr = info.external_attr
                    zi.create_system = info.create_system
                    z.writestr(zi, daten)
                for n in self.anhang:
                    zi = zipfile.ZipInfo(n, self.infos[0].date_time if self.infos else (1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    z.writestr(zi, self.neu[n].encode("utf-8"))
            if ziel.exists():
                raise DateiVorhanden("Die Ergebnis-Datei gibt es schon und wird nicht überschrieben.")
            os.replace(tmp, ziel)
        finally:
            if tmp.exists():
                tmp.unlink()


# ===================================================================================== Blatt-XML
class _Zeile:
    __slots__ = ("r", "vor", "roh")

    def __init__(self, r: int, vor: str, roh: str):
        self.r, self.vor, self.roh = r, vor, roh

    @property
    def tag(self) -> str:
        return _ZEILE_TAG.match(self.roh).group(0)

    def zellen(self) -> list[list[Any]]:
        """[[Spalte, Zell-XML], ...] in Datei-Reihenfolge."""
        tag = self.tag
        if tag.endswith("/>"):
            return []
        innen = self.roh[len(tag):-len("</row>")]
        out: list[list[Any]] = []
        pos = 0
        for m in _ZELLE.finditer(innen):
            if innen[pos:m.start()].strip():
                raise PatchFehler("In einer Zeile der Excel steht unerwarteter Inhalt; sie wird nicht verändert.")
            ref = _attrs(_ZELL_TAG.match(m.group(0)).group(0)).get("r", "")
            sm = re.match(r"\$?([A-Za-z]{1,3})\$?\d+$", ref)
            if not sm:
                raise PatchFehler("Eine Zelle der Excel hat keine Adresse; die Datei wird nicht verändert.")
            out.append([column_index_from_string(sm.group(1).upper()), m.group(0)])
            pos = m.end()
        if innen[pos:].strip():
            raise PatchFehler("In einer Zeile der Excel steht unerwarteter Inhalt; sie wird nicht verändert.")
        return out

    def setze_zellen(self, zellen: list[list[Any]], spans_weg: bool = False) -> None:
        tag = self.tag
        if tag.endswith("/>"):
            tag = tag[:-2].rstrip() + ">"
        if spans_weg:
            tag = _setze_attr(tag, "spans", None)
        self.roh = tag + "".join(x for _, x in zellen) + "</row>"


class _BlattXml:
    def __init__(self, name: str, teil: str, text: str):
        self.name, self.teil, self.text_alt = name, teil, text
        m = re.search(r"<sheetData\b[^>]*?/?>", text)
        if m is None:
            raise PatchFehler(f"Das Tabellenblatt '{name}' hat keinen Datenbereich; die Datei wird nicht verändert.")
        if m.group(0).endswith("/>"):
            self.kopf, daten, self.schwanz = text[:m.start()] + "<sheetData>", "", "</sheetData>" + text[m.end():]
        else:
            e = text.find("</sheetData>", m.end())
            if e < 0:
                raise PatchFehler(f"Das Tabellenblatt '{name}' ist unvollständig; die Datei wird nicht verändert.")
            self.kopf, daten, self.schwanz = text[:m.end()], text[m.end():e], text[e:]
        self.zeilen: list[_Zeile] = []
        pos = 0
        for zm in _ZEILE.finditer(daten):
            vor = daten[pos:zm.start()]
            if vor.strip():
                raise PatchFehler(f"Im Tabellenblatt '{name}' steht unerwarteter Inhalt; die Datei wird nicht verändert.")
            r = _attrs(_ZEILE_TAG.match(zm.group(0)).group(0)).get("r", "")
            if not r.isdigit() or (self.zeilen and int(r) <= self.zeilen[-1].r):
                raise PatchFehler(f"Im Tabellenblatt '{name}' fehlen Zeilennummern; die Datei wird nicht verändert.")
            self.zeilen.append(_Zeile(int(r), vor, zm.group(0)))
            pos = zm.end()
        self.ende = daten[pos:]
        if self.ende.strip():
            raise PatchFehler(f"Im Tabellenblatt '{name}' steht unerwarteter Inhalt; die Datei wird nicht verändert.")
        self.geaendert = False

    def nummern(self) -> list[int]:
        return [z.r for z in self.zeilen]

    def zeile(self, r: int, anlegen: bool = False) -> _Zeile | None:
        nr = self.nummern()
        i = bisect.bisect_left(nr, r)
        if i < len(nr) and nr[i] == r:
            return self.zeilen[i]
        if not anlegen:
            return None
        z = _Zeile(r, "", f'<row r="{r}"/>')
        self.zeilen.insert(i, z)
        self.geaendert = True
        return z

    def zelle(self, r: int, spalte: int) -> str | None:
        z = self.zeile(r)
        if z is None:
            return None
        for sp, xml in z.zellen():
            if sp == spalte:
                return xml
        return None

    def ersetze_zelle(self, r: int, spalte: int, xml: str) -> None:
        z = self.zeile(r, anlegen=True)
        zellen = z.zellen()
        for eintrag in zellen:
            if eintrag[0] == spalte:
                eintrag[1] = xml
                z.setze_zellen(zellen)
                break
        else:
            i = bisect.bisect_left([sp for sp, _ in zellen], spalte)
            zellen.insert(i, [spalte, xml])
            z.setze_zellen(zellen, spans_weg=True)
        self.geaendert = True

    def render(self) -> str:
        if not self.geaendert:
            return self.text_alt
        return self.kopf + "".join(z.vor + z.roh for z in self.zeilen) + self.ende + self.schwanz


# ===================================================================================== Stile und Texte
class _Stile:
    def __init__(self, text: str | None):
        self.text = text
        self.fonts: list[str] = []
        self.xfs: list[str] = []
        self.neue_fonts: list[str] = []
        self.neue_xfs: list[str] = []
        self._cache: dict[tuple, int] = {}
        if text is None:
            return
        m = re.search(r"<fonts\b[^>]*>(.*?)</fonts>", text, re.S)
        if m:
            self.fonts = re.findall(r"<font\b[^>]*?(?:/>|>.*?</font>)", m.group(1), re.S)
        m = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", text, re.S)
        if m:
            self.xfs = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", m.group(1), re.S)

    def _xf(self, s: int) -> str:
        alle = self.xfs + self.neue_xfs
        return alle[s] if 0 <= s < len(alle) else (alle[0] if alle else '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>')

    def schrift(self, s: int) -> dict[str, str]:
        """Zellschrift eines Zellstils (cellXfs[s] -> fonts[fontId])."""
        try:
            fid = int(_attrs(_TAG.match(self._xf(s)).group(0)).get("fontId", "0"))
        except ValueError:
            fid = 0
        alle = self.fonts + self.neue_fonts
        return _schrift_aus_xml(alle[fid]) if 0 <= fid < len(alle) else {}

    def variante(self, s: int, farbe: str, quote: bool, cfg: dict | None) -> int:
        """Neuer Zellstil = cellXfs[s] mit gefaerbter Schrift und/oder quotePrefix; wird hinten angefuegt (einmal je Variante)."""
        if self.text is None:
            raise PatchFehler("Die Excel hat keine Formatvorlagen; eine Zelle kann nicht eingefärbt werden.")
        key = (s, farbe, quote)
        if key in self._cache:
            return self._cache[key]
        xf = self._xf(s)
        tag = _TAG.match(xf).group(0)
        rest = xf[len(tag):]
        if farbe in (ALT, NEU):
            schrift = _mit_format(self.schrift(s), farbe, cfg)
            self.neue_fonts.append(_render_schrift(schrift, _FONT_REIHE, "font"))
            tag = _setze_attr(tag, "fontId", str(len(self.fonts) + len(self.neue_fonts) - 1))
            tag = _setze_attr(tag, "applyFont", "1")
        if quote:
            tag = _setze_attr(tag, "quotePrefix", "1")
        self.neue_xfs.append(tag + rest)
        self._cache[key] = len(self.xfs) + len(self.neue_xfs) - 1
        return self._cache[key]

    def render(self) -> str | None:
        if self.text is None or not (self.neue_fonts or self.neue_xfs):
            return None
        t = self.text
        for block, neue, alle in (("fonts", self.neue_fonts, len(self.fonts)), ("cellXfs", self.neue_xfs, len(self.xfs))):
            if not neue:
                continue
            m = re.search(rf"<{block}\b[^>]*>", t)
            e = t.find(f"</{block}>", m.end())
            tag = _setze_attr(m.group(0), "count", str(alle + len(neue)))
            t = t[:m.start()] + tag + t[m.end():e] + "".join(neue) + t[e:]
        return t


class _Texte:
    """xl/sharedStrings.xml: neue Eintraege hinten anfuegen."""

    def __init__(self, text: str | None):
        self.text = text
        self.anzahl = len(re.findall(r"<si\b", text)) if text else 0
        self.neu: list[str] = []
        self.delta_count = 0

    def neu_eintrag(self, si: str) -> int:
        self.neu.append(si)
        return self.anzahl + len(self.neu) - 1

    def render(self) -> str | None:
        if not self.neu and not self.delta_count:
            return None
        if self.text is None:
            n = len(self.neu)
            return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{max(n, self.delta_count)}" '
                    f'uniqueCount="{n}">' + "".join(self.neu) + "</sst>")
        t = self.text
        m = re.search(r"<sst\b[^>]*?/?>", t)
        tag = m.group(0)
        a = _attrs(tag)
        neu_tag = tag
        if "count" in a and a["count"].isdigit():
            neu_tag = _setze_attr(neu_tag, "count", str(max(0, int(a["count"]) + self.delta_count)))
        if "uniqueCount" in a:
            neu_tag = _setze_attr(neu_tag, "uniqueCount", str(self.anzahl + len(self.neu)))
        if tag.endswith("/>"):
            neu_tag = neu_tag[:-2].rstrip() + ">"
            return t[:m.start()] + neu_tag + "".join(self.neu) + "</sst>" + t[m.end():]
        e = t.find("<extLst", m.end())
        if e < 0:
            e = t.rfind("</sst>")
        return t[:m.start()] + neu_tag + t[m.end():e] + "".join(self.neu) + t[e:]


# ===================================================================================== Arbeitsmappe
class _Mappe:
    def __init__(self, paket: _Paket, cfg: dict | None):
        self.paket = paket
        self.cfg = cfg
        wb = next((r["_pfad"] for r in paket.rels("") if r.get("Type", "").endswith(_T_DOKUMENT) and "_pfad" in r),
                  "xl/workbook.xml")
        if not paket.gibt(wb):
            raise PatchFehler("In der Datei fehlt die Arbeitsmappe (workbook.xml).")
        self.wb_teil = wb
        self.wb_rels = paket.rels(wb)
        wbtext = paket.text(wb)
        nach_id = {r.get("Id"): r for r in self.wb_rels}
        self.blaetter: list[tuple[str, str]] = []                 # (Name, Teil) in Mappen-Reihenfolge
        sm = re.search(r"<sheets\b[^>]*>(.*?)</sheets>", wbtext, re.S)
        for m in re.finditer(r"<sheet\b[^>]*?/?>", sm.group(1) if sm else ""):
            a = _attrs(m.group(0))
            rel = nach_id.get(_attr_name(a, "id") or "")
            if rel is not None and rel.get("Type", "").endswith(_T_BLATT) and "_pfad" in rel:
                self.blaetter.append((a.get("name", ""), rel["_pfad"]))
        self._xml: dict[str, _BlattXml] = {}
        tt = next((r["_pfad"] for r in self.wb_rels if r.get("Type", "").endswith(_T_TEXTE) and "_pfad" in r), None)
        self.texte_teil = tt
        self.texte = _Texte(paket.text(tt) if tt and paket.gibt(tt) else None)
        st = next((r["_pfad"] for r in self.wb_rels if r.get("Type", "").endswith(_T_STILE) and "_pfad" in r), None)
        self.stile_teil = st
        self.stile = _Stile(paket.text(st) if st and paket.gibt(st) else None)

    def teil(self, blatt: str) -> str:
        for n, t in self.blaetter:
            if n == blatt:
                return t
        for n, t in self.blaetter:
            if n.casefold() == blatt.casefold():
                return t
        raise PatchFehler(f"Das Tabellenblatt '{blatt}' gibt es in der Excel nicht.")

    def blatt(self, name: str) -> _BlattXml:
        teil = self.teil(name)
        if teil not in self._xml:
            echt = next(n for n, t in self.blaetter if t == teil)
            self._xml[teil] = _BlattXml(echt, teil, self.paket.text(teil))
        return self._xml[teil]

    # ------------------------------------------------------------------ Zellen
    def _si(self, stuecke: list[Stueck], zellschrift: dict[str, str]) -> str:
        teile = []
        for st in stuecke:
            if not st.text:
                continue
            basis = _schrift_aus_basis(st.basis) if st.basis is not None else dict(zellschrift)
            if not st.behalten and st.format in (ALT, NEU):
                basis = _mit_format(basis, st.format, self.cfg)
            neu = st.format == NEU and not st.behalten
            teile.append(f'<r>{rpr(basis)}<t xml:space="preserve">{_text_xml(st.text, neu)}</t></r>')
        return "<si>" + "".join(teile) + "</si>"

    def zell_xml(self, ref: str, s: int | None, inhalt: ZellInhalt, war_text: bool) -> str:
        """Zell-XML fuer neuen Inhalt; zaehlt den Bezug auf sharedStrings mit."""
        stil = s
        if inhalt.farbe in (ALT, NEU) and inhalt.stuecke is None or inhalt.text_praefix:
            stil = self.stile.variante(s or 0, inhalt.farbe if inhalt.stuecke is None else "", inhalt.text_praefix, self.cfg)
        sa = f' s="{stil}"' if stil is not None else ""
        ist_text = False
        if inhalt.stuecke is not None:
            st = [x for x in inhalt.stuecke if x.text]
            if not st:
                out = f'<c r="{ref}"{sa}/>'
            else:
                k = self.texte.neu_eintrag(self._si(st, self.stile.schrift(s or 0)))
                out = f'<c r="{ref}"{sa} t="s"><v>{k}</v></c>'
                ist_text = True
        elif inhalt.formel:
            f = inhalt.formel[1:] if inhalt.formel.startswith("=") else inhalt.formel
            out = f'<c r="{ref}"{sa}><f>{escape(f)}</f></c>'
        else:
            w = inhalt.wert
            if w is None or w == "":
                out = f'<c r="{ref}"{sa}/>'
            elif isinstance(w, bool):
                out = f'<c r="{ref}"{sa} t="b"><v>{int(w)}</v></c>'
            elif isinstance(w, (int, float)):
                out = f'<c r="{ref}"{sa}><v>{_zahl(w)}</v></c>'
            elif isinstance(w, str):
                k = self.texte.neu_eintrag(f'<si><t xml:space="preserve">{_text_xml(w)}</t></si>')
                out = f'<c r="{ref}"{sa} t="s"><v>{k}</v></c>'
                ist_text = True
            else:
                raise PatchFehler("Ein Wert kann nicht in die Excel geschrieben werden (unbekannte Art).")
        self.texte.delta_count += int(ist_text) - int(war_text)
        return out

    def setze_zelle(self, z: ZellInhalt) -> None:
        bx = self.blatt(z.blatt)
        alt = bx.zelle(z.zeile, z.spalte)
        a = _attrs(_ZELL_TAG.match(alt).group(0)) if alt else {}
        s = a.get("s")
        if z.stil_von:
            vx = bx.zelle(z.stil_von, z.spalte)
            s = _attrs(_ZELL_TAG.match(vx).group(0)).get("s") if vx else None
        ref = f"{get_column_letter(z.spalte)}{z.zeile}"
        xml = self.zell_xml(ref, int(s) if s and s.isdigit() else None, z, a.get("t") == "s")
        bx.ersetze_zelle(z.zeile, z.spalte, xml)

    # ------------------------------------------------------------------ neue Zeilen
    def zeilen_einfuegen(self, blatt: str, liste: list[tuple[int, NeueZeile]], erg: PatchErgebnis,
                         analyse: "_Analyse") -> None:
        ziel = self.blatt(blatt)
        name = ziel.name
        meld: list[str] = []
        for _, nz in liste:
            if not 0 <= nz.nach < MAX_ZEILE:
                raise PatchFehler("Die Stelle für eine neue Zeile liegt außerhalb des Blatts.")
            for h in analyse.hindernisse(name, nz.nach + 1):
                if h not in meld:
                    meld.append(h)
        if meld:
            raise PatchFehler("Eine neue Zeile kann nicht sicher eingefügt werden: " + " ".join(meld))
        zusatz = (self.cfg or {}).get("zeilen") or config.lade("zuordnung").get("zeilen", {})
        v = _Verschiebung([(nz.nach, nz.vorlage) for _, nz in liste], bool(zusatz.get("listen_summen_erweitern", True)))
        # Vorlagen merken (vor dem Umnummerieren)
        vorlagen = {nz.vorlage: ziel.zeile(nz.vorlage) for _, nz in liste}
        vorlage_tags = {r: (z.tag if z else None) for r, z in vorlagen.items()}
        vorlage_zellen = {r: ([(sp, _attrs(_ZELL_TAG.match(x).group(0))) for sp, x in z.zellen()] if z else [])
                          for r, z in vorlagen.items()}
        # 1. Formeln in allen Blaettern, die auf das Blatt zeigen
        for bname, teil in self.blaetter:
            text_alt = self.paket.text(teil) if teil not in self._xml else None
            if teil != ziel.teil and text_alt is not None and name.casefold() not in text_alt.casefold() \
                    and escape(name, {"'": "&apos;"}).casefold() not in text_alt.casefold():
                continue
            bx = self.blatt(bname)
            self._formeln_verschieben(bx, name, v, ist_ziel=(bx is ziel))
        # 2. Zeilen des Blatts umnummerieren
        for z in ziel.zeilen:
            nr = v.zeile(z.r)
            if nr != z.r:
                tag = z.tag
                neu_tag = _setze_attr(tag, "r", str(nr))
                rest = _ZELL_R.sub(lambda m: f"{m.group(1)}{m.group(2)}{nr}{m.group(4)}", z.roh[len(tag):])
                z.roh, z.r = neu_tag + rest, nr
                ziel.geaendert = True
        # 3. neue Zeilen einsetzen
        geordnet = sorted(liste, key=lambda t: (t[1].nach, t[0]))
        zaehler: Counter = Counter()
        neue_merges: list[str] = []
        for i, nz in geordnet:
            nr = v.zeile(nz.nach) + 1 + zaehler[nz.nach] if nz.nach else 1 + zaehler[0]
            zaehler[nz.nach] += 1
            erg.neue_zeilen[i] = (name, nr)
            xml = self._neue_zeile_xml(nz, nr, vorlage_tags.get(nz.vorlage), vorlage_zellen.get(nz.vorlage, []), name, v)
            pos = bisect.bisect_left(ziel.nummern(), nr)
            ziel.zeilen.insert(pos, _Zeile(nr, "", xml))
            for m in re.finditer(r'<mergeCell\b[^>]*?\bref="([A-Za-z]{1,3})(\d+):([A-Za-z]{1,3})(\d+)"', ziel.schwanz):
                if int(m.group(2)) == int(m.group(4)) == nz.vorlage:
                    neue_merges.append(f'<mergeCell ref="{m.group(1)}{nr}:{m.group(3)}{nr}"/>')
            ziel.geaendert = True
        # 4. Kopf (dimension) und Schwanz des Blatts
        ziel.kopf = self._dimension(ziel.kopf, v, max(ziel.nummern() or [1]))
        ziel.schwanz = self._schwanz(ziel.schwanz, name, v, neue_merges)
        # 5. Namen in der Arbeitsmappe, Zeichnungs-Anker
        self._namen(name, v, erg)
        for rel in self.paket.rels(ziel.teil):
            if rel.get("Type", "").endswith(_T_ZEICHNUNG) and "_pfad" in rel and self.paket.gibt(rel["_pfad"]):
                t = self.paket.text(rel["_pfad"])
                neu = _anker_verschieben(t, v)
                if neu != t:
                    self.paket.setze(rel["_pfad"], neu)
        if v.gewachsen:
            erg.warnungen.append(f"Blatt '{name}': {v.gewachsen} Bereich(e) (z. B. Formeln, Druckbereich, Filter, bedingte "
                                 "Formatierung) enden mit der Vorlagezeile und umfassen jetzt auch die neue Zeile.")
        erg.zeilen_abbildung.setdefault(name, [])
        erg.zeilen_abbildung[name] = sorted(erg.zeilen_abbildung[name] + [nz.nach for _, nz in liste])

    def _formeln_verschieben(self, bx: _BlattXml, ziel: str, v: _Verschiebung, ist_ziel: bool) -> None:
        zeile_jetzt = [0]

        def f_ersetzen(m: re.Match) -> str:
            start, schluss, text = m.group(1), m.group(2), m.group(3)
            a = _attrs(start)
            if "ref" in a:
                if ist_ziel:
                    neu_ref = _ref_verschieben(a["ref"], v, False)
                    if neu_ref and neu_ref != a["ref"]:
                        start = _setze_attr(start + ">", "ref", neu_ref)[:-1]
            if text is None:
                return start + schluss
            roh = unescape(text, _ENT)
            wachsen = v.formeln_wachsen and a.get("t") != "shared"      # geteilte Formel: alle Zellen gleich verschieben
            neu = _formel_verschieben(roh, bx.name, ziel, v, wachsen, zeile_jetzt[0])
            if neu == roh:
                return start + schluss
            return start + ">" + escape(neu) + "</f>"

        for z in bx.zeilen:
            if "<f" not in z.roh:
                continue
            zeile_jetzt[0] = z.r
            neu = _FORMEL.sub(f_ersetzen, z.roh)
            if neu != z.roh:
                z.roh = neu
                bx.geaendert = True
        if not ist_ziel:                                    # bedingte Formatierung/Datenpruefung/Links anderer Blaetter
            neu = _orte_verschieben(_formeln_im_schwanz(bx.schwanz, bx.name, ziel, v), bx.name, ziel, v)
            if neu != bx.schwanz:
                bx.schwanz = neu
                bx.geaendert = True

    def _neue_zeile_xml(self, nz: NeueZeile, nr: int, vtag: str | None, vzellen: list[tuple[int, dict]], blatt: str,
                        v: _Verschiebung) -> str:
        tag = vtag if vtag else '<row r="0">'
        if tag.endswith("/>"):
            tag = tag[:-2].rstrip() + ">"
        for weg in ("hidden", "outlineLevel", "collapsed"):
            tag = _setze_attr(tag, weg, None)
        tag = _setze_attr(tag, "r", str(nr))
        if nz.hoehe is not None:
            tag = _setze_attr(tag, "ht", _zahl(nz.hoehe))
            tag = _setze_attr(tag, "customHeight", "1")
        zeilen_s = _attrs(tag).get("s") if _attrs(tag).get("customFormat") in ("1", "true") else None
        stile = {sp: a.get("s") for sp, a in vzellen}
        spalten = sorted(set(stile) | set(nz.zellen))
        teile = []
        for sp in spalten:
            ref = f"{get_column_letter(sp)}{nr}"
            s = stile.get(sp, zeilen_s)
            if sp not in nz.zellen:
                if s is not None:
                    teile.append(f'<c r="{ref}" s="{s}"/>')
                continue
            inhalt = nz.zellen[sp]
            if not isinstance(inhalt, ZellInhalt):
                inhalt = ZellInhalt(blatt, 0, sp, wert=inhalt)
            if inhalt.stil_von:
                bx = self.blatt(blatt)
                vx = bx.zelle(v.zeile(inhalt.stil_von), sp)
                s = _attrs(_ZELL_TAG.match(vx).group(0)).get("s") if vx else s
            if inhalt.kopie_von:
                teile.append(self._kopie(blatt, v.zeile(inhalt.kopie_von), sp, ref))
                continue
            if inhalt.formel:
                von = inhalt.formel_von or nz.vorlage
                f = inhalt.formel if inhalt.formel.startswith("=") else "=" + inhalt.formel
                try:
                    f = _formel_verschieben(f, blatt, blatt, v, v.formeln_wachsen, von)
                    f = Translator(f, origin=f"{get_column_letter(sp)}{v.zeile(von)}").translate_formula(ref)
                except (TokenizerError, TranslatorError, ValueError) as e:
                    raise PatchFehler("Eine Formel der Vorlagezeile kann nicht in die neue Zeile übertragen werden.") from e
                inhalt = ZellInhalt(inhalt.blatt, inhalt.zeile, sp, formel=f, farbe=inhalt.farbe)
            teile.append(self.zell_xml(ref, int(s) if s and s.isdigit() else None, inhalt, False))
        tag = _setze_attr(tag, "spans", None) if any(sp not in stile for sp in nz.zellen) else tag
        return tag + "".join(teile) + "</row>"

    def _kopie(self, blatt: str, zeile: int, spalte: int, ref: str) -> str:
        """Zelle (schon umnummeriert) mit neuer Adresse kopieren: Stil, Typ und Wert bleiben, eine Formel faellt weg."""
        alt = self.blatt(blatt).zelle(zeile, spalte)
        if alt is None:
            return f'<c r="{ref}"/>'
        tag = _ZELL_TAG.match(alt).group(0)
        neu_tag = _setze_attr(tag, "r", ref)
        if tag.endswith("/>"):
            return neu_tag
        innen = alt[len(tag):-len("</c>")]
        innen = re.sub(r"<f\b[^>]*?(?:/>|>.*?</f>)", "", innen, flags=re.S)
        if _attrs(tag).get("t") == "s":
            self.texte.delta_count += 1
        return neu_tag + innen + "</c>"

    @staticmethod
    def _dimension(kopf: str, v: _Verschiebung, max_zeile: int) -> str:
        m = re.search(r'(<dimension\b[^>]*?\bref=")([^"]*)(")', kopf)
        if not m:
            return kopf
        ref = m.group(2)
        teile = ref.split(":")
        z = _ZELL_REF.match(teile[-1])
        if not z:
            return kopf
        neu_ende = max(v.zeile(int(z.group(4))), max_zeile)
        if len(teile) == 1:
            a = _ZELL_REF.match(teile[0])
            neu = f"{teile[0]}:{a.group(2)}{neu_ende}" if neu_ende != int(a.group(4)) else ref
        else:
            a = _ZELL_REF.match(teile[0])
            anfang = f"{a.group(1)}{a.group(2)}{a.group(3)}{v.zeile(int(a.group(4)))}" if a else teile[0]
            neu = f"{anfang}:{z.group(1)}{z.group(2)}{z.group(3)}{neu_ende}"
        return kopf[:m.start(2)] + neu + kopf[m.end(2):]

    def _schwanz(self, s: str, name: str, v: _Verschiebung, neue_merges: list[str]) -> str:
        # verbundene Zellen: Excel-Regel (keine Erweiterung an der Vorlage); Zeilen-Verbindungen der Vorlage kopieren
        def merge(m: re.Match) -> str:
            neu = _ref_verschieben(m.group(2), v, False)
            return m.group(1) + (neu or m.group(2)) + m.group(3)
        s = re.sub(r'(<mergeCell\b[^>]*?\bref=")([^"]*)(")', merge, s)
        if neue_merges:
            m = re.search(r"<mergeCells\b[^>]*>", s)
            e = s.find("</mergeCells>", m.end())
            n = len(re.findall(r"<mergeCell\b", s[m.end():e])) + len(neue_merges)
            s = s[:m.start()] + _setze_attr(m.group(0), "count", str(n)) + s[m.end():e] + "".join(neue_merges) + s[e:]
        # Filter und Sortierung: wachsen mit der Vorlage
        for el in ("autoFilter", "sortState"):
            s = re.sub(rf'(<{el}\b[^>]*?\bref=")([^"]*)(")',
                       lambda m: m.group(1) + (_ref_verschieben(m.group(2), v, True) or m.group(2)) + m.group(3), s)
        # Hyperlinks: Excel-Regel
        s = re.sub(r'(<hyperlink\b[^>]*?\bref=")([^"]*)(")',
                   lambda m: m.group(1) + (_ref_verschieben(m.group(2), v, False) or m.group(2)) + m.group(3), s)
        # sqref (bedingte Formatierung, Datenpruefung, geschuetzte Bereiche ...), auch x14 (xm:sqref)
        s = re.sub(r'(\ssqref=")([^"]*)(")', lambda m: m.group(1) + _sqref_verschieben(m.group(2), v) + m.group(3), s)
        s = re.sub(r"(<([A-Za-z_][\w.-]*:)?sqref>)(.*?)(</\2sqref>)",
                   lambda m: m.group(1) + _sqref_verschieben(m.group(3), v) + m.group(4), s, flags=re.S)
        # Formeln der bedingten Formatierung und Datenpruefung, interne Hyperlinks
        s = _orte_verschieben(_formeln_im_schwanz(s, name, name, v), name, name, v)
        # Seitenumbrueche: id = letzte Zeile vor dem Umbruch
        def brk(m: re.Match) -> str:
            return m.group(1) + str(v.zeile(int(m.group(2)) + 1) - 1) + m.group(3)
        rb = re.search(r"<rowBreaks\b.*?</rowBreaks>", s, re.S)
        if rb:
            neu = re.sub(r'(<brk\b[^>]*?\bid=")(\d+)(")', brk, rb.group(0))
            s = s[:rb.start()] + neu + s[rb.end():]
        # Steuerelement-Anker (xdr:row, 0-basiert)
        return _anker_verschieben(s, v)

    def _namen(self, name: str, v: _Verschiebung, erg: PatchErgebnis) -> None:
        t = self.paket.text(self.wb_teil)

        def ersetzen(m: re.Match) -> str:
            a = _attrs(m.group(1))
            lokal = a.get("localSheetId")
            eigenes = self.blaetter[int(lokal)][0] if lokal and lokal.isdigit() and int(lokal) < len(self.blaetter) else None
            roh = unescape(m.group(2), _ENT)
            name_ = a.get("name", "")                  # Druckbereich und Filter wachsen wie der Autofilter; Drucktitel nie
            wachsen = name_ in ("_xlnm.Print_Area", "_xlnm._FilterDatabase") or (v.formeln_wachsen and not name_.startswith("_xlnm."))
            try:
                neu = _formel_verschieben(roh, eigenes, name, v, wachsen)
            except (TokenizerError, ValueError, IndexError):
                return m.group(0)
            if neu == roh:
                return m.group(0)
            return m.group(1) + escape(neu, {"'": "&apos;"} if "&apos;" in m.group(2) else {}) + m.group(3)

        neu = re.sub(r"(<definedName\b[^>]*>)(.*?)(</definedName>)", ersetzen, t, flags=re.S)
        if neu != t:
            self.paket.setze(self.wb_teil, neu)

    # ------------------------------------------------------------------ calcChain
    def kette_entfernen(self) -> None:
        for r in self.wb_rels:
            if r.get("Type", "").endswith(_T_KETTE):
                rn = _rels_name(self.wb_teil)
                self.paket.setze(rn, self.paket.text(rn).replace(r["_xml"], "", 1))
                if "_pfad" in r and self.paket.gibt(r["_pfad"]):
                    self.paket.entferne(r["_pfad"])
                    ct = self.paket.text("[Content_Types].xml")
                    neu = re.sub(r'<Override\b[^>]*?PartName="/' + re.escape(r["_pfad"]) + r'"[^>]*?/>', "", ct)
                    if neu != ct:
                        self.paket.setze("[Content_Types].xml", neu)

    def texte_sichern(self) -> None:
        neu = self.texte.render()
        if neu is None:
            return
        if self.texte_teil is None:                        # Teil anlegen: Verweis und Inhaltstyp dazu
            self.texte_teil = posixpath.join(posixpath.dirname(self.wb_teil), "sharedStrings.xml")
            rn = _rels_name(self.wb_teil)
            rt = self.paket.text(rn)
            ids = {int(x) for x in re.findall(r'\bId="rId(\d+)"', rt)}
            rid = f"rId{max(ids | {0}) + 1}"
            rel = f'<Relationship Id="{rid}" Type="{_TEXTE_TYP_URI}" Target="sharedStrings.xml"/>'
            self.paket.setze(rn, rt.replace("</Relationships>", rel + "</Relationships>", 1))
            ct = self.paket.text("[Content_Types].xml")
            ov = f'<Override PartName="/{self.texte_teil}" ContentType="{_TEXTE_INHALT}"/>'
            self.paket.setze("[Content_Types].xml", ct.replace("</Types>", ov + "</Types>", 1))
        self.paket.setze(self.texte_teil, neu)

    def sichern(self) -> None:
        for teil, bx in self._xml.items():
            if bx.geaendert:
                self.paket.setze(teil, bx.render())
        self.texte_sichern()
        st = self.stile.render()
        if st is not None:
            self.paket.setze(self.stile_teil, st)


def _formeln_im_schwanz(s: str, eigenes: str, ziel: str, v: _Verschiebung) -> str:
    def ersetzen(m: re.Match) -> str:
        roh = unescape(m.group(3), _ENT)
        try:
            neu = _formel_verschieben(roh, eigenes, ziel, v, False)
        except (TokenizerError, ValueError, IndexError):
            return m.group(0)
        return m.group(0) if neu == roh else m.group(1) + escape(neu) + m.group(4)
    return re.sub(r"(<(formula[12]?|[A-Za-z_][\w.-]*:f)>)(.*?)(</\2>)", ersetzen, s, flags=re.S)


def _orte_verschieben(s: str, eigenes: str, ziel: str, v: _Verschiebung) -> str:
    """Interne Hyperlinks (location="Blatt!A1") zeigen nach dem Einfuegen weiter auf dieselbe Zelle."""
    def ersetzen(m: re.Match) -> str:
        roh = unescape(m.group(2), _ENT)
        try:
            neu = _formel_verschieben(roh, eigenes, ziel, v, False)
        except (TokenizerError, ValueError, IndexError):
            return m.group(0)
        return m.group(0) if neu == roh else m.group(1) + escape(neu, _ATTR_ESC) + m.group(3)
    return re.sub(r'(<hyperlink\b[^>]*?\slocation=")([^"]*)(")', ersetzen, s)


def _anker_verschieben(text: str, v: _Verschiebung) -> str:
    """<xdr:row> (0-basiert) in Zeichnungen und Steuerelementen."""
    return re.sub(r"(<([A-Za-z_][\w.-]*:)row>)(\d+)(</\2row>)",
                  lambda m: m.group(1) + str(v.zeile(int(m.group(3)) + 1) - 1) + m.group(4), text)


# ===================================================================================== Hindernisse
class _Analyse:
    """Was in einer Datei beim Einfuegen von Zeilen nicht sicher mitgezogen werden kann (einmal je Datei gelesen)."""

    def __init__(self, paket: _Paket):
        mappe = _Mappe(paket, None)                 # nur zum Lesen; das Paket selbst wird nicht aufbewahrt (Speicher)
        p = paket
        self.namen = [n for n, _ in mappe.blaetter]
        self.extern = any(r.get("Type", "").endswith(_T_EXTERN) for r in mappe.wb_rels)
        self.pivot_quellen: set[str] = set()
        for r in mappe.wb_rels:
            if r.get("Type", "").endswith(_T_PIVOT_CACHE) and "_pfad" in r and p.gibt(r["_pfad"]):
                for m in re.finditer(r"<worksheetSource\b[^>]*?/?>", p.text(r["_pfad"])):
                    self.pivot_quellen.add(_attrs(m.group(0)).get("sheet", "").casefold())
        wbtext = p.text(mappe.wb_teil)
        self.namens_formeln = [(unescape(m.group(1), _ENT)) for m in
                               re.finditer(r"<definedName\b[^>]*>(.*?)</definedName>", wbtext, re.S)]
        self.blatt: dict[str, dict[str, Any]] = {}
        self.formeln: list[tuple[str, int, int, dict[str, str], str]] = []    # (Blatt, Zeile, Spalte, f-Attribute, Text)
        self.schwanz_formeln: list[tuple[str, str]] = []                      # (Blatt, Formel) aus CF/DV
        for name, teil in mappe.blaetter:
            self._blatt_lesen(p, name, teil)
        self._info: dict[str, dict[str, Any]] = {}

    def _blatt_lesen(self, p: _Paket, name: str, teil: str) -> None:
        text = p.text(teil)
        e: dict[str, Any] = {"tabellen": bool(re.search(r"<tableParts\b[^>]*count=\"[1-9]", text))
                             or bool(re.search(r"<tablePart\b", text)),
                             "merges": [], "versteckt": set(), "max": 0, "kommentare": [], "vml": [],
                             "diagramme": [], "pivot": False, "geteilt": [], "matrix": [], "datentabelle": False}
        for m in re.finditer(r'<mergeCell\b[^>]*?\bref="([^"]*)"', text):
            z = _bezug_zeilen(m.group(1))
            if z and len(z) == 2:
                e["merges"].append((z[0][0], z[1][0]))
        for m in _ZEILE_TAG.finditer(text):
            a = _attrs(m.group(0))
            r = int(a.get("r", "0") or 0)
            e["max"] = max(e["max"], r)
            if a.get("hidden") in ("1", "true") or (a.get("outlineLevel", "0") not in ("0", "")):
                e["versteckt"].add(r)
        for m in re.finditer(r'(<c\b[^>]*?\sr="[A-Za-z]{1,3}(\d+)"[^>]*(?<!/)>)(?:(?!</c>).)*?(<f\b[^>]*?(?:/>|>(.*?)</f>))',
                             text, re.S):
            zeile = int(m.group(2))
            spalte_m = re.search(r'\sr="([A-Za-z]{1,3})', m.group(1))
            spalte = column_index_from_string(spalte_m.group(1).upper())
            fa = _attrs(re.match(r"<f\b[^>]*", m.group(3)).group(0))
            ftext = unescape(m.group(4), _ENT) if m.group(4) is not None else ""
            self.formeln.append((name, zeile, spalte, fa, ftext))
            if fa.get("t") == "shared" and "ref" in fa:
                z = _bezug_zeilen(fa["ref"])
                if z:
                    e["geteilt"].append((z[0][0], z[-1][0], zeile, ftext))
            elif fa.get("t") == "array" and "ref" in fa:
                z = _bezug_zeilen(fa["ref"])
                if z:
                    e["matrix"].append((z[0][0], z[-1][0]))
            elif fa.get("t") == "dataTable":
                e["datentabelle"] = True
        for m in re.finditer(r"<(formula[12]?|[A-Za-z_][\w.-]*:f)>(.*?)</\1>", text, re.S):
            self.schwanz_formeln.append((name, unescape(m.group(2), _ENT)))
        for rel in p.rels(teil):
            typ, pf = rel.get("Type", ""), rel.get("_pfad")
            if pf is None or not p.gibt(pf):
                continue
            if typ.endswith(_T_KOMMENTARE) or typ.endswith(_T_THREAD):
                for m in re.finditer(r'<(?:\w+:)?(?:comment|threadedComment)\b[^>]*?\bref="([^"]*)"', p.text(pf)):
                    z = _bezug_zeilen(m.group(1))
                    if z:
                        e["kommentare"].append(z[0][0])
            elif typ.endswith(_T_VML):
                vt = p.text(pf)
                e["vml"] += [int(x) + 1 for x in re.findall(r"<x:Row>\s*(\d+)\s*</x:Row>", vt)]
                for a in re.findall(r"<x:Anchor>\s*([\d,\s]+)</x:Anchor>", vt):
                    teile = [x.strip() for x in a.split(",")]
                    if len(teile) == 8:
                        e["vml"] += [int(teile[2]) + 1, int(teile[6]) + 1]
            elif typ.endswith(_T_PIVOT):
                e["pivot"] = True
            elif typ.endswith(_T_ZEICHNUNG):
                for dr in p.rels(pf):
                    if dr.get("Type", "").endswith(_T_DIAGRAMM) and dr.get("_pfad") and p.gibt(dr["_pfad"]):
                        for fm in re.finditer(r"<(?:\w+:)?f>(.*?)</(?:\w+:)?f>", p.text(dr["_pfad"]), re.S):
                            e["diagramme"].append(unescape(fm.group(1), _ENT))
        self.blatt[name.casefold()] = e

    def info(self, formel: str) -> dict[str, Any]:
        if formel not in self._info:
            self._info[formel] = _formel_info(formel)
        return self._info[formel]

    def hindernisse(self, blatt: str, zeile: int) -> list[str]:
        """Einfache deutsche Saetze: was beim Einfuegen einer Zeile an `zeile` (sie bekommt diese Nummer, alles ab hier
        rutscht nach unten) nicht sicher mitgezogen werden kann. Leer = unbedenklich."""
        b = blatt.casefold()
        e = self.blatt.get(b)
        if e is None:
            return [f"Das Tabellenblatt '{blatt}' gibt es in der Excel nicht."]
        z = zeile
        out: list[str] = []

        def add(t: str) -> None:
            if t not in out:
                out.append(t)

        if e["tabellen"]:
            add("Im Blatt gibt es eine Excel-Tabelle (Tabellenobjekt); ihr Bereich wächst beim Einfügen nicht sicher mit.")
        if e["pivot"] or b in self.pivot_quellen:
            add("Eine Pivot-Tabelle nutzt dieses Blatt; sie würde beim Einfügen nicht angepasst.")
        if self.extern:
            add("Die Excel ist mit anderen Dateien verknüpft; ob Bezüge dort verrutschen, lässt sich nicht prüfen.")
        if any(r0 < z <= r1 for r0, r1 in e["merges"]):
            add("Hier laufen verbundene Zellen durch die Zeile; sie würden zerrissen oder verrutschen.")
        if z - 1 in e["versteckt"] and z in e["versteckt"]:
            add("Die neue Zeile käme mitten in eine Folge ausgeblendeter oder gruppierter Zeilen; die Folge würde unterbrochen.")
        if any(r >= z for r in e["kommentare"]):
            add("Unterhalb gibt es Kommentare (Notizen); sie werden beim Einfügen nicht mitverschoben.")
        if any(r >= z for r in e["vml"]):
            add("Unterhalb gibt es Steuerelemente oder alte Zeichnungen; sie werden beim Einfügen nicht mitverschoben.")
        if e["max"] >= MAX_ZEILE:
            add("Die letzte Zeile des Blatts ist belegt; sie ginge beim Einfügen verloren.")
        if e["datentabelle"]:
            add("Im Blatt gibt es eine Datentabelle (Mehrfachoperation); sie wird beim Einfügen nicht angepasst.")
        for r0, r1 in e["matrix"]:
            if r0 < z <= r1:
                add("Eine Matrixformel läuft über die Einfügestelle; sie kann nicht sicher angepasst werden.")
        for r0, r1, master, ftext in e["geteilt"]:
            if r0 < z <= r1:
                add("Eine geteilte Formel läuft über die Einfügestelle; sie kann nicht sicher angepasst werden.")
                continue
            if not self._gleichmaessig(ftext, blatt, blatt, r0 - master, r1 - master, z):
                add("Eine geteilte Formel zeigt über die Einfügestelle hinweg; sie kann nicht sicher angepasst werden.")
        for fname, fz, _, fa, ftext in self.formeln:
            if not ftext:
                continue
            eigen = fname.casefold() == b
            if not eigen and b not in ftext.casefold():
                continue
            i = self.info(ftext)
            if i["fehler"]:
                add("Eine Formel kann nicht gelesen werden; ihre Bezüge würden beim Einfügen nicht angepasst.")
            if i["dynamisch"]:
                add("Eine Formel nutzt INDIREKT oder BEREICH.VERSCHIEBEN; ihre Bezüge können beim Einfügen nicht angepasst werden.")
            if i["drei_d"]:
                add("Eine Formel nutzt einen Bezug über mehrere Blätter (3D-Bezug); er wird beim Einfügen nicht angepasst.")
            if not eigen and fa.get("t") == "shared" and "ref" in fa:
                g = _bezug_zeilen(fa["ref"])
                if g and not self._gleichmaessig(ftext, fname, blatt, g[0][0] - fz, g[-1][0] - fz, z):
                    add("Eine geteilte Formel zeigt über die Einfügestelle hinweg; sie kann nicht sicher angepasst werden.")
        for fname, ftext in self.schwanz_formeln:
            if fname.casefold() != b and b not in ftext.casefold():
                continue
            i = self.info(ftext)
            if i["dynamisch"] or i["drei_d"] or i["fehler"]:
                add("Eine bedingte Formatierung oder Datenprüfung nutzt einen Bezug, der beim Einfügen nicht angepasst werden kann.")
        for ftext in self.namens_formeln:
            i = self.info(ftext)
            if i["drei_d"] or (i["dynamisch"] and b in ftext.casefold()):
                add("Ein benannter Bereich nutzt einen Bezug, der beim Einfügen nicht angepasst werden kann.")
        for ftext in e["diagramme"]:
            i = self.info(ftext)
            for bl, zeilen in i["bezuege"]:
                if (bl or "").casefold() == b and any(r >= z for r, _ in zeilen):
                    add("Ein Diagramm zeigt auf Zeilen ab dieser Stelle; sein Bezug wird beim Einfügen nicht angepasst.")
        for name in self.namen:                             # Diagramme in anderen Blaettern
            if name.casefold() == b:
                continue
            for ftext in self.blatt[name.casefold()]["diagramme"]:
                for bl, zeilen in self.info(ftext)["bezuege"]:
                    if (bl or "").casefold() == b and any(r >= z for r, _ in zeilen):
                        add("Ein Diagramm zeigt auf Zeilen ab dieser Stelle; sein Bezug wird beim Einfügen nicht angepasst.")
        return out

    def _gleichmaessig(self, ftext: str, eigenes: str, ziel: str, von: int, bis: int, z: int) -> bool:
        """Verschieben sich alle Zellen einer geteilten Formel gleich? (relative Zeilen r+von .. r+bis duerfen die Stelle z
        nicht ueberqueren)."""
        i = self.info(ftext)
        if i["fehler"]:
            return False
        for bl, zeilen in i["bezuege"]:
            if (bl if bl is not None else eigenes).casefold() != ziel.casefold():
                continue
            for r, absolut in zeilen:
                lo, hi = (r, r) if absolut else (r + von, r + bis)
                if lo < z <= hi:
                    return False
        return True


_ANALYSEN: dict[tuple[str, int, int], _Analyse] = {}


def _analyse(quelle: Path) -> _Analyse:
    st = quelle.stat()
    key = (str(quelle.resolve()), st.st_mtime_ns, st.st_size)
    if key not in _ANALYSEN:
        if len(_ANALYSEN) >= 4:
            _ANALYSEN.clear()
        _ANALYSEN[key] = _Analyse(_Paket(quelle))
    return _ANALYSEN[key]


def hindernisse(quelle: str | Path, blatt: str, zeile: int) -> list[str]:
    """Was beim Einfuegen einer Zeile an `zeile` (die neue Zeile bekommt diese Nummer; alle Zeilen ab hier rutschen
    nach unten -- wie ExcelModell.einfuege_hindernisse) im Blatt `blatt` nicht sicher mitgezogen werden kann.
    Leere Liste = der Patch-Schreiber kann es sauber. Alles andere (Formeln, bedingte Formatierung, Datenpruefung,
    Filter, Druckbereich, Namen, verbundene Zellen darunter, ausgeblendete Zeilen, Hyperlinks, Seitenumbrueche,
    Bilder/Formen) wird beim Einfuegen mitgezogen."""
    try:
        return _analyse(Path(quelle)).hindernisse(blatt, zeile)
    except PatchFehler as e:
        return [e.meldung]


# ===================================================================================== Schreiben
def schreibe(quelle: str | Path, ziel: str | Path, zellen: Iterable[ZellInhalt] = (),
             neue_zeilen: Iterable[NeueZeile] = (), cfg: dict | None = None) -> PatchErgebnis:
    """Schreibt eine Kopie von `quelle` nach `ziel`, in der nur die genannten Zellen und neuen Zeilen anders sind.

    Eine vorhandene Zieldatei wird nie ueberschrieben (Regel 29). Was nicht sicher geht, bricht mit PatchFehler ab;
    dann gibt es keine Ausgabedatei."""
    quelle, ziel = Path(quelle), Path(ziel)
    if ziel.exists():
        raise DateiVorhanden("Die Ergebnis-Datei gibt es schon und wird nicht überschrieben. "
                             "Bitte einen neuen Dateinamen oder einen neuen Lauf nehmen.")
    zellen, neue = list(zellen), list(neue_zeilen)
    erg = PatchErgebnis(ausgabe=str(ziel), neue_zeilen=[("", 0)] * len(neue))
    ziel.parent.mkdir(parents=True, exist_ok=True)
    if not zellen and not neue:
        paket = _Paket(quelle)
        paket.schreibe(ziel)
        return erg
    paket = _Paket(quelle)
    mappe = _Mappe(paket, cfg)
    gesehen: set[tuple[str, int, int]] = set()
    for z in zellen:
        key = (mappe.teil(z.blatt), z.zeile, z.spalte)
        if key in gesehen:
            raise PatchFehler("Eine Zelle sollte zweimal geschrieben werden; die Datei wird nicht verändert.")
        if not (1 <= z.zeile <= MAX_ZEILE and 1 <= z.spalte <= 16384):
            raise PatchFehler("Eine Zelle liegt außerhalb des Blatts.")
        gesehen.add(key)
        mappe.setze_zelle(z)
    if neue:
        analyse = _analyse(quelle)
        je_blatt: dict[str, list[tuple[int, NeueZeile]]] = {}
        for i, nz in enumerate(neue):
            je_blatt.setdefault(mappe.teil(nz.blatt), []).append((i, nz))
        for teil, liste in je_blatt.items():
            name = next(n for n, t in mappe.blaetter if t == teil)
            mappe.zeilen_einfuegen(name, liste, erg, analyse)
    mappe.sichern()
    mappe.kette_entfernen()
    paket.schreibe(ziel)
    erg.geaenderte_teile = paket.geaendert()
    return erg


# ===================================================================================== Pruefen
def _zip_inhalt(pfad: Path) -> tuple[list[str], dict[str, bytes]]:
    with zipfile.ZipFile(pfad) as z:
        namen = z.namelist()
        return namen, {n: z.read(n) for n in namen}


def pruefe_treue(quelle: str | Path, ziel: str | Path, geaenderte_teile: Iterable[str] | None = None) -> list[str]:
    """Leer = alle unberuehrten Teile sind entpackt bytegleich, die Teile-Liste ist gleich (bis auf die entfernte calcChain
    und neu angelegte, als geaendert gemeldete Teile) und jeder geaenderte XML-Teil ist wohlgeformt.
    Ohne `geaenderte_teile` wird jeder veraenderte Teil gemeldet."""
    na, a = _zip_inhalt(Path(quelle))
    nb, b = _zip_inhalt(Path(ziel))
    erlaubt = set(geaenderte_teile or ())
    out: list[str] = []
    kette = {n for n in na if n.endswith("calcChain.xml")}
    for n in na:
        if n not in b and n not in kette:
            out.append(f"Teil fehlt im Ergebnis: {n}")
    for n in nb:
        if n not in a and n not in erlaubt:
            out.append(f"Unerwarteter neuer Teil: {n}")
    gemeinsam_a = [n for n in na if n in b]
    gemeinsam_b = [n for n in nb if n in a]
    if gemeinsam_a != gemeinsam_b:
        out.append("Die Reihenfolge der Teile hat sich geändert.")
    for n in gemeinsam_a:
        if a[n] != b[n] and n not in erlaubt:
            out.append(f"Teil verändert: {n}")
    for n in nb:
        if n in erlaubt and (n.endswith(".xml") or n.endswith(".rels")):
            try:
                ET.fromstring(b[n])
            except ET.ParseError:
                out.append(f"Teil nicht wohlgeformt: {n}")
    return out
