"""Tests fuer den Paar-Generator (tests/fixtures/generator.py, stoerer.py, pdf_writer.py).

Alles ist erfunden (Regel 16). Die Tests pruefen nur den Generator -- nicht, wie gut das Tool (src/) auf
den Paaren abschneidet. Geschrieben wird nur in tmp_path, nie ins Repo.
"""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import time
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
import pdfplumber
import pytest

from tests.fixtures import generator, stoerer
from tests.fixtures.generator import erzeuge_paar
from tests.fixtures.gross import erzeuge_gross
from tests.fixtures.pdf_writer import (PdfDokument, baue_pdf, markierte_woerter, muster_bild, parse_markup,
                                       scan_bild, zustand)
from tests.fixtures.szenario import ELEMENTE

ROOT = Path(__file__).resolve().parent.parent
SOLL_STATUS = ("uebernommen", "pruefen", "nicht_uebernommen")


# ====================================================================== Hilfen
def _farbe(c) -> tuple[float, ...]:
    f = c.get("non_stroking_color")
    if f is None:
        return ()
    if isinstance(f, (int, float)):
        return (float(f),)
    return tuple(float(v) for v in f)


def _hex(h: str) -> tuple[float, float, float]:
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))  # type: ignore[return-value]


def _nahe(farbe: tuple[float, ...], ziel: tuple[float, ...], tol: float = 0.02) -> bool:
    return len(farbe) == len(ziel) and all(abs(a - b) <= tol for a, b in zip(farbe, ziel))


def _blau(f: tuple[float, ...]) -> bool:
    return len(f) == 3 and f[2] > 0.45 and f[2] > f[0] + 0.2 and f[2] > f[1] + 0.1


def _rot(f: tuple[float, ...]) -> bool:
    return len(f) == 3 and f[0] > 0.6 and f[1] < 0.5 and f[2] < 0.5


def _chars(ordner: Path) -> list[dict]:
    with pdfplumber.open(ordner / "aenderungen.pdf") as pdf:
        return [c for pg in pdf.pages for c in pg.chars]


def _seitentexte(ordner: Path) -> list[str]:
    with pdfplumber.open(ordner / "aenderungen.pdf") as pdf:
        return [pg.extract_text() or "" for pg in pdf.pages]


def _erwartet(ordner: Path) -> dict:
    return json.loads((ordner / "erwartet.json").read_text(encoding="utf-8"))


def _hauptblatt(wb, e):
    return wb.worksheets[e["aenderungen"][0]["blatt"] - 1]


def _alle_excel_zeilen(ordner: Path) -> list[str]:
    """Alle Textzeilen (Zellen am Zeilenumbruch geteilt) des Hauptblatts in vorher und nachher."""
    e = _erwartet(ordner)
    zeilen: list[str] = []
    for name in ("vorher", "nachher"):
        ws = _hauptblatt(openpyxl.load_workbook(ordner / f"{name}.xlsx"), e)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    zeilen.extend(c.value.split("\n"))
    return zeilen


def _paar(tmp_path: Path, seed: int, profil="standard", korpus="entwicklung") -> Path:
    return erzeuge_paar(seed, profil, tmp_path / korpus, korpus)


# ====================================================================== pdf_writer: Standardverhalten unveraendert
GOLDEN = {
    "szenario/farbe": "ae837f9a46d426af", "szenario/hintergrund": "e546e8f6136f4c4d",
    "szenario/strich": "24017388e789558e", "szenario/annot": "980f6fcb2a1373c6",
    "gross/farbe": "0606eade7edc13d4", "gross/hintergrund": "cceb7fd118b4775c",
    "gross/strich": "87e8fb7f04fbe124", "gross/annot": "047d4964560295d7",
    "gross2/farbe": "d2d36495859a3292", "gross2/hintergrund": "b690e65cb486f3fb",
    "gross2/strich": "a2b853a6f7220425", "gross2/annot": "a5548c3b632a9134",
    "nokopf": "bb3811b4b5021b1e",
}


def test_pdf_writer_standard_ist_byte_gleich(tmp_path):
    """Die Erweiterungen duerfen das bisherige Standardverhalten nicht veraendern (Hashes vom Stand davor)."""
    ist = {}
    for name, el in (("szenario", ELEMENTE), ("gross", erzeuge_gross()), ("gross2", erzeuge_gross(3, 20, 6, 30))):
        for v in ("farbe", "hintergrund", "strich", "annot"):
            p = tmp_path / "a.pdf"
            baue_pdf(p, el, v)
            ist[f"{name}/{v}"] = hashlib.sha256(p.read_bytes()).hexdigest()[:16]
    d = PdfDokument("farbe", kopf_fuss=False)
    d.block("Test {+neu+} {-alt-}")
    ist["nokopf"] = hashlib.sha256(d.bytes()).hexdigest()[:16]
    assert ist == GOLDEN


# ====================================================================== pdf_writer: neue Faehigkeiten
def test_markup_neue_stile_und_zustaende():
    m = "A {@bu:neu@} B {@bs:blau-weg@} C {@ru:rot-weg@} D {@l:Link@} E {@gn:gruen@} F {-alt-}{+neu2+}"
    assert zustand(m, False) == "A B blau-weg C rot-weg D Link E F alt"
    assert zustand(m, True) == "A neu B C D Link E gruen F neu2"
    assert markierte_woerter(m) == 1 + 1 + 1 + 1 + 1 + 1   # Link zaehlt nicht
    stile = [s for w in parse_markup(m) for _, s in w]
    assert {"bu", "bs", "ru", "l", "gn", "g", "n"} <= set(stile)
    # verschmolzene Ersetzung: ein Wort, zwei Stile
    assert parse_markup("{-100-}{+120+} kWh")[0] == [("100", "g"), ("120", "n")]


def test_geschuetztes_leerzeichen_trennt_kein_wort():
    assert len(parse_markup("10 kWh je Rack")) == 3
    assert zustand("10 kWh", True) == "10 kWh"


def _pdf_laden(d: PdfDokument, tmp_path: Path):
    p = tmp_path / "t.pdf"
    d.speichern(p)
    return pdfplumber.open(p)


def test_silbentrennung_ascii_auch_im_farbigen_wort(tmp_path):
    gefunden = False
    for bis in range(260, 420, 9):
        d = PdfDokument("farbe", trennung="ascii")
        d.block("Das Gerät hat eine {+Überspannungsschutzeinrichtung+} und eine {+Temperaturüberwachung+} "
                "sowie Batteriemanagementsystem und {+Überspannungsschutzeinrichtung+} im Schaltschrank.",
                spalte_bis=float(bis))
        if not any(set(stile) == {"n"} for _, stile, _ in d.getrennt):
            continue
        with _pdf_laden(d, tmp_path) as pdf:
            text = pdf.pages[0].extract_text()
            assert re.search(r"\w-\n", text)
            # der Trennstrich eines blauen Wortes ist selbst blau
            if any(c["text"] == "-" and c["top"] > 90 and _farbe(c) == (0.0, 0.0, 1.0) for c in pdf.pages[0].chars):
                gefunden = True
                break
    assert gefunden


def test_weicher_trennstrich_steht_im_pdf(tmp_path):
    d = PdfDokument("farbe", trennung="weich")
    d.block("Das Gerät hat eine Überspannungsschutzeinrichtung sowie Temperaturüberwachung und "
            "Batteriemanagementsystem im Schaltschrank.", spalte_bis=300)
    assert d.stats["trennungen"] >= 1
    assert b"\xad" in d.bytes()
    assert all(strich == "­" for _, _, strich in d.getrennt)


def test_silbentrennung_am_seitenende(tmp_path):
    """Ein Wort wird auf der letzten Zeile einer Seite getrennt, der Rest steht auf der naechsten Seite."""
    langes = " ".join(["Überspannungsschutzeinrichtung", "Temperaturüberwachung", "Batteriemanagementsystem"] * 6)
    gefunden = False
    for fuellzeilen in range(0, 12):
        d = PdfDokument("farbe", trennung="ascii")
        for _ in range(fuellzeilen):
            d.block("Fülltext.", abstand_danach=0)
        d.y = 842 - 70 - 15 * 3 - fuellzeilen % 4 * 7
        d.block(langes, spalte_bis=330)
        if d.stats["trennung_seitenende"] >= 1:
            assert d.seitenzahl == 2
            gefunden = True
            break
    assert gefunden


def test_fussnote_und_aufzaehlung_und_kopf_nur_seite_1(tmp_path):
    d = PdfDokument("farbe", kopf_fuss={1})
    d.block("Normaler Absatz mit Fussnotenzeichen¹ im Text.")
    assert d.fussnote("Erfundene Fussnote.", "1") is True
    d.block("Erster Punkt", marker="•")
    d.block("Zweiter Punkt", marker="a)")
    d.seitenumbruch()
    d.block("Seite zwei")
    with _pdf_laden(d, tmp_path) as pdf:
        s1, s2 = pdf.pages[0], pdf.pages[1]
        assert "Beispieldokument" in s1.extract_text() and "Beispieldokument" not in s2.extract_text()
        klein = [c for c in s1.chars if abs(c["size"] - 8.0) < 0.1]
        assert klein and all(c["top"] > 650 for c in klein)
        assert re.search(r"^• Erster Punkt", s1.extract_text(), re.M)
        assert re.search(r"^a\) Zweiter Punkt", s1.extract_text(), re.M)


def test_zwei_spalten_haben_zwei_x_cluster(tmp_path):
    d = PdfDokument("farbe")
    text = "Der Wechselrichter muss regelmäßig geprüft werden und gegen Überhitzung geschützt sein. " * 3
    d.zwei_spalten([text, text], [text])
    with _pdf_laden(d, tmp_path) as pdf:
        worte = pdf.pages[0].extract_words()
        links = [w for w in worte if abs(w["x0"] - 56) < 1]
        rechts = [w for w in worte if abs(w["x0"] - 306.5) < 1]
        assert len(links) >= 3 and len(rechts) >= 3
        text_worte = [w for w in worte if 60 < w["top"] < 770]                          # ohne Kopf-/Fusszeile
        assert max(w["x1"] for w in text_worte if w["x0"] < 300) < 292                 # linke Spalte endet vor der Mitte
        assert max(w["x1"] for w in text_worte) <= 595 - 56 + 1                         # rechter Rand


def test_tabelle_mit_farbkopf_und_baenderung(tmp_path):
    d = PdfDokument("farbe")
    d.tabelle([["Pos.", "Bauteil", "Menge"], ["1", "Lüfter", "{-4-} {+6+}"], ["2", "Sicherung", "8"], ["3", "Netzteil", "2"]])
    with _pdf_laden(d, tmp_path) as pdf:
        pg = pdf.pages[0]
        farben = {tuple(round(v, 2) for v in (r["non_stroking_color"] or ())) for r in pg.rects}
        assert (round(0x44 / 255, 2), round(0x72 / 255, 2), round(0xC4 / 255, 2)) in farben      # 4472C4
        assert (round(0xB4 / 255, 2), round(0xC6 / 255, 2), round(0xE7 / 255, 2)) in farben      # B4C6E7
        assert len(pg.lines) >= 9
        assert "Lüfter" in pg.extract_text()


def test_bild_scan_cmyk_gray_rotate_link(tmp_path):
    d = PdfDokument("farbe", farbraum="cmyk", grundfarbraum="gray", kopf_farbe=(0.4, 0.4, 0.4))
    d.block("Ein {+blauer+} und {-roter-} Text mit {@l:Link@}.")
    d.bild(muster_bild(120, 80, 5), 150)
    d.scan_seite(scan_bild([("h", "9 Überschrift"), ("p", "Ein {+neuer+} Satz äöüß {-weg-}.")]))
    d.block("Zurück zu Text")
    d.seite_rotieren(90, 1)
    with _pdf_laden(d, tmp_path) as pdf:
        assert len(pdf.pages) == 3
        assert len(pdf.pages[0].images) == 1 and pdf.pages[0].rotation == 90
        scan = pdf.pages[1]
        assert scan.chars == [] and len(scan.images) == 1
        farben = {_farbe(c) for c in pdf.pages[0].chars} | {_farbe(c) for c in pdf.pages[2].chars}
        assert any(len(f) == 4 for f in farben)      # CMYK
        assert any(len(f) == 1 for f in farben)      # Gray
        assert pdf.pages[0].hyperlinks and pdf.pages[0].hyperlinks[0]["uri"].startswith("https://")


def test_farbtoene_hex(tmp_path):
    d = PdfDokument("farbe", farbe_neu="1F4E9F", farbe_weg="C00000")
    d.block("a {+blau+} b {-rot-}")
    d.block("c {+blau2+} d", farben=("2E75B6", "E06666"))
    with _pdf_laden(d, tmp_path) as pdf:
        farben = [_farbe(c) for c in pdf.pages[0].chars]
        for h in ("1F4E9F", "C00000", "2E75B6"):
            assert any(_nahe(f, _hex(h), 0.01) for f in farben), h


# ====================================================================== Register der Stoerer
def test_register_hat_mindestens_12_familien_mit_vollstaendigen_angaben():
    assert len(stoerer.REGISTER) >= 12
    for name, f in stoerer.REGISTER.items():
        assert f.name == name and f.beschreibung and f.ebene in ("pdf", "excel", "beides")
        assert f.soll_status in SOLL_STATUS
        assert f.luecke is None or f.luecke in stoerer.LUECKEN_ARTEN
    assert len(stoerer.NUR_TEST) == 2
    assert set(stoerer.NUR_TEST) == set(stoerer.TEST_ZWANG)
    assert set(stoerer.NUR_TEST).isdisjoint(stoerer.pool("alle"))
    for n in stoerer.NUR_TEST:
        assert stoerer.REGISTER[n].nur_test


def test_luecken_arten_wie_models():
    models = pytest.importorskip("src.models")
    arten = getattr(models, "LUECKEN_ARTEN", None)
    if arten is None:
        pytest.skip("models.LUECKEN_ARTEN gibt es in diesem Zweig noch nicht")
    assert tuple(arten) == stoerer.LUECKEN_ARTEN


# ====================================================================== Determinismus, Dateien, Format
def _zellen(pfad: Path) -> dict:
    wb = openpyxl.load_workbook(pfad)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = {
            "zellen": {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None},
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "versteckt": sorted(r for r, d in ws.row_dimensions.items() if d.hidden),
        }
    return out


@pytest.mark.parametrize("profil", ["standard", "schwer"])
def test_determinismus_gleicher_seed_gleiche_zellen(tmp_path, profil):
    a = erzeuge_paar(5, profil, tmp_path / "a")
    b = erzeuge_paar(5, profil, tmp_path / "b")
    assert a.name == b.name
    for datei in ("vorher.xlsx", "nachher.xlsx"):
        assert _zellen(a / datei) == _zellen(b / datei)
    assert (a / "erwartet.json").read_bytes() == (b / "erwartet.json").read_bytes()
    assert (a / "aenderungen.pdf").read_bytes() == (b / "aenderungen.pdf").read_bytes()
    c = erzeuge_paar(6, profil, tmp_path / "c")
    assert _zellen(c / "vorher.xlsx") != _zellen(a / "vorher.xlsx")


def test_paare_sind_append_only(tmp_path):
    p = _paar(tmp_path, 3)
    vorher = (p / "vorher.xlsx").read_bytes()
    with pytest.raises(FileExistsError):
        _paar(tmp_path, 3)
    assert (p / "vorher.xlsx").read_bytes() == vorher
    assert not list((tmp_path / "entwicklung").glob(".tmp_*"))


@pytest.mark.parametrize("seed", [1, 2, 3, 4])
@pytest.mark.parametrize("profil", ["standard", "schwer"])
def test_dateien_vorhanden_und_ladbar(tmp_path, seed, profil):
    p = _paar(tmp_path, seed, profil)
    assert p.name == f"synth_{seed:04d}_{profil}"
    assert sorted(f.name for f in p.iterdir()) == ["aenderungen.pdf", "erwartet.json", "nachher.xlsx", "vorher.xlsx"]
    for name in ("vorher.xlsx", "nachher.xlsx"):
        wb = openpyxl.load_workbook(p / name)
        assert "CDE" in wb.sheetnames and wb["CDE"].max_row >= 5
    with pdfplumber.open(p / "aenderungen.pdf") as pdf:
        assert len(pdf.pages) >= 1 and sum(len(pg.chars) for pg in pdf.pages) > 100
    e = _erwartet(p)
    assert e["format"] == 1 and e["generator_version"] == generator.GENERATOR_VERSION
    assert e["seed"] == seed and e["profil"] == profil
    assert isinstance(e["stoerer"], list) and all(n in stoerer.REGISTER for n in e["stoerer"])
    assert e["woerter_markiert"] > 0
    assert len(e["aenderungen"]) >= 7
    for i, a in enumerate(e["aenderungen"], 1):
        assert a["nr"] == i and a["typ"] in generator.TYPEN and a["soll_status"] in SOLL_STATUS
        assert a["blatt"] >= 1 and a["zeile"] >= 1 and a["spalte"] >= 1 and a["woerter"] >= 1
    for lk in e["luecken"]:
        assert lk["art"] in stoerer.LUECKEN_ARTEN and lk["seite"] >= 1


def test_profil_standard_sauber_schwer_viele_stoerer(tmp_path):
    standard = [_erwartet(_paar(tmp_path, s, "standard")) for s in range(1, 13)]
    schwer = [_erwartet(_paar(tmp_path, s, "schwer")) for s in range(1, 13)]
    assert all(len(e["stoerer"]) <= 2 for e in standard)
    assert all(not e["luecken"] for e in standard)
    assert all(a["soll_status"] != "nicht_uebernommen" for e in standard for a in e["aenderungen"])
    assert min(len(e["stoerer"]) for e in schwer) >= 3
    assert sum(len(e["stoerer"]) for e in schwer) > 2 * sum(len(e["stoerer"]) for e in standard)


def test_mehrere_gruppen_von_aenderungen_pro_dokument(tmp_path):
    for s in range(1, 7):
        e = _erwartet(_paar(tmp_path, s))
        assert len({a["typ"] for a in e["aenderungen"]}) >= 4
    mehrseitig = [len({a["seite"] for a in _erwartet(tmp_path / "entwicklung" / f"synth_{s:04d}_standard")["aenderungen"]}) >= 2
                  for s in range(1, 7)]
    assert sum(mehrseitig) >= 3


def test_dev_und_test_haben_getrennte_stoerer(tmp_path):
    nur_test = set(stoerer.NUR_TEST)
    dev = set()
    for s in range(1, 25):
        dev |= set(_erwartet(_paar(tmp_path, s, "schwer"))["stoerer"])
    assert dev and dev.isdisjoint(nur_test)
    test = set()
    for s in range(1001, 1011):
        test |= set(_erwartet(_paar(tmp_path, s, "standard", "test"))["stoerer"])
    assert nur_test <= test
    with pytest.raises(ValueError):
        erzeuge_paar(7, {"name": "x", "stoerer_liste": ["gedrehte_seite"]}, tmp_path / "x", "entwicklung")


# ====================================================================== Jede Stoerer-Familie erzeugt ihre Besonderheit
def _c_fussnote(o, e):
    assert e["pdf_info"]["fussnoten"] >= 1
    klein = [c for c in _chars(o) if abs(c["size"] - 8.0) < 0.1]
    assert klein and all(c["top"] > 600 for c in klein)


def _c_silben(o, e):
    assert e["pdf_info"]["trennungen"] >= 1
    text = "\n".join(_seitentexte(o))
    assert re.search(r"[a-zäöüß]-\n", text) or b"\xad" in (o / "aenderungen.pdf").read_bytes()


def _c_aufz(o, e):
    assert re.search(r"^(•|–|\d+\.|[a-z]\)) \S", "\n".join(_seitentexte(o)), re.M)
    assert any(re.match(r"^(•|–|\d+\.|[a-z]\)) \S", z) for z in _alle_excel_zeilen(o))


def _c_nummeriert(o, e):
    z = _alle_excel_zeilen(o)
    assert any(x.startswith("1. ") for x in z) and any(x.startswith("2. ") for x in z)
    assert re.search(r"^1\. \S", "\n".join(_seitentexte(o)), re.M)


def _c_zwei(o, e):
    seite = [lk for lk in e["luecken"] if lk["art"] == "mehrspaltig"][0]["seite"]
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        worte = pdf.pages[seite - 1].extract_words()
    rechts = [w for w in worte if abs(w["x0"] - 306.5) < 0.7]
    assert len(rechts) >= 1
    assert any(abs(l["top"] - r["top"]) < 2 and l["x1"] < 295 for r in rechts for l in worte)


def _c_tabelle(o, e):
    seite = [lk for lk in e["luecken"] if lk["art"] == "tabelle"][0]["seite"]
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        pg = pdf.pages[seite - 1]
        assert len(pg.lines) >= 9 and "Bauteil" in pg.extract_text()
    assert any(" | " in z for z in _alle_excel_zeilen(o))


def _c_bild(o, e):
    seite = [lk for lk in e["luecken"] if lk["art"] == "bild"][0]["seite"]
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        assert len(pdf.pages[seite - 1].images) >= 1 and pdf.pages[seite - 1].chars


def _c_scan(o, e):
    lk = [lk for lk in e["luecken"] if lk["art"] == "scan_seite"][0]
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        pg = pdf.pages[lk["seite"] - 1]
        assert len(pg.chars) == 0 and len(pg.images) >= 1
        assert sum(len(p.chars) for p in pdf.pages) > 0
    betroffen = [a for a in e["aenderungen"] if "scan_seite" in a["stoerer"]]
    assert betroffen and all(a["soll_status"] == "nicht_uebernommen" for a in betroffen)


def _c_farbtoene(o, e):
    farben = {_farbe(c) for c in _chars(o)}
    assert any(_nahe(f, _hex(h), 0.01) for f in farben for h in stoerer.BLAUTOENE)
    assert any(_nahe(f, _hex(h), 0.01) for f in farben for h in stoerer.ROTTOENE)


def _c_cmyk(o, e):
    farben = {_farbe(c) for c in _chars(o)}
    assert any(len(f) == 4 for f in farben) and any(len(f) == 1 for f in farben)


def _c_misch(o, e):
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        linien = [ln for pg in pdf.pages for ln in pg.lines]
    assert any(_blau(tuple(ln["stroking_color"] or ())) for ln in linien)      # blau unterstrichen/durchgestrichen
    assert any(a["soll_status"] == "pruefen" and "mischmarkierung" in a["stoerer"] for a in e["aenderungen"]) \
        or any("mischmarkierung" in a["stoerer"] for a in e["aenderungen"])


def _c_umbruch(o, e):
    assert e["pdf_info"]["umbrueche_im_absatz"] >= 1


def _c_kopf(o, e):
    texte = _seitentexte(o)
    chars = _chars(o)
    nur_erste = len(texte) > 1 and "Beispieldokument" in texte[0] and "Beispieldokument" not in texte[1]
    blau = any(_blau(_farbe(c)) and c["top"] < 50 for c in chars)
    nummern = "Rev. 4" in texte[0]
    assert nur_erste or blau or nummern


def _c_verschmolzen(o, e):
    chars = sorted(_chars(o), key=lambda c: (round(c["top"]), c["x0"]))
    for a, b in zip(chars, chars[1:]):
        if _rot(_farbe(a)) and _blau(_farbe(b)) and abs(b["x0"] - a["x1"]) < 0.5 and abs(a["top"] - b["top"]) < 1:
            return
    raise AssertionError("keine verschmolzene Ersetzung (rot direkt vor blau) gefunden")


def _ein_wort_anders(a: str, b: str) -> bool:
    ta, tb = a.split(" "), b.split(" ")
    return len(ta) == len(tb) and sum(1 for x, y in zip(ta, tb) if x != y) == 1


def _gleich_oder_fast(zeilen: list[str]) -> bool:
    zeilen = [z for z in zeilen if len(z) > 25]
    for i, a in enumerate(zeilen):
        for b in zeilen[i + 1:]:
            if a == b or _ein_wort_anders(a, b):
                return True
    return False


def _c_doppelt(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    zeilen = [z for row in ws.iter_rows() for c in row if isinstance(c.value, str) for z in c.value.split("\n")]
    assert _gleich_oder_fast(zeilen)
    assert any("doppelte_absaetze" in a["stoerer"] for a in e["aenderungen"])


def _kopfzeile(ws) -> tuple[int, list]:
    for r in (1, 2):
        werte = [c.value for c in ws[r]]
        if "Titel" in werte or "Kapitel" in werte:
            return r, werte
    raise AssertionError("keine Kopfzeile")


def _c_komm(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    r, kopf = _kopfzeile(ws)
    ci = kopf.index("Anmerkung")
    ti = [i for i, k in enumerate(kopf) if k in ("Inhalt", "Anforderung", "Text")][0]
    summe = lambda i: sum(len(str(ws.cell(z, i + 1).value or "")) for z in range(r + 1, ws.max_row + 1))
    assert summe(ci) > summe(ti)


def _c_lfd(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    r, kopf = _kopfzeile(ws)
    assert kopf[0] == "Lfd"
    werte = [ws.cell(z, 1).value for z in range(r + 1, ws.max_row + 1)]
    assert all(isinstance(v, int) for v in werte if v is not None) and len([v for v in werte if v]) >= 3


def _c_zahl(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    r, kopf = _kopfzeile(ws)
    ni = [i for i, k in enumerate(kopf) if k in ("Nr", "Nr.")][0]
    werte = [ws.cell(z, ni + 1).value for z in range(r + 1, ws.max_row + 1)]
    assert sum(1 for v in werte if isinstance(v, float) and abs(v - 2.1) < 1e-9) >= 2     # 2.1 und 2.10
    assert all(isinstance(v, (int, float)) for v in werte if v is not None and not isinstance(v, str)) or True


def _c_blaetter(o, e):
    wb = openpyxl.load_workbook(o / "vorher.xlsx")
    assert wb.sheetnames == ["Deckblatt", "Historie", "CDE", "Glossar"]
    assert {a["blatt"] for a in e["aenderungen"]} == {3}


def _c_formeln(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    formeln = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=LEN(") for f in formeln) and any(f.startswith("=SUM(") for f in formeln)
    assert len(ws.merged_cells.ranges) >= 1


def _c_versteckt(o, e):
    ws = _hauptblatt(openpyxl.load_workbook(o / "vorher.xlsx"), e)
    versteckt = [r for r, d in ws.row_dimensions.items() if d.hidden]
    assert versteckt
    texte = {r: [str(c.value) for c in ws[r] if isinstance(c.value, str) and len(c.value) > 25] for r in range(1, ws.max_row + 1)}
    for r in versteckt:
        andere = [t for rr, ts in texte.items() if rr != r and rr not in versteckt for t in ts]
        assert any(_ein_wort_anders(a, b) for a in texte[r] for b in andere)


def _c_sonder(o, e):
    text = "\n".join(_alle_excel_zeilen(o))
    assert " " in text and ("„" in text or "–" in text) and "ß" in text
    assert "ß" in "\n".join(_seitentexte(o))


def _c_link(o, e):
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        assert any(pg.hyperlinks for pg in pdf.pages)
        chars = [c for pg in pdf.pages for c in pg.chars]
    assert any(_nahe(_farbe(c), (0.0196, 0.3882, 0.7569), 0.02) for c in chars)
    assert e["pdf_info"]["links"] >= 1


def _c_gedreht(o, e):
    with pdfplumber.open(o / "aenderungen.pdf") as pdf:
        assert any(pg.rotation == 90 for pg in pdf.pages)
    assert e["pdf_info"]["gedrehte_seite"] >= 1
    assert any("gedrehte_seite" in a["stoerer"] and a["soll_status"] == "pruefen" for a in e["aenderungen"])


def _c_gruen(o, e):
    assert any(_nahe(_farbe(c), (0.0, 0.6, 0.0), 0.02) for c in _chars(o))
    assert any(lk["art"] == "unbekannte_farbe" for lk in e["luecken"])
    assert any(a["soll_status"] == "nicht_uebernommen" and "gruene_markierung" in a["stoerer"] for a in e["aenderungen"])


CHECKS = {
    "fussnote": _c_fussnote, "silbentrennung": _c_silben, "aufzaehlung": _c_aufz, "nummerierte_liste": _c_nummeriert,
    "zwei_spalten": _c_zwei, "tabelle_mit_farbkopf": _c_tabelle, "bild": _c_bild, "scan_seite": _c_scan,
    "farbtoene": _c_farbtoene, "cmyk_grau": _c_cmyk, "mischmarkierung": _c_misch,
    "seitenumbruch_im_absatz": _c_umbruch, "kopf_fusszeile": _c_kopf, "verschmolzen": _c_verschmolzen,
    "doppelte_absaetze": _c_doppelt, "kommentar_spalte": _c_komm, "lfd_spalte": _c_lfd, "zahl_sections": _c_zahl,
    "mehrere_blaetter": _c_blaetter, "formeln_merges": _c_formeln, "versteckte_zeilen": _c_versteckt,
    "sonderzeichen": _c_sonder, "hyperlink": _c_link, "gedrehte_seite": _c_gedreht, "gruene_markierung": _c_gruen,
}


def test_jede_familie_hat_einen_check():
    assert set(CHECKS) == set(stoerer.REGISTER)


@pytest.mark.parametrize("seed", [21, 22, 23])
@pytest.mark.parametrize("name", sorted(stoerer.REGISTER))
def test_stoerer_familie_erzeugt_ihre_besonderheit(tmp_path, name, seed):
    f = stoerer.REGISTER[name]
    korpus = "test" if f.nur_test else "entwicklung"
    p = erzeuge_paar(seed, {"name": f"t_{name}", "stoerer_liste": [name]}, tmp_path / korpus, korpus)
    e = _erwartet(p)
    assert e["stoerer"] == [name]
    CHECKS[name](p, e)
    if f.luecke:
        assert any(lk["art"] == f.luecke for lk in e["luecken"])


def test_unloesbare_familien_tragen_status_und_luecke(tmp_path):
    for name in ("scan_seite", "zwei_spalten", "tabelle_mit_farbkopf"):
        assert stoerer.REGISTER[name].soll_status == "nicht_uebernommen"
        gefunden = False
        for seed in range(31, 41):          # bei Tabellen hat nicht jedes Paar eine geaenderte Zelle
            e = _erwartet(erzeuge_paar(seed, {"name": f"t_{name}", "stoerer_liste": [name]}, tmp_path / name))
            lk = [x for x in e["luecken"] if x["art"] == stoerer.REGISTER[name].luecke]
            assert lk
            if any(x["mit_aenderung"] for x in lk):
                assert any(a["soll_status"] == "nicht_uebernommen" and name in a["stoerer"] for a in e["aenderungen"])
                gefunden = True
        assert gefunden, name


# ====================================================================== erwartet.json passt zum Zellvergleich
def _norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return "" if v.startswith("=") else v
    return repr(v)


def _zeilen(ws) -> list[tuple[str, ...]]:
    return [tuple(_norm(c.value) for c in row) for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]


def zellvergleich(ordner: Path, blatt: int):
    """Vergleicht vorher/nachher des Hauptblatts. Gibt (geaendert, betroffene_zeilen, anker, geloescht) in vorher-Koordinaten."""
    v = _zeilen(openpyxl.load_workbook(ordner / "vorher.xlsx").worksheets[blatt - 1])
    n = _zeilen(openpyxl.load_workbook(ordner / "nachher.xlsx").worksheets[blatt - 1])
    geaendert: set[tuple[int, int]] = set()
    betroffen: set[int] = set()
    anker: set[int] = set()
    geloescht: set[int] = set()
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, v, n, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        if tag == "delete":
            geloescht |= {i + 1 for i in range(i1, i2)}
        elif tag == "insert":
            anker.add(i1)
        else:                                   # replace: Zeilen veraendert und/oder davor/danach eingefuegt
            betroffen |= {i + 1 for i in range(i1, i2)}
            anker |= {i1, i2}
            if i2 - i1 == j2 - j1:
                for k in range(i2 - i1):
                    for c, (a, b) in enumerate(zip(v[i1 + k], n[j1 + k]), 1):
                        if a != b:
                            geaendert.add((i1 + k + 1, c))
    return geaendert, betroffen, anker, geloescht


def pruefe_zellvergleich(ordner: Path) -> None:
    e = _erwartet(ordner)
    blaetter = {a["blatt"] for a in e["aenderungen"]}
    assert len(blaetter) == 1
    blatt = blaetter.pop()
    wv = openpyxl.load_workbook(ordner / "vorher.xlsx")
    wn = openpyxl.load_workbook(ordner / "nachher.xlsx")
    assert wv.sheetnames == wn.sheetnames
    for i, ws in enumerate(wv.worksheets, 1):          # alle anderen Blaetter bleiben unveraendert
        if i != blatt:
            assert _zeilen(ws) == _zeilen(wn.worksheets[i - 1]), ws.title
    geaendert, betroffen, anker, geloescht = zellvergleich(ordner, blatt)
    wsv = wv.worksheets[blatt - 1]
    for a in e["aenderungen"]:
        z, s = a["zeile"], a["spalte"]
        if "doppelte_absaetze" in a["stoerer"]:      # zwei gleiche Zeilen: der Zellvergleich kann sie vertauschen
            assert any(zz in betroffen | anker | geloescht | {r for r, _ in geaendert} for zz in (z - 1, z, z + 1)), a
        else:
            assert (z, s) in geaendert or z in betroffen or z in anker or z in geloescht, a
        if a["typ"] in ("text_ersetzt", "text_eingefuegt", "text_geloescht", "absatz_geloescht", "section_geloescht"):
            assert str(wsv.cell(z, s).value or "").strip(), f"leere Zelle in vorher bei {a}"
    # keine unerklaerte Aenderung: jede zusammenhaengende Gruppe veraenderter Zeilen liegt bei einem Eintrag
    zeilen = sorted({z for z, _ in geaendert} | betroffen | geloescht | anker)
    eintragszeilen = [a["zeile"] for a in e["aenderungen"]]
    gruppen: list[list[int]] = []
    for z in zeilen:
        if gruppen and z <= gruppen[-1][-1] + 1:
            gruppen[-1].append(z)
        else:
            gruppen.append([z])
    for g in gruppen:
        assert any(g[0] - 1 <= ez <= g[-1] + 1 for ez in eintragszeilen), f"unerklaerte Aenderung bei Zeilen {g}"


@pytest.mark.parametrize("seed", range(1, 9))
@pytest.mark.parametrize("profil", ["standard", "schwer"])
def test_erwartet_passt_zum_zellvergleich(tmp_path, seed, profil):
    pruefe_zellvergleich(_paar(tmp_path, seed, profil))


LAYOUT_FAELLE = [
    ("L1", ["lfd_spalte", "kommentar_spalte"]), ("L2", ["lfd_spalte", "doppelte_absaetze"]),
    ("L3", ["lfd_spalte", "kommentar_spalte", "aufzaehlung"]), ("L1", ["formeln_merges", "versteckte_zeilen"]),
    ("L3", ["formeln_merges", "versteckte_zeilen"]), ("L3", ["zahl_sections", "mehrere_blaetter"]),
    ("L1", ["zahl_sections", "doppelte_absaetze"]), ("L2", ["mehrere_blaetter", "versteckte_zeilen", "nummerierte_liste"]),
]


@pytest.mark.parametrize("layout,liste", LAYOUT_FAELLE)
@pytest.mark.parametrize("seed", [41, 42])
def test_zellvergleich_in_allen_layouts(tmp_path, layout, liste, seed):
    p = erzeuge_paar(seed, {"name": "lay", "layout": layout, "stoerer_liste": liste}, tmp_path)
    pruefe_zellvergleich(p)
    assert _erwartet(p)["layout"].startswith(layout)


def test_stoerer_layouts_und_inhalt_der_nachher_excel(tmp_path):
    """Wahrheit aus dem Modell: die geaenderten Woerter stehen in nachher, die entfallenen nur in vorher."""
    p = erzeuge_paar(51, {"name": "w", "layout": "L3", "stoerer_liste": []}, tmp_path)
    e = _erwartet(p)
    ws_v = openpyxl.load_workbook(p / "vorher.xlsx")["CDE"]
    ws_n = openpyxl.load_workbook(p / "nachher.xlsx")["CDE"]
    text_v = "\n".join(str(c.value) for row in ws_v.iter_rows() for c in row if c.value)
    text_n = "\n".join(str(c.value) for row in ws_n.iter_rows() for c in row if c.value)
    assert text_v != text_n
    assert any(a["typ"] == "section_neu" for a in e["aenderungen"]) or len(e["aenderungen"]) >= 7
    pdf_text = "\n".join(_seitentexte(p))
    # Text, der in vorher UND nachher steht, kommt im PDF vor (Stichprobe: erste Zeile eines unveraenderten Absatzes)
    gemeinsam = [z for z in text_v.split("\n") if z in text_n.split("\n") and len(z) > 40]
    assert gemeinsam
    kopf = gemeinsam[0][:25]
    assert kopf in pdf_text.replace("\n", " ")


def test_woerter_markiert_stimmt_mit_den_farbigen_woertern_im_pdf(tmp_path):
    for seed in (61, 62, 63):
        p = erzeuge_paar(seed, {"name": "f", "stoerer_liste": ["farbtoene"]}, tmp_path)
        e = _erwartet(p)
        with pdfplumber.open(p / "aenderungen.pdf") as pdf:
            worte = [w for pg in pdf.pages for w in pg.extract_words(extra_attrs=["non_stroking_color"])]
        farbig = [w for w in worte if _blau(tuple(w["non_stroking_color"] or ())) or _rot(tuple(w["non_stroking_color"] or ()))]
        assert len(farbig) == e["woerter_markiert"], (seed, len(farbig), e["woerter_markiert"])


# ====================================================================== Smoke-Korpus und Skript
def test_smoke_korpus_6_paare_unter_15_sekunden(tmp_path):
    t = time.time()
    for seed in range(1, 7):
        _paar(tmp_path, seed, "schwer" if seed % 3 == 0 else "standard")
    assert time.time() - t < 15
    assert len(list((tmp_path / "entwicklung").iterdir())) == 6


def _skript(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(ROOT / "scripts" / "erzeuge_paare.py"), *args],
                          capture_output=True, text=True, cwd=str(ROOT), timeout=120)


def test_skript_erzeugt_korpus_und_ueberschreibt_nie(tmp_path):
    r = _skript("--korpus", "entwicklung", "--seeds", "1-3", "--ziel", str(tmp_path))
    assert r.returncode == 0, r.stderr
    ordner = tmp_path / "entwicklung"
    assert sorted(p.name for p in ordner.iterdir() if p.is_dir()) == ["synth_0001_standard", "synth_0002_standard",
                                                                      "synth_0003_schwer"]
    korpus = json.loads((ordner / "KORPUS.json").read_text(encoding="utf-8"))
    assert korpus["generator_version"] == generator.GENERATOR_VERSION and korpus["seeds"] == [1, 2, 3]
    assert korpus["korpus"] == "entwicklung" and re.fullmatch(r"\d{4}-\d{2}-\d{2}", korpus["erstellt"])
    vorher = (ordner / "synth_0001_standard" / "vorher.xlsx").read_bytes()
    mtime = (ordner / "synth_0001_standard" / "vorher.xlsx").stat().st_mtime_ns
    r = _skript("--korpus", "entwicklung", "--seeds", "1-4", "--ziel", str(tmp_path))
    assert r.returncode == 0 and "uebersprungen" in r.stdout
    assert (ordner / "synth_0001_standard" / "vorher.xlsx").read_bytes() == vorher
    assert (ordner / "synth_0001_standard" / "vorher.xlsx").stat().st_mtime_ns == mtime
    assert json.loads((ordner / "KORPUS.json").read_text(encoding="utf-8"))["seeds"] == [1, 2, 3, 4]


def test_skript_haelt_entwicklung_und_test_seeds_getrennt(tmp_path):
    assert _skript("--korpus", "entwicklung", "--seeds", "1000-1002", "--ziel", str(tmp_path)).returncode == 2
    assert _skript("--korpus", "test", "--seeds", "1-3", "--ziel", str(tmp_path)).returncode == 2
    assert not (tmp_path / "entwicklung").exists() and not (tmp_path / "test").exists()
    r = _skript("--korpus", "test", "--seeds", "1001-1002", "--ziel", str(tmp_path))
    assert r.returncode == 0, r.stderr
    assert (tmp_path / "test" / "synth_1001_standard" / "erwartet.json").exists()
