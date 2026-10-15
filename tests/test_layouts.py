"""Layout-Profile (src/layouts.py, ADR 0003): Pruefung, Hash, 'nie direkt live', Haken im Excel-Modell.

Nur erfundene Daten (Regel 16). Alles in tmp_path; conftest.py setzt CDE_DATA_ROOT, die Profil-Dateien
werden pro Test auf Dateien in tmp_path umgebogen (layouts.profil_dateien).
"""
import json

import pytest
import yaml
from openpyxl import Workbook

from src import config, layouts, runs
from src.excel_model import ExcelModell
from src.layouts import ProfilFehler
from tests.fixtures.pdf_writer import baue_pdf
from tests.fixtures.szenario import baue_excel, zustaende

# Markante erfundene Woerter: duerfen nie in einer Vorschau auftauchen
MARKANT = ["quokkablatt", "flamingokopf", "pinguinzelle", "teichtitel"]

ELEMENTE = [
    ("h", "1 Eins"),
    ("p", "Der Quokka frisst {-Gras-} {+Laub+} am Morgen. Dann ruht er lange im Schatten."),
    ("h", "2 Zwei"),
    ("p", "Ein Flamingo steht {-links-} {+rechts+} im Teich und putzt sein rosa Gefieder."),
]

PROFIL = {
    "name": "test_typ",
    "aktiv": True,
    "blaetter": [
        {"blatt": "Deckblatt", "ignorieren": True},
        {"blatt": "CDE", "section_spalte": "A", "titel_spalte": "B", "inhalt_spalte": "C",
         "modell": "zeile_pro_section", "kopfzeile": 1},
    ],
}


def schreibe_profile(tmp_path, monkeypatch, profile, lokal=None):
    """Legt die Vorlage (layouts.yaml) und optional die lokale Datei in tmp_path an und biegt layouts darauf um."""
    vorlage = tmp_path / "cfg" / "layouts.yaml"
    vorlage.parent.mkdir(exist_ok=True)
    vorlage.write_text(yaml.safe_dump({"profile": profile}, allow_unicode=True), encoding="utf-8")
    dateien = [vorlage, tmp_path / "cfg" / "layouts.local.yaml"]
    if lokal is not None:
        dateien[1].write_text(yaml.safe_dump({"profile": lokal}, allow_unicode=True), encoding="utf-8")
    monkeypatch.setattr(layouts, "profil_dateien", lambda: dateien)
    return dateien


def ersetze_datei(datei, profile):
    datei.write_text(yaml.safe_dump({"profile": profile}, allow_unicode=True), encoding="utf-8")


@pytest.fixture
def excel(tmp_path):
    p = tmp_path / "vorher.xlsx"
    baue_excel(p, zustaende(ELEMENTE, False), "L1")
    return p


@pytest.fixture
def lauf_dateien(tmp_path, excel):
    pdf = tmp_path / "a.pdf"
    baue_pdf(pdf, ELEMENTE)
    return pdf, excel


def kopie(**aenderungen):
    d = json.loads(json.dumps(PROFIL))
    d.update(aenderungen)
    return d


# ---------------------------------------------------------------- Pruefung der YAML
def fehler_von(d):
    with pytest.raises(ProfilFehler) as e:
        layouts.profil_aus_dict(d, "test.yaml")
    return e.value.meldung


def test_gueltiges_profil_wird_gelesen():
    p = layouts.profil_aus_dict(PROFIL, "test.yaml")
    assert p.name == "test_typ" and p.aktiv
    assert p.blaetter["CDE"].inhalt_spalte == 3 and p.blaetter["CDE"].kopfzeile == 1
    assert p.blaetter["Deckblatt"].ignorieren and p.hash


def test_unbekannter_schluessel_wird_benannt():
    d = kopie()
    d["blaetter"][1]["sektion_spalte"] = "A"
    m = fehler_von(d)
    assert "sektion_spalte" in m and "section_spalte" in m          # Tippfehler-Hinweis
    assert "'farbe' ist unbekannt" in fehler_von(kopie(farbe="rot"))


@pytest.mark.parametrize("wert", ["1", "AAAA", "", 3, True, "A1"])
def test_ungueltige_spalte(wert):
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = wert
    assert "inhalt_spalte" in fehler_von(d)


def test_spalte_klein_geschrieben_wird_gross():
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = "c"
    assert layouts.profil_aus_dict(d).blaetter["CDE"].inhalt_spalte == 3


def test_pflichtangaben_und_widersprueche():
    d = kopie()
    del d["blaetter"][1]["inhalt_spalte"]
    assert "inhalt_spalte" in fehler_von(d)
    d = kopie()
    d["blaetter"][1]["modell"] = "irgendwas"
    assert "modell" in fehler_von(d)
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = "A"
    assert "verschiedene Spalten" in fehler_von(d)
    d = kopie()
    d["blaetter"][0]["inhalt_spalte"] = "C"
    assert "ignoriert" in fehler_von(d)
    d = kopie()
    del d["blaetter"][1]["section_spalte"]
    assert "section_spalte" in fehler_von(d)                       # zeile_pro_section ohne Section-Spalte
    d = kopie()
    d["blaetter"][1]["kopfzeile"] = -3
    assert "kopfzeile" in fehler_von(d)


def test_doppeltes_blatt_im_profil():
    d = kopie()
    d["blaetter"].append({"blatt": "CDE", "inhalt_spalte": "D", "modell": "zeile_pro_absatz"})
    assert "doppelt" in fehler_von(d)


def test_alle_probleme_auf_einmal():
    d = kopie(farbe=1)
    d["blaetter"][1]["inhalt_spalte"] = 5
    m = fehler_von(d)
    assert "farbe" in m and "inhalt_spalte" in m


def test_doppelte_profilnamen_werden_abgelehnt(tmp_path, monkeypatch):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL], lokal=[kopie()])     # gleicher Name in beiden Dateien
    profile, fehler = layouts.lade_alle()
    assert "test_typ" not in profile                                       # keines der beiden wird benutzt
    assert any("mehrfach" in m for _, m in fehler)
    with pytest.raises(ProfilFehler):
        layouts.bestaetigtes_profil("test_typ")


def test_doppelter_name_in_einer_datei(tmp_path, monkeypatch):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL, kopie()])
    profile, fehler = layouts.lade_alle()
    assert not profile and any("mehrfach" in m for _, m in fehler)


def test_kaputte_yaml_datei_wird_gemeldet(tmp_path, monkeypatch):
    dateien = schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    dateien[1].write_text("profile: [\n  - name: x\n   kaputt", encoding="utf-8")
    profile, fehler = layouts.lade_alle()
    assert "test_typ" in profile                                           # die gute Datei bleibt benutzbar
    assert any("layouts.local.yaml" in m for _, m in fehler)


def test_vorlage_im_repo_ist_gueltig_und_inaktiv():
    roh = yaml.safe_load((config.CONFIG_DIR / "layouts.yaml").read_text(encoding="utf-8"))
    ps = [layouts.profil_aus_dict(e, "layouts.yaml") for e in roh["profile"]]
    assert ps and all(not p.aktiv for p in ps)


# ---------------------------------------------------------------- Hash
def test_hash_aendert_sich_mit_dem_aufbau_aber_nicht_mit_aktiv_oder_beschreibung():
    a = layouts.profil_aus_dict(PROFIL)
    b = layouts.profil_aus_dict(kopie(aktiv=False, beschreibung="anderer Kommentar"))
    assert a.hash == b.hash
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = "D"
    assert layouts.profil_aus_dict(d).hash != a.hash
    d = kopie()
    d["blaetter"].reverse()                                                # Reihenfolge der Blaetter ist egal
    assert layouts.profil_aus_dict(d).hash == a.hash
    assert layouts.profil_aus_dict(kopie(name="anderer")).hash != a.hash


# ---------------------------------------------------------------- Haken im Excel-Modell
def test_ohne_profil_alles_wie_vorher(excel):
    xl = ExcelModell(excel)
    assert xl.profil is None
    for b in xl.beschreibung():
        assert "profil" not in b and "ignoriert" not in b                  # Beschreibung unveraendert
    lay = xl.layouts["CDE"]
    assert (lay.section_spalte, lay.titel_spalte, lay.inhalt_spalte, lay.modell) == (1, 2, 3, "zeile_pro_section")
    assert lay.profil == "" and not lay.ignoriert


def test_profil_legt_aufbau_fest_und_ignoriert_blaetter(excel):
    p = layouts.profil_aus_dict(PROFIL)
    xl = ExcelModell(excel, p)
    lay = xl.layouts["CDE"]
    assert (lay.section_spalte, lay.titel_spalte, lay.inhalt_spalte, lay.modell) == (1, 2, 3, "zeile_pro_section")
    assert lay.sicherheit == "hoch" and lay.sicherheit_gruende == []
    beschr = {b["blatt"]: b for b in xl.beschreibung()}
    assert beschr["CDE"]["profil"] == "Profil: test_typ"
    assert beschr["Deckblatt"]["ignoriert"] is True
    assert not [k for k in xl.kandidaten if k.blatt == "Deckblatt"]        # ignoriertes Blatt: keine Kandidaten
    assert [k for k in xl.kandidaten if k.blatt == "CDE"]
    assert any("test_typ" in w for w in xl.warnungen)


def test_profil_ersetzt_unsichere_erkennung(tmp_path):
    """Zwei Textspalten: die Heuristik ist unsicher ('mittel'); das Profil legt Spalte D fest -> 'hoch'."""
    pfad = tmp_path / "zwei_spalten.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Liste"
    ws.append(["Nr", "Titel", "Text eins", "Text zwei"])
    lang = "Dieser Satz ist lang genug, damit er als Inhaltstext zaehlt und ruhig bleibt."
    for i in range(1, 5):
        ws.append([f"{i}", f"Abschnitt {i}", lang + f" Eins {i}", lang + f" Zwei {i}"])
    wb.save(str(pfad))
    auto = ExcelModell(pfad)
    assert auto.layouts["Liste"].sicherheit != "hoch"
    d = {"name": "zwei", "aktiv": True, "blaetter": [
        {"blatt": "Liste", "section_spalte": "A", "titel_spalte": "B", "inhalt_spalte": "D",
         "modell": "zeile_pro_section", "kopfzeile": 1}]}
    xl = ExcelModell(pfad, layouts.profil_aus_dict(d))
    assert xl.layouts["Liste"].inhalt_spalte == 4 and xl.layouts["Liste"].sicherheit == "hoch"
    assert xl is not auto                                                  # Cache-Schluessel beruecksichtigt das Profil
    assert ExcelModell(pfad).layouts["Liste"].sicherheit != "hoch"          # und ohne Profil bleibt es wie es war


def test_cache_trennt_profile_mit_verschiedenem_inhalt(excel):
    p1 = layouts.profil_aus_dict(PROFIL)
    d = kopie()
    d["blaetter"][1]["modell"] = "zeile_pro_absatz"
    p2 = layouts.profil_aus_dict(d)
    a, b = ExcelModell(excel, p1), ExcelModell(excel, p2)
    assert a is not b and a.layouts["CDE"].modell != b.layouts["CDE"].modell
    assert ExcelModell(excel, p1) is a


def test_spalte_ausserhalb_des_blatts(excel):
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = "Z"
    with pytest.raises(ProfilFehler) as e:
        ExcelModell(excel, layouts.profil_aus_dict(d))
    assert "Spalte Z" in e.value.meldung and "Blatt 2" in e.value.meldung
    assert "CDE" not in e.value.meldung                                    # keine Blattnamen in Meldungen des Modells


def test_blatt_im_profil_fehlt_in_der_excel(excel):
    d = kopie()
    d["blaetter"].append({"blatt": "Gibt es nicht", "ignorieren": True})
    xl = ExcelModell(excel, layouts.profil_aus_dict(d))
    assert any("nicht gibt" in w for w in xl.warnungen)


# ---------------------------------------------------------------- Nie direkt live
def vorschau_und_bestaetigung(excel, name="test_typ"):
    p = layouts.finde(name)
    vid = layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    layouts.bestaetige(vid)
    return vid


def test_unbestaetigtes_profil_wird_abgelehnt(tmp_path, monkeypatch, lauf_dateien):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    run_id = runs.neuer_lauf(pdf, excel)
    with pytest.raises(runs.LaufFehler) as e:
        runs.analysiere(run_id, "test_typ")
    assert "noch nicht bestätigt" in e.value.meldung
    assert not (runs.lauf_dir(run_id) / "plan.json").exists()             # nichts wurde analysiert/geschrieben


def test_vorschau_allein_reicht_nicht(tmp_path, monkeypatch, lauf_dateien):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    p = layouts.finde("test_typ")
    layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    with pytest.raises(runs.LaufFehler):
        runs.analysiere(runs.neuer_lauf(pdf, excel), "test_typ")


def test_bestaetigtes_profil_wird_benutzt_und_in_plan_json_vermerkt(tmp_path, monkeypatch, lauf_dateien):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    vorschau_und_bestaetigung(excel)
    run_id = runs.neuer_lauf(pdf, excel)
    daten = runs.analysiere(run_id, "test_typ")
    p = layouts.finde("test_typ")
    assert daten["profil"] == {"name": "test_typ", "hash": p.hash}
    assert json.loads((runs.lauf_dir(run_id) / "plan.json").read_text(encoding="utf-8"))["profil"]["hash"] == p.hash
    assert any(b.get("profil") == "Profil: test_typ" for b in daten["layout"])
    res = runs.uebernehmen_lauf(run_id)                                    # Eintragen mit demselben Aufbau
    assert res["bilanz"]["uebernommen"] >= 1


def test_ohne_profil_steht_null_in_plan_json(lauf_dateien):
    pdf, excel = lauf_dateien
    daten = runs.analysiere(runs.neuer_lauf(pdf, excel))
    assert daten["profil"] is None
    assert not any("profil" in b for b in daten["layout"])


def test_geaendertes_profil_nach_bestaetigung_wird_abgelehnt(tmp_path, monkeypatch, lauf_dateien):
    dateien = schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    vorschau_und_bestaetigung(excel)
    geaendert = kopie()
    geaendert["blaetter"][1]["modell"] = "zeile_pro_absatz"                 # anderer Inhalt -> anderer Hash
    ersetze_datei(dateien[0], [geaendert])
    with pytest.raises(runs.LaufFehler) as e:
        runs.analysiere(runs.neuer_lauf(pdf, excel), "test_typ")
    assert "geändert" in e.value.meldung


def test_kommentar_aendern_ist_kein_neuer_inhalt_aber_inaktiv_sperrt(tmp_path, monkeypatch, lauf_dateien):
    dateien = schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    vorschau_und_bestaetigung(excel)
    ersetze_datei(dateien[0], [kopie(beschreibung="nur ein Kommentar")])
    runs.analysiere(runs.neuer_lauf(pdf, excel), "test_typ")               # Kommentar aendert den Hash nicht
    ersetze_datei(dateien[0], [kopie(aktiv=False)])
    with pytest.raises(runs.LaufFehler) as e:
        runs.analysiere(runs.neuer_lauf(pdf, excel), "test_typ")
    assert "nicht aktiv" in e.value.meldung


def test_unbekanntes_profil(tmp_path, monkeypatch, lauf_dateien):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    with pytest.raises(runs.LaufFehler) as e:
        runs.analysiere(runs.neuer_lauf(pdf, excel), "gibt_es_nicht")
    assert "kein Profil" in e.value.meldung


def test_uebernehmen_nutzt_den_bestaetigten_aufbau_auch_wenn_die_yaml_sich_aendert(tmp_path, monkeypatch, lauf_dateien):
    dateien = schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    pdf, excel = lauf_dateien
    vorschau_und_bestaetigung(excel)
    run_id = runs.neuer_lauf(pdf, excel)
    runs.analysiere(run_id, "test_typ")
    ersetze_datei(dateien[0], [kopie(aktiv=False)])
    assert runs.uebernehmen_lauf(run_id)["bilanz"]["uebernommen"] >= 1     # Bestaetigung des Laufs gilt weiter


def test_bestaetigung_nur_zum_gezeigten_inhalt(tmp_path, monkeypatch, excel):
    dateien = schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    p = layouts.finde("test_typ")
    vid = layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    geaendert = kopie()
    geaendert["blaetter"][1]["kopfzeile"] = 2
    ersetze_datei(dateien[0], [geaendert])
    with pytest.raises(ProfilFehler) as e:
        layouts.bestaetige(vid)
    assert "seit der Vorschau geändert" in e.value.meldung
    assert layouts.status(layouts.finde("test_typ")) == "unbestaetigt"


def test_inaktives_profil_kann_nicht_bestaetigt_werden(tmp_path, monkeypatch, excel):
    schreibe_profile(tmp_path, monkeypatch, [kopie(aktiv=False)])
    p = layouts.finde("test_typ")
    vid = layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    with pytest.raises(ProfilFehler) as e:
        layouts.bestaetige(vid)
    assert "nicht aktiv" in e.value.meldung


@pytest.mark.parametrize("vid", ["../../x", "abc", "", "20261015-143000-zzzzzz"])
def test_ungueltige_vorschau_id(tmp_path, monkeypatch, vid):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    with pytest.raises(ProfilFehler):
        layouts.bestaetige(vid)


def test_unbekannte_vorschau_id(tmp_path, monkeypatch):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    with pytest.raises(ProfilFehler) as e:
        layouts.bestaetige("20261015-143000-abcdef")
    assert "nicht gefunden" in e.value.meldung


def test_append_only_nichts_wird_ueberschrieben(tmp_path, monkeypatch, excel):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    p = layouts.finde("test_typ")
    v1 = layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    v2 = layouts.schreibe_vorschau(p, layouts.erstelle_vorschau(excel, p))
    assert v1 != v2
    _, neu = layouts.bestaetige(v1)
    _, nochmal = layouts.bestaetige(v2)
    assert neu and not nochmal                                             # zweites Bestaetigen aendert nichts
    namen = sorted(f.name for f in layouts.profile_dir().iterdir())
    assert len([n for n in namen if n.startswith("vorschau_")]) == 2
    assert len([n for n in namen if n.startswith("bestaetigt_")]) == 1
    vorher = (layouts.profile_dir() / f"vorschau_{v1}.json").read_bytes()
    layouts.bestaetige(v1)
    assert (layouts.profile_dir() / f"vorschau_{v1}.json").read_bytes() == vorher
    with pytest.raises(FileExistsError):
        layouts._schreibe_neu(layouts.profile_dir() / f"vorschau_{v1}.json", {})


def test_profil_ordner_liegt_unter_dem_daten_ordner(tmp_path):
    assert layouts.profile_dir() == tmp_path / "data" / "profile"          # CDE_DATA_ROOT (conftest) gilt auch hier


def test_veraenderte_bestaetigungsdatei_zaehlt_nicht(tmp_path, monkeypatch, excel):
    schreibe_profile(tmp_path, monkeypatch, [PROFIL])
    vorschau_und_bestaetigung(excel)
    datei = next(layouts.profile_dir().glob("bestaetigt_*.json"))
    d = json.loads(datei.read_text(encoding="utf-8"))
    d["profil"]["kanon"]["blaetter"][1]["inhalt_spalte"] = "D"             # Inhalt passt nicht mehr zum Hash
    datei.write_text(json.dumps(d), encoding="utf-8")
    assert layouts.status(layouts.finde("test_typ")) == "unbestaetigt"
    with pytest.raises(ProfilFehler):
        layouts.bestaetigtes_profil("test_typ")


# ---------------------------------------------------------------- Vorschau: keine Inhalte
def baue_mit_markern(pfad):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quokkablatt"
    ws.append(["Flamingokopf", "Teichtitel", "Pinguinzelle"])
    lang = "Dieser Satz ist lang genug, damit er als Inhaltstext zaehlt, und enthaelt Quokka und Flamingo."
    for i in range(1, 4):
        ws.append([f"{i}", f"Abschnitt {i} Quokka", lang])
    wb.create_sheet("Deckblatt").append(["Pinguinzelle verborgen", "x"])
    wb.save(str(pfad))


def test_vorschau_enthaelt_keine_zellinhalte_und_keine_namen(tmp_path, monkeypatch):
    pfad = tmp_path / "marker.xlsx"
    baue_mit_markern(pfad)
    d = {"name": "marker_typ", "aktiv": True, "blaetter": [
        {"blatt": "Quokkablatt", "section_spalte": "A", "titel_spalte": "B", "inhalt_spalte": "C",
         "modell": "zeile_pro_section", "kopfzeile": 1},
        {"blatt": "Deckblatt", "ignorieren": True}]}
    schreibe_profile(tmp_path, monkeypatch, [d])
    p = layouts.finde("marker_typ")
    daten = layouts.erstelle_vorschau(pfad, p)
    text = "\n".join(layouts.formatiere_vorschau(daten, p)) + json.dumps(daten, ensure_ascii=False)
    vid = layouts.schreibe_vorschau(p, daten)
    gespeichert = (layouts.profile_dir() / f"vorschau_{vid}.json").read_text(encoding="utf-8")
    ohne_profil = layouts.erstelle_vorschau(pfad, None)
    text_auto = "\n".join(layouts.formatiere_vorschau(ohne_profil, None)) + json.dumps(ohne_profil, ensure_ascii=False)
    for roh in (text, text_auto):
        low = roh.casefold()
        for wort in MARKANT + ["quokka", "flamingo", "pinguin", "abschnitt", "teich"]:
            assert wort not in low, wort
    # gespeichert werden die Profil-Angaben des Nutzers (Blattnamen aus SEINER yaml), aber keine Zellinhalte
    low = gespeichert.casefold()
    for wort in ("abschnitt", "flamingokopf", "pinguin", "teichtitel", "enthaelt"):
        assert wort not in low, wort
    assert "Blatt 1" in text and "Section-Spalte A" in text and "Kandidaten" in text


def test_vorschau_zeigt_abweichung_zur_automatik(tmp_path, monkeypatch, excel):
    d = kopie()
    d["blaetter"][1]["inhalt_spalte"] = "D"                                # anders als die Erkennung (C)
    schreibe_profile(tmp_path, monkeypatch, [d])
    p = layouts.finde("test_typ")
    text = "\n".join(layouts.formatiere_vorschau(layouts.erstelle_vorschau(excel, p), p))
    assert "Die automatische Erkennung hätte gewählt" in text and "Inhalt C" in text


def test_erkennung_schlaegt_nur_vor(tmp_path, monkeypatch, excel):
    d = kopie(erkennung={"blattnamen": ["cde"], "kopfzeilen_begriffe": ["Titel", "inhalt"]})
    d2 = kopie(name="anderes", erkennung={"blattnamen": ["gibt es nicht"]})
    d3 = kopie(name="ohne_muster")
    schreibe_profile(tmp_path, monkeypatch, [d, d2, d3])
    daten = layouts.erstelle_vorschau(excel, None)
    assert daten["vorschlaege"] == ["test_typ"]
    assert layouts.status(layouts.finde("test_typ")) == "unbestaetigt"      # Vorschlag wendet nichts an
