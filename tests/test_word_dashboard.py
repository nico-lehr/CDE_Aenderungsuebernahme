"""Word-Datei durch das Dashboard (Upload -> Analyse -> Uebernehmen) mit dem ECHTEN Word-Leser, ohne Ersatz-Leser.

Nur erfundene Daten (Regel 16): Dokument und Register-Excel kommen aus den Bauhelfern der Tests."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import runs
from src.api.app import app
from tests._docx_bau import register_docx
from tests.test_register import GRUND, excel_absaetze, register_mappe


def _paar(tmp_path: Path) -> tuple[Path, Path]:
    absaetze = list(GRUND)
    absaetze[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    absaetze[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    docx = register_docx(tmp_path / "Aenderungen.docx", absaetze)
    xlsx = register_mappe(tmp_path / "cde.xlsx", excel_absaetze(absaetze))
    return docx, xlsx


@pytest.mark.markierung_an
def test_word_datei_ueber_api_hoch_und_uebernehmen(tmp_path):
    docx, xlsx = _paar(tmp_path)
    client = app.test_client()
    with open(docx, "rb") as d, open(xlsx, "rb") as x:
        r = client.post("/api/lauf", data={"datei": (d, docx.name), "excel": (x, "cde.xlsx")},
                        content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    plan = r.get_json()
    assert plan["quelle_art"] == "docx" and plan["legende_modus"] == "docx" and plan["seiten"] == 0
    assert plan["aenderungen"] and all(a["seite"] == 0 for a in plan["aenderungen"])
    assert plan["bilanz"]["stimmt"] and plan["bilanz"]["uebernommen"] == len(plan["aenderungen"])
    assert plan["stichprobe"]                                        # Vertrauensmodus: Stichprobe ist da
    res = runs.uebernehmen_lauf(plan["run_id"], [])
    assert res["bilanz"]["stimmt"]
    ergebnis = load_workbook(str(runs.lauf_dir(plan["run_id"]) / res["ordner"] / next(
        p.name for p in (runs.lauf_dir(plan["run_id"]) / res["ordner"]).glob("*.xlsx") if "protokoll" not in p.name)))
    assert ergebnis.sheetnames == load_workbook(str(xlsx)).sheetnames
    texte = [str(c.value) for zeile in ergebnis["Abgleich"].iter_rows() for c in zeile if c.value]
    assert any("hebt" in t for t in texte) and any("2041" in t for t in texte)     # die Aenderungen aus Word stehen drin
