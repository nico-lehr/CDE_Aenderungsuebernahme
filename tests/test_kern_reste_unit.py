"""Unit-Tests K-REST: kleine Restarbeiten in Excel-Modell, Matcher, Planer und PDF-Leser (nur erfundene Daten, Regel 16).

Die Red-Team-Tests (tests/redteam/test_korpus_analyse.py) pruefen die ganzen Faelle; hier stehen die Bausteine und
die Gegenproben ("was darf NICHT blockieren").
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from src import config
from src.excel_model import ExcelModell
from src.matcher import Matcher
from src.models import MITTEL, HOCH
from src.textutil import norm


def _tokens(text: str) -> list[tuple[int, str]]:
    return [(i, norm(w)) for i, w in enumerate(text.split()) if norm(w)]


def _excel(tmp_path, zeilen, name="a.xlsx", blatt="CDE", nachher=None):
    wb = Workbook()
    ws = wb.active
    ws.title = blatt
    for z in zeilen:
        ws.append(z)
    if nachher:
        nachher(ws)
    p = tmp_path / name
    wb.save(str(p))
    return p


_AUFBAU = [["Nr.", "Titel", "Inhalt"],
           ["1", "Allgemeines", "Der Betrieb erfolgt jederzeit durch geschultes Personal vor Ort."],
           ["2", "Stoermeldungen", "Die Reaktionszeit darf 150 ms nicht ueberschreiten."],
           ["3", "Normen", "Die Norm gilt fuer alle Anlagen im Betrieb."]]


# ===========================================================================================================
# K-XL 1: Nummern- und Titelspalte sind kein Ziel fuer Fliesstext (RT-120/121)
# ===========================================================================================================
def test_kandidaten_in_nummern_und_titelspalte_sind_markiert(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    lay = xl.layouts["CDE"]
    assert (lay.section_spalte, lay.titel_spalte, lay.inhalt_spalte) == (1, 2, 3)
    for k in xl.kandidaten:
        assert k.kopf_spalte == (k.spalte in (1, 2)), (k.art, k.spalte)
    assert any(k.kopf_spalte for k in xl.kandidaten) and any(not k.kopf_spalte for k in xl.kandidaten)


def test_listennummer_trifft_nummernzelle_nur_als_vorschlag(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    m = Matcher(xl, config.lade("zuordnung"))
    t = m.finde(_tokens("2"), "2", ist_kopf=False)
    assert t is not None and t.kandidat.spalte == 1 and t.kandidat.kopf_spalte
    assert t.stufe == MITTEL
    assert any("Nummern- oder Titelspalte" in g for g in t.gruende)


def test_ueberschrift_darf_weiter_automatisch_zugeordnet_werden(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    m = Matcher(xl, config.lade("zuordnung"))
    t = m.finde(_tokens("2 Stoermeldungen"), "2", ist_kopf=True)
    assert t is not None and t.stufe == HOCH
    assert not any("Nummern- oder Titelspalte" in g for g in t.gruende)


def test_fliesstext_im_inhalt_bleibt_automatisch(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    m = Matcher(xl, config.lade("zuordnung"))
    t = m.finde(_tokens("Die Reaktionszeit darf 150 ms nicht ueberschreiten."), "2", ist_kopf=False)
    assert t is not None and t.stufe == HOCH and not t.kandidat.kopf_spalte


# ===========================================================================================================
# K-XL 2: einfuege_risiken entschaerfen -- einfuege_hindernisse (nur, was beim Einfuegen wirklich brechen kann)
# ===========================================================================================================
def _eigenschaften(tmp_path):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws.title = "Sauber"
    for z in _AUFBAU:
        ws.append(z)
    ws = wb.create_sheet("Merge")
    ws.append(["a", "b"])
    ws.merge_cells("A3:B5")
    ws.merge_cells("A8:A9")
    ws = wb.create_sheet("Formel")
    ws["B2"] = "=SUM(A1:A9)"
    ws["A20"] = 1
    ws = wb.create_sheet("Bedingt")
    ws["A1"] = 1
    ws.conditional_formatting.add("A1:A20", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", bgColor="FF0000")))
    ws = wb.create_sheet("Pruefung")
    dv = DataValidation(type="list", formula1='"a,b"')
    ws.add_data_validation(dv)
    dv.add("B1:B20")
    ws = wb.create_sheet("Filter")
    ws.append(["x", "y"])
    ws.append([1, 2])
    ws.auto_filter.ref = "A1:B20"
    ws = wb.create_sheet("Tab")
    ws.append(["Name", "Wert"])
    for i in range(8):
        ws.append([f"n{i}", i])
    ws.add_table(Table(displayName="Werte", ref="A1:B9"))
    ws = wb.create_sheet("Namen")
    ws["A1"] = 1
    wb.defined_names["Bereich"] = DefinedName("Bereich", attr_text="Namen!$A$1:$A$20")
    ws = wb.create_sheet("Links")
    ws["A10"] = "Link"
    ws["A10"].hyperlink = "https://example.org/"
    ws = wb.create_sheet("Gruppe")
    for i in range(1, 13):
        ws.cell(i, 1, i)
    ws.row_dimensions[8].hidden = True
    ws.row_dimensions[9].outlineLevel = 1
    ws.row_dimensions[10].hidden = True
    ws = wb.create_sheet("Bild")
    ws["A1"], ws["A2"] = "x", 1
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=1, min_row=1, max_row=2), titles_from_data=True)
    ws.add_chart(ch, "E10")
    p = tmp_path / "eig.xlsx"
    wb.save(str(p))
    return ExcelModell(p)


def test_hindernisse_verbundene_zellen_unterhalb_sind_unbedenklich(tmp_path):
    xl = _eigenschaften(tmp_path)
    assert xl.einfuege_hindernisse("Merge", 1) == []              # beide Verbindungen liegen darunter
    assert xl.einfuege_hindernisse("Merge", 3) == []              # A3:B5 beginnt genau an der Einfuegestelle: wandert mit
    assert xl.einfuege_hindernisse("Merge", 6) == []              # A8:A9 liegt darunter, A3:B5 darueber
    assert xl.einfuege_hindernisse("Merge", 10) == []
    # streng (Schnittstelle von frueher) bleibt unveraendert
    assert any("Unterhalb sind verbundene" in r for r in xl.einfuege_risiken("Merge", 1))


@pytest.mark.schreiber_openpyxl          # ADR 0007: prueft das fruehere Speichern mit openpyxl
def test_hindernisse_verbundene_zellen_durch_die_einfuegestelle_bleiben_ein_hindernis(tmp_path):
    xl = _eigenschaften(tmp_path)
    for z in (4, 5, 9):                                           # mitten durch A3:B5 bzw. A8:A9
        r = xl.einfuege_hindernisse("Merge", z)
        assert any("durch die Zeile" in x for x in r), (z, r)
    r = xl.einfuege_hindernisse("Merge", 9)                       # nichts liegt darunter: nur das eine Hindernis
    assert len(r) == 1
    r = xl.einfuege_hindernisse("Merge", 4)                       # A8:A9 liegt darunter: als Hinweis dabei, wenn es ein Hindernis gibt
    assert len(r) == 2 and "Unterhalb sind verbundene" in r[1]


def test_hindernisse_ausgeblendete_zeilen_nur_wenn_mitten_in_der_folge(tmp_path):
    xl = _eigenschaften(tmp_path)                                 # Zeilen 8 (hidden), 9 (Gruppe), 10 (hidden)
    assert xl.einfuege_hindernisse("Gruppe", 3) == []             # alles darunter
    assert xl.einfuege_hindernisse("Gruppe", 8) == []             # Folge beginnt genau hier: wandert mit
    assert xl.einfuege_hindernisse("Gruppe", 11) == []            # direkt hinter der Folge
    assert xl.einfuege_hindernisse("Gruppe", 20) == []
    for z in (9, 10):                                             # Zeile davor UND Zeile dahinter sind ausgeblendet/gruppiert
        r = xl.einfuege_hindernisse("Gruppe", z)
        assert len(r) == 1 and "mitten in eine Folge" in r[0], (z, r)
    assert any("gruppierte Zeilen" in r for r in xl.einfuege_risiken("Gruppe", 3))      # streng: wie frueher


@pytest.mark.schreiber_openpyxl          # ADR 0007: prueft das fruehere Speichern mit openpyxl
def test_hindernisse_alles_andere_bleibt_wie_in_einfuege_risiken(tmp_path):
    xl = _eigenschaften(tmp_path)
    for blatt, z in [("Formel", 1), ("Formel", 5), ("Formel", 30), ("Bedingt", 5), ("Bedingt", 21), ("Pruefung", 5),
                     ("Pruefung", 21), ("Filter", 5), ("Filter", 21), ("Tab", 5), ("Tab", 10), ("Namen", 5), ("Namen", 21),
                     ("Links", 5), ("Links", 11), ("Bild", 5), ("Bild", 11), ("Sauber", 2), ("Sauber", 99)]:
        assert xl.einfuege_hindernisse(blatt, z) == xl.einfuege_risiken(blatt, z), (blatt, z)
    assert any("Formeln" in r for r in xl.einfuege_hindernisse("Formel", 5))
    assert any("bedingte Formatierung" in r for r in xl.einfuege_hindernisse("Bedingt", 5))
    assert any("Datenüberprüfung" in r for r in xl.einfuege_hindernisse("Pruefung", 5))
    assert any("Autofilter" in r for r in xl.einfuege_hindernisse("Filter", 5))
    assert any("Excel-Tabelle 'Werte'" in r for r in xl.einfuege_hindernisse("Tab", 5))
    assert any("benannte Bereich 'Bereich'" in r for r in xl.einfuege_hindernisse("Namen", 5))
    assert any("Hyperlinks" in r for r in xl.einfuege_hindernisse("Links", 5))
    assert any("Bilder oder Diagramme" in r for r in xl.einfuege_hindernisse("Bild", 5))
    assert "gibt es in der Excel nicht" in xl.einfuege_hindernisse("Gibt es nicht", 1)[0]


def test_hindernisse_patch_schreiber_zieht_das_meiste_selbst_mit(tmp_path):
    """ADR 0007 (Standard-Schreiber patch): Formeln, bedingte Formatierung, Datenpruefung, Filter, Namen, Hyperlinks,
    verbundene Zellen und Bilder darunter wandern mit. Hindernis bleibt, was xlsx_patch nicht sicher kann."""
    xl = _eigenschaften(tmp_path)
    for blatt in ("Formel", "Bedingt", "Pruefung", "Filter", "Namen", "Links", "Bild", "Sauber"):
        assert xl.einfuege_hindernisse(blatt, 5) == [], (blatt, xl.einfuege_hindernisse(blatt, 5))
    assert xl.einfuege_hindernisse("Merge", 6) == [] and xl.einfuege_hindernisse("Gruppe", 8) == []
    assert any("durch die Zeile" in r for r in xl.einfuege_hindernisse("Merge", 4))
    assert any("mitten in eine Folge" in r for r in xl.einfuege_hindernisse("Gruppe", 9))
    assert any("Excel-Tabelle" in r for r in xl.einfuege_hindernisse("Tab", 5))
    assert any("Diagramm" in r for r in xl.einfuege_hindernisse("Bild", 2))      # Diagramm zeigt auf A1:A2
    assert "gibt es in der Excel nicht" in xl.einfuege_hindernisse("Gibt es nicht", 1)[0]
    # die Weiche ueberschreibt die Einstellung: so sieht der openpyxl-Schreiber dieselbe Stelle
    assert any("Formeln" in r for r in xl.einfuege_hindernisse("Formel", 5, schreiber="openpyxl"))


# --- Ende zu Ende: Planer + Applier --------------------------------------------------------------------------
_EL = [("h", "1 Eins"), ("p", "Absatz eins mit genug Text darin."),
       ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin."),
       ("h", "{+2.5 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}"),
       ("h", "3 Drei"), ("p", "Absatz drei mit genug Text darin."),
       ("h", "4 Vier"), ("p", "Absatz vier mit genug Text darin."),
       ("h", "5 Fuenf"), ("p", "Absatz fuenf mit genug Text darin."),
       ("h", "6 Sechs"), ("p", "Absatz sechs mit genug Text darin.")]


def _lauf_mit(tmp_path, aendere):
    from tests.redteam._bau import analysiere, baue_alt, bearbeite, uebernimm
    alt = baue_alt(tmp_path, _EL, "L1", nachher=lambda p: bearbeite(p, aendere))
    l = uebernimm(analysiere(tmp_path, _EL, alt))
    neu = [a for a in l.aen if a.typ == "section_neu"]
    assert len(neu) == 1
    return l, neu[0]


def test_verbundene_zelle_und_versteckte_zeile_darunter_lassen_die_neue_section_automatisch_zu(tmp_path):
    from openpyxl import load_workbook
    from tests.redteam._bau import zeile_mit

    def eig(wb):
        wb["CDE"].merge_cells("D5:D6")                      # Section 4/5 -- beide unterhalb der Einfuegestelle (Zeile 4)
        wb["CDE"].row_dimensions[7].hidden = True           # Section 6
    l, neu = _lauf_mit(tmp_path, eig)
    assert (neu.status, neu.konfidenz) == ("uebernommen", "hoch"), (neu.status, neu.grund, neu.hinweise)
    assert neu.nachkontrolle == "ok"
    ws = load_workbook(str(l.out))["CDE"]
    r4, r5, r6 = (zeile_mit(l.out, t) for t in ("Absatz vier", "Absatz fuenf", "Absatz sechs"))
    assert [str(x) for x in ws.merged_cells.ranges] == [f"D{r4}:D{r5}"]                  # Verbindung wanderte mit
    assert ws.row_dimensions[r6].hidden and not ws.row_dimensions[r5].hidden and not ws.row_dimensions[r4].hidden


def test_neue_zeile_mitten_in_ausgeblendeter_folge_bleibt_pruefen(tmp_path):
    def eig(wb):
        wb["CDE"].row_dimensions[3].hidden = True           # Zeile 3 und 4 versteckt, die neue Zeile kaeme bei 4 dazwischen
        wb["CDE"].row_dimensions[4].hidden = True
    l, neu = _lauf_mit(tmp_path, eig)
    assert neu.status == "pruefen"
    assert any("mitten in eine Folge" in h for h in neu.hinweise), neu.hinweise


def test_neue_zeile_durch_verbundene_zellen_bleibt_pruefen(tmp_path):
    def eig(wb):
        wb["CDE"].merge_cells("D2:D4")                      # reicht ueber die Einfuegestelle (Zeile 4) hinweg
    l, neu = _lauf_mit(tmp_path, eig)
    assert neu.status == "pruefen"
    assert any("durch die Zeile" in h for h in neu.hinweise), neu.hinweise


# ===========================================================================================================
# K-PL 1: gestrichene Section und gestrichener Inhalt sind eine Gruppe (RT-122)
# ===========================================================================================================
_PDF_LOESCH = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
               ("h", "{-2 Zwei-}"), ("p", "{-Die Reaktionszeit darf 150 ms nicht ueberschreiten.-}"),
               ("h", "3 Drei"), ("p", "Alle Steckverbindungen sind eindeutig beschriftet."),
               ("h", "4 Vier"), ("p", "Die Lieferung umfasst alles Zubehoer fuer die Montage.")]


def _excel_elemente(ersetze: dict[str, str] | None = None):
    out = []
    for art, m in _PDF_LOESCH:
        m = m.replace("{-", "").replace("-}", "")
        out.append((art, (ersetze or {}).get(m, m)))
    return out


def _loesch_lauf(tmp_path, excel_el):
    from tests.redteam._bau import baue_alt, lauf
    alt = baue_alt(tmp_path, excel_el, "L1")
    l = lauf(tmp_path, _PDF_LOESCH, "L1", alt=alt)
    kopf = next(a for a in l.aen if a.typ == "section_geloescht")
    inhalt = next(a for a in l.aen if a.typ == "absatz_geloescht")
    return l, kopf, inhalt


def test_section_und_inhalt_sicher_werden_zusammen_geloescht(tmp_path):
    from tests.redteam._bau import zeile_mit
    l, kopf, inhalt = _loesch_lauf(tmp_path, _excel_elemente())
    assert (kopf.status, inhalt.status) == ("uebernommen", "uebernommen")
    assert zeile_mit(l.out, "Die Reaktionszeit") is None                      # der Inhalt ist weg
    assert zeile_mit(l.out, "2 Zwei") is None
    assert all(a.nachkontrolle == "ok" for a in (kopf, inhalt))


def test_inhalt_unsicher_haelt_die_ueberschrift_auf_pruefen(tmp_path):
    from tests.redteam._bau import zelle, zeile_mit
    # Excel: mehr Text im Absatz, als im PDF gestrichen ist -> der Inhalt ist nur ein Vorschlag
    el = _excel_elemente({"Die Reaktionszeit darf 150 ms nicht ueberschreiten.":
                          "Die Reaktionszeit darf 150 ms nicht ueberschreiten. Die Messung erfolgt bei Nennlast."})
    l, kopf, inhalt = _loesch_lauf(tmp_path, el)
    assert inhalt.status == "pruefen"
    assert kopf.status == "pruefen" and kopf.konfidenz == "mittel"
    assert "zusammen" in kopf.grund and kopf.hinweise
    assert kopf.ops, "der Vorschlag bleibt stehen, damit ein Mensch ihn freigeben kann"
    z = zeile_mit(l.out, "Die Reaktionszeit")
    assert (zelle(l.out, z, 1), zelle(l.out, z, 2)) == ("2", "Zwei")           # Nummer und Titel stehen noch: keine verwaiste Zeile


def test_ein_unsicherer_absatz_haelt_die_ganze_gruppe_auf_pruefen(tmp_path):
    from tests.redteam._bau import baue_alt, lauf
    pdf = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
           ("h", "{-2 Zwei-}"), ("p", "{-Die Reaktionszeit darf 150 ms nicht ueberschreiten.-}"),
           ("p", "{-Alle Steckverbindungen sind eindeutig beschriftet.-}"),
           ("h", "3 Drei"), ("p", "Die Lieferung umfasst alles Zubehoer fuer die Montage.")]
    xl = [(a, m.replace("{-", "").replace("-}", "")) for a, m in pdf]
    xl[3] = ("p", xl[3][1] + " Die Messung erfolgt bei Nennlast.")            # nur dieser Absatz ist unsicher
    alt = baue_alt(tmp_path, xl, "L1")
    l = lauf(tmp_path, pdf, "L1", alt=alt)
    gruppe = [a for a in l.aen if a.typ in ("section_geloescht", "absatz_geloescht")]
    assert len(gruppe) == 3, [(a.typ, a.status) for a in l.aen]
    assert {a.status for a in gruppe} == {"pruefen"}, [(a.id, a.typ, a.status, a.grund) for a in gruppe]


def test_ueberschrift_unsicher_haelt_auch_den_inhalt_auf_pruefen(tmp_path):
    from tests.redteam._bau import zeile_mit
    # Excel: die Ueberschrift hat mehr Woerter, als im PDF gestrichen sind -> nur ein Vorschlag
    el = _excel_elemente({"2 Zwei": "2 Zwei und Drei"})
    l, kopf, inhalt = _loesch_lauf(tmp_path, el)
    assert kopf.status == "pruefen"
    assert inhalt.status == "pruefen" and inhalt.konfidenz == "mittel"
    assert "zusammen" in inhalt.grund
    assert zeile_mit(l.out, "Die Reaktionszeit") is not None                  # der Inhalt wurde nicht allein geloescht


# ===========================================================================================================
# K-PL 2: Neutext mit Trennmuster "Wort- Rest" oder Tabellen-Spaltentrenner wird nie automatisch uebernommen
# ===========================================================================================================
@pytest.mark.parametrize("text", [
    "Die Netzanschlussbe- dingungen sind einzuhalten.",
    "Das Betrie- bsführungskonzept liegt vor.",
    "Das Betrie­ bsführungskonzept liegt vor.",                       # weiches Trennzeichen
    "Der Schalt- schrank ist beschriftet.",
    "Siehe oben.\nDie Anlage wird jaehrlich ueberpru- ft.",                 # in der zweiten Zeile
])
def test_trennmuster_wird_erkannt(text):
    from src.planner import neutext_probleme
    p = neutext_probleme([text])
    assert any("Trennmuster" in h for h in p["verdaechtig"]), p


@pytest.mark.parametrize("text", [
    "Der Ein- und Ausgang ist beschriftet.",                               # Ergaenzungsstrich
    "Be- oder Entladen ist verboten.",
    "Kabel- bzw. Leitungsschutz ist vorzusehen.",
    "Wartungs-, Inspektions- sowie Pruefarbeiten sind zu dokumentieren.",
    "Ein-, Aus- und Umbau sind getrennt zu melden.",
    "Sowohl Ein- als auch Ausgaenge sind gesichert.",
    "Der kWh-Zähler und die E-Mail-Adresse stehen im Nord-Süd-Plan.",      # echte Bindestriche ohne Leerzeichen
    "Kennlinie A-B und Zeitraum 2020-2024 sind gleich.",
    "Der Wert - wie gemessen - liegt bei 5 V.",                            # Bindestrich mit Leerzeichen davor
    "Der Soll- Ist- Vergleich entfaellt.",                                 # Grossbuchstabe nach dem Strich
    "Trennung a|b ohne Leerzeichen ist kein Tabellenrand.",
])
def test_echte_bindestriche_blockieren_nicht(text):
    from src.planner import neutext_probleme
    assert neutext_probleme([text])["verdaechtig"] == []


@pytest.mark.parametrize("text", [
    "Pos. | Bauteil | Menge | Einheit",
    "| Pos. | Bauteil |\n| 1 | Netzteil |",
    "Netzteil | 24 Stueck",
])
def test_tabellen_pipes_werden_erkannt(text):
    from src.planner import neutext_probleme
    p = neutext_probleme([text])
    assert any("Spaltentrenner" in h for h in p["verdaechtig"]), p


def test_muster_kommen_aus_der_config():
    from src.planner import neutext_probleme
    cfg = {"planer": {"trennmuster": [], "tabellen_muster": []}}
    assert neutext_probleme(["Schalt- schrank", "A | B"], cfg)["verdaechtig"] == []         # ohne Muster in der Config nichts
    cfg = {"planer": {"trennmuster": [r"(?<=\w)- (?P<rest>[a-z]\w*)"], "trennung_ausnahmen": ["schrank"]}}
    assert neutext_probleme(["Schalt- schrank"], cfg)["verdaechtig"] == []                   # Ausnahme aus der Config
    assert neutext_probleme(["Schalt- kasten"], cfg)["verdaechtig"]


def _neu_lauf(tmp_path, neuer_text: str):
    from tests.redteam._bau import baue_alt, lauf
    xl = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
          ("h", "2 Zwei"), ("p", "Die Reaktionszeit darf 150 ms nicht ueberschreiten."),
          ("h", "3 Drei"), ("p", "Alle Steckverbindungen sind eindeutig beschriftet.")]
    pdf = xl[:4] + [("p", "{+" + neuer_text + "+}")] + xl[4:]
    alt = baue_alt(tmp_path, xl, "L1")
    l = lauf(tmp_path, pdf, "L1", alt=alt)
    neu = [a for a in l.aen if a.typ == "absatz_neu"]
    assert len(neu) == 1
    return l, neu[0]


def test_neuer_absatz_mit_trennmuster_bleibt_auf_pruefen_mit_hinweis(tmp_path):
    l, neu = _neu_lauf(tmp_path, "Die Netzanschlussbe- dingungen sind einzuhalten und zu dokumentieren.")
    assert neu.status == "pruefen" and neu.konfidenz == "mittel"
    assert any("Trennmuster" in h for h in neu.hinweise) and "verfälscht" in neu.grund


def test_neuer_absatz_mit_tabellen_pipes_bleibt_auf_pruefen_mit_hinweis(tmp_path):
    l, neu = _neu_lauf(tmp_path, "Pos. | Bauteil | Menge | Einheit")
    assert neu.status == "pruefen"
    assert any("Spaltentrenner" in h for h in neu.hinweise)


def test_neuer_absatz_mit_ergaenzungsstrich_wird_weiter_automatisch_uebernommen(tmp_path):
    l, neu = _neu_lauf(tmp_path, "Der Ein- und Ausgang ist beschriftet und kWh-Zähler sind plombiert.")
    assert (neu.status, neu.konfidenz) == ("uebernommen", "hoch"), (neu.status, neu.grund, neu.hinweise)


# ===========================================================================================================
# K-PDF: Silbentrennung mit kurzem ersten Stueck (RT-115 Seed 51) und Zwei-Spalten-Satz (RT-119 Seed 42)
# ===========================================================================================================
def _tok(alt: str, neu: str, status: str):
    from src.models import PdfToken
    return PdfToken(old=alt, new=neu, status=status, page=1, bbox=(0, 0, 1, 1))


def test_weicher_trennstrich_nach_zwei_buchstaben_wird_zusammengefuegt():
    from src.pdf_reader import _verbinden
    z = _verbinden(_tok("", "ko", "add"), _tok("", "ntrolliert", "add"), True, set(), 2)
    assert z is not None and z.new == "kontrolliert"
    z = _verbinden(_tok("Ne", "Ne", "keep"), _tok("nnleistung", "nnleistung", "keep"), True, set(), 2)
    assert z is not None and (z.old, z.new) == ("Nennleistung", "Nennleistung")


def test_trennstrich_nach_zwei_buchstaben_ascii_wie_vorher():
    from src.pdf_reader import _verbinden
    z = _verbinden(_tok("", "ko-", "add"), _tok("", "ntrolliert", "add"), False, set(), 2)
    assert z is not None and z.new == "kontrolliert"


def test_zu_kurzes_stueck_oder_fehlender_strich_wird_nicht_zusammengefuegt():
    from src.pdf_reader import _verbinden
    assert _verbinden(_tok("", "k", "add"), _tok("", "ontrolliert", "add"), True, set(), 2) is None     # ein Buchstabe
    assert _verbinden(_tok("", "k-", "add"), _tok("", "ontrolliert", "add"), False, set(), 2) is None
    assert _verbinden(_tok("", "ko", "add"), _tok("", "ntrolliert", "add"), False, set(), 2) is None    # kein Trennstrich (nicht weich)
    assert _verbinden(_tok("", "ko", "add"), _tok("", "Ntrolliert", "add"), True, set(), 2) is None     # Grossbuchstabe dahinter
    assert _verbinden(_tok("", "Ein", "add"), _tok("", "und", "add"), True, {"und"}, 2) is None          # Ausnahmewort
    assert _verbinden(_tok("", "ko", "add"), _tok("Ntrolliert", "", "del"), True, set(), 2) is None     # verschiedener Status


def test_mindestlaenge_der_silbe_steht_in_der_config():
    from src import config
    assert config.lade("farben")["layout"]["silbentrennung_min_buchstaben"] == 2


def _spalten_lauf(tmp_path, cfg=None):
    from src.pdf_reader import analysiere_pdf
    from src.planner import plane
    from tests.redteam._bau import Lauf, baue_alt, uebernimm
    from tests.redteam._bau_pdf import dok
    kopf = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
            ("h", "2 Zwei"), ("p", "Die Reaktionszeit darf 150 ms nicht ueberschreiten.")]
    links = ["Der Nennwert von 20 Monate gilt bei 40 Grad Umgebungstemperatur. Die Betriebsspannung darf 120 V nicht "
             "ueberschreiten.", "Bei Luefterausfall wird ein Bericht an die Leitstelle uebermittelt."]
    rechts = "Bei Ueberstrom wird ein Alarm an die Leitstelle uebermittelt."
    el_excel = kopf + [("h", "3 Drei"), ("p", links[0]), ("p", links[1]), ("h", "4 Vier"),
                       ("p", "Die Lieferung umfasst alles Zubehoer fuer die Montage.")]
    d = dok()
    for art, markup in kopf + [("h", "3 Drei")]:
        d.ueberschrift(markup) if art == "h" else d.block(markup)
    d.zwei_spalten(links, ["{+" + rechts + "+}"])
    d.ueberschrift("4 Vier")
    d.block("Die Lieferung umfasst alles Zubehoer fuer die Montage.")
    pdf = tmp_path / "a.pdf"
    d.speichern(pdf)
    alt = baue_alt(tmp_path, el_excel, "L1")
    analyse = analysiere_pdf(pdf)
    from src.excel_model import ExcelModell
    xl = ExcelModell(alt)
    l = Lauf(plan=plane(analyse, xl, cfg) if cfg else plane(analyse, xl), xl=xl, alt=alt)
    return analyse, uebernimm(l), rechts


def test_geaenderte_woerter_im_zwei_spalten_bereich_tragen_den_hinweis(tmp_path):
    analyse, l, _ = _spalten_lauf(tmp_path)
    assert any(lu["art"] == "mehrspaltig" for lu in analyse.luecken)
    geaendert = [t for b in analyse.blocks for t in b.tokens if t.status != "keep"]
    assert geaendert and all("mehrspaltig" in t.hinweis for t in geaendert)
    unveraendert = [t for b in analyse.blocks for t in b.tokens if t.status == "keep"]
    assert not any("mehrspaltig" in t.hinweis for t in unveraendert)
    assert any("mehreren Spalten" in h for b in analyse.blocks for h in b.hinweise)


def test_zwei_spalten_aenderungen_werden_nie_automatisch_eingetragen(tmp_path):
    from tests.redteam._bau import zelle, zeile_mit
    analyse, l, rechts = _spalten_lauf(tmp_path)
    assert l.aen and all(a.status == "pruefen" for a in l.aen), [(a.id, a.status) for a in l.aen]
    assert all(any("Spalten" in h for h in a.hinweise) and "nicht sicher gelesen" in a.grund for a in l.aen)
    z = zeile_mit(l.out, "Der Nennwert von 20 Monate")
    assert "Alarm" not in zelle(l.out, z, 3)                       # nichts zerrissen in der Excel


def test_zwei_spalten_hinweis_greift_nur_wenn_die_config_ihn_nennt(tmp_path):
    from src import config
    import copy
    cfg = copy.deepcopy(config.lade("zuordnung"))
    cfg["planer"]["unsichere_hinweise"] = []
    analyse, l, _ = _spalten_lauf(tmp_path, cfg)
    assert any(a.status == "uebernommen" for a in l.aen)          # ohne den Eintrag in der Config gilt das alte Verhalten
