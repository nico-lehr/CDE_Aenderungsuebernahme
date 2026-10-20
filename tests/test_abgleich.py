"""Tests fuer den Absatz-Abgleich (src/abgleich.py) und seine Verdrahtung im Planer (ADR 0004).

Nur erfundene Mini-Faelle (Regel 16). Die PDF-Seite wird direkt als PdfAnalyse gebaut (ohne PDF-Datei), die Excel mit openpyxl.
"""
from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from src.abgleich import (EINDEUTIG, KEINER, MEHRDEUTIG, REIHENFOLGE, WIDERSPRUCH, ExaktIndex, Fundstelle,
                          laeufe_aus_tokens, ops_fuer_spur, reissverschluss, spuren)
from src.applier import _text, uebernehmen
from src.excel_model import ExcelModell
from src.gegenprobe import pruefe_absatz, zustand
from src.models import ADD, DEL, KEEP, MOD, NICHT, PRUEFEN, UEBERNOMMEN, PdfBlock, PdfToken
from src.pdf_reader import PdfAnalyse
from src.planner import plane
from tests.redteam._bau import schreibe_excel


# ------------------------------------------------------------------------------------------------ Bau-Helfer
def tok(old, new, status, hinweis="", trenner=" ", teile=None, trennstelle=-1, page=1):
    return PdfToken(old=old, new=new, status=status, page=page, bbox=(0, 0, 1, 1), hinweis=hinweis, trenner=trenner,
                    teile=list(teile or []), trennstelle=trennstelle)


def woerter(markup: str, page: int = 1) -> list[PdfToken]:
    """'Das {-alte-} {+neue+} Wort' -> Tokens (je Leerzeichen ein Wort)."""
    out = []
    for m in re.finditer(r"\{\+(.+?)\+\}|\{-(.+?)-\}|(\S+)", markup):
        if m.group(1):
            out += [tok("", w, ADD, page=page) for w in m.group(1).split()]
        elif m.group(2):
            out += [tok(w, "", DEL, page=page) for w in m.group(2).split()]
        else:
            out.append(tok(m.group(3), m.group(3), KEEP, page=page))
    return out


def pdf_aus(elemente) -> PdfAnalyse:
    """[(art, markup) | (art, markup, seite, fortsetzung) | (art, [tokens])] -> PdfAnalyse mit Sections wie der Leser."""
    blocks = []
    sec_alt = sec_neu = None
    for i, e in enumerate(elemente):
        art, inhalt = e[0], e[1]
        seite = e[2] if len(e) > 2 else 1
        toks = inhalt if isinstance(inhalt, list) else woerter(inhalt, seite)
        kind = "ueberschrift" if art == "h" else "absatz"
        if kind == "ueberschrift":
            alt = [t.old for t in toks if t.old]
            neu = [t.new for t in toks if t.new]
            sec_alt = alt[0] if alt else None
            sec_neu = neu[0] if neu else None
        b = PdfBlock(idx=i, page=seite, kind=kind, tokens=toks, bbox=(0, 0, 1, 1), section_alt=sec_alt, section_neu=sec_neu)
        if len(e) > 3:
            b.fortsetzung = e[3]
        blocks.append(b)
    return PdfAnalyse(pfad="x.pdf", seiten=max(b.page for b in blocks), blocks=blocks)


def excel(tmp_path, zeilen, name="a.xlsx", nachher=None):
    p = schreibe_excel(tmp_path / name, [["Nr.", "Titel", "Inhalt"], *zeilen])
    if nachher:
        wb = load_workbook(str(p))
        nachher(wb)
        wb.save(str(p))
    return p


def zelle(pfad, zeile, spalte, blatt="CDE"):
    return _text(load_workbook(str(pfad), rich_text=True)[blatt].cell(zeile, spalte))


def lauf(tmp_path, pdf, alt, name="out.xlsx"):
    xl = ExcelModell(alt)
    plan = plane(pdf, xl)
    out = tmp_path / name
    uebernehmen(plan.aenderungen, alt, out, None, absaetze=plan.absaetze)
    return plan, out


def a3_ok(plan, alt, out):
    """Gegenprobe A3: jeder Absatz mit automatisch angewendeter Aenderung steht so in der gespeicherten Excel, wie das PDF ihn
    verlangt (gegenprobe.pruefe_absatz liefert None)."""
    wb_alt = load_workbook(str(alt), rich_text=True)
    wb_neu = load_workbook(str(out), rich_text=True)
    angewendet = {a.id for a in plan.aenderungen if a.status == UEBERNOMMEN}
    for ab in plan.absaetze:
        if not (ab.aenderungen & angewendet):
            continue
        zellen = {}
        for z, s, _, _ in ab.zellen:
            ops = [o for a in plan.aenderungen if a.id in angewendet for o in a.ops
                   if (o.blatt, o.zeile, o.spalte) == (ab.blatt, z, s)]
            zellen[(z, s)] = (_text(wb_alt[ab.blatt].cell(z, s)), _text(wb_neu[ab.blatt].cell(z, s)), ops)
        assert pruefe_absatz(ab, zellen, angewendet) is None, ab.id


def fund(text, zeile=2, spalte=3, start=0, blatt="CDE", art="absatz", **kw):
    return Fundstelle(blatt=blatt, zellen=((zeile, spalte, start, start + len(text)),), art=art,
                      schluessel=(0, zeile, 2, spalte, start), text=text,
                      end_schluessel=(0, zeile, 2, spalte, start + len(text)), **kw)


def anwenden(text, ops):
    for o in sorted(ops, key=lambda o: o.start, reverse=True):
        assert text[o.start:o.ende] is not None
        text = text[:o.start] + o.text + text[o.ende:]
    return text


def b_ops(laeufe, excel_text):
    so = ops_fuer_spur(laeufe, fund(excel_text), lambda b, z, s: excel_text)
    return so, anwenden(excel_text, [o for ops in so.ops.values() for o in ops])


# ------------------------------------------------------------------------------------------------ 1. Laeufe
def test_laeufe_ganze_woerter_und_ids_je_gruppe():
    l = laeufe_aus_tokens(woerter("Das System {-muss-} {+soll+} die Daten {+sicher+} speichern."))
    ids = {x.aenderung for x in l if x.aenderung}
    assert ids == {"G1", "G2"}
    assert zustand(l) == "Das System muss die Daten speichern."
    assert zustand(l, ids) == "Das System soll die Daten sicher speichern."
    assert zustand(l, {"G1"}) == "Das System soll die Daten speichern."


def test_laeufe_zeichenlaeufe_und_mod_und_trenner():
    t = [tok("Anlage", "Anlage", KEEP), tok("", "insgesamt.", ADD, teile=[("add", "insgesamt"), ("keep", ".")])]
    l = laeufe_aus_tokens(t)
    assert [(x.art, x.text) for x in l] == [("k", "Anlage"), ("s", " "), ("a", "insgesamt"), ("k", ".")]
    m = laeufe_aus_tokens([tok("Rack", "Racks", MOD, trenner="\xa0"), tok("da", "da", KEEP)])
    assert [(x.art, x.text) for x in m] == [("d", "Rack"), ("a", "Racks"), ("s", "\xa0"), ("k", "da")]
    g = laeufe_aus_tokens([tok("ab", "ab", KEEP, trenner=""), tok("", "cd", ADD)])
    assert [x.art for x in g] == ["k", "a"]


def test_laeufe_trennstrich_ist_optional():
    l = laeufe_aus_tokens([tok("E-Mail", "E-Mail", KEEP, trennstelle=1)])
    assert [(x.text, x.optional) for x in l] == [("E", False), ("-", True), ("Mail", False)]


# ------------------------------------------------------------------------------------------------ 5. A1 + B (zeichengenau)
def test_satzzeichen_blaues_wort_vor_schwarzem_punkt():
    l = laeufe_aus_tokens([tok("die", "die", KEEP), tok("Anlage", "Anlage", KEEP),
                           tok("", "insgesamt.", ADD, teile=[("add", "insgesamt"), ("keep", ".")])])
    so, neu = b_ops(l, "die Anlage.")
    assert so.vergleich.gleich and neu == "die Anlage insgesamt."
    assert all(o.art == "ersetzen" for ops in so.ops.values() for o in ops)


def test_satzzeichen_rotes_wort_vor_schwarzem_punkt():
    l = laeufe_aus_tokens([tok("Anlage", "Anlage", KEEP),
                           tok("insgesamt.", ".", DEL, teile=[("del", "insgesamt"), ("keep", ".")])])
    assert b_ops(l, "Anlage insgesamt.")[1] == "Anlage."


@pytest.mark.parametrize("excel_text, erwartet", [
    ("die\xa0Daten speichern", "die\xa0neuen Daten speichern"),        # Excel-eigene Naht bleibt
    ("die Daten\tspeichern", "die neuen Daten\tspeichern"),
])
def test_nbsp_und_tab_an_der_naht_bleiben(excel_text, erwartet):
    l = laeufe_aus_tokens(woerter("die {+neuen+} Daten speichern"))
    assert b_ops(l, excel_text)[1] == erwartet


def test_ersetzen_loeschen_und_rand():
    l = laeufe_aus_tokens(woerter("{-Zudem-} gilt: darf {-12-} {+7+} ms {+dauern+}"))
    assert b_ops(l, "Zudem gilt: darf 12 ms")[1] == "gilt: darf 7 ms dauern"
    l2 = laeufe_aus_tokens(woerter("Text {-weg-} bleibt"))
    assert b_ops(l2, "Text\xa0weg bleibt")[1] == "Text\xa0bleibt"


def test_a1_abweichung_gibt_keine_operationen():
    l = laeufe_aus_tokens(woerter("Die Druckwerte {-werden-} {+sind+} gespeichert."))
    so = ops_fuer_spur(l, fund("Die Druckwarte werden gespeichert."), lambda b, z, s: "Die Druckwarte werden gespeichert.")
    assert not so.vergleich.gleich and so.vergleich.art == "buchstabe" and not so.ops


def test_zeilenumbruch_in_der_spanne_ist_ein_problem():
    l = laeufe_aus_tokens(woerter("eins {-zwei drei-} vier"))
    so = ops_fuer_spur(l, fund("eins zwei\ndrei vier", art="zelle"), lambda b, z, s: "eins zwei\ndrei vier")
    assert so.vergleich.gleich and "excel_umbruch" in next(iter(so.probleme.values()))


def test_ueberschrift_ueber_zwei_zellen():
    """Nummer und Titel in zwei Zellen: die Aenderung der Nummer landet nur in der Nummernzelle."""
    texte = {(2, 1): "2.3", (2, 2): "Sicherheit"}
    f = Fundstelle(blatt="CDE", zellen=((2, 1, 0, 3), (2, 2, 0, 10)), art="kopf", schluessel=(0, 2, 0, 1, 0),
                   text="2.3 Sicherheit", end_schluessel=(0, 2, 1, 2, 10))
    l = laeufe_aus_tokens(woerter("{-2.3-} {+2.2+} Sicherheit"))
    so = ops_fuer_spur(l, f, lambda b, z, s: texte[(z, s)])
    (op,) = next(iter(so.ops.values()))
    assert (op.zeile, op.spalte, op.start, op.ende, op.text) == (2, 1, 0, 3, "2.2")
    l2 = laeufe_aus_tokens(woerter("2.3 {+Neue+} Sicherheit"))
    (op2,) = next(iter(ops_fuer_spur(l2, f, lambda b, z, s: texte[(z, s)]).ops.values()))
    assert (op2.spalte, op2.start, op2.ende, op2.text) == (2, 0, 0, "Neue ")


# ------------------------------------------------------------------------------------------------ 4. Reissverschluss
def _f(zeile, blatt=0, **kw):
    return Fundstelle(blatt="CDE", zellen=((zeile, 3, 0, 5),), art="absatz", schluessel=(blatt, zeile, 2, 3, 0), text="x",
                      end_schluessel=(blatt, zeile, 2, 3, 5), **kw)


def test_reissverschluss_dublette_durch_reihenfolge():
    a, b, c = _f(2), _f(5), _f(9)
    dubl = [(_f(3), "alt"), (_f(7), "alt")]          # gleicher Absatz zweimal; nur Zeile 3 liegt zwischen den Ankern 2 und 5
    z = reissverschluss([[(a, "alt")], dubl, [(b, "alt")], [(c, "alt")]])
    assert [x.status for x in z] == [EINDEUTIG, REIHENFOLGE, EINDEUTIG, EINDEUTIG]
    assert z[1].fund.zeile == 3


def test_reissverschluss_mehrdeutig_widerspruch_keiner():
    z = reissverschluss([[(_f(2), "alt")], [(_f(3), "alt"), (_f(4), "alt")], [(_f(9), "alt")], [], [(_f(1), "alt")]])
    assert [x.status for x in z] == [EINDEUTIG, MEHRDEUTIG, EINDEUTIG, KEINER, WIDERSPRUCH]


def test_reissverschluss_sichtbar_schlaegt_versteckt_bei_mehrdeutig():
    z = reissverschluss([[(_f(3, versteckt="zeile"), "alt"), (_f(4), "alt")]])
    assert z[0].status == MEHRDEUTIG and not z[0].fund.versteckt


# ------------------------------------------------------------------------------------------------ 2. Spuren
def test_spur_seitenumbruch_sicher_wird_ein_absatz():
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage wird geprüft und", 1, "sicher"), ("p", "danach {+sorgfältig+} dokumentiert.", 2)])
    sp = spuren(pdf)
    assert len(sp) == 2 and sp[1].alt == "Die Anlage wird geprüft und danach dokumentiert."


def test_spur_moeglich_mit_ganz_neuem_text_ist_unklar():
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage wird geprüft und {+zusätzlich+}", 1, "moeglich"),
                   ("p", "{+jede Prüfung wird vermerkt.+}", 2)])
    sp = spuren(pdf)
    assert len(sp) == 3 and "fortsetzung_unklar" in sp[1].codes and "fortsetzung_unklar" in sp[2].codes


def test_spur_moeglich_wird_per_a1_entschieden(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage wird geprüft und danach dokumentiert."]])
    idx = ExaktIndex(ExcelModell(alt))
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage wird geprüft und", 1, "moeglich"), ("p", "danach {-sofort-} dokumentiert.", 2)])
    assert len(spuren(pdf, idx)) == 3           # verbunden steht der alte Absatz nicht in der Excel ("sofort" fehlt) -> getrennt
    pdf2 = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage wird geprüft und", 1, "moeglich"), ("p", "danach {+bald+} dokumentiert.", 2)])
    assert len(spuren(pdf2, idx)) == 2


def test_seitenumbruch_blauer_satz_wird_ans_absatzende_angehaengt(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage wird geprüft."], ["2", "Zwei", "Zweiter Absatz mit genug Text darin."]])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage wird geprüft. {+Jede Prüfung+}", 1, "sicher"),
                   ("p", "{+wird vermerkt.+}", 2), ("h", "2 Zwei", 2), ("p", "Zweiter Absatz mit genug Text darin.", 2)])
    plan, out = lauf(tmp_path, pdf, alt)
    assert [a.status for a in plan.aenderungen] == [UEBERNOMMEN]
    assert zelle(out, 2, 3) == "Die Anlage wird geprüft. Jede Prüfung wird vermerkt."
    a3_ok(plan, alt, out)


# ------------------------------------------------------------------------------------------------ Planer Ende-zu-Ende
ZWEI = ["2", "Zwei", "Zweiter Absatz mit genug Text darin."]


def test_ende_zu_ende_hoch_mit_absaetzen_und_a3(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Das System muss die Daten speichern.\nDie Anlage läuft."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Das System {-muss-} {+soll+} die Daten speichern."),
                   ("p", "Die Anlage läuft {+stabil+}."), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert all(a.status == UEBERNOMMEN and a.konfidenz == "hoch" for a in plan.aenderungen)
    assert zelle(out, 2, 3) == "Das System soll die Daten speichern.\nDie Anlage läuft stabil."
    assert len(plan.absaetze) == 2 and {a.id for a in plan.aenderungen} == {i for ab in plan.absaetze for i in ab.aenderungen}
    a = plan.aenderungen[0]
    assert a.soll_absatz == "Das System soll die Daten speichern." and a.ist_absatz == "Das System muss die Daten speichern."
    a3_ok(plan, alt, out)


def test_a1_excel_weicht_ab_ist_pruefen_mit_code(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Druckwarte werden täglich gespeichert."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Druckwerte werden {-täglich-} {+stündlich+} gespeichert."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "excel_weicht_ab" in a.pruefcodes and "a1_buchstabe" in a.pruefcodes
    assert "steht so nicht" in a.grund
    assert zelle(out, 2, 3) == "Die Druckwarte werden täglich gespeichert."


def test_historie_blatt_mit_zitiertem_alttext_stoert_nicht(tmp_path):
    def historie(wb):
        ws = wb.create_sheet("Historie")
        ws.append(["Datum", "Text"])
        ws.append(["01.01.", "Das System muss die Daten speichern."])
    alt = excel(tmp_path, [["1", "Eins", "Das System muss die Daten speichern."], ZWEI], nachher=historie)
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Das System {-muss-} {+soll+} die Daten speichern."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert plan.aenderungen[0].status == UEBERNOMMEN
    assert zelle(out, 2, 3) == "Das System soll die Daten speichern."
    assert zelle(out, 2, 2, "Historie") == "Das System muss die Daten speichern."


def test_dublette_im_dokument_wird_ueber_die_reihenfolge_eindeutig(tmp_path):
    satz = "Die Prüfung erfolgt jährlich durch den Betreiber."
    alt = excel(tmp_path, [["1", "Eins", "Erster Absatz mit genug Text darin."], ["", "", satz],
                           ["2", "Zwei", "Zweiter Absatz mit genug Text darin."], ["", "", satz]])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("p", satz),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin."),
                   ("p", "Die Prüfung erfolgt {-jährlich-} {+monatlich+} durch den Betreiber.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert plan.aenderungen[0].status == UEBERNOMMEN
    assert zelle(out, 3, 3) == satz and zelle(out, 5, 3) == satz.replace("jährlich", "monatlich")


def test_versteckte_koeder_zeile_wird_nie_gewaehlt(tmp_path):
    satz = "Die Prüfung erfolgt jährlich durch den Betreiber."

    def verstecke(wb):
        wb["CDE"].row_dimensions[3].hidden = True
    alt = excel(tmp_path, [["1", "Eins", "Erster Absatz mit genug Text darin."], ["", "", satz], ["", "", satz], ZWEI],
                nachher=verstecke)
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."),
                   ("p", "Die Prüfung erfolgt {-jährlich-} {+monatlich+} durch den Betreiber."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    (a,) = plan.aenderungen
    assert a.ops[0].zeile == 4                     # die sichtbare Stelle


def test_absatz_sperre_ueber_alle_woerter(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage meldet Fehler sofort an die Leitstelle."], ZWEI])
    toks = woerter("Die Anlage meldet {+alle+} Fehler sofort an die Leitstelle.")
    toks[-1].hinweis = "unbekannte_farbe"            # ein UNGEAENDERTES Wort mit unerklaerter Farbe
    pdf = pdf_aus([("h", "1 Eins"), ("p", toks), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "sperre:unbekannte_farbe" in a.pruefcodes and a.hinweise
    assert zelle(out, 2, 3) == "Die Anlage meldet Fehler sofort an die Leitstelle."


def test_nur_hinweis_codes_sperren_nicht(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage meldet Fehler sofort an die Leitstelle."], ZWEI])
    toks = woerter("Die Anlage meldet {+alle+} Fehler sofort an die Leitstelle.")
    toks[1].hinweis = "link"
    pdf = pdf_aus([("h", "1 Eins"), ("p", toks), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, _ = lauf(tmp_path, pdf, alt)
    assert plan.aenderungen[0].status == UEBERNOMMEN and plan.aenderungen[0].hinweise


@pytest.mark.parametrize("markup", ["Das System {-muss-} {+soll+} die Daten speichern.",
                                    "Das System muss die {-alten-} Daten speichern.",
                                    "Das System muss die Daten {+sicher+} speichern."])
def test_idempotenz_ersetzen_loeschen_einfuegen(tmp_path, markup):
    vorher = "Das System muss die alten Daten speichern." if "alten" in markup else "Das System muss die Daten speichern."
    alt = excel(tmp_path, [["1", "Eins", vorher], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", markup), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert plan.aenderungen[0].status == UEBERNOMMEN
    plan2, out2 = lauf(tmp_path, pdf, out, "out2.xlsx")
    (a,) = plan2.aenderungen
    assert a.status == NICHT and "steht schon so" in a.grund and "schon_eingetragen" in a.pruefcodes
    assert zelle(out2, 2, 3) == zelle(out, 2, 3)


def test_ganzer_absatz_geloescht_nimmt_den_zeilenumbruch_mit(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Erster Satz bleibt stehen.\nDieser Absatz entfällt ganz.\nLetzter bleibt."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Satz bleibt stehen."), ("p", "{-Dieser Absatz entfällt ganz.-}"),
                   ("p", "Letzter bleibt."), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert plan.aenderungen[0].status == UEBERNOMMEN
    assert zelle(out, 2, 3) == "Erster Satz bleibt stehen.\nLetzter bleibt."
    a3_ok(plan, alt, out)


def test_neuer_absatz_ersetzt_geloeschte_erste_zeile(tmp_path):
    """Erste Zeile gestrichen, dahinter ein neuer Absatz: beides automatisch, ohne Konflikt, und die Gegenprobe A3 stimmt."""
    alt = excel(tmp_path, [["1", "Eins", "Alter erster Absatz fällt weg.\nZweiter Absatz bleibt stehen."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "{-Alter erster Absatz fällt weg.-}"), ("p", "{+Neuer erster Absatz steht hier.+}"),
                   ("p", "Zweiter Absatz bleibt stehen."), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    assert [a.status for a in plan.aenderungen] == [UEBERNOMMEN, UEBERNOMMEN]
    assert zelle(out, 2, 3) == "Neuer erster Absatz steht hier.\nZweiter Absatz bleibt stehen."
    a3_ok(plan, alt, out)


def test_gestrichene_section_und_inhalt_sind_eine_gruppe(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Erster Absatz mit genug Text darin."], ["2", "Weg", "Inhalt der weg kommt."],
                           ["3", "Drei", "Dritter Absatz mit genug Text darin."]])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "{-2 Weg-}"),
                   ("p", "{-Inhalt der weg kommt.-}"), ("h", "3 Drei"), ("p", "Dritter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    kopf, inhalt = plan.aenderungen
    assert kopf.typ == "section_geloescht" and kopf.gruppe == inhalt.gruppe == kopf.id
    a3_ok(plan, alt, out)


def test_unscharfer_treffer_ist_hoechstens_mittel(tmp_path):
    """Frueher 'hoch' ab 90 % gleicher Woerter; jetzt nur noch Vorschlag (Regel 27)."""
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage meldet alle Fehler sofort an die zentrale Leitstelle des Betreibers."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage meldet alle Fehler {-sofort-} {+umgehend+} an die Leitstelle des Betreibers."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, _ = lauf(tmp_path, pdf, alt)
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and a.konfidenz == "mittel" and "excel_weicht_ab" in a.pruefcodes


# ------------------------------------------------------------------------------------------------ Nacharbeit
def test_ersatzweg_bekommt_absatz_wenn_die_laeufe_passen(tmp_path):
    """Hochzahl in der Excel: der Exakt-Index (Wort-Vorauswahl) findet die Stelle nicht, der alte Matcher schon. A1 ist in der
    Normalform gleich -> der Vorschlag bekommt einen PDF-Absatz; nach Freigabe besteht er die Gegenprobe A3."""
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage¹ meldet Fehler sofort an die Leitstelle."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage meldet {+alle+} Fehler sofort an die Leitstelle."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    xl = ExcelModell(alt)
    plan = plane(pdf, xl)
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and any(a.id in ab.aenderungen for ab in plan.absaetze)
    out = tmp_path / "frei.xlsx"
    uebernehmen(plan.aenderungen, alt, out, {a.id}, absaetze=plan.absaetze)
    assert zelle(out, 2, 3) == "Die Anlage¹ meldet alle Fehler sofort an die Leitstelle."
    a3_ok(plan, alt, out)


def test_ersatzweg_ohne_absatz_wenn_die_excel_abweicht(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Die Anlage meldet Fehler sofort an die zentrale Leitstelle."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Anlage meldet {+alle+} Fehler sofort an die Leitstelle."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan = plane(pdf, ExcelModell(alt))
    assert plan.aenderungen[0].status == PRUEFEN and not plan.absaetze


def test_gestrichener_listenpunkt_und_neuer_direkt_danach_ist_eine_ersetzung_in_der_zelle(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", None], ["", "", "Erster Absatz mit genug Text darin."], ["", "", "a) Alter Listenpunkt hier."],
                           ["", "", "b) Bleibt stehen hier."], ["2", "Zwei", None], ["", "", "Zweiter Absatz mit genug Text darin."]])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("p", "{-a) Alter Listenpunkt hier.-}"),
                   ("p", "{+a) Neuer Listenpunkt hier.+}"), ("p", "b) Bleibt stehen hier."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    loesch, neu = plan.aenderungen
    assert (loesch.status, neu.status) == (UEBERNOMMEN, UEBERNOMMEN) and loesch.gruppe == neu.gruppe != ""
    assert neu.ops[0].art == "ersetzen" and neu.ops[0].zeile == 4
    assert zelle(out, 4, 3) == "a) Neuer Listenpunkt hier." and zelle(out, 5, 3) == "b) Bleibt stehen hier."
    assert load_workbook(str(out))["CDE"].max_row == load_workbook(str(alt))["CDE"].max_row      # keine leere Zeile, keine neue
    a3_ok(plan, alt, out)


def test_neue_section_vor_seitenumbruch_ohne_satzende_ist_unklar(tmp_path):
    alt = excel(tmp_path, [["1", "Eins", "Erster Absatz mit genug Text darin."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "{+1.1 Neu+}"),
                   ("p", "{+Die Anlage wird geprüft, sofern keine+}"), ("p", "Fußnote eins."),
                   ("p", "{+Störung ansteht.+}", 2), ("h", "2 Zwei", 2), ("p", "Zweiter Absatz mit genug Text darin.", 2)])
    plan, _ = lauf(tmp_path, pdf, alt)
    sec = next(a for a in plan.aenderungen if a.typ == "section_neu")
    folge = next(a for a in plan.aenderungen if a.typ == "absatz_neu" and "geprüft" in a.neu)
    assert sec.status == folge.status == PRUEFEN
    assert "sperre:fortsetzung_unklar" in sec.pruefcodes and "sperre:fortsetzung_unklar" in folge.pruefcodes


def test_unklare_fortsetzung_gibt_keinen_vorschlag_und_eine_freigabe_schreibt_nichts(tmp_path):
    """Auftrag 004: Geht ein neuer Absatz vielleicht ueber den Seitenumbruch weiter, ist jeder fertige Vorschlag geraten (eine neue
    Zeile fuer den Rest waere falsch, wenn er zum Satz davor gehoert; im Entwicklungskorpus war er nur in 1 von 10 Faellen
    richtig). Die Aenderungen bleiben auf "pruefen", aber ohne Operationen."""
    alt = excel(tmp_path, [["1", "Eins", "Erster Absatz mit genug Text darin."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "{+1.1 Neu+}"),
                   ("p", "{+Die Anlage wird geprüft, sofern keine+}"), ("p", "Fußnote eins."),
                   ("p", "{+Störung ansteht.+}", 2), ("h", "2 Zwei", 2), ("p", "Zweiter Absatz mit genug Text darin.", 2)])
    plan, _ = lauf(tmp_path, pdf, alt)
    sec = next(a for a in plan.aenderungen if a.typ == "section_neu")
    folge = next(a for a in plan.aenderungen if a.typ == "absatz_neu" and "geprüft" in a.neu)
    for a in (sec, folge):
        assert a.status == PRUEFEN and a.ops == [] and "kein_vorschlag" in a.pruefcodes
        assert "sperre:fortsetzung_unklar" in a.pruefcodes
    frei = tmp_path / "frei.xlsx"
    uebernehmen(plan.aenderungen, alt, frei, {sec.id, folge.id}, absaetze=plan.absaetze)
    assert load_workbook(str(frei))["CDE"].max_row == load_workbook(str(alt))["CDE"].max_row      # keine neue Zeile


# ------------------------------------------------------------------------------------------------ Welle 3: markierter Leerraum
def test_markierter_leerraum_wird_d_bzw_a_lauf():
    """PdfToken.trenner_art (PDF-Leser, ADR 0006): ein gestrichenes bzw. unterstrichenes Leerzeichen gehoert nur zur alten bzw.
    neuen Fassung -- in den Laeufen ein d- bzw. a-Lauf (kein Trenner "s"), am markierten Nachbarwort."""
    t_alt = tok("alt", "", DEL)
    t_alt.trenner_art = "d"
    l = laeufe_aus_tokens([tok("Pumpe", "Pumpe", KEEP), t_alt, tok("Becken", "Becken", KEEP)])
    assert [(x.art, x.text) for x in l] == [("k", "Pumpe"), ("s", " "), ("d", "alt"), ("d", " "), ("k", "Becken")]
    assert zustand(l) == "Pumpe alt Becken" and zustand(l, {"G1"}) == "Pumpe Becken"
    assert b_ops(l, "Pumpe alt Becken")[1] == "Pumpe Becken"
    # eingefuegtes Leerzeichen zwischen zwei unveraenderten Teilen: nur die neue Fassung hat es
    links = tok("Pumpen", "Pumpen", KEEP)
    links.trenner_art = "a"
    l2 = laeufe_aus_tokens([links, tok("haus", "haus", KEEP)])
    assert [(x.art, x.text) for x in l2] == [("k", "Pumpen"), ("a", " "), ("k", "haus")]
    assert zustand(l2) == "Pumpenhaus" and zustand(l2, {"G1"}) == "Pumpen haus"
    assert b_ops(l2, "Pumpenhaus")[1] == "Pumpen haus"


def test_hochgestellte_einheiten_ziffern_gelten_als_ziffern():
    """Welle 3: "mm²" in der Excel und "mm2" aus dem PDF sind gleich ("²"/"³" werden "2"/"3", nicht entfernt); das Zeichen der
    Excel bleibt beim Ersetzen daneben stehen. Andere Hochzahlen (Fussnoten-Marken) zaehlen weiter nicht."""
    from src.gegenprobe import gleich
    assert gleich("Querschnitt 4 mm2", "Querschnitt 4 mm²") and gleich("Volumen 3 m3", "Volumen 3 m³")
    assert not gleich("Querschnitt 4 mm", "Querschnitt 4 mm²")             # "²" faellt nicht mehr weg
    assert gleich("Hinweis", "Hinweis¹")
    l = laeufe_aus_tokens(woerter("Querschnitt {-4-} {+6+} mm2 je Ader"))
    assert b_ops(l, "Querschnitt 4 mm² je Ader")[1] == "Querschnitt 6 mm² je Ader"


def test_excel_weicht_anderswo_ab_ergibt_fertigen_vorschlag(tmp_path):
    """Nicos Entscheidung (ADR 0006 Punkt 7): die Excel weicht im Absatz vom PDF ab (hier fehlt ein Leerzeichen direkt am
    geaenderten Wert). Die Stelle ist bekannt -> "pruefen" mit fertigem Vorschlag, der nur die markierte Stelle aendert und
    die Abweichung der Excel stehen laesst; nie "nicht uebernommen" ohne Grund."""
    alt = excel(tmp_path, [["1", "Eins", "Die Pumpe arbeitet bis [-35]degrees C im Freien ohne Heizung."], ZWEI])
    pdf = pdf_aus([("h", "1 Eins"), ("p", "Die Pumpe arbeitet bis {-[-35]-} {+[-40]+} degrees C im Freien ohne Heizung."),
                   ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")])
    plan, out = lauf(tmp_path, pdf, alt)
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "excel_weicht_ab" in a.pruefcodes and any(c.startswith("a1_") for c in a.pruefcodes)
    assert a.ziel == "CDE!C2" and [(o.start, o.ende, o.text) for o in a.ops] == [(23, 28, "[-40]")]
    frei = tmp_path / "frei.xlsx"
    uebernehmen(plan.aenderungen, alt, frei, {a.id}, absaetze=plan.absaetze)
    assert zelle(frei, 2, 3) == "Die Pumpe arbeitet bis [-40]degrees C im Freien ohne Heizung."


# ------------------------------------------------------------------------------------------------ Leerraum an Einfuegung/Streichung
def _word_laeufe(tmp_path, zeilen):
    """Ein Word-artiges Test-PDF (jede Zeile ein Eintrag, Umbruch zwischen den Zeilen; am Umbruch steht kein Leerzeichen im PDF)
    -> Laeufe des ersten Absatzes, wie `plan.absaetze` sie traegt."""
    from src.pdf_reader import analysiere_pdf
    from tests._word_pdf import WordDok
    d = WordDok()
    d.absatz(zeilen)
    block = analysiere_pdf(d.speichern(tmp_path / "w.pdf")).blocks[0]
    return laeufe_aus_tokens(block.tokens)


def _marke(laeufe, art):
    """Alle Laeufe der Art `art` als ein Text."""
    return "".join(l.text for l in laeufe if l.art == art)


@pytest.mark.legende_neu
@pytest.mark.parametrize("zeilen,art,markiert", [
    # Einfuegung mit Leerzeichen VORN (Word markiert das Leerzeichen vor dem neuen Wort mit)
    ([[("Die Anlage gilt", "k"), (" nachweislich", "n"), (" heute.", "k")]], "a", " nachweislich"),
    # ... mit Leerzeichen HINTEN
    ([[("Die Anlage gilt ", "k"), ("nachweislich ", "n"), ("heute.", "k")]], "a", "nachweislich "),
    # Streichung mit Leerzeichen davor
    ([[("Der Behaelter", "k"), (" erfolgen", "d"), (" halbjaehrlich.", "k")]], "d", " erfolgen"),
    # am Zeilenumbruch fehlt das Leerzeichen im PDF: Einfuegung am Zeilenanfang -> das Leerzeichen davor gehoert ihr
    ([[("Die Anlage gilt", "k")], [("nachweislich", "n"), (" heute.", "k")]], "a", " nachweislich"),
    # ... Streichung am Zeilenanfang
    ([[("Der Behaelter", "k")], [("erfolgen", "d"), (" halbjaehrlich.", "k")]], "d", " erfolgen"),
    # ... Einfuegung am Zeilenende, davor ein sichtbares unmarkiertes Leerzeichen -> das Leerzeichen dahinter gehoert ihr
    ([[("Die Anlage gilt ", "k"), ("nachweislich", "n")], [("heute.", "k")]], "a", "nachweislich "),
    # ... Satz am Absatzende, erstes Wort am Zeilenanfang
    ([[("Die Anlage gilt heute.", "k")], [("Neu ist dies.", "n")]], "a", " Neu ist dies."),
], ids=["vorn", "hinten", "streichung", "umbruch_vorn", "umbruch_streichung", "umbruch_hinten", "umbruch_satz_ende"])
def test_markierter_leerraum_gehoert_zur_einfuegung_bzw_streichung(tmp_path, zeilen, art, markiert):
    l = _word_laeufe(tmp_path, zeilen)
    assert _marke(l, art) == markiert
    assert _marke(l, "d" if art == "a" else "a") == ""


@pytest.mark.legende_neu
def test_ersetzung_am_zeilenumbruch_markiert_keinen_leerraum(tmp_path):
    """Wort ersetzt ("altneu", ohne Leerzeichen dazwischen) und das alte Wort steht am Zeilenanfang: Word markiert hier kein
    Leerzeichen, auch wenn das Zeichen am Umbruch im PDF fehlt."""
    l = _word_laeufe(tmp_path, [[("Der Behaelter", "k")], [("alt", "d"), ("neu", "n"), (" halbjaehrlich.", "k")]])
    assert _marke(l, "d") == "alt" and _marke(l, "a") == "neu"
    assert [(x.art, x.text) for x in l[:4]] == [("k", "Der"), ("s", " "), ("k", "Behaelter"), ("s", " ")]
    # "alt" und "neu" mit einem unmarkierten Leerzeichen dazwischen, der Umbruch liegt dort: das Leerzeichen bleibt unmarkiert
    l2 = _word_laeufe(tmp_path, [[("Der Behaelter ", "k"), ("alt", "d")], [("neu", "n"), (" halbjaehrlich.", "k")]])
    assert _marke(l2, "d") == "alt" and _marke(l2, "a") == "neu"


@pytest.mark.legende_neu
def test_ersetzung_mit_leerzeichen_am_umbruch_streicht_es_mit(tmp_path):
    """"alt " gestrichen und "neu " neu geschrieben, der Umbruch liegt dazwischen: das Leerzeichen hinter "alt" gehoert zur
    Streichung (sonst haette die neue Fassung ein Leerzeichen zu viel)."""
    l = _word_laeufe(tmp_path, [[("Der Behaelter ", "k"), ("alt", "d")], [("neu ", "n"), ("halbjaehrlich.", "k")]])
    assert _marke(l, "d") == "alt " and _marke(l, "a") == "neu "
    assert zustand(l, {"G1"}) == "Der Behaelter neu halbjaehrlich."


@pytest.mark.legende_neu
def test_mehrzeilige_einfuegung_bleibt_ein_markierter_lauf(tmp_path):
    """Eine eingefuegte Wortgruppe ueber den Zeilenumbruch: das Leerzeichen am Umbruch gehoert zur Einfuegung (ein Lauf)."""
    l = _word_laeufe(tmp_path, [[("Die Anlage gilt", "k"), (" ganz neue", "n")], [("Worte", "n"), (" heute.", "k")]])
    assert _marke(l, "a") == " ganz neue Worte" and _marke(l, "d") == ""
    assert zustand(l) == "Die Anlage gilt heute." and zustand(l, {"G1"}) == "Die Anlage gilt ganz neue Worte heute."
