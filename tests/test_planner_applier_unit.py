"""Unit-Tests K-PL (Planen / Anwenden / Protokoll): je eine Reparatur, nur erfundene Daten (Regel 16).

Die Red-Team-Tests (tests/redteam/test_pl_planer_applier.py) pruefen die ganzen Faelle. Hier stehen die kleinen
Bausteine und die Randfaelle, die dort nicht vorkommen.
"""
from __future__ import annotations

import copy

import pytest
from openpyxl import load_workbook

from src import config, runs
from src.applier import _wende_text_an, _wortgrenze, uebernehmen
from src.models import (NICHT, PRUEFEN, UEBERNOMMEN, WIRKUNG_ABWEICHUNG, WIRKUNG_KEINE, WIRKUNG_OK, Aenderung, Operation,
                        luecke)
from src.planner import neutext_probleme
from src.protokoll import bilanz, bilanz_satz, schreibe_xlsx
from tests.redteam._bau import (alle_texte, analysiere, baue_alt, blatt, lauf, schreibe_excel, uebernimm, zelle,
                                zeile_mit)


def _basis_excel(tmp_path, text="Das System muss laufen und stabil bleiben."):
    return schreibe_excel(tmp_path / "a.xlsx", [["Nr.", "Titel", "Inhalt"], ["1", "Eins", text],
                                                ["2", "Zwei", "Zweiter Absatz mit genug Text darin."]])


def _ae(aid, ops, status=UEBERNOMMEN, typ="text_ersetzt", alt="", neu="x", ziel="CDE!C2", **kw):
    return Aenderung(id=aid, typ=typ, seite=1, bbox=(0, 0, 1, 1), section="1", alt=alt, neu=neu, woerter=1,
                     status=status, konfidenz="hoch", ops=ops, ziel=ziel, **kw)


# ---------------------------------------------------------------------------------------------
# 1. Wirkung: nie "uebernommen" ohne nachweisbare Wirkung
# ---------------------------------------------------------------------------------------------
def test_wirkung_ok_wird_gesetzt_und_nachkontrolle_ok(tmp_path):
    p = _basis_excel(tmp_path)
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    a = _ae("A1", [rep], alt="System", neu="Gerät")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das Gerät muss laufen und stabil bleiben."
    assert (a.status, a.wirkung, a.nachkontrolle) == (UEBERNOMMEN, WIRKUNG_OK, "ok")


def test_operation_nur_mit_leerzeichen_hat_keine_wirkung_und_ist_nie_uebernommen(tmp_path):
    p = _basis_excel(tmp_path)
    ins = Operation("einfuegen", "CDE", 2, 3, 3, 3, " ", 1, aenderung="A1")
    a = _ae("A1", [ins], typ="text_eingefuegt", neu="")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.wirkung == WIRKUNG_KEINE
    assert a.status == NICHT and a.grund and a.hinweise
    assert a.nachkontrolle == "abweichung"


def test_aenderung_ohne_operationen_ist_nie_uebernommen(tmp_path):
    p = _basis_excel(tmp_path)
    a = _ae("A1", [], typ="absatz_neu")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == NICHT and a.wirkung == WIRKUNG_KEINE


def test_einfuegung_mitten_im_wort_ist_abweichung_und_pruefen(tmp_path):
    p = _basis_excel(tmp_path, "Das Gehäuse steht hier.")
    ins = Operation("einfuegen", "CDE", 2, 3, 8, 8, " s", 1, aenderung="A1")      # mitten in "Gehäuse"
    a = _ae("A1", [ins], typ="text_eingefuegt", neu="s")
    res = uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.wirkung == WIRKUNG_ABWEICHUNG and a.status == PRUEFEN and a.konfidenz == "niedrig"
    assert a.hinweise and res.warnungen


def test_wortgrenze():
    assert _wortgrenze("ab cd", 0) and _wortgrenze("ab cd", 2) and _wortgrenze("ab cd", 3) and _wortgrenze("ab cd", 5)
    assert not _wortgrenze("abcd", 2)


def test_zwei_aenderungen_in_einer_zelle_beide_mit_wirkung(tmp_path):
    p = _basis_excel(tmp_path)
    r1 = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    r2 = Operation("ersetzen", "CDE", 2, 3, 11, 15, "darf", 2, aenderung="A2", erwartet="muss")
    a1, a2 = _ae("A1", [r1]), _ae("A2", [r2])
    uebernehmen([a1, a2], p, tmp_path / "out.xlsx")
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das Gerät darf laufen und stabil bleiben."
    assert a1.wirkung == a2.wirkung == WIRKUNG_OK


def test_abgedeckte_aenderung_in_geleerter_zelle_bleibt_uebernommen(tmp_path):
    p = _basis_excel(tmp_path)
    leer = Operation("zelle_leeren", "CDE", 2, 3, aenderung="A1",
                     erwartet="Das System muss laufen und stabil bleiben.")
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 2, aenderung="A2", erwartet="System")
    a1 = _ae("A1", [leer], typ="absatz_geloescht", alt="Das System muss laufen und stabil bleiben.", neu="")
    a2 = _ae("A2", [rep])
    uebernehmen([a1, a2], p, tmp_path / "out.xlsx")
    assert zelle(tmp_path / "out.xlsx", 2, 3) == ""
    assert a1.wirkung == WIRKUNG_OK and a2.wirkung == WIRKUNG_OK


# ---------------------------------------------------------------------------------------------
# 2. Folgeabsaetze neuer Sections: Status und Freigabe gekoppelt
# ---------------------------------------------------------------------------------------------
EL_SEC = [("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "2 Zwei"),
          ("p", "Bekannter Absatz mit genug Text darin."),
          ("p", "Unbekannter Absatz der in der Excel gar nicht vorkommt."),
          ("h", "{+3 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}")]
EX_SEC = [e for e in EL_SEC if "Unbekannter" not in e[1]]
FOLGE = "Folgeabsatz der neuen Section steht hier."


def _sec(tmp_path, freigaben):
    alt = baue_alt(tmp_path, EX_SEC, "L1")
    l = analysiere(tmp_path, EL_SEC, alt)
    kopf = next(a for a in l.aen if a.typ == "section_neu")
    folge = next(a for a in l.aen if a.typ == "absatz_neu")
    return l, kopf, folge, {{"kopf": kopf.id, "folge": folge.id}[f] for f in freigaben}


def test_folgeabsatz_hat_markierung_statt_leerer_operationen(tmp_path):
    l, kopf, folge, _ = _sec(tmp_path, [])
    assert kopf.status == folge.status == PRUEFEN
    assert len(folge.ops) == 1 and folge.ops[0].abgedeckt and folge.ops[0].art == "zeile_einfuegen"
    assert folge.ops[0].seq == kopf.ops[0].seq and not folge.ops[0].zellen


@pytest.mark.parametrize("freigeben", [["kopf"], ["folge"], ["kopf", "folge"]])
def test_gruppe_ist_gemeinsam_wirksam(tmp_path, freigeben):
    l, kopf, folge, ids = _sec(tmp_path, freigeben)
    uebernimm(l, ids)
    assert kopf.status == folge.status == UEBERNOMMEN
    assert kopf.wirkung == folge.wirkung == WIRKUNG_OK
    assert sum(FOLGE in t for t in alle_texte(l.out)) == 1


def test_ohne_freigabe_bleibt_die_gruppe_unberuehrt(tmp_path):
    l, kopf, folge, ids = _sec(tmp_path, [])
    uebernimm(l, ids)
    assert kopf.status == folge.status == PRUEFEN
    assert not any(FOLGE in t for t in alle_texte(l.out))


def test_folgeabsatz_allein_uebernommen_ohne_section_wird_zurueckgestuft(tmp_path):
    """Alter/fremder Plan: Folgeabsatz 'uebernommen', Section nur 'pruefen' -> nichts wird geschrieben, kein falsches 'ok'."""
    l, kopf, folge, _ = _sec(tmp_path, [])
    folge.status = UEBERNOMMEN
    uebernimm(l, set())
    assert folge.status == PRUEFEN and folge.hinweise
    assert not any(FOLGE in t for t in alle_texte(l.out))


def test_schema1_folgeabsatz_ohne_markierung_wird_ueber_die_zielzeile_gekoppelt(tmp_path):
    l, kopf, folge, _ = _sec(tmp_path, ["kopf"])
    folge.ops = []                               # so sah ein Plan vor der Markierung aus
    uebernimm(l, {kopf.id})
    assert folge.status == UEBERNOMMEN and folge.wirkung == WIRKUNG_OK


def test_section_mit_steuerzeichen_nimmt_den_folgeabsatz_mit(tmp_path):
    l, kopf, folge, _ = _sec(tmp_path, ["kopf"])
    kopf.ops[0].zellen[2] = "Neue\x07 Section"
    uebernimm(l, {kopf.id})
    assert kopf.status == NICHT and folge.status == NICHT
    assert not any(FOLGE in t for t in alle_texte(l.out))


# ---------------------------------------------------------------------------------------------
# 3. Idempotenz
# ---------------------------------------------------------------------------------------------
def _zweimal(tmp_path, el, layout):
    l1 = lauf(tmp_path / "eins", el, layout)
    l2 = uebernimm(analysiere(tmp_path / "zwei", el, l1.out), name="out2.xlsx")
    return l1, l2


def test_zweiter_lauf_neue_section_wird_nicht_doppelt_angelegt(tmp_path):
    el = [("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "{+2 Neue Section+}"),
          ("p", "{+Folgeabsatz der neuen Section steht hier.+}"), ("h", "3 Drei"), ("p", "Dritter Absatz mit genug Text darin.")]
    for layout in ("L1", "L3"):
        l1, l2 = _zweimal(tmp_path / layout, el, layout)
        assert all(a.status == UEBERNOMMEN for a in l1.aen), layout
        assert all(a.status == NICHT and "steht schon" in a.grund for a in l2.aen), (layout, [(a.typ, a.status, a.grund) for a in l2.aen])
        assert sum("Neue Section" in t for t in alle_texte(l2.out)) == 1


def test_zweiter_lauf_ersetzen_sagt_dass_der_neue_text_schon_da_ist(tmp_path):
    el = [("h", "1 Eins"), ("p", "Das System {-muss-} {+soll+} die Daten speichern."), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    l1, l2 = _zweimal(tmp_path, el, "L1")
    assert "steht schon" in l2.aen[0].grund and l2.aen[0].status == NICHT


def test_kurzer_doppelter_absatz_ist_nur_ein_vorschlag(tmp_path):
    """Unter `planer.doppelt_min_woerter` Woerter gilt ein gleicher Absatz nicht sicher als Doppelung -> pruefen mit Hinweis."""
    el = [("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("p", "{+Ja gut.+}"), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    alt = baue_alt(tmp_path, [("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("p", "Ja gut."), ("h", "2 Zwei"),
                              ("p", "Zweiter Absatz mit genug Text darin.")], "L1")
    l = analysiere(tmp_path, el, alt)
    neu = next(a for a in l.aen if a.typ == "absatz_neu")
    assert neu.status == PRUEFEN and any("kurzer Absatz" in h for h in neu.hinweise)


def test_neuer_absatz_der_nur_woanders_vorkommt_wird_normal_geplant(tmp_path):
    """Gleicher Text in einer ANDEREN Section ist kein Duplikat."""
    el = [("h", "1 Eins"), ("p", "Erster Absatz mit genug Text darin."), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin."), ("p", "{+Erster Absatz mit genug Text darin.+}")]
    alt = baue_alt(tmp_path, [e for e in el if "{+" not in e[1]], "L3")
    l = analysiere(tmp_path, el, alt)
    neu = next(a for a in l.aen if a.typ == "absatz_neu")
    assert neu.status != NICHT


# ---------------------------------------------------------------------------------------------
# 4. Woerter ohne Normform ("•", ".")
# ---------------------------------------------------------------------------------------------
def _el_punkt(text):
    return [("h", "1 Eins"), ("p", text), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]


def test_gestrichener_punkt_am_ende_wird_entfernt(tmp_path):
    l = lauf(tmp_path, _el_punkt("Das System läuft stabil im Betrieb und liefert Daten {-•-}"), "L1",
             alt=baue_alt(tmp_path, _el_punkt("Das System läuft stabil im Betrieb und liefert Daten •"), "L1"))
    assert zelle(l.out, 2, 3) == "Das System läuft stabil im Betrieb und liefert Daten"
    assert l.aen[0].status == UEBERNOMMEN and l.aen[0].wirkung == WIRKUNG_OK


def test_gestrichenes_zeichen_das_in_der_excel_fehlt_wird_nicht_geschrieben(tmp_path):
    el = _el_punkt("Das System läuft stabil im Betrieb {-•-} und liefert Daten.")
    alt = baue_alt(tmp_path, _el_punkt("Das System läuft stabil im Betrieb und liefert Daten."), "L1")
    l = uebernimm(analysiere(tmp_path, el, alt))
    ae = l.aen[0]
    assert ae.status == NICHT and ae.grund and ae.wirkung == ""
    assert zelle(l.out, 2, 3) == "Das System läuft stabil im Betrieb und liefert Daten."


def test_gestrichener_punkt_samt_woertern_am_rand(tmp_path):
    """'{-• Das System-}': das Aufzaehlungszeichen vor dem ersten Wort wird mit entfernt."""
    el = _el_punkt("{-• Das System-} Die Anlage läuft stabil im Betrieb und liefert Daten.")
    alt = baue_alt(tmp_path, _el_punkt("• Das System Die Anlage läuft stabil im Betrieb und liefert Daten."), "L1")
    l = uebernimm(analysiere(tmp_path, el, alt))
    assert zelle(l.out, 2, 3) == "Die Anlage läuft stabil im Betrieb und liefert Daten."
    assert l.aen[0].wirkung == WIRKUNG_OK


def test_gestrichenes_zeichen_am_rand_fehlt_in_der_excel_gibt_nicht_uebernommen(tmp_path):
    el = _el_punkt("{-• Das System-} Die Anlage läuft stabil im Betrieb und liefert Daten.")
    alt = baue_alt(tmp_path, _el_punkt("Das System Die Anlage läuft stabil im Betrieb und liefert Daten."), "L1")
    l = analysiere(tmp_path, el, alt)
    assert l.aen[0].status == NICHT and "Zeichen" in l.aen[0].grund


def test_ersetztes_zeichen_am_absatzanfang_wird_ersetzt(tmp_path):
    el = _el_punkt("{-•-} {+-+} Das System läuft stabil im Betrieb und liefert Daten.")
    alt = baue_alt(tmp_path, _el_punkt("• Das System läuft stabil im Betrieb und liefert Daten."), "L1")
    l = uebernimm(analysiere(tmp_path, el, alt))
    assert zelle(l.out, 2, 3) == "- Das System läuft stabil im Betrieb und liefert Daten."


# ---------------------------------------------------------------------------------------------
# 5. Gefaehrlicher Neutext
# ---------------------------------------------------------------------------------------------
def test_neutext_probleme_erkennt_die_faelle():
    assert neutext_probleme(["Alles gut hier."])["verdaechtig"] == []
    assert neutext_probleme(["a\x07b"])["steuerzeichen"]
    assert neutext_probleme(["Teil (cid:12) da"])["verdaechtig"]
    assert neutext_probleme(["Die ﬁnale Prüfung"])["verdaechtig"]
    assert neutext_probleme(["Müll � hier"])["verdaechtig"]
    assert neutext_probleme(["=SUMME(A1)"])["formel"] and not neutext_probleme(["a =b"])["formel"]
    assert neutext_probleme(["\tTab und\nZeilenumbruch sind erlaubt"])["steuerzeichen"] is False


def test_formel_text_bleibt_als_text_mit_quote_prefix(tmp_path):
    p = _basis_excel(tmp_path)
    rep = Operation("ersetzen", "CDE", 2, 3, 0, 3, "=Das", 1, aenderung="A1", erwartet="Das")
    a = _ae("A1", [rep], alt="Das", neu="=Das")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    z = blatt(tmp_path / "out.xlsx").cell(2, 3)
    assert z.data_type == "s" and z.value.startswith("=Das System") and z.quotePrefix
    assert a.status == UEBERNOMMEN and a.wirkung == WIRKUNG_OK and a.hinweise


def test_formel_in_der_zelle_wird_nie_veraendert(tmp_path):
    p = schreibe_excel(tmp_path / "f.xlsx", [["Nr.", "Titel", "Inhalt"], ["1", "Eins", "=A1&B1"]])
    rep = Operation("ersetzen", "CDE", 2, 3, 0, 3, "x", 1, aenderung="A1", erwartet="=A1")
    a = _ae("A1", [rep])
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == NICHT and "Formel" in a.grund


def test_steuerzeichen_im_neutext_stuerzt_nicht_ab_applier(tmp_path):
    p = _basis_excel(tmp_path)
    ins = Operation("einfuegen", "CDE", 2, 3, 3, 3, " Neu\x07er Text", 1, aenderung="A1")
    a = _ae("A1", [ins], typ="text_eingefuegt")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == NICHT and "Steuerzeichen" in a.grund
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das System muss laufen und stabil bleiben."


def test_steuerzeichen_im_neutext_planer_nicht_uebernommen():
    """Der Planer schickt Steuerzeichen nie durch (hier ueber die Pruefung, die `_setze` benutzt)."""
    assert neutext_probleme(["x", "y\x00"], config.lade("zuordnung"))["steuerzeichen"]


@pytest.mark.parametrize("text", ["Teil (cid:12) hier", "Die ﬁnale Prüfung", "Kaputt � hier"])
def test_verdaechtiger_neutext_wird_im_applier_nicht_automatisch_geschrieben(tmp_path, text):
    p = _basis_excel(tmp_path)
    ins = Operation("einfuegen", "CDE", 2, 3, 3, 3, " " + text, 1, aenderung="A1")
    a = _ae("A1", [ins], typ="text_eingefuegt", neu=text)
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == PRUEFEN and a.hinweise
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das System muss laufen und stabil bleiben."


def test_verdaechtiger_neutext_wird_nach_menschlicher_freigabe_geschrieben(tmp_path):
    p = _basis_excel(tmp_path)
    ins = Operation("einfuegen", "CDE", 2, 3, 3, 3, " Die ﬁnale", 1, aenderung="A1")
    a = _ae("A1", [ins], status=PRUEFEN, typ="text_eingefuegt")
    uebernehmen([a], p, tmp_path / "out.xlsx", {"A1"})
    assert a.status == UEBERNOMMEN and a.bestaetigt and a.wirkung == WIRKUNG_OK
    assert "ﬁnale" in zelle(tmp_path / "out.xlsx", 2, 3)


# ---------------------------------------------------------------------------------------------
# 6. Zelle leeren nur, wenn ALLE Woerter gestrichen sind
# ---------------------------------------------------------------------------------------------
def test_zelle_leeren_mit_ungestrichenem_wort_wird_nicht_geleert(tmp_path):
    p = _basis_excel(tmp_path)
    leer = Operation("zelle_leeren", "CDE", 2, 3, aenderung="A1", erwartet="Das System muss laufen und stabil bleiben.")
    a = _ae("A1", [leer], typ="absatz_geloescht", alt="Das System muss laufen bleiben.", neu="")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == PRUEFEN and a.hinweise
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das System muss laufen und stabil bleiben."


def test_zelle_leeren_bei_geaenderter_zelle_ist_konflikt(tmp_path):
    p = _basis_excel(tmp_path)
    leer = Operation("zelle_leeren", "CDE", 2, 3, aenderung="A1", erwartet="Etwas ganz anderes")
    a = _ae("A1", [leer], typ="absatz_geloescht", alt="Etwas ganz anderes", neu="")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == NICHT and "geändert" in a.grund


def test_planer_leert_zelle_nur_bei_vollstaendig_gestrichenem_absatz(tmp_path):
    el = [("h", "1 Eins"), ("p", "{-Alle Racks werden vor dem Versand geprüft.-}"), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    alt = baue_alt(tmp_path, [("h", "1 Eins"), ("p", "Alle Racks werden vor dem Versand geprüft."), ("h", "2 Zwei"),
                              ("p", "Zweiter Absatz mit genug Text darin.")], "L3")
    l = uebernimm(analysiere(tmp_path, el, alt))
    ae = l.aen[0]
    assert ae.status == UEBERNOMMEN and ae.ops[0].art == "zelle_leeren" and ae.ops[0].erwartet
    assert zeile_mit(l.out, "Alle Racks") is None


def test_teil_loeschung_mit_zusatzwoertern_ist_nur_vorschlag_und_entfernt_gezielt(tmp_path):
    el = [("h", "1 Eins"), ("p", "{-Alle Racks werden jährlich vor dem Versand geprüft.-}"), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    alt = baue_alt(tmp_path, [("h", "1 Eins"), ("p", "Alle Racks werden zusätzlich jährlich noch vor dem Versand geprüft."),
                              ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")], "L3")
    l = analysiere(tmp_path, el, alt)
    ae = l.aen[0]
    assert ae.status == PRUEFEN and ae.hinweise and all(o.art == "ersetzen" for o in ae.ops)
    uebernimm(l, {ae.id})
    assert zelle(l.out, 3, 3) == "zusätzlich noch"      # nur die gestrichenen Woerter sind weg
    assert ae.wirkung == WIRKUNG_OK


# ---------------------------------------------------------------------------------------------
# 7. Reihenfolge bei gleicher Startstelle
# ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("seq_ins, seq_rep", [(1, 2), (2, 1), (5, 5)])
def test_gleicher_start_einfuegen_steht_vor_dem_ersetzten(seq_ins, seq_rep):
    ins = Operation("einfuegen", "CDE", 2, 3, 4, 4, "neue ", seq_ins, aenderung="A1")
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", seq_rep, aenderung="A2", erwartet="System")
    assert _wende_text_an("Das System muss laufen.", [ins, rep]) == "Das neue Gerät muss laufen."
    assert _wende_text_an("Das System muss laufen.", [rep, ins]) == "Das neue Gerät muss laufen."


def test_zwei_einfuegungen_gleicher_start_nach_seq():
    a = Operation("einfuegen", "CDE", 2, 3, 4, 4, "eins ", 1, aenderung="A1")
    b = Operation("einfuegen", "CDE", 2, 3, 4, 4, "zwei ", 2, aenderung="A2")
    assert _wende_text_an("Das System.", [b, a]) == "Das eins zwei System."


# ---------------------------------------------------------------------------------------------
# 8. Format neuer Zeilen
# ---------------------------------------------------------------------------------------------
def test_stilzeile_einer_neuen_absatzzeile_ist_nie_die_ueberschrift(tmp_path):
    el = [("h", "1 Eins"), ("p", "{+Neu direkt nach der Überschrift eins.+}"), ("p", "Absatz eins mit genug Text darin."),
          ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin.")]
    alt = baue_alt(tmp_path, [e for e in el if "{+" not in e[1]], "L3")
    l = analysiere(tmp_path, el, alt)
    op = l.aen[0].ops[0]
    kopf_zeilen = l.xl.layouts["CDE"].kopf_zeilen
    assert op.stil_von not in kopf_zeilen and l.xl.text("CDE", op.stil_von, 3).startswith("Absatz eins")


# ---------------------------------------------------------------------------------------------
# 9. Protokoll
# ---------------------------------------------------------------------------------------------
def _ae_einfach(status=UEBERNOMMEN, **kw):
    return Aenderung(id="A0001", typ="text_ersetzt", seite=1, bbox=(0, 0, 1, 1), section="1", alt="a", neu="b", woerter=2,
                     status=status, **kw)


def test_bilanz_satz_nur_bei_vollstaendig_und_stimmig():
    ok = bilanz([_ae_einfach()], 2, [])
    assert "nichts ging still verloren" in bilanz_satz(ok)[1]
    mit = bilanz([_ae_einfach()], 2, [], [luecke("scan_seite", 2, "Seite 2 enthält keinen lesbaren Text.")])
    assert mit["stimmt"] and not mit["vollstaendig"]
    assert "nichts ging still verloren" not in bilanz_satz(mit)[1] and "UNVOLLSTÄNDIG" in bilanz_satz(mit)[1]
    falsch = bilanz([_ae_einfach()], 5, [])
    assert "nichts ging still verloren" not in bilanz_satz(falsch)[1] and "NEIN" in bilanz_satz(falsch)[1]


def _zellen_text(pfad):
    wb = load_workbook(str(pfad))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def test_protokoll_xlsx_unvollstaendig_mit_luecken_hinweisen_und_wirkung(tmp_path):
    a = _ae_einfach(hinweise=["Bitte mit dem PDF vergleichen."], wirkung=WIRKUNG_OK, nachkontrolle="ok")
    lk = [luecke("scan_seite", 3, "Seite 3 enthält keinen lesbaren Text.", 0)]
    bil = bilanz([a], 2, [], lk)
    pfad = tmp_path / "p.xlsx"
    schreibe_xlsx(pfad, {"pdf_name": "a.pdf", "excel_name": "b.xlsx"}, [a], bil, [], [], None, lk)
    inhalt = _zellen_text(pfad)
    bilanz_zellen = [str(c) for row in inhalt["Bilanz"] for c in row if c is not None]
    assert not any("nichts ging still verloren" in c for c in bilanz_zellen)
    assert any("UNVOLLSTÄNDIG" in c for c in bilanz_zellen)
    assert any("Seite 3 enthält keinen lesbaren Text." in c for c in bilanz_zellen)
    kopf = inhalt["Änderungen"][0]
    zeile = dict(zip(kopf, inhalt["Änderungen"][1]))
    assert zeile["Wirkung"] == "ok" and "Bitte mit dem PDF vergleichen." in zeile["Warnungen"]


def test_protokoll_xlsx_vollstaendig_sagt_nichts_ging_verloren(tmp_path):
    a = _ae_einfach()
    bil = bilanz([a], 2, [])
    pfad = tmp_path / "p.xlsx"
    schreibe_xlsx(pfad, {}, [a], bil, [], [], None)
    zellen = [str(c) for row in _zellen_text(pfad)["Bilanz"] for c in row if c is not None]
    assert any("nichts ging still verloren" in c for c in zellen)


def test_protokoll_xlsx_ueberlebt_steuerzeichen_und_formeltext(tmp_path):
    a = _ae_einfach()
    a.neu, a.alt, a.grund = "Text\x07 mit Steuerzeichen", "=SUMME(A1)", "=1+1"
    bil = bilanz([a], 2, [])
    pfad = tmp_path / "p.xlsx"
    schreibe_xlsx(pfad, {"pdf_name": "=evil"}, [a], bil, [], ["=Warnung\x01"], None)
    wb = load_workbook(str(pfad))
    formeln = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert not formeln


def test_bilanz_zaehlt_wirkung_als_abweichung():
    a = _ae_einfach(status=NICHT, wirkung=WIRKUNG_KEINE)
    assert bilanz([a], 2, [])["nachkontrolle_abweichungen"] == 1


def test_markierung_gelb_liest_farbe_aus_der_konfiguration(tmp_path):
    p = _basis_excel(tmp_path)
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    cfg = copy.deepcopy(config.lade("zuordnung"))
    cfg["markierung"]["geaenderte_zellen"] = "gelb"
    cfg["anwenden"]["markierung_farbe"] = "AABBCC"
    uebernehmen([_ae("A1", [rep])], p, tmp_path / "out.xlsx", cfg=cfg)
    assert blatt(tmp_path / "out.xlsx").cell(2, 3).fill.fgColor.rgb.endswith("AABBCC")


def test_kommentar_markierung_liest_autor_und_ueberschreibt_keinen_kommentar(tmp_path):
    from openpyxl.comments import Comment
    p = _basis_excel(tmp_path)
    wb = load_workbook(str(p))
    wb["CDE"]["C2"].comment = Comment("Alter Kommentar von Hand", "Nico")
    wb.save(str(p))
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    cfg = copy.deepcopy(config.lade("zuordnung"))
    cfg["markierung"]["geaenderte_zellen"] = "kommentar"
    cfg["anwenden"]["kommentar_autor"] = "Test-Tool"
    uebernehmen([_ae("A1", [rep], alt="System", neu="Gerät")], p, tmp_path / "out.xlsx", cfg=cfg)
    k = blatt(tmp_path / "out.xlsx").cell(2, 3).comment
    assert "Alter Kommentar von Hand" in k.text and "PDF S.1" in k.text
    # Zelle ohne vorhandenen Kommentar: Autor kommt aus der Konfiguration
    rep2 = Operation("ersetzen", "CDE", 3, 3, 0, 7, "Dritter", 1, aenderung="A2", erwartet="Zweiter")
    uebernehmen([_ae("A2", [rep2], alt="Zweiter", neu="Dritter")], p, tmp_path / "out2.xlsx", cfg=cfg)
    assert blatt(tmp_path / "out2.xlsx").cell(3, 3).comment.author == "Test-Tool"


def test_einstellung_neue_section_gibt_es_nicht_mehr():
    assert "neue_section" not in config.lade("zuordnung")["einfuegen"]
    assert config.lade("zuordnung")["planer"]["doppelt_min_woerter"] >= 1


def test_kommentar_markierung_ueberlebt_steuerzeichen_im_alten_text(tmp_path):
    p = _basis_excel(tmp_path)
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    cfg = copy.deepcopy(config.lade("zuordnung"))
    cfg["markierung"]["geaenderte_zellen"] = "kommentar"
    a = _ae("A1", [rep], alt="Sys\x07tem", neu="Gerät")
    uebernehmen([a], p, tmp_path / "out.xlsx", cfg=cfg)
    assert a.status == UEBERNOMMEN and "\x07" not in blatt(tmp_path / "out.xlsx").cell(2, 3).comment.text


def test_lauf_schreibt_wirkung_und_hinweise_in_json_und_protokoll(tmp_path):
    import json

    from tests.fixtures.pdf_writer import baue_pdf
    ordner = tmp_path / "fx"
    ordner.mkdir()
    el = [("h", "1 Eins"), ("p", "Erster Absatz."), ("p", "{+Teil (cid:12) hier unten.+}"), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    alt = baue_alt(ordner, el, "L3")
    baue_pdf(ordner / "a.pdf", el, "farbe")
    run_id = runs.neuer_lauf(ordner / "a.pdf", alt)
    runs.analysiere(run_id)
    res = runs.uebernehmen_lauf(run_id, [])
    ae = res["aenderungen"][0]
    assert ae["status"] == PRUEFEN and ae["hinweise"] and ae["wirkung"] == ""
    erg = json.loads((runs.lauf_dir(run_id) / res["ordner"] / "ergebnis.json").read_text(encoding="utf-8"))
    assert erg["aenderungen"][0]["hinweise"] == ae["hinweise"]
    kopf = _zellen_text(runs.lauf_dir(run_id) / res["ordner"] / "aenderungsprotokoll.xlsx")["Änderungen"][0]
    assert "Wirkung" in kopf and "Warnungen" in kopf


# ---------------------------------------------------------------------------------------------
# Runde 2: Excel-Sicherheit (Zeilen einfuegen), Zahl-Zellen, "%g", Ueberlappung, Anker, xl-Parameter
# ---------------------------------------------------------------------------------------------
def _bearbeite(pfad, fn):
    wb = load_workbook(str(pfad))
    fn(wb)
    wb.save(str(pfad))


_EL_SECTION = [("h", "1 Eins"), ("p", "Absatz eins mit genug Text darin."),
               ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin."),
               ("h", "{+2.5 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}"),
               ("h", "3 Drei"), ("p", "Absatz drei mit genug Text darin."),
               ("h", "4 Vier"), ("p", "Absatz vier mit genug Text darin.")]


def _mit_formel(wb):
    wb["CDE"]["F2"] = "=C7"          # zeigt auf eine Zeile unterhalb der Einfuegestelle


def test_zeile_einfuegen_ohne_risiko_bleibt_automatisch(tmp_path):
    l = lauf(tmp_path, _EL_SECTION, "L1")
    neu = next(a for a in l.aen if a.typ == "section_neu")
    assert (neu.status, neu.konfidenz) == (UEBERNOMMEN, "hoch")
    assert zeile_mit(l.out, "Neue Section") is not None


def test_zeile_einfuegen_mit_formel_wird_pruefen_mit_saetzen(tmp_path):
    l = lauf(tmp_path, _EL_SECTION, "L1", nachher=lambda p: _bearbeite(p, _mit_formel))
    kopf = next(a for a in l.aen if a.typ == "section_neu")
    folge = next(a for a in l.aen if a.typ == "absatz_neu")
    assert kopf.status == PRUEFEN and kopf.konfidenz == "mittel"
    assert any("Formeln" in h for h in kopf.hinweise)
    assert "Zeile" in kopf.grund
    assert folge.status == PRUEFEN and folge.grund          # der Folgeabsatz geht nur zusammen mit der Section
    assert zeile_mit(l.out, "Neue Section") is None         # nichts wurde geschrieben
    assert l.bil["stimmt"]


def test_zeile_einfuegen_mit_formel_nach_freigabe_mit_warnung(tmp_path):
    alt = baue_alt(tmp_path, _EL_SECTION, "L1", nachher=lambda p: _bearbeite(p, _mit_formel))
    l0 = analysiere(tmp_path, _EL_SECTION, alt)
    kopf = next(a for a in l0.aen if a.typ == "section_neu")
    l = uebernimm(l0, {kopf.id})
    assert kopf.status == UEBERNOMMEN and kopf.bestaetigt and kopf.nachkontrolle == "ok"
    assert zeile_mit(l.out, "Neue Section") is not None
    assert any(kopf.id in w and "Formeln" in w for w in l.res.warnungen)
    assert blatt(l.out)["F2"].value == "=C7"            # Formeln werden nie veraendert


def test_applier_stuft_uebernommene_zeile_mit_risiko_zurueck(tmp_path):
    """Auch ein von Hand gebauter Plan ('uebernommen') darf bei Risiko nicht automatisch Zeilen einfuegen."""
    alt = baue_alt(tmp_path, _EL_SECTION, "L1", nachher=lambda p: _bearbeite(p, _mit_formel))
    plan = analysiere(tmp_path, _EL_SECTION, alt).plan
    for a in plan.aenderungen:                           # so tun, als haette der Planer nichts bemerkt
        a.status, a.konfidenz = UEBERNOMMEN, "hoch"
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, alt, out)
    assert all(a.status == PRUEFEN for a in plan.aenderungen if a.typ in ("section_neu", "absatz_neu"))
    assert zeile_mit(out, "Neue Section") is None


def test_versteckte_zeile_und_hyperlink_wandern_mit_der_zeile(tmp_path):
    def eig(wb):
        ws = wb["CDE"]
        z = next(c.row for row in ws.iter_rows() for c in row if c.value and "Absatz drei" in str(c.value))
        ws.row_dimensions[z].hidden = True
        ws.row_dimensions[z].outlineLevel = 1
        ws.cell(z, 3).hyperlink = "https://example.org/x"
    alt = baue_alt(tmp_path, _EL_SECTION, "L1", nachher=lambda p: _bearbeite(p, eig))
    l0 = analysiere(tmp_path, _EL_SECTION, alt)
    kopf = next(a for a in l0.aen if a.typ == "section_neu")
    l = uebernimm(l0, {kopf.id})                          # Freigabe: Risiko wird gemeldet, aber alles wandert mit
    ws = blatt(l.out)
    alt_r = zeile_mit(alt, "Absatz drei")
    neu_r = zeile_mit(l.out, "Absatz drei")
    assert alt_r is not None and neu_r == alt_r + 1
    assert ws.row_dimensions[neu_r].hidden and ws.row_dimensions[neu_r].outlineLevel == 1
    assert not ws.row_dimensions[alt_r].hidden
    assert ws.cell(neu_r, 3).hyperlink is not None and ws.cell(neu_r, 3).hyperlink.ref == f"C{neu_r}"
    assert ws.cell(alt_r, 3).hyperlink is None


def test_leere_zeile_wird_bei_risiko_nicht_entfernt(tmp_path):
    cfg = copy.deepcopy(config.lade("zuordnung"))
    cfg["geloeschtes"]["inhalt"] = "zeile_entfernen"
    p = schreibe_excel(tmp_path / "a.xlsx", [["Nr.", "Titel", "Inhalt"], [None, None, "Erster Absatz mit genug Text."],
                                             ["2", "Zwei", "Zweiter Absatz mit genug Text."],
                                             ["3", "Drei", "Dritter Absatz mit genug Text."]])
    _bearbeite(p, lambda wb: wb["CDE"].__setitem__("F4", "=C4"))          # Formel unterhalb der Loeschstelle
    op = Operation("zelle_leeren", "CDE", 2, 3, seq=1, aenderung="A1", erwartet="Erster Absatz mit genug Text.")
    a = _ae("A1", [op], typ="absatz_geloescht", alt="Erster Absatz mit genug Text.", neu="")
    res = uebernehmen([a], p, tmp_path / "out.xlsx", cfg=cfg)
    ws = blatt(tmp_path / "out.xlsx")
    assert ws.max_row == 4 and ws["C2"].value is None               # Zeile bleibt (leer) stehen
    assert a.status == UEBERNOMMEN and any("Zeile" in h for h in a.hinweise)
    assert res.warnungen


def test_ankerstufe_mittel_macht_neuen_absatz_zu_pruefen(tmp_path):
    ex = [("h", "1 Eins"),
          ("p", "Die Anlage wird vor dem Versand einzeln geprüft und danach dokumentiert; anschließend wird ein "
                "Prüfbericht erstellt und dem Kunden übergeben."),
          ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]
    el = [("h", "1 Eins"), ("p", "Die Anlage wird vor dem Versand einzeln geprüft und danach"),
          ("p", "{+Neuer Absatz nach dem Anfang.+}"), ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]
    l = uebernimm(analysiere(tmp_path, el, baue_alt(tmp_path, ex, "L1")))
    ae = next(a for a in l.aen if a.typ == "absatz_neu")
    assert ae.status == PRUEFEN and ae.konfidenz == "mittel"
    assert "unsicher" in ae.grund and ae.hinweise


def _zahlen_excel(pfad, fmt="0.0", wert=2.2):
    p = schreibe_excel(pfad, [["Nr.", "Titel", "Inhalt"], [1, "Eins", "Erster Abschnitt mit genug Text hier drin."],
                              [wert, "Zwei", "Zweiter Abschnitt mit genug Text hier drin."],
                              [3, "Drei", "Dritter Abschnitt mit genug Text hier drin."]])
    _bearbeite(p, lambda wb: setattr(wb["CDE"]["A3"], "number_format", fmt))
    return p


def test_zahl_zelle_wird_als_zahl_ersetzt_zahlenformat_bleibt(tmp_path):
    p = _zahlen_excel(tmp_path / "z.xlsx")
    op = Operation("ersetzen", "CDE", 3, 1, 0, 3, "2.3", 1, aenderung="A1", erwartet="2.2")
    a = _ae("A1", [op], alt="2.2", neu="2.3", ziel="CDE!A3")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    z = blatt(tmp_path / "out.xlsx")["A3"]
    assert (a.status, a.nachkontrolle) == (UEBERNOMMEN, "ok")
    assert z.data_type == "n" and z.value == 2.3 and z.number_format == "0.0"


def test_zahl_zelle_text_wird_pruefen_zahlenformat_geht_verloren(tmp_path):
    p = _zahlen_excel(tmp_path / "z.xlsx")
    op = Operation("ersetzen", "CDE", 3, 1, 0, 3, "zwei", 1, aenderung="A1", erwartet="2.2")
    a = _ae("A1", [op], alt="2.2", neu="zwei", ziel="CDE!A3")
    uebernehmen([a], p, tmp_path / "out.xlsx")
    assert a.status == PRUEFEN and "Zahlenformat geht verloren" in a.grund and a.hinweise
    assert blatt(tmp_path / "out.xlsx")["A3"].value == 2.2           # nichts geschrieben
    # nach Freigabe: Text, aber laute Warnung
    a2 = _ae("A1", [Operation("ersetzen", "CDE", 3, 1, 0, 3, "zwei", 1, aenderung="A1", erwartet="2.2")],
             status=PRUEFEN, alt="2.2", neu="zwei", ziel="CDE!A3")
    res = uebernehmen([a2], p, tmp_path / "out2.xlsx", {"A1"})
    assert a2.status == UEBERNOMMEN and blatt(tmp_path / "out2.xlsx")["A3"].value == "zwei"
    assert any("Zahlenformat" in w for w in res.warnungen)


def test_zahl_zelle_mit_zwei_nachkommastellen_nur_gleiche_anzeige(tmp_path):
    p = _zahlen_excel(tmp_path / "z.xlsx", fmt="0.00", wert=2.1)       # Anzeige "2.10"
    ok = _ae("A1", [Operation("ersetzen", "CDE", 3, 1, 0, 4, "2.20", 1, aenderung="A1", erwartet="2.10")], alt="2.10", neu="2.20")
    uebernehmen([ok], p, tmp_path / "out.xlsx")
    assert ok.status == UEBERNOMMEN and blatt(tmp_path / "out.xlsx")["A3"].value == 2.2
    abw = _ae("A2", [Operation("ersetzen", "CDE", 3, 1, 0, 4, "2.2", 1, aenderung="A2", erwartet="2.10")], alt="2.10", neu="2.2")
    uebernehmen([abw], p, tmp_path / "out2.xlsx")
    assert abw.status == PRUEFEN                                        # "2.2" sieht anders aus als das Format "0.00"


def test_text_einer_grossen_zahl_ohne_prozent_g(tmp_path):
    """'%g' machte aus 1234567.5 den Text '1.23457e+06'; die Stelle im Plan haette nie gepasst."""
    from src.applier import _text
    p = schreibe_excel(tmp_path / "g.xlsx", [["Nr.", "Titel", "Inhalt"], ["1", "Eins", 1234567.5]])
    assert _text(load_workbook(str(p))["CDE"]["C2"]) == "1234567.5"


def test_zwei_aufeinanderfolgende_gestrichene_absaetze_ueberlappen_nicht(tmp_path):
    text = "Zeile eins\nPunkt aaa bbb. Ccc ddd."
    p = schreibe_excel(tmp_path / "a.xlsx", [["Nr.", "Titel", "Inhalt"], ["1", "Eins", text]])
    # Jede Loeschung nimmt das Leerzeichen dazwischen mit (Stelle 25): [11, 26) und [25, 34)
    o1 = Operation("ersetzen", "CDE", 2, 3, 11, 26, "", 1, aenderung="A1", erwartet=text[11:26])
    o2 = Operation("ersetzen", "CDE", 2, 3, 25, 34, "", 2, aenderung="A2", erwartet=text[25:34])
    a1, a2 = _ae("A1", [o1], typ="text_geloescht", alt="Punkt aaa bbb."), _ae("A2", [o2], typ="text_geloescht", alt="Ccc ddd.")
    uebernehmen([a1, a2], p, tmp_path / "out.xlsx")
    assert (a1.status, a2.status) == (UEBERNOMMEN, UEBERNOMMEN)
    assert (a1.nachkontrolle, a2.nachkontrolle) == ("ok", "ok")
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Zeile eins"


def test_wende_text_an_loescht_gemeinsames_leerzeichen_nur_einmal():
    o1 = Operation("ersetzen", "CDE", 1, 1, 4, 8, "", 1, aenderung="A1")
    o2 = Operation("ersetzen", "CDE", 1, 1, 7, 12, "", 2, aenderung="A2")
    assert _wende_text_an("aaa bbb ccc ddd", [o1, o2]) == "aaa ddd"


def test_ueberschneidung_mit_buchstaben_bleibt_konflikt(tmp_path):
    p = schreibe_excel(tmp_path / "a.xlsx", [["Nr.", "Titel", "Inhalt"], ["1", "Eins", "aaa bbb ccc ddd"]])
    o1 = Operation("ersetzen", "CDE", 2, 3, 0, 9, "", 1, aenderung="A1", erwartet="aaa bbb c")
    o2 = Operation("ersetzen", "CDE", 2, 3, 7, 15, "", 2, aenderung="A2", erwartet=" ccc ddd")
    a1, a2 = _ae("A1", [o1], typ="text_geloescht"), _ae("A2", [o2], typ="text_geloescht")
    uebernehmen([a1, a2], p, tmp_path / "out.xlsx")
    assert a2.status == NICHT and "Konflikt" in a2.grund


def test_uebernehmen_mit_fertigem_modell_liest_die_excel_nicht_noch_einmal(tmp_path, monkeypatch):
    from src import applier
    from src.excel_model import ExcelModell
    p = _basis_excel(tmp_path)
    xl = ExcelModell(p)

    def nie(*a, **k):
        raise AssertionError("ExcelModell darf nicht noch einmal gebaut werden")
    monkeypatch.setattr(applier, "ExcelModell", nie)
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 1, aenderung="A1", erwartet="System")
    a = _ae("A1", [rep], alt="System", neu="Gerät")
    uebernehmen([a], p, tmp_path / "out.xlsx", xl=xl)
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das Gerät muss laufen und stabil bleiben."
