"""Tests fuer Generator v3 `word_markup` (Korpus nachbau_word) und die Mess-Erweiterungen von Auftrag 003:
Legende/nachher-Format je Korpus, Format-Treue, Rueckbau-Messung (messen --nachher).

Alles ist erfunden (Regel 16); geschrieben wird nur in tmp_path. Die Paare sind klein (40-60 Zeilen), damit die Tests schnell bleiben.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

import pdfplumber
import pytest
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

from src import config, messen
from src.models import PRUEFEN, UEBERNOMMEN
from tests.fixtures import generator, pdf_word as pw, word_markup as wm
from tests.fixtures.generator import erzeuge_paar

ROOT = Path(__file__).resolve().parent.parent
KLEIN = {"basis": "word_markup", "name": "klein", "zeilen": (45, 60)}
SEED = 371                       # Entwicklungs-Seed ausserhalb des Korpus 301-360 (Regel 11)


def _nahe(a, b, tol: float = 0.01) -> bool:
    return a is not None and len(a) == len(b) and all(abs(x - y) <= tol for x, y in zip(a, b))


def _laeufe(zelle) -> list[tuple[str, str]]:
    """[(Klasse, Text)] einer Zelle: alt = FF00B0F0 + Strich, neu = FFFF0000, normal sonst (Rich Text oder Zellschrift)."""
    def klasse(f) -> str:
        if f is None:
            return "normal"
        rgb = getattr(getattr(f, "color", None), "rgb", None)
        rgb = rgb[-6:].upper() if isinstance(rgb, str) else ""
        if getattr(f, "strike", None) and rgb == "00B0F0":
            return "alt"
        if not getattr(f, "strike", None) and rgb == "FF0000":
            return "neu"
        return "normal"
    v = zelle.value
    if v is None:
        return []
    if isinstance(v, CellRichText):
        return [(klasse(b.font), b.text) if isinstance(b, TextBlock) else (klasse(zelle.font), str(b)) for b in v]
    return [(klasse(zelle.font), str(v))]


def _ansicht(laeufe, welche: str) -> str:
    return "".join(t for k, t in laeufe if k in ("normal", welche))


@pytest.fixture(scope="module")
def paar(tmp_path_factory) -> Path:
    """Ein kleines Paar in einem Korpus-Ordner mit KORPUS.json (Legende auto, nachher markiert) -- so liest messen die Einstellungen
    wie im echten Korpus (und unabhaengig von der Test-Weiche in conftest.py)."""
    ziel = tmp_path_factory.mktemp("nachbau") / "nachbau_word"
    p = erzeuge_paar(SEED, {**KLEIN, "leiste": True, "zwei_autoren": True, "alte_markierung": True}, ziel, "nachbau_word")
    (ziel / "KORPUS.json").write_text(json.dumps({"generator_version": "3", "korpus": "nachbau_word", "seeds": [SEED],
                                                  **generator.korpus_einstellungen("nachbau_word")}), encoding="utf-8")
    return p


@pytest.fixture(scope="module")
def erwartet(paar) -> dict:
    return json.loads((paar / "erwartet.json").read_text(encoding="utf-8"))


# ====================================================================== Generator: Determinismus und Weichen
def test_determinismus_gleicher_seed_bytegleiche_dateien(tmp_path):
    a = erzeuge_paar(SEED + 1, KLEIN, tmp_path / "a", "nachbau_word")
    b = erzeuge_paar(SEED + 1, KLEIN, tmp_path / "b", "nachbau_word")
    for n in ("aenderungen.pdf", "vorher.xlsx", "nachher.xlsx", "erwartet.json"):
        assert (a / n).read_bytes() == (b / n).read_bytes(), n
    c = erzeuge_paar(SEED + 2, KLEIN, tmp_path / "c", "nachbau_word")
    assert (c / "vorher.xlsx").read_bytes() != (a / "vorher.xlsx").read_bytes()
    with pytest.raises(FileExistsError):                                  # append-only (Regel 29)
        erzeuge_paar(SEED + 1, KLEIN, tmp_path / "a", "nachbau_word")


def test_weichen_profil_korpus_erzeuger(tmp_path):
    with pytest.raises(ValueError):
        erzeuge_paar(SEED, "word_markup", tmp_path / "x", "entwicklung")      # v3 nur im Korpus nachbau_word
    with pytest.raises(ValueError):
        erzeuge_paar(SEED, "standard", tmp_path / "x", "nachbau_word")       # und umgekehrt
    with pytest.raises(ValueError):
        erzeuge_paar(SEED, "word_markup", tmp_path / "x", "nachbau_word", erzeuger="libreoffice")
    with pytest.raises(ValueError):
        erzeuge_paar(1000, "word_markup", tmp_path / "x", "nachbau_word")    # Regel 11: Entwicklungs-Seeds < 1000
    assert not (tmp_path / "x").exists() or not any((tmp_path / "x").iterdir())
    assert generator.version_fuer("nachbau_word") == "3" and generator.version_fuer("entwicklung") == "2"
    assert generator.korpus_einstellungen("nachbau_word") == {"legende": "auto", "nachher_format": "markiert"}
    assert generator.korpus_einstellungen("ex_test") == {"legende": "alt", "nachher_format": "sauber"}


def test_skript_korpus_nachbau_word(tmp_path):
    def skript(*args):
        return subprocess.run([sys.executable, str(ROOT / "scripts" / "erzeuge_paare.py"), *args], capture_output=True, text=True,
                              cwd=str(ROOT), timeout=180)
    assert skript("--korpus", "nachbau_word", "--seeds", "999-1000", "--ziel", str(tmp_path)).returncode == 2
    assert skript("--korpus", "nachbau_word", "--seeds", "373", "--erzeuger", "wechselnd", "--ziel", str(tmp_path)).returncode == 2
    assert not (tmp_path / "nachbau_word").exists()
    r = skript("--korpus", "nachbau_word", "--seeds", "373", "--ziel", str(tmp_path))
    assert r.returncode == 0, r.stderr
    k = json.loads((tmp_path / "nachbau_word" / "KORPUS.json").read_text(encoding="utf-8"))
    assert k["generator_version"] == "3" and k["legende"] == "auto" and k["nachher_format"] == "markiert" and k["seeds"] == [373]
    e = json.loads((tmp_path / "nachbau_word" / "synth_0373_word_markup" / "erwartet.json").read_text(encoding="utf-8"))
    assert e["generator_version"] == "3" and e["profil"] == "word_markup" and 300 <= e["excel"]["zeilen_vorher"] <= 560
    r = skript("--korpus", "nachbau_word", "--seeds", "373", "--ziel", str(tmp_path))
    assert r.returncode == 0 and "uebersprungen" in r.stdout


# ====================================================================== PDF-Struktur
def test_pdf_struktur_wie_word_markup(paar, erwartet):
    spez = None
    with pdfplumber.open(paar / "aenderungen.pdf") as pdf:
        assert len(pdf.pages) == erwartet["pdf_info"]["seiten"] >= 4
        kopf: set[str] = set()
        farben: Counter = Counter()
        linien_geprueft = 0
        for nr, pg in enumerate(pdf.pages, 1):
            assert (pg.width, pg.height) == (612, 792) and not pg.annots                     # keine PDF-Annotationen
            leiste = [r for r in pg.rects if _nahe(r["non_stroking_color"], pw.LEISTE_GRAU)]
            assert leiste and abs(leiste[0]["x0"] - 415) < 0.5 and abs(leiste[0]["x1"] - 612) < 0.5
            assert abs(leiste[0]["top"] - 95) < 0.5 and abs(leiste[0]["bottom"] - 696) < 0.5
            assert not [c for c in pg.chars if 396 < c["x0"] < 414]                         # Text bis ~395, Leiste ab 415
            for c in pg.chars:                                                               # Sprechblasen-Text in der Leiste
                if c["x0"] >= 415:
                    assert c["x0"] >= 433 and abs(c["size"] - wm.pw.BLASE_SCHRIFT) < 0.05 and "Helvetica" in c["fontname"]
            kopf.add("".join(c["text"] for c in sorted(pg.chars, key=lambda c: (round(c["top"]), c["x0"]))
                             if 120 < c["top"] < 148 and c["x0"] < 400))
            fuss = "".join(c["text"] for c in pg.chars if 625 < c["top"] < 665 and c["x0"] < 400)
            m = re.search(r"(\d\d \d\d \d\d) - (\d+)$", fuss)
            assert m and int(m.group(2)) == nr
            spez = spez or m.group(1)
            balken = [r for r in pg.rects if abs(r["x0"] - pw.BALKEN_X) < 0.05]
            assert all(abs(r["width"] - pw.BALKEN_B) < 0.02 and _nahe(r["non_stroking_color"], pw.SCHWARZ) for r in balken)
            for c in pg.chars:
                f = c["non_stroking_color"]
                if c["x0"] >= 415 or c["text"].isspace():
                    continue
                for name, farbe in (("rot", pw.WORD_ROT), ("blau", pw.WORD_BLAU), ("gruen", pw.WORD_GRUEN)):
                    if not _nahe(f, farbe):
                        continue
                    farben[name] += 1
                    mitte = (c["x0"] + c["x1"]) / 2
                    h = c["bottom"] - c["top"]
                    lagen = [((r["top"] + r["bottom"]) / 2 - c["top"]) / h for r in pg.rects
                             if _nahe(r["non_stroking_color"], farbe) and r["x0"] - 0.01 <= mitte <= r["x1"] + 0.01
                             and c["top"] - 1 < r["top"] < c["bottom"] + 1]
                    dicken = [r["height"] for r in pg.rects
                              if _nahe(r["non_stroking_color"], farbe) and r["x0"] - 0.01 <= mitte <= r["x1"] + 0.01
                              and c["top"] - 1 < r["top"] < c["bottom"] + 1]
                    # neu = Unterstreichung (~90 % der Zeichenhoehe, 0,36 pt), entfaellt = Durchstreichung (50-60 %, 0,46 pt)
                    assert len(lagen) == 1, (nr, c["text"])
                    assert 0.85 <= lagen[0] <= 0.95 and abs(dicken[0] - 0.36) < 0.02 or \
                        0.5 <= lagen[0] <= 0.6 and abs(dicken[0] - 0.46) < 0.02
                    linien_geprueft += 1
        assert len(kopf) == 1                                           # Kopfzeile auf jeder Seite gleich
        assert farben["rot"] and farben["blau"] and farben["gruen"] and linien_geprueft > 50
        seite = pdf.pages[-1]
        alle = [o for p in pdf.pages for o in p.curves]
        gestrichelt = [o for o in alle if o.get("dash") and abs(o["linewidth"] - pw.LINIE_STAERKE) < 0.01]
        assert gestrichelt and all(abs(o["x1"] - pw.BLASE_X0) < 0.5 for o in gestrichelt)
        blasen = [o for o in alle if o["fill"] and o["stroke"] and _nahe(o["non_stroking_color"], pw.WEISS)]
        assert blasen and all(o["x0"] >= 432.9 for o in blasen)
        gelb = [r for p in pdf.pages for r in p.rects if _nahe(r["non_stroking_color"], pw.GELB)]
        assert gelb and all(abs(r["height"] - pw.GELB_HOEHE) < 0.05 for r in gelb)
        text = "".join(c["text"] for p in pdf.pages for c in p.chars if c["x0"] >= 415)
        assert "formatiert:" in text or "Formatiert:" in text or "Kommentiert" in text
        assert seite.chars
    assert spez and erwartet["excel"]["datenstart"] == 8


def test_pdf_nur_inline_ohne_leiste(tmp_path):
    p = erzeuge_paar(SEED + 3, {**KLEIN, "leiste": False}, tmp_path, "nachbau_word")
    e = json.loads((p / "erwartet.json").read_text(encoding="utf-8"))
    assert "nur_inline" in e["stoerer"] and not e["pdf_info"]["leiste"] and e["pdf_info"]["sprechblasen"] == 0
    with pdfplumber.open(p / "aenderungen.pdf") as pdf:
        for pg in pdf.pages:
            assert not [r for r in pg.rects if _nahe(r["non_stroking_color"], pw.LEISTE_GRAU)]
            assert not [c for c in pg.chars if c["x0"] >= 400] and not [o for o in pg.curves if o.get("dash")]


# ====================================================================== Excel-Struktur
def test_excel_register_struktur_und_zusatzteile(paar, erwartet):
    for name in ("vorher.xlsx", "nachher.xlsx"):
        wb = load_workbook(paar / name, rich_text=True)
        assert wb.sheetnames[:3] == ["Hinweise", "Deckblatt", "Unterschriften"] and wb.sheetnames[4] == "Listen"
        assert wb["Listen"].sheet_state == "hidden"
        ws = wb[erwartet["hauptblatt"]]
        assert wb.sheetnames.index(ws.title) + 1 == wm.HAUPTBLATT_NR
        verbund = {str(r) for r in ws.merged_cells.ranges}
        assert {"A1:A4", "B1:B4", "C1:C4", "D1:F4", "G1:G4", "H1:P4", "G5:G6", "H5:H6", "I5:I6"} <= verbund
        assert str(ws["U2"].value).startswith("=COUNTIF(H:H") and str(ws["V2"].value).startswith("=COUNTIFS(H:H")
        assert ws["A7"].value == "0" and ws["C7"].value == "Rev 0." and ws["E6"].value and ws["F6"].value
        letzte = ws.max_row
        n = erwartet["excel"]["zeilen_vorher" if name == "vorher.xlsx" else "zeilen_nachher"]
        assert letzte == wm.DATEN_START - 1 + n
        assert ws.auto_filter.ref == f"A6:AB{letzte}" and ws.print_title_rows == "$5:$6"
        assert ws.print_area.endswith(f"$A$5:$P${letzte}")
        cf = {str(c.sqref) for c in ws.conditional_formatting}
        assert "R5 R8:R1048576" in cf and "V2:V4" in cf and f"O8:O{letzte}" in cf
        assert any(len(s.split()) == 2 and s.startswith("F") for s in cf)                  # zerstueckelter Bereich
        dv = {str(d.sqref) for d in ws.data_validations.dataValidation}
        assert f"J7:N{letzte} K{letzte + 6}" in dv
        assert ws.protection.formatCells is False
        labels = 0
        for r in range(wm.DATEN_START, letzte + 1):
            g = ws.cell(r, 7)
            text = "".join(t for _k, t in _laeufe(g))
            assert text and ws.cell(r, 2).value and ws.cell(r, 5).value
            if re.match(r"^(\d{1,2}[.)]|[a-z]{1,2}[.)])(\t|\xa0{13,19} )", text):
                labels += 1
                assert ws.cell(r, 6).value                                                # Listenpunkt: F = Sub-Section
            if ws.cell(r, 9).value == wm.KENNWORT:                                         # Ueberschriftzeile: fett, ohne Status
                assert ws.cell(r, 8).value is None
                assert all(getattr(b.font, "b", None) for b in g.value if isinstance(b, TextBlock)) \
                    if isinstance(g.value, CellRichText) else g.font.b
        assert labels >= 5
        wb.close()
        with zipfile.ZipFile(paar / name) as z:
            teile = set(z.namelist())
            assert set(erwartet["excel"]["zusatzteile"]) <= teile
            ct = z.read("[Content_Types].xml").decode()
            for teil in ("/xl/sharedStrings.xml", "/xl/calcChain.xml", "/xl/drawings/drawing1.xml", "/customXml/itemProps1.xml",
                         "/docMetadata/LabelInfo.xml"):
                assert f'PartName="{teil}"' in ct
            assert 'Extension="png"' in ct and 'Extension="bin"' in ct
            assert "docMetadata/LabelInfo.xml" in z.read("_rels/.rels").decode()
            wbr = z.read("xl/_rels/workbook.xml.rels").decode()
            assert "../customXml/item1.xml" in wbr and "calcChain.xml" in wbr and "sharedStrings.xml" in wbr
            blatt = next(n for n in teile if n.startswith("xl/worksheets/sheet4"))
            x = z.read(blatt).decode()
            assert 'inlineStr' not in x and '<sheetProtection formatCells="0"/>' in x and '<drawing r:id=' in x
            assert 'r:id="rIdP1"' in x
            zeichnung = z.read("xl/drawings/drawing1.xml").decode()
            assert "<xdr:sp " in zeichnung and "<xdr:pic>" in zeichnung and "<xdr:row>2</xdr:row>" in zeichnung
            assert "<xdr:row>6</xdr:row>" in zeichnung
            assert 'r="U2"' in z.read("xl/calcChain.xml").decode()


# ====================================================================== nachher-Markierung und erwartet.json
def test_nachher_markierung_wie_nico(paar, erwartet):
    v = load_workbook(paar / "vorher.xlsx", rich_text=True)[erwartet["hauptblatt"]]
    n = load_workbook(paar / "nachher.xlsx", rich_text=True)[erwartet["hauptblatt"]]
    cde = [a for a in erwartet["aenderungen"] if a["bereich"] == "cde"]
    arten = Counter(a["art"] for a in cde)
    assert arten["absatz_neu"] + arten["subsection_neu"] and arten["absatz_geloescht"] and arten["listennummer"]
    direkt_hintereinander = 0
    for a in cde:
        g = n.cell(a["nachher_zeile"], 7)
        laeufe = _laeufe(g)
        assert any(k != "normal" for k, _t in laeufe), a
        if a["typ"] == "absatz_neu":
            assert {k for k, _t in laeufe} == {"neu"}
            continue
        if a["typ"] == "absatz_geloescht":
            assert {k for k, _t in laeufe} == {"alt"}
        # alt-Ansicht = vorher (nichts verloren) -- auch wenn die Excel anderswo vom PDF abweicht: die Abweichung bleibt stehen
        assert _ansicht(laeufe, "alt") == v.cell(a["zeile"], 7).value
        if a["soll_status"] == "pruefen":
            assert a["vorschlag"] == [list(x) for x in _zus(laeufe)]                   # fertiger Vorschlag = Zelle in nachher
        for (k1, _t1), (k2, _t2) in zip(laeufe, laeufe[1:]):
            direkt_hintereinander += k1 == "alt" and k2 == "neu"
        if isinstance(g.value, CellRichText):                                            # Grundschrift in jedem Lauf
            for b in g.value:
                assert isinstance(b, TextBlock) and b.font.rFont == "Arial" and float(b.font.sz) == 11
    assert direkt_hintereinander >= 1                                                    # kein Trenner zwischen alt und neu
    # Markierungen frueherer Runden in C/D bleiben unangetastet (und sind keine Markierung dieser Runde)
    frueher = [r for r in range(wm.DATEN_START, v.max_row + 1) if isinstance(v.cell(r, 3).value, CellRichText)]
    assert frueher
    for r in frueher:
        farben = {b.font.color.rgb for b in v.cell(r, 3).value if isinstance(b, TextBlock)}
        assert farben == {wm.FRUEHER_ALT, wm.FRUEHER_NEU} and all(k == "normal" for k, _t in _laeufe(v.cell(r, 3)))


def _zus(laeufe):
    out: list[tuple[str, str]] = []
    for k, t in laeufe:
        if out and out[-1][0] == k:
            out[-1] = (k, out[-1][1] + t)
        else:
            out.append((k, t))
    return out


def test_erwartet_konsistent(paar, erwartet):
    assert erwartet["generator_version"] == "3" and erwartet["korpus"] == "nachbau_word"
    assert erwartet["legende"] == "auto" and erwartet["nachher_format"] == "markiert" and erwartet["luecken"] == []
    assert erwartet["woerter_markiert"] > 0 and erwartet["format"] == 1
    for a in erwartet["aenderungen"]:
        assert a["typ"] in generator.TYPEN and a["soll_status"] in ("uebernommen", "pruefen", "ausserhalb")
        assert (a["bereich"] == "ausserhalb") == (a["soll_status"] == "ausserhalb")
        if a["bereich"] == "ausserhalb":
            assert a["blatt"] == 0 and a["zeile"] == 0
        else:
            assert a["blatt"] == wm.HAUPTBLATT_NR and a["spalte"] == 7 and a["zeile"] >= wm.DATEN_START - 1
            assert ("vorschlag" in a) == (a["soll_status"] == "pruefen")
    assert {"kopf_fuss", "vorspann"} <= {a["art"] for a in erwartet["aenderungen"] if a["bereich"] == "ausserhalb"}
    assert any(a["soll_status"] == "pruefen" for a in erwartet["aenderungen"])          # Excel weicht anderswo ab
    v = load_workbook(paar / "vorher.xlsx")[erwartet["hauptblatt"]]
    n = load_workbook(paar / "nachher.xlsx", rich_text=True)[erwartet["hauptblatt"]]
    for x in erwartet["abweichungen_ohne_aenderung"]:                                   # unveraendert, Abweichung bleibt
        assert any("".join(t for _k, t in _laeufe(n.cell(r, 7))) == v.cell(x["zeile"], 7).value
                   for r in range(x["zeile"], x["zeile"] + 12))


# ====================================================================== messen: Einstellungen je Korpus
def test_korpus_einstellungen_und_konfig_ueberschreibung(tmp_path):
    assert messen.korpus_einstellungen(tmp_path, "entwicklung") == {"legende": "alt", "nachher_format": "sauber"}
    assert messen.korpus_einstellungen(tmp_path, "nachbau_word") == {"legende": "auto", "nachher_format": "markiert"}
    assert messen.korpus_einstellungen(tmp_path) == {}
    (tmp_path / "KORPUS.json").write_text(json.dumps({"generator_version": "2"}), encoding="utf-8")
    assert messen.korpus_einstellungen(tmp_path, "nachbau_word") == {"legende": "alt", "nachher_format": "sauber"}
    (tmp_path / "KORPUS.json").write_text(json.dumps({"legende": "linie", "nachher_format": "markiert"}), encoding="utf-8")
    assert messen.korpus_einstellungen(tmp_path) == {"legende": "linie", "nachher_format": "markiert"}
    vorher = (config.lade("farben")["legende"]["modus"], config.lade("zuordnung")["markierung"]["aenderungen"])
    with messen.konfig_ueberschreibung({"legende": "auto", "nachher_format": "markiert"}):
        assert config.lade("farben")["legende"]["modus"] == "auto"
        assert config.lade("zuordnung")["markierung"]["aenderungen"] == "farbe"
        assert messen._wirksam() == {"legende": "auto", "nachher_format": "markiert"}
    assert (config.lade("farben")["legende"]["modus"], config.lade("zuordnung")["markierung"]["aenderungen"]) == vorher
    with pytest.raises(RuntimeError):
        with messen.konfig_ueberschreibung({"nachher_format": "sauber"}):
            raise RuntimeError
    assert config.lade("zuordnung")["markierung"]["aenderungen"] == vorher[1]


# ====================================================================== messen: Format-Treue (kleine, von Hand gebaute Zellen)
def _ae(id_, zeile, alt="", neu="", status=UEBERNOMMEN):
    op = SimpleNamespace(blatt="CDE", art="ersetzen", zeile=zeile, spalte=2, zellen={})
    return SimpleNamespace(id=id_, status=status, ops=[op], ziel="", alt=alt, neu=neu, pruefcodes=[], nachkontrolle="",
                           typ="text_ersetzt", grund="g", hinweise=[])


def _auswertung(v, n, e, aend):
    plan = SimpleNamespace(layout=[{"blatt": "CDE", "inhalt_spalte": "B", "section_spalte": "A"}], luecken=[])
    bil = {"stimmt": True, "vollstaendig": True, "woerter_im_pdf": 0}
    aw = messen._Auswertung([v], [n], [e], None, ["CDE"], aend, plan, SimpleNamespace(blocks=[]), bil, bil, True, None,
                            {a.id: a.status for a in aend}, None, None, markiert=True)
    return aw.auswerten()


F = messen._FText


def _ft(*teile: tuple[str, str]) -> "messen._FText":
    return F("".join(t for _k, t in teile), "".join(k * len(t) for k, t in teile))

VORHER = [["Nr", "Text"], ["1", "Der Quokka frisst Gras am Morgen."], ["2", "Der Flamingo steht im Teich."],
          ["3", "Dieser Absatz entfällt ganz."]]
SOLL_2 = F("Der Quokka frisst GrasLaub am Morgen.", "-" * 18 + "aaaa" + "nnnn" + "-" * 11)
SOLL_3 = F("Der Flamingo steht im kalten Teich.", "-" * 22 + "nnnnnnn" + "-" * 6)
SOLL_4 = F("Dieser Absatz entfällt ganz.", "a" * 28)
NACHHER = [VORHER[0], ["1", SOLL_2], ["2", SOLL_3], ["3", SOLL_4]]


def test_format_treue_richtig_markiert():
    aend = [_ae("a1", 2, "Gras", "Laub"), _ae("a2", 3, "", "kalten"), _ae("a3", 4, "Dieser Absatz entfällt ganz.", "")]
    r = _auswertung(VORHER, NACHHER, NACHHER, aend)
    assert r["soll_zellen"] == 3 and r["erreicht_auto"] == 3                   # auch der nur durchgestrichene Absatz ist Soll
    assert r["stille_fehler"] == 0 and r["format_treue_zellen_geprueft"] == 3 and r["format_treue_zellen_gleich"] == 3
    assert r["format_treue_zeichen"] == r["format_treue_zeichen_gleich"] == len(SOLL_2) + len(SOLL_3) + len(SOLL_4)


def test_format_treue_falsche_klasse_ist_stiller_fehler():
    aend = [_ae("a1", 2, "Gras", "Laub"), _ae("a2", 3, "", "kalten"), _ae("a3", 4, "Dieser Absatz entfällt ganz.", "")]
    falsch_2 = F(str(SOLL_2), "-" * 18 + "aaaa" + "----" + "-" * 11)                # "Laub" nicht rot
    ergebnis = [VORHER[0], ["1", falsch_2], ["2", SOLL_3], VORHER[3]]               # gestrichener Absatz nicht markiert
    r = _auswertung(VORHER, NACHHER, ergebnis, aend)
    # Zeile 2: Text richtig, "Laub" nicht rot; Zeile 4: Text richtig, aber nicht durchgestrichen (und gar nicht beruehrt)
    assert r["erreicht_auto"] == 1 and r["stille_fehler_nach_art"].get("format_falsch") == 2
    assert r["stille_fehler_nach_art"].get("uebernommen_aber_falsch") == 2
    assert r["stille_fehler_nach_art"].get("keine_wirkung") == 1                   # Ziel a3 unberuehrt
    assert r["format_treue_zellen_geprueft"] == 2 and r["format_treue_zellen_gleich"] == 1
    assert r["format_treue_zeichen"] - r["format_treue_zeichen_gleich"] == 4


def test_format_wartet_auf_freigabe_einer_offenen_aenderung():
    """Zwei Aenderungen in einer Zelle, eine angewendet (markiert), eine offen: kein Fehler, die Zelle wartet."""
    vorher = [["Nr", "Text"], ["1", "Der Quokka frisst Gras am Morgen und schläft."]]
    soll = _ft(("-", "Der Quokka frisst "), ("a", "Gras"), ("n", "Laub"), ("-", " am Morgen und "), ("n", "tief"), ("-", "schläft."))
    teil = _ft(("-", "Der Quokka frisst "), ("a", "Gras"), ("n", "Laub"), ("-", " am Morgen und schläft."))
    aend = [_ae("a1", 2, "Gras", "Laub"), _ae("a2", 2, "schläft", "tiefschläft", status=PRUEFEN)]
    r = _auswertung(vorher, [vorher[0], ["1", soll]], [vorher[0], ["1", teil]], aend)
    assert r["stille_fehler"] == 0 and r["zellen_warten_auf_freigabe"] == 1


def test_lesen_mit_formatklassen(tmp_path):
    from openpyxl import Workbook
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws["A1"] = CellRichText(TextBlock(InlineFont(rFont="Arial", sz=11, color="FF00B0F0", strike=True), "alt"),
                            TextBlock(InlineFont(rFont="Arial", sz=11, color="FFFF0000"), "neu"), " Rest")
    ws["A2"] = "ganz neu"
    ws["A2"].font = Font(name="Arial", color="FFFF0000")
    ws["A3"] = CellRichText(TextBlock(InlineFont(color="FF0070C0"), "A.2.0"), "\n",
                            TextBlock(InlineFont(color="FFC00000"), "A.3.0"))            # fruehere Runde: normal
    ws["A4"] = "schlicht"
    wb.save(tmp_path / "f.xlsx")
    z = [r[0] for r in messen._lese_excel(tmp_path / "f.xlsx", True)[0]]
    assert z[0] == "altneu Rest" and z[0].klassen == "aaannn-----"
    assert z[1].klassen == "n" * 8 and not getattr(z[2], "klassen", None) and type(z[3]) is str
    assert messen._s(z[0]) != messen._s("altneu Rest") and messen._s_text(z[0]) == "altneu Rest"
    assert messen._ansicht(z[0], "alt") == "alt Rest" and messen._ansicht(z[0], "neu") == "neu Rest"
    assert [r[0] for r in messen._lese_excel(tmp_path / "f.xlsx")[0]][0] == "altneu Rest"   # ohne formate: wie bisher


# ====================================================================== messen am Nachbau-Paar (echtes Tool)
@pytest.fixture(scope="module")
def gemessen(paar):
    return messen.messe_paar(paar, schablonen=True)


def test_messe_paar_nachbau_liest_legende_und_format_aus_korpus(paar, gemessen, tmp_path):
    r = gemessen
    assert r["legende"] == "auto" and r["nachher_format"] == "markiert" and r["markiert"] is True
    assert r["soll_zellen"] > 0 and r["idempotent"] is not False and "abgestuerzt" not in r
    assert r["original_unveraendert"] is True
    assert r["legende_erkannt"] in (*messen.LEGENDEN, "unbekannt", "sonstige") and isinstance(r["markup_leiste_erkannt"], bool)
    # derselbe Paar-Ordner ohne KORPUS.json: Konfiguration wie sie ist (hier: Test-Weiche alt/keine)
    kopie = tmp_path / "ohne" / paar.name
    shutil.copytree(paar, kopie)
    r2 = messen.messe_paar(kopie, einstellungen={})
    assert r2["legende"] == "alt" and r2["nachher_format"] == "sauber" and r2["markiert"] is False


def test_messen_bericht_nennt_legende_und_format_ohne_inhalte(paar, tmp_path):
    eltern = tmp_path / "nachbau_word"
    shutil.copytree(paar, eltern / paar.name)
    (eltern / "KORPUS.json").write_text(json.dumps({"generator_version": "3", "korpus": "nachbau_word", "legende": "auto",
                                                    "nachher_format": "markiert"}), encoding="utf-8")
    assert messen.main(["--korpus", "nachbau_word", "--ordner", str(eltern), "--schablonen"]) in (0, 1)
    ordner = sorted(config.pfad("messung_dir").iterdir())[-1]
    b = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    assert b["legende"] == "auto" and b["nachher_format"] == "markiert" and b["gesamt"]["paare_markiert"] == 1
    assert any(k["name"].startswith("Format-Treue") for k in b["kriterien"])
    md = (ordner / "bericht.md").read_text(encoding="utf-8")
    assert "Legende: auto · nachher-Format: markiert" in md
    _keine_inhalte(paar, (ordner / "bericht.json").read_text(encoding="utf-8") + md)


def _keine_inhalte(paar: Path, text: str) -> None:
    """Regel 16: kein Inhalt aus PDF/Excel im Bericht -- Stichprobe mit den (erfundenen) Fachwoertern des Generators, die in
    einem Bericht nie vorkommen duerfen, und den Zelltexten von nachher.xlsx."""
    vokabel = {w for liste in (wm.PLURAL, wm.ZUBEHOER, wm.ZUSTAND) for w in liste if len(w) >= 7}
    assert not [w for w in vokabel if w in text]
    wb = load_workbook(paar / "nachher.xlsx")
    zellen = {v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str) and len(v) >= 25}
    assert not [z for z in zellen if z in text]


def test_perfektes_tool_volle_format_treue_und_sabotage(paar, monkeypatch, tmp_path):
    """Liefert das Tool genau nachher, stimmen Text und Format jeder Soll-Zelle; ein falscher Strich faellt auf."""
    def perfekt(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(paar / "nachher.xlsx" if Path(quelle).name == "vorher.xlsx" else quelle, ziel)
    monkeypatch.setattr(messen, "uebernehmen", perfekt)
    r = messen.messe_paar(paar)
    assert r["format_treue_zellen_geprueft"] == r["soll_zellen"] == r["format_treue_zellen_gleich"] > 0
    assert r["kollateral"] == 0 and r["idempotent"] is True

    # Sabotage: in einer ersetzten Stelle ist das Neue nicht rot -- Text richtig, Formatklasse falsch
    from openpyxl.cell.text import InlineFont
    erwartet = json.loads((paar / "erwartet.json").read_text(encoding="utf-8"))
    ziel_zeile = next(a["nachher_zeile"] for a in erwartet["aenderungen"]
                      if a["typ"] == "text_ersetzt" and a["art"] != "listennummer" and a["soll_status"] == "uebernommen")
    falsch = tmp_path / "falsch.xlsx"
    wb = load_workbook(paar / "nachher.xlsx", rich_text=True)
    zelle = wb[erwartet["hauptblatt"]].cell(ziel_zeile, 7)
    zelle.value = CellRichText(*[TextBlock(InlineFont(rFont="Arial", sz=11, color="FF00B0F0", strike=True) if k == "alt"
                                           else InlineFont(rFont="Arial", sz=11), t) for k, t in _laeufe(zelle)])
    wb.save(falsch)

    def sabotiert(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(falsch if Path(quelle).name == "vorher.xlsx" else quelle, ziel)
    monkeypatch.setattr(messen, "uebernehmen", sabotiert)
    r2 = messen.messe_paar(paar)
    assert r2["format_treue_zellen_geprueft"] == r["format_treue_zellen_geprueft"]
    assert r2["format_treue_zellen_gleich"] == r["format_treue_zellen_gleich"] - 1
    assert r2["format_treue_zeichen_gleich"] < r["format_treue_zeichen_gleich"]


# ====================================================================== Rueckbau-Messung (messen --nachher)
def test_rueckbau_baut_vorher_nach(paar, erwartet, tmp_path):
    ziel = tmp_path / "rueck.xlsx"
    stats = messen.baue_vorher_zurueck(paar / "nachher.xlsx", ziel)
    neue = sum(1 for a in erwartet["aenderungen"] if a["typ"] == "absatz_neu" and a["bereich"] == "cde")
    assert stats["zeilen_entfernt"] == neue and stats["zellen_zurueckgebaut"] > neue
    v = load_workbook(paar / "vorher.xlsx", rich_text=True)[erwartet["hauptblatt"]]
    r = load_workbook(ziel, rich_text=True)[erwartet["hauptblatt"]]
    assert r.max_row == v.max_row
    for z in range(1, v.max_row + 1):
        for c in (2, 3, 4, 5, 6, 7):
            a, b = v.cell(z, c).value, r.cell(z, c).value
            assert str(a) == str(b), (z, c)                                  # auch Zeilen mit Excel-Abweichung
            assert not _laeufe(r.cell(z, c)) or all(k == "normal" for k, _t in _laeufe(r.cell(z, c)))


def test_messen_nachher_kategorien(paar, monkeypatch, tmp_path):
    eltern = tmp_path / "echt_artig"
    nur = eltern / "echt_0001"
    nur.mkdir(parents=True)
    for n in ("aenderungen.pdf", "nachher.xlsx"):                           # vorher fehlt (oder wuerde ignoriert)
        shutil.copy2(paar / n, nur / n)
    einst = {"legende": "auto"}
    original = messen.uebernehmen

    def perfekt(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(paar / "nachher.xlsx", ziel)
    monkeypatch.setattr(messen, "uebernehmen", perfekt)
    r = messen.messe_nachher(nur, einst)
    alle = sum(r["zellen"].values())
    assert alle > 0 and r["zellen"]["gleich"] == alle and r["zellen"]["tool_extra"] == 0 and not r["schablonen"]
    assert r["original_unveraendert"] and r["nachher_format"] == "markiert"

    def nichts(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(quelle, ziel)
    monkeypatch.setattr(messen, "uebernehmen", nichts)
    r = messen.messe_nachher(nur, einst)
    fehlt = r["zellen"]["tool_fehlt_vorschlag_offen"] + r["zellen"]["tool_fehlt_gar_nicht"]
    assert fehlt == alle and r["zellen"]["gleich"] == 0 and r["schablonen"]
    assert all(set(s) == {"kategorie", "blatt", "zeile", "spalte", "neue_zeile", "soll", "ist"} for s in r["schablonen"])
    monkeypatch.setattr(messen, "uebernehmen", original)       # nicht monkeypatch.undo(): das hebt auch die conftest-Weichen auf
    assert messen.main(["--nachher", "--ordner", str(eltern), "--legende", "auto"]) == 0
    ordner = sorted(config.pfad("messung_dir").iterdir())[-1]
    assert ordner.name.endswith("_rueckbau")
    b = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    assert b["art"] == "rueckbau" and set(b["gesamt"]["zellen"]) == set(messen.RUECKBAU_KATEGORIEN)
    _keine_inhalte(paar, (ordner / "bericht.json").read_text(encoding="utf-8") + (ordner / "bericht.md").read_text(encoding="utf-8"))
    assert messen.main(["--nachher", "--ordner", str(tmp_path / "gibt_es_nicht")]) == 2


# ====================================================================== Welle 3: Fehler-Orte, neue Zeilen, mehrere Specs
def test_jeder_gemeldete_stille_fehler_ist_eine_echte_abweichung(paar, monkeypatch, tmp_path):
    """Selbsttest der Fehler-Orte: `zeile` ist die Zeile in vorher; `zeile_nachher`/`zeile_ergebnis` zeigen, wo Soll und Ist stehen.
    An jedem gemeldeten Ort muss sich Soll und Ist unterscheiden. Sabotage unterhalb einer eingefuegten Zeile wird genau dort
    gemeldet (kein Versatz um die eingefuegten Zeilen)."""
    erwartet = json.loads((paar / "erwartet.json").read_text(encoding="utf-8"))
    neu = sorted(a["nachher_zeile"] for a in erwartet["aenderungen"] if a["typ"] == "absatz_neu" and a["bereich"] == "cde")
    unter = [a for a in erwartet["aenderungen"] if a["bereich"] == "cde" and a["typ"] == "text_ersetzt"
             and a["soll_status"] == "uebernommen" and a["nachher_zeile"] > neu[0] and a["art"] != "listennummer"]
    ziel = unter[-1]
    falsch = tmp_path / "falsch.xlsx"
    wb = load_workbook(paar / "nachher.xlsx", rich_text=True)
    zelle = wb[erwartet["hauptblatt"]].cell(ziel["nachher_zeile"], 7)
    zelle.value = "".join(t for k, t in _laeufe(zelle) if k != "alt") + " Zusatz"     # falsch: Alt fehlt, Wort zu viel
    from openpyxl.styles import Font
    zelle.font = Font(name="Arial", size=11)
    wb.save(falsch)

    def sabotiert(aend, quelle, ziel_, freigaben=None, cfg=None, **_):
        shutil.copy2(falsch if Path(quelle).name == "vorher.xlsx" else quelle, ziel_)
    monkeypatch.setattr(messen, "uebernehmen", sabotiert)
    kopie = tmp_path / "ergebnis.xlsx"
    r = messen.messe_paar(paar, ergebnis_kopie=kopie)
    # Regel 27 = richtige Zelle ohne Freigabe (Soll und Ist absichtlich gleich) -- alle anderen stillen Fehler sind Abweichungen
    falsche_orte = [f for f in r["fehler"] if f["stumm"] and f["zeile"] != 0 and "regel27_verstoss" not in f["stille_arten"]]
    assert any(f["zeile_ergebnis"] == ziel["nachher_zeile"] == f["zeile_nachher"] and f["zeile"] == ziel["zeile"]
               and f["spalte"] == 7 for f in falsche_orte)
    assert ziel["nachher_zeile"] > ziel["zeile"]                          # die Zeile liegt wirklich unter einer neuen Zeile
    neu_rgb = messen._markierungsfarben()[1]
    wn, we = load_workbook(paar / "nachher.xlsx", rich_text=True), load_workbook(kopie, rich_text=True)
    for f in falsche_orte:
        s_ = wn.worksheets[f["blatt"] - 1].cell(f["zeile_nachher"], f["spalte"]) if f["zeile_nachher"] else None
        i_ = we.worksheets[f["blatt"] - 1].cell(f["zeile_ergebnis"], f["spalte"]) if f["zeile_ergebnis"] else None
        form = lambda z: None if z is None else messen._s(messen._zelltext_formatiert(z, neu_rgb))   # noqa: E731
        assert form(s_) != form(i_), f


def test_neue_zeilen_im_soll_nur_textspalte_rot(paar, erwartet):
    n = load_workbook(paar / "nachher.xlsx", rich_text=True)[erwartet["hauptblatt"]]
    for a in erwartet["aenderungen"]:
        if a["typ"] != "absatz_neu" or a["bereich"] != "cde":
            continue
        z = a["nachher_zeile"]
        assert {k for k, _t in _laeufe(n.cell(z, 7))} == {"neu"}
        for c in (2, 3, 4):                                                   # B/C/D wie die Zeile davor
            assert str(n.cell(z, c).value) == str(n.cell(z - 1, c).value), (z, c)
        for c in (5, 6):                                                      # E/F aus der Gliederung, normal
            assert n.cell(z, c).value and all(k == "normal" for k, _t in _laeufe(n.cell(z, c)))
        assert all(n.cell(z, c).value is None for c in range(8, wm.LETZTE_SPALTE + 1))   # rechts der Textspalte leer
        assert n.cell(z, 1).value is None


MEHRERE = {**KLEIN, "name": "spec", "mehrere_specs": True, "fremd_zeilen": (15, 25)}


@pytest.fixture(scope="module")
def paar_specs(tmp_path_factory) -> Path:
    ziel = tmp_path_factory.mktemp("specs") / "nachbau_word"
    p = erzeuge_paar(SEED + 4, MEHRERE, ziel, "nachbau_word")
    (ziel / "KORPUS.json").write_text(json.dumps({"generator_version": "3", **generator.korpus_einstellungen("nachbau_word")}),
                                      encoding="utf-8")
    return p


def test_mehrere_specs_im_register(paar_specs):
    e = json.loads((paar_specs / "erwartet.json").read_text(encoding="utf-8"))
    assert "mehrere_specs" in e["stoerer"] and "dubletten_fremde_spec" in e["stoerer"]
    rollen = [s["rolle"] for s in e["specs"]]
    i = rollen.index("ziel")
    assert rollen.count("ziel") == 1 and 0 < i < len(rollen) - 1 and 3 <= len(rollen) <= 4    # fremde Specs davor UND danach
    v = load_workbook(paar_specs / "vorher.xlsx", rich_text=True)[e["hauptblatt"]]
    n = load_workbook(paar_specs / "nachher.xlsx", rich_text=True)[e["hauptblatt"]]
    texte_ziel: set[str] = set()
    a, b = e["specs"][i]["zeilen_vorher"]
    for z in range(a, b + 1):
        texte_ziel.add(re.sub(r"^\S+[\t\xa0 ]+", "", str(v.cell(z, 7).value)))
    dubletten = 0
    norm = {re.sub(r"[\s_]", "", s["spez"]) for s in e["specs"]}
    assert len(norm) == len(e["specs"])                                       # jede Spec mit eigener Nummer
    for s in e["specs"]:
        a, b = s["zeilen_vorher"]
        an, bn = s["zeilen_nachher"]
        werte = {re.sub(r"[\s_]", "", str(v.cell(z, 2).value)) for z in range(a, b + 1)}
        assert len(werte) == 1                                                # Block: eine Spec (Schreibweise darf schwanken)
        assert {str(v.cell(z, 3).value) for z in range(a, b + 1) if isinstance(v.cell(z, 3).value, str)} <= {s["revision"]}
        assert str(v.cell(a, 5).value).strip() == "Part 1"                   # eigene Gliederung
        if s["rolle"] == "fremd":
            assert b - a == bn - an                                           # Soll: fremder Block unveraendert
            for z in range(a, b + 1):
                for c in range(1, wm.LETZTE_SPALTE + 1):
                    assert str(v.cell(z, c).value) == str(n.cell(z - a + an, c).value)
                dubletten += re.sub(r"^\S+[\t\xa0 ]+", "", str(v.cell(z, 7).value)) in texte_ziel
    assert dubletten >= 1                                                     # gleiche Saetze wie in der Ziel-Spec
    schreibweisen = {str(v.cell(s["zeilen_vorher"][0], 2).value) for s in e["specs"]}
    assert any(" " in x.strip() for x in schreibweisen) or any("_" in x for x in schreibweisen)
    for x in e["aenderungen"]:                                                # alle Aenderungen im Ziel-Block
        if x["bereich"] == "cde":
            assert e["specs"][i]["zeilen_nachher"][0] <= x["nachher_zeile"] <= e["specs"][i]["zeilen_nachher"][1]


def test_messen_erkennt_fremde_specs_und_zaehlt_beruehrung(paar_specs, monkeypatch, tmp_path):
    e = json.loads((paar_specs / "erwartet.json").read_text(encoding="utf-8"))
    fremd = sum(s["zeilen_vorher"][1] - s["zeilen_vorher"][0] + 1 for s in e["specs"] if s["rolle"] == "fremd")

    def perfekt(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(paar_specs / "nachher.xlsx" if Path(quelle).name == "vorher.xlsx" else quelle, ziel)
    monkeypatch.setattr(messen, "uebernehmen", perfekt)
    r = messen.messe_paar(paar_specs)
    assert r["fremde_spec_zeilen"] == fremd and r["fremde_spec_beruehrt"] == 0 and r["kollateral"] == 0
    # Sabotage: das "Tool" aendert eine Dublette in einer fremden Spec
    s = next(x for x in e["specs"] if x["rolle"] == "fremd")
    falsch = tmp_path / "falsch.xlsx"
    wb = load_workbook(paar_specs / "nachher.xlsx", rich_text=True)
    zelle = wb[e["hauptblatt"]].cell(s["zeilen_nachher"][0] + 2, 7)
    zelle.value = str(zelle.value) + " zusätzlich"
    wb.save(falsch)

    def sabotiert(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        shutil.copy2(falsch if Path(quelle).name == "vorher.xlsx" else quelle, ziel)
    monkeypatch.setattr(messen, "uebernehmen", sabotiert)
    r2 = messen.messe_paar(paar_specs)
    assert r2["fremde_spec_beruehrt"] == 1 and r2["kollateral"] == 1
    assert r2["stille_fehler_nach_art"].get("fremde_spec") == 1 and r2["stille_fehler"] == r["stille_fehler"] + 1
    b = messen.fasse_zusammen([dict(r2, name="x")], "nachbau_word", "3", False)
    k = next(x for x in b["kriterien"] if x["name"] == "Fremde Spec berührt")
    assert k["ampel"] == "NICHT ok" and b["gesamt"]["paare_mit_fremder_spec"] == 1
