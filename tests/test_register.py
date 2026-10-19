"""CDE-Register (ADR 0006 Punkte 5 und 7, Auftrag 003 ZUORDNEN): Aufbau-Erkennung nur aus der Struktur, Zuordnung PDF-Absatz
<-> Registerzeile (Listennummer, Sub-Section, Section, Part), Neu-Ansicht markierter Zellen, Dubletten ueber die Reihenfolge,
feine Operationen, ganz gestrichene Zeile, neue Zeilen mit Section/Sub-Section/Polster.

Nur erfundene Daten (Regel 16): die Mappe hat die gemessenen FORMEN einer echten CDE (5 Blaetter, Kopfbereich, Kopfzeile 6,
Beispielzeile 7, Daten ab Zeile 8, Spalten E/F/G/I), aber ausgedachte Woerter, Blattnamen und ein ausgedachtes Kennwort.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from src import markierung as mk
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.models import NICHT, PRUEFEN, UEBERNOMMEN
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests._word_pdf import WORT_GRUEN, WordDok

KENNWORT = "Gliederungszeile"           # erfunden
NB = "\xa0"
POLSTER = {1: {"zahl.": "\t", "klein.": NB * 15 + " ", "zahl)": NB * 13 + " ", "klein)": NB * 14 + " "},
           2: {"zahl.": "\t", "klein.": "\t", "zahl)": "\t", "klein)": "\t"}}
E, F, G, I = 5, 6, 7, 9
DATEN_AB = 8


def _label_art(label: str) -> str:
    kern = label[:-1]
    return ("zahl" if kern.isdigit() else "klein") + label[-1]


# ================================================================================================ Mappe (Excel)
def register_mappe(pfad: Path, absaetze: list[tuple], variante: int = 1, e_leer: bool = False,
                   g_wert: dict[int, object] | None = None, nebenblaetter: bool = True, spec_von=None) -> Path:
    """Erfundene CDE-Mappe. absaetze: (art, label, text[, fett]) mit art part | section | sub | sub_titel | liste |
    liste_titel | absatz. Zeile des i-ten Absatzes: DATEN_AB + i. g_wert: Index -> eigener Zellwert fuer G (z. B. Rich Text).
    spec_von: Index -> Wert der Spec-Spalte B (Standard: eine Spec fuer alle Zeilen)."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Hinweise"
    for i, t in enumerate(["Hinweise zur Bearbeitung (erfunden)", "Je Absatz gibt es eine Zeile.",
                           "Status in Spalte H waehlen.", "Kennwort in Spalte I bei Ueberschriften."], 1):
        ws0.cell(i, 1, t)
    if nebenblaetter:
        deck = wb.create_sheet("Deckblatt")
        for i, (k, v) in enumerate([("Projekt", "Zebrafink Nord"), ("Ort", "Pumpwerk Ostfeld"), ("Titel", "Pumpenhaus Leitfaden")], 2):
            deck.cell(i, 2, k)
            deck.cell(i, 3, v)
        unt = wb.create_sheet("Unterschriften")
        for i, z in enumerate([["Rolle", "Stelle", "Datum"], ["Erstellt", "Planung", ""], ["Geprueft", "Pruefstelle", ""]], 1):
            for j, v in enumerate(z, 1):
                if v:
                    unt.cell(i, j, v)
    ws = wb.create_sheet("Abgleich")
    for rng in ("A1:A4", "B1:B4", "D1:F4", "G1:G4", "H1:P4", "E5:F5", "G5:G6"):
        ws.merge_cells(rng)
    ws["B1"] = "CDE"
    ws["D1"] = "Abgleich Spezifikation 40 05 13"
    ws["G1"] = "Projekt Zebrafink - Pumpenhaus"
    ws["H1"] = "Hinweise: Spalten J bis N ankreuzen, Status in Spalte R pflegen, Begruendung in Spalte P eintragen."
    ws["U2"] = '=COUNTIF(H:H,"D")'
    for adr, t in {"A5": "Nr.", "B5": "Spezifikation", "C5": "Revision", "D5": "Datum", "E5": "Gliederung",
                   "G5": "Anforderung", "H5": "Status", "I5": "Beschreibung", "J5": "Bewertung", "P5": "Begruendung",
                   "A6": "Pos.", "B6": "Spez.", "C6": "Rev.", "D6": "Datum", "E6": "Abschnitt:", "F6": "Unterabschnitt:",
                   "J6": "K1", "K6": "K2", "R6": "Stand"}.items():
        ws[adr] = t
    for sp, t in {"A": "0", "B": "00-00-00", "C": "Rev 0.", "D": "00-00-00", "E": "Section", "F": "Sub-Section",
                  "G": "Beispiel: eine Anforderung wie in der Spezifikation mit Listennummer im Text", "H": "D",
                  "I": "Beispiel fuer eine Beschreibung der Zeile", "R": "Offen"}.items():
        ws[f"{sp}7"] = t
    sec = ""
    sub = ""
    for i, a in enumerate(absaetze):
        art, label, text = a[0], a[1], a[2]
        fett = a[3] if len(a) > 3 else art in ("part", "section", "sub_titel", "liste_titel")
        r = DATEN_AB + i
        if art == "part":
            sec, sub = f"Part {label}", ""
        elif art == "section":
            sec, sub = label, ""
        elif art in ("sub", "sub_titel"):
            sub = label
        ws.cell(r, 2, spec_von(i) if spec_von else "40 05 13")
        ws.cell(r, 3, "A.3.0")
        ws.cell(r, 4, "2026_03_01")
        ws.cell(r, E, sec + (" " if e_leer else ""))
        if sub and art not in ("part", "section"):
            ws.cell(r, F, sub)
        if art in ("liste", "liste_titel"):
            g = label + POLSTER[variante][_label_art(label)] + text
        else:
            g = text
        z = ws.cell(r, G, g)
        z.font = Font(name="Arial", size=11, bold=fett or None)
        if g_wert and i in g_wert:
            z.value = g_wert[i]
        if art in ("part", "section", "sub_titel", "liste_titel"):
            ws.cell(r, I, KENNWORT)
        else:
            ws.cell(r, 8, "DEC"[i % 3])
            ws.cell(r, 18, "Offen")
            if i % 4 == 1:
                ws.cell(r, 10, "X")
            if i % 5 == 2:
                ws.cell(r, I, f"Beschreibung {i}: Kontrolle durch die Fachplanung")
    ws.auto_filter.ref = f"A6:R{DATEN_AB + len(absaetze) - 1}"
    aus = wb.create_sheet("Auswahl")
    for i, v in enumerate("DEC", 1):
        aus.cell(i, 1, v)
    aus.sheet_state = "hidden"
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(pfad))
    return pfad


# ================================================================================================ PDF (Word-Markup)
def stuecke(markup: str, farbe=None) -> list[tuple]:
    """'Die {-alte-}{+neue+} Pumpe' -> Stuecke fuer WordDok (k schwarz, n neu, d entfaellt)."""
    out = []
    for m in re.finditer(r"\{\+(.+?)\+\}|\{-(.+?)-\}|([^{]+)", markup):
        if m.group(1):
            out.append((m.group(1), "n") + ((farbe,) if farbe else ()))
        elif m.group(2):
            out.append((m.group(2), "d") + ((farbe,) if farbe else ()))
        else:
            out.append((m.group(3), "k"))
    return out


_EBENE = {"zahl.": 2, "klein.": 3, "zahl)": 4, "klein)": 5}


def wort_pdf(pfad: Path, absaetze: list[tuple], vorspann: list[str] | None = None) -> Path:
    """Word-artiges PDF (Markup-Leiste, Linie entscheidet). absaetze: (art, label_markup, text_markup); vorspann: Zeilen
    (Markup) einer ersten Seite (Deckblatt, Revisionstabelle), danach beginnt die Spezifikation auf Seite 2."""
    d = WordDok()
    if vorspann:
        for z in vorspann:
            d.absatz([stuecke(z)], ebene=1, abstand=6)
        d.neue_seite()
    for art, label, text in absaetze:
        if d.top > 600:
            d.neue_seite()
        gruen = WORT_GRUEN if "{" in label else None
        if art == "part":
            d.kapitel(stuecke(f"PART {label} ") + stuecke(text))
        elif art == "section":
            d.ueberschrift(stuecke(label), stuecke(text))
        elif art in ("sub", "sub_titel"):
            d.listenpunkt(1, stuecke(label, gruen), [stuecke(text)], abstand=0)
        elif art in ("liste", "liste_titel"):
            roh = re.sub(r"\{[-+](.+?)[-+]\}", r"\1", label)
            d.listenpunkt(_EBENE[_label_art(roh)], stuecke(label, gruen), [stuecke(text)], abstand=0)
        else:
            d.absatz([stuecke(text)], ebene=1)
    return d.speichern(pfad)


def alt_von(markup: str) -> str:
    return re.sub(r"\{\+.+?\+\}", "", markup).replace("{-", "").replace("-}", "")


def neu_von(markup: str) -> str:
    return re.sub(r"\{-.+?-\}", "", markup).replace("{+", "").replace("+}", "")


def excel_absaetze(pdf_absaetze: list[tuple]) -> list[tuple]:
    """Die alte Fassung (Excel vorher) zu PDF-Absaetzen; ganz neue Absaetze fehlen."""
    out = []
    for art, label, text in pdf_absaetze:
        if text.startswith("{+") and text.endswith("+}") and "{" not in text[2:-2]:
            continue
        out.append((art, alt_von(label).rstrip(".") if art in ("sub", "sub_titel") else alt_von(label), alt_von(text)))
    return out


# Ein kleines Register (Teil 1 der Spezifikation), wie es echte CDEs zeigen
GRUND = [
    ("part", "1", "ALLGEMEINES"),
    ("section", "1.1", "PUMPENHAUS"),
    ("sub", "A.", "Die Pumpe foerdert das Wasser in das obere Becken."),
    ("sub_titel", "B.", "Bezugsnormen:"),
    ("liste", "1.", "Werknorm WN 1001 fuer Ventile."),
    ("liste", "2.", "Werknorm WN 2040 fuer Leitungen."),
    ("liste", "a.", "Pruefung der Dichtungen nach Plan."),
    ("liste", "b.", "Pruefung der Lager nach Plan."),
    ("sub", "C.", "Der Zebrafink wird im Nistkasten beobachtet."),
    ("section", "1.2", "LAGERUNG"),
    ("sub", "A.", "Ersatzteile werden trocken und frostfrei gelagert."),
    ("sub_titel", "B.", "Kennzeichnung:"),
    ("liste", "1.", "Jedes Teil traegt ein Typenschild."),
    ("liste", "2.", "Der Kranich steht auf dem Etikett."),
]


def _plan(tmp_path, pdf_absaetze, excel_abs=None, vorspann=None, **mappe):
    pdf = wort_pdf(tmp_path / "a.pdf", pdf_absaetze, vorspann)
    xlsx = register_mappe(tmp_path / "alt.xlsx", excel_abs if excel_abs is not None else excel_absaetze(pdf_absaetze), **mappe)
    xl = ExcelModell(xlsx)
    return plane(analysiere_pdf(pdf), xl), xl, xlsx


def _ops(ae):
    return [(o.art, o.zeile, o.spalte, o.start, o.ende, o.text) for o in ae.ops]


def _ae(plan, alt: str):
    return next(a for a in plan.aenderungen if a.alt == alt)


# ================================================================================================ 1. Aufbau-Erkennung
def test_register_wird_aus_der_struktur_erkannt(tmp_path):
    xl = ExcelModell(register_mappe(tmp_path / "a.xlsx", excel_absaetze(GRUND)))
    lay = xl.layouts["Abgleich"]
    assert lay.register and lay.modell == "zeile_pro_absatz"
    assert (lay.section_spalte, lay.untergliederung_spalte, lay.inhalt_spalte, lay.kennwort_spalte) == (E, F, G, I)
    assert lay.titel_spalte is None
    assert (lay.kopfzeile, lay.daten_ab) == (6, 8)                      # Kopfzeile 6, Beispielzeile 7 uebersprungen
    assert lay.kennwort == KENNWORT and lay.sicherheit == "hoch"
    assert lay.kopf_zeilen == [8, 9, 17]                                # Part 1, 1.1, 1.2
    assert not any(l.register for n, l in xl.layouts.items() if n != "Abgleich")
    # Kandidaten nur aus der Textspalte der Datenzeilen (kein Kopfbereich, keine Beispielzeile, keine Pruefer-Spalten)
    eigene = {(k.zeile, k.spalte) for k in xl.kandidaten if k.blatt == "Abgleich"}
    assert eigene == {(r, G) for r in range(8, 8 + len(GRUND))}
    arten = [xl.register_zeile("Abgleich", r).art for r in range(8, 8 + len(GRUND))]
    assert arten == ["part", "section", "sub", "sub_titel", "liste", "liste", "liste", "liste", "sub", "section", "sub",
                     "sub_titel", "liste", "liste"]
    rz = xl.register_zeile("Abgleich", 14)
    assert (rz.label, rz.polster, rz.sub, rz.section) == ("a.", NB * 15 + " ", "B", "1.1")
    assert xl.zeilen_section[("Abgleich", 8)] == "Part 1" and xl.zeilen_section[("Abgleich", 7)] is None
    b = next(d for d in xl.beschreibung() if d["blatt"] == "Abgleich")
    assert b["register"] and b["untergliederung_spalte"] == "F" and b["kennwort_spalte"] == "I"
    assert KENNWORT not in repr(xl.beschreibung())                    # das Kennwort ist ein Zellinhalt


def test_zweite_variante_tab_polster_viele_ueberschriften_und_sub_sections_bis_bb(tmp_path):
    subs = [chr(65 + i) for i in range(26)] + ["AA", "BB"]
    abs_ = [("part", "2", "PRODUKTE"), ("section", "2.1", "PUMPEN")]
    for s in subs:
        abs_.append(("sub_titel", s, f"Gruppe {s} Bauteile:"))
        abs_.append(("liste_titel", "1.", f"Pruefungen fuer Gruppe {s}:"))
        abs_.append(("liste", "a.", f"Sichtpruefung der Gruppe {s} vor der Montage."))
        abs_.append(("liste", "b.", f"Druckpruefung der Gruppe {s} nach der Montage."))
    xl = ExcelModell(register_mappe(tmp_path / "b.xlsx", abs_, variante=2, e_leer=True))
    lay = xl.layouts["Abgleich"]
    assert lay.register and (lay.section_spalte, lay.untergliederung_spalte, lay.inhalt_spalte, lay.kennwort_spalte) == (E, F, G, I)
    assert lay.daten_ab == 8
    letzte = xl.register_zeile("Abgleich", 8 + len(abs_) - 1)
    assert (letzte.sub, letzte.label, letzte.polster, letzte.section, letzte.section_roh) == ("BB", "b.", "\t", "2.1", "2.1 ")
    arten = {xl.register_zeile("Abgleich", r).art for r in range(8, 8 + len(abs_))}
    assert arten == {"part", "section", "sub_titel", "liste_titel", "liste"}
    assert xl.zeilen_section[("Abgleich", 8)] == "Part 2"              # "Part 2 " mit Leerzeichen am Ende


def test_fruehere_layouts_sind_kein_register(tmp_path):
    from tests.fixtures.szenario import baue_excel
    for layout in ("L1", "L2", "L3"):
        p = tmp_path / f"{layout}.xlsx"
        baue_excel(p, [("1", "Eins", ["Erster Absatz mit Text.", "Zweiter Absatz."]), ("1.1", "Zwei", ["Dritter."]),
                       ("2", "Drei", ["Vierter Absatz."]), ("2.1", "Vier", ["Fuenfter."]), ("3", "Fuenf", ["Sechster."]),
                       ("3.1", "Sechs", ["Siebter."]), ("4", "Sieben", ["Achter."])], layout)
        assert not any(l.register for l in ExcelModell(p).layouts.values()), layout


def test_profil_mit_modell_register(tmp_path):
    from src import layouts
    xlsx = register_mappe(tmp_path / "a.xlsx", excel_absaetze(GRUND))
    p = layouts.profil_aus_dict({"name": "reg", "aktiv": True, "blaetter": [
        {"blatt": "Abgleich", "section_spalte": "E", "inhalt_spalte": "G", "untergliederung_spalte": "F",
         "kennwort_spalte": "I", "modell": "register", "kopfzeile": 6}]})
    xl = ExcelModell(xlsx, p)
    lay = xl.layouts["Abgleich"]
    assert lay.register and lay.daten_ab == 8 and lay.kennwort == KENNWORT and lay.profil == "reg"
    assert xl.register_zeile("Abgleich", 14).art == "liste"
    # fruehere Profile behalten ihren Hash (die neuen Schluessel stehen nur im Kanon, wenn sie gesetzt sind)
    alt = layouts.profil_aus_dict({"name": "x", "blaetter": [{"blatt": "B", "inhalt_spalte": "C", "modell": "zeile_pro_absatz"}]})
    assert "daten_ab" not in alt.blaetter["B"].kanon()
    with pytest.raises(layouts.ProfilFehler, match="nur mit modell"):
        layouts.profil_aus_dict({"name": "y", "blaetter": [{"blatt": "B", "inhalt_spalte": "C", "modell": "zeile_pro_absatz",
                                                            "kennwort_spalte": "D"}]})


# ================================================================================================ 2. Zuordnung
@pytest.mark.legende_neu
def test_zuordnung_label_ueberschrift_part_und_sub_section(tmp_path):
    pdf = list(GRUND)
    pdf[0] = ("part", "1", "{-ALLGEMEINES-}{+ALLGEMEINE ANGABEN+}")
    pdf[1] = ("section", "1.1", "PUMPEN{+HAUS+}")
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    pdf[7] = ("liste", "{-b.-}{+c.+}", "Pruefung der Lager nach Plan.")
    excel = excel_absaetze(pdf)
    plan, xl, _ = _plan(tmp_path, pdf, excel)
    sub = _ae(plan, "foerdert")
    assert sub.status == UEBERNOMMEN and _ops(sub) == [("ersetzen", 10, G, 10, 18, "hebt")]   # G ohne "A."
    liste = _ae(plan, "2040")
    assert liste.status == UEBERNOMMEN and _ops(liste) == [("ersetzen", 13, G, 15, 19, "2041")]
    label = _ae(plan, "b.")
    assert label.status == UEBERNOMMEN and _ops(label) == [("ersetzen", 15, G, 0, 2, "c.")]   # nur die Listennummer
    part = _ae(plan, "ALLGEMEINES")
    assert part.status == UEBERNOMMEN and _ops(part) == [("ersetzen", 8, G, 0, 11, "ALLGEMEINE ANGABEN")]
    titel = next(a for a in plan.aenderungen if a.neu == "HAUS")
    assert titel.status == UEBERNOMMEN and [o[:3] for o in _ops(titel)] == [("ersetzen", 9, G)]
    assert all(a.grund_einfach for a in plan.aenderungen)
    assert all(a.bereich == "" for a in plan.aenderungen)


@pytest.mark.legende_neu
def test_geaenderte_section_nummer_steht_in_der_section_spalte_und_ist_pruefen(tmp_path):
    pdf = list(GRUND)
    pdf[9] = ("section", "{-1.2-}{+1.3+}", "LAGERUNG")
    plan, xl, _ = _plan(tmp_path, pdf)
    ae = _ae(plan, "1.2")
    assert ae.status == PRUEFEN and "label_spalte" in ae.pruefcodes
    assert _ops(ae) == [("ersetzen", 17, E, 0, 3, "1.3")]
    assert any(ab.zellen == [(17, E, 0, 3)] for ab in plan.absaetze)


@pytest.mark.legende_neu
def test_dublette_in_zwei_sub_sections_wird_ueber_die_reihenfolge_eindeutig(tmp_path):
    grund = GRUND[:9] + [("sub_titel", "D.", "Pruefungen:"), ("liste", "1.", "Pruefung nach Plan."),
                         ("sub_titel", "E.", "Abnahme:"), ("liste", "1.", "Pruefung nach Plan.")] + GRUND[9:]
    pdf = list(grund)
    pdf[12] = ("liste", "1.", "Pruefung nach {-Plan-}{+Zeitplan+}.")
    plan, xl, _ = _plan(tmp_path, pdf)
    ae = _ae(plan, "Plan")
    assert ae.status == UEBERNOMMEN and _ops(ae)[0][1] == 8 + 12        # die zweite Pruefzeile, nicht die erste


def _rich_g() -> CellRichText:
    alt_lauf = InlineFont(rFont="Arial", sz=11, color="FF0070C0", strike=True)     # Markierung einer frueheren Runde
    return CellRichText("Die Pumpe ", TextBlock(alt_lauf, "saugt "), "foerdert das Wasser in das obere Becken.")


@pytest.mark.legende_neu
def test_rich_text_mit_alter_markierung_wird_ueber_die_neu_ansicht_gelesen(tmp_path):
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    plan, xl, xlsx = _plan(tmp_path, pdf, g_wert={2: _rich_g()})
    assert xl.text("Abgleich", 10, G) == "Die Pumpe foerdert das Wasser in das obere Becken."
    assert xl.roh_position("Abgleich", 10, G, 10) == 16                   # "foerdert" steht roh hinter "saugt "
    ae = _ae(plan, "foerdert")
    assert ae.status == UEBERNOMMEN and _ops(ae) == [("ersetzen", 10, G, 10, 18, "hebt")]   # Position in der Neu-Ansicht


@pytest.mark.markierung_an
@pytest.mark.legende_neu
def test_markiert_eintragen_und_zweiter_lauf_ist_schon_eingetragen(tmp_path):
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    pdf[7] = ("liste", "{-b.-}", "{-Pruefung der Lager nach Plan.-}")
    plan, xl, xlsx = _plan(tmp_path, pdf)
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze)
    assert all(a.status == UEBERNOMMEN and a.wirkung == "ok" for a in plan.aenderungen), \
        [(a.id, a.status, a.pruefcodes) for a in plan.aenderungen]
    ws = load_workbook(str(out), rich_text=True)["Abgleich"]
    st = mk.stuecke_aus_zelle(ws.cell(10, G))
    assert mk.ansicht(st, "neu") == "Die Pumpe hebt das Wasser in das obere Becken."
    assert [[s.format, s.text] for s in st if s.format != mk.NORMAL] == [["alt", "foerdert"], ["neu", "hebt"]]
    # Vorschau (Aenderung.stuecke) = spaeteres Ergebnis
    assert _ae(plan, "foerdert").stuecke == [[s.format, s.text] for s in st]
    # zweiter Lauf auf dem markierten Ergebnis: das Modell liest die Neu-Ansicht, alles steht schon da, nichts wird geaendert
    xl2 = ExcelModell(out)
    assert xl2.text("Abgleich", 10, G) == "Die Pumpe hebt das Wasser in das obere Becken."
    assert xl2.text("Abgleich", 15, G) == ""                             # ganz gestrichen: in der Neu-Ansicht leer
    plan2 = plane(analysiere_pdf(tmp_path / "a.pdf"), xl2)
    assert len(plan2.aenderungen) == 3 and all(a.status == NICHT and a.pruefcodes == ["schon_eingetragen"] and not a.ops
                                                for a in plan2.aenderungen)
    out2 = tmp_path / "out2.xlsx"
    uebernehmen(plan2.aenderungen, out, out2, xl=xl2, absaetze=plan2.absaetze)
    ws2 = load_workbook(str(out2), rich_text=True)["Abgleich"]
    for r in range(8, 8 + len(GRUND)):
        assert [[s.format, s.text] for s in mk.stuecke_aus_zelle(ws2.cell(r, G))] == \
            [[s.format, s.text] for s in mk.stuecke_aus_zelle(ws.cell(r, G))]


@pytest.mark.markierung_an
@pytest.mark.legende_neu
def test_zelle_mit_alter_markierung_wird_nie_still_falsch_eingetragen(tmp_path):
    """Eine Zelle mit durchgestrichenem Text einer frueheren Runde: die Operation zeigt in die Neu-Ansicht; beim Eintragen ist
    das Ergebnis entweder richtig (alte Runde bleibt, neue Markierung dazu) oder ein Fall zum Pruefen -- nie still falsch."""
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    plan, xl, xlsx = _plan(tmp_path, pdf, g_wert={2: _rich_g()})
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze)
    ae = _ae(plan, "foerdert")
    st = mk.stuecke_aus_zelle(load_workbook(str(out), rich_text=True)["Abgleich"].cell(10, G))
    if ae.status == UEBERNOMMEN:
        assert mk.ansicht(st, "neu") == "Die Pumpe hebt das Wasser in das obere Becken."
        assert "".join(s.text for s in st if s.format == mk.ALT) == "saugt foerdert"
    else:
        assert ae.status == PRUEFEN and any(c.startswith("gegenprobe_") for c in ae.pruefcodes)
        assert mk.ansicht(st, "alt") == "Die Pumpe saugt foerdert das Wasser in das obere Becken."    # unveraendert


# ================================================================================================ 4. feine Operationen
@pytest.mark.legende_neu
def test_label_aenderung_behaelt_das_polster_der_excel(tmp_path):
    pdf = list(GRUND)
    pdf[6] = ("liste", "{-a.-}{+b.+}", "Pruefung der {-Dichtungen-}{+Lager+} nach Plan.")
    pdf[7] = ("liste", "{-b.-}", "{-Pruefung der Lager nach Plan.-}")
    plan, xl, xlsx = _plan(tmp_path, pdf)
    ae = _ae(plan, "a.")
    assert ae.status == UEBERNOMMEN
    assert [o[3:] for o in _ops(ae)] == [(0, 2, "b.")]                    # nur die Nummer; das Polster (NBSP) bleibt
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze)
    g = load_workbook(str(out))["Abgleich"].cell(14, G).value
    assert g == "b." + NB * 15 + " Pruefung der Lager nach Plan."


def _anwenden(text: str, ops) -> str:
    for o in sorted(ops, key=lambda o: -o.start):
        text = text[:o.start] + o.text + text[o.ende:]
    return text


def test_feine_operationen_lassen_excel_zeichen_zwischen_zwei_aenderungen():
    from src.abgleich import Fundstelle, laeufe_aus_tokens, ops_fuer_spur
    from src.models import ADD, DEL, KEEP, MOD, PdfToken

    def t(o, n, st, teile=()):
        return PdfToken(old=o, new=n, status=st, page=1, bbox=(0, 0, 1, 1), teile=list(teile))

    def fund(x):
        return Fundstelle(blatt="B", zellen=((2, 3, 0, len(x)),), art="absatz", schluessel=(0, 2, 2, 3, 0), text=x,
                          end_schluessel=(0, 2, 2, 3, len(x)))
    toks = [t("Druck", "Druck", KEEP), t("10", "16", MOD, [(DEL, "10"), (ADD, "16")]),
            t("bar", "kPa", MOD, [(DEL, "bar"), (ADD, "kPa")]), t("fest.", "fest.", KEEP)]
    excel = "Druck 10\xa0bar fest."                                       # geschuetztes Leerzeichen zwischen Zahl und Einheit
    so = ops_fuer_spur(laeufe_aus_tokens(toks), fund(excel), lambda b, z, s: excel)
    (ops,) = so.ops.values()
    assert [(o.start, o.ende, o.text) for o in ops] == [(6, 8, "16"), (9, 12, "kPa")]   # zwei Operationen, NBSP bleibt
    assert _anwenden(excel, ops) == "Druck 16\xa0kPa fest."
    # Tab der Excel zwischen einer Ersetzung und einer Streichung bleibt; nur die Streichung nimmt ihre Naht mit
    toks2 = [t("Die", "Die", KEEP), t("x", "y", MOD, [(DEL, "x"), (ADD, "y")]), t("z", "", DEL), t("Ende.", "Ende.", KEEP)]
    excel2 = "Die x\tz Ende."
    so2 = ops_fuer_spur(laeufe_aus_tokens(toks2), fund(excel2), lambda b, z, s: excel2)
    assert _anwenden(excel2, [o for v in so2.ops.values() for o in v]) == "Die y\tEnde."
    # reine Einfuegung ohne Bereich: das geschuetzte Leerzeichen der Excel wird nicht mit ersetzt
    toks3 = [t("die", "die", KEEP), t("", "neuen", ADD), t("Daten", "Daten", KEEP)]
    excel3 = "die\xa0Daten"
    (ops3,) = ops_fuer_spur(laeufe_aus_tokens(toks3), fund(excel3), lambda b, z, s: excel3).ops.values()
    assert [(o.start, o.ende, o.text) for o in ops3] == [(4, 4, "neuen ")]


# ================================================================================================ ganz gestrichene Zeile
@pytest.mark.legende_neu
def test_ganz_gestrichener_listenpunkt_und_sub_section_absatz_markieren_die_ganze_zeile(tmp_path):
    pdf = list(GRUND)
    pdf[7] = ("liste", "{-b.-}", "{-Pruefung der Lager nach Plan.-}")
    pdf[8] = ("sub", "{-C.-}", "{-Der Zebrafink wird im Nistkasten beobachtet.-}")
    plan, xl, xlsx = _plan(tmp_path, pdf)
    liste = next(a for a in plan.aenderungen if "Lager" in a.alt)
    assert liste.typ == "absatz_geloescht" and liste.status == UEBERNOMMEN
    assert _ops(liste) == [("ersetzen", 15, G, 0, len(xl.text("Abgleich", 15, G)), "")]   # ganze Zelle (Nummer im Text)
    sub = next(a for a in plan.aenderungen if "Zebrafink" in a.alt)
    # nur die Textzelle; der Sub-Section-Buchstabe in F bleibt (so markiert Nico)
    assert _ops(sub) == [("ersetzen", 16, G, 0, len(xl.text("Abgleich", 16, G)), "")]


@pytest.mark.markierung_an
@pytest.mark.legende_neu
def test_ganz_gestrichene_zeile_bleibt_hellblau_durchgestrichen(tmp_path):
    pdf = list(GRUND)
    pdf[7] = ("liste", "{-b.-}", "{-Pruefung der Lager nach Plan.-}")
    plan, xl, xlsx = _plan(tmp_path, pdf)
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze)
    ws = load_workbook(str(out), rich_text=True)["Abgleich"]
    st = mk.stuecke_aus_zelle(ws.cell(15, G))
    assert [s.format for s in st] == [mk.ALT] and mk.ansicht(st, "alt") == xl.text("Abgleich", 15, G)
    assert ws.cell(16, G).value == "Der Zebrafink wird im Nistkasten beobachtet."    # die Zeile bleibt, nichts rueckt


# ================================================================================================ 5. neue Zeilen
@pytest.mark.legende_neu
def test_neue_zeilen_mit_section_sub_section_polster_und_kennwort(tmp_path):
    pdf = GRUND[:8] + [("liste", "{+c.+}", "{+Pruefung der Wellen nach Plan.+}")] + GRUND[8:9] + \
        [("sub_titel", "{+D.+}", "{+Ersatzteile:+}")] + GRUND[9:] + \
        [("section", "{+1.3+}", "{+UEBERGABE+}"), ("sub", "{+A.+}", "{+Die Anlage wird dem Betreiber uebergeben.+}")]
    plan, xl, xlsx = _plan(tmp_path, pdf)
    neu = [a for a in plan.aenderungen if a.typ in ("absatz_neu", "section_neu")]
    werte = {a.neu.split()[-1]: a.ops[0] for a in neu}
    liste = werte["Plan."]
    assert (liste.art, liste.zeile, liste.vor) == ("zeile_einfuegen", 15, False)          # direkt nach "b."
    assert liste.zellen == {E: "1.1", F: "B", G: "c." + NB * 15 + " Pruefung der Wellen nach Plan."}
    assert xl.register_zeile("Abgleich", liste.stil_von).art == "liste"
    sub = werte["Ersatzteile:"]
    assert sub.zeile == 16 and sub.zellen == {E: "1.1", F: "D", G: "Ersatzteile:", I: KENNWORT}
    assert xl.register_zeile("Abgleich", sub.stil_von).art == "sub_titel"
    sec = werte["UEBERGABE"]
    assert sec.zeile == 8 + len(GRUND) - 1 and sec.zellen == {E: "1.3", G: "UEBERGABE", I: KENNWORT}
    assert xl.register_zeile("Abgleich", sec.stil_von).art == "section"
    folge = werte["uebergeben."]
    assert folge.zellen == {E: "1.3", F: "A", G: "Die Anlage wird dem Betreiber uebergeben."} and folge.seq > sec.seq
    assert all(a.stuecke == [[mk.NEU, a.ops[0].zellen[G]]] for a in neu)


@pytest.mark.legende_neu
def test_neue_zeile_hinter_abweichendem_vorgaenger_ist_ueber_die_reihenfolge_sicher(tmp_path):
    """Welle 3 (Korpus nachbau_word, Paar 303): der Vorgaenger des neuen Listenpunkts steht in der Excel etwas anders (Excel
    weicht ab), liegt aber als einzige offene Zeile zwischen zwei sicheren Ankern -> der neue Absatz ist sicher."""
    pdf = GRUND[:6] + [("liste", "{+3.+}", "{+Werknorm WN 3000 fuer Pumpen.+}")] + GRUND[6:]
    excel = excel_absaetze(GRUND)
    excel[5] = ("liste", "2.", "Werknorm WN 2040 fur Leitungen.")              # Abweichung in der Excel
    plan, xl, _ = _plan(tmp_path, pdf, excel)
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and not ae.pruefcodes
    assert (ae.ops[0].zeile, ae.ops[0].vor) == (13, False)
    # zwei offene Zeilen zwischen den Ankern: nicht eindeutig -> weiter nur ein Vorschlag
    excel[4] = ("liste", "1.", "Werknorm WN 1001 fur Ventile.")
    plan2, _, _ = _plan(tmp_path / "b", pdf, excel)
    (ae2,) = plan2.aenderungen
    assert ae2.status == PRUEFEN and "anker_unsicher" in ae2.pruefcodes


@pytest.mark.legende_neu
def test_neue_zeile_mit_tab_polster_wie_die_nachbarn(tmp_path):
    pdf = GRUND[:6] + [("liste", "{+3.+}", "{+Werknorm WN 3000 fuer Pumpen.+}")] + GRUND[6:]
    plan, xl, _ = _plan(tmp_path, pdf, variante=2)
    (ae,) = [a for a in plan.aenderungen if a.typ == "absatz_neu"]
    assert ae.ops[0].zellen[G] == "3.\tWerknorm WN 3000 fuer Pumpen." and ae.ops[0].zeile == 13


# ================================================================================================ 8. farben-Bericht
@pytest.mark.legende_neu
def test_farben_bericht_richtet_sich_nach_der_gewaehlten_legende(tmp_path):
    from src.cli import farben_bericht
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    a = analysiere_pdf(wort_pdf(tmp_path / "a.pdf", pdf))
    assert a.legende_modus == "linie"
    text = "\n".join(farben_bericht(a))
    assert "Legende: Linie entscheidet" in text
    assert "gedeutet als unverändert (es zählt nur die Linie)" in text and "gedeutet als entfällt" not in text.split("Wörter")[0]
    assert "neu: 1 |" in text or "geändert: 1 |" in text
    for wort in ("Pumpe", "foerdert", "hebt", "Becken", "Zebrafink"):
        assert wort not in text


# ================================================================================================ 6./7. Grund und Stuecke
@pytest.mark.legende_neu
def test_stuecke_und_grund_einfach(tmp_path):
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")
    plan, _, _ = _plan(tmp_path, pdf)
    ae = _ae(plan, "foerdert")
    assert ae.stuecke == [["normal", "Die Pumpe "], ["alt", "foerdert"], ["neu", "hebt"],
                          ["normal", " das Wasser in das obere Becken."]]
    assert ae.grund_einfach.startswith("Wird eingetragen")
