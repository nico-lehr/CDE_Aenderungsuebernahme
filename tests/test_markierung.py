"""Excel-Markierung (src/markierung.py, ADR 0005): Zustand mit Markierung, Lesen, Schreiben, Rundreise. Nur erfundene Woerter."""
import random
import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.gegenprobe import normalform, zustand
from src.markierung import (ALT, NEU, NORMAL, Stueck, aktiv, ansichten, ansicht, faerbe_zelle, format_von_schrift,
                            schreibe, zustand_markiert)
from src.models import Lauf
from tests.test_gegenprobe import L


def stuecke(spec: str, angewendet=("D", "N")):
    return [(s.format, s.text) for s in zustand_markiert(L(spec), set(angewendet))]


def test_standard_ist_farbe_in_der_konfiguration_aber_tests_laufen_ohne():
    assert not aktiv()          # conftest schaltet fuer die frueheren Tests ab


@pytest.mark.markierung_an
def test_markierung_an_marker_nutzt_den_standard():
    assert aktiv()


def test_ersetzung_alt_direkt_vor_neu():
    assert stuecke("darf {-12-} {+7+} ms") == [(NORMAL, "darf "), (ALT, "12"), (NORMAL, " "), (NEU, "7"), (NORMAL, " ms")]


def test_verschmolzene_ersetzung_ohne_trenner_wie_im_pdf():
    # Nico, 18.10.: Alt und Neu stehen genau so hintereinander wie im PDF (kein zusaetzliches Leerzeichen)
    assert stuecke("x {-alt-}|{+neu+} y") == [(NORMAL, "x "), (ALT, "alt"), (NEU, "neu"), (NORMAL, " y")]


def test_konfigurierter_trenner_wird_eingefuegt():
    cfg = {"markierung": {"aenderungen": "farbe", "trenner": " "}}
    st = [(s.format, s.text) for s in zustand_markiert(L("x {-alt-}|{+neu+} y"), {"D", "N"}, cfg)]
    assert st == [(NORMAL, "x "), (ALT, "alt"), (NORMAL, " "), (NEU, "neu"), (NORMAL, " y")]


def test_einfuegung_rot_loeschung_blau_durchgestrichen_stehen_lassen():
    assert stuecke("Die Anlage {+insgesamt+}|.", {"N"}) == [(NORMAL, "Die Anlage "), (NEU, "insgesamt"), (NORMAL, ".")]
    assert stuecke("Die Anlage {-insgesamt-}|.", {"D"}) == [(NORMAL, "Die Anlage "), (ALT, "insgesamt"), (NORMAL, ".")]


def test_nicht_angewendet_alt_bleibt_normal_neu_fehlt():
    assert stuecke("a {-b-} {+c+} d", set()) == [(NORMAL, "a b d")]


def test_ansichten_alt_ist_das_original_neu_ist_der_neue_absatz():
    laeufe = L("Die {-alte-} {+neue+} Pumpe {-und-} {+oder+} Leitung.")
    st = zustand_markiert(laeufe, {"D", "N"})
    assert normalform(ansicht(st, ALT))[0] == normalform(zustand(laeufe, set()))[0]
    assert normalform(ansicht(st, NEU))[0] == normalform(zustand(laeufe, {"D", "N"}))[0]


def test_zufall_beide_ansichten_stimmen_fuer_jede_teilmenge():
    rnd = random.Random(11)
    woerter = ["Ventil", "Pumpe", "Druck", "Wert", "Anlage", "Prüfung", "Leitung"]
    for _ in range(300):
        laeufe: list[Lauf] = []
        for i in range(rnd.randint(1, 9)):
            if laeufe:
                laeufe.append(Lauf("s", rnd.choice([" ", " ", " "])))
            art = rnd.choice("kkkda")
            laeufe.append(Lauf(art, rnd.choice(woerter), f"A{rnd.randint(1, 3)}" if art in "da" else ""))
        ids = sorted({l.aenderung for l in laeufe if l.aenderung})
        for k in range(len(ids) + 1):
            S = set(rnd.sample(ids, k))
            st = zustand_markiert(laeufe, S)
            # neue Ansicht == Absatz laut PDF fuer S; alte Ansicht == ganzer alter Absatz (nichts geht verloren)
            assert normalform(ansicht(st, NEU))[0] == normalform(zustand(laeufe, S))[0]
            assert normalform(ansicht(st, ALT))[0] == normalform(zustand(laeufe, set()))[0]


def test_format_von_schrift_strich_ist_alt_jede_farbe_rot_ohne_strich_ist_neu():
    assert format_von_schrift(Font(strike=True, color="FF000000")) == ALT
    assert format_von_schrift(Font(strike=True, color="FFFF0000")) == ALT
    assert format_von_schrift(Font(color="FFFF0000")) == NEU
    assert format_von_schrift(Font(color="FFC00000")) == NEU
    assert format_von_schrift(Font(color="FF0000FF")) == NORMAL      # blau ohne Strich ist in der Excel nichts Besonderes
    assert format_von_schrift(Font(color="FF000000")) == NORMAL
    assert format_von_schrift(None) == NORMAL


def _zelle(ws="A1", **kw):
    wb = Workbook()
    z = wb.active[ws]
    z.font = Font(name="Arial", sz=12, b=kw.get("fett", False))
    z.fill = PatternFill("solid", fgColor="EEEEEE")
    z.alignment = Alignment(wrap_text=True, vertical="top")
    return wb, z


def test_schreiben_rundreise_behaelt_schrift_fuellung_und_formate(tmp_path):
    wb, z = _zelle(fett=True)
    st = [Stueck(NORMAL, "Die Anlage "), Stueck(ALT, "alt"), Stueck(NORMAL, " "), Stueck(NEU, "neu"), Stueck(NORMAL, " läuft.")]
    schreibe(z, st)
    p = tmp_path / "m.xlsx"
    wb.save(p)
    z2 = load_workbook(p, rich_text=True).active["A1"]
    alt, neu, zs = ansichten(z2)
    assert alt == "Die Anlage alt  läuft." and neu == "Die Anlage  neu läuft."
    assert [s.format for s in zs] == [NORMAL, ALT, NORMAL, NEU, NORMAL]
    assert z2.fill.fgColor.rgb.endswith("EEEEEE")
    blocke = [b for b in z2.value]
    assert all(getattr(b.font, "rFont", "Arial") == "Arial" and getattr(b.font, "sz", 12) == 12 for b in blocke if hasattr(b, "font"))
    assert blocke[1].font.strike and blocke[1].font.color.rgb.endswith("00B0F0")
    assert blocke[3].font.color.rgb.endswith("FF0000") and not blocke[3].font.strike


def test_ganze_zelle_ein_format_wird_zellschrift_kein_rich_text(tmp_path):
    wb, z = _zelle()
    schreibe(z, [Stueck(NEU, "Ganz neue Zeile")])
    assert z.value == "Ganz neue Zeile" and z.font.color.rgb.endswith("FF0000") and z.font.name == "Arial"
    wb2, z2 = _zelle("B2")
    schreibe(z2, [Stueck(ALT, "Ganz weg")])
    assert z2.font.strike and z2.font.color.rgb.endswith("00B0F0")
    assert ansichten(z2)[:2] == ("Ganz weg", "")


def test_faerbe_zelle_zahl_bleibt_zahl():
    wb, z = _zelle()
    z.value = 2.1
    z.number_format = "0.00"
    faerbe_zelle(z, ALT)
    assert z.value == 2.1 and z.number_format == "0.00" and z.font.strike


def test_leere_und_normale_zelle():
    wb, z = _zelle()
    schreibe(z, [])
    assert z.value is None
    schreibe(z, [Stueck(NORMAL, "nur "), Stueck(NORMAL, "Text")])
    assert z.value == "nur Text" and not z.font.strike


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice nicht installiert")
def test_libreoffice_zeigt_rot_blau_und_strich(tmp_path):
    """Beweis, dass Excel-Programme die Markierung wirklich so darstellen: xlsx -> PDF -> Farbe/Strich je Zeichen."""
    import pdfplumber
    wb, z = _zelle()
    schreibe(z, [Stueck(NORMAL, "Die "), Stueck(ALT, "alt"), Stueck(NORMAL, " "), Stueck(NEU, "neu"), Stueck(NORMAL, " Pumpe.")])
    wb.save(tmp_path / "m.xlsx")
    profil = tmp_path / "lo"
    subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{profil}", "--convert-to", "pdf",
                    "--outdir", str(tmp_path), str(tmp_path / "m.xlsx")], check=True, capture_output=True, timeout=180)
    with pdfplumber.open(tmp_path / "m.pdf") as pdf:
        seite = pdf.pages[0]
        farben = {c["text"]: c["non_stroking_color"] for c in seite.chars if c["text"] in "altneu"}
        assert tuple(round(v, 2) for v in farben["a"]) == (0.0, 0.69, 0.94)   # alt = Hellblau 00B0F0
        assert tuple(farben["n"]) == (1.0, 0.0, 0.0)      # neu = rot
        assert len(seite.lines) + len(seite.rects) >= 1   # Durchstreichung


# ================================================================ Auftrag 003: markiere_ops (End-Stuecke aus Zellinhalt + Operationen)
from src.gegenprobe import pruefe_absatz_markiert, zustand_markiert_zeichen    # noqa: E402
from src.markierung import (SICHT_NEU, klassen_stuecke, markiere_ops, sicht_text, trenner_aus_laeufen,  # noqa: E402
                            zusammenfassen)
from src.models import Absatz, Operation                                       # noqa: E402


def _ops(*spec):
    """("e", start, ende, text, id) = ersetzen; ("i", pos, text, id) = einfuegen; ("l", id) = zelle_leeren."""
    out = []
    for k, s in enumerate(spec):
        if s[0] == "e":
            out.append(Operation("ersetzen", "CDE", 2, 3, s[1], s[2], s[3], seq=k, aenderung=s[4]))
        elif s[0] == "i":
            out.append(Operation("einfuegen", "CDE", 2, 3, s[1], s[1], s[2], seq=k, aenderung=s[3]))
        else:
            out.append(Operation("zelle_leeren", "CDE", 2, 3, seq=k, aenderung=s[1]))
    return out


def _fs(stuecke):
    return [(s.format, s.text) for s in stuecke]


def test_markiere_ersetzung_einfuegung_loeschung():
    t = "Die Pumpe läuft leise."
    m = markiere_ops([Stueck(NORMAL, t)], _ops(("e", 4, 9, "Kessel", "A1"), ("i", 15, " heute", "A2"),
                                                ("e", 15, 21, "", "A3")))
    # Leerraum am Rand von Einfuegung und Loeschung ist nie markiert (Nacharbeit: wie im PDF)
    assert _fs(m.stuecke) == [(NORMAL, "Die "), (ALT, "Pumpe"), (NEU, "Kessel"), (NORMAL, " läuft "), (NEU, "heute"),
                              (NORMAL, " "), (ALT, "leise"), (NORMAL, ".")]
    assert normalform(ansicht(m.stuecke, ALT))[0] == t and normalform(ansicht(m.stuecke, NEU))[0] == "Die Kessel läuft heute."
    assert not any(s.behalten for s in m.stuecke if s.format != NORMAL)
    assert [z[2] for z in m.zeichen if z[2] >= 0] == list(range(len(t)))       # jedes Original-Zeichen steht noch da


def test_markiere_mit_pdf_leerraum_zwischen_alt_und_neu():
    zw = trenner_aus_laeufen(L("darf {-12-} {+7+} ms"))
    assert zw == {"D": " "} or zw == {}                     # L() gibt d und a verschiedene IDs (D/N): kein Paar
    laeufe = [Lauf("k", "darf"), Lauf("s", " "), Lauf("d", "12", "A1"), Lauf("s", " "), Lauf("a", "7", "A1"),
              Lauf("s", " "), Lauf("k", "ms")]
    assert trenner_aus_laeufen(laeufe) == {"A1": " "}
    assert trenner_aus_laeufen([Lauf("d", "alt", "A1"), Lauf("a", "neu", "A1")]) == {"A1": ""}
    m = markiere_ops([Stueck(NORMAL, "darf 12 ms")], _ops(("e", 5, 7, "7", "A1")), zwischen=trenner_aus_laeufen(laeufe))
    assert _fs(m.stuecke) == [(NORMAL, "darf "), (ALT, "12"), (NORMAL, " "), (NEU, "7"), (NORMAL, " ms")]


def test_markiere_behaelt_alte_markierungen_und_schrift_der_umgebung():
    fett, rot_alt, strich_alt = object(), object(), object()
    vorher = [Stueck(NORMAL, "Die Pumpe ", fett), Stueck(NEU, "läuft ", rot_alt), Stueck(ALT, "lief ", strich_alt),
              Stueck(NORMAL, "leise.", fett)]
    t = sicht_text(vorher)
    s = t.index("leise")
    m = markiere_ops(vorher, _ops(("e", s, s + 5, "laut", "A1")))
    assert [(x.format, x.text, x.basis, x.behalten) for x in m.stuecke] == [
        (NORMAL, "Die Pumpe ", fett, True), (NEU, "läuft ", rot_alt, True), (ALT, "lief ", strich_alt, True),
        (ALT, "leise", fett, False), (NEU, "laut", fett, False), (NORMAL, ".", fett, True)]


def test_markiere_in_der_neu_sicht_ueberspringt_durchgestrichenes():
    vorher = [Stueck(NORMAL, "Die "), Stueck(ALT, "alte "), Stueck(NORMAL, "Pumpe läuft.")]
    t = sicht_text(vorher, SICHT_NEU)
    assert t == "Die Pumpe läuft."
    s = t.index("läuft")
    m = markiere_ops(vorher, _ops(("e", s, s + 5, "steht", "A1")), sicht=SICHT_NEU)
    assert _fs(m.stuecke) == [(NORMAL, "Die "), (ALT, "alte "), (NORMAL, "Pumpe "), (ALT, "läuft"), (NEU, "steht"),
                              (NORMAL, ".")]
    assert [z[2] for z in m.zeichen if z[2] >= 0] == list(range(len(t)))


def test_markiere_zelle_leeren_und_reihenfolge_an_derselben_stelle():
    m = markiere_ops([Stueck(NORMAL, "Ganz weg.")], _ops(("l", "A1")))
    assert _fs(m.stuecke) == [(ALT, "Ganz weg.")]
    # an derselben Stelle wie im Applier: Ersatztext der dort endenden Ersetzung, dann Einfuegungen
    m = markiere_ops([Stueck(NORMAL, "a b c")], _ops(("i", 3, " y", "A2"), ("e", 2, 3, "B", "A1"), ("i", 2, "x ", "A3")))
    assert ansicht(m.stuecke, NEU) == "a x B y c"
    with pytest.raises(ValueError):
        markiere_ops([Stueck(NORMAL, "kurz")], _ops(("e", 2, 9, "x", "A1")))


def test_klassen_stuecke_fuer_die_anzeige():
    m = markiere_ops([Stueck(NORMAL, "x alt y")], _ops(("e", 2, 5, "neu", "A1")))
    assert klassen_stuecke(m.zeichen) == [["normal", "x "], ["alt", "alt"], ["neu", "neu"], ["normal", " y"]]


def test_zusammenfassen_trennt_behaltene_und_markierte_stuecke():
    # alt: schon durchgestrichen (behalten) und neu durchgestrichen bleiben getrennt -- sie sehen verschieden aus
    st = zusammenfassen([Stueck(ALT, "a", None, True), Stueck(ALT, "b", None, False), Stueck(ALT, "c", None, False)])
    assert [(s.text, s.behalten) for s in st] == [("a", True), ("bc", False)]
    # normal mit gleicher Schrift sieht gleich aus (vorhandener Text + neuer Leerraum): ein Lauf
    st = zusammenfassen([Stueck(NORMAL, "a", None, True), Stueck(NORMAL, " ", None, False, "A1"), Stueck(NORMAL, "b", None, True)])
    assert [(s.text, s.behalten, s.aenderung) for s in st] == [("a b", False, "A1")]


# ---------------------------------------------------------------- A3 mit Markierung (src/gegenprobe.py)
def _a3(text, laeufe, ops, start=0, ende=None, angewendet=("A1",), zwischen=None):
    m = markiere_ops([Stueck(NORMAL, text)], ops, zwischen=zwischen or trenner_aus_laeufen(laeufe))
    ab = Absatz("P1", "CDE", [(2, 3, start, len(text) if ende is None else ende)], laeufe)
    return pruefe_absatz_markiert(ab, {(2, 3): (text, m.zeichen)}, set(angewendet))


def _lauf(spec: str, aid: str = "A1") -> list[Lauf]:
    out = L(spec)
    return [Lauf(x.art, x.text, aid if x.aenderung else "", x.optional) for x in out]


def test_a3_markiert_richtig_besteht():
    assert _a3("Das System muss laufen.", _lauf("Das {-System-} {+Gerät+} muss laufen."), _ops(("e", 4, 10, "Gerät", "A1"))) is None
    assert _a3("Die Anlage arbeitet.", _lauf("Die Anlage arbeitet {+leise+}|."), _ops(("i", 19, " leise", "A1"))) is None
    t = "Eins bleibt. Zwei weg. Vier."
    s, e = t.index("Zwei"), t.index("Zwei") + len("Zwei weg.")
    assert _a3(t, _lauf("{-Zwei-} {-weg.-}"), _ops(("e", s - 1, e, "", "A1")), s, e) is None


def test_a3_markiert_erkennt_fehler():
    assert _a3("Der Wert und der Wert.", _lauf("Der Wert und der {-Wert-} {+Pegel+}|."),
               _ops(("e", 4, 8, "Pegel", "A1"))) == "gegenprobe_text"
    assert _a3("Die Anlage arbeitet.", _lauf("Die Anlage arbeitet {+leise+}|."),
               _ops(("i", 20, " leise", "A1"))) == "gegenprobe_text"                  # hinter dem Punkt
    assert _a3("Erster Teil zweiter Teil.", _lauf("Erster Teil {+neu+} zweiter Teil."),
               _ops(("i", 11, "\nneu", "A1"))) == "gegenprobe_umbruch"
    # richtiger Text, aber falsche Formatklasse: das Wort steht als neu da, ist im PDF aber unveraendert
    assert _a3("x y", _lauf("x {+y+}"), _ops(("e", 2, 3, "y", "A1"))) in ("gegenprobe_alt", "gegenprobe_format")


def test_a3_markiert_meldet_verlust_und_eigenes_ausserhalb():
    laeufe = _lauf("Die Pumpe {-läuft-}|.")
    m = markiere_ops([Stueck(NORMAL, "Die Pumpe läuft.")], _ops(("e", 10, 15, "", "A1")))
    zeichen = [z for z in m.zeichen if z[0] != "P"]                          # ein Original-Zeichen fehlt
    ab = Absatz("P1", "CDE", [(2, 3, 0, 16)], laeufe)
    assert pruefe_absatz_markiert(ab, {(2, 3): ("Die Pumpe läuft.", zeichen)}, {"A1"}) == "gegenprobe_verlust"
    t = "Erster Satz.\nZweiter Satz alt."
    m = markiere_ops([Stueck(NORMAL, t)], _ops(("e", 0, 6, "", "A1")))         # eigene Loeschung im FREMDEN Absatz
    ab = Absatz("P1", "CDE", [(2, 3, 13, len(t))], _lauf("Zweiter Satz {-alt-}|."))
    assert pruefe_absatz_markiert(ab, {(2, 3): (t, m.zeichen)}, {"A1"}) == "gegenprobe_aussen"


def test_soll_zeichen_je_klasse():
    z = zustand_markiert_zeichen(_lauf("x {-alt-}|{+neu+} y"), {"A1"})
    assert "".join(c for c, _, _ in z) == "x altneu y"
    assert [k for c, k, _ in z if c != " "] == ["normal", "alt", "alt", "alt", "neu", "neu", "neu", "normal"]


def test_zufall_markierte_ops_bestehen_die_gegenprobe_und_verlieren_nichts():
    """Wie der Planer: ein Wort ersetzen, ein Wort samt Leerzeichen loeschen, ein Wort einfuegen -- je Absatz eine
    Aenderung. Ergebnis: A3 besteht, alt = Original, neu = PDF-neu, kein Zeichen geht verloren."""
    rnd = random.Random(7)
    woerter = ["Ventil", "Pumpe", "Druck", "Wert", "Anlage", "Prüfung", "Leitung", "Zebrafink"]
    for _ in range(300):
        w = [rnd.choice(woerter) for _ in range(rnd.randint(2, 7))]
        t = " ".join(w) + "."
        k = rnd.randrange(len(w))
        pos = sum(len(x) + 1 for x in w[:k])
        art = rnd.choice("eli")
        laeufe: list[Lauf] = []
        for i, x in enumerate(w):
            if i:
                laeufe.append(Lauf("s", " "))
            if i == k and art == "e":
                laeufe += [Lauf("d", x, "A1"), Lauf("s", " "), Lauf("a", "Neu", "A1")]
            elif i == k and art == "l":
                laeufe.append(Lauf("d", x, "A1"))
            else:
                laeufe.append(Lauf("k", x))
            if i == k and art == "i":
                laeufe += [Lauf("s", " "), Lauf("a", "Neu", "A1")]
        laeufe.append(Lauf("k", "."))
        if art == "e":
            ops = _ops(("e", pos, pos + len(w[k]), "Neu", "A1"))
        elif art == "l":
            s, e = (pos - 1, pos + len(w[k])) if k else (pos, pos + len(w[k]) + (1 if len(w) > 1 else 0))
            ops = _ops(("e", s, e, "", "A1"))
        else:
            ops = _ops(("i", pos + len(w[k]), " Neu", "A1"))
        m = markiere_ops([Stueck(NORMAL, t)], ops, zwischen=trenner_aus_laeufen(laeufe))
        ab = Absatz("P1", "CDE", [(2, 3, 0, len(t))], laeufe)
        assert pruefe_absatz_markiert(ab, {(2, 3): (t, m.zeichen)}, {"A1"}) is None, (t, art, k)
        assert "".join(z[0] for z in m.zeichen if z[2] >= 0) == t                  # Original vollstaendig, in Reihenfolge
        assert normalform(ansicht(m.stuecke, ALT))[0] == normalform(t)[0]        # (Leerraum aus dem PDF zwischen Alt und Neu)
        assert normalform(ansicht(m.stuecke, NEU))[0] == normalform(zustand(laeufe, {"A1"}))[0]


# ================================================================ Nacharbeit: Markierung aus den PDF-Laeufen projizieren
from src.markierung import projiziere                                     # noqa: E402

# Die fuenf Muster aus der Messung an echten Paaren (nur Formen; Woerter und Zahlen erfunden).
# (Excel-Text, PDF-Absatz in Kurzschrift, Soll mit {-alt-} / {+neu+})
MUSTER = [
    ("nach NORM 4711.", "nach {+ZUSATZ-EN+} NORM 4711.", "nach {+ZUSATZ-EN+} NORM 4711."),
    # doppelter Leerraum: der neue Text kommt hinter den ganzen Excel-Leerraum (Welle 3, Punkt 2)
    ("nach  NORM 4711.", "nach {+ZUSATZ-EN+} NORM 4711.", "nach  {+ZUSATZ-EN+} NORM 4711."),
    ("Feuchte: [0 bis 100] Prozent", "Feuchte: {-[-}|0 bis 100|{-]-} Prozent", "Feuchte: {-[-}0 bis 100{-]-} Prozent"),
    ("und NORM 4712. Danach", "und {-NORM-} {-4712-}|. Danach", "und {-NORM 4712-}. Danach"),
    ("Teil 11223‑6 mit Pumpe alt.", "Teil 11223-6 mit Pumpe {-alt-}|{+neu+}|.", "Teil 11223‑6 mit Pumpe {-alt-}{+neu+}."),
    ("Temperatur [20] Grad", "Temperatur {-[20]-}|{+21+} Grad", "Temperatur {-[20]-}{+21+} Grad"),
]


def _kurz(stuecke) -> str:
    return "".join({"normal": s.text, "alt": "{-" + s.text + "-}", "neu": "{+" + s.text + "+}"}[s.format] for s in stuecke)


@pytest.mark.parametrize("excel, pdf, soll", MUSTER)
def test_projektion_markiert_genau_was_das_pdf_markiert(excel, pdf, soll):
    laeufe = _lauf(pdf)
    proj = projiziere([excel], laeufe, {"A1"})
    assert proj is not None
    m = markiere_ops([Stueck(NORMAL, excel)], [], bereiche=[(0, len(excel), proj[0])])
    assert _kurz(m.stuecke) == soll
    assert "".join(z[0] for z in m.zeichen if z[2] >= 0) == excel              # jedes Excel-Zeichen steht unveraendert da
    ab = Absatz("P1", "CDE", [(2, 3, 0, len(excel))], laeufe)
    assert pruefe_absatz_markiert(ab, {(2, 3): (excel, m.zeichen)}, {"A1"}) is None    # A3 markiert gilt streng weiter


def test_projektion_ueber_zwei_zellen_und_ohne_zuordnung():
    laeufe = _lauf("{-2.3-} {+2.4+} Pumpe {+neu+}")
    proj = projiziere(["2.3", "Pumpe"], laeufe, {"A1"})
    nummer = markiere_ops([Stueck(NORMAL, "2.3")], [], bereiche=[(0, 3, proj[0])])
    titel = markiere_ops([Stueck(NORMAL, "Pumpe")], [], bereiche=[(0, 5, proj[1])])
    assert _kurz(nummer.stuecke) == "{-2.3-} {+2.4+}" and _kurz(titel.stuecke) == "Pumpe {+neu+}"
    assert projiziere(["Ganz anderer Text"], laeufe, {"A1"}) is None           # passt nicht: bei den Operationen bleiben


def test_projektion_nicht_angewendete_aenderung_bleibt_normal():
    laeufe = [Lauf("k", "Die"), Lauf("s", " "), Lauf("d", "alte", "A1"), Lauf("s", " "), Lauf("a", "neue", "A1"),
              Lauf("s", " "), Lauf("k", "Pumpe"), Lauf("s", " "), Lauf("a", "heute", "A2")]
    proj = projiziere(["Die alte Pumpe"], laeufe, {"A2"})
    m = markiere_ops([Stueck(NORMAL, "Die alte Pumpe")], [], bereiche=[(0, 14, proj[0])])
    assert _kurz(m.stuecke) == "Die alte Pumpe {+heute+}"


def test_projektion_neben_operationen_in_derselben_zelle():
    t = "Erster Satz alt.\nZweiter Satz."
    proj = projiziere([t[:16]], _lauf("Erster Satz {-alt-}|{+neu+}|."), {"A1"})
    op = Operation("ersetzen", "CDE", 2, 3, len(t), len(t), "\nDritter Satz.", seq=1, aenderung="A2")  # neuer Absatz ohne PDF-Absatz
    m = markiere_ops([Stueck(NORMAL, t)], [op], bereiche=[(0, 16, proj[0])])
    assert _kurz(m.stuecke) == "Erster Satz {-alt-}{+neu+}.\nZweiter Satz.\n{+Dritter Satz.+}"
    with pytest.raises(ValueError):                                         # Operation mitten im projizierten Bereich
        markiere_ops([Stueck(NORMAL, t)], [Operation("ersetzen", "CDE", 2, 3, 3, 5, "x", aenderung="A3")],
                     bereiche=[(0, 16, proj[0])])


def _pruefe_projektion(excel: str, laeufe: list[Lauf], soll: str) -> None:
    proj = projiziere([excel], laeufe, {"A1"})
    assert proj is not None, excel
    m = markiere_ops([Stueck(NORMAL, excel)], [], bereiche=[(0, len(excel), proj[0])])
    assert _kurz(m.stuecke) == soll, excel
    assert "".join(z[0] for z in m.zeichen if z[2] >= 0) == excel              # jedes Excel-Zeichen steht unveraendert da
    ab = Absatz("P1", "CDE", [(2, 3, 0, len(excel))], laeufe)
    assert pruefe_absatz_markiert(ab, {(2, 3): (excel, m.zeichen)}, {"A1"}) is None


def test_a3_markiert_ignoriert_aeltere_unsichtbare_markierungen():
    """Welle 3, Punkt 4: eine Zelle mit durchgestrichenem Text einer frueheren Runde (in der Neu-Ansicht unsichtbar). A3
    vergleicht nur, was zur Neu-Ansicht gehoert -- das alte Gestrichene darf kein gegenprobe_alt ausloesen."""
    vorher = [Stueck(NORMAL, "Die "), Stueck(ALT, "alte "), Stueck(NORMAL, "Pumpe läuft "), Stueck(ALT, "sehr "),
              Stueck(NORMAL, "leise.")]
    t = sicht_text(vorher, SICHT_NEU)
    assert t == "Die Pumpe läuft leise."
    s = t.index("leise")
    laeufe = _lauf("Die Pumpe läuft {-leise-}|{+laut+}|.")
    m = markiere_ops(vorher, _ops(("e", s, s + 5, "laut", "A1")), sicht=SICHT_NEU)
    assert [z for z in m.zeichen if z[2] < 0 and not z[4]]                  # die alten Zeichen stehen mit in der Liste
    ab = Absatz("P1", "CDE", [(2, 3, 0, len(t))], laeufe)
    assert pruefe_absatz_markiert(ab, {(2, 3): (t, m.zeichen)}, {"A1"}) is None
    proj = projiziere([t], laeufe, {"A1"})
    m2 = markiere_ops(vorher, [], sicht=SICHT_NEU, bereiche=[(0, len(t), proj[0])])
    assert pruefe_absatz_markiert(ab, {(2, 3): (t, m2.zeichen)}, {"A1"}) is None
    # ein eigener Fehler faellt weiter auf: der neue Text fehlt
    m3 = markiere_ops(vorher, _ops(("e", s, s + 5, "", "A1")), sicht=SICHT_NEU)
    assert pruefe_absatz_markiert(ab, {(2, 3): (t, m3.zeichen)}, {"A1"}) == "gegenprobe_text"


def test_projektion_einfuegung_hinter_den_ganzen_excel_leerraum():
    """Welle 3, Punkt 2: Listennummer + viele geschuetzte Leerzeichen; das PDF fuegt am Absatzanfang ein. Der neue Text kommt
    hinter den ganzen Excel-Leerraum (vor das naechste behaltene Zeichen) -- ausser das PDF fuegt ausdruecklich vor dem
    Trenner ein (der neue Text klebt am Zeichen davor)."""
    nb = "\xa0" * 17
    excel = "3)" + nb + " Allgemeine Regeln."
    k, a, sp = (lambda t: Lauf("k", t)), (lambda t: Lauf("a", t, "A1")), Lauf("s", " ")
    # markierter Leerraum hinter dem neuen Wort (Teil des a-Laufs)
    _pruefe_projektion(excel, [k("3)"), sp, a("Weitere "), k("Allgemeine"), sp, k("Regeln.")],
                       "3)" + nb + " {+Weitere +}Allgemeine Regeln.")
    # unmarkierter Trenner auf beiden Seiten: er bleibt normal, nichts wird mitten in den Excel-Leerraum gesetzt
    _pruefe_projektion(excel, [k("3)"), sp, a("Weitere"), sp, k("Allgemeine"), sp, k("Regeln.")],
                       "3)" + nb + " {+Weitere+} Allgemeine Regeln.")
    # das PDF fuegt ausdruecklich VOR dem Trenner ein
    _pruefe_projektion(excel, [k("3)"), a("a"), sp, k("Allgemeine"), sp, k("Regeln.")],
                       "3){+a+}" + nb + " Allgemeine Regeln.")


def test_projektion_ersetzte_listennummer_mit_polster_und_einfuegung():
    """Nachauftrag: die Listennummer selbst wird ersetzt und das PDF fuegt Text dahinter ein. Die neue Nummer steht direkt
    hinter der alten, der neue Text HINTER dem Polster (geschuetzte Leerzeichen bzw. Tab), nicht an der Nummer."""
    k, sp = (lambda t: Lauf("k", t)), Lauf("s", " ")
    a, d = (lambda t: Lauf("a", t, "A1")), (lambda t: Lauf("d", t, "A1"))
    nb = "\xa0" * 12
    for excel, laeufe, soll in [
        ("d." + nb + " Ausführung des Ventils.", [d("d."), sp, a("f."), a(" "), a("Besondere"), a(" "), k("Ausführung"), sp,
                                                  k("des"), sp, k("Ventils.")],
         "{-d.-}{+f.+}" + nb + " {+Besondere +}Ausführung des Ventils."),
        ("3.\tAllgemeine Regeln.", [d("3."), sp, a("4. Weitere "), k("Allgemeine"), sp, k("Regeln.")],
         "{-3.-}{+4.+}\t{+Weitere +}Allgemeine Regeln."),
    ]:
        proj = projiziere([excel], laeufe, {"A1"})
        assert proj is not None
        m = markiere_ops([Stueck(NORMAL, excel)], [], bereiche=[(0, len(excel), proj[0])])
        assert _kurz(m.stuecke) == soll
        assert "".join(z[0] for z in m.zeichen if z[2] >= 0) == excel
        ab = Absatz("P1", "CDE", [(2, 3, 0, len(excel))], laeufe)
        assert pruefe_absatz_markiert(ab, {(2, 3): (excel, m.zeichen)}, {"A1"}) is None
    # nur ein einzelnes Leerzeichen hinter der Nummer ist kein Polster: der Text bleibt an der Nummer
    excel = "3. Allgemeine Regeln."
    proj = projiziere([excel], [d("3."), sp, a("4. Weitere "), k("Allgemeine"), sp, k("Regeln.")], {"A1"})
    m = markiere_ops([Stueck(NORMAL, excel)], [], bereiche=[(0, len(excel), proj[0])])
    assert _kurz(m.stuecke) == "{-3.-} {+4. Weitere +}Allgemeine Regeln."


def test_projektion_markierter_leerraum_aus_dem_pdf():
    """Welle 3, Punkt 3: LESEN liefert markierten Leerraum als Teil der d/a-Laeufe. Er wird genau so markiert: ein
    gestrichenes Leerzeichen ist das vorhandene Excel-Leerzeichen (hellblau + Strich), ein eingefuegtes ist rot -- ein
    vorhandenes Excel-Leerzeichen wird nie gedoppelt. Unmarkierter Leerraum bleibt unmarkiert. Woerter erfunden."""
    k, sp = (lambda t: Lauf("k", t)), Lauf("s", " ")
    a, d = (lambda t: Lauf("a", t, "A1")), (lambda t: Lauf("d", t, "A1"))
    faelle = [
        ("Für Rückschlagventile gilt das.", [k("Für"), sp, a("zusätzlich "), k("Rückschlagventile"), sp, k("gilt"), sp,
                                              k("das.")], "Für {+zusätzlich +}Rückschlagventile gilt das."),
        ("Ersatzdichtungen für Lager.", [k("Ersatzdichtungen"), d(" für"), sp, k("Lager.")],
         "Ersatzdichtungen{- für-} Lager."),
        ("Förderhöhe: [40] m.", [k("Förderhöhe:"), d(" [40] m.")], "Förderhöhe:{- [40] m.-}"),
        ("Das Rohr wird gespült.", [k("Das"), sp, k("Rohr"), sp, k("wird"), sp, k("gespült."),
                                    a(" Die Leistungsaufnahme der Pumpe ist zu messen.")],
         "Das Rohr wird gespült.{+ Die Leistungsaufnahme der Pumpe ist zu messen.+}"),
        # nie doppeln: Excel hat den Leerraum schon (davor bzw. dahinter), der markierte Rand-Leerraum faellt weg
        ("Für Rückschlagventile.", [k("Für"), sp, a(" zusätzlich"), sp, k("Rückschlagventile.")],
         "Für {+zusätzlich+} Rückschlagventile."),
        ("Das Rohr wird gespült. ", [k("Das"), sp, k("Rohr"), sp, k("wird"), sp, k("gespült."), a(" Neu ist das.")],
         "Das Rohr wird gespült.{+ Neu ist das.+} "),
        # unmarkierter Leerraum bleibt unmarkiert
        ("Für Rückschlagventile.", [k("Für"), sp, a("zusätzlich"), sp, k("Rückschlagventile.")],
         "Für {+zusätzlich+} Rückschlagventile."),
        # markierter Trenner als eigener d/a-Lauf (trenner_art des linken Worts): dasselbe Ergebnis
        ("Für Rückschlagventile gilt.", [k("Für"), sp, a("zusätzlich"), a(" "), k("Rückschlagventile"), sp, k("gilt.")],
         "Für {+zusätzlich +}Rückschlagventile gilt."),
        ("Ersatzdichtungen für Lager.", [k("Ersatzdichtungen"), d(" "), d("für"), sp, k("Lager.")],
         "Ersatzdichtungen{- für-} Lager."),
    ]
    for excel, laeufe, soll in faelle:
        _pruefe_projektion(excel, laeufe, soll)


def test_operationen_ohne_projektion_markieren_keinen_rand_leerraum():
    # Rueckfallweg (kein PDF-Absatz): gemeinsamer Anfang/Ende an Leerraum-Grenzen und Leerraum am Rand bleiben unmarkiert
    fall = [("nach  NORM 4711.", ("e", 4, 6, " ZUSATZ-EN ", "A1"), "nach {+ZUSATZ-EN+} NORM 4711."),
            ("Feuchte: [0 bis 100] Prozent", ("e", 8, 10, " ", "A1"), "Feuchte: {-[-}0 bis 100] Prozent"),
            ("und NORM 4712. Danach", ("e", 3, 13, "", "A1"), "und {-NORM 4712-}. Danach"),
            ("Teil 11223‑6 mit Pumpe alt.", ("e", 5, 26, "11223-6 mit Pumpe neu", "A1"),
             "Teil 11223‑6 mit Pumpe {-alt-}{+neu+}."),
            ("Wert und Rest", ("e", 0, 4, "Werte", "A1"), "{-Wert-}{+Werte+} und Rest"),        # ganzes Wort wie im PDF
            ("a  b", ("e", 1, 3, " ", "A1"), "a  b")]                                       # reiner Leerraum: nie markiert
    for text, spec, soll in fall:
        assert _kurz(markiere_ops([Stueck(NORMAL, text)], _ops(spec)).stuecke) == soll, text


def test_zufall_projektion_excel_bleibt_pdf_markierung_genau():
    """Zufaellige Absaetze mit Excel-eigenen Schreibweisen (doppelter Leerraum, U+2011, geschuetztes Leerzeichen): jedes
    Excel-Zeichen bleibt, markiert ist genau der Text der PDF-Laeufe, kein Leerraum am Rand einer Markierung, A3 besteht."""
    rnd = random.Random(23)
    woerter = ["Ventil", "Pumpe", "Druck", "11223-6", "Anlage", "[20]", "Zebrafink", "Wert:", "Prüfung"]
    for _ in range(400):
        laeufe: list[Lauf] = []
        for i in range(rnd.randint(1, 8)):
            if laeufe:
                laeufe.append(Lauf("s", " "))
            art = rnd.choice("kkkkda")
            if art == "d" and rnd.random() < 0.4:                # Ersetzung, mal mit, mal ohne Leerraum dazwischen
                laeufe += [Lauf("d", rnd.choice(woerter), "A1")] + ([Lauf("s", " ")] if rnd.random() < 0.5 else []) \
                    + [Lauf("a", rnd.choice(woerter), "A1")]
            else:
                laeufe.append(Lauf(art, rnd.choice(woerter), "A1" if art in "da" else ""))
        if not any(l.art in "da" for l in laeufe):
            continue
        alt_pdf = zustand(laeufe, set())
        excel = "".join(rnd.choice([" ", "  ", " "]) if c == " " else ("‑" if c == "-" and rnd.random() < 0.5
                                                                             else c) for c in alt_pdf)
        proj = projiziere([excel], laeufe, {"A1"})
        assert proj is not None, (excel, alt_pdf)
        m = markiere_ops([Stueck(NORMAL, excel)], [], bereiche=[(0, len(excel), proj[0])])
        assert "".join(z[0] for z in m.zeichen if z[2] >= 0) == excel
        neu = "".join(z[0] for z in m.zeichen if z[1] == NEU)
        assert neu.replace(" ", "") == "".join(l.text for l in laeufe if l.art == "a")
        assert "".join(z[0] for z in m.zeichen if z[1] == ALT).replace(" ", "").replace(" ", "").replace("‑", "-") \
            == "".join(l.text for l in laeufe if l.art == "d")
        for s in m.stuecke:
            if s.format != NORMAL:
                assert s.text == s.text.strip(), (excel, _kurz(m.stuecke))
        ab = Absatz("P1", "CDE", [(2, 3, 0, len(excel))], laeufe)
        assert pruefe_absatz_markiert(ab, {(2, 3): (excel, m.zeichen)}, {"A1"}) is None, (excel, _kurz(m.stuecke))
