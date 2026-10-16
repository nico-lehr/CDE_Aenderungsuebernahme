"""Steckbrief (src/steckbrief.py): nur Zahlen, Farbcodes und Formen. Nur erfundene Daten (Regel 16).

Wichtigste Tests: die Positivliste (`pruefe`), die Kanarienwort-Tests (kein markantes Wort aus PDF oder Excel darf im Ergebnis,
in der Markdown-Datei, in der Zusammenfassung oder in der Konsole auftauchen), die Zaehlung der Muster/Legende an handgebauten
Beispielen und `messen --schablonen`.
"""
import json
import random
import re
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from src import cli, config, messen, steckbrief
from src.markierung import ALT, NEU, NORMAL, Stueck, schreibe
from src.steckbrief import FORM_ZEICHEN, fenster, formel_muster, pruefe, schablone
from tests.fixtures.pdf_writer import PdfDokument
from tests.fixtures.szenario import baue_excel, zustaende
from tests.test_messen import MIT_NEUER_SECTION, SAUBER, baue_paar, sabotiere_ergebnis, zellen_aendern

NBSP = chr(0xA0)
NNBSP = chr(0x202F)


def steckbrief_von(p: Path, name: str = "probe") -> dict:
    """Baut den Steckbrief eines Paar-Ordners im Speicher."""
    return steckbrief.erzeuge(p / "aenderungen.pdf", p / "vorher.xlsx", p / "nachher.xlsx", name)


def lauf(argv: list[str]) -> int:
    return cli.main(["steckbrief", *argv])


def kanarien_treffer(text: str, woerter: list[str]) -> list[str]:
    """Kanarienwoerter (auch Teile ab 4 Zeichen, auch kleingeschrieben) im Text."""
    t = text.lower()
    out = []
    for w in woerter:
        w = w.lower()
        teile = {w} | {w[i:i + 4] for i in range(len(w) - 3)}
        out += [x for x in sorted(teile) if x in t]
    return out


# ---------------------------------------------------------------------------------------------------------------
# Formen
# ---------------------------------------------------------------------------------------------------------------
def test_schablone_buchstaben_ziffern_umlaute():
    assert schablone("Der Zug 12, ÄÖÜ äöü ß!") == "Aaa Aaa 99, AAA aaa a!"
    assert schablone("kWh-Zähler (3,5 Mm²)") == "aAa-Aaaaaa (9,9 Aa?)"


def test_schablone_leerraum_bleibt_sichtbar():
    assert schablone("a" + NBSP + "b" + NNBSP + "c") == "a·a·a"
    assert schablone("a\tb") == "a⇥a"
    assert schablone("a\nb\r\nc\rd") == "a¶a¶a¶a"
    assert schablone("a  b") == "a  a"                        # doppeltes Leerzeichen bleibt doppelt


def test_schablone_andere_zeichen_und_lange_woerter():
    assert schablone("é€中😀") == "????"
    assert schablone("x" * 29) == "a" * 29
    assert schablone("x" * 30 + " ab") == "W… aa"
    assert schablone("eins zwei drei vier", max_woerter=2) == "aaaa aaaa …"


def test_schablone_kennt_nur_den_formen_zeichensatz():
    rnd = random.Random(5)
    for _ in range(300):
        t = "".join(chr(rnd.choice([rnd.randint(0, 0x2FF), rnd.randint(0x2000, 0x20FF), rnd.randint(0x1F300, 0x1F3FF), 32, 10, 160])) for _ in range(rnd.randint(0, 600)))
        s = schablone(t)
        assert set(s) <= FORM_ZEICHEN and len(s) <= steckbrief.FORM_MAX


def test_fenster_begrenzt_auf_zwoelf_woerter_und_kennzeichnet_die_stellen():
    st = [(NORMAL, "Der Quokka frisst "), (ALT, "Gras"), (NORMAL, " "), (NEU, "Laub"), (NORMAL, " am Morgen. Dann ruht er lange im Schatten und noch mehr dahinter.")]
    laeufe = fenster(st, 18)
    assert [k for k, _ in laeufe] == ["-", "a", "-", "n", "-"]
    form = "".join(f for _, f in laeufe)
    assert form.endswith("…") and not form.startswith("…")
    assert len(re.findall(r"[Aa9]+", form)) <= 12
    assert laeufe[1][1] == "Aaaa" and laeufe[3][1] == "Aaaa"


def test_fenster_kuerzt_lange_woerter_und_kennt_den_anfang():
    laeufe = fenster([(NORMAL, "a" * 40 + " b")], 0)
    assert laeufe == [("-", "W… a")]
    text = " ".join(f"w{i:02d}" for i in range(30))
    laeufe = fenster([(NORMAL, text)], text.index("w15"), davor=1)
    assert laeufe[0][1].startswith("…") and laeufe[-1][1].endswith("…")
    assert laeufe[0][1] == "…" + " ".join("a99" for _ in range(12)) + "…".replace("…", "…") or len(laeufe) == 1


def test_formel_muster_relativ_zur_zeile():
    assert formel_muster("=LEN(C5)", 5) == ("=LEN(C{r})", ["LEN"])
    assert formel_muster("=SUM(D6:D9)", 5)[0] == "=SUM(D{r+1}:D{r+4})"
    assert formel_muster("=$A$1+B3", 5)[0] == "=$A$9+B{r-2}"


def test_formel_muster_enthaelt_nie_ein_wort_aus_der_formel():
    m, f = formel_muster('=IF(A5="Zebrafink";Quarzblatt!B2*1,19;Zebrafunk(3))+Tresorkiste', 5)
    assert not kanarien_treffer(m, ["Zebrafink", "Quarzblatt", "Zebrafunk", "Tresorkiste"])
    assert m == '=IF(A{r}="T";@!B{r-3}*9;FUNC(9))+NAME' and f == ["IF", "FUNC"]
    assert steckbrief.Formel().pruefe(m, "x", e := []) is None and not e


# ---------------------------------------------------------------------------------------------------------------
# Positivliste
# ---------------------------------------------------------------------------------------------------------------
GUELTIG = {"art": "steckbrief", "format": 1, "name": "probe", "pdf": {"seiten": 2, "schriften": {"Arial": 10}},
           "excel_vorher": {"blaetter": 1}, "excel_nachher": {"muster": {"normal,alt,neu,normal": 3}, "trenner": {" ": 2, "kein": 1}},
           "schablonen": [{"typ": "text_ersetzt", "status": "pruefen", "pdf_alt": [{"f": "-", "form": "Aaa a:a"}],
                           "pdf_neu": [], "excel_vorher": [], "excel_nachher": []}]}


def kopie(**ersatz):
    d = json.loads(json.dumps(GUELTIG))
    d.update(ersatz)
    return d


def test_gueltiger_steckbrief_besteht():
    assert pruefe(GUELTIG) == []


def test_sabotage_unbekanntes_textfeld_schlaegt_an():
    d = kopie()
    d["pdf"]["titel"] = "Zebrafink Quarzuhr"
    fehler = pruefe(d)
    assert fehler and not kanarien_treffer("\n".join(fehler), ["Zebrafink", "Quarzuhr", "titel"])
    d = kopie(notiz="Zebrafink")
    assert pruefe(d)


@pytest.mark.parametrize("aendere", [
    lambda d: d["pdf"]["schriften"].update({"Die Quokka frisst Gras.": 1}),                        # Schluessel ausserhalb des Musters
    lambda d: d["pdf"]["schriften"].update({"Zebrafink": "viel"}),                                 # Text statt Zahl
    lambda d: d["pdf"].update(seiten="2"),                                                          # Zahl als Text
    lambda d: d["pdf"].update(seiten=True),                                                         # Wahrheitswert ist keine Zahl
    lambda d: d["pdf"].update(seiten=-1),
    lambda d: d["excel_nachher"]["muster"].update({"normal,zebra": 1}),
    lambda d: d["excel_nachher"]["trenner"].update({"Zebra": 1}),                                   # Kleinbuchstaben sind keine Form
    lambda d: d["schablonen"][0]["pdf_alt"][0].update(form="Zebrafink"),
    lambda d: d["schablonen"][0]["pdf_alt"][0].update(form="A" * 401),
    lambda d: d["schablonen"][0]["pdf_alt"][0].update(form="Aaé"),                             # ein Zeichen ausserhalb des Zeichensatzes
    lambda d: d["schablonen"][0]["pdf_alt"][0].update(f="x"),
    lambda d: d["schablonen"][0].update(typ="Zebrafink"),
    lambda d: d["schablonen"][0].update(status="irgendwas"),
    lambda d: d.update(name="Zebrafink"),
    lambda d: d.update(art="anderes"),
    lambda d: d["excel_nachher"].update(muster=["Zebrafink"]),
    lambda d: d["pdf"].update(farben={"DeviceRGB|#12345": {"woerter": 1}}),
    lambda d: d["excel_vorher"].update(formeln={"muster": {"=ZEBRAFINK(A{r})": 1}}),
    lambda d: d["excel_vorher"].update(konstanten_spalten=[{"blatt": 1, "spalte": "AAAA", "anteil": 0.9, "zellen": 5, "form": "A"}]),
    lambda d: d["excel_vorher"].update(konstanten_spalten=[{"blatt": 1, "spalte": "A", "anteil": float("nan"), "zellen": 5, "form": "A"}]),
])
def test_positivliste_lehnt_ab_ohne_den_inhalt_zu_wiederholen(aendere):
    d = kopie()
    aendere(d)
    fehler = pruefe(d)
    assert fehler, "Verstoß nicht erkannt"
    assert not kanarien_treffer("\n".join(fehler), ["Zebrafink", "Quokka", "Gras"])


def test_pruefe_ist_kein_text_filter_sondern_positivliste_auch_fuer_zahlen_und_typen():
    assert pruefe([]) and pruefe("text") and pruefe(None)
    d = kopie()
    d["pdf"]["seite"] = [{"zeichen": 1, "woerter": 2, "bilder": 0, "text": "x"}]
    assert pruefe(d)


def test_schreiber_fuer_formen_besteht_nur_mit_erlaubtem_zeichensatz():
    assert not [c for c in schablone("Ünï©ödé 𝔘𝔫𝔦 中文 ​­") if c not in FORM_ZEICHEN]


# ---------------------------------------------------------------------------------------------------------------
# Steckbrief eines Paars
# ---------------------------------------------------------------------------------------------------------------
def test_steckbrief_eines_einfachen_paars_besteht_die_positivliste_und_hat_die_haupt_teile(tmp_path):
    p = baue_paar(tmp_path, "paar", MIT_NEUER_SECTION)
    sb = steckbrief_von(p)
    assert pruefe(sb) == []
    assert set(sb) >= {"art", "format", "name", "pdf", "excel_vorher", "excel_nachher", "tool", "schablonen"}
    assert sb["pdf"]["seiten"] == 1 and sb["pdf"]["seiten_masse"] == {"595x842": 1}
    assert "DeviceRGB|#0000ff" in sb["pdf"]["farben"] and "DeviceRGB|#ff0000" in sb["pdf"]["farben"]
    assert sb["pdf"]["farben"]["DeviceRGB|#ff0000"]["durchgestrichen"] == sb["pdf"]["farben"]["DeviceRGB|#ff0000"]["woerter"] > 0
    assert sb["pdf"]["farben"]["DeviceRGB|#0000ff"]["durchgestrichen"] == 0
    assert sb["pdf"]["striche"]["anzahl"] > 0
    assert sb["tool"]["aenderungen"] == 5 and sb["tool"]["je_typ"]["text_ersetzt"] == 2
    assert sb["excel_vorher"]["blaetter"] == 2 and sb["excel_vorher"]["zellen_mit_mehreren_absaetzen"] >= 1
    blatt = {b["nr"]: b for b in sb["excel_vorher"]["blatt"]}[2]
    assert blatt["rollen"] == {"section": "A", "titel": "B", "inhalt": "C"} and blatt["sicherheit"] == "hoch"
    assert sb["excel_nachher"]["zeilen"]["neu"] == 1 and sb["excel_nachher"]["zellen"]["sauber_ersetzt"] == 3
    assert 1 <= len(sb["schablonen"]) <= 40


def test_schablonen_beispiele_sind_deterministisch_gestreut_und_hoechstens_vierzig(tmp_path):
    elemente = [("h", f"{i} Titel {i}") if k == 0 else ("p", f"Ein Absatz {i} mit {{-alt{i}-}} {{+neu{i}+}} im Satz.")
                for i in range(1, 31) for k in (0, 1)]
    p = baue_paar(tmp_path, "viele", elemente)
    a, b = steckbrief_von(p)["schablonen"], steckbrief_von(p)["schablonen"]
    assert a == b and len(a) <= 40
    assert len(a) == 30                                              # weniger Aenderungen als Platz: alle
    elemente = elemente + [("p", "{+Neu " + str(i) + " Absatz im Text.+}") for i in range(30)]
    p2 = baue_paar(tmp_path, "mehr", elemente)
    s2 = steckbrief_von(p2)["schablonen"]
    assert len(s2) == 40 and {x["typ"] for x in s2} >= {"text_ersetzt", "absatz_neu"}   # Typen gemischt, nicht nur der erste


def test_konstanten_spalten_und_neue_zeile_uebernimmt_konstante(tmp_path):
    elemente = [(a, b) for i in range(1, 7) for a, b in (("h", f"{i} Titel"), ("p", f"Absatz {i} bleibt hier stehen."))]
    elemente += [("h", "{+7 Neu+}"), ("p", "{+Nur ein neuer Absatz.+}")]
    p = baue_paar(tmp_path, "konst", elemente)
    sb = steckbrief_von(p)
    assert pruefe(sb) == []
    k = [x for x in sb["excel_vorher"]["konstanten_spalten"] if x["spalte"] == "D"]
    assert k and k[0]["anteil"] == 1.0 and k[0]["form"] == "Aaaaaa aaaaa"
    assert sb["excel_nachher"]["neue_zeilen"]["konstanten_uebernommen"] == 1
    assert sb["excel_nachher"]["neue_zeilen"]["konstanten_fehlen"] == 0


def test_formeln_und_excel_eigenschaften_werden_gezaehlt(tmp_path):
    p = baue_paar(tmp_path, "form", SAUBER)
    for name in ("vorher.xlsx", "nachher.xlsx"):
        wb = load_workbook(str(p / name))
        ws = wb["CDE"]
        ws["F2"] = "=LEN(C2)"
        ws["F3"] = "=LEN(C3)"
        ws["F5"] = "=SUM(F2:F3)"
        ws.merge_cells("G2:H3")
        ws.row_dimensions[3].hidden = True
        wb.save(str(p / name))
    ev = steckbrief_von(p)["excel_vorher"]
    assert ev["formeln"]["anzahl"] == 3 and ev["formeln"]["muster"]["=LEN(C{r})"] == 2
    assert ev["formeln"]["muster"]["=SUM(F{r-3}:F{r-2})"] == 1 and ev["formeln"]["funktionen"] == {"LEN": 2, "SUM": 1}
    assert ev["verbundene_zellen"] == {"2x2": 1} and ev["versteckte_zeilen"] == 1


# ---------------------------------------------------------------------------------------------------------------
# Wie sind Aenderungen in der Excel nachher formatiert?
# ---------------------------------------------------------------------------------------------------------------
def markiere_nachher(p: Path) -> None:
    """Handgebaute Markierung in nachher.xlsx: alt = blau + durchgestrichen, neu = rot (Standardfarben aus config/zuordnung.yaml)."""
    vor = load_workbook(str(p / "vorher.xlsx"))["CDE"]
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    v = vor["C2"].value
    i = v.index("Gras")
    schreibe(ws["C2"], [Stueck(NORMAL, v[:i]), Stueck(ALT, "Gras"), Stueck(NORMAL, " "), Stueck(NEU, "Laub"), Stueck(NORMAL, v[i + 4:])])
    v = vor["C3"].value
    i = v.index("links")
    schreibe(ws["C3"], [Stueck(NORMAL, v[:i]), Stueck(NEU, "rechts"), Stueck(NORMAL, "\n"), Stueck(ALT, "links"), Stueck(NORMAL, v[i + 5:])])
    for adr in ("A4", "B4", "C4"):                                         # ganz gestrichene Section: Zeile bleibt
        schreibe(ws[adr], [Stueck(ALT, str(vor[adr].value))])
    ws["E2"] = "Wert 12 kWh"                                               # Spalte ohne Rolle, ohne Trenner
    schreibe(ws["E2"], [Stueck(NORMAL, "Wert "), Stueck(ALT, "12"), Stueck(NEU, "15"), Stueck(NORMAL, " kWh")])
    wb.save(str(p / "nachher.xlsx"))
    wv = load_workbook(str(p / "vorher.xlsx"))
    wv["CDE"]["E2"] = "Wert 12 kWh"
    wv.save(str(p / "vorher.xlsx"))


def test_muster_trenner_reihenfolge_und_farben_an_handgebauter_nachher_excel(tmp_path):
    p = baue_paar(tmp_path, "markiert", SAUBER)
    markiere_nachher(p)
    sb = steckbrief_von(p)
    assert pruefe(sb) == []
    en = sb["excel_nachher"]
    assert en["muster"]["normal,alt,normal,neu,normal"] == 1
    assert en["muster"]["normal,neu,normal,alt,normal"] == 1
    assert en["muster"]["normal,alt,neu,normal"] == 1
    assert en["muster"]["alt"] == 3                                        # A4, B4, C4
    assert en["trenner"] == {" ": 1, "¶": 1, "kein": 1}
    assert en["trenner_ort"] == {"eigenes_stueck": 2, "kein": 1}
    assert en["reihenfolge"] == {"alt_vor_neu": 2, "neu_vor_alt": 1}
    assert en["stueck_farben"]["alt"] == {"#0000ff": 6} and en["stueck_farben"]["neu"] == {"#ff0000": 3}
    assert en["strich_anteil"]["blau"] == 1.0 and en["strich_anteil"]["rot"] == 0.0
    assert en["farbe_strich"]["blau|strich"] == 6 and en["farbe_strich"]["rot|-"] == 3
    assert en["geaenderte_rollen"] == {"section": 1, "titel": 1, "inhalt": 3, "andere": 1}
    assert en["geaenderte_spalten"]["E"] == 1 and en["geaenderte_spalten"]["C"] == 3
    assert en["zellen"]["alt_ansicht_stimmt"] == {"ja": 4, "nur_leerraum": 2}     # C2 und C3: der Trenner zwischen neu und alt ist zusaetzlicher Leerraum
    g = en["gestrichene_zeilen"]
    assert g["anzahl"] == 1 and g["alle_durchgestrichen"] == 1 and g["inhalt_noch_da"] == 1
    assert g["zellfarben"] == {"#0000ff": 3}
    assert en["zeilen"]["gestrichen_markiert"] == 1 and en["zeilen"]["entfernt"] == 0
    assert en["masse"]["zeilenhoehe_geaendert_ja"] is False and en["masse"]["spaltenbreite_geaendert_ja"] is False


def test_neue_zeile_in_rot_und_sauber_ersetzte_zellen_und_andere_formatierung(tmp_path):
    p = baue_paar(tmp_path, "neuzeile", MIT_NEUER_SECTION)
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    for c in ws[5]:
        if c.value is not None:
            schreibe(c, [Stueck(NEU, str(c.value))])
    from openpyxl.styles import PatternFill
    ws["C3"].fill = PatternFill("solid", fgColor="FFFF00")                  # gelb hinterlegt
    ws["B2"].font = Font(color="FF00B050")                                  # gruene Schrift ohne Markierung
    wb.save(str(p / "nachher.xlsx"))
    en = steckbrief_von(p)["excel_nachher"]
    n = en["neue_zeilen"]
    assert n["anzahl"] == 1 and n["alle_stuecke_neu"] == 1
    assert n["zellfarben"] == {"#ff0000": 4} and n["spalten_befuellt"] == {"A": 1, "B": 1, "C": 1, "D": 1}
    assert en["formatierung"]["zellen_neu_gefuellt"] == 1 and en["formatierung"]["fuellfarben"] == {"#ffff00": 1}
    assert en["formatierung"]["zellen_andere_schriftfarbe"] == 1 and en["formatierung"]["schriftfarben"] == {"#00b050": 1}
    assert en["zellen"]["sauber_ersetzt"] == 3 and en["zellen"]["mit_markierung"] == 0
    assert en["muster"] == {} and n["muster"] == {"neu": 4}


def test_zeilenhoehe_und_spaltenbreite_aenderung_wird_erkannt(tmp_path):
    p = baue_paar(tmp_path, "mass", SAUBER)
    wb = load_workbook(str(p / "nachher.xlsx"))
    wb["CDE"].row_dimensions[2].height = 40
    wb["CDE"].column_dimensions["C"].width = 99
    wb.save(str(p / "nachher.xlsx"))
    m = steckbrief_von(p)["excel_nachher"]["masse"]
    assert m == {"zeilenhoehe_geaendert": 1, "spaltenbreite_geaendert": 1, "zeilenhoehe_geaendert_ja": True,
                 "spaltenbreite_geaendert_ja": True}


# ---------------------------------------------------------------------------------------------------------------
# Legende im PDF
# ---------------------------------------------------------------------------------------------------------------
def pdf_legende(tmp_path: Path, markup: str, variante: str = "farbe") -> dict:
    d = PdfDokument(variante=variante)
    d.ueberschrift("1 Eins")
    d.block(markup)
    pfad = tmp_path / f"legende_{variante}.pdf"
    d.speichern(pfad)
    return steckbrief._pdf_zahlen(pfad)


def test_legende_zaehlt_farbklasse_mal_strich_mal_unterstreichung(tmp_path):
    z = pdf_legende(tmp_path, "Vorweg {+blauneu+} {@r:rotneu@} {@bs:blaustrich@} {-rotstrich-} {@s:schwarzstrich@} {@l:linkwort@} Ende")
    leg = z["legende"]
    assert leg["blau|-|-|-"] == 1                 # blau ohne Strich = neu
    assert leg["rot|-|-|-"] == 1                  # rot ohne Strich
    assert leg["blau|strich|-|-"] == 1            # blau + Strich = entfaellt
    assert leg["rot|strich|-|-"] == 1             # rot + Strich = entfaellt
    assert leg["schwarz|strich|-|-"] == 1         # schwarz + Strich = entfaellt
    assert leg["blau|-|unterstrichen|-"] == 1     # Link
    assert "schwarz|-|-|-" not in leg             # gewoehnlicher schwarzer Text gehoert nicht in die Legende
    f = z["farben"]
    assert f["DeviceRGB|#ff0000"]["woerter"] == 2 and f["DeviceRGB|#ff0000"]["durchgestrichen"] == 1
    assert f["DeviceRGB|#0000ff"]["woerter"] == 2 and f["DeviceRGB|#0000ff"]["durchgestrichen"] == 1
    assert z["striche"]["anzahl"] == 3 and z["unterstreichungen"]["anzahl"] == 1
    assert z["striche"]["farbe_wie_text"] == 3 and z["striche"]["farbe_anders"] == 0       # jeder Strich hat die Farbe seines Wortes
    assert set(z["striche"]["hoehenband"]) == {"0.5"} and z["unterstreichungen"]["hoehenband"].keys() == {"0.9"}
    assert pruefe(kopie(pdf=z)) == []


def test_legende_hinterlegung(tmp_path):
    z = pdf_legende(tmp_path, "Vorweg {+hintergrundneu+} {-hintergrundweg-} Ende", "hintergrund")
    leg = z["legende"]
    assert leg["schwarz|-|-|hinterlegt"] == 1 and leg["schwarz|strich|-|hinterlegt"] == 1
    assert len(z["hinterlegungen"]) == 2 and sum(z["hinterlegungen"].values()) == 2


def test_pdf_hyperlink_annotation_trennung_und_aufzaehlung_werden_gezaehlt(tmp_path):
    d = PdfDokument(trennung="ascii")
    d.ueberschrift("1 Eins")
    d.block("{@l:Verweis@} im Text " + "Netzanschlussbedingungen " * 12 + "und Schluss.")
    d.block("• Punkt eins")
    pfad = tmp_path / "x.pdf"
    d.speichern(pfad)
    z = steckbrief._pdf_zahlen(pfad)
    assert z["hyperlinks"] == 1 and z["annotationen"].get("Link") == 1
    assert z["silbentrennung"] >= 1
    assert z["aufzaehlungszeichen"].get("punkt", 0) >= 1


# ---------------------------------------------------------------------------------------------------------------
# Kommandozeile: Schreiben, append-only, --pruefe, mehrere Paare
# ---------------------------------------------------------------------------------------------------------------
def test_cli_schreibt_append_only_json_und_markdown_und_pruefe_bestaetigt(tmp_path, capsys):
    p = baue_paar(tmp_path, "paar", SAUBER)
    args = ["--pdf", str(p / "aenderungen.pdf"), "--vorher", str(p / "vorher.xlsx"), "--nachher", str(p / "nachher.xlsx"),
            "--name", "projekt1"]
    assert lauf(args) == 0
    assert lauf(args) == 0                                                  # zweiter Lauf: neuer Ordner, nichts ueberschrieben
    basis = config.pfad("steckbrief_dir")
    assert sorted(d.name for d in basis.iterdir()) == ["001_projekt1", "002_projekt1"]
    j = basis / "001_projekt1" / "steckbrief.json"
    assert j.is_file() and (basis / "001_projekt1" / "steckbrief.md").is_file()
    assert json.loads(j.read_text(encoding="utf-8"))["name"] == "projekt1"
    capsys.readouterr()
    assert lauf(["--pruefe", str(j)]) == 0
    assert "OK: enthält nur Zahlen, Farbcodes und Formen" in capsys.readouterr().out


def test_cli_pruefe_findet_geschmuggelten_text_und_wiederholt_ihn_nicht(tmp_path, capsys):
    p = baue_paar(tmp_path, "paar", SAUBER)
    sb = steckbrief_von(p)
    sb["pdf"]["notiz"] = "Zebrafink Quarzuhr"
    sb["excel_nachher"]["muster"]["Zebrafink,Quarzuhr"] = 2
    f = tmp_path / "sabotiert.json"
    f.write_text(json.dumps(sb), encoding="utf-8")
    assert lauf(["--pruefe", str(f)]) == 1
    out = capsys.readouterr().out
    assert "NICHT in Ordnung" in out and "NICHT" in out.upper()
    assert not kanarien_treffer(out, ["Zebrafink", "Quarzuhr", "notiz"])
    assert lauf(["--pruefe", str(tmp_path / "gibtsnicht.json")]) == 2
    kaputt = tmp_path / "kaputt.json"
    kaputt.write_text("{nicht json", encoding="utf-8")
    assert lauf(["--pruefe", str(kaputt)]) == 2


def test_cli_schreiber_schreibt_bei_verstoss_nichts(tmp_path, monkeypatch, capsys):
    p = baue_paar(tmp_path, "paar", SAUBER)
    original = steckbrief.erzeuge

    def schmuggeln(*a, **kw):
        d = original(*a, **kw)
        d["pdf"]["schriften"]["Zebrafink Quarzuhr!"] = 1
        return d
    monkeypatch.setattr(steckbrief, "erzeuge", schmuggeln)
    rc = lauf(["--pdf", str(p / "aenderungen.pdf"), "--vorher", str(p / "vorher.xlsx"), "--nachher", str(p / "nachher.xlsx")])
    out = capsys.readouterr().out
    assert rc == 3 and "NICHT geschrieben" in out
    assert not kanarien_treffer(out, ["Zebrafink", "Quarzuhr"])
    assert not config.pfad("steckbrief_dir").exists() or not list(config.pfad("steckbrief_dir").iterdir())


@pytest.mark.parametrize("name", ["Zebrafink", "mit leerzeichen", "ä", "x" * 41, "a/b", ""])
def test_cli_name_nur_kleinbuchstaben_ziffern_unterstrich_bindestrich(tmp_path, capsys, name):
    p = baue_paar(tmp_path, "paar", SAUBER)
    rc = lauf(["--pdf", str(p / "aenderungen.pdf"), "--vorher", str(p / "vorher.xlsx"), "--nachher", str(p / "nachher.xlsx"),
               "--name", name])
    out = capsys.readouterr().out
    assert rc == 2 and (name not in out or name == "")
    assert not config.pfad("steckbrief_dir").exists()


def test_cli_mehrere_paare_schreiben_zusammenfassung_mit_anonymen_namen(tmp_path, capsys):
    eltern = tmp_path / "Zebrafink_Kunde"
    baue_paar(eltern, "Quarzuhr_Projekt_A", SAUBER)
    baue_paar(eltern, "Quarzuhr_Projekt_B", MIT_NEUER_SECTION)
    (eltern / "kein_paar").mkdir()
    assert lauf(["--ordner", str(eltern), "--name", "sammlung1"]) == 0
    out = capsys.readouterr().out
    assert not kanarien_treffer(out.replace(str(config.pfad("steckbrief_dir")), ""), ["Zebrafink", "Quarzuhr"])
    ziel = config.pfad("steckbrief_dir") / "001_sammlung1"
    assert sorted(d.name for d in ziel.iterdir() if d.is_dir()) == ["paar_01", "paar_02"]
    z = json.loads((ziel / "zusammenfassung.json").read_text(encoding="utf-8"))
    assert z["art"] == "zusammenfassung" and z["paare"] == 2
    assert z["pdf"]["seiten"] == 2
    a = json.loads((ziel / "paar_01" / "steckbrief.json").read_text(encoding="utf-8"))
    b = json.loads((ziel / "paar_02" / "steckbrief.json").read_text(encoding="utf-8"))
    assert z["pdf"]["seiten"] == a["pdf"]["seiten"] + b["pdf"]["seiten"]
    assert z["tool"]["aenderungen"] == a["tool"]["aenderungen"] + b["tool"]["aenderungen"]
    assert z["excel_vorher"]["blaetter"] == a["excel_vorher"]["blaetter"] + b["excel_vorher"]["blaetter"]
    assert z["excel_nachher"]["blaetter_gleich"] == 2                      # Wahrheitswerte werden zu "Anzahl der Paare"
    assert z["excel_vorher"]["section_zahlenformat"]["General"] >= 2
    assert pruefe(z) == [] and lauf(["--pruefe", str(ziel / "zusammenfassung.json")]) == 0
    assert (ziel / "zusammenfassung.md").is_file()


def test_cli_fehlende_oder_kaputte_dateien_geben_einfache_meldungen_ohne_namen(tmp_path, capsys):
    p = baue_paar(tmp_path, "paar", SAUBER)
    kaputt = tmp_path / "Zebrafink_kaputt.pdf"
    kaputt.write_bytes(b"das ist kein pdf Quarzuhr")
    assert lauf(["--pdf", str(kaputt), "--vorher", str(p / "vorher.xlsx"), "--nachher", str(p / "nachher.xlsx")]) == 2
    assert lauf(["--pdf", str(tmp_path / "Zebrafink_fehlt.pdf"), "--vorher", str(p / "vorher.xlsx"),
                 "--nachher", str(p / "nachher.xlsx")]) == 2
    assert lauf(["--zebrafink", "x"]) == 2
    assert lauf([]) == 2
    assert lauf(["--ordner", str(tmp_path / "Zebrafink_gibtsnicht")]) == 2
    out = capsys.readouterr().out
    assert out.strip() and not kanarien_treffer(out, ["Zebrafink", "Quarzuhr"])
    assert not config.pfad("steckbrief_dir").exists()


def test_laufzeit_unter_zehn_sekunden_je_paar(tmp_path):
    p = baue_paar(tmp_path, "paar", MIT_NEUER_SECTION)
    t0 = time.monotonic()
    steckbrief_von(p)
    assert time.monotonic() - t0 < 10


# ---------------------------------------------------------------------------------------------------------------
# Kanarienwoerter: nichts aus den Dokumenten darf im Ergebnis stehen
# ---------------------------------------------------------------------------------------------------------------
KANARIEN = ["Zebrafink", "Quarzuhr", "Xyloquint", "Plumbaxo", "Vrokzeld"]


def pdf_mit_metadaten(pfad: Path, titel: str, autor: str) -> None:
    b = pfad.read_bytes()
    i = b.rindex(b"trailer")
    b = b[:i] + b[i:].replace(b"/Root", f"/Info << /Title ({titel}) /Author ({autor}) /Subject (Vrokzeld) >> /Root".encode("latin-1"), 1)
    pfad.write_bytes(b)


def baue_kanarien_paar(tmp_path: Path) -> Path:
    ordner = tmp_path / "Zebrafink_Kunde" / "Quarzuhr_Projekt"
    ordner.mkdir(parents=True)
    d = PdfDokument(kopf_text="Zebrafink Quarzuhr Kopf", link_url="https://zebrafink.example/quarzuhr")
    d.ueberschrift("1 Zebrafink Titel")
    d.block("Der Zebrafink frisst {-Gras-} {+Quarzuhr+} am Morgen. {@l:Vrokzeld@} bleibt. Der Xyloquint ist {-offen-} {+Plumbaxo+}.")
    d.block("{-Dieser Zebrafink Absatz entfällt.-}")
    d.ueberschrift("2 Xyloquint Zweiter Titel")
    d.block("Ein Absatz mit Vrokzeld bleibt unverändert.")
    d.ueberschrift("{+3 Plumbaxo+}")
    d.block("{+Ein neuer Absatz mit Quarzuhr und Xyloquint.+}")
    pdf = ordner / "Quarzuhr_aenderungen.pdf"
    d.speichern(pdf)
    pdf_mit_metadaten(pdf, "Zebrafink Titel", "Quarzuhr Autor")
    elemente = [("h", "1 Zebrafink Titel"),
                ("p", "Der Zebrafink frisst {-Gras-} {+Quarzuhr+} am Morgen. Vrokzeld bleibt. Der Xyloquint ist {-offen-} {+Plumbaxo+}."),
                ("p", "{-Dieser Zebrafink Absatz entfällt.-}"), ("h", "2 Xyloquint Zweiter Titel"),
                ("p", "Ein Absatz mit Vrokzeld bleibt unverändert."), ("h", "{+3 Plumbaxo+}"),
                ("p", "{+Ein neuer Absatz mit Quarzuhr und Xyloquint.+}")]
    for name, neu in (("Zebrafink_vorher.xlsx", False), ("Quarzuhr_nachher.xlsx", True)):
        baue_excel(ordner / name, zustaende(elemente, neu), "L1")
        wb = load_workbook(str(ordner / name))
        wb.properties.creator = "Quarzuhr Autor"
        wb.properties.title = "Zebrafink Titel"
        wb.properties.description = "Vrokzeld"
        ws = wb["CDE"]
        ws.title = "Xyloquint Blatt"
        ws["A1"], ws["B1"], ws["C1"] = "Zebrafink Kopf", "Quarzuhr Kopf", "Plumbaxo Kopf"
        ws["C2"].comment = Comment("Zebrafink Kommentar", "Quarzuhr")
        ws["D2"].hyperlink = "https://zebrafink.example/walfischbein"
        ws["D2"].value = "Vrokzeld Link"
        ws["A2"].number_format = '0.0" Zebrafink"'
        ws["B2"].font = Font(name="Quarzuhrschrift", size=13)
        ws["E2"] = '=IF(C2="Zebrafink";"Quarzuhr";"Plumbaxo")'
        ws["E3"] = '=LEN(C3)&"Vrokzeld"'
        ws["E4"] = "=Xyloquint_Blatt!A1"
        ws["F5"] = CellRichText("Zebrafink ", TextBlock(InlineFont(b=True), "Quarzuhr"), " Vrokzeld")      # Rich Text
        dv = DataValidation(type="list", formula1='"Zebrafink,Quarzuhr"')
        ws.add_data_validation(dv)
        dv.add("D3")
        wb.defined_names["Zebrafinkname"] = DefinedName("Zebrafinkname", attr_text="'Xyloquint Blatt'!$A$1")
        wb.create_sheet("Plumbaxo Zweitblatt")["A1"] = "Vrokzeld im zweiten Blatt"
        wb.save(str(ordner / name))
    wb = load_workbook(str(ordner / "Quarzuhr_nachher.xlsx"))     # markierte Aenderungen mit Kanarienwoertern
    ws = wb["Xyloquint Blatt"]
    v = load_workbook(str(ordner / "Zebrafink_vorher.xlsx"))["Xyloquint Blatt"]["C2"].value
    i = v.index("Gras")
    schreibe(ws["C2"], [Stueck(NORMAL, v[:i]), Stueck(ALT, "Gras"), Stueck(NORMAL, " "), Stueck(NEU, "Quarzuhr"), Stueck(NORMAL, v[i + 4:])])
    wb.save(str(ordner / "Quarzuhr_nachher.xlsx"))
    (ordner / "aenderungen.pdf").write_bytes(pdf.read_bytes())          # Dateinamen fuer --ordner
    (ordner / "vorher.xlsx").write_bytes((ordner / "Zebrafink_vorher.xlsx").read_bytes())
    (ordner / "nachher.xlsx").write_bytes((ordner / "Quarzuhr_nachher.xlsx").read_bytes())
    return ordner


def test_kanarienwoerter_stehen_nirgends_im_ergebnis_einzelnes_paar(tmp_path, capsys):
    o = baue_kanarien_paar(tmp_path)
    rc = lauf(["--pdf", str(o / "Quarzuhr_aenderungen.pdf"), "--vorher", str(o / "Zebrafink_vorher.xlsx"),
               "--nachher", str(o / "Quarzuhr_nachher.xlsx"), "--name", "kanarie1"])
    konsole = capsys.readouterr()
    assert rc == 0
    ziel = config.pfad("steckbrief_dir") / "001_kanarie1"
    texte = {n: (ziel / n).read_text(encoding="utf-8") for n in ("steckbrief.json", "steckbrief.md")}
    sb = json.loads(texte["steckbrief.json"])
    assert sb["tool"]["aenderungen"] >= 3 and sb["excel_nachher"]["muster"]
    ausgabe = konsole.out.replace(str(config.pfad("steckbrief_dir")), "") + konsole.err
    for name, t in {**texte, "konsole": ausgabe}.items():
        assert not kanarien_treffer(t, KANARIEN), name
    assert pruefe(sb) == []
    # Formeln, Kommentar, Hyperlink, Datenpruefung, Namen, Zahlenformat, Schrift: gezaehlt oder als Form, nie als Text
    ev = sb["excel_vorher"]
    assert ev["formeln"]["anzahl"] == 3 and ev["kommentare"] == 1 and ev["hyperlinks"] >= 1 and ev["datenpruefung"] == 1
    assert ev["namen"] >= 1 and ev["blaetter"] == 3
    assert ev["section_zahlenformat"] == {"anderes": 1, "General": 1}
    assert "Quarzuhrschrift" not in json.dumps(sb)


def test_kanarienwoerter_stehen_nirgends_im_ergebnis_mehrere_paare(tmp_path, capsys):
    o = baue_kanarien_paar(tmp_path)
    assert lauf(["--ordner", str(o.parent), "--name", "kanarie2"]) == 0
    konsole = capsys.readouterr()
    ziel = config.pfad("steckbrief_dir") / "001_kanarie2"
    dateien = [p for p in ziel.rglob("*") if p.is_file()]
    assert {p.name for p in dateien} == {"steckbrief.json", "steckbrief.md", "zusammenfassung.json", "zusammenfassung.md"}
    for p in dateien:
        assert not kanarien_treffer(p.read_text(encoding="utf-8"), KANARIEN), p.name
        assert not kanarien_treffer(str(p.relative_to(ziel)), KANARIEN)
    assert not kanarien_treffer(konsole.out.replace(str(config.pfad("steckbrief_dir")), "") + konsole.err, KANARIEN)


def test_kanarienwoerter_auch_nach_dem_zweiten_pruefen_und_bei_fehlern(tmp_path, capsys):
    o = baue_kanarien_paar(tmp_path)
    (o / "Zebrafink_kaputt.xlsx").write_bytes(b"Vrokzeld kein xlsx")
    rc = lauf(["--pdf", str(o / "Quarzuhr_aenderungen.pdf"), "--vorher", str(o / "Zebrafink_kaputt.xlsx"),
               "--nachher", str(o / "Quarzuhr_nachher.xlsx")])
    out = capsys.readouterr()
    assert rc == 2 and not kanarien_treffer(out.out + out.err, KANARIEN)


# ---------------------------------------------------------------------------------------------------------------
# Zusammenfassung und Laufzeit-Schutz
# ---------------------------------------------------------------------------------------------------------------
def test_zusammenfassung_summiert_zaehlt_und_mittelt_und_besteht_die_positivliste():
    a = {"art": "steckbrief", "format": 1, "name": "a", "pdf": {"seiten": 2, "schriften": {"Arial": 5}},
         "excel_nachher": {"blaetter_gleich": True, "strich_anteil": {"blau": 1.0}, "trenner": {" ": 3}},
         "tool": {"lauf_ok": True, "aenderungen": 4}, "schablonen": []}
    b = {"art": "steckbrief", "format": 1, "name": "b", "pdf": {"seiten": 3, "schriften": {"Arial": 1, "Calibri": 2}},
         "excel_nachher": {"blaetter_gleich": False, "strich_anteil": {"blau": 0.5}, "trenner": {"kein": 1}},
         "tool": {"lauf_ok": True, "aenderungen": 1}, "schablonen": []}
    assert pruefe(a) == [] and pruefe(b) == []
    z = steckbrief.fasse_zusammen([a, b])
    assert z["paare"] == 2 and z["pdf"]["seiten"] == 5 and z["pdf"]["schriften"] == {"Arial": 6, "Calibri": 2}
    assert z["excel_nachher"]["blaetter_gleich"] == 1                                # Wahrheitswert -> Anzahl der Paare mit "ja"
    assert z["excel_nachher"]["strich_anteil"]["blau"] == 0.75                       # Anteil -> Mittelwert
    assert z["excel_nachher"]["trenner"] == {" ": 3, "kein": 1}
    assert z["tool"]["lauf_ok"] == 2 and z["tool"]["aenderungen"] == 5
    assert "schablonen" not in z and pruefe(z) == []
    z["pdf"]["schriften"]["Zebrafink!"] = 1                                          # auch die Zusammenfassung hat eine Positivliste
    assert pruefe(z)


def test_zeilen_vergleich_stellt_messen_nach_dem_aufruf_wieder_her_und_bleibt_bei_grossen_bloecken_schnell():
    block, fn = messen.BLOCK_MAX, messen._zell_aehnlich
    v = [[f"Zeile {i}", " ".join(f"wort{(i * 7 + j) % 50}" for j in range(400))] for i in range(60)]
    n = [[f"Zeile {i}", " ".join(f"wort{(i * 11 + j) % 50}" for j in range(400))] for i in range(60)]
    t0 = time.monotonic()
    steckbrief._zeilen_vergleich(messen, messen._zeilen(v), messen._zeilen(n), {1, 2})
    assert time.monotonic() - t0 < 5
    assert messen.BLOCK_MAX == block and messen._zell_aehnlich is fn


# ---------------------------------------------------------------------------------------------------------------
# Generator-Paare (writer + LibreOffice)
# ---------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("erzeuger, seed", [("writer", 7), ("libreoffice", 102)])
def test_steckbrief_laeuft_an_generator_paaren(tmp_path, capsys, erzeuger, seed):
    from tests.fixtures import generator as g
    from tests.fixtures import pdf_lo
    if erzeuger == "libreoffice" and not pdf_lo.verfuegbar():
        pytest.skip("soffice nicht verfuegbar")
    ordner = g.erzeuge_paar(seed, "standard", tmp_path / "paare", "entwicklung", erzeuger=erzeuger)
    t0 = time.monotonic()
    rc = lauf(["--pdf", str(ordner / "aenderungen.pdf"), "--vorher", str(ordner / "vorher.xlsx"),
               "--nachher", str(ordner / "nachher.xlsx"), "--name", "synth"])
    assert rc == 0 and time.monotonic() - t0 < 10
    j = config.pfad("steckbrief_dir") / "001_synth" / "steckbrief.json"
    sb = json.loads(j.read_text(encoding="utf-8"))
    assert pruefe(sb) == [] and sb["pdf"]["seiten"] >= 1 and sb["tool"]["aenderungen"] >= 1
    assert "synth_" not in j.read_text(encoding="utf-8") and ordner.name not in capsys.readouterr().out


# ---------------------------------------------------------------------------------------------------------------
# messen --schablonen
# ---------------------------------------------------------------------------------------------------------------
def messen_lauf(eltern: Path, *optionen: str) -> tuple[int, Path]:
    rc = messen.main(["--ordner", str(eltern), *optionen])
    berichte = sorted(config.pfad("messung_dir").iterdir())
    return rc, berichte[-1]


def test_messen_ohne_schablonen_bleibt_der_bericht_unveraendert(tmp_path):
    eltern = tmp_path / "paare"
    baue_paar(eltern, "echt_01", SAUBER)
    messen_lauf(eltern, "--mit-vorschlaegen")
    bericht = json.loads((config.pfad("messung_dir") / "001_entwicklung" / "bericht.json").read_text(encoding="utf-8"))
    md = (config.pfad("messung_dir") / "001_entwicklung" / "bericht.md").read_text(encoding="utf-8")
    assert all("fehler_schablonen" not in p for p in bericht["paare"]) and "Fehler-Schablonen" not in md


def test_messen_schablonen_formen_fuer_stille_fehler_und_ohne_klartext(tmp_path, monkeypatch, capsys):
    eltern = tmp_path / "paare"
    baue_paar(eltern, "echt_01", SAUBER)
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"C2": "Zebrafink Plumbaxo " + NBSP + "steht hier 12 mal."}))
    rc, ordner = messen_lauf(eltern, "--mit-vorschlaegen", "--schablonen")
    konsole = capsys.readouterr().out
    bericht = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    md = (ordner / "bericht.md").read_text(encoding="utf-8")
    eintraege = bericht["paare"][0]["fehler_schablonen"]
    assert eintraege and bericht["paare"][0]["fehler_schablonen_verworfen"] == 0
    e = eintraege[0]
    assert e["quelle"] == "stiller_fehler" and e["fehlerart"] == "plan_falsch" and e["aenderungstyp"] == "text_ersetzt"
    assert e["erzeuger"] == "unbekannt"
    soll, ist = "".join(x["form"] for x in e["soll"]), "".join(x["form"] for x in e["ist"])
    assert soll.startswith("Aaa Aaaaaa aaaaaa Aaaa") and ist.startswith("Aaaaaaaaa Aaaaaaaa ·aaaaa aaaa 99 aaa.")   # Soll vs Ist: Form, nicht Text
    assert "## Fehler-Schablonen" in md and "stiller_fehler" in md
    for t in (json.dumps(bericht), md, konsole):
        assert not kanarien_treffer(t, ["Zebrafink", "Plumbaxo", "Quokka", "Flamingo"]), "Klartext im Bericht"
    assert steckbrief.pruefe_fehler_schablonen(eintraege) == []


def test_messen_schablonen_auch_fuer_einzelnes_paar_und_falsche_vorschlaege(tmp_path, monkeypatch):
    eltern = tmp_path / "paare"
    baue_paar(eltern, "echt_01", SAUBER)
    baue_paar(eltern, "echt_02", SAUBER)
    rc, ordner = messen_lauf(eltern, "--paar", "echt_01", "--schablonen")
    b = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    assert [p["name"] for p in b["paare"]] == ["echt_01"] and b["paare"][0]["fehler_schablonen"] == []
    # falscher Vorschlag: der Plan meldet "pruefen", die freigegebene Aenderung landet an falscher Stelle
    original = messen.uebernehmen
    zaehler = {"n": 0}

    def falsch(aenderungen, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aenderungen, quelle, ziel, freigaben, cfg)
        zaehler["n"] += 1
        if freigaben:                                                    # die Variante "alle Vorschlaege freigegeben"
            zellen_aendern(Path(ziel), {"C3": "Zebrafink Plumbaxo"})
        return res
    monkeypatch.setattr(messen, "uebernehmen", falsch)
    original_plane = messen.plane

    def plane_pruefen(pa, xl, *a, **kw):
        plan = original_plane(pa, xl, *a, **kw)
        for ae in plan.aenderungen:
            if ae.id == "A0002":
                ae.status = "pruefen"
        return plan
    monkeypatch.setattr(messen, "plane", plane_pruefen)
    rc, ordner = messen_lauf(eltern, "--paar", "echt_02", "--schablonen", "--mit-vorschlaegen")
    b = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    quellen = {e["quelle"] for e in b["paare"][0]["fehler_schablonen"]}
    assert "falscher_vorschlag" in quellen
    for t in ((ordner / "bericht.json").read_text(encoding="utf-8"), (ordner / "bericht.md").read_text(encoding="utf-8")):
        assert not kanarien_treffer(t, ["Zebrafink", "Plumbaxo"])
