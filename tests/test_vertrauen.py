"""Vertrauensmodus (Auftrag Vertrauen): Eingabe Word oder PDF, Stichprobe, Dokument-Ampel, Freigabe-Protokoll, Auswertung.

Nur erfundene Daten (Regel 16). Das Word-Lesen (src/docx_reader.py) kommt von einem anderen Helfer: Die Tests ersetzen das Modul
per sys.modules durch einen Fake, der ein erfundenes PDF liest und es wie eine Analyse aus Word aussehen laesst
(legende_modus "docx", Seite 0, keine Bildkoordinaten). Ein echter Test mit einer echten .docx folgt nach dem Merge.
"""
from __future__ import annotations

import io
import json
import subprocess
import sys
import types
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import cli, config, protokoll, runs, vertrauen
from src.api.app import app
from src.pdf_reader import analysiere_pdf
from tests.fixtures.szenario import erzeuge

DATEI = "entscheidungen.jsonl"
ERLAUBTE_FELDER = {"zeit", "run_id", "ergebnis", "id", "art", "entscheidung", "codes"}


# ---------------------------------------------------------------------------------------------------------------- Hilfen
@pytest.fixture
def paar(tmp_path):
    """Erfundenes Paar: 14 automatische Aenderungen (A0004 ist der einzige Vorschlag) -> Stichprobe von 3."""
    return erzeuge(tmp_path / "fx", "hintergrund")


@pytest.fixture
def lauf(paar):
    rid = runs.neuer_lauf(paar["pdf"], paar["alt"])
    return rid, runs.analysiere(rid)


def _docx_datei(pfad: Path, mit_inhalt: bool = True) -> Path:
    """Kleinste Datei, die wie ein Word-Dokument aufgebaut ist (ZIP mit [Content_Types].xml und word/document.xml)."""
    with zipfile.ZipFile(pfad, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        if mit_inhalt:
            z.writestr("word/document.xml", "<w:document/>")
    return pfad


@pytest.fixture
def docx_leser(monkeypatch, paar):
    """Fake fuer src.docx_reader: liest das erfundene PDF und gibt es als Analyse aus Word zurueck (Seite 0, keine bbox)."""
    aufrufe: list[str] = []

    def analysiere_docx(pfad):
        aufrufe.append(str(pfad))
        pa = analysiere_pdf(paar["pdf"])
        pa.seiten, pa.legende_modus = 0, "docx"
        for b in pa.blocks:
            b.page, b.bbox = 0, (0.0, 0.0, 0.0, 0.0)
            for t in b.tokens:
                t.page, t.bbox = 0, (0.0, 0.0, 0.0, 0.0)
        return pa

    modul = types.ModuleType("src.docx_reader")
    modul.analysiere_docx = analysiere_docx
    monkeypatch.setitem(sys.modules, "src.docx_reader", modul)
    return aufrufe


def _upload(client, dok: Path, xlsx: Path, feld: str = "datei", name: str | None = None, **form):
    with open(dok, "rb") as d, open(xlsx, "rb") as x:
        return client.post("/api/lauf", data={feld: (d, name or dok.name), "excel": (x, "cde.xlsx"), **form},
                           content_type="multipart/form-data")


def _zeilen(rid: str) -> list[dict]:
    z, kaputt = protokoll.lese_entscheidungen(runs.lauf_dir(rid) / DATEI)
    assert kaputt == 0
    return z


# ---------------------------------------------------------------------------------------------------------------- 1. Eingabe
def test_pruefe_aenderungsdatei_pdf_und_docx(tmp_path, paar):
    assert runs.pruefe_aenderungsdatei(paar["pdf"]) == "pdf"
    assert runs.pruefe_aenderungsdatei(_docx_datei(tmp_path / "a.docx")) == "docx"
    assert runs.pruefe_aenderungsdatei(_docx_datei(tmp_path / "x.bin"), "GROSS.DOCX") == "docx"      # die Endung nennt der Aufrufer


@pytest.mark.parametrize("name, inhalt, erwartet", [
    ("a.pdf", b"kein pdf", "kein gültiges PDF"),
    ("a.docx", b"kein zip", "kein gültiges Word-Dokument"),
    ("a.docx", b"%PDF-1.4 ein PDF mit falscher Endung", "kein gültiges Word-Dokument"),
    ("a.doc", b"alt", ".docx speichern"),
    ("a.txt", b"x", "Word-Dokument (.docx) oder ein PDF"),
    ("ohne_endung", b"x", "Word-Dokument (.docx) oder ein PDF"),
])
def test_pruefe_aenderungsdatei_lehnt_ab_in_einfachem_deutsch(tmp_path, name, inhalt, erwartet):
    p = tmp_path / name
    p.write_bytes(inhalt)
    with pytest.raises(runs.NutzerFehler) as e:
        runs.pruefe_aenderungsdatei(p)
    assert erwartet in e.value.meldung and e.value.status == 400


def test_zip_ohne_word_teile_ist_kein_docx(tmp_path):
    p = _docx_datei(tmp_path / "a.docx", mit_inhalt=False)          # z. B. eine umbenannte Excel
    with pytest.raises(runs.NutzerFehler, match="kein gültiges Word-Dokument"):
        runs.pruefe_aenderungsdatei(p)


def test_lade_analyse_waehlt_nach_endung(tmp_path, paar, docx_leser):
    assert runs.lade_analyse(paar["pdf"]).seiten >= 1
    docx = _docx_datei(tmp_path / "w.docx")
    pa = runs.lade_analyse(docx)
    assert pa.legende_modus == "docx" and pa.seiten == 0 and docx_leser == [str(docx)]
    with pytest.raises(runs.NutzerFehler):
        runs.lade_analyse(tmp_path / "x.txt")


def test_fehlender_word_leser_gibt_einfache_meldung(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "src.docx_reader", None)         # wie ein nicht vorhandenes Modul
    with pytest.raises(runs.NutzerFehler) as e:
        runs.lade_analyse(_docx_datei(tmp_path / "w.docx"))
    assert "noch nicht lesen" in e.value.meldung


def test_lauf_aus_word_speichert_quelle_art_und_den_namen(paar, tmp_path, docx_leser):
    docx = _docx_datei(tmp_path / "Änderungen.docx")
    rid = runs.neuer_lauf(docx, paar["alt"])
    assert json.loads((runs.lauf_dir(rid) / "lauf.json").read_text(encoding="utf-8"))["quelle_art"] == "docx"
    plan = runs.analysiere(rid)
    assert plan["quelle_art"] == "docx" and plan["pdf_name"].endswith(".docx") and plan["legende_modus"] == "docx" and plan["seiten"] == 0
    assert all(a["seite"] == 0 and list(a["bbox"]) == [0, 0, 0, 0] for a in plan["aenderungen"])
    assert (runs.lauf_dir(rid) / "eingabe" / plan["pdf_name"]).is_file()
    res = runs.uebernehmen_lauf(rid, [])
    assert res["bilanz"]["stimmt"] and res["bilanz"]["uebernommen"] == plan["bilanz"]["uebernommen"]
    bilanz = load_workbook(runs.lauf_dir(rid) / res["ordner"] / "aenderungsprotokoll.xlsx")["Bilanz"]
    beschriftungen = {z[0] for z in bilanz.iter_rows(values_only=True) if z and z[0]}
    assert "Word-Dokument mit Änderungen" in beschriftungen and "PDF mit Änderungen" not in beschriftungen


def test_api_upload_word_ueber_datei_und_docx_feld(paar, tmp_path, docx_leser):
    client = app.test_client()
    docx = _docx_datei(tmp_path / "w.docx")
    for feld in ("datei", "docx"):
        r = _upload(client, docx, paar["alt"], feld=feld)
        assert r.status_code == 200, r.get_json()
        plan = r.get_json()
        assert plan["quelle_art"] == "docx" and plan["legende_modus"] == "docx" and plan["stichprobe"]
        # keine PDF-Vorschau fuer Word: einfache Meldung statt Absturz
        a = client.get(f"/api/lauf/{plan['run_id']}/vorschau/{plan['aenderungen'][0]['id']}.png")
        assert a.status_code == 404 and a.is_json and "keinen PDF-Ausschnitt" in a.get_json()["fehler"]
    assert [e["status"] for e in client.get("/api/lauf").get_json()] == ["analysiert", "analysiert"]


def test_api_upload_pdf_bleibt_moeglich_alter_feldname_und_neuer(paar):
    client = app.test_client()
    for feld in ("pdf", "datei"):
        r = _upload(client, Path(paar["pdf"]), paar["alt"], feld=feld)
        assert r.status_code == 200, r.get_json()
        assert r.get_json()["quelle_art"] == "pdf" and r.get_json()["legende_modus"] != "docx"
        png = client.get(f"/api/lauf/{r.get_json()['run_id']}/vorschau/A0001.png")
        assert png.status_code == 200 and png.data[:4] == b"\x89PNG"


@pytest.mark.parametrize("name, inhalt, erwartet", [
    ("a.doc", b"alt", ".docx"), ("a.txt", b"x", "Word-Dokument (.docx) oder ein PDF"),
    ("a.docx", b"kein zip", "Word-Dokument"), ("a.pdf", b"kein pdf", "kein gültiges PDF"),
    ("a.pdf", b"PK\x03\x04 doch ein zip", "kein gültiges PDF")])
def test_api_upload_lehnt_falsche_dateien_ab(paar, tmp_path, name, inhalt, erwartet):
    p = tmp_path / name
    p.write_bytes(inhalt)
    r = _upload(app.test_client(), p, paar["alt"])
    assert r.status_code == 400 and r.is_json and erwartet in r.get_json()["fehler"]
    assert runs.liste() == []                                          # nichts angelegt


def test_cli_docx_und_pdf_und_beides_nicht_gleichzeitig(paar, tmp_path, docx_leser, capsys):
    docx = _docx_datei(tmp_path / "w.docx")
    assert cli.main(["analyse", "--docx", str(docx), "--excel", str(paar["alt"])]) == 0
    out = capsys.readouterr().out
    assert "Lauf:" in out and "Dokument-Ampel:" in out and "Wörter im Word-Dokument markiert" in out and "das Word-Dokument wurde vollständig" in out
    assert cli.main(["analyse", "--pdf", str(paar["pdf"]), "--excel", str(paar["alt"])]) == 0           # wie bisher
    assert cli.main(["analyse", "--pdf", str(docx), "--excel", str(paar["alt"])]) == 0                   # beide Endungen an beiden Flaggen
    capsys.readouterr()
    n = len(runs.liste())
    assert cli.main(["analyse", "--pdf", str(paar["pdf"]), "--docx", str(docx), "--excel", str(paar["alt"])]) == 2
    assert "genau eine Änderungsdatei" in capsys.readouterr().out
    assert cli.main(["analyse", "--excel", str(paar["alt"])]) == 2
    kaputt = tmp_path / "kaputt.docx"
    kaputt.write_bytes(b"kein zip")
    assert cli.main(["analyse", "--docx", str(kaputt), "--excel", str(paar["alt"])]) == 2
    assert "kein gültiges Word-Dokument" in capsys.readouterr().out
    assert len(runs.liste()) == n                                                                         # nichts angelegt


def test_cli_farben_sagt_bei_word_dass_es_keine_farben_gibt(tmp_path, capsys):
    assert cli.main(["farben", "--pdf", str(_docx_datei(tmp_path / "w.docx"))]) == 2
    assert "keine Farb-Auswertung" in capsys.readouterr().out


# ---------------------------------------------------------------------------------------------------------------- 2. Stichprobe
def _auto(n: int) -> list[dict]:
    return [{"id": f"A{i:04d}", "status": "uebernommen"} for i in range(1, n + 1)]


@pytest.mark.parametrize("anzahl, erwartet", [(0, 0), (1, 1), (2, 2), (3, 3), (25, 3), (30, 3), (31, 4), (100, 10), (101, 11)])
def test_stichprobe_groesse_anteil_mindestens_drei_hoechstens_alle(anzahl, erwartet):
    assert vertrauen.stichprobe_groesse(anzahl) == erwartet
    assert len(vertrauen.ziehe_stichprobe("20260101-000000-abcdef", _auto(anzahl))) == erwartet


def test_stichprobe_ist_deterministisch_und_haengt_von_der_lauf_id_ab():
    aen = _auto(60)
    a = vertrauen.ziehe_stichprobe("20260101-000000-abcdef", aen)
    assert a == vertrauen.ziehe_stichprobe("20260101-000000-abcdef", aen) and len(a) == 6
    assert len(set(a)) == 6 and a == sorted(a)                                  # in der Reihenfolge des Plans, ohne Doppelte
    andere = {tuple(vertrauen.ziehe_stichprobe(f"20260101-000000-{i:06x}", aen)) for i in range(8)}
    assert len(andere) > 1, "die Auswahl muss von der Lauf-ID abhaengen"
    assert set(vertrauen.ziehe_stichprobe("20260101-000000-abcdef", list(reversed(aen)))) == set(a)      # Reihenfolge der Eingabe egal


def test_stichprobe_ist_unabhaengig_vom_hash_seed_des_interpreters():
    code = ("from src import vertrauen;"
            "print(vertrauen.ziehe_stichprobe('20260101-000000-abcdef', "
            "[{'id': 'A%04d' % i, 'status': 'uebernommen'} for i in range(1, 61)]))")
    ergebnisse = set()
    for seed in ("1", "12345"):
        r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=120,
                           env={**__import__("os").environ, "PYTHONHASHSEED": seed}, cwd=Path(__file__).resolve().parents[1])
        assert r.returncode == 0, r.stderr
        ergebnisse.add(r.stdout.strip())
    assert len(ergebnisse) == 1


def test_stichprobe_nur_aus_automatisch_eingetragenen():
    aen = _auto(10) + [{"id": "V1", "status": "pruefen"}, {"id": "N1", "status": "nicht_uebernommen"}]
    assert not {"V1", "N1"} & set(vertrauen.ziehe_stichprobe("20260101-000000-abcdef", aen))
    assert vertrauen.ziehe_stichprobe("20260101-000000-abcdef", [{"id": "V1", "status": "pruefen"}]) == []


def test_stichprobe_aus_der_config_und_abschaltbar(monkeypatch):
    cfg = config.lade("vertrauen")["vertrauen"]
    assert cfg["aktiv"] is True and cfg["stichprobe_anteil"] == 0.10 and cfg["stichprobe_min"] == 3
    monkeypatch.setitem(cfg, "stichprobe_anteil", 0.5)
    monkeypatch.setitem(cfg, "stichprobe_min", 1)
    assert len(vertrauen.ziehe_stichprobe("20260101-000000-abcdef", _auto(25))) == 13
    monkeypatch.setitem(cfg, "aktiv", False)
    assert vertrauen.ziehe_stichprobe("20260101-000000-abcdef", _auto(25)) == []


def test_kaputte_config_werte_fallen_auf_den_standard(monkeypatch):
    cfg = config.lade("vertrauen")["vertrauen"]
    monkeypatch.setitem(cfg, "stichprobe_anteil", "viel")
    monkeypatch.setitem(cfg, "stichprobe_min", None)
    assert vertrauen.stichprobe_groesse(100) == 10


def test_plan_enthaelt_die_stichprobe_und_wiederholung_aendert_sie_nicht(paar, lauf):
    rid, plan = lauf
    auto = {a["id"] for a in plan["aenderungen"] if a["status"] == "uebernommen"}
    assert len(plan["stichprobe"]) == 3 and set(plan["stichprobe"]) <= auto
    assert runs.lade_plan(rid)["stichprobe"] == plan["stichprobe"]
    gespeichert = json.loads((runs.lauf_dir(rid) / "plan.json").read_text(encoding="utf-8"))
    assert gespeichert["stichprobe"] == plan["stichprobe"] and gespeichert["quelle_art"] == "pdf"
    # gleiche Lauf-ID + gleiche Aenderungen -> gleiche Stichprobe (auch wenn man sie neu zieht)
    assert vertrauen.ziehe_stichprobe(rid, plan["aenderungen"]) == plan["stichprobe"]


# ---------------------------------------------------------------------------------------------------------------- 3. Ampel
def _plan(auto=5, pruefen=0, nicht=0, ausserhalb=0, stimmt=True, vollstaendig=True, luecken=(), spec=None, stichprobe=(), codes=None):
    aen = [{"id": f"A{i}", "status": "uebernommen", "bereich": "", "pruefcodes": []} for i in range(auto)]
    aen += [{"id": f"P{i}", "status": "pruefen", "bereich": "", "pruefcodes": []} for i in range(pruefen)]
    aen += [{"id": f"N{i}", "status": "nicht_uebernommen", "bereich": "", "pruefcodes": []} for i in range(nicht)]
    aen += [{"id": f"X{i}", "status": "nicht_uebernommen", "bereich": "ausserhalb", "pruefcodes": []} for i in range(ausserhalb)]
    for i, c in (codes or {}).items():
        aen[i]["pruefcodes"] = list(c)
    return {"aenderungen": aen, "bilanz": {"stimmt": stimmt, "vollstaendig": vollstaendig}, "luecken": list(luecken),
            "spec": spec, "stichprobe": list(stichprobe)}


def test_ampel_gruen_ohne_offene_punkte():
    a = vertrauen.ampel(_plan())
    assert a["stufe"] == "gruen" and a["satz"] == "Alles sicher eingetragen." and a["gruende"] == [] and a["offen"] == 0


def test_ampel_gelb_bei_offenen_vorschlaegen_und_wird_beim_entscheiden_gruen():
    plan = _plan(auto=20, pruefen=2)
    a = vertrauen.ampel(plan)
    assert a["stufe"] == "gelb" and a["satz"] == "Bitte ein paar Punkte ansehen." and a["offen"] == 2
    assert vertrauen.ampel(plan, offene_vorschlaege=1)["stufe"] == "gelb"
    assert vertrauen.ampel(plan, offene_vorschlaege=0)["stufe"] == "gruen"
    assert vertrauen.ampel(plan, offene_vorschlaege=99)["offen"] == 2                 # nie mehr als vorhanden


def test_ampel_gelb_solange_die_stichprobe_nicht_bestaetigt_ist():
    plan = _plan(stichprobe=["A0", "A1", "A2"])
    a = vertrauen.ampel(plan)
    assert a["stufe"] == "gelb" and a["stichprobe_offen"] == 3 and any("Stichprobe" in g for g in a["gruende"])
    assert vertrauen.ampel(plan, stichprobe_offen=1)["stufe"] == "gelb"
    assert vertrauen.ampel(plan, stichprobe_offen=0)["stufe"] == "gruen"


def test_ampel_nicht_eingetragenes_bleibt_gelb_ausserhalb_der_cde_zaehlt_nicht():
    assert vertrauen.ampel(_plan(nicht=1), offene_vorschlaege=0, stichprobe_offen=0)["stufe"] == "gelb"
    assert vertrauen.ampel(_plan(ausserhalb=3))["stufe"] == "gruen"


@pytest.mark.parametrize("plan", [
    _plan(stimmt=False),
    _plan(vollstaendig=False),
    _plan(luecken=[{"art": "tabelle", "seite": 2}]),
    _plan(codes={0: ["spec_unklar"]}),
    _plan(spec={"erkannt": None, "kandidaten": ["11 22 33", "44 55 66"]}),
    _plan(auto=5, pruefen=5),                                                             # 50 offene Punkte je 100 > 40
], ids=["bilanz", "unvollstaendig", "luecke", "spec_code", "spec_ohne_wahl", "zu_viele_offene"])
def test_ampel_rot(plan):
    a = vertrauen.ampel(plan)
    assert a["stufe"] == "rot" and a["satz"] == "Hier bitte genau hinsehen." and a["gruende"]


def test_ampel_rot_wegen_vieler_offener_punkte_wird_beim_entscheiden_besser():
    plan = _plan(auto=5, pruefen=5)
    assert vertrauen.ampel(plan)["offen_je_100"] == 50.0
    assert vertrauen.ampel(plan, offene_vorschlaege=1)["stufe"] == "gelb"                 # 10 je 100
    assert vertrauen.ampel(plan, offene_vorschlaege=0)["stufe"] == "gruen"
    # Bilanzfehler und Luecken bleiben rot, egal wie viel entschieden ist
    assert vertrauen.ampel(_plan(pruefen=2, stimmt=False), offene_vorschlaege=0, stichprobe_offen=0)["stufe"] == "rot"


def test_ampel_schwelle_kommt_aus_der_config(monkeypatch):
    assert config.lade("vertrauen")["ampel"]["rot_offene_je_100"] == 40
    monkeypatch.setitem(config.lade("vertrauen")["ampel"], "rot_offene_je_100", 25)
    assert vertrauen.ampel(_plan(auto=8, pruefen=2))["stufe"] == "gelb"                   # 20 offene Punkte je 100, Schwelle 25
    monkeypatch.setitem(config.lade("vertrauen")["ampel"], "rot_offene_je_100", 19)
    assert vertrauen.ampel(_plan(auto=8, pruefen=2))["stufe"] == "rot"                    # 20 > 19


def test_ampel_konsens_code_zaehlt_wie_ein_offener_punkt_unbekannte_codes_nicht():
    a = vertrauen.ampel(_plan(codes={0: ["konsens_abweichung"]}), stichprobe_offen=0)
    assert a["stufe"] == "gelb" and a["offen"] == 1
    # ein Vorschlag mit konsens_-Code zaehlt nur einmal
    b = vertrauen.ampel(_plan(auto=20, pruefen=1, codes={20: ["konsens_zwei_leser"]}))
    assert b["offen"] == 1
    c = vertrauen.ampel(_plan(codes={0: ["irgendein_neuer_code", "sperre:neu", "a1_unbekannt"]}))
    assert c["stufe"] == "gruen"


def test_ampel_aelterer_abgleich_ohne_lueckenangabe_ist_nie_gruen():
    plan = _plan()
    plan.pop("luecken")
    plan["bilanz"].pop("vollstaendig")
    a = vertrauen.ampel(plan)
    assert a["stufe"] == "gelb" and any("älteren" in g for g in a["gruende"])
    assert vertrauen.ampel({"aenderungen": [], "bilanz": {}})["stufe"] == "gelb"           # nichts bekannt


def test_ampel_im_plan_der_api_und_der_endpunkt(paar):
    client = app.test_client()
    plan = _upload(client, Path(paar["pdf"]), paar["alt"]).get_json()
    rid = plan["run_id"]
    assert plan["ampel"]["stufe"] == "gelb" and plan["ampel"]["offen"] == 1 and plan["ampel"]["stichprobe_offen"] == 3
    assert plan["stichprobe"] and set(plan["stichprobe"]) <= {a["id"] for a in plan["aenderungen"] if a["kategorie"] == "automatisch"}
    assert client.get(f"/api/lauf/{rid}").get_json()["ampel"] == plan["ampel"]
    fertig = client.get(f"/api/lauf/{rid}/ampel?offen=0&stichprobe_offen=0")
    assert fertig.status_code == 200 and fertig.get_json()["stufe"] == "gruen" and fertig.get_json()["satz"] == "Alles sicher eingetragen."
    assert client.get(f"/api/lauf/{rid}/ampel").get_json()["stufe"] == "gelb"
    for schlecht in ("offen=-1", "offen=abc", "stichprobe_offen=1.5", "offen=" + "9" * 9):
        r = client.get(f"/api/lauf/{rid}/ampel?{schlecht}")
        assert r.status_code == 400 and r.is_json
    assert client.get("/api/lauf/20260101-000000-abcdef/ampel").status_code == 404


# ---------------------------------------------------------------------------------------------------------------- 4. Freigabe-Protokoll
def test_entscheidung_zeile_hat_nur_kennungen_und_ersetzt_fremden_text():
    z = protokoll.entscheidung_zeile("20260101-000000-abcdef", "ergebnis_001", "A0001", "pruefen", "uebernommen",
                                     ["a1_leerraum", "Das ist ein Satz aus dem Dokument", "sperre:teilweise_markiert", "a1_leerraum"])
    assert set(z) == ERLAUBTE_FELDER
    assert z["codes"] == ["a1_leerraum", "sperre:teilweise_markiert", "unbekannt"]
    with pytest.raises(ValueError):
        protokoll.entscheidung_zeile("r", "e", "A1", "frei_erfunden", "uebernommen", [])
    with pytest.raises(ValueError):
        protokoll.entscheidung_zeile("r", "e", "A1", "pruefen", "vielleicht", [])


def test_protokoll_enthaelt_alle_entscheidungen_mit_art_und_ohne_dokumenttext(lauf):
    rid, plan = lauf
    s1, s2, s3 = plan["stichprobe"]
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["status"] == "pruefen"]
    res = runs.uebernehmen_lauf(rid, [vorschlag], [s1], stichprobe_passt=[s2], stichprobe_falsch=[s1])
    zeilen = _zeilen(rid)
    assert len(zeilen) == 15 and all(set(z) == ERLAUBTE_FELDER for z in zeilen)
    nach_id = {z["id"]: z for z in zeilen}
    assert (nach_id[s1]["art"], nach_id[s1]["entscheidung"]) == ("stichprobe", "abgelehnt")           # "Falsch"
    assert (nach_id[s2]["art"], nach_id[s2]["entscheidung"]) == ("stichprobe", "bestaetigt")          # "Passt"
    assert (nach_id[s3]["art"], nach_id[s3]["entscheidung"]) == ("stichprobe", "uebernommen")         # nicht angesehen
    assert (nach_id[vorschlag]["art"], nach_id[vorschlag]["entscheidung"]) == ("pruefen", "uebernommen")
    assert nach_id[vorschlag]["codes"] == ["sperre:teilweise_markiert"]
    uebrige = [z for i, z in nach_id.items() if i not in (s1, s2, s3, vorschlag)]
    assert len(uebrige) == 11 and all((z["art"], z["entscheidung"]) == ("automatisch", "uebernommen") for z in uebrige)
    assert {z["run_id"] for z in zeilen} == {rid} and {z["ergebnis"] for z in zeilen} == {res["ordner"]}
    # Regel 16: kein Text aus Dokument oder Excel in der Datei
    text = (runs.lauf_dir(rid) / DATEI).read_text(encoding="utf-8")
    fremd = {plan["pdf_name"], plan["excel_name"]}
    for a in plan["aenderungen"]:
        fremd |= {a[k] for k in ("alt", "neu", "excel_text", "kontext", "soll_absatz", "ist_absatz", "grund", "grund_einfach", "ziel", "section")
                  if isinstance(a.get(k), str)}
        fremd |= {s[1] for s in a["stuecke"] if isinstance(s, list) and len(s) > 1}
    for f in sorted(x for x in fremd if len(x) >= 5):
        assert f not in text, f"Dokumenttext im Protokoll: {f!r}"


def test_protokoll_abgelehnter_vorschlag_pauschale_freigabe_und_abwahl_ohne_stichprobe(paar, lauf):
    rid, plan = lauf
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["status"] == "pruefen"]
    uebrig = [i for i in (a["id"] for a in plan["aenderungen"] if a["status"] == "uebernommen") if i not in plan["stichprobe"]]
    runs.uebernehmen_lauf(rid, [], [uebrig[0]], abgelehnt=[vorschlag])
    z = {x["id"]: x for x in _zeilen(rid)}
    assert (z[vorschlag]["art"], z[vorschlag]["entscheidung"]) == ("pruefen", "abgelehnt")
    assert (z[uebrig[0]]["art"], z[uebrig[0]]["entscheidung"]) == ("automatisch", "abgewaehlt")
    # zweite Excel: pauschal freigegeben -> art vorschlag; ein nicht entschiedener Vorschlag steht gar nicht im Protokoll
    rid2 = runs.neuer_lauf(paar["pdf"], paar["alt"])
    runs.analysiere(rid2)
    assert len(_zeilen(rid)) == 15 and not (runs.lauf_dir(rid2) / DATEI).exists()
    runs.uebernehmen_lauf(rid2, [vorschlag], [], pauschal=[vorschlag])
    assert {x["id"]: x["art"] for x in _zeilen(rid2)}[vorschlag] == "vorschlag"
    rid3 = runs.neuer_lauf(paar["pdf"], paar["alt"])
    runs.analysiere(rid3)
    runs.uebernehmen_lauf(rid3, [])
    assert vorschlag not in {x["id"] for x in _zeilen(rid3)} and len(_zeilen(rid3)) == 14


def test_protokoll_ist_append_only(lauf):
    rid, plan = lauf
    datei = runs.lauf_dir(rid) / DATEI
    runs.uebernehmen_lauf(rid, [])
    erste = datei.read_bytes()
    runs.uebernehmen_lauf(rid, [], [plan["stichprobe"][0]], stichprobe_falsch=[plan["stichprobe"][0]])
    zweite = datei.read_bytes()
    assert zweite.startswith(erste) and len(zweite) > len(erste)                      # alte Zeilen unveraendert
    assert {z["ergebnis"] for z in _zeilen(rid)} == {"ergebnis_001", "ergebnis_002"}
    assert len(_zeilen(rid)) == 14 + 14                                              # die Ausgangslage hat 14 Eintraege je Excel
    # abgebrochene letzte Zeile: die naechste Zeile beginnt trotzdem sauber
    datei.write_bytes(zweite + b'{"abgebrochen')
    runs.uebernehmen_lauf(rid, [])
    zeilen, kaputt = protokoll.lese_entscheidungen(datei)
    assert kaputt == 1 and len(zeilen) == 14 * 3 and datei.read_bytes().startswith(zweite)


def test_fehlgeschlagene_uebernahme_schreibt_nichts_ins_protokoll(lauf, monkeypatch):
    rid, plan = lauf
    monkeypatch.setattr(runs, "uebernehmen", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("Abbruch")))
    with pytest.raises(RuntimeError):
        runs.uebernehmen_lauf(rid, [])
    assert not (runs.lauf_dir(rid) / DATEI).exists()


def test_schreibfehler_im_protokoll_laesst_die_excel_bestehen_und_warnt(lauf, monkeypatch):
    rid, _ = lauf

    def kaputt(*a, **k):
        raise OSError("Platte voll")
    monkeypatch.setattr(runs, "haenge_entscheidungen_an", kaputt)
    res = runs.uebernehmen_lauf(rid, [])
    assert any("Freigabe-Protokoll" in w for w in res["warnungen"])
    assert (runs.lauf_dir(rid) / res["ordner"] / res["excel"]).is_file() and len(runs.ergebnisse(rid)) == 1


@pytest.mark.parametrize("fall, erwartet", [
    ("falsch_ohne_abwahl", "auch abgewählt"), ("passt_ausserhalb", "nicht zur Stichprobe"), ("passt_und_falsch", "zugleich"),
    ("abgelehnt_automatisch", "nur bei Vorschlägen"), ("abgelehnt_und_freigegeben", "zugleich"), ("pauschal_ohne_freigabe", "nur bei freigegebenen"),
    ("passt_aber_abgewaehlt", "abgewählt"), ("unbekannte_id", "Unbekannte")])
def test_ungueltige_entscheidungen_legen_nichts_an(lauf, fall, erwartet):
    rid, plan = lauf
    s1, s2, s3 = plan["stichprobe"]
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["status"] == "pruefen"]
    ausserhalb = next(a["id"] for a in plan["aenderungen"] if a["status"] == "uebernommen" and a["id"] not in plan["stichprobe"])
    args = {
        "falsch_ohne_abwahl": dict(stichprobe_falsch=[s1]),
        "passt_ausserhalb": dict(stichprobe_passt=[ausserhalb]),
        "passt_und_falsch": dict(abwahl=[s1], stichprobe_passt=[s1], stichprobe_falsch=[s1]),
        "abgelehnt_automatisch": dict(abgelehnt=[s1]),
        "abgelehnt_und_freigegeben": dict(freigaben=[vorschlag], abgelehnt=[vorschlag]),
        "pauschal_ohne_freigabe": dict(pauschal=[vorschlag]),
        "passt_aber_abgewaehlt": dict(abwahl=[s1], stichprobe_passt=[s1]),
        "unbekannte_id": dict(stichprobe_passt=["A9999"]),
    }[fall]
    mit = {"freigaben": args.pop("freigaben", None), "abwahl": args.pop("abwahl", None)}
    with pytest.raises(runs.NutzerFehler, match=erwartet) as e:
        runs.uebernehmen_lauf(rid, mit["freigaben"], mit["abwahl"], **args)
    assert e.value.status == 400
    assert list(runs.lauf_dir(rid).glob("ergebnis_*")) == [] and not (runs.lauf_dir(rid) / DATEI).exists()


def test_api_uebernehmen_mit_stichprobe_und_ablehnung_wird_gespeichert_und_vorbelegt(paar):
    client = app.test_client()
    plan = _upload(client, Path(paar["pdf"]), paar["alt"]).get_json()
    rid = plan["run_id"]
    s1, s2, _s3 = plan["stichprobe"]
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["kategorie"] == "pruefen"]
    r = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": [], "abwahl": [s1], "abgelehnt": [vorschlag],
                                                          "stichprobe": {"passt": [s2], "falsch": [s1]}})
    assert r.status_code == 200, r.get_json()
    assert r.get_json()["stichprobe"] == {"passt": [s2], "falsch": [s1]} and r.get_json()["abgelehnt"] == [vorschlag]
    erg = client.get(f"/api/lauf/{rid}").get_json()["ergebnisse"][-1]
    assert erg["stichprobe"] == {"passt": [s2], "falsch": [s1]} and erg["abgelehnt"] == [vorschlag] and erg["abwahl"] == [s1]
    z = {x["id"]: x for x in _zeilen(rid)}
    assert z[s1]["entscheidung"] == "abgelehnt" and z[s2]["entscheidung"] == "bestaetigt" and z[vorschlag]["entscheidung"] == "abgelehnt"
    # ungueltige Formen -> 400 mit einfachem Satz, nichts Neues angelegt
    for body in ({"stichprobe": [s1]}, {"stichprobe": {"passt": "A1"}}, {"abgelehnt": "A1"}, {"pauschal": [1]},
                 {"stichprobe": {"falsch": [s1]}}):
        antwort = client.post(f"/api/lauf/{rid}/uebernehmen", json=body)
        assert antwort.status_code == 400 and antwort.is_json, body
    assert len(client.get(f"/api/lauf/{rid}").get_json()["ergebnisse"]) == 1


# ---------------------------------------------------------------------------------------------------------------- 5. Auswertung
def _zeile(run, ergebnis, i, art, entscheidung, codes):
    return protokoll.entscheidung_zeile(run, ergebnis, i, art, entscheidung, codes)


def _lege_protokoll(run_id: str, zeilen: list[dict]) -> Path:
    d = runs.runs_dir() / run_id
    d.mkdir()
    protokoll.haenge_entscheidungen_an(d / DATEI, zeilen)
    return d


def test_statistik_quoten_und_empfehlungen():
    r1, r2 = "20260101-000000-aaaaaa", "20260102-000000-bbbbbb"
    z = []
    z += [_zeile(r1, "ergebnis_001", f"A{i}", "pruefen", "uebernommen", ["a1_leerraum"]) for i in range(25)]       # 25 Freigaben
    z += [_zeile(r1, "ergebnis_001", f"B{i}", "pruefen", "uebernommen", ["excel_weicht_ab"]) for i in range(20)]
    z += [_zeile(r1, "ergebnis_001", f"B{i}", "pruefen", "abgelehnt", ["excel_weicht_ab"]) for i in range(20, 30)]  # 20 von 30
    z += [_zeile(r2, "ergebnis_001", f"C{i}", "stichprobe", "bestaetigt", []) for i in range(22)]
    z += [_zeile(r2, "ergebnis_001", f"D{i}", "automatisch", "uebernommen", []) for i in range(100)]            # unbesehen
    z += [_zeile(r2, "ergebnis_001", "E1", "pruefen", "uebernommen", ["neuer_code"])]
    stat = vertrauen.statistik(z)
    s = stat["a1_leerraum"]
    assert (s["anzahl"], s["frei"], s["abgelehnt"]) == (25, 25, 0) and vertrauen.quoten(s) == (25, 1.0, 0.0)
    assert vertrauen.empfehlung(s) == "sicher genug zum Automatisieren"
    w = stat["excel_weicht_ab"]
    assert vertrauen.quoten(w)[0] == 30 and round(vertrauen.quoten(w)[1], 2) == 0.67
    assert vertrauen.empfehlung(w) == "noch nicht automatisieren: zu oft abgelehnt"
    ohne = stat[vertrauen.OHNE_CODE]
    assert ohne["anzahl"] == 122 and vertrauen.quoten(ohne)[0] == 22                 # unbesehene Eintraege zaehlen nicht als angesehen
    assert vertrauen.empfehlung(ohne) == "Automatik bewährt sich"
    assert vertrauen.empfehlung(stat["neuer_code"]).startswith("noch zu wenig angesehen (1 von mindestens 20")
    # "abgewaehlt" ist eine Ablehnung, auch ohne Stichprobe
    x = vertrauen.statistik([_zeile(r1, "ergebnis_001", "A1", "automatisch", "abgewaehlt", ["a1_leerraum"])])["a1_leerraum"]
    assert vertrauen.quoten(x) == (1, 0.0, 1.0)


def test_statistik_zaehlt_je_lauf_nur_die_letzte_excel(tmp_path):
    r = "20260101-000000-aaaaaa"
    alt = [_zeile(r, "ergebnis_001", f"A{i}", "pruefen", "abgelehnt", ["x"]) for i in range(5)]
    neu = [_zeile(r, "ergebnis_002", f"A{i}", "pruefen", "uebernommen", ["x"]) for i in range(5)]
    assert vertrauen.statistik(alt + neu)["x"]["frei"] == 5 and vertrauen.statistik(alt + neu)["x"]["abgelehnt"] == 0


def test_empfehlung_schwellen_stehen_in_der_config(monkeypatch):
    z = {"anzahl": 10, "frei": 10, "abgelehnt": 0, "je_art": {"pruefen": 10}}
    assert vertrauen.empfehlung(z).startswith("noch zu wenig angesehen")
    monkeypatch.setitem(config.lade("vertrauen")["auswertung"], "min_angesehen", 10)
    assert vertrauen.empfehlung(z) == "sicher genug zum Automatisieren"
    z2 = {"anzahl": 10, "frei": 9, "abgelehnt": 1, "je_art": {"pruefen": 10}}
    assert vertrauen.empfehlung(z2) == "noch nicht automatisieren: zu oft abgelehnt"
    monkeypatch.setitem(config.lade("vertrauen")["auswertung"], "ablehnung_max", 0.2)
    monkeypatch.setitem(config.lade("vertrauen")["auswertung"], "freigabe_ab", 0.95)
    assert vertrauen.empfehlung(z2).startswith("weiter ansehen")


def _verzeichnis_stand(wurzel: Path) -> dict[str, bytes]:
    return {str(p.relative_to(wurzel)): p.read_bytes() for p in sorted(wurzel.rglob("*")) if p.is_file()}


def test_cli_vertrauen_zeigt_je_pruefcode_und_schreibt_nichts(capsys):
    r1 = "20260101-000000-aaaaaa"
    z = [_zeile(r1, "ergebnis_001", f"A{i}", "pruefen", "uebernommen", ["a1_leerraum"]) for i in range(25)]
    z += [_zeile(r1, "ergebnis_001", f"B{i}", "stichprobe", "bestaetigt", []) for i in range(3)]
    z += [_zeile(r1, "ergebnis_001", "C1", "pruefen", "abgelehnt", ["erfundener_code"])]
    _lege_protokoll(r1, z)
    _lege_protokoll("20260102-000000-bbbbbb", [_zeile("20260102-000000-bbbbbb", "ergebnis_001", "A1", "pruefen", "uebernommen", ["a1_leerraum"])])
    vorher = _verzeichnis_stand(runs.runs_dir())
    assert cli.main(["vertrauen"]) == 0
    out = capsys.readouterr().out
    assert "a1_leerraum: Anzahl 26 · angesehen 26 · Freigabe 100 % · Ablehnung 0 %" in out
    assert "Empfehlung: sicher genug zum Automatisieren" in out
    assert "(ohne Prüfcode): Anzahl 3" in out and "erfundener_code: Anzahl 1 · angesehen 1 · Freigabe 0 % · Ablehnung 100 %" in out
    assert "Code ohne Beschreibung" in out and "Läufen" in out and "Das ist nur eine Empfehlung" in out
    assert _verzeichnis_stand(runs.runs_dir()) == vorher                       # nur gelesen


def test_cli_vertrauen_ohne_protokolle_und_mit_kaputter_zeile(capsys):
    assert cli.main(["vertrauen"]) == 0
    assert "Noch keine Entscheidungen gespeichert" in capsys.readouterr().out
    d = _lege_protokoll("20260101-000000-aaaaaa", [_zeile("20260101-000000-aaaaaa", "ergebnis_001", "A1", "pruefen", "uebernommen", ["x"])])
    with open(d / DATEI, "ab") as f:
        f.write(b"das ist keine json-zeile\n{\"art\": \"unsinn\"}\n")
    assert cli.main(["vertrauen"]) == 0
    out = capsys.readouterr().out
    assert "x: Anzahl 1" in out and "2 Zeilen in den Protokollen waren nicht lesbar" in out
    assert cli.main(["vertrauen", "--unsinn"]) == 2


def test_ende_zu_ende_dashboard_entscheidungen_landen_in_der_auswertung(paar, capsys):
    client = app.test_client()
    plan = _upload(client, Path(paar["pdf"]), paar["alt"]).get_json()
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["kategorie"] == "pruefen"]
    assert client.post(f"/api/lauf/{plan['run_id']}/uebernehmen",
                       json={"freigaben": [vorschlag], "stichprobe": {"passt": plan["stichprobe"][:2]}}).status_code == 200
    assert cli.main(["vertrauen"]) == 0
    out = capsys.readouterr().out
    assert "sperre:teilweise_markiert: Anzahl 1 · angesehen 1 · Freigabe 100 %" in out
    assert "(ohne Prüfcode): Anzahl 14 · angesehen 2 · Freigabe 100 %" in out


# ---------------------------------------------------------------------------------------------------------------- 6. Alte Laeufe
def test_alter_lauf_ohne_neue_felder_bleibt_lesbar_und_nutzbar(paar):
    client = app.test_client()
    plan = _upload(client, Path(paar["pdf"]), paar["alt"]).get_json()
    rid = plan["run_id"]
    d = runs.lauf_dir(rid)
    pj = json.loads((d / "plan.json").read_text(encoding="utf-8"))
    for k in ("stichprobe", "quelle_art"):
        pj.pop(k)
    (d / "plan.json").write_text(json.dumps(pj), encoding="utf-8")
    lj = json.loads((d / "lauf.json").read_text(encoding="utf-8"))
    lj.pop("quelle_art")
    (d / "lauf.json").write_text(json.dumps(lj), encoding="utf-8")
    g = client.get(f"/api/lauf/{rid}")
    assert g.status_code == 200
    alt = g.get_json()
    assert alt["stichprobe"] == [] and alt["quelle_art"] == "pdf" and alt["ampel"]["stichprobe_offen"] == 0
    assert alt["ampel"]["stufe"] == "gelb"                                    # der eine Vorschlag ist offen
    assert client.get(f"/api/lauf/{rid}/ampel?offen=0").get_json()["stufe"] == "gruen"
    assert client.get(f"/api/lauf/{rid}/vorschau/A0001.png").status_code == 200
    assert client.get("/api/lauf").get_json()[0]["status"] == "analysiert"
    u = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []})
    assert u.status_code == 200
    # auch ein Ergebnis aus der Zeit davor (ohne abgelehnt/stichprobe/pauschal) ist lesbar
    e = d / u.get_json()["ordner"] / "ergebnis.json"
    ej = json.loads(e.read_text(encoding="utf-8"))
    for k in ("abgelehnt", "pauschal", "stichprobe"):
        ej.pop(k)
    e.write_text(json.dumps(ej), encoding="utf-8")
    erg = client.get(f"/api/lauf/{rid}").get_json()["ergebnisse"][0]
    assert erg["abgelehnt"] == [] and erg["stichprobe"] == {"passt": [], "falsch": []}
    # Ein Lauf ohne entscheidungen.jsonl stoert die Auswertung nicht
    assert vertrauen.lese_alle(runs.runs_dir())[1] >= 0


def test_stichprobe_nur_im_plan_gueltig_alter_plan_ohne_stichprobe_lehnt_passt_ab(paar):
    rid = runs.neuer_lauf(paar["pdf"], paar["alt"])
    plan = runs.analysiere(rid)
    pj = json.loads((runs.lauf_dir(rid) / "plan.json").read_text(encoding="utf-8"))
    pj.pop("stichprobe")
    (runs.lauf_dir(rid) / "plan.json").write_text(json.dumps(pj), encoding="utf-8")
    with pytest.raises(runs.NutzerFehler, match="nicht zur Stichprobe"):
        runs.uebernehmen_lauf(rid, [], stichprobe_passt=[plan["stichprobe"][0]])
