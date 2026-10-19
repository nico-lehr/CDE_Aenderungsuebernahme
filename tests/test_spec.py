"""Spec-Filter (Welle 3, Nico: eine CDE enthaelt mehrere Specs; er filtert nach der Spec des PDFs und bearbeitet nur deren
Zeilen). Erfundene CDE mit drei Specs: gleiche Saetze in allen drei Specs (Dubletten-Falle), uneinheitliche Schreibweise
der Spec-Werte ("44 55 66", "44_55_66 ", "44.55.66 EMEA"). Nur erfundene Daten (Regel 16).
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from src import markierung as mk
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.models import NICHT, PRUEFEN, UEBERNOMMEN
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests.test_register import DATEN_AB, GRUND, NB, E, F, G, excel_absaetze, register_mappe, wort_pdf

pytestmark = pytest.mark.legende_neu

N = len(GRUND)
SCHREIBWEISEN = {0: ["11_22_33"], 1: ["44 55 66", "44_55_66 ", "44.55.66 EMEA"], 2: ["77-88-99"]}


def _spec_von(i: int) -> str:
    werte = SCHREIBWEISEN[i // N]
    return werte[i % len(werte)]


def _mappe(tmp_path, name="cde.xlsx"):
    return register_mappe(tmp_path / name, excel_absaetze(GRUND) * 3, spec_von=_spec_von)


def _zeile(block: int, i: int) -> int:
    return DATEN_AB + block * N + i


def _plan(tmp_path, pdf_absaetze, titel="SECTION 44 55 66 – PUMPENANLAGE", spec=None):
    xlsx = _mappe(tmp_path)
    pdf = wort_pdf(tmp_path / "a.pdf", pdf_absaetze, [titel] if titel else None)
    xl = ExcelModell(xlsx)
    return plane(analysiere_pdf(pdf), xl, spec=spec), xl, xlsx


def _werte(pfad, zeilen):
    ws = load_workbook(str(pfad), rich_text=True)["Abgleich"]
    return {(r, c): repr(ws.cell(r, c).value) for r in zeilen for c in range(1, 19)}


def test_spec_spalte_und_bloecke_werden_erkannt(tmp_path):
    xl = ExcelModell(_mappe(tmp_path))
    lay = xl.layouts["Abgleich"]
    assert lay.register and lay.spec_spalte == 2
    specs = xl.specs("Abgleich")
    assert sorted(specs) == ["11 22 33", "44 55 66", "77 88 99"]
    assert specs["44 55 66"] == list(range(_zeile(1, 0), _zeile(2, 0)))          # alle Schreibweisen gehoeren zusammen


@pytest.mark.markierung_an
def test_aenderung_landet_nur_in_der_passenden_spec(tmp_path):
    pdf = list(GRUND)
    pdf[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")    # steht in allen drei Specs
    plan, xl, xlsx = _plan(tmp_path, pdf)
    assert plan.spec["erkannt"] == "44 55 66" and plan.spec["quelle"] == "pdf"
    assert plan.spec["zeilen"] == [_zeile(1, 0), _zeile(1, N - 1)]
    assert plan.spec["kandidaten"] == ["11 22 33", "44 55 66", "77 88 99"]
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and ae.ziel == f"Abgleich!G{_zeile(1, 2)}"            # eindeutig, nicht mehrdeutig
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze)
    fremd = [r for r in range(DATEN_AB, DATEN_AB + 3 * N) if not _zeile(1, 0) <= r < _zeile(2, 0)]
    assert _werte(out, fremd) == _werte(xlsx, fremd)                                       # fremde Specs unveraendert
    st = mk.stuecke_aus_zelle(load_workbook(str(out), rich_text=True)["Abgleich"].cell(_zeile(1, 2), G))
    assert mk.ansicht(st, "neu") == "Die Pumpe hebt das Wasser in das obere Becken."


def test_neue_zeile_im_richtigen_block(tmp_path):
    pdf = GRUND[:8] + [("liste", "{+c.+}", "{+Pruefung der Wellen nach Plan.+}")] + GRUND[8:] + \
        [("liste", "{+3.+}", "{+Der Kranich steht auch auf der Rueckseite.+}")]
    plan, xl, _ = _plan(tmp_path, pdf)
    neu = {a.neu.split()[-1]: a for a in plan.aenderungen}
    mitte = neu["Plan."].ops[0]
    assert neu["Plan."].status == UEBERNOMMEN and (mitte.zeile, mitte.vor) == (_zeile(1, 7), False)
    assert mitte.zellen[G] == "c." + NB * 15 + " Pruefung der Wellen nach Plan." and mitte.zellen[F] == "B"
    assert _zeile(1, 0) <= mitte.stil_von < _zeile(2, 0)                                   # Vorlage aus derselben Spec
    ende = neu["Rueckseite."].ops[0]
    assert (ende.zeile, ende.vor) == (_zeile(1, N - 1), False)                              # am Ende des eigenen Blocks
    assert ende.zellen[E] == "1.2" and _zeile(1, 0) <= ende.stil_von < _zeile(2, 0)


def test_unbekannte_spec_warnt_und_traegt_nichts_automatisch_ein(tmp_path):
    pdf = list(GRUND)
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    plan, _, _ = _plan(tmp_path, pdf, titel="SECTION 12 12 12 – ANDERE ANLAGE")
    assert plan.spec["erkannt"] is None and plan.spec["kandidaten"] == ["11 22 33", "44 55 66", "77 88 99"]
    assert any("Spec des PDFs nicht in der CDE gefunden" in w for w in plan.warnungen)
    assert plan.aenderungen and all(a.status != UEBERNOMMEN for a in plan.aenderungen)
    assert all(a.status in (PRUEFEN, NICHT) for a in plan.aenderungen)


def test_spec_vom_nutzer_vorgegeben(tmp_path):
    pdf = list(GRUND)
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    plan, _, _ = _plan(tmp_path, pdf, titel=None, spec="77.88.99")
    assert plan.spec["erkannt"] == "77 88 99" and plan.spec["quelle"] == "vorgabe"
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and ae.ziel == f"Abgleich!G{_zeile(2, 5)}"
    plan2, _, _ = _plan(tmp_path / "x", pdf, titel=None, spec="99 99 99")
    assert plan2.spec["erkannt"] is None and all(a.status != UEBERNOMMEN for a in plan2.aenderungen)
    assert any("vorgegebene Spec" in w for w in plan2.warnungen)


def test_fusszeile_mit_spec_genuegt(tmp_path):
    pdf = list(GRUND)
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    plan, _, _ = _plan(tmp_path, pdf, titel="11 22 33 Pumpenhaus Leitfaden - 1")
    assert plan.spec["erkannt"] == "11 22 33"
    (ae,) = plan.aenderungen
    assert ae.ziel == f"Abgleich!G{_zeile(0, 5)}"


def test_eine_spec_in_der_cde_wie_bisher(tmp_path):
    from tests.test_register import _plan as plan_einfach
    pdf = list(GRUND)
    pdf[5] = ("liste", "2.", "Werknorm WN {-2040-}{+2041+} fuer Leitungen.")
    plan, _, _ = plan_einfach(tmp_path, pdf)
    assert plan.spec == {"erkannt": "40 05 13", "quelle": "einzige", "zeilen": [DATEN_AB, DATEN_AB + N - 1],
                         "kandidaten": ["40 05 13"], "pdf": []}
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and not any("Specs" in w for w in plan.warnungen)
