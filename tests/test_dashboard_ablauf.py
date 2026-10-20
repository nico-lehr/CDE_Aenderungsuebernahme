"""Ablauf des Dashboards mit erfundenen Daten (Regel 16).

1. API: Upload eines erfundenen Paars -> Plan mit Schema-4-Feldern -> Freigabe + Abwahl -> Ergebnis enthält die abgewählte Änderung
   NICHT, die freigegebene schon -> Download von Excel und Protokoll; die Original-Datei bleibt bytegleich.
2. Kommandozeile: `lauf --abwahl ID` reicht die Abwahl durch.
3. Browser (Chromium headless über Playwright für Node, nur wenn vorhanden, sonst übersprungen): Seite lädt ohne Konsolenfehler,
   Karte per Taste „A“ freigeben, „Excel erstellen“ -> Schritt 3.
4. Spec (Welle 3): erfundene CDE mit drei Specs. Spec des PDFs erkannt/vorgegeben, `neu-analysieren` (neuer Lauf, alter bleibt),
   Änderung landet nur in der gewählten Spec, fremde Spec-Zeilen bleiben gleich; unklare Spec: nichts automatisch, im Browser
   „Welche Spec gehört zu diesem PDF?“ und „Excel erstellen“ blockiert, bis eine Spec gewählt ist.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import subprocess
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import cli, config, protokoll, runs
from src.api.app import app
from tests.fixtures.szenario import erzeuge
from tests.test_register import GRUND, wort_pdf
from tests.test_spec import N, _mappe, _werte, _zeile


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


@pytest.fixture
def paar(tmp_path):
    """Erfundenes Paar mit genau einem Vorschlag ("Rack{+s+}", teilweise hinterlegt) und sonst sicheren Änderungen."""
    return erzeuge(tmp_path / "fx", "hintergrund")


def _cde_text(daten: bytes) -> str:
    ws = load_workbook(io.BytesIO(daten))["CDE"]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


# ---------------------------------------------------------------------------------------------------------------- 1. API
@pytest.mark.parametrize("markierung", ["aus", pytest.param("an", marks=pytest.mark.markierung_an)])
def test_ablauf_freigabe_und_abwahl_bis_zum_download(paar, markierung):
    original_vorher = _sha(paar["alt"])
    client = app.test_client()
    with open(paar["pdf"], "rb") as p, open(paar["alt"], "rb") as x:
        r = client.post("/api/lauf", data={"pdf": (p, "aenderungen.pdf"), "excel": (x, "cde.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    plan = r.get_json()
    rid = plan["run_id"]

    # Plan fürs Dashboard: Schema 4 (bereich, stuecke, grund_einfach) ist da, die Einordnung kommt vom Server
    for a in plan["aenderungen"]:
        assert isinstance(a["bereich"], str) and isinstance(a["stuecke"], list) and isinstance(a["grund_einfach"], str)
    z = plan["zusammenfassung"]
    assert z["gesamt"] == len(plan["aenderungen"]) and z["pruefen"] == 1 and z["automatisch"] == z["gesamt"] - 1
    (vorschlag,) = [a for a in plan["aenderungen"] if a["kategorie"] == "pruefen"]
    (auto,) = [a for a in plan["aenderungen"] if a["kategorie"] == "automatisch" and (a["alt"], a["neu"]) == ("zwei", "drei")]

    u = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": [vorschlag["id"]], "abwahl": [auto["id"]]})
    assert u.status_code == 200, u.get_json()
    res = u.get_json()
    nach_id = {a["id"]: a for a in res["aenderungen"]}
    assert nach_id[vorschlag["id"]]["status"] == "uebernommen" and nach_id[vorschlag["id"]]["bestaetigt"]
    assert nach_id[auto["id"]]["status"] == "nicht_uebernommen" and nach_id[auto["id"]]["kategorie"] == "nicht"
    assert res["bilanz"]["stimmt"] and res["bilanz"]["uebernommen"] == z["automatisch"]       # -1 abgewählt, +1 freigegeben

    # Download: die Excel enthält die freigegebene Änderung, die abgewählte NICHT
    excel = client.get(f"/api/lauf/{rid}/download/{res['ordner']}/{res['excel']}")
    assert excel.status_code == 200
    text = _cde_text(excel.data)
    assert "Die Racks werden" in text
    assert "dauert zwei Wochen" in text and "drei Wochen" not in text
    proto = client.get(f"/api/lauf/{rid}/download/{res['ordner']}/aenderungsprotokoll.xlsx")
    assert proto.status_code == 200
    zeilen = list(load_workbook(io.BytesIO(proto.data))["Änderungen"].iter_rows(values_only=True))
    (zeile,) = [z_ for z_ in zeilen if z_[0] == auto["id"]]
    assert zeile[4] == "nicht übernommen" and "abgewählt" in str(zeile[10])                # nichts geht still verloren

    # Ergebnisliste nennt Freigabe und Abwahl; die Original-Datei (und ihre Kopie im Lauf) bleibt bytegleich
    erg = client.get(f"/api/lauf/{rid}").get_json()["ergebnisse"]
    assert erg[-1]["freigaben"] == [vorschlag["id"]] and erg[-1]["abwahl"] == [auto["id"]]
    assert _sha(paar["alt"]) == original_vorher == _sha(runs.lauf_dir(rid) / "eingabe" / "cde.xlsx") == plan["excel_sha256"]


# ---------------------------------------------------------------------------------------------------------------- 2. CLI
def test_cli_lauf_mit_abwahl(paar, capsys):
    def neuer() -> str:                                 # der eben angelegte Lauf (Laeufe derselben Sekunde sortieren zufaellig)
        (rid,) = {e["run_id"] for e in runs.liste()} - bekannt
        bekannt.add(rid)
        return rid
    bekannt: set[str] = set()
    assert cli.main(["analyse", "--pdf", str(paar["pdf"]), "--excel", str(paar["alt"])]) == 0
    plan = runs.lade_plan(neuer())
    (auto,) = [a["id"] for a in plan["aenderungen"] if (a["alt"], a["neu"]) == ("zwei", "drei")]
    capsys.readouterr()
    assert cli.main(["lauf", "--pdf", str(paar["pdf"]), "--excel", str(paar["alt"]), "--abwahl", f" {auto} ,"]) == 0
    out = capsys.readouterr().out
    assert f"Abgewählt (nicht eingetragen): {auto}" in out
    assert runs.ergebnisse(neuer())[0]["abwahl"] == [auto]
    # unbekannte ID: klare Meldung, kein Ergebnis
    assert cli.main(["lauf", "--pdf", str(paar["pdf"]), "--excel", str(paar["alt"]), "--abwahl", "A9999"]) == 2
    assert "Unbekannte" in capsys.readouterr().out
    assert runs.ergebnisse(neuer()) == []


# ---------------------------------------------------------------------------------------------------------------- 3. Browser
_NODE_SKRIPT = r"""
const { chromium } = require('playwright');
const [basis, pdf, xlsx] = process.argv.slice(2);
(async () => {
  let b;
  try { b = await chromium.launch(); } catch (e) { console.log(JSON.stringify({ kein_browser: String(e).slice(0, 200) })); return; }
  const p = await b.newPage({ viewport: { width: 1280, height: 900 } });
  const fehler = [];
  p.on('console', (m) => { if (m.type() === 'error') fehler.push(m.text()); });
  p.on('pageerror', (e) => fehler.push(String(e)));
  const aus = { fehler };
  await p.goto(basis + '/');
  await p.waitForFunction(() => document.getElementById('pill-text').textContent.includes('Port'));
  aus.schritt1 = await p.isVisible('#schritt-dateien');
  aus.aktuell1 = await p.getAttribute('#schritt-knopf-dateien', 'aria-current');
  await p.setInputFiles('#file-pdf', pdf);
  await p.setInputFiles('#file-xlsx', xlsx);
  await p.click('#btn-abgleichen');
  await p.waitForSelector('#schritt-pruefen:not([hidden]) #karten-bitte .karte', { timeout: 60000 });
  aus.karten = await p.locator('#karten-bitte .karte').count();
  aus.satz_vorher = await p.textContent('#fortschritt-satz');
  // Vertrauensmodus: Dokument-Ampel oben, Stichprobe (eingeklappter Block nur wenn vorhanden) steht VOR „Bitte prüfen“
  aus.ampel_vorher = await p.getAttribute('#ampel', 'data-stufe');
  aus.ampel_satz = await p.textContent('#ampel-satz');
  aus.sp_karten = await p.locator('#karten-stichprobe .karte').count();
  aus.sp_offen = await p.evaluate(() => document.getElementById('klapp-stichprobe').open);
  aus.sp_zuerst = await p.evaluate(() => !!(document.getElementById('klapp-stichprobe').compareDocumentPosition(document.getElementById('bereich-bitte')) & Node.DOCUMENT_POSITION_FOLLOWING));
  await p.keyboard.press('?');
  aus.hilfe_offen = await p.evaluate(() => document.getElementById('hilfe').open);
  await p.keyboard.press('Escape');
  aus.hilfe_zu = await p.evaluate(() => !document.getElementById('hilfe').open);
  await p.keyboard.press('j');                                  // erste Karte wählen: die erste der Stichprobe
  aus.fokus_karte = await p.evaluate(() => document.activeElement.classList.contains('karte'));
  aus.erste_ist_stichprobe = await p.evaluate(() => !!document.activeElement.closest('#karten-stichprobe'));
  await p.keyboard.press('a');                                  // Stichprobe: Passt (springt zur nächsten offenen Karte)
  aus.sp1 = await p.getAttribute('#karten-stichprobe li:nth-child(1) .karte', 'data-entscheid');
  aus.sp_stand = await p.textContent('#stichprobe-stand');
  await p.keyboard.press('n');                                  // zweite Karte: Falsch (wählt die Änderung ab)
  aus.sp2 = await p.getAttribute('#karten-stichprobe li:nth-child(2) .karte', 'data-entscheid');
  aus.sp2_abgewaehlt = await p.evaluate(() => {
    const id = document.querySelector('#karten-stichprobe li:nth-child(2) .karte').dataset.id;
    return [...document.querySelectorAll('#karten-auto .karte')].find((k) => k.dataset.id === id).classList.contains('abgewaehlt');
  });
  aus.ampel_mitte = await p.getAttribute('#ampel', 'data-stufe');
  await p.keyboard.press('a');                                  // dritte Karte: Passt -> weiter zur ersten offenen Karte unter „Bitte prüfen“
  aus.sp3 = await p.getAttribute('#karten-stichprobe li:nth-child(3) .karte', 'data-entscheid');
  aus.fokus_bitte = await p.evaluate(() => !!document.activeElement.closest('#karten-bitte'));
  await p.keyboard.press('a');                                  // übernehmen
  aus.gedrueckt = await p.getAttribute('#karten-bitte .karte .entscheid.ja', 'aria-pressed');
  aus.entscheid = await p.getAttribute('#karten-bitte .karte', 'data-entscheid');
  aus.satz_nachher = await p.textContent('#fortschritt-satz');
  aus.balken = await p.getAttribute('#fortschritt', 'aria-valuenow');
  // alles entschieden und die Stichprobe angesehen: die Ampel (Regeln kommen vom Server) wird grün
  await p.waitForFunction(() => document.getElementById('ampel').dataset.stufe === 'gruen', null, { timeout: 15000 });
  aus.ampel_nachher = await p.textContent('#ampel-satz');
  await p.click('#btn-excel');
  await p.waitForSelector('#schritt-fertig:not([hidden])', { timeout: 60000 });
  aus.fertig_titel = await p.textContent('#titel-fertig');
  aus.fertig_satz = await p.textContent('#fertig-satz');
  aus.aktuell3 = await p.getAttribute('#schritt-knopf-fertig', 'aria-current');
  aus.links = await p.$$eval('#fertig-knoepfe a', (l) => l.map((a) => a.getAttribute('href')));
  aus.fertig_fokus = await p.evaluate(() => document.activeElement.id);
  await b.close();
  console.log(JSON.stringify(aus));
})().catch((e) => { console.log(JSON.stringify({ absturz: String(e) })); process.exit(1); });
"""


def _node() -> str | None:
    return shutil.which("node") or (str(p) if (p := Path("/opt/node22/bin/node")).exists() else None)


def _node_umgebung() -> dict[str, str]:
    env = dict(os.environ)
    pfade = [x for x in env.get("NODE_PATH", "").split(os.pathsep) if x]
    knoten = _node()
    if knoten:                                                  # globale Module neben node (z. B. /opt/node22/lib/node_modules)
        pfade.append(str(Path(knoten).resolve().parent.parent / "lib" / "node_modules"))
    env["NODE_PATH"] = os.pathsep.join(pfade)
    return env


def _playwright_da() -> bool:
    if not _node():
        return False
    r = subprocess.run([_node(), "-e", "require.resolve('playwright')"], capture_output=True, env=_node_umgebung(), timeout=30)
    return r.returncode == 0


@pytest.mark.skipif(not _playwright_da(), reason="node mit Playwright nicht vorhanden")
def test_browser_karte_per_taste_freigeben_und_excel_erstellen(paar, tmp_path, monkeypatch):
    from werkzeug.serving import make_server

    server = make_server("127.0.0.1", 0, app, threaded=True)
    monkeypatch.setitem(config.lade("server"), "port", server.server_port)     # Origin-Schutz: der Test-Port ist der eigene
    faden = threading.Thread(target=server.serve_forever, daemon=True)
    faden.start()
    skript = tmp_path / "ablauf.js"
    skript.write_text(_NODE_SKRIPT, encoding="utf-8")
    try:
        r = subprocess.run([_node(), str(skript), f"http://127.0.0.1:{server.server_port}", str(paar["pdf"]), str(paar["alt"])],
                           capture_output=True, text=True, env=_node_umgebung(), timeout=180)
    finally:
        server.shutdown()
        faden.join(10)
    zeilen = [z for z in r.stdout.splitlines() if z.startswith("{")]
    assert zeilen, f"keine Ausgabe vom Browser-Skript: {r.stdout[-500:]} {r.stderr[-500:]}"
    aus = json.loads(zeilen[-1])
    if "kein_browser" in aus:
        pytest.skip(f"Chromium nicht startbar: {aus['kein_browser']}")
    assert "absturz" not in aus, aus
    assert aus["fehler"] == [], f"Konsolenfehler: {aus['fehler']}"
    assert aus["schritt1"] and aus["aktuell1"] == "step"
    assert aus["karten"] == 1 and "1 bitte prüfen (0 erledigt)" in aus["satz_vorher"] and "Stichprobe 0 von 3 angesehen" in aus["satz_vorher"]
    assert aus["ampel_vorher"] == "gelb" and aus["ampel_satz"] == "Bitte ein paar Punkte ansehen."
    assert aus["sp_karten"] == 3 and aus["sp_offen"] and aus["sp_zuerst"]
    assert aus["hilfe_offen"] and aus["hilfe_zu"]
    # J wählt die erste Karte (die Stichprobe steht zuerst); A = Passt, N = Falsch (wählt ab), springt jeweils weiter
    assert aus["fokus_karte"] and aus["erste_ist_stichprobe"] and aus["sp1"] == "passt" and aus["sp_stand"] == "1 von 3 angesehen"
    assert aus["sp2"] == "falsch" and aus["sp2_abgewaehlt"] and aus["ampel_mitte"] == "gelb" and aus["sp3"] == "passt" and aus["fokus_bitte"]
    assert aus["gedrueckt"] == "true" and aus["entscheid"] == "ja" and aus["ampel_nachher"] == "Alles sicher eingetragen."
    assert "(1 erledigt)" in aus["satz_nachher"] and aus["balken"] == "1"
    assert aus["aktuell3"] == "step" and aus["fertig_titel"] == "Fertig." and "Änderungen eingetragen" in aus["fertig_satz"]
    assert aus["fertig_fokus"] == "titel-fertig"
    assert len(aus["links"]) == 2 and all(l.startswith("/api/lauf/") for l in aus["links"])
    # Auf dem Server: die per Taste freigegebene Änderung ist als Freigabe gespeichert
    (lauf,) = runs.liste()
    plan = runs.lade_plan(lauf["run_id"])
    (vorschlag,) = [a["id"] for a in plan["aenderungen"] if a["status"] == "pruefen"]
    erg = runs.ergebnisse(lauf["run_id"])[0]
    assert erg["freigaben"] == [vorschlag] and len(erg["abwahl"]) == 1 and erg["stichprobe"]["falsch"] == erg["abwahl"]
    assert len(erg["stichprobe"]["passt"]) == 2 and set(erg["stichprobe"]["passt"] + erg["stichprobe"]["falsch"]) == set(plan["stichprobe"])
    zeilen, _ = protokoll.lese_entscheidungen(runs.lauf_dir(lauf["run_id"]) / "entscheidungen.jsonl")
    assert {z["id"]: z["entscheidung"] for z in zeilen}[erg["abwahl"][0]] == "abgelehnt"           # „Falsch“ im Freigabe-Protokoll
    assert sum(1 for z in zeilen if z["entscheidung"] == "bestaetigt") == 2


# ---------------------------------------------------------------------------------------------------------------- 4. Spec
_PDF_SPEC = list(GRUND)
_PDF_SPEC[2] = ("sub", "A.", "Die Pumpe {-foerdert-}{+hebt+} das Wasser in das obere Becken.")     # steht in allen drei Specs
_TITEL_44 = "SECTION 44 55 66 – PUMPENANLAGE"


@pytest.fixture
def spec_paar(tmp_path):
    """Erfundene CDE mit drei Specs (11 22 33, 44 55 66, 77 88 99) und ein PDF, das im Titel die Spec 44 55 66 nennt."""
    return {"xlsx": _mappe(tmp_path), "pdf": wort_pdf(tmp_path / "a.pdf", _PDF_SPEC, [_TITEL_44]),
            "pdf_unklar": wort_pdf(tmp_path / "unklar.pdf", _PDF_SPEC, None)}


def _hochladen_spec(client, pdf, xlsx, **form):
    with open(pdf, "rb") as p, open(xlsx, "rb") as x:
        return client.post("/api/lauf", data={"pdf": (p, "aenderungen.pdf"), "excel": (x, "cde.xlsx"), **form},
                           content_type="multipart/form-data")


@pytest.mark.legende_neu
@pytest.mark.markierung_an
def test_spec_erkannt_neu_analysieren_und_nur_die_gewaehlte_spec_wird_bearbeitet(spec_paar):
    client = app.test_client()
    original = spec_paar["xlsx"].read_bytes()
    plan = _hochladen_spec(client, spec_paar["pdf"], spec_paar["xlsx"]).get_json()
    rid = plan["run_id"]
    assert plan["spec"]["erkannt"] == "44 55 66" and plan["spec"]["quelle"] == "pdf"
    assert plan["spec"]["zeilen"] == [_zeile(1, 0), _zeile(1, N - 1)]
    assert [(x["name"], x["zeilen"], x["von"], x["bis"]) for x in plan["specs"]] == [
        ("11 22 33", N, _zeile(0, 0), _zeile(0, N - 1)), ("44 55 66", N, _zeile(1, 0), _zeile(1, N - 1)),
        ("77 88 99", N, _zeile(2, 0), _zeile(2, N - 1))]
    g = client.get(f"/api/lauf/{rid}/specs").get_json()
    assert g == {"specs": plan["specs"], "spec": plan["spec"]}
    (ae,) = [a for a in plan["aenderungen"] if a["alt"] == "foerdert"]
    assert ae["kategorie"] == "automatisch" and ae["ziel"] == f"Abgleich!G{_zeile(1, 2)}"
    assert any("Specs" in w for w in plan["warnungen"])

    # Andere Spec waehlen: neuer Lauf mit denselben Eingabedateien, der erste bleibt unveraendert (Regel 29)
    plan_vorher = (runs.lauf_dir(rid) / "plan.json").read_bytes()
    r = client.post(f"/api/lauf/{rid}/neu-analysieren", json={"spec": "77.88.99"})
    assert r.status_code == 200, r.get_json()
    neu = r.get_json()
    assert neu["run_id"] != rid and neu["spec"]["erkannt"] == "77 88 99" and neu["spec"]["quelle"] == "vorgabe"
    (ae2,) = [a for a in neu["aenderungen"] if a["alt"] == "foerdert"]
    assert ae2["kategorie"] == "automatisch" and ae2["ziel"] == f"Abgleich!G{_zeile(2, 2)}"
    assert (runs.lauf_dir(rid) / "plan.json").read_bytes() == plan_vorher
    assert (runs.lauf_dir(neu["run_id"]) / "eingabe" / "cde.xlsx").read_bytes() == original == spec_paar["xlsx"].read_bytes()
    liste = client.get("/api/lauf").get_json()
    assert {e["run_id"]: e["spec"] for e in liste} == {rid: "44 55 66", neu["run_id"]: "77 88 99"}

    # Eintragen: nur die Zeilen der gewaehlten Spec aendern sich, alle anderen Zeilen bleiben gleich
    u = client.post(f"/api/lauf/{neu['run_id']}/uebernehmen", json={}).get_json()
    ziel = runs.lauf_dir(neu["run_id"]) / u["ordner"] / u["excel"]
    fremd = [r_ for r_ in range(8, 8 + 3 * N) if not _zeile(2, 0) <= r_ < _zeile(2, 0) + N]
    assert _werte(ziel, fremd) == _werte(spec_paar["xlsx"], fremd)
    assert _werte(ziel, [_zeile(2, 2)]) != _werte(spec_paar["xlsx"], [_zeile(2, 2)])
    assert "hebt" in "".join(str(v) for v in _werte(ziel, [_zeile(2, 2)]).values())
    proto = client.get(f"/api/lauf/{neu['run_id']}/download/{u['ordner']}/aenderungsprotokoll.xlsx").data
    zeilen = list(load_workbook(io.BytesIO(proto))["Bilanz"].iter_rows(values_only=True))
    assert any(z and z[0] == "Spec dieses Abgleichs" and "77 88 99" in str(z[1]) for z in zeilen)
    assert spec_paar["xlsx"].read_bytes() == original


@pytest.mark.legende_neu
def test_spec_unklar_traegt_nichts_automatisch_ein_und_ungueltige_spec_wird_abgelehnt(spec_paar):
    client = app.test_client()
    plan = _hochladen_spec(client, spec_paar["pdf_unklar"], spec_paar["xlsx"]).get_json()
    assert plan["spec"]["erkannt"] is None and len(plan["spec"]["kandidaten"]) == 3
    assert plan["zusammenfassung"]["automatisch"] == 0 and plan["zusammenfassung"]["pruefen"] >= 1
    assert not [a for a in plan["aenderungen"] if a["status"] == "uebernommen"]
    assert any("Spec des PDFs nicht in der CDE gefunden" in w for w in plan["warnungen"])
    # unbekannte Spec: kein Fehler, aber Warnung und nichts automatisch
    r = client.post(f"/api/lauf/{plan['run_id']}/neu-analysieren", json={"spec": "99 99 99"})
    assert r.status_code == 200 and r.get_json()["spec"]["erkannt"] is None
    assert r.get_json()["zusammenfassung"]["automatisch"] == 0
    assert any("vorgegebene Spec" in w for w in r.get_json()["warnungen"])
    # Ungueltiges: nur Zeichen und Laenge, kein Freitext; es entsteht kein neuer Lauf
    vorher = len(client.get("/api/lauf").get_json())
    for schlecht in ("a;b", "11 22 33'; x", "x" * 41, "<b>1</b>", "../etc"):
        r = client.post(f"/api/lauf/{plan['run_id']}/neu-analysieren", json={"spec": schlecht})
        assert r.status_code == 400 and r.is_json and "Spec" in r.get_json()["fehler"], schlecht
        r = _hochladen_spec(client, spec_paar["pdf"], spec_paar["xlsx"], spec=schlecht)
        assert r.status_code == 400 and r.is_json, schlecht
    for leer in ({}, {"spec": ""}, {"spec": "  "}, {"spec": 5}, [1]):
        r = client.post(f"/api/lauf/{plan['run_id']}/neu-analysieren", json=leer)
        assert r.status_code == 400 and r.is_json
    assert len(client.get("/api/lauf").get_json()) == vorher
    assert client.post("/api/lauf/20260101-000000-abcdef/neu-analysieren", json={"spec": "11 22 33"}).status_code == 404
    # Spec als Form-Feld beim Hochladen
    r = _hochladen_spec(client, spec_paar["pdf_unklar"], spec_paar["xlsx"], spec="11-22-33")
    assert r.status_code == 200 and r.get_json()["spec"]["erkannt"] == "11 22 33" and r.get_json()["spec"]["quelle"] == "vorgabe"


@pytest.mark.legende_neu
def test_cli_lauf_mit_spec(spec_paar, capsys):
    assert cli.main(["analyse", "--pdf", str(spec_paar["pdf_unklar"]), "--excel", str(spec_paar["xlsx"])]) == 0
    assert "nicht eindeutig" in capsys.readouterr().out
    assert cli.main(["analyse", "--pdf", str(spec_paar["pdf_unklar"]), "--excel", str(spec_paar["xlsx"]), "--spec", "77 88 99"]) == 0
    assert "Spec: 77 88 99 (Zeilen" in capsys.readouterr().out
    n = len(runs.liste())
    assert cli.main(["analyse", "--pdf", str(spec_paar["pdf"]), "--excel", str(spec_paar["xlsx"]), "--spec", "a;b"]) == 2
    assert "Spec darf nur" in capsys.readouterr().out and len(runs.liste()) == n          # vor dem Anlegen abgelehnt


@pytest.mark.legende_neu
def test_eine_spec_in_der_cde_hat_hoechstens_eine_spec(tmp_path):
    """Eine CDE mit nur einer Spec: das Dashboard blendet die Spec-Zeile aus (specs/kandidaten haben hoechstens einen Eintrag)."""
    from tests.test_register import excel_absaetze, register_mappe
    client = app.test_client()
    xlsx = register_mappe(tmp_path / "eine.xlsx", excel_absaetze(GRUND))
    pdf = wort_pdf(tmp_path / "a.pdf", _PDF_SPEC, [_TITEL_44])
    plan = _hochladen_spec(client, pdf, xlsx).get_json()
    assert len(plan["specs"]) <= 1 and (plan["spec"] is None or len(plan["spec"]["kandidaten"]) <= 1)


# ---- Browser: unklare Spec
_NODE_SPEC = r"""
const { chromium } = require('playwright');
const [basis, pdf, xlsx] = process.argv.slice(2);
(async () => {
  let b;
  try { b = await chromium.launch(); } catch (e) { console.log(JSON.stringify({ kein_browser: String(e).slice(0, 200) })); return; }
  const p = await b.newPage({ viewport: { width: 1280, height: 900 } });
  const fehler = [];
  p.on('console', (m) => { if (m.type() === 'error') fehler.push(m.text()); });
  p.on('pageerror', (e) => fehler.push(String(e)));
  const aus = { fehler };
  await p.goto(basis + '/');
  await p.waitForFunction(() => document.getElementById('pill-text').textContent.includes('Port'));
  await p.setInputFiles('#file-pdf', pdf);
  await p.setInputFiles('#file-xlsx', xlsx);
  await p.click('#btn-abgleichen');
  await p.waitForSelector('#schritt-pruefen:not([hidden])', { timeout: 60000 });
  aus.zone = await p.isVisible('#spec-zone');
  aus.wahl = await p.isVisible('#spec-wahl');
  aus.zeile_aus = !(await p.isVisible('#spec-zeile'));
  aus.titel = await p.textContent('#spec-titel');
  aus.rolle = await p.getAttribute('#spec-wahl', 'role');
  aus.knoepfe = await p.$$eval('#spec-liste button', (l) => l.map((x) => x.getAttribute('aria-label')));
  aus.hinweis = await p.textContent('#excel-hinweis');
  await p.click('#btn-excel');
  await p.waitForFunction(() => document.getElementById('toast-fehler').textContent.length > 0);
  aus.fehlertext = await p.textContent('#toast-fehler');
  aus.noch_pruefen = await p.isVisible('#schritt-pruefen');
  await p.click('#toast-schliessen');
  await p.click('#spec-liste button[data-spec="44 55 66"]');
  await p.waitForFunction(() => document.getElementById('spec-text').textContent.includes('44 55 66'), null, { timeout: 60000 });
  aus.spec_text = await p.textContent('#spec-text');
  aus.wahl_zu = !(await p.isVisible('#spec-wahl'));
  aus.andere = await p.getAttribute('#btn-spec-andere', 'aria-expanded');
  await p.click('#btn-spec-andere');
  aus.andere_offen = await p.getAttribute('#btn-spec-andere', 'aria-expanded');
  aus.wahl_offen = await p.isVisible('#spec-wahl');
  aus.aktuell = await p.getAttribute('#spec-liste button[aria-current="true"]', 'data-spec');
  aus.satz = await p.textContent('#fortschritt-satz');
  await p.click('#btn-excel');
  await p.waitForSelector('#schritt-fertig:not([hidden])', { timeout: 60000 });
  aus.fertig = await p.textContent('#fertig-satz');
  await b.close();
  console.log(JSON.stringify(aus));
})().catch((e) => { console.log(JSON.stringify({ absturz: String(e) })); process.exit(1); });
"""


@pytest.mark.legende_neu
@pytest.mark.skipif(not _playwright_da(), reason="node mit Playwright nicht vorhanden")
def test_browser_unklare_spec_blockiert_excel_erstellen_bis_eine_gewaehlt_ist(spec_paar, tmp_path, monkeypatch):
    from werkzeug.serving import make_server

    server = make_server("127.0.0.1", 0, app, threaded=True)
    monkeypatch.setitem(config.lade("server"), "port", server.server_port)
    faden = threading.Thread(target=server.serve_forever, daemon=True)
    faden.start()
    skript = tmp_path / "spec.js"
    skript.write_text(_NODE_SPEC, encoding="utf-8")
    try:
        r = subprocess.run([_node(), str(skript), f"http://127.0.0.1:{server.server_port}", str(spec_paar["pdf_unklar"]), str(spec_paar["xlsx"])],
                           capture_output=True, text=True, env=_node_umgebung(), timeout=240)
    finally:
        server.shutdown()
        faden.join(10)
    zeilen = [z for z in r.stdout.splitlines() if z.startswith("{")]
    assert zeilen, f"keine Ausgabe vom Browser-Skript: {r.stdout[-500:]} {r.stderr[-500:]}"
    aus = json.loads(zeilen[-1])
    if "kein_browser" in aus:
        pytest.skip(f"Chromium nicht startbar: {aus['kein_browser']}")
    assert "absturz" not in aus, aus
    assert aus["fehler"] == [], f"Konsolenfehler: {aus['fehler']}"
    assert aus["zone"] and aus["wahl"] and aus["zeile_aus"] and aus["rolle"] == "alert"
    assert aus["titel"] == "Welche Spec gehört zu diesem PDF?" and len(aus["knoepfe"]) == 3
    assert all("Zeilen" in k for k in aus["knoepfe"]) and "Bitte zuerst die Spec wählen" in aus["hinweis"]
    assert "welche Spec" in aus["fehlertext"] and aus["noch_pruefen"]
    erste = sorted(e["run_id"] for e in runs.liste())[0]
    assert runs.ergebnisse(erste) == []                                                      # der unklare Lauf bekam kein Ergebnis
    assert "44 55 66" in aus["spec_text"] and "Zeilen" in aus["spec_text"] and aus["wahl_zu"]
    assert aus["andere"] == "false" and aus["andere_offen"] == "true" and aus["wahl_offen"] and aus["aktuell"] == "44 55 66"
    assert "eingetragen" in aus["fertig"]


# ---------------------------------------------------------------------------------------------------------------- 5. Word-Eingabe
# Das Word-Lesen kommt von einem anderen Helfer (src/docx_reader.py); hier ein Fake wie in tests/test_vertrauen.py.
from tests.test_vertrauen import _docx_datei, docx_leser  # noqa: E402,F401  (docx_leser ist eine Fixture)

_NODE_DOCX = r"""
const { chromium } = require('playwright');
const [basis, docx, xlsx] = process.argv.slice(2);
(async () => {
  let b;
  try { b = await chromium.launch(); } catch (e) { console.log(JSON.stringify({ kein_browser: String(e).slice(0, 200) })); return; }
  const p = await b.newPage({ viewport: { width: 1280, height: 900 } });
  const fehler = [];
  p.on('console', (m) => { if (m.type() === 'error') fehler.push(m.text()); });
  p.on('pageerror', (e) => fehler.push(String(e)));
  const aus = { fehler };
  await p.goto(basis + '/');
  await p.waitForFunction(() => document.getElementById('pill-text').textContent.includes('Port'));
  aus.accept = await p.getAttribute('#file-pdf', 'accept');
  await p.setInputFiles('#file-pdf', docx);
  aus.dz_sub = await p.textContent('#dz-pdf-sub');
  await p.setInputFiles('#file-xlsx', xlsx);
  await p.click('#btn-abgleichen');
  await p.waitForSelector('#schritt-pruefen:not([hidden]) #karten-bitte .karte', { timeout: 60000 });
  aus.lauf_name = await p.textContent('#lauf-name');
  aus.bilder = await p.locator('.im-pdf, #schritt-pruefen img').count();
  aus.text = await p.evaluate(() => document.getElementById('schritt-pruefen').innerText);
  aus.sp_karten = await p.locator('#karten-stichprobe .karte').count();
  aus.vorschau_da = await p.locator('#karten-bitte .karte .vorschau').count();
  aus.orte = await p.$$eval('#karten-bitte .karte h3', (l) => l.map((x) => x.textContent));
  await p.click('#fachleute summary');
  aus.fachleute = await p.textContent('#fachleute-inhalt');
  await p.keyboard.press('Escape');
  await p.click('#btn-excel');
  await p.waitForSelector('#schritt-fertig:not([hidden])', { timeout: 60000 });
  aus.fertig = await p.textContent('#fertig-satz');
  await b.close();
  console.log(JSON.stringify(aus));
})().catch((e) => { console.log(JSON.stringify({ absturz: String(e) })); process.exit(1); });
"""


@pytest.mark.skipif(not _playwright_da(), reason="node mit Playwright nicht vorhanden")
def test_browser_word_dokument_zeigt_nur_den_text_ohne_pdf_ausschnitt(paar, tmp_path, monkeypatch, docx_leser):
    from werkzeug.serving import make_server

    docx = _docx_datei(tmp_path / "Änderungen.docx")
    server = make_server("127.0.0.1", 0, app, threaded=True)
    monkeypatch.setitem(config.lade("server"), "port", server.server_port)
    faden = threading.Thread(target=server.serve_forever, daemon=True)
    faden.start()
    skript = tmp_path / "docx.js"
    skript.write_text(_NODE_DOCX, encoding="utf-8")
    try:
        r = subprocess.run([_node(), str(skript), f"http://127.0.0.1:{server.server_port}", str(docx), str(paar["alt"])],
                           capture_output=True, text=True, env=_node_umgebung(), timeout=180)
    finally:
        server.shutdown()
        faden.join(10)
    zeilen = [z for z in r.stdout.splitlines() if z.startswith("{")]
    assert zeilen, f"keine Ausgabe vom Browser-Skript: {r.stdout[-500:]} {r.stderr[-500:]}"
    aus = json.loads(zeilen[-1])
    if "kein_browser" in aus:
        pytest.skip(f"Chromium nicht startbar: {aus['kein_browser']}")
    assert "absturz" not in aus, aus
    assert aus["fehler"] == [], f"Konsolenfehler: {aus['fehler']}"
    assert ".docx" in aus["accept"] and "✓ Änderungen.docx" in aus["dz_sub"] and aus["lauf_name"].startswith("Änderungen.docx")
    # Word: nur der Text (Ergebnis wie in der Excel), kein PDF-Ausschnitt, keine „Seite 0“, kein „PDF“ in den Karten
    assert aus["bilder"] == 0 and "Seite 0" not in aus["text"] and "Im PDF" not in aus["text"] and "PDF-Ausschnitt" not in aus["text"]
    assert aus["vorschau_da"] == 1 and aus["sp_karten"] == 3 and all(o.startswith("Section") or o.startswith("Im Word-Dokument") for o in aus["orte"])
    assert "Word-Dokument" in aus["fachleute"] and "PDF-Seiten" not in aus["fachleute"] and "Word-Änderungsverfolgung direkt gelesen" in aus["fachleute"]
    assert "eingetragen" in aus["fertig"]
    (lauf,) = runs.liste()
    assert runs.lade_plan(lauf["run_id"])["quelle_art"] == "docx"
