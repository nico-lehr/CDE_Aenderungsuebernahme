"""Vorschlaege (Status "pruefen") bei abweichender Excel und bei unsicherer Fortsetzung (Auftrag 004).

Grundsatz (Regel 27): Ein Vorschlag entsteht nur, wenn er mit hoher Sicherheit stimmt -- sonst bleibt die Aenderung auf "pruefen",
aber OHNE Vorschlag (keine Operationen, Code `kein_vorschlag`). Die PDFs baut tests/_word_pdf.py in der Form der Word-Aenderungs-
verfolgung; die Excel hat einen Buchstabendreher (die Excel weicht anderswo im Absatz vom PDF ab). Nur erfundene Daten (Regel 16).
"""
from __future__ import annotations

import copy

import pytest
from openpyxl import load_workbook

from src import config, messen
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.models import PRUEFEN
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests._word_pdf import WordDok
from tests.redteam._bau import schreibe_excel

pytestmark = [pytest.mark.legende_neu, pytest.mark.markierung_an]


def _lauf(tmp_path, zeilen_pdf, excel_text, name="a"):
    """Word-PDF (Ueberschrift 1.1 + ein Absatz) und Excel (eine Zeile mit `excel_text`) -> (Plan, Zellstuecke nach Freigabe aller
    Vorschlaege als [(format, text)], Excel-Datei vorher)."""
    d = WordDok()
    d.ueberschrift([("1.1", "k")], [("ANLAGE", "k")])
    d.absatz(zeilen_pdf)
    pdf = d.speichern(tmp_path / f"{name}.pdf")
    alt = schreibe_excel(tmp_path / f"{name}.xlsx", [["Nr.", "Titel", "Inhalt"], ["1.1", "ANLAGE", None], ["", "", excel_text]])
    plan = plane(analysiere_pdf(pdf), ExcelModell(alt))
    out = tmp_path / f"{name}_frei.xlsx"
    uebernehmen(copy.deepcopy(plan.aenderungen), alt, out, {a.id for a in plan.aenderungen if a.status == PRUEFEN},
                absaetze=plan.absaetze)
    zelle = messen._lese_excel(out, True)[0][2][2]
    return plan, messen._stuecke(zelle), alt


def _klassen_text(stuecke, klasse):
    return "".join(t for k, t in stuecke if k == klasse)


def test_klammern_bei_abweichender_excel_werden_zeichengenau_gestrichen(tmp_path):
    """Im PDF sind nur "[" und "]" gestrichen, der Wert "63" bleibt. Der alte Matcher machte daraus "[63]" -> "63" (Wert gestrichen
    und neu geschrieben). Jetzt: genau die Klammern sind gestrichen, der Wert bleibt, die Abweichung der Excel bleibt stehen."""
    pdf_zeile = [[("Die Lieferung umfasst ", "k"), ("[", "d"), ("63", "k"), ("]", "d"),
                  (" Kupplungen sowie Ersatzdichtungen. Das gilt nur fuer Pumpen.", "k")]]
    plan, stuecke, _ = _lauf(tmp_path, pdf_zeile,
                             "Die Lieferung umfasst [63] Kupplungen sowie Ersatzdichtungen. Das gilt nur fur Pumpen.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "excel_weicht_ab" in a.pruefcodes and "kein_vorschlag" not in a.pruefcodes
    assert len(a.ops) == 2
    assert stuecke == [("normal", "Die Lieferung umfasst "), ("alt", "["), ("normal", "63"), ("alt", "]"),
                       ("normal", " Kupplungen sowie Ersatzdichtungen. Das gilt nur fur Pumpen.")]


def test_neues_wort_vor_dem_punkt_bei_abweichender_excel_laesst_den_punkt_stehen(tmp_path):
    """Nur "jeweils" ist neu; der Punkt dahinter gehoert nicht zur Aenderung (der alte Matcher ersetzte "Pumpen." als Wort)."""
    pdf_zeile = [[("Die Anlage gilt fuer alle Pumpen", "k"), (" jeweils", "n"), (". Das gilt nur fuer Armaturen im Keller.", "k")]]
    plan, stuecke, _ = _lauf(tmp_path, pdf_zeile, "Die Anlage gilt fuer alle Pumpen. Das gilt nur fur Armaturen im Keller.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "excel_weicht_ab" in a.pruefcodes and a.ops
    assert _klassen_text(stuecke, "alt") == ""
    assert _klassen_text(stuecke, "neu").strip() == "jeweils"
    assert "".join(t for _, t in stuecke) == "Die Anlage gilt fuer alle Pumpen jeweils. Das gilt nur fur Armaturen im Keller."


def test_neuer_satz_bei_abweichender_excel_ist_samt_leerzeichen_markiert_wie_im_pdf(tmp_path):
    """Word unterstreicht das Leerzeichen vor dem neuen Satz mit. Ohne PDF-Absatz markierte der Applier nur nach den Operationen
    (Leerraum am Rand unmarkiert); jetzt legt der Planer die PDF-Laeufe an die abweichende Excel an, und die Projektion markiert
    wie das PDF -- das Leerzeichen gehoert zum neuen Satz, der Punkt davor bleibt unmarkiert."""
    pdf_zeile = [[("Die Anlage gilt fuer alle Pumpen.", "k"), (" Abweichungen sind schriftlich zu beantragen.", "n")]]
    plan, stuecke, _ = _lauf(tmp_path, pdf_zeile, "Die Anlage gilt fur alle Pumpen.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and "excel_weicht_ab" in a.pruefcodes and a.ops
    assert len(plan.absaetze) == 1 and a.id in plan.absaetze[0].aenderungen
    assert stuecke == [("normal", "Die Anlage gilt fur alle Pumpen."), ("neu", " Abweichungen sind schriftlich zu beantragen.")]


def test_ersetztes_wort_bei_abweichender_excel_laesst_trenner_und_abweichung_stehen(tmp_path):
    """Alt direkt vor Neu (ohne Leerzeichen dazwischen), davor und dahinter unmarkierter Text; die Abweichung der Excel
    ("fur" statt "fuer") bleibt unveraendert und unmarkiert."""
    pdf_zeile = [[("Die Anlage gilt fuer alle ", "k"), ("Pumpen", "d"), ("Armaturen", "n"), (" im Keller und im Hof des Werks.", "k")]]
    plan, stuecke, _ = _lauf(tmp_path, pdf_zeile, "Die Anlage gilt fur alle Pumpen im Keller und im Hof des Werks.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and a.ops and len(plan.absaetze) == 1
    assert stuecke == [("normal", "Die Anlage gilt fur alle "), ("alt", "Pumpen"), ("neu", "Armaturen"),
                       ("normal", " im Keller und im Hof des Werks.")]


def test_ohne_zeichengenaue_stelle_gibt_es_keinen_vorschlag(tmp_path):
    """Der Text direkt hinter dem gestrichenen Wort ist in der Excel ein anderer ("xx" statt "im"): die Stelle ist nicht
    zeichengenau bestimmbar. Der alte Matcher hat das Wort trotzdem ersetzt (geraten); jetzt bleibt die Aenderung auf "pruefen"
    OHNE Vorschlag, und eine Freigabe schreibt nichts."""
    pdf_zeile = [[("Die Anlage gilt fuer alle ", "k"), ("Pumpen", "d"), ("Armaturen", "n"),
                  (" im Keller und im Hof des Werks.", "k")]]
    plan, stuecke, alt = _lauf(tmp_path, pdf_zeile, "Die Anlage gilt fuer alle Pumpen xx Keller und im Hof des Werks.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and a.ops == [] and "kein_vorschlag" in a.pruefcodes
    assert "excel_weicht_ab" in a.pruefcodes
    assert stuecke == [("normal", "Die Anlage gilt fuer alle Pumpen xx Keller und im Hof des Werks.")]      # Freigabe schreibt nichts
    assert load_workbook(str(alt))["CDE"].cell(3, 3).value == "Die Anlage gilt fuer alle Pumpen xx Keller und im Hof des Werks."


def test_wortgenauer_vorschlag_nur_auf_ausdrueckliche_einstellung(tmp_path, monkeypatch):
    """`planer.vorschlag_wortgenau: true` bringt den frueheren Wort-fuer-Wort-Vorschlag des alten Matchers zurueck."""
    monkeypatch.setitem(config.lade("zuordnung")["planer"], "vorschlag_wortgenau", True)
    pdf_zeile = [[("Die Anlage gilt fuer alle ", "k"), ("Pumpen", "d"), ("Armaturen", "n"),
                  (" im Keller und im Hof des Werks.", "k")]]
    plan, _, _ = _lauf(tmp_path, pdf_zeile, "Die Anlage gilt fuer alle Pumpen xx Keller und im Hof des Werks.")
    (a,) = plan.aenderungen
    assert a.status == PRUEFEN and a.ops and "kein_vorschlag" not in a.pruefcodes
