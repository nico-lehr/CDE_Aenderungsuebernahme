"""Unit-Tests K-FINAL: Funde der unabhaengigen Pruefer (nur erfundene Daten, Regel 16)."""
from __future__ import annotations

from openpyxl import Workbook

from src import config
from src.excel_model import ExcelModell
from src.matcher import Matcher
from src.models import HOCH, MITTEL, PRUEFEN, UEBERNOMMEN
from src.textutil import norm
from tests.redteam._bau import analysiere, baue_alt, bearbeite, lauf, uebernimm, blatt


def _tokens(text: str) -> list[tuple[int, str]]:
    return [(i, norm(w)) for i, w in enumerate(text.split()) if norm(w)]


def _excel(tmp_path, zeilen, name="a.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "CDE"
    for z in zeilen:
        ws.append(z)
    p = tmp_path / name
    wb.save(str(p))
    return p


_ZWILLING = "Die Reaktionszeit der Anlage darf im Normalbetrieb hundertfuenfzig Millisekunden nicht ueberschreiten."
_AUFBAU = [["Nr.", "Titel", "Inhalt"],
           ["1", "Allgemeines", "Der Betrieb erfolgt jederzeit durch geschultes Personal vor Ort."],
           ["2", "Stoermeldungen", _ZWILLING],
           ["3", "Normen", "Die Norm gilt fuer alle Anlagen im Betrieb."],
           ["4", "Reaktion", "Im Normalbetrieb darf die Anlage hoechstens hundertfuenfzig Millisekunden brauchen."]]


# ===========================================================================================================
# 1. Matcher: eine falsche Section ist nie automatisch
# ===========================================================================================================
def test_zwilling_in_anderer_section_ist_nur_mittel(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    m = Matcher(xl, config.lade("zuordnung"))
    t = m.finde(_tokens(_ZWILLING), "4", ist_kopf=False)         # PDF-Section 4, der Text steht aber in Section 2
    assert t is not None and t.kandidat.section == "2" and t.section_passt is False and t.eindeutig
    assert t.stufe == MITTEL
    assert any("Section passt nicht" in g for g in t.gruende)


def test_passende_section_bleibt_hoch(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    t = Matcher(xl, config.lade("zuordnung")).finde(_tokens(_ZWILLING), "2", ist_kopf=False)
    assert t is not None and t.section_passt is True and t.stufe == HOCH


def test_ohne_section_im_pdf_bleibt_hoch(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    t = Matcher(xl, config.lade("zuordnung")).finde(_tokens(_ZWILLING), None, ist_kopf=False)
    assert t is not None and t.section_passt is None and t.stufe == HOCH


def test_schalter_section_unpassend_stufe_wirkt(tmp_path):
    xl = ExcelModell(_excel(tmp_path, _AUFBAU))
    cfg = config.lade("zuordnung")
    cfg = {**cfg, "matcher": {**cfg["matcher"], "section_unpassend_stufe": "hoch"}}
    t = Matcher(xl, cfg).finde(_tokens(_ZWILLING), "4", ist_kopf=False)
    assert t is not None and t.stufe == HOCH


_SATZ = "Die Reaktionszeit der Anlage {-darf-} {+soll+} im Normalbetrieb hundertfuenfzig Millisekunden nicht ueberschreiten."
_ELEM_ZWILLING = [
    ("h", "1 Einleitung"),
    ("p", "Der Betrieb erfolgt jederzeit durch geschultes Personal vor Ort."),
    ("h", "2 Stoermeldungen"),
    ("p", _ZWILLING),                       # Section 2 bleibt im PDF unveraendert
    ("h", "3 Normen"),
    ("p", "Die Norm gilt fuer alle Anlagen im Betrieb."),
    ("h", "4 Reaktion"),
    ("p", _SATZ),
]


def test_zwilling_in_section_2_wird_nicht_fuer_section_4_geschrieben(tmp_path):
    """Nachbau des Pruefer-Falls: Die Aenderung gehoert zu Section 4, deren Excel-Text abweicht (umformuliert);
    Section 2 hat denselben alten Satz. Vorher wurde die Aenderung automatisch in Section 2 geschrieben."""
    def umformulieren(pfad):
        def fn(wb):
            ws = wb["CDE"]
            for row in ws.iter_rows(min_row=2):
                if str(row[0].value).strip() == "4":
                    row[2].value = "Im Normalbetrieb darf die Anlage hoechstens hundertfuenfzig Millisekunden brauchen."
        bearbeite(pfad, fn)

    alt = baue_alt(tmp_path, _ELEM_ZWILLING, "L1", umformulieren)
    l = uebernimm(analysiere(tmp_path, _ELEM_ZWILLING, alt))            # ohne Freigabe
    # Der Satz in Section 2 steht in der Ausgabe unveraendert da (nur ein Mensch darf ihn aendern)
    zeilen = {str(r[0].value).strip(): r[2].value for r in blatt(l.out).iter_rows(min_row=2)}
    assert "darf im Normalbetrieb" in zeilen["2"] and "soll" not in zeilen["2"], zeilen["2"]
    zu_4 = [a for a in l.aen if a.section == "4" and a.typ != "section_neu" and "soll" in a.neu]
    assert zu_4 and all(a.status == PRUEFEN for a in zu_4), [(a.status, a.konfidenz, a.ziel) for a in l.aen]


def test_umnummerierung_bleibt_automatisch(tmp_path):
    """Die Section wird umnummeriert (alt 2.3 -> neu 2.2): der Vergleich nutzt die ALTE Nummer, es bleibt automatisch."""
    elem = [
        ("h", "1 Einleitung"),
        ("p", "Dieses Dokument beschreibt die Anforderungen an das Speichersystem fuer alle Standorte des Kunden."),
        ("h", "2.1 Leistung"),
        ("p", "Die Nennleistung betraegt 50 kW und darf kurzzeitig ueberschritten werden."),
        ("h", "{-2.3-} {+2.2+} Sicherheit"),
        ("p", "Das System muss gegen {-Ueberhitzung-} {+Ueberspannung+} geschuetzt sein und dies im Betrieb nachweisen."),
        ("h", "3 Lieferung"),
        ("p", "Die Lieferung erfolgt frei Haus und wird vom Kunden abgenommen."),
    ]
    l = lauf(tmp_path, elem)
    a = [x for x in l.aen if "Ueberspannung" in x.neu]
    assert a and a[0].status == UEBERNOMMEN and a[0].konfidenz == HOCH, [(x.typ, x.status, x.grund) for x in l.aen]


# ===========================================================================================================
# 2. Applier: Gruppen-Mitfreigabe nur durch einen Menschen; keine Ueberschreibung; kein Ersetzen von plan.json
# ===========================================================================================================
_ELEM_GRUPPE = [
    ("h", "1 Einleitung"),
    ("p", "Dieses Dokument beschreibt die Anforderungen an das Speichersystem fuer alle Standorte."),
    ("h", "{+2 Dokumentation+}"),
    ("p", "{+Alle Unterlagen werden in deutscher Sprache geliefert.+}"),
    ("p", "{+Pruefberichte liegen dem Angebot bei.+}"),
    ("h", "3 Lieferung"),
    ("p", "Die Lieferung erfolgt frei Haus und wird vom Kunden abgenommen."),
]


def _gruppe(tmp_path):
    alt = baue_alt(tmp_path, _ELEM_GRUPPE, "L1")
    l = analysiere(tmp_path, _ELEM_GRUPPE, alt)
    kopf = next(a for a in l.aen if a.typ == "section_neu")
    folgen = [a for a in l.aen if a.typ == "absatz_neu"]
    assert len(folgen) == 2
    return l, kopf, folgen


def _texte(pfad):
    return [str(c.value) for row in blatt(pfad).iter_rows() for c in row if c.value is not None]


def test_gemischte_gruppe_ohne_freigabe_wird_nicht_mitgenommen(tmp_path):
    """Handgebauter Plan: neue Section 'uebernommen', ein Folgeabsatz 'pruefen', keine Freigabe -> nichts wird geschrieben."""
    l, kopf, folgen = _gruppe(tmp_path)
    for a in [kopf] + folgen:
        a.status = UEBERNOMMEN
    folgen[1].status = PRUEFEN
    uebernimm(l)                                              # keine Freigaben
    assert not any("Pruefberichte" in t or "Dokumentation" in t for t in _texte(l.out))
    assert [a.status for a in [kopf] + folgen] == [PRUEFEN, PRUEFEN, PRUEFEN]
    assert all(a.hinweise for a in [kopf] + folgen if a is not folgen[1]), "Hinweis fehlt"


def test_gemischte_gruppe_mit_menschlicher_freigabe_kommt_zusammen(tmp_path):
    l, kopf, folgen = _gruppe(tmp_path)
    for a in [kopf] + folgen:
        a.status = UEBERNOMMEN
    folgen[1].status = PRUEFEN
    uebernimm(l, freigaben={folgen[1].id})
    texte = " ".join(_texte(l.out))
    assert "Pruefberichte" in texte and "Dokumentation" in texte and "Unterlagen" in texte
    assert all(a.status == UEBERNOMMEN for a in [kopf] + folgen)


def test_uebernehmen_ueberschreibt_nie_eine_vorhandene_zieldatei(tmp_path):
    import pytest
    from src.applier import uebernehmen
    l, kopf, folgen = _gruppe(tmp_path)
    ziel = tmp_path / "ergebnis.xlsx"
    ziel.write_bytes(b"alter Inhalt")
    with pytest.raises(FileExistsError, match="nicht überschrieben"):
        uebernehmen(l.aen, l.alt, ziel)
    assert ziel.read_bytes() == b"alter Inhalt"


# ===========================================================================================================
# 2c / 3. Runs und API: Meldungen, Rohtext, Schema-1-Bilanz, Server-Einstellungen
# ===========================================================================================================
import io
import json

import pytest

from src import runs
from src.api.app import app
from src.fehler import NutzerFehler
from tests.fixtures.szenario import erzeuge


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


def _hochladen(client, d, excel_bytes=None):
    excel = io.BytesIO(excel_bytes) if excel_bytes is not None else open(d["alt"], "rb")
    return client.post("/api/lauf", data={"pdf": (open(d["pdf"], "rb"), "a.pdf"), "excel": (excel, "cde.xlsx")},
                       content_type="multipart/form-data")


def _fehler_dateien():
    return [p.read_text(encoding="utf-8") for p in runs.runs_dir().rglob("analyse_fehler.json")]


def test_analyse_mit_rohtext_valueerror_zeigt_dem_client_nie_den_rohtext(client, tmp_path, monkeypatch):
    def kaputt(*a, **k):
        raise ValueError("GEHEIMER_ROHTEXT C:\\intern\\x.py")
    monkeypatch.setattr(runs, "analysiere", kaputt)
    r = _hochladen(client, erzeuge(tmp_path / "fx"))
    assert r.status_code == 500 and r.is_json
    assert "GEHEIMER_ROHTEXT" not in r.get_data(as_text=True)
    assert all("GEHEIMER_ROHTEXT" not in t for t in _fehler_dateien()) and _fehler_dateien()


def test_analyse_mit_gewollter_meldung_kommt_beim_client_an(client, tmp_path, monkeypatch):
    def gewollt(*a, **k):
        raise NutzerFehler("Das ist eine einfache Meldung.", 400)
    monkeypatch.setattr(runs, "analysiere", gewollt)
    r = _hochladen(client, erzeuge(tmp_path / "fx"))
    assert r.status_code == 400 and r.get_json()["fehler"] == "Das ist eine einfache Meldung."


def test_beschaedigte_excel_gibt_einfache_meldung_400(client, tmp_path):
    r = _hochladen(client, erzeuge(tmp_path / "fx"), excel_bytes=b"PK\x03\x04 kein echtes zip")
    assert r.status_code == 400 and "beschädigt" in r.get_json()["fehler"]
    assert "Traceback" not in r.get_data(as_text=True)


@pytest.mark.parametrize("pfad", ["/api/lauf/{rid}", "/api/lauf/{rid}/download/eingabe/a.pdf"])
def test_valueerror_in_anderen_routen_ist_500_ohne_rohtext(client, tmp_path, monkeypatch, pfad):
    rid = _hochladen(client, erzeuge(tmp_path / "fx")).get_json()["run_id"]

    def kaputt(*a, **k):
        raise ValueError("GEHEIMER_ROHTEXT")
    monkeypatch.setattr(runs, "lade_plan", kaputt)
    monkeypatch.setattr(runs, "datei", kaputt)
    r = client.get(pfad.format(rid=rid))
    assert r.status_code == 500 and "GEHEIMER_ROHTEXT" not in r.get_data(as_text=True)


def test_unbekannter_lauf_bleibt_404(client):
    r = client.get("/api/lauf/20260101-000000-abcdef")
    assert r.status_code == 404 and r.get_json()["fehler"] == "Lauf nicht gefunden."


def test_analysiere_ersetzt_nie_ein_vorhandenes_plan_json(tmp_path):
    d = erzeuge(tmp_path / "fx")
    rid = runs.neuer_lauf(d["pdf"], d["alt"])
    runs.analysiere(rid)
    plan = runs.lauf_dir(rid) / "plan.json"
    vorher = plan.read_bytes()
    with pytest.raises(runs.LaufFehler, match="schon analysiert"):
        runs.analysiere(rid)
    assert plan.read_bytes() == vorher


def test_schema1_plan_behauptet_nie_vollstaendig(tmp_path):
    d = erzeuge(tmp_path / "fx")
    rid = runs.neuer_lauf(d["pdf"], d["alt"])
    runs.analysiere(rid)
    p = runs.lauf_dir(rid) / "plan.json"
    plan = json.loads(p.read_text(encoding="utf-8"))
    for k in ("luecken", "schema_version"):                     # so sah ein Schema-1-Plan aus
        plan.pop(k)
    for k in ("luecken", "vollstaendig"):
        plan["bilanz"].pop(k)
    p.write_text(json.dumps(plan), encoding="utf-8")
    res = runs.uebernehmen_lauf(rid, [])
    assert res["bilanz"]["vollstaendig"] is None and res["bilanz"]["luecken_geprueft"] is False
    from openpyxl import load_workbook
    ws = load_workbook(runs.lauf_dir(rid) / res["ordner"] / "aenderungsprotokoll.xlsx").active
    zellen = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "JA – nichts ging still verloren" not in zellen


def test_server_ohne_host_oder_port_bricht_mit_klarer_meldung_ab(monkeypatch):
    from src.api import einstellungen
    monkeypatch.setattr(einstellungen.config, "lade", lambda name: {"port": 5100})
    with pytest.raises(SystemExit, match="'host'"):
        einstellungen.srv()
    monkeypatch.setattr(einstellungen.config, "lade", lambda name: {"host": "127.0.0.1"})
    with pytest.raises(SystemExit, match="'port'"):
        einstellungen.srv()
    monkeypatch.setattr(einstellungen.config, "lade", lambda name: {"host": "127.0.0.1", "port": 5100})
    assert einstellungen.srv()["upload_max_mb"] == 200               # Upload-Grenzen behalten Standardwerte


# ===========================================================================================================
# 4. CLI analyse/lauf: Luecken sichtbar, einfache Fehlermeldungen, Hinweise einmal, Lauf-Hinweis
# ===========================================================================================================
from src import cli
from tests.redteam._bau_pdf import RohSeite, roh_pdf, zeile

_TEXT = ("Dieser Satz steht auf der Seite und ist lang genug damit die Seite als Textseite gilt "
         "und niemand sie fuer einen Scan haelt weil genug Woerter darauf stehen und weiter und weiter")
_BILD = "q 595 0 0 842 0 0 cm /Im1 Do Q\n"


def _cli(capsys, *args):
    rc = cli.main(list(args))
    return rc, capsys.readouterr().out


@pytest.fixture
def paar_mit_luecke(tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(roh_pdf([RohSeite(inhalt=zeile(56, 64, _TEXT, 11.0, "0 0 0 rg")), RohSeite(inhalt=_BILD, bild=True)]))
    return pdf, _excel(tmp_path, _AUFBAU)


def test_cli_zeigt_unvollstaendig_und_die_luecken_liste(capsys, paar_mit_luecke):
    pdf, excel = paar_mit_luecke
    rc, out = _cli(capsys, "analyse", "--pdf", str(pdf), "--excel", str(excel))
    assert rc == 0
    assert "UNVOLLSTÄNDIG" in out and "Bilanz stimmt: ja" not in out
    assert "Seite 2 | Scan-Seite |" in out                               # Seite, Art, Beschreibung


def test_cli_lauf_zeigt_luecken_nur_einmal_und_hinweise_nur_einmal(capsys, paar_mit_luecke):
    pdf, excel = paar_mit_luecke
    rc, out = _cli(capsys, "lauf", "--pdf", str(pdf), "--excel", str(excel))
    assert rc == 0 and out.count("UNVOLLSTÄNDIG") == 2                   # einmal je Bilanz (Analyse und Ergebnis)
    assert out.count("Seite 2 | Scan-Seite |") == 1
    hinweise = [z for z in out.splitlines() if z.startswith("Hinweis:")]
    assert hinweise and len(hinweise) == len(set(hinweise))


def test_cli_analyse_sagt_dass_ein_lauf_angelegt_und_keine_excel_geschrieben_wird(capsys, paar_mit_luecke):
    pdf, excel = paar_mit_luecke
    rc, out = _cli(capsys, "analyse", "--pdf", str(pdf), "--excel", str(excel))
    assert "Die Analyse liegt als Lauf in" in out and "nichts geschrieben" in out
    assert len(runs.liste()) == 1


def test_cli_vollstaendig_zeigt_ja(capsys, tmp_path):
    pdf = tmp_path / "ok.pdf"
    pdf.write_bytes(roh_pdf([RohSeite(inhalt=zeile(56, 64, _TEXT, 11.0, "0 0 0 rg"))]))
    rc, out = _cli(capsys, "analyse", "--pdf", str(pdf), "--excel", str(_excel(tmp_path, _AUFBAU)))
    assert rc == 0 and "Bilanz stimmt: ja" in out and "UNVOLLSTÄNDIG" not in out


def test_cli_fehlende_datei_gibt_meldung_und_exit_2(capsys, tmp_path, paar_mit_luecke):
    pdf, excel = paar_mit_luecke
    rc, out = _cli(capsys, "analyse", "--pdf", str(tmp_path / "gibt_es_nicht.pdf"), "--excel", str(excel))
    assert rc == 2 and "PDF-Datei wurde nicht gefunden" in out and "Traceback" not in out
    rc, out = _cli(capsys, "lauf", "--pdf", str(pdf), "--excel", str(tmp_path / "gibt_es_nicht.xlsx"))
    assert rc == 2 and "Excel-Datei wurde nicht gefunden" in out


@pytest.mark.parametrize("welche", ["pdf", "excel"])
def test_cli_kaputte_datei_gibt_einfache_meldung_und_exit_2(capsys, tmp_path, paar_mit_luecke, welche):
    pdf, excel = paar_mit_luecke
    kaputt = tmp_path / ("kaputt.pdf" if welche == "pdf" else "kaputt.xlsx")
    kaputt.write_bytes(b"das ist keine echte Datei")
    args = ["--pdf", str(kaputt if welche == "pdf" else pdf), "--excel", str(kaputt if welche == "excel" else excel)]
    rc, out = _cli(capsys, "analyse", *args)
    assert rc == 2 and ("konnte nicht" in out or "beschädigt" in out) and "Traceback" not in out and "Error" not in out
