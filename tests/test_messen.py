"""Selbsttest des Messwerkzeugs (src/messen.py): Es muss Fehler finden, die wir absichtlich einbauen.

Nur erfundene Daten (Regel 16). Alles in tmp_path; conftest.py setzt CDE_DATA_ROOT.
Die Paare werden hier selbst gebaut (pdf_writer + openpyxl), nicht mit dem Generator.
"""
import hashlib
import json
import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import cli, config, messen, runs
from src.models import PRUEFEN, SCHEMA_VERSION
from tests.fixtures.pdf_writer import baue_pdf
from tests.fixtures.szenario import baue_excel, zustaende

# Markante erfundene Woerter: duerfen nie in Bericht oder Konsole auftauchen.
MARKANT = ["quokka", "flamingo", "pinguin", "laub", "gras", "links", "rechts", "teich", "gefieder", "schatten",
           "entfällt", "zweiter absatz", "letzter absatz", "komplett", "ersatz"]

SAUBER = [
    ("h", "1 Eins"),
    ("p", "Der Quokka frisst {-Gras-} {+Laub+} am Morgen. Dann ruht er lange im Schatten."),
    ("p", "Zweiter Absatz bleibt unverändert hier stehen."),
    ("h", "2 Zwei"),
    ("p", "Ein Flamingo steht {-links-} {+rechts+} im Teich und putzt sein rosa Gefieder."),
    ("p", "Der Teich ist sehr tief."),
    ("h", "3 Drei"),
    ("p", "{-Dieser Absatz entfällt komplett im Dokument ohne Ersatz.-}"),
    ("p", "Letzter Absatz bleibt erhalten und ruhig."),
]
MIT_NEUER_SECTION = SAUBER + [("h", "{+4 Vier+}"), ("p", "{+Nur ein neuer Absatz mit einem Pinguin darin.+}")]


def baue_paar(eltern: Path, name: str, elemente=SAUBER, layout: str = "L1") -> Path:
    p = eltern / name
    p.mkdir(parents=True)
    baue_pdf(p / "aenderungen.pdf", elemente, "farbe")
    baue_excel(p / "vorher.xlsx", zustaende(elemente, False), layout)
    baue_excel(p / "nachher.xlsx", zustaende(elemente, True), layout)
    return p


def korpus_dir(korpus: str = "entwicklung") -> Path:
    d = config.pfad("paare_dir") / korpus
    d.mkdir(parents=True, exist_ok=True)
    return d


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def zellen_aendern(pfad: Path, aenderungen: dict[str, str]) -> None:
    wb = load_workbook(str(pfad))
    ws = wb["CDE"]
    for adresse, text in aenderungen.items():
        ws[adresse].value = text
    wb.save(str(pfad))


def sabotiere_ergebnis(monkeypatch, wirkung):
    """Der ERSTE uebernehmen()-Aufruf (Automatik) wird nachtraeglich veraendert: wirkung(ziel_pfad, aenderungen)."""
    original = messen.uebernehmen
    zaehler = {"n": 0}

    def falsch(aenderungen, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aenderungen, quelle, ziel, freigaben, cfg)
        zaehler["n"] += 1
        if zaehler["n"] == 1:
            wirkung(Path(ziel), aenderungen, Path(quelle))
        return res
    monkeypatch.setattr(messen, "uebernehmen", falsch)


def arten(res) -> set[str]:
    return {f["art"] for f in res["fehler"]}


# ---------------------------------------------------------------------------------------
# (1) Sauberes Paar
# ---------------------------------------------------------------------------------------
def test_sauberes_paar_hat_keine_stillen_fehler(tmp_path):
    p = baue_paar(tmp_path, "synth_0001_standard")
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 0 and res["fehler"] == []
    assert res["soll_zellen"] == 3                       # drei Zellen aendern sich laut nachher
    assert res["erreicht_auto"] == 3
    assert res["auto_gesamt"] == 3 and res["auto_richtig"] == 3
    assert res["kollateral"] == 0
    assert res["bilanz_stimmt"] and res["vollstaendig"]
    assert res["idempotent"] is True
    assert res["original_unveraendert"] is True
    assert res["lese_luecke_woerter"] == 0
    assert res["pruefen_zellen"] == 0 and res["nicht_uebernommen_zellen"] == 0
    assert res["vorschlaege_gesamt"] is None             # ohne --mit-vorschlaegen nicht gemessen


@pytest.mark.parametrize("layout", ["L1", "L2", "L3"])
def test_sauber_in_allen_layouts(tmp_path, layout):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_0002", layout=layout))
    assert res["stille_fehler"] == 0 and res["fehler"] == []
    assert res["erreicht_auto"] == res["soll_zellen"] > 0


def test_neue_section_wird_als_neue_zeile_gemessen(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_0003", MIT_NEUER_SECTION))
    assert res["stille_fehler"] == 0 and res["fehler"] == []
    assert res["erreicht_auto"] == res["soll_zellen"] and res["soll_zellen"] > 3


def test_original_wird_nie_veraendert(tmp_path):
    p = baue_paar(tmp_path, "synth_0004")
    vorher = {n: sha(p / n) for n in messen.DATEIEN}
    messen.messe_paar(p)
    assert {n: sha(p / n) for n in messen.DATEIEN} == vorher


def test_original_veraenderung_wird_gemeldet(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_0005")
    original = messen.uebernehmen

    def boese(aenderungen, quelle, ziel, freigaben=None, cfg=None, **_):
        Path(quelle).write_bytes(Path(quelle).read_bytes() + b"x")      # veraendert das Original
        return original(aenderungen, quelle, ziel, freigaben, cfg)
    monkeypatch.setattr(messen, "uebernehmen", boese)
    res = messen.messe_paar(p)
    assert res["original_unveraendert"] is False


# ---------------------------------------------------------------------------------------
# (2) Sabotage: das Messwerkzeug muss die eingebauten Fehler finden
# ---------------------------------------------------------------------------------------
def test_nichts_uebernommen_wird_als_luecke_erkannt(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1000")
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, quelle: shutil.copy2(quelle, ziel))   # E = vorher
    res = messen.messe_paar(p)
    assert res["erreicht_auto"] == 0 and res["soll_zellen"] == 3
    assert res["stille_fehler"] >= 3                                    # "uebernommen" gemeldet, Zelle unveraendert
    assert "keine_wirkung" in res["stille_fehler_nach_art"]
    assert "plan_falsch" in arten(res)                                  # Ziel stimmt, Ergebnis nicht


def test_manipulierte_zielzelle_ist_stiller_fehler(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1001")
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"C2": "Etwas ganz anderes steht hier."}))
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["erreicht_auto"] == 2
    assert res["stille_fehler_nach_art"].get("uebernommen_aber_falsch") == 1
    assert res["auto_richtig"] == 2 and res["auto_gesamt"] == 3
    assert [f["art"] for f in res["fehler"]] == ["plan_falsch"]
    assert res["fehler"][0]["aenderung"].startswith("A") and res["fehler"][0]["stumm"] is True
    assert (res["fehler"][0]["blatt"], res["fehler"][0]["zeile"], res["fehler"][0]["spalte"]) == (2, 2, 3)   # Blatt 2 = "CDE"


def test_veraenderte_fremde_zelle_ist_kollateralschaden(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1002")
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"B3": "Fremder Titel"}))
    res = messen.messe_paar(p)
    assert res["kollateral"] == 1 and res["stille_fehler"] == 1
    assert res["erreicht_auto"] == 3                                    # die eigentlichen Aenderungen sind richtig
    assert res["auto_richtig"] == 3 and res["auto_gesamt"] == 4


def test_nachkontrolle_ok_trotz_falschem_inhalt(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1003")

    def wirkung(ziel, aend, quelle):
        zellen_aendern(ziel, {"C2": "Falscher Inhalt, aber das Tool sagt ok."})
        for a in aend:
            a.nachkontrolle = "ok"
    sabotiere_ergebnis(monkeypatch, wirkung)
    res = messen.messe_paar(p)
    assert res["stille_fehler_nach_art"].get("nachkontrolle_ok_aber_falsch") == 1
    assert res["stille_fehler"] == 1


def test_zweiter_lauf_der_etwas_aendert_wird_erkannt(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1010")
    original = messen.uebernehmen
    zaehler = {"n": 0}

    def unsauber(aenderungen, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aenderungen, quelle, ziel, freigaben, cfg)
        zaehler["n"] += 1
        if zaehler["n"] == 2:                      # zweiter Lauf (Idempotenz-Pruefung) aendert wieder etwas
            zellen_aendern(Path(ziel), {"C3": "Noch eine Aenderung beim zweiten Lauf."})
        return res
    monkeypatch.setattr(messen, "uebernehmen", unsauber)
    res = messen.messe_paar(p)
    assert res["idempotent"] is False and res["stille_fehler"] == 0


def test_bilanz_stimmt_false_ist_bilanz_kaputt(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1004")
    original = messen.bilanz
    monkeypatch.setattr(messen, "bilanz", lambda *a, **k: {**original(*a, **k), "stimmt": False})
    res = messen.messe_paar(p)
    assert res["bilanz_stimmt"] is False
    assert "bilanz_kaputt" in arten(res)
    assert res["stille_fehler_nach_art"].get("bilanz") == 1 and res["stille_fehler"] >= 1


def test_falsche_zuordnung_gleiche_spalte(tmp_path, monkeypatch):
    """Zwei Zellen in derselben Spalte vertauscht -> richtige Spalte, falsche Zeile."""
    p = baue_paar(tmp_path, "synth_1005")

    def wirkung(ziel, aend, quelle):
        ws = load_workbook(str(ziel))["CDE"]
        a, b = ws["C2"].value, ws["C3"].value
        zellen_aendern(ziel, {"C2": b, "C3": a})
    sabotiere_ergebnis(monkeypatch, wirkung)
    res = messen.messe_paar(p)
    assert res["erreicht_auto"] == 1
    assert arten(res) == {"zuordnung_falsch"}
    assert res["stille_fehler"] >= 2


def test_falsche_spalte_ist_layout_falsch(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1006")

    def wirkung(ziel, aend, quelle):
        ws = load_workbook(str(ziel))["CDE"]
        neu = ws["C2"].value
        zellen_aendern(ziel, {"D2": neu, "C2": load_workbook(str(quelle))["CDE"]["C2"].value})
    sabotiere_ergebnis(monkeypatch, wirkung)
    res = messen.messe_paar(p)
    assert "layout_falsch" in arten(res)


def test_falscher_plan_ziel_stimmt_ergebnis_nicht(tmp_path, monkeypatch):
    """Ein Wort doppelt eingetragen: Ziel und Woerter stimmen, das Ergebnis nicht -> plan_falsch."""
    p = baue_paar(tmp_path, "synth_1007")

    def wirkung(ziel, aend, quelle):
        ws = load_workbook(str(ziel))["CDE"]
        zellen_aendern(ziel, {"C2": ws["C2"].value.replace("Laub", "Laub Laub")})
    sabotiere_ergebnis(monkeypatch, wirkung)
    res = messen.messe_paar(p)
    assert arten(res) == {"plan_falsch"} and res["stille_fehler"] == 1


def test_pdf_falsch_gelesen(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1008")
    original = messen.analysiere_pdf

    def blind(pfad):
        pa = original(pfad)
        for b in pa.blocks:                                  # ein neues Wort "verschluckt" der Leser
            b.tokens = [t for t in b.tokens if not (t.status == "add" and t.new == "Laub")]
        return pa
    monkeypatch.setattr(messen, "analysiere_pdf", blind)
    res = messen.messe_paar(p)
    assert "pdf_falsch_gelesen" in arten(res)
    assert res["lese_luecke_woerter"] >= 1 and res["erreicht_auto"] < res["soll_zellen"]


def test_abgestuerzte_paare_zaehlen_als_unklar_ohne_text(tmp_path, monkeypatch):
    korpus = korpus_dir()
    baue_paar(korpus, "synth_2000")
    baue_paar(korpus, "synth_2001")
    original = messen.plane
    zaehler = {"n": 0}

    def kaputt(pdf, xl, cfg=None):
        zaehler["n"] += 1
        if zaehler["n"] == 1:
            raise RuntimeError("GEHEIMER-DOKUMENTTEXT-quokka")
        return original(pdf, xl, cfg)
    monkeypatch.setattr(messen, "plane", kaputt)
    assert messen.main(["--korpus", "entwicklung"]) == 0
    bericht = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    assert bericht["fehlerarten"]["unklar"] == 1
    assert bericht["gesamt"]["paare"] == 2 and bericht["gesamt"]["paare_abgestuerzt"] == 1
    assert bericht["paare"][0]["fehler"][0]["grund"] == "RuntimeError"
    text = json.dumps(bericht)
    assert "GEHEIMER" not in text and "quokka" not in text.lower()


# ---------------------------------------------------------------------------------------
# Vorschlaege (Status "pruefen")
# ---------------------------------------------------------------------------------------
def test_vorschlaege_getrennt_ausgewiesen(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_3000")
    original = messen.plane

    def mit_vorschlag(pdf, xl, cfg=None):
        plan = original(pdf, xl, cfg)
        ae = next(a for a in plan.aenderungen if a.ziel.endswith("C2"))
        ae.status, ae.konfidenz = PRUEFEN, "mittel"
        return plan
    monkeypatch.setattr(messen, "plane", mit_vorschlag)
    res = messen.messe_paar(p, mit_vorschlaegen=True)
    assert res["pruefen_zellen"] == 1 and res["erreicht_auto"] == 2
    assert res["fehler"] == [] and res["stille_fehler"] == 0              # "pruefen" ist kein Fehler
    assert res["auto_gesamt"] == 2 and res["auto_richtig"] == 2
    assert res["vorschlaege_gesamt"] == 1 and res["vorschlaege_richtig"] == 1


# ---------------------------------------------------------------------------------------
# erwartet.json (optional)
# ---------------------------------------------------------------------------------------
def test_erwartet_json_stoerer_und_luecken(tmp_path):
    korpus = korpus_dir()
    p = baue_paar(korpus, "synth_4000")
    (p / "erwartet.json").write_text(json.dumps({
        "format": 1, "profil": "standard", "stoerer": ["fussnote"], "woerter_markiert": 99,
        "luecken": [{"art": "scan_seite", "seite": 2}]}), encoding="utf-8")
    baue_paar(korpus, "synth_4001")                               # ohne erwartet.json: muss ebenfalls laufen
    assert messen.main(["--korpus", "entwicklung"]) == 1          # nicht gemeldete Luecke = stiller Fehler -> Ziel verfehlt (Exit-Code 1)
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    assert b["je_stoerer"]["fussnote"]["paare"] == 1
    erst = b["paare"][0]
    assert erst["luecken_nicht_gemeldet"] == 1                    # erwartete Luecke nicht gemeldet
    assert erst["stille_fehler_nach_art"].get("luecke_nicht_gemeldet") == 1
    assert erst["lese_luecke_erwartet"] > 0
    assert b["paare"][1]["stille_fehler"] == 0


# ---------------------------------------------------------------------------------------
# (3) Keine Inhalte in Bericht und Konsole
# ---------------------------------------------------------------------------------------
def test_bericht_und_konsole_enthalten_keine_inhalte(tmp_path, monkeypatch, capsys):
    korpus = korpus_dir()
    baue_paar(korpus, "synth_5000", MIT_NEUER_SECTION)
    baue_paar(korpus, "synth_5001")
    # ein Paar mit Fehlern, damit auch die Fehlerpfade geprueft werden
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"C2": "Quokka Pinguin Flamingo anders"}))
    assert messen.main(["--korpus", "entwicklung", "--mit-vorschlaegen"]) == 1     # stille Fehler -> Ziel verfehlt (Exit-Code 1)
    konsole = capsys.readouterr().out.lower()
    ordner = next(config.pfad("messung_dir").glob("001_entwicklung"))
    texte = [konsole, (ordner / "bericht.json").read_text(encoding="utf-8").lower(),
             (ordner / "bericht.md").read_text(encoding="utf-8").lower()]
    for text in texte:
        for wort in MARKANT:
            assert wort not in text, wort
    assert "stille" in konsole                                   # Hinweis bei stillen Fehlern


def test_bericht_nur_erlaubte_felder_in_fehlern(tmp_path):
    korpus = korpus_dir()
    baue_paar(korpus, "synth_5002")
    messen.main(["--korpus", "entwicklung"])
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    for p in b["paare"]:
        for f in p["fehler"]:
            assert set(f) <= {"blatt", "zeile", "spalte", "aenderung", "art", "stumm", "grund"}
            assert f["art"] in messen.FEHLERARTEN


# ---------------------------------------------------------------------------------------
# (4) Append-only
# ---------------------------------------------------------------------------------------
def test_bericht_ist_append_only(tmp_path):
    baue_paar(korpus_dir(), "synth_6000")
    assert messen.main(["--korpus", "entwicklung"]) == 0
    basis = config.pfad("messung_dir")
    erste = (basis / "001_entwicklung" / "bericht.json").read_bytes()
    assert messen.main(["--korpus", "entwicklung"]) == 0
    assert (basis / "002_entwicklung" / "bericht.json").exists() and (basis / "002_entwicklung" / "bericht.md").exists()
    assert (basis / "001_entwicklung" / "bericht.json").read_bytes() == erste


def test_nummer_zaehlt_ueber_korpora_und_weicht_belegten_aus(tmp_path):
    baue_paar(korpus_dir("test"), "synth_6001")
    basis = config.pfad("messung_dir")
    (basis / "007_entwicklung").mkdir(parents=True)
    assert messen.main(["--korpus", "test"]) == 0
    assert (basis / "008_test" / "bericht.md").exists()


def test_schreibe_bericht_ueberschreibt_nie(tmp_path):
    basis = config.pfad("messung_dir")
    (basis / "001_entwicklung").mkdir(parents=True)
    (basis / "001_entwicklung" / "bericht.json").write_text("alt", encoding="utf-8")
    ziel = messen.schreibe_bericht(messen.fasse_zusammen([], "entwicklung", "", False), "entwicklung")
    assert ziel.name == "002_entwicklung"
    assert (basis / "001_entwicklung" / "bericht.json").read_text(encoding="utf-8") == "alt"


# ---------------------------------------------------------------------------------------
# (5) Gleicher Weg wie ueber runs
# ---------------------------------------------------------------------------------------
@pytest.mark.parametrize("elemente", [SAUBER, MIT_NEUER_SECTION])
def test_weg_ueber_runs_ist_identisch(tmp_path, elemente):
    p = baue_paar(tmp_path / "paare", "synth_7000", elemente)
    kopie = tmp_path / "messen_ergebnis.xlsx"
    res = messen.messe_paar(p, ergebnis_kopie=kopie)

    run_id = runs.neuer_lauf(p / "aenderungen.pdf", p / "vorher.xlsx")
    daten = runs.analysiere(run_id)
    ergebnis = runs.uebernehmen_lauf(run_id, [])
    runs_excel = runs.lauf_dir(run_id) / ergebnis["ordner"] / ergebnis["excel"]

    assert messen._lese_excel(kopie) == messen._lese_excel(runs_excel)
    assert res["bilanz"] == {k: v for k, v in ergebnis["bilanz"].items() if isinstance(v, (int, bool))}
    assert res["aenderungen"] == daten["bilanz"]["aenderungen"]
    assert [res["aenderungen_nach_status"][s] for s in ("uebernommen", "pruefen", "nicht_uebernommen")] == \
        [ergebnis["bilanz"]["uebernommen"], ergebnis["bilanz"]["pruefen"], ergebnis["bilanz"]["nicht_uebernommen"]]


# ---------------------------------------------------------------------------------------
# Robustheit / Eingabe
# ---------------------------------------------------------------------------------------
def test_kaputte_eingabe_exit_2_ohne_rohtext(tmp_path, capsys):
    assert messen.main(["--korpus", "entwicklung"]) == 2          # Ordner fehlt
    korpus = korpus_dir()
    assert messen.main(["--korpus", "entwicklung"]) == 2          # Ordner leer
    out = capsys.readouterr().out
    assert "Traceback" not in out and "Error" not in out
    baue_paar(korpus, "synth_8000")
    assert messen.main(["--korpus", "entwicklung", "--paar", "gibt_es_nicht"]) == 2
    assert messen.main(["--korpus", "entwicklung", "--paar", "../synth_8000"]) == 2
    assert messen.main(["--ordner", str(tmp_path / "nirgends")]) == 2
    assert messen.main(["--unbekannt"]) == 2
    assert not config.pfad("messung_dir").exists() or not list(config.pfad("messung_dir").iterdir())


def test_kaputtes_paar_bricht_messung_nicht_ab(tmp_path, capsys):
    korpus = korpus_dir()
    baue_paar(korpus, "synth_8100")
    (korpus / "synth_8101").mkdir()                                # leerer Ordner: Dateien fehlen
    kaputt = baue_paar(korpus, "synth_8102")
    (kaputt / "vorher.xlsx").write_bytes(b"PK das ist keine Excel")
    kaputt2 = baue_paar(korpus, "synth_8103")
    (kaputt2 / "aenderungen.pdf").write_bytes(b"%PDF-1.4 kein echtes PDF")
    assert messen.main(["--korpus", "entwicklung"]) == 0
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    assert b["gesamt"]["paare"] == 4 and b["gesamt"]["paare_abgestuerzt"] == 3
    assert b["fehlerarten"]["unklar"] == 3
    grund = {p["name"]: p["fehler"][0].get("grund") for p in b["paare"] if p["fehler"]}
    assert grund["synth_8101"] == "FileNotFoundError"
    assert all(re.fullmatch(r"\w+", g) for g in grund.values())      # nur Klassennamen
    assert b["paare"][0]["fehler"] == [] and b["paare"][0]["soll_zellen"] == 3


def test_paar_ordner_und_korpus_hinweise(tmp_path, capsys):
    eigener = tmp_path / "irgendwo"
    baue_paar(eigener, "echt_0001")
    baue_paar(eigener, "synth_9")
    (eigener / "KORPUS.json").write_text(json.dumps({"generator_version": "7"}), encoding="utf-8")
    assert messen.main(["--korpus", "test", "--ordner", str(eigener), "--paar", "echt_0001"]) == 0
    out = capsys.readouterr().out
    assert "Regel 11" in out
    b = json.loads(next(config.pfad("messung_dir").glob("001_test/bericht.json")).read_text(encoding="utf-8"))
    assert [p["name"] for p in b["paare"]] == ["echt_0001"] and b["generator_version"] == "7"
    assert b["korpus"] == "test" and b["schema_version"] == SCHEMA_VERSION and b["format"] == 1
    assert {"paare", "soll_zellen", "bilanz_stimmig", "vollstaendig", "stille_fehler", "auto_richtig", "auto_gesamt",
            "abdeckung_ohne_nachfrage", "lese_luecke_woerter", "idempotent", "kollateral"} <= set(b["gesamt"])
    assert set(b["fehlerarten"]) == set(messen.FEHLERARTEN)


def test_ampel_und_kriterien_in_der_konsole(tmp_path, capsys):
    baue_paar(korpus_dir(), "synth_9100")
    assert cli.main(["messen", "--korpus", "entwicklung"]) == 0     # ueber die Weiche in cli.main
    out = capsys.readouterr().out
    assert "ok" in out and "Stille Fehler" in out and "Abdeckung ohne Nachfrage" in out
    assert "--mit-vorschlaegen" in out                              # Hinweis: Vorschlaege nicht gemessen
    assert "NICHT ok" not in out


# ---------------------------------------------------------------------------------------
# Ausrichtung (Soll/Ist) ohne Excel-Dateien
# ---------------------------------------------------------------------------------------
def _zeilen(*zeilen):
    return messen._zeilen([list(z) for z in zeilen])


def test_diff_zelle_zeile_neu_zeile_weg():
    v = _zeilen(("1", "Eins", "alpha beta"), ("2", "Zwei", "gamma delta epsilon"), ("3", "Drei", "zeta eta theta"))
    n = _zeilen(("1", "Eins", "alpha beta gamma"), ("2", "Zwei", "gamma delta epsilon"), ("2.5", "Neu", "iota kappa"),
                ("3", "Drei", "zeta eta theta"))
    d = messen._diff(v, n)
    assert d.zellen == {(1, 3): "alpha beta gamma"} and d.weg == set()
    assert d.neue == [(2, ("2.5", "Neu", "iota kappa"), ("2.5", "Neu", "iota kappa"))]
    d2 = messen._diff(v, _zeilen(("1", "Eins", "alpha beta"), ("3", "Drei", "zeta eta theta")))
    assert d2.weg == {2} and set(d2.zellen) == {(2, 1), (2, 2), (2, 3)} and d2.neue == []


def test_diff_leere_zeilen_zaehlen_nicht_und_geleerte_zeile_wird_gepaart():
    v = _zeilen(("1", "Eins", "alpha beta", "Rest"), ("", "", ""), ("2", "Zwei", "gamma delta", "Rest"))
    n = _zeilen(("1", "Eins", "alpha beta", "Rest"), ("", "", "", "Rest"))
    d = messen._diff(v, n)
    assert set(d.zellen) == {(3, 1), (3, 2), (3, 3)} and d.weg == set()          # Zeile 3 geleert, Rest bleibt
    assert d.zellen[(3, 3)] == ""


def test_zellvergleich_ignoriert_nur_leerraum():
    assert messen._k("  a  b\n c ") == "a b c"
    assert messen._k("A") != messen._k("a")


# ---------------------------------------------------------------------------------------
# (6) Bestehende CLI-Befehle bleiben unveraendert
# ---------------------------------------------------------------------------------------
def test_cli_analyse_und_lauf_unveraendert(tmp_path, capsys):
    p = baue_paar(tmp_path, "paar")
    excel = tmp_path / "meine_cde.xlsx"
    shutil.copy2(p / "vorher.xlsx", excel)
    vorher = sha(excel)

    assert cli.main(["analyse", "--pdf", str(p / "aenderungen.pdf"), "--excel", str(excel)]) == 0
    out = capsys.readouterr().out
    assert "Lauf:" in out and "Änderungen: 3" in out and "Bilanz stimmt: ja" in out
    lauf1 = runs.liste()
    assert len(lauf1) == 1 and lauf1[0]["ergebnisse"] == 0

    assert cli.main(["lauf", "--pdf", str(p / "aenderungen.pdf"), "--excel", str(excel)]) == 0
    out = capsys.readouterr().out
    assert "Ergebnis:" in out and "übernommen: 3" in out
    neu = [l for l in runs.liste() if l["run_id"] != lauf1[0]["run_id"]][0]
    assert neu["ergebnisse"] == 1
    ergebnis = next((runs.lauf_dir(neu["run_id"]) / "ergebnis_001").glob("*_aktualisiert.xlsx"))
    ws = load_workbook(str(ergebnis))["CDE"]
    assert "Laub" in ws["C2"].value and "Gras" not in ws["C2"].value
    assert sha(excel) == vorher                                    # Original unberuehrt

    with pytest.raises(SystemExit):
        cli.main(["unbekannt"])
    assert cli.main(["analyse", "--pdf", str(p / "aenderungen.pdf"), "--excel", str(excel), "--alle-vorschlaege"]) == 0



# =======================================================================================
# Reparatur des Messwerkzeugs (Phantom-Fehler entfernt, versteckte Fehler sichtbar)
# Jeder Befund hat hier einen eigenen Selbsttest. Alles erfunden (Regel 16).
# =======================================================================================
from collections import Counter                      # noqa: E402

from openpyxl import Workbook                        # noqa: E402
from openpyxl.styles import Font, PatternFill        # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation   # noqa: E402

from tests.fixtures.pdf_writer import PdfDokument, scan_bild   # noqa: E402

ZWEI_AENDERUNGEN = [("h", "1 Eins"),
                    ("p", "Der Quokka frisst {-Gras-} {+Laub+} am Morgen und {-ruht-} {+schlaeft+} danach im Schatten."),
                    ("h", "2 Zwei"), ("p", "Ein Flamingo steht im Teich.")]


def zweite_aenderung_offen(monkeypatch):
    original = messen.plane

    def offen(pdf, xl, cfg=None):
        plan = original(pdf, xl, cfg)
        plan.aenderungen[-1].status, plan.aenderungen[-1].konfidenz = PRUEFEN, "mittel"
        return plan
    monkeypatch.setattr(messen, "plane", offen)


# ---- strenger Zellvergleich (RT-135) ---------------------------------------------------------
def test_strenger_zellvergleich_sieht_leerraum_und_zeichenformen():
    assert messen._s("a  b") != messen._s("a b")
    assert messen._s("a b") != messen._s("a b")
    assert messen._s("Mess­protokoll") != messen._s("Messprotokoll")
    assert messen._s("ﬁnden") != messen._s("finden")
    assert messen._s("a\nb") != messen._s("a b")
    assert messen._s("a\r\nb") == messen._s("a\nb")
    assert messen._s("   ") == "" and messen._s("​") == ""                 # ohne sichtbaren Inhalt = leer


def test_nur_leerraum_unterschied_ist_stiller_fehler_und_als_textform_ausgewiesen(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1100")
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(
        ziel, {"C2": load_workbook(str(ziel))["CDE"]["C2"].value.replace("Laub am", "Laub  am")}))
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["stille_fehler_nur_textform"] == 1
    assert res["erreicht_auto"] == 2 and res["auto_richtig"] == 2 and res["auto_gesamt"] == 3


def test_geschuetztes_leerzeichen_nur_im_soll_ist_nicht_darstellbar_im_ergebnis_ist_es_ein_fehler(tmp_path, monkeypatch):
    """Das PDF traegt kein geschuetztes Leerzeichen: steht es nur im Soll (nachher), kann kein Tool es kennen -> ausgewiesen,
    kein Fehler. Steht es im ERGEBNIS und nicht im Soll, ist es ein Fehler."""
    p = baue_paar(tmp_path, "synth_1101")
    zellen_aendern(p / "nachher.xlsx", {"C2": load_workbook(str(p / "nachher.xlsx"))["CDE"]["C2"].value.replace("am Morgen", "am Morgen")})
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 0 and res["erreicht_auto"] == res["soll_zellen"]
    assert res["soll_zellen_pdf_blind"] == 1 and res["soll_zellen_pdf_blind_erreicht"] == 1
    q = baue_paar(tmp_path, "synth_1102")
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, quelle: zellen_aendern(
        ziel, {"C2": load_workbook(str(ziel))["CDE"]["C2"].value.replace("am Morgen", "am Morgen")}))
    res2 = messen.messe_paar(q)
    assert res2["stille_fehler"] == 1 and res2["soll_zellen_pdf_blind"] == 0


# ---- je Aenderung bewerten (MA-1) -------------------------------------------------------------
def test_offene_aenderung_in_derselben_zelle_ist_kein_fehler_und_wird_je_aenderung_gezaehlt(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1110", ZWEI_AENDERUNGEN)
    zweite_aenderung_offen(monkeypatch)
    res = messen.messe_paar(p, mit_vorschlaegen=True)
    assert res["stille_fehler"] == 0 and res["fehler"] == []
    assert res["zellen_warten_auf_freigabe"] == 1 and res["pruefen_zellen"] == 1 and res["erreicht_auto"] == 0
    assert res["aenderungen_uebernommen_korrekt"] == 1 and res["aenderungen_uebernommen_falsch"] == 0
    assert res["aenderungen_zu_pruefen"] == 1 and res["aenderungen_nicht_uebernommen"] == 0
    assert res["vorschlaege_gesamt"] == res["vorschlaege_richtig"] == 1


def test_falsche_angewendete_aenderung_neben_offener_wird_je_aenderung_gefunden(tmp_path, monkeypatch):
    """Die angewendete Aenderung ist falsch (anderes Wort), die offene fehlt noch: Fehler, und zwar an der angewendeten."""
    p = baue_paar(tmp_path, "synth_1111", ZWEI_AENDERUNGEN)
    zweite_aenderung_offen(monkeypatch)
    original = messen.uebernehmen

    def falsch(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aend, quelle, ziel, freigaben, cfg)
        if Path(ziel).name == "ergebnis_auto.xlsx":
            zellen_aendern(Path(ziel), {"C2": load_workbook(str(ziel))["CDE"]["C2"].value.replace("Laub", "Kraut")})
        return res
    monkeypatch.setattr(messen, "uebernehmen", falsch)
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["zellen_warten_auf_freigabe"] == 0
    assert res["aenderungen_uebernommen_falsch"] == 1 and res["aenderungen_uebernommen_korrekt"] == 0
    assert res["aenderungen_zu_pruefen"] == 1


def test_unerklaert_nennt_nur_woerter_ausserhalb_der_offenen_aenderung():
    erlaubt = Counter(["jaehrlich", "monatlich"])
    ist, soll = "Die Wartung erfolgt jaehrlich durch uns.", "Die Wartung erfolgt monatlich durch uns."
    assert not messen._Auswertung._unerklaert(ist, soll, erlaubt)
    assert messen._Auswertung._unerklaert(ist, soll.replace("uns", "alle"), erlaubt)
    assert messen._Auswertung._unerklaert(ist, soll.replace("erfolgt ", "erfolgt  "), erlaubt)       # zusaetzlicher Leerraum
    nur_form = messen._Auswertung._unerklaert("Der Rest.\nEnde", "Der Rest. Ende", Counter())
    assert "\x00format" in nur_form and "rest." not in nur_form and len(nur_form) >= 2          # Nachbarwoerter fuer die Zuordnung


# ---- Zeilenpaarung / Anker (MA-2, MA-3) --------------------------------------------------------
def test_teilweise_geleerte_zeile_wird_wiedergefunden():
    v = _zeilen(("1", "Eins", "Absatz A.", "Status offen"),
                ("2", "Zwei", "Satz eins ist hier. Satz zwei ist hier. Satz drei ist hier.\nAbsatz zwei ist hier.", "Status offen"),
                ("3", "Drei", "Absatz C.", "Status offen"))
    x = _zeilen(("1", "Eins", "Absatz A.", "Status offen"),
                ("", "", "Satz eins ist hier. Satz zwei ist hier. Satz drei ist hier.", "Status offen"),
                ("3", "Drei", "Absatz C.", "Status offen"))
    d = messen._diff(v, x)
    assert d.neue == [] and d.weg == set() and d.zuordnung == {1: 1, 2: 2, 3: 3}
    assert set(d.zellen) == {(2, 1), (2, 2), (2, 3)}


def test_anker_neuer_zeilen_ignoriert_geleerte_und_entfernte_zeilen():
    leer = lambda z: not any(z.keys[:3])                                   # noqa: E731  (Bemerkung = 4. Spalte unverwaltet)
    v = _zeilen(("1", "Eins", "A", "Rest"), ("2", "Zwei", "B", "Rest"), ("3", "Drei", "C", "Rest"))
    soll = messen._diff(v, _zeilen(("1", "Eins", "A", "Rest"), ("7", "Neu", "ganz anderer Text hier", "Rest"),
                                   ("3", "Drei", "C", "Rest")), leer)
    x_ist = _zeilen(("1", "Eins", "A", "Rest"), ("", "", "", "Rest"), ("7", "Neu", "ganz anderer Text hier", "Rest"),
                    ("3", "Drei", "C", "Rest"))
    # a) Das Layout ist bekannt: die geleerte Zeile 2 gilt auf beiden Seiten als entfernt
    ist = messen._diff(v, x_ist, leer)
    assert ist.geleert == {2} and soll.weg == {2}
    tot = soll.weg | soll.geleert | ist.weg | ist.geleert
    messen._anker_neu(soll, v, tot)
    messen._anker_neu(ist, v, tot)
    assert [a for a, _, _ in soll.neue] == [a for a, _, _ in ist.neue] == [1]
    assert messen._neu_abgleich(soll.neue, ist.neue) == ({0: 0}, {0: True})
    # b) Das Layout ist unklar (keine Leer-Erkennung): die Zeile, die das SOLL entfernt, zaehlt trotzdem nicht als Anker
    soll2 = messen._diff(v, _zeilen(("1", "Eins", "A", "Rest"), ("7", "Neu", "ganz anderer Text hier", "Rest"),
                                    ("3", "Drei", "C", "Rest")))
    ist2 = messen._diff(v, x_ist)
    assert [a for a, _, _ in ist2.neue] == [2] and [a for a, _, _ in soll2.neue] == [1]        # ohne Abgleich: verschieden
    tot2 = soll2.weg | ist2.weg
    messen._anker_neu(soll2, v, tot2)
    messen._anker_neu(ist2, v, tot2)
    assert [a for a, _, _ in ist2.neue] == [a for a, _, _ in soll2.neue] == [1]


# ---- Regel 27 (RT-137/138) ---------------------------------------------------------------------
def test_regel27_pruefen_ohne_freigabe_angewendet_ist_stiller_fehler(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1120", ZWEI_AENDERUNGEN)
    zweite_aenderung_offen(monkeypatch)
    original = messen.uebernehmen

    def mit_heimlicher_freigabe(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        if Path(ziel).name == "ergebnis_auto.xlsx":
            freigaben = {a.id for a in aend if a.status == PRUEFEN}
        return original(aend, quelle, ziel, freigaben, cfg)
    monkeypatch.setattr(messen, "uebernehmen", mit_heimlicher_freigabe)
    res = messen.messe_paar(p)
    assert res["regel27_verstoesse"] == 1 and res["stille_fehler"] == 1
    assert res["stille_fehler_nach_art"].get("regel27_verstoss") == 1
    assert res["erreicht_auto"] == 0 and res["auto_richtig"] == 0                          # der Treffer zaehlt nicht
    assert res["aenderungen_uebernommen_falsch"] == 1 and res["aenderungen_uebernommen_korrekt"] == 1
    bericht = messen.fasse_zusammen([dict(res, name="synth_1120")], "entwicklung", "", False)
    assert [k["ampel"] for k in bericht["kriterien"] if k["name"].startswith("Regel 27")] == ["NICHT ok"]


def test_regel27_wird_auch_an_der_wirkung_erkannt_wenn_der_status_nichts_verraet(tmp_path, monkeypatch):
    """Das Ergebnis enthaelt den Text der offenen Aenderung, obwohl ihr Status 'pruefen' geblieben ist."""
    p = baue_paar(tmp_path, "synth_1121", ZWEI_AENDERUNGEN)
    zweite_aenderung_offen(monkeypatch)
    nachher_text = load_workbook(str(p / "nachher.xlsx"))["CDE"]["C2"].value
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"C2": nachher_text}))
    res = messen.messe_paar(p)
    assert res["regel27_verstoesse"] == 1 and res["stille_fehler"] >= 1


def test_ohne_verstoss_keine_regel27_meldung(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_1122"))
    assert res["regel27_verstoesse"] == 0 and res["stille_fehler_nach_art"].get("regel27_verstoss", 0) == 0


# ---- Struktur (RT-136), direkt mit gebauten Arbeitsmappen ---------------------------------------------
def _mappe(pfad: Path, zeilen: list[list[str]], hidden=(), hoehe=None, merges=(), dv=None, fett=(), fmt=None, fuellung=(), extra=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "CDE"
    for z in zeilen:
        ws.append(z)
    for r in hidden:
        ws.row_dimensions[r].hidden = True
    if hoehe:
        ws.row_dimensions[hoehe[0]].height = hoehe[1]
    for m in merges:
        ws.merge_cells(m)
    if dv:
        v = DataValidation(type="list", formula1='"a,b"')
        v.add(dv)
        ws.add_data_validation(v)
    for adr in fett:
        ws[adr].font = Font(bold=True)
    for adr in fuellung:
        ws[adr].fill = PatternFill("solid", fgColor="FFFF00")
    if fmt:
        ws[fmt[0]].number_format = fmt[1]
    if extra:
        extra(ws)
    wb.save(str(pfad))
    return pfad


def _befund(tmp_path: Path, vz, ez, **kw) -> dict:
    ev = {k[2:]: v for k, v in kw.items() if k.startswith("v_")}
    ee = {k[2:]: v for k, v in kw.items() if k.startswith("e_")}
    v = _mappe(tmp_path / "v.xlsx", vz, **ev)
    e = _mappe(tmp_path / "e.xlsx", ez, **ee)
    diff = messen._diff(messen._zeilen(messen._lese_excel(v)[0]), messen._zeilen(messen._lese_excel(e)[0]))
    return messen._struktur_befund(v, e, [diff])


KOPF4 = ["Nr", "Titel", "Inhalt", "Bem"]
VZ = [KOPF4] + [[str(i), f"Titel {i}", f"Text zu Eintrag {i} steht hier", "x"] for i in range(1, 8)]
EINGEFUEGT = VZ[:5] + [["4.1", "Neu A", "Ein ganz neuer Absatz steht jetzt hier", "x"],
                       ["4.2", "Neu B", "Noch ein ganz neuer Absatz steht hier", "x"]] + VZ[5:]


def test_struktur_einfuegen_verschiebt_hoehe_verbund_versteckt_und_validierung_legitim(tmp_path):
    b = _befund(tmp_path, VZ, EINGEFUEGT, v_hidden=[6], v_hoehe=(7, 40), v_merges=["D6:D7"], v_dv="D6:D7",
                e_hidden=[8], e_hoehe=(9, 40), e_merges=["D8:D9"], e_dv="D8:D9")
    assert b["struktur_geaendert"] == 0, b


def test_struktur_nicht_mitgewanderte_dinge_sind_befunde(tmp_path):
    b = _befund(tmp_path, VZ, EINGEFUEGT, v_hidden=[6], v_merges=["D6:D7"], v_dv="D6:D7",
                e_hidden=[6], e_merges=["D6:D7"], e_dv="D6:D7")
    assert b["struktur_nach_art"]["zeile_versteckt"] >= 1 and b["struktur_nach_art"]["verbunden"] >= 1
    assert b["struktur_nach_art"]["validierung"] >= 1


def test_struktur_neue_zeile_darf_nicht_ausgeblendet_sein(tmp_path):
    b = _befund(tmp_path, VZ, EINGEFUEGT, e_hidden=[6])
    assert b["struktur_nach_art"]["zeile_versteckt"] == 1


def test_struktur_zahlenformat_und_schrift_unveraenderter_zellen(tmp_path):
    b = _befund(tmp_path, VZ, VZ, e_fmt=("A3", "0.00"), e_fett=["C4"])
    assert b["struktur_nach_art"]["zahlenformat"] == 1 and b["struktur_nach_art"]["schrift"] == 1


def test_struktur_geaenderte_zelle_darf_markiert_werden_aber_nicht_fett(tmp_path):
    nachher = [z[:] for z in VZ]
    nachher[2][2] = "Text zu Eintrag 2 wurde ganz neu gefasst"
    b = _befund(tmp_path, VZ, nachher, e_fuellung=["C3"])                      # gelbe Markierung der Aenderung: erlaubt
    assert b["struktur_geaendert"] == 0
    b2 = _befund(tmp_path, VZ, nachher, e_fuellung=["C3"], e_fett=["C3"])      # aber nicht fett
    assert b2["struktur_nach_art"]["schrift"] == 1


def test_struktur_blatt_umbenannt_und_ausgeblendet(tmp_path):
    v = _mappe(tmp_path / "v.xlsx", VZ)
    wb = load_workbook(str(v))
    wb.create_sheet("Zwei")
    wb.save(str(v))
    wb["CDE"].title = "Anders"
    wb["Zwei"].sheet_state = "hidden"
    e = tmp_path / "e.xlsx"
    wb.save(str(e))
    diff = messen._diff(messen._zeilen(messen._lese_excel(v)[0]), messen._zeilen(messen._lese_excel(e)[0]))
    b = messen._struktur_befund(v, e, [diff, messen.Diff()])
    assert b["struktur_nach_art"]["blatt"] == 1 and b["struktur_nach_art"]["sichtbarkeit"] == 1


def test_sauberes_paar_hat_keine_strukturbefunde(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_1130", MIT_NEUER_SECTION))
    assert res["struktur_geaendert"] == 0 and set(res["struktur_nach_art"]) == set(messen.STRUKTUR_ARTEN)


# ---- Absturz im Wiederholungslauf (RT-139/140) --------------------------------------------------
def test_absturz_im_wiederholungslauf_behaelt_die_funde_des_ersten_laufs(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_1140")
    original = messen.uebernehmen
    zaehler = {"n": 0}

    def falsch_dann_absturz(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aend, quelle, ziel, freigaben, cfg)
        zaehler["n"] += 1
        if zaehler["n"] == 1:
            zellen_aendern(Path(ziel), {"C2": "Etwas ganz anderes steht hier."})
        if zaehler["n"] == 2:
            raise ValueError("GEHEIMER-DOKUMENTTEXT-quokka")
        return res
    monkeypatch.setattr(messen, "uebernehmen", falsch_dann_absturz)
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["idempotent"] is None
    assert res["abgestuerzt"] == "ValueError" and res["teilweise_gemessen"] is True
    assert res["ausgefallene_messteile"] == [{"teil": "idempotenz", "grund": "ValueError"}]
    assert res["fehler"][-1]["art"] == "unklar" and res["fehler"][-1]["grund"] == "ValueError"
    assert "GEHEIMER" not in json.dumps(res)


def test_abgestuerztes_paar_setzt_keine_ampel_auf_ok():
    sauber = {"name": "a", "stille_fehler": 0, "soll_zellen": 3, "erreicht_auto": 3, "auto_gesamt": 3, "auto_richtig": 3,
              "bilanz_stimmt": True, "vollstaendig": True, "idempotent": True, "original_unveraendert": True,
              "kollateral": 0, "fehler": [], "luecken_nicht_gemeldet": 0, "luecken_pruefbar": True, "soll_woerter": 10,
              "lese_luecke_woerter": 0, "struktur_geaendert": 0, "plausibilitaet_hinweise": [], "plausibilitaet_ungemeldet": 0}
    kaputt = messen._fehlgeschlagen("b", RuntimeError("x"))
    b = messen.fasse_zusammen([sauber, kaputt], "entwicklung", "", False)
    ampeln = {k["name"]: k["ampel"] for k in b["kriterien"]}
    for name in ("Stille Fehler", "Kollateralschaden", "Bilanz stimmig"):
        assert ampeln[name] == "nicht vollständig gemessen", name
    ganz_ok = messen.fasse_zusammen([sauber], "entwicklung", "", False)
    assert {k["name"]: k["ampel"] for k in ganz_ok["kriterien"]}["Stille Fehler"] == "ok"


# ---- keine erwartet.json: Luecken nicht pruefbar, Plausibilitaet direkt am PDF (RT-141) -----------------
def test_ohne_erwartet_json_ist_kriterium_1_nicht_gemessen(tmp_path, capsys):
    baue_paar(korpus_dir(), "echt_9200")
    assert messen.main(["--korpus", "entwicklung"]) == 0
    out = capsys.readouterr().out
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    k1 = next(k for k in b["kriterien"] if k["name"] == "Bilanz stimmig")
    assert k1["ampel"] == "nicht gemessen" and b["gesamt"]["luecken_nicht_pruefbar_paare"] == 1
    assert "erwartet.json" in out and "nicht prüfbar" in out
    md = next(config.pfad("messung_dir").glob("001_entwicklung/bericht.md")).read_text(encoding="utf-8")
    assert "erwartet.json" in md


def test_plausibilitaet_findet_scan_seite_aber_nicht_normale_seiten(tmp_path):
    d = PdfDokument(variante="farbe")
    d.ueberschrift("1 Eins")
    d.block("Ein ganz normaler Absatz mit {+neuem Text+} hier.")
    d.scan_seite(scan_bild([("h", "2 Scan"), ("p", "Hier steht {+Text+} nur als Pixel.")]))
    d.speichern(tmp_path / "a.pdf")
    hinweise, zeichen = messen._pdf_pruefung(tmp_path / "a.pdf")
    assert "keine_textzeichen" in {h["art"] for h in hinweise} and all(h["seite"] >= 2 for h in hinweise)
    assert "a" in zeichen and " " not in zeichen
    d2 = PdfDokument(variante="farbe")
    d2.ueberschrift("1 Eins")
    d2.block("Nur Text ohne Bild.")
    d2.speichern(tmp_path / "b.pdf")
    assert messen._pdf_pruefung(tmp_path / "b.pdf")[0] == []
    assert messen._pdf_pruefung(tmp_path / "nicht_da.pdf") == (None, None)


def test_plausibilitaet_ist_nur_ein_hinweis_nie_ein_fehler(tmp_path):
    ordner = tmp_path / "paar"
    ordner.mkdir()
    d = PdfDokument(variante="farbe")
    d.ueberschrift("1 Eins")
    d.block("Der Nachweis erfolgt durch ein Messprotokoll der Abnahme.")
    d.scan_seite(scan_bild([("h", "2 Scan"), ("p", "Hier steht {+Text+} nur als Pixel.")]))
    d.speichern(ordner / "aenderungen.pdf")
    elemente = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme.")]
    baue_excel(ordner / "vorher.xlsx", zustaende(elemente, False), "L1")
    baue_excel(ordner / "nachher.xlsx", zustaende(elemente, True), "L1")
    res = messen.messe_paar(ordner)
    scan = [h for h in res["plausibilitaet_hinweise"] if h["art"] == "keine_textzeichen"]
    assert scan and scan[0]["gemeldet"] == (res["luecken_gemeldet"] > 0)
    assert res["stille_fehler"] == 0 and res["luecken_nicht_gemeldet"] == 0


# ---- nicht verwaltete Soll-Zellen (RT-142) ----------------------------------------------------------
def test_klassen_der_soll_zellen_formel_haupt_harmlos_sonstige(tmp_path):
    p = baue_paar(tmp_path, "synth_1150")
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    ws["E2"] = "=LEN(C2)"
    ws["D3"] = "Neue Bemerkung mit mehr als drei Woertern hier"
    wb.save(str(p / "nachher.xlsx"))
    res = messen.messe_paar(p)
    assert res["soll_zellen_formel"] == 1 and res["soll_zellen_formel_verfehlt"] == 1
    assert res["soll_zellen_unverwaltet_sonstige"] == 1 and res["soll_zellen_unverwaltet_sonstige_verfehlt"] == 1
    assert res["soll_zellen"] == 3 and res["erreicht_auto"] == 3
    assert [(o["klasse"], o["zeile"], o["spalte"]) for o in res["nicht_verwaltet_verfehlt_orte"]] == \
        [("formel", 2, 5), ("sonstige", 3, 4)]
    assert res["lese_luecke_woerter"] == 0                                    # Formel-Woerter zaehlen nicht


# ---- Zaehleinheiten und Bericht -----------------------------------------------------------------------
def test_bericht_nennt_zaehleinheiten_und_neue_kennzahlen_ohne_inhalt(tmp_path, capsys):
    baue_paar(korpus_dir(), "synth_9300", ZWEI_AENDERUNGEN)
    assert messen.main(["--korpus", "entwicklung", "--mit-vorschlaegen"]) == 0
    ordner = next(config.pfad("messung_dir").glob("001_entwicklung"))
    b = json.loads((ordner / "bericht.json").read_text(encoding="utf-8"))
    assert all(k.get("einheit") for k in b["kriterien"])
    g = b["gesamt"]
    for feld in ("aenderungen_uebernommen_korrekt", "aenderungen_uebernommen_falsch", "aenderungen_zu_pruefen",
                 "aenderungen_nicht_uebernommen", "regel27_verstoesse", "struktur_geaendert", "soll_zellen_formel",
                 "soll_zellen_unverwaltet_sonstige", "plausibilitaet_ungemeldet", "stille_fehler_zellen",
                 "stille_fehler_paar_befunde", "luecken_nicht_pruefbar_paare", "zaehleinheiten",
                 "abdeckung_alle", "erreicht_alle", "struktur_geprueft", "struktur_nicht_geprueft", "kriterien_nicht_ok",
                 "struktur_vorschlaege_geaendert", "standard_paare_angenommen"):
        assert feld in g, feld
    md = (ordner / "bericht.md").read_text(encoding="utf-8")
    assert "| Einheit |" in md and "Einheit: Änderungen" in md
    text = (capsys.readouterr().out + md + json.dumps(b)).lower()
    for wort in MARKANT:
        assert wort not in text, wort


# =======================================================================================
# K-MESS2: blinde Flecken der Struktur-Pruefung und des Datentyps schliessen
# Jede Eigenschaft hat drei Faelle: korrekt mitverschoben (kein Befund), unerwartet nicht verschoben/veraendert (Befund), entfernt (Befund).
# =======================================================================================
import datetime                                        # noqa: E402

from openpyxl.comments import Comment                  # noqa: E402
from openpyxl.formatting.rule import CellIsRule        # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.table import Table            # noqa: E402


def _hoehe(zeile, wert):
    return lambda ws: setattr(ws.row_dimensions[zeile], "height", wert)


def _breite(buchstabe, wert):
    return lambda ws: setattr(ws.column_dimensions[buchstabe], "width", wert)


def _fix(adresse):
    return lambda ws: setattr(ws, "freeze_panes", adresse)


def _filter(bereich):
    return lambda ws: setattr(ws.auto_filter, "ref", bereich)


def _tabelle(bereich):
    return lambda ws: ws.add_table(Table(displayName="Tab1", ref=bereich))


def _name(text):
    return lambda ws: ws.parent.defined_names.add(DefinedName("Liste", attr_text=text))


def _link(adresse, url="https://example.org/a"):
    def f(ws):
        ws[adresse].hyperlink = url
    return f


def _komm(*adressen, text="Hinweis"):
    def f(ws):
        for adr in adressen:
            ws[adr].comment = Comment(text, "Autor")
    return f


def _bedingt(bereich):
    return lambda ws: ws.conditional_formatting.add(
        bereich, CellIsRule(operator="equal", formula=['"x"'], fill=PatternFill("solid", bgColor="FFFF00")))


def _alle(*funktionen):
    def f(ws):
        for fn in funktionen:
            fn(ws)
    return f


def _art(b, art):
    return b["struktur_nach_art"][art]


def test_struktur_zeilenhoehe_wandert_mit_zurueckbleiben_ist_befund(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_hoehe(7, 40), e_extra=_hoehe(9, 40))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_hoehe(7, 40), e_extra=_hoehe(7, 40)), "zeilenhoehe") >= 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_hoehe(7, 40)), "zeilenhoehe") == 1                 # Hoehe verloren
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_hoehe(3, 20), e_extra=_hoehe(3, 35)), "zeilenhoehe") == 1  # unveraenderte Zeile
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_hoehe(3, 35)), "zeilenhoehe") == 1                         # Hoehe dazugekommen


def test_struktur_zeilenhoehe_einer_geaenderten_zeile_ist_keine_aussage(tmp_path):
    nachher = [z[:] for z in VZ]
    nachher[2][2] = "Text zu Eintrag 2 wurde ganz neu gefasst"
    assert _befund(tmp_path, VZ, nachher, v_extra=_hoehe(3, 20), e_extra=_hoehe(3, 35))["struktur_geaendert"] == 0


def test_struktur_spaltenbreite(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_breite("C", 40), e_extra=_breite("C", 40))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_breite("C", 40), e_extra=_breite("C", 55)), "spaltenbreite") == 1
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_breite("C", 40)), "spaltenbreite") == 1                       # Breite verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_breite("B", 30)), "spaltenbreite") == 1                       # Breite dazugekommen


def test_struktur_fixierte_zeilen(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_fix("A2"), e_extra=_fix("A2"))["struktur_geaendert"] == 0
    # Fixierung unterhalb der Einfuegestelle: die fixierten Zeilen 1 bis 6 werden zu 1 bis 8, die erste bewegliche Zeile ist dann 9
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_fix("A7"), e_extra=_fix("A9"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_fix("A7"), e_extra=_fix("A7")), "fixierung") == 1
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_fix("A2")), "fixierung") == 1                               # Fixierung verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_fix("A2")), "fixierung") == 1                               # Fixierung dazugekommen
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_fix("A2"), e_extra=_fix("C2")), "fixierung") == 1          # Spalten dazu


def test_struktur_autofilter(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_filter("A1:D8"), e_extra=_filter("A1:D10"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_filter("A1:D8"), e_extra=_filter("A1:D8")), "autofilter") == 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_filter("A1:D8")), "autofilter") == 1                # Filter verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_filter("A1:D8")), "autofilter") == 1                        # Filter dazugekommen


def test_struktur_excel_tabellen(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_tabelle("A1:D8"), e_extra=_tabelle("A1:D10"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_tabelle("A1:D8"), e_extra=_tabelle("A1:D8")), "tabelle") >= 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_tabelle("A1:D8")), "tabelle") == 1                  # Tabelle verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_tabelle("A1:D8")), "tabelle") == 1                          # Tabelle dazugekommen


def test_struktur_benannte_bereiche(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_name("CDE!$C$2:$C$8"), e_extra=_name("CDE!$C$2:$C$10"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_name("CDE!$C$2:$C$8"), e_extra=_name("CDE!$C$2:$C$8")), "benannter_bereich") == 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_name("CDE!$C$2:$C$8")), "benannter_bereich") == 1  # Name verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_name("CDE!$C$2:$C$8")), "benannter_bereich") == 1          # Name dazugekommen


def test_struktur_hyperlinks(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_link("B7"), e_extra=_link("B9"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_link("B7"), e_extra=_link("B7")), "hyperlink") >= 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_link("B7")), "hyperlink") == 1                      # Link verloren
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_link("B3"), e_extra=_link("B3", "https://example.org/b")), "hyperlink") >= 1
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_link("B3")), "hyperlink") == 1                              # Link dazugekommen


def test_struktur_kommentare(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_komm("C7"), e_extra=_komm("C9"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_komm("C3")), "kommentar") == 1                              # Kommentar verloren
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=_komm("C3"), e_extra=_komm("C3", text="Anders")), "kommentar") == 1
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_komm("D4")), "kommentar") == 1                              # Kommentar dazugekommen
    # in einer NEUEN Zeile darf das Tool einen Kommentar anlegen (Markierungsart "kommentar")
    assert _befund(tmp_path, VZ, EINGEFUEGT, e_extra=_komm("C6"))["struktur_geaendert"] == 0
    # an einer veraenderten Zelle darf der Kommentar ergaenzt werden, verloren gehen darf er nicht
    nachher = [z[:] for z in VZ]
    nachher[2][2] = "Text zu Eintrag 2 wurde ganz neu gefasst"
    assert _befund(tmp_path, VZ, nachher, v_extra=_komm("C3"), e_extra=_komm("C3", text="Hinweis und mehr"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, nachher, v_extra=_komm("C3")), "kommentar") == 1


def test_struktur_bedingte_formatierung(tmp_path):
    assert _befund(tmp_path, VZ, EINGEFUEGT, v_extra=_bedingt("D2:D8"), e_extra=_bedingt("D2:D10"))["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_bedingt("D2:D8"), e_extra=_bedingt("D2:D8")), "bedingte_formatierung") >= 1
    assert _art(_befund(tmp_path, VZ, EINGEFUEGT, v_extra=_bedingt("D2:D8")), "bedingte_formatierung") == 1   # Regel verloren
    assert _art(_befund(tmp_path, VZ, VZ, e_extra=_bedingt("D2:D8")), "bedingte_formatierung") == 1           # Regel dazugekommen


def test_struktur_alle_neuen_eigenschaften_korrekt_mitverschoben_sind_kein_befund(tmp_path):
    """Gegenprobe: alle neuen Eigenschaften zusammen, korrekt mitverschoben, ergeben 0 (sonst waere jede Einfuegung ein Phantom-Fehler)."""
    vorher = _alle(_hoehe(7, 40), _breite("C", 40), _fix("A2"), _filter("A1:D8"), _tabelle("A1:D8"), _name("CDE!$C$2:$C$8"),
                   _link("B7"), _komm("C7"), _bedingt("D2:D8"))
    nachher = _alle(_hoehe(9, 40), _breite("C", 40), _fix("A2"), _filter("A1:D10"), _tabelle("A1:D10"), _name("CDE!$C$2:$C$10"),
                    _link("B9"), _komm("C9"), _bedingt("D2:D10"))
    b = _befund(tmp_path, VZ, EINGEFUEGT, v_extra=vorher, e_extra=nachher)
    assert b["struktur_geaendert"] == 0, b["struktur_nach_art"]
    # und derselbe Aufbau ohne Mitverschieben: jede Eigenschaft meldet sich
    b2 = _befund(tmp_path, VZ, EINGEFUEGT, v_extra=vorher, e_extra=vorher)
    for art in ("zeilenhoehe", "autofilter", "tabelle", "benannter_bereich", "hyperlink", "kommentar", "bedingte_formatierung"):
        assert _art(b2, art) >= 1, art


# ---- Datentyp ---------------------------------------------------------------------------------------------------
ZAHLZEILEN = [KOPF4] + [[i, f"Titel {i}", f"Text zu Eintrag {i} steht hier", "x"] for i in range(1, 8)]    # Nr als Zahl


def test_typ_von_unterscheidet_zahl_text_formel_datum():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"], ws["A4"], ws["A5"], ws["A6"] = 31, "31", "=1+1", datetime.date(2026, 10, 15), True, None
    assert [messen._typ_von(ws[f"A{i}"]) for i in range(1, 7)] == ["zahl", "text", "formel", "datum", "bool", "leer"]
    assert messen._zelltext(ws["A1"].value) == messen._zelltext(ws["A2"].value)          # der Text allein sieht keinen Unterschied


def test_struktur_datentyp_an_unveraendertem_text(tmp_path):
    assert _befund(tmp_path, ZAHLZEILEN, ZAHLZEILEN)["struktur_geaendert"] == 0
    b = _befund(tmp_path, ZAHLZEILEN, VZ)                                    # alle sieben Nr-Zellen: Zahl -> Text
    assert _art(b, "datentyp") == 7 and len(b["datentyp_orte"]) == 7
    assert messen._lese_excel(tmp_path / "v.xlsx") == messen._lese_excel(tmp_path / "e.xlsx")     # der Text ist identisch
    assert _art(_befund(tmp_path, VZ, ZAHLZEILEN), "datentyp") == 7          # und umgekehrt: Text -> Zahl


def test_struktur_datentyp_formel_und_datum_werden_zu_text(tmp_path):
    def formel(ws):
        ws["C3"].value = "=1+1"

    def formel_als_text(ws):
        ws["C3"].value = "=1+1"
        ws["C3"].data_type = "s"
    assert _befund(tmp_path, VZ, VZ, v_extra=formel, e_extra=formel)["struktur_geaendert"] == 0
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=formel, e_extra=formel_als_text), "datentyp") == 1
    assert _art(_befund(tmp_path, VZ, VZ, v_extra=lambda ws: setattr(ws["D2"], "value", datetime.date(2026, 10, 15)),
                        e_extra=lambda ws: setattr(ws["D2"], "value", "2026-10-15 00:00:00")), "datentyp") == 1


def test_zahl_wird_text_in_neuer_zeile_ist_stiller_fehler(tmp_path):
    """Soll: die neue Section hat die NUMMER 4 (Zahl). Das Tool schreibt "4" als Text. Der Text ist gleich, der Datentyp nicht."""
    p = baue_paar(tmp_path, "synth_6100", MIT_NEUER_SECTION)
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    ws["A5"].value = 4
    wb.save(str(p / "nachher.xlsx"))
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["stille_fehler_nach_art"].get("datentyp_gewechselt") == 1
    assert res["erreicht_auto"] == res["soll_zellen"] - 1 and res["auto_richtig"] == res["auto_gesamt"] - 1
    assert (res["fehler"][0]["blatt"], res["fehler"][0]["spalte"], res["fehler"][0]["stumm"]) == (2, 1, True)
    sauber = messen.messe_paar(baue_paar(tmp_path, "synth_6101", MIT_NEUER_SECTION))                  # Gegenprobe: Text bleibt Text
    assert sauber["stille_fehler"] == 0 and sauber["erreicht_auto"] == sauber["soll_zellen"]


def test_zahl_wird_text_an_unveraenderter_zelle_in_verwalteter_spalte_ist_stiller_fehler(monkeypatch, tmp_path):
    p = baue_paar(tmp_path, "synth_6102")
    for datei in ("vorher.xlsx", "nachher.xlsx"):
        zellen_aendern(p / datei, {"A4": 3})                                  # Section-Nummer der dritten Section als Zahl
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"A4": "3"}))
    res = messen.messe_paar(p)
    assert res["stille_fehler"] == 1 and res["stille_fehler_nach_art"].get("datentyp_gewechselt") == 1
    assert res["struktur_nach_art"]["datentyp"] == 0                          # verwaltete Spalte: stiller Fehler, kein Strukturbefund
    monkeypatch.undo()
    ohne = baue_paar(tmp_path, "synth_6103")
    for datei in ("vorher.xlsx", "nachher.xlsx"):
        zellen_aendern(ohne / datei, {"A4": 3})
    sauber = messen.messe_paar(ohne)                                           # Gegenprobe: das Tool laesst die Zahl in Ruhe
    assert sauber["stille_fehler"] == 0 and sauber["struktur_geaendert"] == 0


def test_zahl_wird_text_in_nicht_verwalteter_spalte_ist_strukturbefund(monkeypatch, tmp_path):
    p = baue_paar(tmp_path, "synth_6104")
    for datei in ("vorher.xlsx", "nachher.xlsx"):
        zellen_aendern(p / datei, {"D4": 5})                                  # Bemerkung als Zahl
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"D4": "5"}))
    res = messen.messe_paar(p)
    assert res["struktur_nach_art"]["datentyp"] == 1 and res["struktur_geaendert"] == 1
    assert res["stille_fehler"] == 0


# ---- Struktur der Excel mit allen freigegebenen Vorschlaegen ----------------------------------------------------
def _sabotiere_vorschlags_excel(monkeypatch, wirkung):
    original = messen.uebernehmen

    def mit(aend, quelle, ziel, freigaben=None, cfg=None, **_):
        res = original(aend, quelle, ziel, freigaben, cfg)
        if Path(ziel).name == "ergebnis_alle.xlsx":
            wb = load_workbook(str(ziel))
            wirkung(wb["CDE"])
            wb.save(str(ziel))
        return res
    monkeypatch.setattr(messen, "uebernehmen", mit)


def test_vorschlags_excel_wird_auf_struktur_geprueft_unerwartete_abweichung_ist_befund(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_6200")
    _sabotiere_vorschlags_excel(monkeypatch, lambda ws: setattr(ws, "freeze_panes", "A5"))
    res = messen.messe_paar(p, mit_vorschlaegen=True)
    assert res["struktur_vorschlaege_geaendert"] == 1 and res["struktur_vorschlaege_nach_art"]["fixierung"] == 1
    assert res["struktur_vorschlaege_angekuendigt"] == 0
    assert res["struktur_geaendert"] == 0                                      # die Excel der Automatik ist in Ordnung
    b = messen.fasse_zusammen([dict(res, name="synth_6200")], "entwicklung", "1", True)
    k = next(k for k in b["kriterien"] if k["name"].startswith("Vorschlags-Excel"))
    assert k["ampel"] == "NICHT ok" and b["gesamt"]["struktur_vorschlaege_geaendert"] == 1


def test_vorschlags_excel_abweichung_ist_zulaessig_wenn_der_plan_sie_vorab_gemeldet_hat(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_6201")
    original = messen.plane

    def mit_warnung(pdf, xl, cfg=None):
        plan = original(pdf, xl, cfg)
        plan.aenderungen[0].hinweise.append("Hier laufen verbundene Zellen durch die Zeile; sie würden zerrissen oder verrutschen.")
        return plan
    monkeypatch.setattr(messen, "plane", mit_warnung)
    _sabotiere_vorschlags_excel(monkeypatch, lambda ws: ws.merge_cells("G2:H3"))
    res = messen.messe_paar(p, mit_vorschlaegen=True)
    assert res["struktur_vorschlaege_geaendert"] == 0 and res["struktur_vorschlaege_angekuendigt"] == 1
    assert res["struktur_vorschlaege_angekuendigt_nach_art"]["verbunden"] == 1
    # dieselbe Abweichung OHNE Vorab-Warnung ist ein Befund
    monkeypatch.setattr(messen, "plane", original)
    ohne = messen.messe_paar(baue_paar(tmp_path, "synth_6202"), mit_vorschlaegen=True)
    assert ohne["struktur_vorschlaege_geaendert"] == 1 and ohne["struktur_vorschlaege_angekuendigt"] == 0


def test_ohne_vorschlaege_ist_die_vorschlags_struktur_nicht_gemessen(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_6203"))
    assert res["struktur_vorschlaege_geaendert"] is None
    b = messen.fasse_zusammen([dict(res, name="synth_6203")], "entwicklung", "1", False)
    assert next(k for k in b["kriterien"] if k["name"].startswith("Vorschlags-Excel"))["ampel"] == "nicht gemessen"


def test_sauberes_paar_mit_vorschlaegen_hat_keine_strukturbefunde(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_6204", MIT_NEUER_SECTION), mit_vorschlaegen=True)
    assert res["struktur_geaendert"] == 0 and res["struktur_vorschlaege_geaendert"] == 0


# ---- Ampel nennt, was sie prueft; Abdeckung ueber alle Soll-Zellen; 90-%-Regel ohne Profil -----------------------------
def _paar_dict(**kw) -> dict:
    d = {"name": "p", "stille_fehler": 0, "soll_zellen": 4, "erreicht_auto": 4, "auto_gesamt": 4, "auto_richtig": 4,
         "bilanz_stimmt": True, "vollstaendig": True, "idempotent": True, "original_unveraendert": True, "kollateral": 0,
         "fehler": [], "luecken_nicht_gemeldet": 0, "luecken_pruefbar": True, "soll_woerter": 10, "lese_luecke_woerter": 0,
         "struktur_geaendert": 0, "plausibilitaet_hinweise": [], "plausibilitaet_ungemeldet": 0}
    d.update(kw)
    return d


def _ampeln(paare, vorschlaege=False) -> dict:
    return {k["name"]: k["ampel"] for k in messen.fasse_zusammen(paare, "entwicklung", "1", vorschlaege)["kriterien"]}


def test_struktur_ampel_sagt_was_sie_prueft_und_was_nicht(tmp_path):
    res = messen.messe_paar(baue_paar(tmp_path, "synth_6300"))
    b = messen.fasse_zusammen([dict(res, name="synth_6300")], "entwicklung", "1", False)
    k = next(k for k in b["kriterien"] if k["name"].startswith("Struktur unverändert"))
    assert "geprüft:" in k["detail"] and "nicht geprüft:" in k["detail"]
    geprueft, nicht_geprueft = k["detail"].split("nicht geprüft:")
    for wort in ("bedingte Formatierung", "Hyperlink", "Kommentare", "benannte Bereiche", "Autofilter", "Tabellen", "fixierte",
                 "Spaltenbreiten", "Zeilenhöhen", "Datentyp"):
        assert wort in geprueft, wort
    for wort in ("Rahmen", "Ausrichtung", "Diagramme", "Druckbereich"):
        assert wort in nicht_geprueft, wort
    md = messen._md(b)
    assert "Struktur ok heißt nur" in md and "Nicht geprüft wurde" in md
    assert b["gesamt"]["struktur_geprueft"] and b["gesamt"]["struktur_nicht_geprueft"]


def test_abdeckung_ueber_alle_soll_zellen_ist_eine_eigene_ampel_mit_grenze_70_prozent():
    name = "Abdeckung über ALLE Soll-Zellen (inkl. Füllspalten, Formeln)"
    # Haupt-Spalten 100 %, aber nur 4 von 10 Soll-Zellen insgesamt (Fuellung und Formeln nicht erreicht): Haupt-Zeile ok, Gesamt-Zeile NICHT ok
    knapp = _paar_dict(soll_zellen_unverwaltet=4, soll_zellen_formel=2, soll_zellen_formel_verfehlt=2, erreicht_alle=4)
    a = _ampeln([knapp])
    assert a[name] == "NICHT ok" and a["Abdeckung ohne Nachfrage (nur Haupt-Spalten)"] == "ok"
    b = messen.fasse_zusammen([knapp], "entwicklung", "1", False)
    assert b["gesamt"]["soll_zellen_alle"] == 10 and b["gesamt"]["erreicht_alle"] == 4 and b["gesamt"]["abdeckung_alle"] == 0.4
    k = next(k for k in b["kriterien"] if k["name"] == name)
    assert k["einheit"] == "Zellen" and "Formelzellen" in k["detail"] and "0 von 2" in k["detail"]
    # 7 von 10 = 70 %: ok (Grenze gilt "mindestens")
    assert _ampeln([_paar_dict(soll_zellen_unverwaltet=6, erreicht_alle=7)])[name] == "ok"
    assert _ampeln([_paar_dict(soll_zellen_unverwaltet=6, erreicht_alle=6)])[name] == "NICHT ok"


def test_abdeckung_alle_misst_die_echte_zaehlung_im_paar(tmp_path):
    p = baue_paar(tmp_path, "synth_6301")
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    ws["E2"] = "=LEN(C2)"                                  # Formelzelle: das Tool fasst sie nicht an
    ws["D3"] = "Neue Bemerkung mit mehr als drei Woertern hier"      # sonstige nicht verwaltete Zelle
    wb.save(str(p / "nachher.xlsx"))
    res = messen.messe_paar(p)
    assert res["soll_zellen"] == 3 and res["soll_zellen_alle"] == 5 and res["erreicht_alle"] == 3 and res["erreicht_formel"] == 0
    assert messen.fasse_zusammen([dict(res, name="p")], "entwicklung", "1", False)["gesamt"]["abdeckung_alle"] == 0.6


def test_neunzig_prozent_regel_gilt_auch_ohne_profil():
    name = "Abdeckung ohne Nachfrage (nur Haupt-Spalten)"
    knapp = dict(soll_zellen=4, erreicht_auto=3)                       # 75 %: unter 90 %, ueber 70 %
    ohne_profil = _ampeln([_paar_dict(**knapp)])[name]
    standard = _ampeln([_paar_dict(profil="standard", **knapp)])[name]
    assert ohne_profil == standard == "NICHT ok"                       # frueher: ohne Profil "ok", mit Profil "NICHT ok"
    assert _ampeln([_paar_dict(profil="schwer", **knapp)])[name] == "ok"
    b = messen.fasse_zusammen([_paar_dict(**knapp), _paar_dict(profil="standard", soll_zellen=4, erreicht_auto=4)], "entwicklung", "1", False)
    assert b["gesamt"]["standard_paare"] == 2 and b["gesamt"]["standard_paare_angenommen"] == 1
    k = next(k for k in b["kriterien"] if k["name"] == name)
    assert "Standard-Aufbau angenommen" in k["detail"]


def test_paar_ohne_erwartet_json_wird_als_standard_angenommen_gemeldet(tmp_path, capsys):
    res = messen.messe_paar(baue_paar(tmp_path, "echt_6400"))
    assert res["profil_angenommen"] is True
    korpus = korpus_dir()
    baue_paar(korpus, "echt_6401")
    messen.main(["--korpus", "entwicklung"])
    assert "Standard-Aufbau angenommen" in capsys.readouterr().out


# ---- Exit-Code ------------------------------------------------------------------------------------------------------
def test_exit_code_0_ohne_verfehltes_kriterium_1_bei_nicht_ok_2_bei_kaputter_eingabe(tmp_path, monkeypatch, capsys):
    korpus = korpus_dir()
    baue_paar(korpus, "synth_6500")
    assert messen.main(["--korpus", "entwicklung"]) == 0
    assert "Exit-Code 0" in capsys.readouterr().out
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    assert b["gesamt"]["kriterien_nicht_ok"] == 0 and b["gesamt"]["ziel_verfehlt"] is False
    sabotiere_ergebnis(monkeypatch, lambda ziel, aend, q: zellen_aendern(ziel, {"C2": "Etwas ganz anderes steht hier."}))
    assert cli.main(["messen", "--korpus", "entwicklung"]) == 1                  # Messung gelungen, Ziel verfehlt
    out = capsys.readouterr().out
    assert "Exit-Code 1" in out and "NICHT ok" in out
    b2 = json.loads(next(config.pfad("messung_dir").glob("002_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    assert b2["gesamt"]["kriterien_nicht_ok"] >= 1 and b2["gesamt"]["ziel_verfehlt"] is True
    assert messen.main(["--korpus", "entwicklung", "--paar", "gibt_es_nicht"]) == 2   # kaputte Eingabe


def test_exit_code_1_auch_wenn_nur_die_gesamtabdeckung_verfehlt_wird(tmp_path):
    korpus = korpus_dir()
    p = baue_paar(korpus, "synth_6501")
    wb = load_workbook(str(p / "nachher.xlsx"))
    ws = wb["CDE"]
    ws["E2"] = "=LEN(C2)"                                                     # eine Formelzelle, die das Tool bewusst nicht anpasst
    ws["D3"] = "Neue Bemerkung mit mehr als drei Woertern hier"               # eine sonstige nicht verwaltete Zelle (3 von 5 = 60 %)
    wb.save(str(p / "nachher.xlsx"))
    assert messen.main(["--korpus", "entwicklung"]) == 1
    b = json.loads(next(config.pfad("messung_dir").glob("001_entwicklung/bericht.json")).read_text(encoding="utf-8"))
    ampeln = {k["name"]: k["ampel"] for k in b["kriterien"]}
    assert ampeln["Abdeckung über ALLE Soll-Zellen (inkl. Füllspalten, Formeln)"] == "NICHT ok"
    assert ampeln["Abdeckung ohne Nachfrage (nur Haupt-Spalten)"] == "ok" and ampeln["Stille Fehler"] == "ok"


# ---- Ende zu Ende: das echte Tool fuegt Zeilen ein, die neuen Eigenschaften duerfen keinen Phantom-Befund ergeben -----------------
ELEMENTE_EINFUEGEN = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
                      ("h", "2 Zwei"), ("p", "Die Reaktionszeit darf 150 ms nicht ueberschreiten."),
                      ("h", "{+2.5 Neu+}"), ("p", "{+Ein neuer Absatz mit Zusatz.+}"),
                      ("h", "3 Drei"), ("p", "Alle Steckverbindungen sind eindeutig beschriftet."),
                      ("h", "4 Vier"), ("p", "Die Lieferung umfasst alles Zubehoer.")]


def _paar_zum_einfuegen(tmp_path: Path, unten: bool) -> Path:
    ordner = tmp_path / "paar"
    ordner.mkdir()
    baue_pdf(ordner / "aenderungen.pdf", ELEMENTE_EINFUEGEN, "farbe")
    baue_excel(ordner / "nachher.xlsx", zustaende(ELEMENTE_EINFUEGEN, True), "L3", True)
    baue_excel(ordner / "vorher.xlsx", zustaende(ELEMENTE_EINFUEGEN, False), "L3", True)
    wb = load_workbook(str(ordner / "vorher.xlsx"))
    ws = wb["CDE"]
    ws.freeze_panes = "A2"
    ws["B2"].hyperlink = "https://example.org/oben"
    ws["C3"].comment = Comment("Hinweis oben", "Autor")
    ws.conditional_formatting.add("D2:D3", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", bgColor="FFFF00")))
    if unten:        # Dinge unterhalb der Einfuegestelle: die Automatik fasst sie nicht an (pruefen), die Vorschlags-Excel verschiebt sie nicht
        ws["B8"].hyperlink = "https://example.org/unten"
        ws["C8"].comment = Comment("Hinweis unten", "Autor")
        ws.conditional_formatting.add("D6:D9", CellIsRule(operator="equal", formula=['"B"'], fill=PatternFill("solid", bgColor="FFFF00")))
        ws.auto_filter.ref = "A1:D9"
        wb.defined_names.add(DefinedName("Liste", attr_text="CDE!$C$2:$C$9"))
    else:
        ws.row_dimensions[7].height = 40
    wb.save(str(ordner / "vorher.xlsx"))
    return ordner


def test_ende_zu_ende_zeilen_einfuegen_mit_eigenschaften_oberhalb_und_mitgezogenem_ist_kein_befund(tmp_path):
    res = messen.messe_paar(_paar_zum_einfuegen(tmp_path, unten=False), mit_vorschlaegen=True)
    assert res["auto_gesamt"] >= 2                                              # die Zeilen wurden wirklich eingefuegt
    assert res["struktur_geaendert"] == 0 and res["struktur_vorschlaege_geaendert"] == 0, res["struktur_nach_art"]
    assert res["stille_fehler"] == 0


def test_ende_zu_ende_nicht_mitgezogenes_wird_in_der_vorschlags_excel_als_angekuendigt_gezaehlt(tmp_path):
    res = messen.messe_paar(_paar_zum_einfuegen(tmp_path, unten=True), mit_vorschlaegen=True)
    assert res["pruefen_zellen"] >= 1 and res["struktur_geaendert"] == 0          # die Automatik hat nichts angefasst
    assert res["struktur_vorschlaege_geaendert"] == 0                             # der Plan hatte jede Abweichung vorab gemeldet
    assert {a for a, n in res["struktur_vorschlaege_angekuendigt_nach_art"].items() if n} == \
        {"bedingte_formatierung", "autofilter", "benannter_bereich"}


# ---- Ende zu Ende mit Sabotage am Ergebnis des Tools: jede neue Eigenschaft meldet sich im Paar-Ergebnis und in der Ampel ---------
def _paar_mit_eigenschaften_oben(tmp_path: Path) -> Path:
    ordner = _paar_zum_einfuegen(tmp_path, unten=False)
    wb = load_workbook(str(ordner / "vorher.xlsx"))
    ws = wb["CDE"]
    ws.auto_filter.ref = "A1:D3"                                       # beide oberhalb der Einfuegestelle
    wb.defined_names.add(DefinedName("Liste", attr_text="CDE!$C$2:$C$3"))
    wb.save(str(ordner / "vorher.xlsx"))
    return ordner


SABOTAGE_STRUKTUR_NEU = {
    "hyperlink": lambda ws: setattr(ws["B2"], "hyperlink", None),
    "kommentar": lambda ws: setattr(ws["C3"], "comment", None),
    "bedingte_formatierung": lambda ws: setattr(ws.conditional_formatting, "_cf_rules", {}),
    "autofilter": lambda ws: setattr(ws.auto_filter, "ref", None),
    "benannter_bereich": lambda ws: ws.parent.defined_names.pop("Liste"),
    "fixierung": lambda ws: setattr(ws, "freeze_panes", None),
    "spaltenbreite": lambda ws: setattr(ws.column_dimensions["C"], "width", 99),
    "zeilenhoehe": lambda ws: setattr(ws.row_dimensions[3], "height", 77),
}


def test_ende_zu_ende_ausgangslage_fuer_die_sabotage_ist_ohne_befund(tmp_path):
    res = messen.messe_paar(_paar_mit_eigenschaften_oben(tmp_path))
    assert res["auto_gesamt"] >= 2 and res["struktur_geaendert"] == 0, res["struktur_nach_art"]


@pytest.mark.parametrize("art", sorted(SABOTAGE_STRUKTUR_NEU))
def test_ende_zu_ende_sabotage_wird_als_strukturbefund_der_richtigen_art_gemeldet(tmp_path, monkeypatch, art):
    ordner = _paar_mit_eigenschaften_oben(tmp_path)

    def wirkung(ziel, aend, quelle):
        wb = load_workbook(str(ziel))
        SABOTAGE_STRUKTUR_NEU[art](wb["CDE"])
        wb.save(str(ziel))
    sabotiere_ergebnis(monkeypatch, wirkung)
    res = messen.messe_paar(ordner)
    assert res["struktur_nach_art"][art] >= 1, (art, res["struktur_nach_art"])
    b = messen.fasse_zusammen([dict(res, name="synth_6600")], "entwicklung", "1", False)
    assert next(k for k in b["kriterien"] if k["name"].startswith("Struktur unverändert"))["ampel"] == "NICHT ok"
    assert b["gesamt"]["kriterien_nicht_ok"] >= 1


# =======================================================================================
# Messung v2 (Welle 1): Soll-Zellen-Klassen, Restarbeit, Pruefcodes/Sperr-Quote, Vergleich, Korpora, je Erzeuger
# =======================================================================================
def test_klassen_v2_fuell_konstant_formel_nicht_ableitbar(tmp_path):
    """Neue Section in L1: Bemerkung "Status offen" steht in vorher in ALLEN Section-Zeilen -> fuell_konstant (ableitbar).
    Ein individueller Wert in derselben Spalte ist nicht ableitbar; eine Formelzelle ist "formel"."""
    p = baue_paar(tmp_path, "synth_9400", MIT_NEUER_SECTION)
    res = messen.messe_paar(p)
    k = res["soll_klassen"]
    assert set(k) == set(messen.SOLL_KLASSEN)
    assert k["fuell_konstant"] == 1 and k["nicht_ableitbar"] == 0 and k["haupt"] == res["soll_zellen"]
    assert res["soll_klassen_erreicht"]["haupt"] == res["erreicht_auto"]

    q = baue_paar(tmp_path, "synth_9401", MIT_NEUER_SECTION)
    wb = load_workbook(str(q / "nachher.xlsx"))
    ws = wb["CDE"]
    neue = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("4"))
    ws.cell(neue, 4).value = "Sonderfall mit Rückfrage beim Kunden"
    ws["E2"] = "=LEN(C2)"
    wb.save(str(q / "nachher.xlsx"))
    res = messen.messe_paar(q)
    assert res["soll_klassen"]["nicht_ableitbar"] == 1 and res["soll_klassen"]["fuell_konstant"] == 0
    assert res["soll_klassen"]["formel"] == 1

    b = messen.fasse_zusammen([dict(res, name="a")], "entwicklung", "2", False)
    g = b["gesamt"]
    assert g["soll_zellen_ableitbar"] == res["soll_klassen"]["haupt"] + 1          # haupt + formel (fuell_konstant 0)
    k7 = next(x for x in b["kriterien"] if x["name"].startswith("Abdeckung ableitbare"))
    assert k7["ampel"] in ("ok", "Hinweis") and k7["einheit"] == "Zellen"


@pytest.mark.parametrize("layout", ["L3"])
def test_klassen_v2_prio_spalte_je_zeilentyp(tmp_path, layout):
    """L3: Section-Zeilen haben Prio "A", Absatz-Zeilen "B" -- je Zeilentyp konstant, also fuell_konstant (neu UND entfernt)."""
    res = messen.messe_paar(baue_paar(tmp_path, "synth_9402", MIT_NEUER_SECTION, layout))
    assert res["soll_klassen"]["fuell_konstant"] >= 2 and res["soll_klassen"]["nicht_ableitbar"] == 0


def test_restarbeit_formel_und_hinweis_ampel():
    paar = {"name": "x", "aenderungen": 50, "restarbeit": {"pruefen": 3, "nicht_uebernommen": 1,
                                                          "leere_zellen_neue_zeilen": 2, "luecken_woerter": 40}}
    b = messen.fasse_zusammen([{**messen._fehlgeschlagen("x", ValueError()), **paar}], "entwicklung", "2", False)
    g = b["gesamt"]
    assert g["restarbeit_je_100"] == 20.0                     # 100 * (3 + 1 + 2 + 40/10) / 50
    assert "÷ 10" in g["restarbeit_formel"] and g["restarbeit"]["luecken_woerter"] == 40
    k = next(x for x in b["kriterien"] if x["name"].startswith("Restarbeit"))
    assert k["ampel"] == "Hinweis" and k["wert"] == "20,0"     # ueber dem Richtwert 12: nur Hinweis, Exit-Code unberuehrt
    assert all(x["ampel"] != "NICHT ok" for x in b["kriterien"] if x["nr"] == 7)


def test_pruefcodes_histogramm_und_sperr_quote(tmp_path, monkeypatch):
    p = baue_paar(tmp_path, "synth_9403")
    original = messen.plane

    def gesperrt(pdf, xl, cfg=None):
        plan = original(pdf, xl, cfg)
        ae = next(a for a in plan.aenderungen if a.ziel.endswith("C2"))
        ae.status, ae.konfidenz = PRUEFEN, "mittel"
        ae.pruefcodes = ["a1_abweichung", "Freitext mit Quokka-Inhalt"]
        return plan
    monkeypatch.setattr(messen, "plane", gesperrt)
    res = messen.messe_paar(p, mit_vorschlaegen=True)
    assert res["pruefcodes"] == {"a1_abweichung": 1, "sonstiger_code": 1}          # Freitext wird nie berichtet
    assert res["sperr_quote"]["a1_abweichung"] == {"gesperrt": 1, "messbar": 1, "richtig_nach_freigabe": 1}
    assert res["restarbeit"]["pruefen"] == 1
    ohne = messen.messe_paar(p)                                                     # ohne Vorschlags-Excel: nicht messbar
    assert ohne["sperr_quote"]["a1_abweichung"]["messbar"] == 0
    b = messen.fasse_zusammen([dict(res, name="a")], "entwicklung", "2", True)
    assert b["gesamt"]["sperr_quote"]["a1_abweichung"]["sperr_quote"] == 1.0
    assert "quokka" not in json.dumps(b).lower() and "quokka" not in messen._md(b).lower()


def test_vergleich_mit_letztem_bericht_und_gegen(tmp_path, capsys):
    baue_paar(korpus_dir(), "synth_9404")
    assert messen.main(["--korpus", "entwicklung"]) == 0
    assert messen.main(["--korpus", "entwicklung", "--vergleich"]) == 0
    basis = config.pfad("messung_dir")
    b = json.loads((basis / "002_entwicklung" / "bericht.json").read_text(encoding="utf-8"))
    v = b["vergleich"]
    assert v["gegen"] == "001_entwicklung" and v["vergleichbar"]
    assert v["kennzahlen"]["stille_fehler"] == {"vorher": 0, "jetzt": 0, "delta": 0}
    assert "Vergleich mit Bericht 001_entwicklung" in (basis / "002_entwicklung" / "bericht.md").read_text(encoding="utf-8")
    assert messen.main(["--korpus", "entwicklung", "--vergleich", "--gegen", "1"]) == 0
    b3 = json.loads((basis / "003_entwicklung" / "bericht.json").read_text(encoding="utf-8"))
    assert b3["vergleich"]["gegen"] == "001_entwicklung"
    assert messen.main(["--korpus", "entwicklung", "--vergleich", "--gegen", "999"]) == 2
    assert not (basis / "004_entwicklung").exists()


def test_vergleich_deltas_je_kennzahl_und_familie():
    alt = {"korpus": "entwicklung", "generator_version": "1",
           "gesamt": {"stille_fehler": 7, "treffergenauigkeit_auto": 0.98, "idempotent": 57, "vorschlags_genauigkeit": None},
           "je_stoerer": {"fussnote": {"paare": 5, "stille_fehler": 2, "abdeckung_ohne_nachfrage": 0.8}}}
    neu = {"korpus": "entwicklung", "generator_version": "2",
           "gesamt": {"stille_fehler": 3, "treffergenauigkeit_auto": 0.99, "idempotent": 60, "vorschlags_genauigkeit": 0.8,
                      "neu_nur_jetzt": 4},
           "je_stoerer": {"fussnote": {"paare": 5, "stille_fehler": 0, "abdeckung_ohne_nachfrage": 0.9},
                          "bild": {"paare": 2, "stille_fehler": 0, "abdeckung_ohne_nachfrage": 1.0}}}
    v = messen.vergleiche(alt, neu, "006_entwicklung")
    assert not v["vergleichbar"]                                                    # andere Generator-Version
    assert v["kennzahlen"]["stille_fehler"]["delta"] == -4 and v["kennzahlen"]["idempotent"]["delta"] == 3
    assert v["kennzahlen"]["vorschlags_genauigkeit"] == {"vorher": None, "jetzt": 0.8, "delta": None}
    assert v["je_stoerer"]["fussnote"]["stille_fehler"]["delta"] == -2
    assert v["je_stoerer"]["bild"]["paare"] == {"vorher": None, "jetzt": 2, "delta": None}


def test_korpus_ex_test_ohne_regel11_warnton_und_je_erzeuger(tmp_path, capsys):
    k = korpus_dir("ex_test")
    p = baue_paar(k, "synth_1001")
    (p / "erwartet.json").write_text(json.dumps({"format": 1, "generator_version": "2", "erzeuger": "libreoffice",
                                                 "luecken": [], "stoerer": []}), encoding="utf-8")
    baue_paar(k, "echt_0002")
    assert messen.main(["--korpus", "ex_test"]) == 0
    out = capsys.readouterr().out
    assert "Regel 11" not in out and "offen als Entwicklung" in out
    b = json.loads(next(config.pfad("messung_dir").glob("001_ex_test/bericht.json")).read_text(encoding="utf-8"))
    assert b["korpus"] == "ex_test"
    assert set(b["je_erzeuger"]) == {"libreoffice", "unbekannt"} and b["je_erzeuger"]["libreoffice"]["paare"] == 1
    assert {"paare", "stille_fehler", "treffergenauigkeit_auto", "abdeckung_ohne_nachfrage"} <= set(b["je_erzeuger"]["libreoffice"])
    assert "Je PDF-Erzeuger" in (config.pfad("messung_dir") / "001_ex_test" / "bericht.md").read_text(encoding="utf-8")


# =======================================================================================
# Nacharbeit Welle 2: Messartefakte eng gefasst -- jede Ausnahme hat einen Sabotage-Test (echter stiller Fehler zaehlt weiter)
# =======================================================================================
from types import SimpleNamespace  # noqa: E402
from collections import Counter as _Counter  # noqa: E402


def _ausw(**felder):
    """Minimale _Auswertung ohne Konstruktor (nur die Felder, die die getesteten Methoden lesen)."""
    a = messen._Auswertung.__new__(messen._Auswertung)
    a.__dict__.update(felder)
    return a


def _ae(id_, alt="", neu="", status=PRUEFEN, codes=()):
    return SimpleNamespace(id=id_, status=status, pruefcodes=list(codes), alt=alt, neu=neu)


def _wo(*aend):
    return {a.id: (_Counter(messen._wort_liste(a.neu)), _Counter(messen._wort_liste(a.alt))) for a in aend}


def test_wirkung_nachbaraenderung_mit_gleichem_wort_ist_kein_regel27():
    offen = _ae("A2", alt="Gras", neu="fachgerecht")           # offen (pruefen)
    nachbar = _ae("A1", alt="ruht", neu="fachgerecht", status="uebernommen")   # angewendet, fuegt zufaellig dasselbe Wort ein
    a = _ausw(ae_woerter=_wo(offen, nachbar))
    vorher = "Der Quokka frisst Gras und ruht"
    ergebnis = "Der Quokka frisst Gras und fachgerecht"       # nur die Nachbaraenderung wirkte
    assert not a._wirkung_vorhanden(offen, vorher, ergebnis, ["A1"])
    # Sabotage: die offene Aenderung hat WIRKLICH gewirkt (Wort neu UND altes Wort weg) -> weiter Regel-27-Verstoss
    assert a._wirkung_vorhanden(offen, vorher, "Der Quokka frisst fachgerecht und fachgerecht", ["A1"])
    # Sabotage: ohne Nachbaraenderung bleibt das Verhalten wie frueher
    assert a._wirkung_vorhanden(offen, vorher, "Der Quokka frisst fachgerecht und ruht", [])
    # Sabotage: reine Loeschung einer offenen Aenderung, Wort ist weg -> Verstoss
    loesch = _ae("A3", alt="frisst", neu="")
    a2 = _ausw(ae_woerter=_wo(loesch))
    assert a2._wirkung_vorhanden(loesch, vorher, "Der Quokka Gras und ruht", [])
    assert not a2._wirkung_vorhanden(loesch, vorher, vorher, [])


def test_neu_abgleich_fuellspalte_verfaelscht_paarung_nicht():
    soll = [(5, ("", "", "Neuer Absatz eins", "B"), ("", "", "Neuer Absatz eins", "B")),
            (7, ("", "", "Neuer Absatz zwei", "B"), ("", "", "Neuer Absatz zwei", "B"))]
    ist = [(7, ("", "", "Neuer Absatz zwei", ""), ("", "", "Neuer Absatz zwei", "")),
           (5, ("", "", "Neuer Absatz eins", ""), ("", "", "Neuer Absatz eins", ""))]
    paar, ort = messen._neu_abgleich(soll, ist, {1, 2, 3})
    assert paar == {0: 1, 1: 0} and ort == {0: True, 1: True}
    # Sabotage: anderer Inhalt in der Haupt-Spalte wird NICHT gleichgesetzt
    falsch = (9, ("", "", "Anderer Text", "B"), ("", "", "Anderer Text", "B"))
    assert messen._neu_abgleich(soll[:1], [falsch], {1, 2, 3})[0] == {}                 # weder Anker noch Inhalt
    assert messen._neu_abgleich(soll[:1], [(5,) + falsch[1:]], {1, 2, 3})[1] == {0: True}   # nur ueber den Anker (Stufe 3)


def test_paare_zeilenaehnlichkeit_nur_haupt_spalten():
    va = [messen.Zeile(1, ("", "", "Alle Anschlussleitungen sind gut", "B"), ("", "", "Alle Anschlussleitungen sind gut", "B"))]
    xb = [messen.Zeile(1, ("", "", "Alle Anschlussleitungen sind gut", "X"), ("", "", "Alle Anschlussleitungen sind gut", "X"))]
    assert messen._paare(va, xb, {3}) == [(0, 0)]
    # Sabotage: voellig anderer Inhalt der Haupt-Spalte bleibt ungepaart, auch wenn die Fuellspalte gleich ist
    xb2 = [messen.Zeile(1, ("", "", "Voellig anderer Text hier", "B"), ("", "", "Voellig anderer Text hier", "B"))]
    assert messen._paare(va, xb2, {3}) == []


def _farbe_ausw(farbe_gemeldet, pdf_neu=()):
    return _ausw(farbe_gemeldet=_Counter(farbe_gemeldet), pdf_neu=_Counter(pdf_neu),
                 ae_woerter={"A1": (_Counter(), _Counter()), "A2": (_Counter(), _Counter())})


def test_gemeldete_gruene_luecke_ist_nicht_still_aber_echter_fehler_schon():
    sperre = _ae("A1", status=PRUEFEN, codes=["sperre:unbekannte_farbe"], alt="a", neu="b")
    a = _farbe_ausw(["gruenes", "wort"], pdf_neu=["blau"])
    vorher = "Das ist ein Satz."
    soll = "Das ist ein gruenes wort Satz."
    assert a._nur_gemeldete_farbe(vorher, soll, vorher, [sperre])                 # Ergebnis = vorher, nur Gruen fehlt
    # Sabotage 1: zusaetzlich fehlt ein BLAUES Wort, das das Tool haette uebernehmen muessen -> weiter still
    assert not a._nur_gemeldete_farbe(vorher, soll.replace("Satz", "blau Satz"), vorher, [sperre])
    # Sabotage 2: im Ergebnis steht ein falsches Wort zu viel -> weiter still
    assert not a._nur_gemeldete_farbe("Das ist ein falsch Satz.", soll, vorher, [sperre])
    # Sabotage 3: keine gemeldete Farb-Luecke -> weiter still
    assert not _farbe_ausw([])._nur_gemeldete_farbe(vorher, soll, vorher, [sperre])
    # Sabotage 4: fehlendes Wort ist nicht in der gemeldeten Farbe -> weiter still
    assert not a._nur_gemeldete_farbe(vorher, "Das ist ein anderes Satz.", vorher, [sperre])
    # Sabotage 5: nur Leerraum-Unterschied bleibt ein Fehler
    assert not a._nur_gemeldete_farbe("Das  ist ein Satz.", vorher, vorher, [sperre])
    # Sabotage 6: Wort steht auch blau im PDF und der Absatz ist NICHT wegen der Farbe gesperrt -> koennte ein verpasstes
    # blaues Wort sein -> weiter still
    ohne_code = _ae("A2", status=PRUEFEN, codes=["a1_ziffer"])
    b = _farbe_ausw(["wort"], pdf_neu=["wort"])
    assert not b._nur_gemeldete_farbe(vorher, "Das ist ein wort Satz.", vorher, [ohne_code])


def test_eingefuegte_zeile_als_gepaarte_zeile_nur_wenn_ergebnis_und_soll_gleich():
    gleiche = SimpleNamespace(zuordnung={3: 4}, zellen={(3, 3): "Neuer Text", (3, 4): "B"})
    falsch = SimpleNamespace(zuordnung={3: 4}, zellen={(3, 3): "Anderer Text", (3, 4): "B"})
    a = _ausw(S=[gleiche], I=[gleiche])
    assert a._neu_als_gepaarte_zeile(0, {3: "Neuer Text"})
    b = _ausw(S=[gleiche], I=[falsch])                                       # Ergebnis weicht vom Soll ab
    assert not b._neu_als_gepaarte_zeile(0, {3: "Neuer Text"})
    c = _ausw(S=[SimpleNamespace(zuordnung={}, zellen={})], I=[gleiche])    # Soll sieht dort keine gepaarte Zeile
    assert not c._neu_als_gepaarte_zeile(0, {3: "Neuer Text"})
