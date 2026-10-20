"""Tests fuer den zweiten Pruefweg (src/konsens.py, Welle 4): Wort-Bildprobe und Konsens-Text.

Nur erfundene Daten (Regel 16): ein Paar aus dem Generator v3 (Seed ausserhalb der Korpora, Regel 11) und ein kleines
Word-PDF aus tests/_word_pdf.py. Geschrieben wird nur in tmp_path.
"""
from __future__ import annotations

import random
import re
import zipfile
from copy import copy, deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles.colors import Color

from src import codes, config, konsens, messen
from src import applier
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.models import ADD, DEL, KEEP, PRUEFEN, UEBERNOMMEN, Aenderung
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests._word_pdf import EBENEN, WordDok
from tests.fixtures.generator import erzeuge_paar

SEED = 372                                             # Entwicklungs-Seed ausserhalb der Korpora 301-360 / 2001-2030
KLEIN = {"basis": "word_markup", "name": "klein", "zeilen": (45, 60)}
EINST = {"legende": "auto", "nachher_format": "markiert"}


@pytest.fixture
def word_markiert():
    """Legende "auto" und farbige Markierung wie im Korpus nachbau_word (die conftest-Weichen stehen sonst auf alt / keine)."""
    with messen.konfig_ueberschreibung(EINST):
        yield


@pytest.fixture(scope="module")
def basen(tmp_path_factory):
    """Vier erfundene Paare (Generator v3) mit Analyse und Plan."""
    out = []
    for seed in range(SEED, SEED + 4):
        ordner = erzeuge_paar(seed, KLEIN, tmp_path_factory.mktemp("konsens"), "nachbau_word")
        with messen.konfig_ueberschreibung(EINST):
            pa = analysiere_pdf(ordner / "aenderungen.pdf")
            plan = plane(pa, ExcelModell(ordner / "vorher.xlsx"))
        out.append(SimpleNamespace(ordner=ordner, vorher=ordner / "vorher.xlsx", pa=pa, plan=plan))
    return out


@pytest.fixture(scope="module")
def basis(basen):
    return basen[0]


def _schreibe(basis, ziel: Path, mit_pdf: bool = True):
    aend = deepcopy(basis.plan.aenderungen)
    res = uebernehmen(aend, basis.vorher, ziel, None, absaetze=basis.plan.absaetze, pdf=basis.pa if mit_pdf else None)
    return aend, res


def _teile(pfad: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(pfad) as z:
        return {n: z.read(n) for n in z.namelist()}


# ====================================================================== Konfiguration, Codes
def test_konfiguration_und_codes():
    k = config.lade("zuordnung")["konsens"]
    assert k["aktiv"] is True and set(k) >= {"aktiv", "bild", "text"}
    for c in (konsens.CODE_TEXT, konsens.CODE_BILD):
        s = codes.satz(c)
        assert s and len(codes.kurz(c).split()) <= codes.KURZ_MAX_WOERTER and "_" not in s


def test_aus_schaltet_beides_ab(basis, tmp_path, word_markiert, monkeypatch):
    monkeypatch.setitem(config.lade("zuordnung")["konsens"], "aktiv", False)
    aend, res = _schreibe(basis, tmp_path / "aus.xlsx")
    assert res.konsens == {} and res.text_pruefung is None
    b = konsens.bildprobe(basis.pa, aend)
    assert not b.gelaufen and b.grund == "ausgeschaltet"


# ====================================================================== Normalfall
def test_normalfall_keine_abweichung_ergebnis_unveraendert(basis, tmp_path, word_markiert, monkeypatch):
    aend_mit, res = _schreibe(basis, tmp_path / "mit.xlsx")
    bild, text = res.konsens["bild"], res.konsens["text"]
    assert bild["gelaufen"] and bild["geprueft"] > 20 and bild["abweichend"] == 0
    assert bild["unpruefbar"] > 0                                          # die Zahl steht im Ergebnis
    assert text["gelaufen"] and text["geprueft"] > 3 and text["abweichend"] == 0 and text["fremde_zellen"] == 0
    assert not [w for w in res.warnungen if "Zweitprüfung" in w]
    auto = [a for a in aend_mit if a.status == UEBERNOMMEN]
    assert auto and not any("konsens_bild" in a.pruefcodes or "konsens_text" in a.pruefcodes for a in aend_mit)
    # ohne Konsens: gleiche Stati und dieselbe Datei
    monkeypatch.setitem(config.lade("zuordnung")["konsens"], "aktiv", False)
    aend_ohne, _ = _schreibe(basis, tmp_path / "ohne.xlsx")
    assert [(a.id, a.status, a.wirkung) for a in aend_mit] == [(a.id, a.status, a.wirkung) for a in aend_ohne]
    assert _teile(tmp_path / "mit.xlsx") == _teile(tmp_path / "ohne.xlsx")


def test_bildprobe_im_plan_aendert_beim_normalfall_nichts(basis, word_markiert):
    aend = deepcopy(basis.plan.aenderungen)
    vorher = [(a.id, a.status) for a in aend]
    erg = konsens.bildprobe_anwenden(aend, basis.pa)
    assert erg.gelaufen and erg.abweichend == 0 and erg.geprueft > 20
    assert sum(erg.gruende.values()) == erg.unpruefbar
    assert [(a.id, a.status) for a in aend] == vorher


# ====================================================================== (b) Konsens-Text: Mutationen der geschriebenen Zellen
def _zellen_mit_markierung(wb, art: str):
    """(Zelle, Block) fuer jeden Textblock einer Rich-Text-Zelle der Klasse art ('alt' | 'neu' | 'normal')."""
    out = []
    for ws in wb:
        for (_, _), z in sorted(ws._cells.items()):
            if not isinstance(z.value, CellRichText):
                continue
            klassen = [("alt" if b.font.strike else "neu" if (b.font.color is not None and b.font.color.rgb == "FFFF0000")
                        else "normal") if isinstance(b, TextBlock) else None for b in z.value]
            if "alt" not in klassen and "neu" not in klassen:
                continue
            for b, k in zip(z.value, klassen):
                if k == art and len(b.text.strip()) >= 3 and any(c.isalpha() for c in b.text):
                    out.append((z, b))
    return out


def _buchstabe_aendern(b: TextBlock):
    """Ersetzt den ersten Buchstaben durch einen anderen; Rueckgabe: Rueckgaengig-Funktion."""
    alt = b.text
    i = next(k for k, c in enumerate(alt) if c.isalpha())
    b.text = alt[:i] + ("x" if alt[i].lower() != "x" else "y") + alt[i + 1:]
    return lambda: setattr(b, "text", alt)


def _font_aendern(b: TextBlock, **kw):
    alt = b.font
    f = copy(alt)
    for k, v in kw.items():
        setattr(f, k, v)
    b.font = f
    return lambda: setattr(b, "font", alt)


def _text_leeren(b: TextBlock):
    alt = b.text
    b.text = ""
    return lambda: setattr(b, "text", alt)


def _nimm(liste, n, zufall):
    liste = list(liste)
    zufall.shuffle(liste)
    return liste[:n]


@pytest.fixture(scope="module")
def ergebnis_datei(basis, tmp_path_factory):
    """Geschriebenes Ergebnis (mit Konsens) und die Mappen zum Abwandeln im Speicher."""
    with messen.konfig_ueberschreibung(EINST):
        ziel = tmp_path_factory.mktemp("ergebnis") / "e.xlsx"
        aend = deepcopy(basis.plan.aenderungen)
        res = uebernehmen(aend, basis.vorher, ziel, None, absaetze=basis.plan.absaetze, pdf=basis.pa)
        return SimpleNamespace(pfad=ziel, res=res, aend=aend, quelle=load_workbook(str(basis.vorher), rich_text=True),
                               ziel=load_workbook(str(ziel), rich_text=True))


def test_konsens_text_erkennt_mindestens_95_prozent_der_eingebauten_fehler(ergebnis_datei, word_markiert):
    e = ergebnis_datei
    pruefe = lambda: e.res.text_pruefung(e.pfad, _wb_quelle=e.quelle, _wb_ziel=e.ziel)      # noqa: E731
    kontrolle = pruefe()
    assert kontrolle.abweichungen == {} and kontrolle.fremd == 0 and kontrolle.geprueft > 3
    zufall = random.Random(7)
    neue = _zellen_mit_markierung(e.ziel, "neu")
    alte = _zellen_mit_markierung(e.ziel, "alt")
    normale = _zellen_mit_markierung(e.ziel, "normal")
    assert len(neue) >= 3 and len(alte) >= 3
    mutationen: list[tuple[str, object]] = []
    for z, b in _nimm(neue, 8, zufall):
        mutationen += [("buchstabe_neu", lambda b=b: _buchstabe_aendern(b)),
                       ("farbe_neu_falsch", lambda b=b: _font_aendern(b, color=Color(rgb="FF0000FF"))),
                       ("strich_zu_viel", lambda b=b: _font_aendern(b, strike=True)),
                       ("lauf_fehlt", lambda b=b: _text_leeren(b))]
    for z, b in _nimm(alte, 8, zufall):
        mutationen += [("buchstabe_alt", lambda b=b: _buchstabe_aendern(b)),
                       ("strich_fehlt", lambda b=b: _font_aendern(b, strike=False)),
                       ("farbe_alt_falsch", lambda b=b: _font_aendern(b, color=Color(rgb="FF0000FF"))),
                       ("lauf_fehlt", lambda b=b: _text_leeren(b))]
    for z, b in _nimm(normale, 6, zufall):
        mutationen.append(("buchstabe_normal", lambda b=b: _buchstabe_aendern(b)))
    # Zelle einer unberuehrten Zeile (z. B. eine andere Spec): ein Buchstabe geaendert
    veraendert = {(ws.title, z.row, z.column) for ws in e.ziel for z in ws._cells.values()
                  if isinstance(z.value, CellRichText) and any(isinstance(b, TextBlock) and (b.font.strike or (
                      b.font.color is not None and b.font.color.rgb == "FFFF0000")) for b in z.value)}
    fremde = [z for ws in e.ziel for z in ws._cells.values()
              if isinstance(z.value, str) and len(z.value) > 6 and any(c.isalpha() for c in z.value)
              and (ws.title, z.row, z.column) not in veraendert]
    for z in _nimm(fremde, 10, zufall):
        def fremd_aendern(z=z):
            alt = z.value
            i = next(k for k, c in enumerate(alt) if c.isalpha())
            z.value = alt[:i] + ("x" if alt[i].lower() != "x" else "y") + alt[i + 1:]
            return lambda: setattr(z, "value", alt)
        mutationen.append(("fremde_zelle", fremd_aendern))
    gefunden: dict[str, list[bool]] = {}
    for art, mach in mutationen:
        zurueck = mach()
        try:
            gefunden.setdefault(art, []).append(bool(pruefe().abweichungen))
        finally:
            zurueck()
    assert pruefe().abweichungen == {}                                      # alles wieder wie vorher
    gesamt = [x for v in gefunden.values() for x in v]
    assert len(gesamt) >= 40
    quote = sum(gesamt) / len(gesamt)
    print("Konsens-Text Erkennung:", f"{sum(gesamt)}/{len(gesamt)}", {k: f"{sum(v)}/{len(v)}" for k, v in gefunden.items()})
    assert quote >= 0.95, {k: f"{sum(v)}/{len(v)}" for k, v in gefunden.items()}
    for art, v in gefunden.items():
        assert sum(v) / len(v) >= 0.75, (art, v)                           # keine Fehlerart bleibt ganz unbemerkt


def test_konsens_text_erkennt_geloeschte_und_verschobene_inhalte(ergebnis_datei, word_markiert):
    e = ergebnis_datei
    ws = next(w for w in e.ziel if any(isinstance(z.value, CellRichText) for z in w._cells.values()))
    z = next(z for z in ws._cells.values() if isinstance(z.value, CellRichText) and z.row > 12)
    alt = z.value
    z.value = None                                                         # Zelle leer
    try:
        assert e.res.text_pruefung(e.pfad, _wb_quelle=e.quelle, _wb_ziel=e.ziel).abweichungen
    finally:
        z.value = alt
    assert e.res.text_pruefung(e.pfad, _wb_quelle=e.quelle, _wb_ziel=e.ziel).abweichungen == {}


# ====================================================================== Neuschreiben ohne die beanstandete Aenderung
def test_abweichung_nach_dem_schreiben_stuft_zurueck_und_schreibt_neu(basis, tmp_path, word_markiert, monkeypatch):
    original = applier.patch_schreiben
    aufrufe = []

    def sabotiert(quelle, ziel, zell_inhalte, *a, **kw):
        erg = original(quelle, ziel, zell_inhalte, *a, **kw)
        aufrufe.append(1)
        if len(aufrufe) == 1:                                               # nur beim ersten Schreiben: ein Buchstabe falsch
            frag = next(st.text[:8] for z in zell_inhalte if z.stuecke for st in z.stuecke
                        if st.format == "neu" and re.fullmatch(r"[A-Za-z]{8,}.*", st.text) and re.fullmatch(r"[A-Za-z]{8}", st.text[:8]))
            teile = _teile(Path(ziel))
            roh = teile["xl/sharedStrings.xml"].decode("utf-8")
            i = roh.rindex(frag)                                              # neu angehaengte Eintraege stehen hinten
            roh = roh[:i] + ("x" if frag[0] != "x" else "y") + roh[i + 1:]
            teile["xl/sharedStrings.xml"] = roh.encode("utf-8")
            with zipfile.ZipFile(ziel) as z:
                namen = z.namelist()
            Path(ziel).unlink()
            with zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as z:
                for n in namen:
                    z.writestr(n, teile[n])
        return erg
    monkeypatch.setattr(applier, "patch_schreiben", sabotiert)
    aend = deepcopy(basis.plan.aenderungen)
    vorher_auto = {a.id for a in aend if a.status == UEBERNOMMEN}
    ziel = tmp_path / "s.xlsx"
    res = uebernehmen(aend, basis.vorher, ziel, None, absaetze=basis.plan.absaetze, pdf=basis.pa)
    assert len(aufrufe) == 2 and res.konsens["neu_geschrieben"] == 1
    zurueck = [a for a in aend if "konsens_text" in a.pruefcodes]
    assert zurueck and all(a.status == PRUEFEN and a.id in vorher_auto for a in zurueck)       # wie bei einer A3-Abweichung
    assert all(a.wirkung == "abweichung" and a.grund_einfach is not None for a in zurueck)
    assert any("Zweitprüfung" in w for w in res.warnungen)
    assert [a for a in aend if a.id in vorher_auto and a.status == UEBERNOMMEN]                # der Rest bleibt automatisch
    # die Ergebnisdatei steht ohne die beanstandete Aenderung: der Konsens-Text findet nichts mehr
    assert res.text_pruefung(ziel).abweichungen == {}
    assert not any(a.status == UEBERNOMMEN and a.id in {z.id for z in zurueck} for a in aend)


# ====================================================================== (a) Wort-Bildprobe
def _wort_dok(mit_leiste: bool = True):
    """Erfundenes Word-PDF: Absatz mit Aenderung mitten in der Zeile, Sprechblase mit Verbindungslinie ueber den Folgeworten,
    gelber Platzhalter."""
    d = WordDok(mit_leiste=mit_leiste)
    tops = d.absatz([[("Der Kessel heizt das Wasser ", "k"), ("alter", "d"), ("neuer", "n"), (" Wert im Keller und haelt die", "k")],
                     [("Temperatur im Kreislauf ueber den ganzen Tag stabil und gleichmaessig.", "k")]], ebene=1)
    if mit_leiste:
        d.sprechblase(tops[0] + 1, "hat formatiert:", "Schriftart", anker=(300.0, tops[0]))      # Linie ueber Folgeworten
    t2 = d.absatz([[("Jede Kontrolle wird im Betriebsbuch ", "k"), ("[Datum]", "d"), ("Montag", "n"),
                    (" eingetragen und bestaetigt.", "k")]], ebene=1)
    d.gelb(EBENEN[1] + 5.0, EBENEN[1] + 60.0, t2[0])
    d.absatz([[("Schwarz unterstrichen bleibt ", "k"), ("Ventil", "u"), (" ohne Aenderung im Text stehen.", "k")]], ebene=1)
    return d


def _fake_aenderungen(pa) -> list[Aenderung]:
    out = []
    for b in pa.blocks:
        if b.hat_aenderung and not b.ignoriert:
            out.append(Aenderung(id=f"A{len(out) + 1:04d}", typ="text_ersetzt", seite=b.page, bbox=b.bbox, section=None, alt="", neu="",
                                 woerter=1, status=UEBERNOMMEN, block_idx=b.idx))
    return out


def test_bildprobe_word_pdf_ohne_fehlalarm_mit_leiste_und_verbindungslinie(tmp_path, word_markiert):
    pdf = _wort_dok().speichern(tmp_path / "w.pdf")
    pa = analysiere_pdf(pdf)
    assert pa.legende_modus == "linie" and pa.markup_leiste
    aend = _fake_aenderungen(pa)
    assert aend
    erg = konsens.bildprobe(pa, aend, pdf=pdf)
    assert erg.gelaufen and erg.abweichend == 0 and erg.abweichungen == {}
    assert erg.geprueft >= 10 and erg.unpruefbar >= 1                          # gelber Platzhalter und kurze Woerter
    # ohne Leiste und Linie genauso
    pdf2 = _wort_dok(mit_leiste=False).speichern(tmp_path / "w2.pdf")
    pa2 = analysiere_pdf(pdf2)
    erg2 = konsens.bildprobe(pa2, _fake_aenderungen(pa2), pdf=pdf2)
    assert erg2.gelaufen and erg2.abweichend == 0


def _nur(pa, status: set[str]):
    """Kopie der Analyse mit nur den Woertern dieser Stati (ohne Hinweis) in den Bloecken mit Aenderung."""
    p = deepcopy(pa)
    for b in p.blocks:
        b.tokens = [t for t in b.tokens if t.status in status and not t.hinweis]
    return p


def _umdeuten(p, abbildung: dict[str, str]):
    for b in p.blocks:
        for t in b.tokens:
            t.status = abbildung[t.status]
    return p


@pytest.mark.parametrize("von,abbildung", [
    ("add_del", {ADD: DEL, DEL: ADD}),
    ("keep_add", {KEEP: ADD}),
    ("keep_del", {KEEP: DEL}),
    ("add_keep", {ADD: KEEP}),
    ("del_keep", {DEL: KEEP}),
])
def test_bildprobe_erkennt_mindestens_95_prozent_falscher_woerter(basen, word_markiert, von, abbildung):
    """Absichtlich falsch gelesene Woerter (ADD <-> DEL vertauscht, KEEP als ADD/DEL, ADD/DEL als KEEP), ueber vier Paare: Von
    den Woertern, die die Probe beurteilen kann, muss sie mindestens 95 % beanstanden; ohne Fehler im Lesen kein Alarm."""
    quelle = {"add_del": {ADD, DEL}, "keep_add": {KEEP}, "keep_del": {KEEP}, "add_keep": {ADD}, "del_keep": {DEL}}[von]
    beurteilbar = erkannt = 0
    for b in basen:
        aend = deepcopy(b.plan.aenderungen)
        echt = _nur(b.pa, quelle)
        kontrolle = konsens.bildprobe(echt, aend)
        assert kontrolle.gelaufen and kontrolle.abweichend == 0                      # richtig gelesen: kein Alarm
        falsch = _umdeuten(deepcopy(echt), {s: abbildung.get(s, s) for s in (ADD, DEL, KEEP)})
        erg = konsens.bildprobe(falsch, aend)
        assert all("Seite" in t for v in erg.abweichungen.values() for t in v)       # Beschreibung ohne Dokumentinhalt
        beurteilbar += kontrolle.geprueft
        erkannt += erg.abweichend
    assert beurteilbar >= 8, (von, beurteilbar)
    assert erkannt / beurteilbar >= 0.95, (von, erkannt, beurteilbar)


def test_bildprobe_im_plan_stuft_betroffene_aenderungen_auf_pruefen(basis, word_markiert):
    aend = deepcopy(basis.plan.aenderungen)
    falsch = _umdeuten(deepcopy(basis.pa), {ADD: DEL, DEL: ADD, KEEP: KEEP, "mod": "mod"})
    erg = konsens.bildprobe_anwenden(aend, falsch)
    assert erg.abweichungen
    for aid in erg.abweichungen:
        a = next(x for x in aend if x.id == aid)
        assert a.status == PRUEFEN and "konsens_bild" in a.pruefcodes and a.grund_einfach
        assert a.konfidenz == "niedrig"


def test_bildprobe_im_applier_schreibt_beanstandete_aenderung_nicht(basis, tmp_path, word_markiert):
    aend = deepcopy(basis.plan.aenderungen)
    vorher_auto = {a.id for a in aend if a.status == UEBERNOMMEN}
    falsch = _umdeuten(deepcopy(basis.pa), {ADD: DEL, DEL: ADD, KEEP: KEEP, "mod": "mod"})
    ziel = tmp_path / "b.xlsx"
    res = uebernehmen(aend, basis.vorher, ziel, None, absaetze=basis.plan.absaetze, pdf=falsch)
    assert res.konsens["bild"]["abweichend"] > 0
    beanstandet = [a for a in aend if "konsens_bild" in a.pruefcodes]
    assert beanstandet and all(a.status == PRUEFEN and a.id in vorher_auto and a.wirkung == "" for a in beanstandet)
    assert res.text_pruefung(ziel).abweichungen == {}                          # der Rest passt
    assert not any(a.status == UEBERNOMMEN for a in beanstandet)


def test_bildprobe_ohne_pdf_oder_bei_word_eingabe_laeuft_nicht(basis, word_markiert):
    aend = deepcopy(basis.plan.aenderungen)
    docx = SimpleNamespace(legende_modus="docx", pfad="x.docx", blocks=[], seiten_masse={})
    erg = konsens.bildprobe(docx, aend)
    assert not erg.gelaufen and "PDF" in erg.grund
    weg = deepcopy(basis.pa)
    weg.pfad = str(Path(basis.pa.pfad).with_name("gibt_es_nicht.pdf"))
    erg = konsens.bildprobe(weg, aend)
    assert not erg.gelaufen and erg.abweichend == 0


def test_ohne_pdf_keine_bildprobe_und_ausfall_der_zweitpruefung_ist_nie_still(basis, tmp_path, word_markiert, monkeypatch):
    aend, res = _schreibe(basis, tmp_path / "a.xlsx", mit_pdf=False)
    assert "bild" not in res.konsens and res.konsens["text"]["gelaufen"]            # Text-Probe braucht das PDF nicht

    def kaputt(*a, **kw):
        raise ValueError("kaputt")
    monkeypatch.setattr(konsens, "pruefe_text", kaputt)
    aend2, res2 = _schreibe(basis, tmp_path / "b.xlsx")
    assert res2.konsens["text"] == {"gelaufen": False, "fehler": "ValueError"}
    assert any("Zweitprüfung" in w for w in res2.warnungen)                           # der Ausfall wird gemeldet
    assert [(x.id, x.status) for x in aend2] == [(x.id, x.status) for x in aend]      # sonst bleibt alles wie es war
