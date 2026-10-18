"""Verlustfreies Schreiben der Excel (src/xlsx_patch.py, ADR 0007). Nur erfundene Daten (Regel 16).

Die Test-Mappe ist ein erfundenes CDE-Register mit den FORMEN echter Dateien (Auftrag 003): Kopf Zeilen 1-6 mit
verbundenen Zellen, Zaehlformeln U2:Z5, Beispielzeile 7, Daten ab Zeile 8, zerstueckelte bedingte Formatierung,
Datenpruefung (auch x14), Autofilter, Druckbereich und -titel, Hyperlink, Seitenumbruch; per ZIP-Nachbearbeitung
customXml, docMetadata/LabelInfo.xml, Druckereinstellungen, eine Zeichnung mit Form und Bild (Anker Zeile 3 und 7),
Blattschutz, calcChain, mc:AlternateContent, x14ac-Attribute und sharedStrings (statt openpyxls Inline-Text).
"""
from __future__ import annotations

import datetime as dt
import io
import re
import shutil
import subprocess
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from src.fehler import DateiVorhanden
from src.markierung import ALT, NEU, NORMAL, Stueck, ansichten, markiere_ops, stuecke_aus_zelle
from src.models import Operation
from src.xlsx_patch import (NeueZeile, PatchFehler, ZellInhalt, hindernisse, pruefe_treue, rpr, schreibe,
                            _schrift_aus_xml)

BLATT = "Register"
N = 550                       # letzte Datenzeile
ENDE = N - 8                  # Ende von Filter und Druckbereich
TEXT = "a.\t\u00a0Die Pumpe {r} läuft mit Wasser."
ALT_BLAU, ALT_ROT = "FF0070C0", "FFC00000"          # Markierungen "frueherer Runden" (andere Farben)
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_X14AC = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
pytestmark = pytest.mark.filterwarnings("ignore:Data Validation extension is not supported")   # openpyxl liest x14 nicht


# ===================================================================================== Test-Mappe bauen
def _png() -> bytes:
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (4, 2), (10, 120, 200)).save(b, "PNG")
    return b.getvalue()


def _openpyxl_register(pfad: Path, n: int = N) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    fett = Font(name="Arial", sz=11, b=True)
    ws["A1"] = "Kennung"
    ws["D1"] = "Register Beispielanlage"
    ws["H1"] = "Hinweis zum Ausfüllen"
    for rng in ("A1:A4", "D1:F4", "G1:G4", "H1:P4", "G5:G6"):
        ws.merge_cells(rng)
    for i, t in enumerate(["Section", "Sub", "Text", "Spezifikation", "Revision", "Datum", "Nr", "Status", "Prüfer"], 1):
        ws.cell(5, i, t).font = Font(b=True)
    for i, (sp, f) in enumerate([("U", 'COUNTIF(H:H,"x")'), ("V", 'COUNTIF(H:H,"y")'), ("W", 'COUNTIF(H:H,"z")')]):
        for r in range(2, 5):
            ws[f"{sp}{r}"] = "=" + f
    ws["V5"] = "=SUM(V2:V4)"
    ws.cell(7, 3, "Beispiel: so wird eine Zeile ausgefüllt.")
    for r in range(8, n + 1):
        ws.cell(r, 1, f"{(r - 8) // 5 + 1}.1")
        ws.cell(r, 2, "ABCDE"[(r - 8) % 5])
        c = ws.cell(r, 3, TEXT.format(r=r))
        c.font = fett
        ws.cell(r, 4, "SPEC-7")
        ws.cell(r, 5, 3)
        ws.cell(r, 6, dt.date(2026, 10, 1)).number_format = "DD.MM.YYYY"
        ws.cell(r, 7, r * 10)
        ws.cell(r, 8, "xyz"[r % 3])
        ws.row_dimensions[r].height = 30
    ws["C12"] = CellRichText(TextBlock(InlineFont(rFont="Arial", sz=11, b=True), "Der Kessel "),
                             TextBlock(InlineFont(rFont="Arial", sz=11, b=True, color=ALT_BLAU), "wurde geprüft "),
                             TextBlock(InlineFont(rFont="Arial", sz=11, b=True, color=ALT_ROT), "und gereinigt "),
                             TextBlock(InlineFont(rFont="Arial", sz=11, b=True), "im Herbst."))
    gelb = PatternFill("solid", bgColor="FFFF00")
    ws.conditional_formatting.add(f"F{n - 22}:F{n - 21} F{n - 18}:F{ENDE}",
                                  CellIsRule(operator="equal", formula=['"x"'], fill=gelb))
    ws.conditional_formatting.add(f"O8:O{n - 23}", FormulaRule(formula=['$O8="x"'], fill=gelb))
    ws.conditional_formatting.add("R5 R8:R1048576", FormulaRule(formula=['R5="?"'], fill=gelb))
    dv = DataValidation(type="list", formula1="Listen!$A$1:$A$3")
    ws.add_data_validation(dv)
    dv.add(f"J7:N{n - 23}")
    dv.add(f"K{n - 17}")
    ws.auto_filter.ref = f"A6:AB{ENDE}"
    ws.print_area = f"A5:P{ENDE}"
    ws.print_title_rows = "5:6"
    ws[f"C{max(9, n - 100)}"].hyperlink = "https://example.org/pumpe"
    ws.row_breaks.append(Break(id=min(200, n - 5)))
    ws.merge_cells(f"P{max(9, n - 50)}:P{max(10, n - 49)}")      # verbundene Zellen im Datenbereich (unterhalb der Stelle)
    listen = wb.create_sheet("Listen")
    for i, w in enumerate(["offen", "erledigt", "entfällt", "Pumpe", "Ventil"], 1):
        listen.cell(i, 1, w)
    ausw = wb.create_sheet("Auswertung")
    ausw["A1"] = f"='{BLATT}'!C400"
    ausw["A2"] = f"=COUNTA({BLATT}!H8:H{ENDE})"
    ausw["A3"] = f"=SUM({BLATT}!G8:G{n})"
    ausw["A4"] = f"={BLATT}!G300*2"
    from openpyxl.worksheet.hyperlink import Hyperlink
    ausw["B1"] = "Sprung"
    ausw["B1"].hyperlink = Hyperlink(ref="B1", location=f"'{BLATT}'!C400")      # interner Link auf das Register
    wb.save(str(pfad))


def _inline_zu_shared(blatt_xml: str, si: list[str]) -> str:
    def ersetzen(m: re.Match) -> str:
        si.append(f"<si>{m.group(3)}</si>")
        return f'<c {m.group(1)}t="s"{m.group(2)}><v>{len(si) - 1}</v></c>'
    return re.sub(r'<c ([^>]*?)t="inlineStr"([^>]*)><is>(.*?)</is></c>', ersetzen, blatt_xml, flags=re.S)


def _nachbearbeiten(pfad: Path, xlsm: bool = False) -> None:
    """ZIP-Nachbearbeitung: Teile, die openpyxl beim Speichern nicht schreibt (echte Dateien haben sie)."""
    with zipfile.ZipFile(pfad) as z:
        teile = {n: z.read(n) for n in z.namelist()}
    t = {n: v.decode("utf-8") for n, v in teile.items() if n.endswith((".xml", ".rels"))}
    si: list[str] = []
    for n in sorted(x for x in t if x.startswith("xl/worksheets/sheet")):
        t[n] = _inline_zu_shared(t[n], si)
    t["xl/sharedStrings.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                                 '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                                 f'count="{len(si)}" uniqueCount="{len(si)}">' + "".join(si) + "</sst>")
    s = t["xl/worksheets/sheet1.xml"]
    s = s.replace('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
                  '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                  f'xmlns:r="{NS_R}" xmlns:mc="{NS_MC}" mc:Ignorable="x14ac" xmlns:x14ac="{NS_X14AC}">', 1)
    s = re.sub(r"<row (r=\"\d+\")", r'<row \1 x14ac:dyDescent="0.25"', s)
    s = s.replace("</sheetData>", '</sheetData><sheetProtection formatCells="0"/>', 1)
    s = re.sub(r"(<pageMargins\b[^>]*/>)", r'\1<pageSetup paperSize="9" orientation="landscape" r:id="rIdPs1"/>', s, 1)
    x14 = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
           'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1" '
           'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidation type="list" allowBlank="1" '
           'showInputMessage="1" showErrorMessage="1"><x14:formula1><xm:f>Listen!$A$1:$A$5</xm:f></x14:formula1>'
           '<xm:sqref>H36:H50</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
    s = s.replace("</worksheet>", '<drawing r:id="rIdDr1"/>' + x14 + "</worksheet>")
    t["xl/worksheets/sheet1.xml"] = s
    rels = t["xl/worksheets/_rels/sheet1.xml.rels"]
    t["xl/worksheets/_rels/sheet1.xml.rels"] = rels.replace("</Relationships>", (
        f'<Relationship Id="rIdPs1" Type="{NS_R}/printerSettings" Target="../printerSettings/printerSettings1.bin"/>'
        f'<Relationship Id="rIdDr1" Type="{NS_R}/drawing" Target="../drawings/drawing1.xml"/></Relationships>'))
    t["xl/drawings/drawing1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0'
        '</xdr:rowOff></xdr:from><xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0'
        '</xdr:rowOff></xdr:to><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Form 1"/><xdr:cNvSpPr/>'
        '</xdr:nvSpPr><xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:sp><xdr:clientData/>'
        '</xdr:twoCellAnchor><xdr:oneCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>6'
        '</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="952500" cy="190500"/><xdr:pic><xdr:nvPicPr>'
        '<xdr:cNvPr id="3" name="Bild 1"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill>'
        f'<a:blip xmlns:r="{NS_R}" r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
        '</xdr:wsDr>')
    t["xl/drawings/_rels/drawing1.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/'
        f'package/2006/relationships"><Relationship Id="rId1" Type="{NS_R}/image" Target="../media/image1.png"/>'
        '</Relationships>')
    t["xl/calcChain.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<calcChain xmlns="http://schemas.'
                             'openxmlformats.org/spreadsheetml/2006/main"><c r="U2" i="1"/><c r="V2"/><c r="V5"/>'
                             '<c r="A1" i="3"/></calcChain>')
    for i in (1, 2):
        t[f"customXml/item{i}.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n<b:Sources xmlns:b="'
                                       'http://schemas.openxmlformats.org/officeDocument/2006/bibliography" '
                                       'SelectedStyle=""/>')
        t[f"customXml/itemProps{i}.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n<ds:datastoreItem '
                                            f'ds:itemID="{{0000000{i}-1111-2222-3333-444444444444}}" xmlns:ds="http://schemas.'
                                            'openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/></ds:datastoreItem>')
        t[f"customXml/_rels/item{i}.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.'
            f'org/package/2006/relationships"><Relationship Id="rId1" Type="{NS_R}/customXmlProps" '
            f'Target="itemProps{i}.xml"/></Relationships>')
    t["docMetadata/LabelInfo.xml"] = ('<?xml version="1.0" encoding="utf-8" standalone="yes"?>\n<clbl:labelList '
                                      'xmlns:clbl="http://schemas.microsoft.com/office/2020/mipLabelMetadata"><clbl:label '
                                      'id="{00000000-0000-0000-0000-00000000000a}" enabled="1" method="Standard" '
                                      'siteId="{00000000-0000-0000-0000-00000000000b}" removed="0"/></clbl:labelList>')
    wb_rels = t["xl/_rels/workbook.xml.rels"]
    zusatz = (f'<Relationship Id="rIdSst" Type="{NS_R}/sharedStrings" Target="sharedStrings.xml"/>'
              f'<Relationship Id="rIdCc" Type="{NS_R}/calcChain" Target="calcChain.xml"/>'
              f'<Relationship Id="rIdCx1" Type="{NS_R}/customXml" Target="../customXml/item1.xml"/>'
              f'<Relationship Id="rIdCx2" Type="{NS_R}/customXml" Target="../customXml/item2.xml"/>')
    if xlsm:
        zusatz += '<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" ' \
                  'Target="vbaProject.bin"/>'
    t["xl/_rels/workbook.xml.rels"] = wb_rels.replace("</Relationships>", zusatz + "</Relationships>")
    t["_rels/.rels"] = t["_rels/.rels"].replace("</Relationships>", (
        '<Relationship Id="rIdLbl" Type="http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels" '
        'Target="docMetadata/LabelInfo.xml"/></Relationships>'))
    w = t["xl/workbook.xml"]
    w = w.replace("<workbookPr />", f'<workbookPr /><mc:AlternateContent xmlns:mc="{NS_MC}"><mc:Choice Requires="x15">'
                  '<x15ac:absPath xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac" '
                  'url="C:\\Beispiel\\Ordner\\"/></mc:Choice></mc:AlternateContent>', 1)
    t["xl/workbook.xml"] = w
    ct = t["[Content_Types].xml"]
    ct = ct.replace('<Default Extension="xml"', '<Default Extension="png" ContentType="image/png" />'
                    '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
                    'printerSettings" /><Default Extension="xml"', 1)
    over = ('<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sharedStrings+xml"/><Override PartName="/xl/calcChain.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.calcChain+xml"/><Override PartName="/xl/drawings/drawing1.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
            '<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'customXmlProperties+xml"/><Override PartName="/customXml/itemProps2.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.customXmlProperties+xml"/><Override PartName="/docMetadata/LabelInfo.xml" '
            'ContentType="application/vnd.ms-office.classificationlabels+xml"/>')
    if xlsm:
        ct = ct.replace("spreadsheetml.sheet.main+xml", "PLATZHALTER")
        ct = ct.replace("application/vnd.openxmlformats-officedocument.PLATZHALTER",
                        "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        over += '<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>'
    t["[Content_Types].xml"] = ct.replace("</Types>", over + "</Types>")
    binaer = {"xl/printerSettings/printerSettings1.bin": bytes(range(256)) * 3, "xl/media/image1.png": _png()}
    if xlsm:
        binaer["xl/vbaProject.bin"] = b"\xd0\xcf\x11\xe0erfundenes VBA-Projekt" + bytes(200)
    reihe = ["[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml", "xl/workbook.xml",
             "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml", "xl/worksheets/_rels/sheet1.xml.rels",
             "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml", "xl/theme/theme1.xml", "xl/styles.xml",
             "xl/sharedStrings.xml", "xl/calcChain.xml", "xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels",
             "xl/media/image1.png", "xl/printerSettings/printerSettings1.bin", "customXml/item1.xml",
             "customXml/itemProps1.xml", "customXml/_rels/item1.xml.rels", "customXml/item2.xml",
             "customXml/itemProps2.xml", "customXml/_rels/item2.xml.rels", "docMetadata/LabelInfo.xml"]
    if xlsm:
        reihe.insert(13, "xl/vbaProject.bin")
    pfad.unlink()
    with zipfile.ZipFile(pfad, "w", zipfile.ZIP_DEFLATED) as z:
        for n in reihe:
            z.writestr(n, t[n].encode("utf-8") if n in t else binaer[n] if n in binaer else teile[n])


def baue_register(pfad: Path, n: int = N, xlsm: bool = False) -> Path:
    _openpyxl_register(pfad, n)
    _nachbearbeiten(pfad, xlsm)
    return pfad


@pytest.fixture(scope="module")
def register(tmp_path_factory) -> Path:
    return baue_register(tmp_path_factory.mktemp("register") / "register.xlsx")


def _teil(pfad: Path, name: str) -> str:
    with zipfile.ZipFile(pfad) as z:
        return z.read(name).decode("utf-8")


def _namen(pfad: Path) -> list[str]:
    with zipfile.ZipFile(pfad) as z:
        return z.namelist()


def _ws(pfad: Path, rich: bool = True):
    return load_workbook(str(pfad), rich_text=rich)[BLATT]


def _markiert(pfad: Path, zeile: int, ops: list[Operation], spalte: int = 3) -> list[Stueck]:
    zelle = load_workbook(str(pfad), rich_text=True)[BLATT].cell(zeile, spalte)
    return markiere_ops(stuecke_aus_zelle(zelle), ops).stuecke


def _ers(zeile: int, start: int, ende: int, text: str, aid: str = "A1") -> Operation:
    return Operation("ersetzen", BLATT, zeile, 3, start, ende, text, aenderung=aid)


# ===================================================================================== Die Mappe selbst
def test_die_test_mappe_hat_die_formen_echter_dateien(register):
    namen = _namen(register)
    for teil in ("customXml/item1.xml", "docMetadata/LabelInfo.xml", "xl/printerSettings/printerSettings1.bin",
                 "xl/calcChain.xml", "xl/drawings/drawing1.xml", "xl/media/image1.png", "xl/sharedStrings.xml"):
        assert teil in namen
    s = _teil(register, "xl/worksheets/sheet1.xml")
    assert '<sheetProtection formatCells="0"/>' in s and "x14ac:dyDescent" in s and "<xm:sqref>H36:H50</xm:sqref>" in s
    assert 't="inlineStr"' not in s and 't="s"' in s
    assert "mc:AlternateContent" in _teil(register, "xl/workbook.xml")
    wb = load_workbook(str(register), rich_text=True)          # openpyxl liest die Mappe
    assert wb[BLATT]["C8"].value.startswith("a.\t")


def test_openpyxl_verliert_beim_speichern_was_echte_dateien_haben(register, tmp_path):
    """Darum ADR 0007: Laden + Speichern mit openpyxl laesst Teile weg."""
    wb = load_workbook(str(register))
    wb.save(str(tmp_path / "openpyxl.xlsx"))
    namen = _namen(tmp_path / "openpyxl.xlsx")
    assert "docMetadata/LabelInfo.xml" not in namen and "xl/printerSettings/printerSettings1.bin" not in namen
    assert not any(n.startswith("customXml/") for n in namen)


# ===================================================================================== Zellen
def test_zelle_markiert_unberuehrte_teile_bytegleich(register, tmp_path):
    t = TEXT.format(r=20)
    s = t.index("Wasser")
    stuecke = _markiert(register, 20, [_ers(20, s, s + 6, "Öl")])
    erg = schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 20, 3, stuecke=stuecke)])
    assert pruefe_treue(register, tmp_path / "out.xlsx", erg.geaenderte_teile) == []
    assert set(erg.geaenderte_teile) == {"xl/worksheets/sheet1.xml", "xl/sharedStrings.xml", "xl/calcChain.xml",
                                         "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
    alt, neu = _namen(register), _namen(tmp_path / "out.xlsx")
    assert neu == [n for n in alt if n != "xl/calcChain.xml"]           # gleiche Teile, gleiche Reihenfolge
    with zipfile.ZipFile(register) as a, zipfile.ZipFile(tmp_path / "out.xlsx") as b:
        for n in neu:
            if n not in erg.geaenderte_teile:
                assert a.read(n) == b.read(n), n
    assert "calcChain" not in _teil(tmp_path / "out.xlsx", "[Content_Types].xml")
    assert "calcChain" not in _teil(tmp_path / "out.xlsx", "xl/_rels/workbook.xml.rels")
    zaehler = lambda pfad: re.search(r'count="(\d+)" uniqueCount="(\d+)"', _teil(pfad, "xl/sharedStrings.xml")).groups()  # noqa: E731
    (c0, u0), (c1, u1) = zaehler(register), zaehler(tmp_path / "out.xlsx")
    assert (int(c1), int(u1)) == (int(c0), int(u0) + 1)      # ein neuer Eintrag, die Zelle zeigt weiter auf EINEN Text
    # nur die eine Zelle hat sich im Blatt geaendert
    a = _teil(register, "xl/worksheets/sheet1.xml")
    b = _teil(tmp_path / "out.xlsx", "xl/worksheets/sheet1.xml")
    za = re.search(r'<row r="20".*?</row>', a).group(0)
    assert a.replace(za, "") == b.replace(re.search(r'<row r="20".*?</row>', b).group(0), "")


def test_markierte_zelle_liest_sich_wie_nicos_soll(register, tmp_path):
    t = TEXT.format(r=20)
    s = t.index("Wasser")
    stuecke = _markiert(register, 20, [_ers(20, s, s + 6, "Öl")])
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 20, 3, stuecke=stuecke)])
    z = _ws(tmp_path / "out.xlsx")["C20"]
    alt, neu, st = ansichten(z)
    assert alt == t and neu == t.replace("Wasser", "Öl")
    assert [(x.format, x.text) for x in st] == [(NORMAL, t[:s]), (ALT, "Wasser"), (NEU, "Öl"), (NORMAL, ".")]
    for b in z.value:                                          # Grundschrift (Arial 11 fett) in JEDEM Lauf
        assert b.font.rFont == "Arial" and b.font.sz == 11 and b.font.b
    blocke = list(z.value)
    assert blocke[1].font.strike and blocke[1].font.color.rgb == "FF00B0F0"
    assert not blocke[2].font.strike and blocke[2].font.color.rgb == "FFFF0000"
    assert blocke[0].text.startswith("a.\t\u00a0")            # Listennummer + Polster bleiben normal


def test_alte_markierungen_bleiben_unveraendert(register, tmp_path):
    zelle = load_workbook(str(register), rich_text=True)[BLATT]["C12"]
    vorher = [(b.text, b.font.color.rgb if b.font.color is not None else None) for b in zelle.value]
    text = str(zelle.value)
    s = text.index("Herbst")
    stuecke = markiere_ops(stuecke_aus_zelle(zelle), [_ers(12, s, s + 6, "Frühjahr")]).stuecke
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 12, 3, stuecke=stuecke)])
    nachher = list(_ws(tmp_path / "out.xlsx")["C12"].value)
    farben = [(b.text, b.font.color.rgb if b.font.color is not None else None, b.font.strike) for b in nachher]
    assert farben[1] == ("wurde geprüft ", ALT_BLAU, None) or farben[1][:2] == ("wurde geprüft ", ALT_BLAU)
    assert farben[2][:2] == ("und gereinigt ", ALT_ROT) and not farben[2][2]
    assert farben[-3][1:] == ("FF00B0F0", True) and farben[-2][:2] == ("Frühjahr", "FFFF0000")
    assert vorher[1][1] == ALT_BLAU


def test_zahl_bleibt_zahl_nur_schriftfarbe(register, tmp_path):
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 30, 7, wert=301, farbe=NEU),
                                               ZellInhalt(BLATT, 31, 5, wert=3, farbe=ALT)])
    ws = _ws(tmp_path / "out.xlsx")
    assert ws["G30"].value == 301 and ws["G30"].data_type == "n" and ws["G30"].font.color.rgb == "FFFF0000"
    assert ws["E31"].value == 3 and ws["E31"].font.strike and ws["E31"].font.color.rgb == "FF00B0F0"
    assert ws["G32"].value == 320 and ws["G32"].font.color is None or ws["G32"].font.color.rgb != "FFFF0000"
    assert ws["F30"].value == dt.datetime(2026, 10, 1) and ws["F30"].number_format == "DD.MM.YYYY"
    styles_alt = _teil(register, "xl/styles.xml")
    styles_neu = _teil(tmp_path / "out.xlsx", "xl/styles.xml")
    alt_fonts = re.search(r"<fonts\b.*?</fonts>", styles_alt, re.S).group(0)
    neu_fonts = re.search(r"<fonts\b.*?</fonts>", styles_neu, re.S).group(0)
    inhalt_alt = re.sub(r"<fonts[^>]*>", "", alt_fonts)[:-len("</fonts>")]
    assert inhalt_alt in neu_fonts                             # vorhandene Schriften unveraendert, neue nur hinten


def test_text_escapes_und_sonderzeichen(register, tmp_path):
    text = 'Größe <10 & "Druck" > 5 bar _x000D_ ende\r'
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 40, 3, stuecke=[Stueck(NORMAL, "Alt "),
                                                                                   Stueck(NEU, text)])])
    z = _ws(tmp_path / "out.xlsx")["C40"]
    assert str(z.value) == "Alt " + text                     # Rundreise in der openpyxl-Lesart (woertliches _x000D_)
    sst = _teil(tmp_path / "out.xlsx", "xl/sharedStrings.xml")
    assert "&lt;10 &amp; " in sst and "_x005F_x000D_" in sst and "&#13;" in sst


def test_leere_und_neue_zellen_in_richtiger_reihenfolge(register, tmp_path):
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 50, 12, wert="neu hinten"),
                                               ZellInhalt(BLATT, 50, 1, wert=None),
                                               ZellInhalt(BLATT, 6, 30, wert=7.5),
                                               ZellInhalt(BLATT, 600, 2, wert="ganz neue Zeile")])
    s = _teil(tmp_path / "out.xlsx", "xl/worksheets/sheet1.xml")
    zeile = re.search(r'<row r="50".*?</row>', s).group(0)
    refs = re.findall(r'<c r="([A-Z]+)50"', zeile)
    assert refs == sorted(refs, key=lambda x: (len(x), x)) and "L" in refs
    ws = _ws(tmp_path / "out.xlsx")
    assert ws["A50"].value is None and ws["L50"].value == "neu hinten" and ws["AD6"].value == 7.5
    assert ws["B600"].value == "ganz neue Zeile"
    nummern = [int(x) for x in re.findall(r'<row r="(\d+)"', s)]
    assert nummern == sorted(nummern)


def test_sharedstrings_werden_angelegt_wenn_es_keine_gibt(tmp_path):
    wb = Workbook()
    wb.active.title = BLATT
    wb.active["A1"] = 1
    wb.active["B2"] = "Text"                                   # openpyxl schreibt Inline-Text, keine sharedStrings
    wb.save(str(tmp_path / "a.xlsx"))
    assert "xl/sharedStrings.xml" not in _namen(tmp_path / "a.xlsx")
    erg = schreibe(tmp_path / "a.xlsx", tmp_path / "b.xlsx", [ZellInhalt(BLATT, 1, 2, stuecke=[Stueck(NEU, "rot")])])
    assert pruefe_treue(tmp_path / "a.xlsx", tmp_path / "b.xlsx", erg.geaenderte_teile) == []
    assert "xl/sharedStrings.xml" in _namen(tmp_path / "b.xlsx")
    assert "sharedStrings" in _teil(tmp_path / "b.xlsx", "[Content_Types].xml")
    ws = load_workbook(str(tmp_path / "b.xlsx"), rich_text=True)[BLATT]
    assert str(ws["B1"].value) == "rot" and ws["B2"].value == "Text" and ws["A1"].value == 1


def test_text_mit_gleichheitszeichen_bleibt_text(register, tmp_path):
    schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 41, 3, wert="=Pumpe", text_praefix=True)])
    z = load_workbook(str(tmp_path / "out.xlsx"))[BLATT]["C41"]
    assert z.data_type == "s" and z.value == "=Pumpe" and z.quotePrefix


def test_zieldatei_wird_nie_ueberschrieben(register, tmp_path):
    (tmp_path / "out.xlsx").write_bytes(b"schon da")
    with pytest.raises(DateiVorhanden):
        schreibe(register, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 41, 3, wert="x")])
    assert (tmp_path / "out.xlsx").read_bytes() == b"schon da"


def test_ohne_aenderung_bytegleiche_kopie(register, tmp_path):
    erg = schreibe(register, tmp_path / "out.xlsx")
    assert erg.geaenderte_teile == [] and pruefe_treue(register, tmp_path / "out.xlsx") == []


def test_pruefe_treue_meldet_veraenderte_und_fehlende_teile(register, tmp_path):
    with zipfile.ZipFile(register) as a, zipfile.ZipFile(tmp_path / "kaputt.xlsx", "w") as b:
        for n in a.namelist():
            if n == "docMetadata/LabelInfo.xml":
                continue
            b.writestr(n, a.read(n) + (b" " if n == "customXml/item1.xml" else b""))
    m = pruefe_treue(register, tmp_path / "kaputt.xlsx")
    assert any("LabelInfo" in x for x in m) and any("item1.xml" in x for x in m)


def test_rpr_reihenfolge_wie_excel():
    s = _schrift_aus_xml('<font><name val="Arial"/><family val="2"/><b/><sz val="11"/><color theme="1"/></font>')
    assert rpr(s) == '<rPr><b/><sz val="11"/><color theme="1"/><rFont val="Arial"/><family val="2"/></rPr>'
    assert _schrift_aus_xml('<font><b val="0"/><i/></font>') == {"i": ""}


# ===================================================================================== neue Zeilen
def _formel(pfad: Path, blatt: str, ref: str) -> str:
    return load_workbook(str(pfad))[blatt][ref].value


def test_zeile_einfuegen_verschiebt_alle_bezuege(register, tmp_path):
    nach = 300
    erg = schreibe(register, tmp_path / "out.xlsx", neue_zeilen=[
        NeueZeile(BLATT, nach, nach, {3: ZellInhalt(BLATT, 0, 3, stuecke=[Stueck(NEU, "Neuer Absatz.")])})])
    out = tmp_path / "out.xlsx"
    assert pruefe_treue(register, out, erg.geaenderte_teile) == []
    assert erg.neue_zeilen == [(BLATT, nach + 1)] and erg.zeile(BLATT, nach) == nach and erg.zeile(BLATT, nach + 1) == nach + 2
    s = _teil(out, "xl/worksheets/sheet1.xml")
    assert '<dimension ref="A1:W551"' in s
    assert f'sqref="F{N - 21}:F{N - 20} F{N - 17}:F{ENDE + 1}"' in s           # beide Stuecke liegen darunter
    assert f'sqref="O8:O{N - 22}"' in s and 'sqref="R5 R8:R1048576"' in s
    assert f'sqref="J7:N{N - 22} K{N - 16}"' in s and "<xm:sqref>H36:H50</xm:sqref>" in s
    assert f'<autoFilter ref="A6:AB{ENDE + 1}"' in s
    assert f'<mergeCell ref="P{N - 49}:P{N - 48}"' in s and '<mergeCell ref="A1:A4"' in s
    assert f'<hyperlink xmlns:r="{NS_R}" ref="C{N - 99}"' in s
    assert '<brk id="200"' in s                                  # Umbruch nach Zeile 200 liegt oberhalb
    w = _teil(out, "xl/workbook.xml")
    assert f"'{BLATT}'!$A$5:$P${ENDE + 1}" in w and f"'{BLATT}'!$A$6:$AB${ENDE + 1}" in w and f"'{BLATT}'!$5:$6" in w
    assert w.replace(f"{ENDE + 1}<", f"{ENDE}<") == _teil(register, "xl/workbook.xml")   # sonst bytegleich (mc:AlternateContent)
    assert _formel(out, "Auswertung", "A1") == f"='{BLATT}'!C401"
    assert _formel(out, "Auswertung", "A2") == f"=COUNTA({BLATT}!H8:H{ENDE + 1})"
    assert _formel(out, "Auswertung", "A4") == f"={BLATT}!G300*2"            # Vorlagezeile 300 liegt oberhalb: bleibt
    assert _formel(out, BLATT, "U2") == '=COUNTIF(H:H,"x")' and _formel(out, BLATT, "V5") == "=SUM(V2:V4)"
    assert f'location="&apos;{BLATT}&apos;!C401"' in _teil(out, "xl/worksheets/sheet3.xml") or \
        f"location=\"'{BLATT}'!C401\"" in _teil(out, "xl/worksheets/sheet3.xml")       # interner Link zieht mit
    d = _teil(out, "xl/drawings/drawing1.xml")
    assert d == _teil(register, "xl/drawings/drawing1.xml")      # Anker (Zeile 3 und 7) liegen oberhalb: unveraendert
    ws = _ws(out)
    assert str(ws.cell(nach + 1, 3).value) == "Neuer Absatz." and ws.cell(nach + 2, 3).value == TEXT.format(r=nach + 1)
    assert ws.cell(N + 1, 3).value == TEXT.format(r=N)
    assert ws.row_dimensions[nach + 1].height == 30                # Hoehe wie die Vorlage
    assert ws.cell(nach + 1, 3).font.b and ws.cell(nach + 1, 3).font.name == "Arial"      # Zellstil wie die Vorlage
    zeile = re.search(rf'<row r="{nach + 1}".*?</row>', s).group(0)
    assert 'x14ac:dyDescent="0.25"' in zeile


def test_zeile_einfuegen_im_kopf_verschiebt_zeichnungs_anker(register, tmp_path):
    assert any("verbundene Zellen durch" in h for h in hindernisse(register, BLATT, 6))     # G5:G6 liefe durch die Stelle
    erg = schreibe(register, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, 6, 6, {3: "Zwischenzeile"})])
    d = _teil(tmp_path / "out.xlsx", "xl/drawings/drawing1.xml")
    anker = re.findall(r"<xdr:row>(\d+)</xdr:row>", d)
    assert anker == ["2", "3", "7"]                              # Form (Zeile 3-4) bleibt, Bild (Zeile 7) rutscht auf 8
    assert "xl/drawings/drawing1.xml" in erg.geaenderte_teile
    assert pruefe_treue(register, tmp_path / "out.xlsx", erg.geaenderte_teile) == []
    w = _teil(tmp_path / "out.xlsx", "xl/workbook.xml")
    assert f"'{BLATT}'!$A$5:$P${ENDE + 1}" in w and f"'{BLATT}'!$5:$6" in w and f"'{BLATT}'!$A$6:$AB${ENDE + 1}" in w
    s = _teil(tmp_path / "out.xlsx", "xl/worksheets/sheet1.xml")
    assert '<mergeCell ref="G5:G6"' in s and 'sqref="R5 R9:R1048576"' in s and 'sqref="J8:N' in s


def test_zeile_am_ende_waechst_mit_der_vorlage(register, tmp_path):
    erg = schreibe(register, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, ENDE, ENDE, {3: "Am Ende."})])
    s = _teil(tmp_path / "out.xlsx", "xl/worksheets/sheet1.xml")
    assert f'<autoFilter ref="A6:AB{ENDE + 1}"' in s and f"F{N - 18}:F{ENDE + 1}" in s
    w = _teil(tmp_path / "out.xlsx", "xl/workbook.xml")
    assert f"$A$5:$P${ENDE + 1}" in w
    assert _formel(tmp_path / "out.xlsx", "Auswertung", "A2") == f"=COUNTA({BLATT}!H8:H{ENDE + 1})"
    assert any("umfassen jetzt auch die neue Zeile" in x for x in erg.warnungen)


def test_mehrere_neue_zeilen_in_reihenfolge(register, tmp_path):
    erg = schreibe(register, tmp_path / "out.xlsx", neue_zeilen=[
        NeueZeile(BLATT, 100, 100, {3: "zweite"}), NeueZeile(BLATT, 20, 20, {3: "erste"}),
        NeueZeile(BLATT, 100, 100, {3: "dritte"})])
    assert erg.neue_zeilen == [(BLATT, 102), (BLATT, 21), (BLATT, 103)]
    ws = _ws(tmp_path / "out.xlsx", rich=False)
    assert [ws.cell(r, 3).value for r in (21, 102, 103)] == ["erste", "zweite", "dritte"]
    assert ws.cell(104, 3).value == TEXT.format(r=101) and ws.cell(22, 3).value == TEXT.format(r=21)
    assert _formel(tmp_path / "out.xlsx", "Auswertung", "A1") == f"='{BLATT}'!C403"


def test_formel_der_vorlage_wird_uebersetzt_und_konstante_kopiert(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    for r in range(1, 8):
        ws.cell(r, 1, f"Text {r}")
        ws.cell(r, 2, f"=LEN(A{r})")
        ws.cell(r, 3, "SPEC-7")
    ws["B10"] = "=SUM(B1:B7)"
    wb.save(str(tmp_path / "a.xlsx"))
    schreibe(tmp_path / "a.xlsx", tmp_path / "b.xlsx", neue_zeilen=[NeueZeile(BLATT, 4, 4, {
        1: "Neu", 2: ZellInhalt(BLATT, 0, 2, formel="=LEN(A4)", formel_von=4), 3: ZellInhalt(BLATT, 0, 3, kopie_von=4)})])
    ws = load_workbook(str(tmp_path / "b.xlsx"))[BLATT]
    assert [ws.cell(r, 2).value for r in range(1, 9)] == [f"=LEN(A{r})" for r in range(1, 9)]
    assert ws["C5"].value == "SPEC-7" and ws["A5"].value == "Neu" and ws["B11"].value == "=SUM(B1:B8)"


# ===================================================================================== Hindernisse
def _blatt_mit(tmp_path, name: str, fn) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    for r in range(1, 21):
        ws.cell(r, 1, f"Zeile {r}")
    fn(wb, ws)
    p = tmp_path / f"{name}.xlsx"
    wb.save(str(p))
    return p


def test_hindernisse_sauberes_blatt_ist_frei(register, tmp_path):
    assert hindernisse(register, BLATT, 301) == []
    assert hindernisse(register, BLATT, ENDE + 1) == []
    assert "gibt es in der Excel nicht" in hindernisse(register, "Fehlt", 3)[0]


def _geteilt(p: Path) -> None:
    """openpyxl schreibt keine geteilten Formeln: D2:D6 als geteilte Formel nachbauen."""
    with zipfile.ZipFile(p) as z:
        teile = {n: z.read(n) for n in z.namelist()}
    s = teile["xl/worksheets/sheet1.xml"].decode()
    for r in range(2, 7):
        f = '<f t="shared" ref="D2:D6" si="0">B2*2</f>' if r == 2 else '<f t="shared" si="0"/>'
        s = re.sub(rf'(<row r="{r}"[^>]*>)', rf'\1<c r="D{r}">{f}<v>0</v></c>', s, 1)
    teile["xl/worksheets/sheet1.xml"] = s.encode()
    with zipfile.ZipFile(p, "w") as z:
        for n, v in teile.items():
            z.writestr(n, v)


def test_geteilte_formel_ueber_die_stelle_ist_ein_hindernis(tmp_path):
    p = _blatt_mit(tmp_path, "geteilt", lambda wb, ws: None)
    _geteilt(p)
    assert any("geteilte Formel" in h for h in hindernisse(p, BLATT, 4))
    assert hindernisse(p, BLATT, 10) == [] and hindernisse(p, BLATT, 2) == []
    with pytest.raises(PatchFehler):
        schreibe(p, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, 3, 3, {1: "x"})])
    assert not (tmp_path / "out.xlsx").exists()                 # keine halbe Datei


def test_geteilte_formel_darunter_wandert_mit(tmp_path):
    p = _blatt_mit(tmp_path, "geteilt2", lambda wb, ws: None)
    _geteilt(p)
    schreibe(p, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, 1, 1, {1: "neu"})])
    ws = load_workbook(str(tmp_path / "out.xlsx"))[BLATT]
    assert [ws.cell(r, 4).value for r in range(3, 8)] == [f"=B{r}*2" for r in range(3, 8)]


@pytest.mark.parametrize("art, erwartet", [
    ("tabelle", "Excel-Tabelle"), ("indirekt", "INDIREKT"), ("bereich", "INDIREKT"), ("drei_d", "3D-Bezug"),
    ("verbunden", "verbundene Zellen durch"), ("versteckt", "ausgeblendeter"), ("matrix", "Matrixformel"),
    ("kommentar", "Kommentare")])
def test_hindernisse_einzeln(tmp_path, art, erwartet):
    from openpyxl.comments import Comment
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.table import Table

    def bau(wb, ws):
        if art == "tabelle":
            ws.add_table(Table(displayName="Tab1", ref="A1:A20"))
        elif art == "indirekt":
            ws["C1"] = '=INDIRECT("A"&5)'
        elif art == "bereich":
            wb.create_sheet("Andere")["A1"] = f"=SUM(OFFSET({BLATT}!A1,2,0,3,1))"
        elif art == "drei_d":
            wb.create_sheet("Zwei")
            wb.create_sheet("Andere")["A1"] = f"=SUM({BLATT}:Zwei!A1)"
        elif art == "verbunden":
            ws.merge_cells("B3:B6")
        elif art == "versteckt":
            ws.row_dimensions[4].hidden = True
            ws.row_dimensions[5].hidden = True
        elif art == "matrix":
            ws["E3"] = ArrayFormula("E3:E6", "=A3:A6")
        elif art == "kommentar":
            ws["A15"].comment = Comment("Notiz", "Prüfer")
    p = _blatt_mit(tmp_path, art, bau)
    h = hindernisse(p, BLATT, 5)
    assert any(erwartet in x for x in h), h


def test_externe_verknuepfung_und_pivot_sind_hindernisse(tmp_path):
    p = _blatt_mit(tmp_path, "extern", lambda wb, ws: None)
    with zipfile.ZipFile(p) as z:
        teile = {n: z.read(n) for n in z.namelist()}
    r = teile["xl/_rels/workbook.xml.rels"].decode()
    r = r.replace("</Relationships>", f'<Relationship Id="rIdE" Type="{NS_R}/externalLink" '
                  'Target="externalLinks/externalLink1.xml"/><Relationship Id="rIdP" Type="'
                  f'{NS_R}/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>')
    teile["xl/_rels/workbook.xml.rels"] = r.encode()
    teile["xl/externalLinks/externalLink1.xml"] = b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'
    teile["xl/pivotCache/pivotCacheDefinition1.xml"] = (
        b'<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><cacheSource '
        b'type="worksheet"><worksheetSource ref="A1:A20" sheet="Register"/></cacheSource></pivotCacheDefinition>')
    with zipfile.ZipFile(p, "w") as z:
        for n, v in teile.items():
            z.writestr(n, v)
    h = hindernisse(p, BLATT, 5)
    assert any("verknüpft" in x for x in h) and any("Pivot" in x for x in h)


# ===================================================================================== .xlsm, LibreOffice, Laufzeit
def test_xlsm_vba_bleibt_bytegleich(tmp_path):
    p = baue_register(tmp_path / "makro.xlsm", n=60, xlsm=True)
    t = TEXT.format(r=20)
    s = t.index("Wasser")
    zelle = load_workbook(str(p), rich_text=True, keep_vba=True)[BLATT]["C20"]
    stuecke = markiere_ops(stuecke_aus_zelle(zelle), [_ers(20, s, s + 6, "Öl")]).stuecke
    erg = schreibe(p, tmp_path / "out.xlsm", [ZellInhalt(BLATT, 20, 3, stuecke=stuecke)],
                   [NeueZeile(BLATT, 30, 30, {3: "neu"})])
    assert pruefe_treue(p, tmp_path / "out.xlsm", erg.geaenderte_teile) == []
    with zipfile.ZipFile(p) as a, zipfile.ZipFile(tmp_path / "out.xlsm") as b:
        assert a.read("xl/vbaProject.bin") == b.read("xl/vbaProject.bin")
    wb = load_workbook(str(tmp_path / "out.xlsm"), keep_vba=True)
    assert wb[BLATT]["C31"].value == "neu"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice nicht installiert")
def test_libreoffice_oeffnet_und_zeigt_hellblau_strich_und_rot(tmp_path):
    import pdfplumber
    p = baue_register(tmp_path / "klein.xlsx", n=40)
    t = TEXT.format(r=20)
    s = t.index("Wasser")
    stuecke = _markiert(p, 20, [_ers(20, s, s + 6, "Öl")])
    schreibe(p, tmp_path / "out.xlsx", [ZellInhalt(BLATT, 20, 3, stuecke=stuecke)],
             [NeueZeile(BLATT, 25, 25, {3: ZellInhalt(BLATT, 0, 3, stuecke=[Stueck(NEU, "Zebrafink")])})])
    profil = tmp_path / "lo"
    subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{profil}", "--convert-to", "pdf",
                    "--outdir", str(tmp_path), str(tmp_path / "out.xlsx")], check=True, capture_output=True, timeout=240)
    with pdfplumber.open(tmp_path / "out.pdf") as pdf:
        zeichen = [c for seite in pdf.pages for c in seite.chars]
        linien = sum(len(seite.lines) + len(seite.rects) for seite in pdf.pages)
        text = "".join(c["text"] for c in zeichen)
    i = text.index("WasserÖl")
    wasser, oel = zeichen[i:i + 6], zeichen[i + 6:i + 8]
    assert all(tuple(round(v, 2) for v in c["non_stroking_color"]) == (0.0, 0.69, 0.94) for c in wasser)
    assert all(tuple(c["non_stroking_color"]) == (1.0, 0.0, 0.0) for c in oel)
    assert linien >= 1                                            # Durchstreichung
    assert all("Bold" in c["fontname"] for c in zeichen[i - 8:i + 8])   # Grundschrift fett in jedem Lauf
    j = text.index("Zebrafink")
    assert all(tuple(c["non_stroking_color"]) == (1.0, 0.0, 0.0) for c in zeichen[j:j + 9])


def test_laufzeit_550_zeilen_60_aenderungen_unter_2_sekunden(register, tmp_path):
    wb = load_workbook(str(register), rich_text=True)
    ws = wb[BLATT]
    zellen = []
    for k in range(60):
        r = 8 + k * 8
        t = TEXT.format(r=r)
        s = t.index("Wasser")
        zellen.append(ZellInhalt(BLATT, r, 3, stuecke=markiere_ops(stuecke_aus_zelle(ws.cell(r, 3)),
                                                                     [_ers(r, s, s + 6, "Öl")]).stuecke))
    neue = [NeueZeile(BLATT, 100 + 50 * k, 100 + 50 * k, {3: ZellInhalt(BLATT, 0, 3, stuecke=[Stueck(NEU, "neu")])})
            for k in range(5)]
    start = time.perf_counter()
    erg = schreibe(register, tmp_path / "out.xlsx", zellen, neue)
    dauer = time.perf_counter() - start
    assert dauer < 2.0, dauer
    assert pruefe_treue(register, tmp_path / "out.xlsx", erg.geaenderte_teile) == []


def test_neue_zeile_uebernimmt_verbundene_zellen_der_vorlagezeile(tmp_path):
    def bau(wb, ws):
        for r in range(2, 8):
            ws.merge_cells(f"B{r}:C{r}")                        # je Zeile verbunden (Form des Registers)
    p = _blatt_mit(tmp_path, "zeilenmerge", bau)
    erg = schreibe(p, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, 4, 4, {1: "neu"})])
    ws = load_workbook(str(tmp_path / "out.xlsx"))[BLATT]
    bereiche = sorted(str(r) for r in ws.merged_cells.ranges)
    assert "B5:C5" in bereiche and "B8:C8" in bereiche and len(bereiche) == 7
    assert pruefe_treue(p, tmp_path / "out.xlsx", erg.geaenderte_teile) == []


def test_laufende_summe_waechst_nicht_summenzeile_schon(tmp_path):
    """Liste D2:D6 mit laufender Summe in E und einer Summenzeile darunter; neue Zeile am Ende (Vorlage 6):
    die Summe ueber die ganze Liste waechst mit (zeilen.listen_summen_erweitern), die laufende Summe der Vorlage nicht."""
    def bau(wb, ws):
        for r in range(2, 7):
            ws.cell(r, 4, r)
            ws.cell(r, 5, f"=SUM($D$2:D{r})")
        ws["D8"] = "=SUM(D2:D6)"
        wb.create_sheet("Andere")["A1"] = f"=SUM({BLATT}!D2:D6)"
    p = _blatt_mit(tmp_path, "summen", bau)
    erg = schreibe(p, tmp_path / "out.xlsx", neue_zeilen=[NeueZeile(BLATT, 6, 6, {
        1: "neu", 5: ZellInhalt(BLATT, 0, 5, formel="=SUM($D$2:D6)", formel_von=6)})])
    wb = load_workbook(str(tmp_path / "out.xlsx"))
    ws = wb[BLATT]
    assert [ws.cell(r, 5).value for r in range(2, 8)] == [f"=SUM($D$2:D{r})" for r in range(2, 8)]
    assert ws["D9"].value == "=SUM(D2:D7)" and wb["Andere"]["A1"].value == f"=SUM({BLATT}!D2:D7)"
    assert any("umfassen jetzt auch die neue Zeile" in w for w in erg.warnungen)
