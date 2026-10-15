"""Erfundenes Testszenario (Regel 16: keine echten Inhalte).

Ein Szenario ist eine Liste von Elementen ('h' = Ueberschrift, 'p' = Absatz) mit Markup:
  {+neu+}  {-entfaellt-}
Daraus entstehen:
  * das Aenderungs-PDF (alle Markierungen sichtbar)
  * die "alte" Excel (Zustand vor den Aenderungen) in verschiedenen Layouts
  * die erwartete "neue" Excel (Loesung), mit demselben Layout gebaut
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.textutil import kopf_zerlegen
from .pdf_writer import baue_pdf, zustand

ELEMENTE: list[tuple[str, str]] = [
    ("h", "1 Einleitung"),
    ("p", "Dieses Dokument beschreibt die Anforderungen an das Speichersystem. {+Es gilt für alle Standorte des Kunden.+}"),
    ("p", "Das System {-muss-} {+soll+} die Betriebsdaten speichern. Die nutzbare Kapazität beträgt {-100 kWh-} {+120 kWh+} je Rack."),
    ("p", "Die Rack{+s+} werden vor dem Versand einzeln geprüft."),
    ("h", "2 Anforderungen"),
    ("h", "2.1 Leistung"),
    ("p", "Die Nennleistung beträgt 50 kW. Die maximale Entladeleistung darf kurzzeitig überschritten werden."),
    ("p", "{+Die Leistung wird im Betrieb ständig überwacht.+}"),
    ("h", "{-2.2 Wartung-}"),
    ("p", "{-Die Wartung erfolgt jährlich durch den Betreiber.-}"),
    ("p", "{-Ersatzteile werden fünf Jahre vorgehalten.-}"),
    ("h", "{-2.3-} {+2.2+} Sicherheit"),
    ("p", "Das System muss gegen Überhitzung geschützt sein. {-Ein Brandschutzkonzept ist nicht erforderlich.-}"),
    ("h", "{+2.3 Dokumentation+}"),
    ("p", "{+Alle Unterlagen werden in deutscher Sprache geliefert.+}"),
    ("p", "{+Prüfberichte liegen dem Angebot bei.+}"),
    ("h", "3 Lieferung"),
    ("p", "Die Lieferung erfolgt frei Haus. {-Der Termin wird später genannt.-} Die Inbetriebnahme dauert {-zwei-} {+drei+} Wochen."),
    ("p", "Hinweise zur Verpackung bleiben unverändert und werden nicht angefasst."),
]


def zustaende(elemente=ELEMENTE, neu: bool = False) -> list[tuple[str, str, list[str]]]:
    """-> [(nummer, titel, [absaetze])] im alten oder neuen Stand."""
    sections: list[tuple[str, str, list[str]]] = []
    for art, markup in elemente:
        text = zustand(markup, neu)
        if not text:
            continue
        if art == "h":
            nr, titel = kopf_zerlegen(text) or ("", text)
            sections.append((nr, titel, []))
        else:
            sections[-1][2].append(text)
    return sections


# ---------------------------------------------------------------- Excel
LAYOUTS = {
    # name: verwaltete Spalten (Inhalt, den das Tool aendern darf)
    "L1": (1, 2, 3),   # Nr | Titel | Inhalt | Bemerkung       (eine Zeile je Section)
    "L2": (1, 2),      # "Nr Titel" | Inhalt | Kommentar        (Nummer und Titel zusammen)
    "L3": (1, 2, 3),   # Nr | Titel | Text | Prio               (eine Zeile je Absatz)
}


def baue_excel(pfad, sections, layout: str = "L1", deckblatt: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    if deckblatt:
        ws.title = "Deckblatt"
        ws["A1"] = "Beispielprojekt (erfunden)"
        ws["A2"] = "Dieses Blatt gehört nicht zur Gliederung und bleibt unverändert."
        ws = wb.create_sheet("CDE")
    else:
        ws.title = "CDE"
    kopf = {"L1": ["Nr.", "Titel", "Inhalt", "Bemerkung"], "L2": ["Kapitel", "Anforderung", "Kommentar"],
            "L3": ["Nr", "Titel", "Text", "Prio"]}[layout]
    ws.append(kopf)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for nr, titel, absaetze in sections:
        if layout == "L1":
            ws.append([nr, titel, "\n".join(absaetze), "Status offen"])
        elif layout == "L2":
            ws.append([f"{nr} {titel}", "\n".join(absaetze), "Status offen"])
        else:
            ws.append([nr, titel, None, "A"])
            for a in absaetze:
                ws.append([None, None, a, "B"])
    breiten = {"L1": (8, 22, 70, 14), "L2": (26, 70, 14), "L3": (8, 22, 70, 8)}[layout]
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = b
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(str(pfad))


def inhalt_zeilen(pfad, layout: str) -> list[tuple[str, ...]]:
    """Alle verwalteten Spalten je Zeile (leere Zeilen weggelassen), Blatt 'CDE', ohne Kopfzeile."""
    from openpyxl import load_workbook
    ws = load_workbook(str(pfad))["CDE"]
    spalten = LAYOUTS[layout]
    out = []
    for r in range(2, ws.max_row + 1):
        werte = tuple(str(ws.cell(r, c).value or "").strip() for c in spalten)
        if any(werte):
            out.append(werte)
    return out


def erzeuge(verzeichnis: Path, variante: str = "farbe", layout: str = "L1", elemente=ELEMENTE,
            deckblatt: bool = True) -> dict[str, Path]:
    verzeichnis.mkdir(parents=True, exist_ok=True)
    pdf, alt, neu = verzeichnis / "aenderungen.pdf", verzeichnis / "cde_alt.xlsx", verzeichnis / "cde_soll.xlsx"
    baue_pdf(pdf, elemente, variante)
    baue_excel(alt, zustaende(elemente, neu=False), layout, deckblatt)
    baue_excel(neu, zustaende(elemente, neu=True), layout, deckblatt)
    return {"pdf": pdf, "alt": alt, "soll": neu}
