"""Generator fuer erfundene Testpaare mit bekannter Wahrheit (Regel 16: nur erfundene Daten).

Aus einem LOGISCHEN MODELL (Sections, Absaetze, Aenderungen) entsteht ein Paar-Ordner nach
docs/paar_format.md:  aenderungen.pdf, vorher.xlsx, nachher.xlsx, erwartet.json

WICHTIG: vorher.xlsx / nachher.xlsx entstehen direkt aus dem Modell (Wahrheit) -- nie durch den
PDF-Leser oder ein anderes Modul aus src/. Dieses Modul importiert nichts aus src/.

Aufruf:
    from tests.fixtures.generator import erzeuge_paar
    ordner = erzeuge_paar(7, "schwer", Path("/irgendwo/entwicklung"), korpus="entwicklung")

Gleicher Seed + gleiches Profil + gleicher Korpus => gleiche Zellen und gleiche erwartet.json.

Aufbau:
  Modell      Section / Absatz (Text steht als Markup {+neu+} {-alt-} {@stil:..@}, siehe pdf_writer)
  Kontext     baut Struktur, plant Aenderungen, wendet die Stoerer-Familien an (stoerer.py)
  Rendern     PDF (pdf_writer) und die beiden Excel-Zustaende (openpyxl)
  Bewerten    soll_status, Luecken, erwartet.json
"""
from __future__ import annotations

import json
import random
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import stoerer as st
from .stoerer import schlimmer
from .pdf_writer import PdfDokument, markierte_woerter, muster_bild, scan_bild, zustand

GENERATOR_VERSION = "1"
SEED_GRENZE = 1000          # Entwicklung: Seeds < 1000, Test: Seeds >= 1000 (getrennt, Regel 11)
FORMAT = 1
TYPEN = ("text_ersetzt", "text_eingefuegt", "text_geloescht", "absatz_neu", "absatz_geloescht",
         "section_neu", "section_geloescht")


# ====================================================================== Textbausteine (alles erfunden)
SUBJ = ["Das Speichersystem", "Die Kühlanlage", "Der Wechselrichter", "Das Batteriemodul", "Die Steuerung",
        "Der Schaltschrank", "Das Gehäuse", "Die Überwachungseinheit", "Der Betreiber", "Der Lieferant",
        "Die Wartungsmannschaft", "Die Dokumentation", "Das Rack", "Die Netzanbindung", "Der Hersteller",
        "Das Batteriemanagementsystem", "Die Brandschutzeinrichtung", "Der Transportbehälter",
        "Die Überspannungsschutzeinrichtung", "Das Fernüberwachungssystem", "Der Netzanschlusspunkt"]
PRAED = ["regelmäßig geprüft werden", "vor der Inbetriebnahme dokumentiert werden", "gegen Überhitzung geschützt sein",
         "jederzeit zugänglich sein", "mit einem Typenschild versehen sein", "den geltenden Normen entsprechen",
         "im Betrieb ständig überwacht werden", "bei Netzausfall weiterlaufen", "vor Feuchtigkeit geschützt werden",
         "einzeln gekennzeichnet sein", "nach jedem Einsatz kontrolliert werden", "ohne Werkzeug erreichbar sein"]
GROESSEN = [("Nennleistung", "kW"), ("nutzbare Kapazität", "kWh"), ("Betriebsspannung", "V"),
            ("Masse je Rack", "kg"), ("Lieferzeit", "Wochen"), ("Gewährleistungsdauer", "Monate"),
            ("maximale Stromstärke", "A"), ("Ladezeit", "Stunden"), ("Wartungsintervall", "Monate"),
            ("Geräuschentwicklung", "dB"), ("Kühlleistung", "kW"), ("Reaktionszeit", "ms")]
ZEITEN = ["jährlich", "halbjährlich", "vierteljährlich", "vor jeder Auslieferung", "nach jedem Einsatz", "monatlich"]
PARTIZIPIEN = ["geprüft", "gewartet", "dokumentiert", "kalibriert", "gereinigt", "protokolliert"]
STOERUNGEN = ["Übertemperatur", "Spannungsabfall", "Kommunikationsausfall", "Lüfterausfall", "Isolationsfehler",
              "Überstrom", "Türkontaktverletzung"]
MELDUNGEN = ["eine Meldung", "ein Alarm", "eine Warnung", "ein Störungsbericht"]
NACHWEISE = ["ein Prüfprotokoll", "eine Herstellererklärung", "ein Messbericht", "eine Abnahme vor Ort",
             "eine Sichtprüfung", "ein Zertifikat einer anerkannten Stelle"]
OBJEKTE = ["Anschlussleitungen", "Kühlmittelleitungen", "Sicherungen", "Steckverbindungen", "Lüfter",
           "Temperaturfühler", "Schaltschränke", "Batteriemodule", "Messeinrichtungen", "Schutzabdeckungen",
           "Befestigungsschrauben"]
ZUSTAENDE = ["eindeutig beschriftet", "gegen Verwechslung geschützt", "staubdicht verschlossen",
             "für Prüfungen zugänglich", "fest montiert", "frei von Beschädigungen", "vor Auslieferung getestet"]
BEDINGUNGEN = ["die Umgebungstemperatur im zulässigen Bereich liegt", "eine Freigabe des Auftraggebers vorliegt",
               "keine Störmeldung ansteht", "die Anlage freigeschaltet ist", "der Zugang gesichert ist"]
EREIGNISSE = ["jedem Eingriff", "einem Spannungsausfall", "der Montage", "jeder Störung", "dem Transport"]
FOLGEN = ["eine Funktionsprüfung durchzuführen", "das Ergebnis zu protokollieren",
          "der Betrieb erst nach Freigabe fortzusetzen", "die Anlage neu zu starten"]
ZUBEHOER = ["Befestigungsmaterial", "Kabelsatz", "Bedienungsanleitung", "Ersatzsicherungen", "Montagewerkzeug",
            "Typenschild"]
TERME = [("Betreiber", "die Stelle, die das System dauerhaft nutzt"), ("Rack", "den Einbaurahmen für mehrere Module"),
         ("Stillstandszeit", "die Zeit ohne nutzbare Leistung"), ("Abnahme", "die förmliche Prüfung vor der Übergabe")]
SONDER = [
    "Maßgeblich ist die Fassung vom {t}.{m}.2026 – abweichende Regelungen bedürfen der Schriftform.",
    "Die Außenmaße und die Größe des Gehäuses sind der Zeichnung zu entnehmen.",
    "Der Nennwert von {n} {e} gilt bei {n2} °C Umgebungstemperatur.",
    "Der Begriff „{term}“ bezeichnet im Sinne dieses Dokuments {defi}.",
    "Für Prüfungen in der Nähe von Schaltschränken gilt: Zutritt nur mit Freigabe – ohne Ausnahme.",
]
TITEL = ["Allgemeines", "Geltungsbereich", "Leistung", "Kapazität", "Kühlung", "Wartung", "Sicherheit", "Brandschutz",
         "Lieferung", "Verpackung", "Transport", "Montage", "Inbetriebnahme", "Dokumentation", "Schnittstellen",
         "Steuerung", "Überwachung", "Prüfung", "Abnahme", "Garantie", "Umgebungsbedingungen", "Netzanschluss",
         "Ersatzteile", "Schulung", "Termine", "Gewährleistung", "Schutzarten", "Beschriftung", "Normen",
         "Störmeldungen", "Anschlüsse", "Gehäuse"]
ADVERBIEN = ["regelmäßig", "halbjährlich", "vollständig", "unverzüglich", "schriftlich", "zusätzlich", "automatisch",
             "dauerhaft", "jederzeit", "einzeln", "gemeinsam", "fachgerecht", "ausnahmslos", "unmittelbar"]
MODALE = ["muss", "soll", "darf", "kann", "wird"]
LANGE_WOERTER = ["Überspannungsschutzeinrichtung", "Batteriemanagementsystem", "Netzanschlusspunkt",
                 "Temperaturüberwachung", "Brandschutzkonzept", "Kühlmittelkreislauf", "Wartungsdokumentation",
                 "Fernüberwachungssystem", "Betriebsführungskonzept", "Inbetriebnahmeprotokoll"]
ZAHLEN = [5, 8, 10, 12, 15, 20, 25, 30, 40, 48, 50, 60, 75, 80, 100, 120, 150, 200, 250, 400, 500]
PUNKT_HANDLUNG = ["Prüfung", "Kontrolle", "Reinigung", "Dokumentation", "Kennzeichnung", "Messung", "Abnahme"]
PUNKT_OBJEKT = ["Anschlussleitungen", "Kühlmittelleitungen", "Sicherungen", "Steckverbindungen", "Lüfter",
                "Temperaturfühler", "Schutzabdeckungen", "Messeinrichtungen"]
KURZ_PUNKTE = ["Kabel prüfen", "Lüfter reinigen", "Sicherung tauschen", "Türkontakt testen", "Filter wechseln",
               "Schrauben nachziehen", "Anzeige ablesen", "Protokoll ausfüllen", "Rack beschriften", "Alarm quittieren"]
BAUTEILE = ["Lüfter", "Sicherung", "Temperaturfühler", "Schütz", "Netzteil", "Messwandler", "Steckverbinder"]
FUSSNOTEN = ["Siehe auch das Beiblatt zur Prüfung.", "Angabe gilt bei Normbedingungen.",
             "Wert nach Herstellerangabe, ohne Gewähr.", "Nähere Erläuterung im Anhang des Lastenhefts.",
             "Bezug: interne Richtlinie in der vereinbarten Fassung."]

_REINES_WORT = re.compile(r"^[A-Za-zÄÖÜäöüß0-9-]+$")


# ====================================================================== Modell
@dataclass
class Absatz:
    id: int
    markup: str = ""
    typ: str | None = None            # Aenderungstyp (siehe TYPEN) oder None
    soll: str = "uebernommen"
    marker: str | None = None         # Aufzaehlungszeichen: "•", "–", "1.", "a)"
    art: str = "text"                 # text | tabelle | bild
    tabelle: list[list[str]] | None = None
    tabelle_stil: tuple[str | None, str | None] = ("4472C4", "B4C6E7")
    bild_seed: int = 0
    fussnote: str | None = None
    umbruch: int | None = None        # Seitenumbruch nach dieser Zeile (mitten im Absatz)
    farben: tuple[str, str] | None = None
    decoy_text: str = ""              # versteckte Zeile mit fast gleichem Text (nur Excel)
    kommentar: str = ""
    lfd: int = 0
    stoerer: list[str] = field(default_factory=list)
    luecke: str | None = None
    seite: int = 0
    seite_ende: int = 0
    woerter: int = 0

    def excel_text(self, neu: bool) -> str:
        if self.art == "bild":
            return ""
        if self.art == "tabelle":
            return "\n".join(" | ".join(zustand(c, neu) for c in zeile) for zeile in (self.tabelle or []))
        t = zustand(self.markup, neu)
        if not t:
            return ""
        return f"{self.marker} {t}" if self.marker else t

    def markiert(self) -> int:
        if self.art == "tabelle":
            return sum(markierte_woerter(c) for z in (self.tabelle or []) for c in z)
        return markierte_woerter(self.markup)


@dataclass
class Section:
    id: int
    nummer: str
    titel: str
    ebene: int
    status: str | None = None         # None | "neu" | "geloescht"
    absaetze: list[Absatz] = field(default_factory=list)
    kommentar: str = ""
    decoy_text: str = ""
    lfd: int = 0
    zwei_spalten: bool = False
    scan: bool = False
    soll: str = "uebernommen"
    stoerer: list[str] = field(default_factory=list)
    seite: int = 0

    @property
    def kopf_markup(self) -> str:
        t = f"{self.nummer} {self.titel}"
        if self.status == "neu":
            return "{+" + t + "+}"
        if self.status == "geloescht":
            return "{-" + t + "-}"
        return t

    def im_zustand(self, neu: bool) -> bool:
        return not ((neu and self.status == "geloescht") or (not neu and self.status == "neu"))


# ====================================================================== Profile
PROFILE: dict[str, dict[str, Any]] = {
    "standard": {"kapitel": (3, 5), "gruppen": (4, 6), "min_aenderungen": 9, "stoerer": (0, 2), "pool": "mild"},
    "schwer": {"kapitel": (4, 6), "gruppen": (5, 7), "min_aenderungen": 11, "stoerer": (4, 7), "pool": "alle"},
}


def _profil(profil: str | dict) -> tuple[str, dict[str, Any]]:
    if isinstance(profil, str):
        if profil not in PROFILE:
            raise ValueError(f"unbekanntes Profil: {profil!r} (erlaubt: {', '.join(PROFILE)})")
        return profil, dict(PROFILE[profil])
    cfg = dict(PROFILE[profil.get("basis", "standard")])
    cfg.update({k: v for k, v in profil.items() if k not in ("name", "basis")})
    return str(profil.get("name", "eigen")), cfg


# ====================================================================== Kontext
class Kontext:
    """Baut ein Dokument-Modell Schritt fuer Schritt. Wird auch von den Stoerer-Hooks benutzt."""

    def __init__(self, seed: int, profil_name: str, cfg: dict[str, Any], korpus: str, versuch: int = 0) -> None:
        self.seed, self.profil_name, self.cfg, self.korpus = seed, profil_name, cfg, korpus
        self.versuch = versuch
        self.rng = random.Random(seed * 1_000_003 + sum(map(ord, korpus + profil_name)) * 7919 + versuch * 104_729)
        self.fehlend: list[str] = []
        self.sektionen: list[Section] = []
        self._zaehler = 0
        self.benutzt: set[str] = set()
        self.gewaehlt: list[str] = []
        self.aktiv: list[str] = []
        self.luecken: list[dict[str, Any]] = []
        self.pdfkw: dict[str, Any] = {}
        self.layout: dict[str, Any] = {"base": None, "lfd": False, "kommentar": False, "kommentar_vor": False,
                                      "zahl": False, "formeln": False, "merges": False, "banner": False,
                                      "versteckt": False}
        self.blaetter = "deckblatt"           # deckblatt | nur | mehrere
        self.opt: dict[str, Any] = {"verschmolzen": False, "sonder": False, "lange_woerter": False,
                                    "stil_neu": ["n"], "stil_weg": ["g"], "misch_p": 0.5, "rotiere": False,
                                    "sub_nummern": None, "farben_pro_absatz": False}
        self.hauptblatt = 1

    FUSSNOTEN = FUSSNOTEN
    BAUTEILE = BAUTEILE

    # ---- Hilfen fuer Hooks ----------------------------------------------------
    def neue_id(self) -> int:
        self._zaehler += 1
        return self._zaehler

    def neuer_absatz(self, markup: str = "", **kw: Any) -> Absatz:
        return Absatz(id=self.neue_id(), markup=markup, **kw)

    def alle_absaetze(self) -> list[tuple[Section, Absatz]]:
        return [(s, a) for s in self.sektionen for a in s.absaetze]

    def zeilen(self, a: Absatz) -> int:
        """Grobe Schaetzung der Zeilen im PDF (84 Zeichen je Zeile)."""
        if a.art == "tabelle":
            return len(a.tabelle or [])
        if a.art == "bild":
            return 9
        n = len(re.sub(r"\{\+|\+\}|\{-|-\}|\{@\w+:|@\}", "", a.markup)) + (3 if a.marker else 0)   # alt UND neu stehen im PDF
        return max(1, -(-n // 84))

    def section_zeilen(self, s: Section) -> int:
        return sum(self.zeilen(a) for a in s.absaetze)

    def markiere(self, ziel: Absatz | Section, name: str, soll: str | None = None) -> None:
        if name not in ziel.stoerer:
            ziel.stoerer.append(name)
        if soll:
            ziel.soll = schlimmer(ziel.soll, soll)

    # ---- Text ---------------------------------------------------------------
    def _zahl(self) -> int:
        return self.rng.choice(ZAHLEN)

    def satz(self) -> str:
        r = self.rng
        sonder = self.opt["sonder"] and r.random() < 0.3
        if sonder:
            t = r.choice(SONDER)
            g = r.choice(GROESSEN)
            term = r.choice(TERME)
            return t.format(t=r.randint(1, 28), m=r.randint(1, 12), n=self._zahl(), e=g[1], n2=r.choice([20, 25, 40]),
                            term=term[0], defi=term[1])
        k = r.randint(1, 10)
        s = r.choice(SUBJ)
        if k == 1:
            return f"{s} muss {r.choice(PRAED)}."
        if k == 2:
            g = r.choice(GROESSEN)
            return f"Die {g[0]} beträgt {self._zahl()} {g[1]}."
        if k == 3:
            g = r.choice(GROESSEN)
            return f"Die {g[0]} darf {self._zahl()} {g[1]} nicht überschreiten."
        if k == 4:
            return f"{s} wird {r.choice(ZEITEN)} {r.choice(PARTIZIPIEN)}."
        if k == 5:
            return f"Bei {r.choice(STOERUNGEN)} wird {r.choice(MELDUNGEN)} an die Leitstelle übermittelt."
        if k == 6:
            return f"Der Nachweis erfolgt durch {r.choice(NACHWEISE)}."
        if k == 7:
            return f"Alle {r.choice(OBJEKTE)} sind {r.choice(ZUSTAENDE)}."
        if k == 8:
            return f"{s} {r.choice(MODALE[:2])} {r.choice(PRAED)}, sofern {r.choice(BEDINGUNGEN)}."
        if k == 9:
            return f"Die Lieferung umfasst {self._zahl()} {r.choice(OBJEKTE)} einschließlich {r.choice(ZUBEHOER)}."
        return f"Nach {r.choice(EREIGNISSE)} ist {r.choice(FOLGEN)}."

    def text(self, saetze: int | None = None) -> str:
        """Ein Absatz-Text (1-4 Saetze), im Dokument einmalig."""
        for _ in range(30):
            n = saetze or self.rng.choices([1, 2, 3, 4], [25, 35, 25, 15])[0]
            t = " ".join(self.satz() for _ in range(n))
            if t not in self.benutzt:
                self.benutzt.add(t)
                return t
        t = t + f" Ergänzung {self._zahl()}."
        self.benutzt.add(t)
        return t

    def kurzer_punkt(self, nummeriert: bool = False) -> str:
        for _ in range(30):
            if nummeriert:
                t = self.rng.choice(KURZ_PUNKTE)
            else:
                t = f"{self.rng.choice(PUNKT_HANDLUNG)} der {self.rng.choice(PUNKT_OBJEKT)}"
            if t not in self.benutzt:
                self.benutzt.add(t)
                return t
        return t + f" {self._zahl()}"

    # ---- Markup fuer Aenderungen --------------------------------------------------
    def _stil(self, art: str) -> str:
        wahl = self.opt["stil_neu" if art == "neu" else "stil_weg"]
        if len(wahl) == 1:
            return wahl[0]
        if self.rng.random() < self.opt["misch_p"]:
            return self.rng.choice(wahl[1:])
        return wahl[0]

    def wrap(self, a: Absatz | None, art: str, text: str) -> str:
        stil = self._stil(art)
        if a is not None:
            a.soll = schlimmer(a.soll, st.STIL_SOLL.get(stil))
            fam = st.STIL_FAMILIE.get(stil)
            if fam and fam not in a.stoerer:
                a.stoerer.append(fam)
        return st.wrap_stil(stil, text)

    def ersatz(self, alt: str) -> str:
        r = self.rng
        if alt.isdigit():
            neu = alt
            while neu == alt:
                neu = str(max(1, int(alt) + r.choice([-20, -10, -5, 5, 10, 20, 50])))
            return neu
        if alt in MODALE:
            return r.choice([m for m in MODALE if m != alt])
        pool = LANGE_WOERTER if (self.opt["lange_woerter"] and r.random() < 0.6) else ADVERBIEN
        neu = r.choice(pool)
        return neu if neu != alt else neu + "e"

    def tausch(self, a: Absatz, alt: str, neu: str) -> str:
        fuse = self.opt["verschmolzen"] and self.rng.random() < 0.7
        sep = "" if fuse else " "
        return self.wrap(a, "weg", alt) + sep + self.wrap(a, "neu", neu)

    def fast_gleich(self, text: str) -> str:
        """Fast gleicher Text: genau ein einfaches Wort ausgetauscht."""
        tokens = text.split(" ")
        rein = [i for i, t in enumerate(tokens) if i > 0 and len(t) >= 3 and _REINES_WORT.match(t)]
        if not rein:
            return text + " (Entwurf)"
        i = self.rng.choice(rein)
        tokens[i] = self.ersatz(tokens[i])
        return " ".join(tokens)

    def verlaengere(self, a: Absatz, mindest_zeilen: int) -> None:
        """Haengt unveraenderte Saetze an einen Absatz, bis er mindestens so viele Zeilen hat."""
        n = 0
        while self.zeilen(a) < mindest_zeilen and n < 4:
            a.markup += " " + self.satz()
            n += 1

    # ---- Aenderungen am Text eines Absatzes (Text darf noch kein Markup enthalten) ----------
    def aendere(self, a: Absatz, typ: str, anzahl: int = 1) -> bool:
        r = self.rng
        tokens = a.markup.split(" ")
        rein = [i for i, t in enumerate(tokens) if i > 0 and len(t) >= 2 and _REINES_WORT.match(t)]
        if typ == "text_ersetzt":
            if not rein:
                return False
            gewaehlt: list[int] = []
            for i in r.sample(rein, len(rein)):
                if all(abs(i - j) > 1 for j in gewaehlt):
                    gewaehlt.append(i)
                if len(gewaehlt) >= anzahl:
                    break
            for i in sorted(gewaehlt):
                alt = tokens[i]
                if (i + 1 < len(tokens) and i + 1 in rein and i + 1 not in gewaehlt and r.random() < 0.25
                        and not alt.isdigit()):
                    alt2 = alt + " " + tokens[i + 1]
                    tokens[i] = self.tausch(a, alt2, self.ersatz(alt) + " " + r.choice(ADVERBIEN))
                    tokens[i + 1] = None  # type: ignore[call-overload]
                else:
                    tokens[i] = self.tausch(a, alt, self.ersatz(alt))
            a.markup = " ".join(t for t in tokens if t is not None)
        elif typ == "text_eingefuegt":
            if r.random() < 0.5 or not rein:
                a.markup = a.markup + " " + self.wrap(a, "neu", self.satz())
            else:
                i = r.choice(rein)
                neu = " ".join(r.sample(ADVERBIEN, r.randint(1, 2)))
                tokens.insert(i + 1, self.wrap(a, "neu", neu))
                a.markup = " ".join(tokens)
        elif typ == "text_geloescht":
            saetze = re.split(r"(?<=\.) ", a.markup)
            if len(saetze) >= 2:
                a.markup = " ".join(saetze[:-1]) + " " + self.wrap(a, "weg", saetze[-1])
            else:
                mitte = [i for i in rein if i < len(tokens) - 2]
                if not mitte:
                    return False
                i = r.choice(mitte)
                j = min(len(tokens) - 1, i + r.randint(1, 3))
                while j > i + 1 and not all(_REINES_WORT.match(t) for t in tokens[i:j]):
                    j -= 1
                if not all(_REINES_WORT.match(t) for t in tokens[i:j]):
                    return False
                tokens[i:j] = [self.wrap(a, "weg", " ".join(tokens[i:j]))]
                a.markup = " ".join(tokens)
        elif typ == "absatz_neu":
            a.markup = self.wrap(a, "neu", a.markup)
        elif typ == "absatz_geloescht":
            a.markup = self.wrap(a, "weg", a.markup)
        else:
            raise ValueError(typ)
        a.typ = typ
        return True

    # ---- Struktur ---------------------------------------------------------------------------
    def neue_section(self, nummer: str, titel: str, ebene: int, absaetze: int = 2) -> Section:
        s = Section(id=self.neue_id(), nummer=nummer, titel=titel, ebene=ebene)
        for _ in range(absaetze):
            s.absaetze.append(self.neuer_absatz(self.text()))
        return s

    def _freier_titel(self) -> str:
        benutzt = {s.titel for s in self.sektionen}
        frei = [t for t in TITEL if t not in benutzt]
        if frei:
            return self.rng.choice(frei)
        return f"Anhang {len(benutzt) + 1}"

    def baue_struktur(self) -> None:
        r = self.rng
        n_kap = r.randint(*self.cfg["kapitel"])
        sub_nr = self.opt["sub_nummern"] or {}
        for k in range(1, n_kap + 1):
            kap = self.neue_section(str(k), self._freier_titel(), 1, r.randint(1, 3))
            self.sektionen.append(kap)
            if k in sub_nr:
                nummern = sub_nr[k]
            else:
                nummern = list(range(1, r.choice([0, 1, 2, 2, 3]) + 1))
            for j in nummern:
                sub = self.neue_section(f"{k}.{j}", self._freier_titel(), 2, r.randint(1, 3))
                self.sektionen.append(sub)
                if r.random() < 0.15:
                    self.sektionen.append(self.neue_section(f"{k}.{j}.1", self._freier_titel(), 3, r.randint(1, 2)))

    def _kandidaten(self) -> list[tuple[Section, Absatz]]:
        return [(s, a) for s, a in self.alle_absaetze()
                if s.status is None and a.typ is None and a.art == "text" and not a.marker]

    def zaehle_aenderungen(self) -> int:
        n = 0
        for s in self.sektionen:
            if s.status:
                n += 1
            else:
                n += sum(1 for a in s.absaetze if a.typ)
        return n

    def _neue_section_ans_kapitelende(self, absaetze: int = 2) -> Section:
        """Neue Section am Ende eines Kapitels oder als neues Kapitel am Dokumentende."""
        r = self.rng
        kapitel = [s for s in self.sektionen if s.ebene == 1]
        if r.random() < 0.35 or not kapitel:
            nr = str(max(int(s.nummer.split(".")[0]) for s in self.sektionen) + 1) if self.sektionen else "1"
            sec = self.neue_section(nr, self._freier_titel(), 1, absaetze)
            self.sektionen.append(sec)
        else:
            kap = r.choice(kapitel)
            top = kap.nummer
            subs = [s for s in self.sektionen if s.nummer.startswith(top + ".") and s.ebene == 2]
            nr_j = max([int(s.nummer.split(".")[1]) for s in subs] + [0]) + 1
            sec = self.neue_section(f"{top}.{nr_j}", self._freier_titel(), 2, absaetze)
            ende = max(i for i, s in enumerate(self.sektionen) if s.nummer == top or s.nummer.startswith(top + "."))
            self.sektionen.insert(ende + 1, sec)
        sec.status = "neu"
        for a in sec.absaetze:
            a.markup = self.wrap(a, "neu", a.markup)
            a.typ = "absatz_neu"
        return sec

    def plane_aenderungen(self) -> None:
        r = self.rng
        arten = ["ersetzen", "ersetzen", "ersetzen", "einfuegen", "einfuegen", "loeschen", "absatz_neu",
                 "absatz_weg", "section_neu", "section_weg"]
        r.shuffle(arten)
        gruppen = arten[:r.randint(*self.cfg["gruppen"])]
        for art in ("ersetzen", "einfuegen", "absatz_neu"):       # Grundausstattung
            if art not in gruppen:
                gruppen.append(art)
        versuche = 0
        while True:
            for art in gruppen:
                self._gruppe(art)
            versuche += 1
            if self.zaehle_aenderungen() >= self.cfg["min_aenderungen"] or versuche > 6:
                break
            gruppen = r.sample(["ersetzen", "einfuegen", "loeschen", "absatz_neu"], 2)

    def _gruppe(self, art: str) -> None:
        r = self.rng
        kand = self._kandidaten()
        if art in ("ersetzen", "einfuegen", "loeschen"):
            if not kand:
                return
            sec = r.choice(kand)[0]
            ziele = [a for s, a in kand if s is sec]
            r.shuffle(ziele)
            typ = {"ersetzen": "text_ersetzt", "einfuegen": "text_eingefuegt", "loeschen": "text_geloescht"}[art]
            for a in ziele[:r.randint(1, 3)]:
                self.aendere(a, typ, r.choice([1, 1, 2]) if typ == "text_ersetzt" else 1)
        elif art == "absatz_neu":
            secs = [s for s in self.sektionen if s.status is None]
            sec = r.choice(secs)
            for _ in range(r.randint(1, 2)):
                a = self.neuer_absatz(self.text())
                self.aendere(a, "absatz_neu")
                sec.absaetze.insert(r.randint(0, len(sec.absaetze)), a)
        elif art == "absatz_weg":
            ziele = [(s, a) for s, a in kand if len([x for x in s.absaetze if x.art == "text"]) >= 2]
            if ziele:
                self.aendere(r.choice(ziele)[1], "absatz_geloescht")
        elif art == "section_neu":
            self._neue_section_ans_kapitelende(r.randint(1, 3))
        elif art == "section_weg":
            ziele = [s for s in self.sektionen
                     if s.status is None and s.ebene >= 2 and s.absaetze
                     and all(a.typ is None and a.art == "text" and not a.marker for a in s.absaetze)
                     and s.nummer not in self.opt.get("geschuetzt", ())
                     and not any(t.nummer.startswith(s.nummer + ".") for t in self.sektionen)]
            if ziele:
                sec = r.choice(ziele)
                sec.status = "geloescht"
                for a in sec.absaetze:
                    self.aendere(a, "absatz_geloescht")

    def kleine_section(self, mit_aenderung: bool = True, absaetze: int = 2) -> Section:
        """Haengt eine kleine neue (unveraenderte) Section am Dokumentende an -- Rueckfall fuer Hooks."""
        nr = str(max(int(s.nummer.split(".")[0]) for s in self.sektionen) + 1)
        sec = self.neue_section(nr, self._freier_titel(), 1, absaetze)
        self.sektionen.append(sec)
        if mit_aenderung:
            self.aendere(sec.absaetze[0], "text_ersetzt")
            if len(sec.absaetze) > 1:
                self.aendere(sec.absaetze[1], "text_eingefuegt")
        return sec

    # ---- Layout und Optionen ------------------------------------------------------------------
    def waehle_stoerer(self) -> None:
        r = self.rng
        ex = self.cfg.get("stoerer_liste")
        if ex is not None:
            for n in ex:
                if n not in st.REGISTER:
                    raise ValueError(f"unbekannte Stoerer-Familie: {n}")
                if st.REGISTER[n].nur_test and self.korpus != "test":
                    raise ValueError(f"{n} ist NUR_TEST und darf nicht im Entwicklungskorpus vorkommen (Regel 11)")
            self.gewaehlt = sorted(ex, key=lambda x: st.REGISTER[x].reihenfolge)
            return
        pool = st.pool(self.cfg["pool"], self.korpus)
        lo, hi = self.cfg["stoerer"]
        n = r.randint(lo, hi)
        if self.cfg["pool"] == "alle" and pool:
            # drei reihum festgelegte Familien (jede Familie kommt so sicher vor), der Rest zufaellig
            fest: list[str] = []
            for j in range(3):
                name = pool[(self.seed * 5 + j * 7) % len(pool)]
                if name not in fest:
                    fest.append(name)
            gew = fest + [p for p in r.sample(pool, len(pool)) if p not in fest][:max(0, n - len(fest))]
        else:
            if r.random() < 0.25 and self.cfg["pool"] == "mild":
                n = 0
            gew = r.sample(pool, min(n, len(pool)))
        if self.korpus == "test":
            for name, (modul, rest) in st.TEST_ZWANG.items():
                if self.seed % modul == rest and name not in gew:
                    gew.append(name)
        self.gewaehlt = sorted(gew, key=lambda x: st.REGISTER[x].reihenfolge)

    def waehle_layout_und_pdf(self) -> None:
        r = self.rng
        lay = self.layout
        if lay["base"] is None:
            lay["base"] = self.cfg.get("layout") or r.choices(["L1", "L2", "L3"], [35, 25, 40])[0]
        if lay["zahl"] and lay["base"] == "L2":
            lay["base"] = "L1"
        if "variante" not in self.pdfkw:
            self.pdfkw["variante"] = r.choices(["farbe", "strich", "hintergrund", "annot"], [60, 15, 15, 10])[0]
        if "deckblatt" in self.cfg and self.blaetter == "deckblatt":
            self.blaetter = "deckblatt" if self.cfg["deckblatt"] else "nur"
        elif self.blaetter == "deckblatt" and r.random() < 0.3:
            self.blaetter = "nur"

    # ---- Abschluss -------------------------------------------------------------------------------
    def finalisieren(self) -> None:
        n = 0
        for s in self.sektionen:
            n += 1
            s.lfd = n
            for a in s.absaetze:
                n += 1
                a.lfd = n
                a.woerter = a.markiert()
        if self.layout["kommentar"]:
            self._kommentare()

    def _kommentare(self) -> None:
        r = self.rng

        def kommentar(text_alt: str) -> str:
            woerter = text_alt.split()
            zitat = " ".join(woerter[:r.randint(5, 9)]) if woerter else "Abschnitt"
            teile = [f"Rückfrage Auftraggeber: Ist die Formulierung „{zitat} …“ so verbindlich?",
                     f"Antwort Lieferant: {self.satz()}",
                     f"Hinweis Projektleitung: {self.satz()} {self.satz()}"]
            t = " ".join(teile)
            while len(t) < len(text_alt) * 1.6 + 40:
                t += " " + self.satz()
            return t

        absatzmodus = self.layout["base"] == "L3"
        for s in self.sektionen:
            if s.status == "neu":
                continue
            if absatzmodus:
                s.kommentar = kommentar(s.titel)
                for a in s.absaetze:
                    if a.art == "text" and zustand(a.markup, False):
                        a.kommentar = kommentar(zustand(a.markup, False))
            else:
                alt = "\n".join(t for t in (a.excel_text(False) for a in s.absaetze) if t)
                s.kommentar = kommentar(alt)

    # ---- Lauf ---------------------------------------------------------------------------------------
    def lauf(self) -> None:
        self.waehle_stoerer()
        for name in self.gewaehlt:
            hook = st.REGISTER[name].vorbereiten
            if hook:
                hook(self)
        self.waehle_layout_und_pdf()
        self.baue_struktur()
        self.plane_aenderungen()
        for name in self.gewaehlt:
            hook = st.REGISTER[name].anwenden
            if hook is None or hook(self) is not False:
                self.aktiv.append(name)
            else:
                self.fehlend.append(name)
        self.finalisieren()


# ====================================================================== Excel
KOPF = {"nr": "Nr.", "titel": "Titel", "nrtitel": "Kapitel", "text": "Inhalt", "status": "Bemerkung",
        "prio": "Prio", "kommentar": "Anmerkung", "lfd": "Lfd", "zeichen": "Zeichen"}
BREITE = {"nr": 8, "titel": 22, "nrtitel": 26, "text": 70, "status": 14, "prio": 8, "kommentar": 60, "lfd": 6,
          "zeichen": 9}
BASIS = {
    "L1": ("sektion", ["nr", "titel", "text", "status"], ["Nr.", "Titel", "Inhalt", "Bemerkung"]),
    "L2": ("sektion", ["nrtitel", "text", "status"], ["Kapitel", "Anforderung", "Kommentar"]),
    "L3": ("absatz", ["nr", "titel", "text", "prio"], ["Nr", "Titel", "Text", "Prio"]),
}


def _spalten(lay: dict[str, Any]) -> tuple[str, list[str], list[str]]:
    modus, kinds, kopf = BASIS[lay["base"]]
    kinds, kopf = list(kinds), list(kopf)
    if lay["lfd"]:
        kinds.insert(0, "lfd")
        kopf.insert(0, "Lfd")
    if lay["kommentar"]:
        i = kinds.index("text")
        pos = i if lay["kommentar_vor"] else i + 1
        kinds.insert(pos, "kommentar")
        kopf.insert(pos, "Anmerkung")
    if lay["formeln"]:
        kinds.append("zeichen")
        kopf.append("Zeichen")
    return modus, kinds, kopf


def _nr_wert(nummer: str, zahl: bool):
    """Section-Nummer als Zahl (2.10 wird zu 2.1!) -- absichtlich wie in echten Tabellen."""
    if not zahl:
        return nummer
    teile = nummer.split(".")
    if len(teile) == 1:
        return int(teile[0])
    if len(teile) == 2:
        return float(nummer)
    return nummer


def _nr_format(nummer: str) -> str | None:
    teile = nummer.split(".")
    if len(teile) == 2:
        return "0." + "0" * len(teile[1])
    return None


class _Pos:
    """Wo steht was im erzeugten Blatt (Zeilen sind 1-basiert)."""

    def __init__(self) -> None:
        self.absatz: dict[int, int] = {}
        self.sec_kopf: dict[int, int] = {}
        self.sec_ende: dict[int, int] = {}
        self.sec_text: dict[int, int] = {}
        self.letzte: int = 0
        self.blatt: int = 1
        self.nr_spalte: int = 1
        self.text_spalte: int = 1


def schreibe_excel(pfad: Path, ctx: Kontext, neu: bool) -> _Pos:
    lay = ctx.layout
    modus, kinds, kopf = _spalten(lay)
    wb = Workbook()
    wb.properties.creator = "erfunden"
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    ws_haupt = wb.active
    namen = {"deckblatt": ["Deckblatt", "CDE"], "nur": ["CDE"],
             "mehrere": ["Deckblatt", "Historie", "CDE", "Glossar"]}[ctx.blaetter]
    ws_haupt.title = namen[0]
    blaetter = {namen[0]: ws_haupt}
    for n in namen[1:]:
        blaetter[n] = wb.create_sheet(n)
    _nebenblaetter(ctx, blaetter, neu)
    ws = blaetter["CDE"]
    pos = _Pos()
    pos.blatt = namen.index("CDE") + 1

    hr = 1
    if lay["banner"]:
        ws.cell(1, 1, "CDE-Anforderungskatalog (erfunden)")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(kinds))
        ws.cell(1, 1).font = Font(bold=True, size=13)
        hr = 2
    for c, text in enumerate(kopf, 1):
        z = ws.cell(hr, c, text)
        z.font = Font(bold=True)
        z.fill = PatternFill("solid", fgColor="DDDDDD")
    pos.nr_spalte = (kinds.index("nr") + 1) if "nr" in kinds else (kinds.index("nrtitel") + 1)
    pos.text_spalte = kinds.index("text") + 1
    zcol = get_column_letter(kinds.index("zeichen") + 1) if "zeichen" in kinds else None
    tcol = get_column_letter(pos.text_spalte)

    zeile = hr
    versteckt: list[int] = []
    merges: list[tuple[int, int]] = []

    def schreibe(werte: dict[str, Any], hidden: bool = False, nrtext: str = "") -> int:
        nonlocal zeile
        zeile += 1
        for c, k in enumerate(kinds, 1):
            if k == "zeichen":
                z = ws.cell(zeile, c, f"=LEN({tcol}{zeile})")
                continue
            v = werte.get(k)
            if v is None:
                continue
            z = ws.cell(zeile, c, v)
            if isinstance(v, str) and v.startswith("="):
                z.data_type = "s"
            if k == "nr" and nrtext and isinstance(v, float):
                f = _nr_format(nrtext)
                if f:
                    z.number_format = f
        if hidden:
            versteckt.append(zeile)
        return zeile

    for s in ctx.sektionen:
        if not s.im_zustand(neu):
            continue
        texte = [(a, a.excel_text(neu)) for a in s.absaetze]
        texte = [(a, t) for a, t in texte if t]
        nr = _nr_wert(s.nummer, lay["zahl"])
        kopfwerte = {"lfd": s.lfd, "nr": nr, "titel": s.titel, "nrtitel": f"{s.nummer} {s.titel}",
                     "kommentar": s.kommentar or None}
        if modus == "sektion":
            inhalt = "\n".join(t for _, t in texte)
            if s.decoy_text:
                schreibe({**kopfwerte, "text": s.decoy_text, "status": "Status offen"}, True, s.nummer)
            r = schreibe({**kopfwerte, "text": inhalt or None, "status": "Status offen"}, False, s.nummer)
            pos.sec_kopf[s.id] = pos.sec_ende[s.id] = pos.sec_text[s.id] = r
            for a, _ in texte:
                pos.absatz[a.id] = r
        else:
            first = schreibe({**kopfwerte, "prio": "A"}, False, s.nummer)
            pos.sec_kopf[s.id] = first
            letzte = first
            for a, t in texte:
                if a.decoy_text:
                    schreibe({"text": a.decoy_text, "prio": None if lay["merges"] else "B"}, True)
                w = {"lfd": a.lfd, "text": t, "kommentar": a.kommentar or None}
                if not lay["merges"]:
                    w["prio"] = "B"
                letzte = schreibe(w)
                pos.absatz[a.id] = letzte
            pos.sec_ende[s.id] = letzte
            if lay["merges"] and letzte > first:
                merges.append((first, letzte))
    pos.letzte = zeile
    if lay["formeln"] and zcol:
        zeile += 1
        ws.cell(zeile, 1, "Summe")
        ws.cell(zeile, kinds.index("zeichen") + 1, f"=SUM({zcol}{hr + 1}:{zcol}{pos.letzte})")
    if merges:
        pc = kinds.index("prio") + 1
        for a, b in merges:
            ws.merge_cells(start_row=a, start_column=pc, end_row=b, end_column=pc)
    for r in versteckt:
        ws.row_dimensions[r].hidden = True
    for c, k in enumerate(kinds, 1):
        ws.column_dimensions[get_column_letter(c)].width = BREITE[k]
    for row in ws.iter_rows(min_row=hr + 1):
        for z in row:
            z.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(str(pfad))
    return pos


def _nebenblaetter(ctx: Kontext, blaetter: dict[str, Any], neu: bool) -> None:
    if "Deckblatt" in blaetter:
        ws = blaetter["Deckblatt"]
        ws["A1"] = "Beispielprojekt (erfunden)"
        ws["A2"] = "Dieses Blatt gehört nicht zur Gliederung und bleibt unverändert."
    if "Historie" in blaetter:
        ws = blaetter["Historie"]
        ws.append(["Datum", "Änderung", "Bearbeiter"])
        ws.append(["2026-01-05", "Erstfassung angelegt", "Team A"])
        ws.append(["2026-02-11", "Abschnitte zur Lieferung ergänzt", "Team B"])
        zitat = [a for s, a in ctx.alle_absaetze()
                 if a.typ in ("text_ersetzt", "text_eingefuegt", "text_geloescht") and a.art == "text"]
        if zitat:
            ws.append(["2026-03-02", "Vorherige Fassung: " + zustand(zitat[0].markup, False), "Team A"])
        ws.append(["2026-04-17", "Redaktionelle Korrekturen", "Team C"])
    if "Glossar" in blaetter:
        ws = blaetter["Glossar"]
        ws.append(["Begriff", "Erklärung"])
        for term, defi in TERME:
            ws.append([term, defi])


# ====================================================================== PDF
def rendere_pdf(ctx: Kontext) -> PdfDokument:
    kw = dict(ctx.pdfkw)
    d = PdfDokument(**kw)
    farbpaar = None
    if ctx.opt["farben_pro_absatz"]:
        farbpaar = (st.BLAUTOENE, st.ROTTOENE)
    for s in ctx.sektionen:
        if s.scan:
            elemente = [("h", s.kopf_markup)] + [("p", ((a.marker + " ") if a.marker else "") + a.markup)
                                                 for a in s.absaetze if a.art == "text"]
            bild = scan_bild(elemente, farbe_neu=_rgb255(d.farbe_neu), farbe_weg=_rgb255(d.farbe_weg),
                             seed=ctx.seed + s.id)
            d.scan_seite(bild)
            s.seite = d.seitenzahl
            for a in s.absaetze:
                a.seite = a.seite_ende = d.seitenzahl
            continue
        d._bereit()
        groesse, abstand = (14.0, 18.0) if s.ebene == 1 else (12.5, 16.0)
        d.ueberschrift(s.kopf_markup, groesse=groesse, zeilenabstand=abstand)
        s.seite = d.seitenzahl
        if s.zwei_spalten:
            texte = [a for a in s.absaetze if a.art == "text"]
            haelfte = (len(texte) + 1) // 2
            links, rechts = texte[:haelfte], texte[haelfte:]
            d.zwei_spalten([_inline(a) for a in links], [_inline(a) for a in rechts])
            for a in texte:
                a.seite = a.seite_ende = d.seitenzahl
            continue
        for a in s.absaetze:
            d._bereit()
            if a.art == "tabelle":
                kopf, band = a.tabelle_stil
                d.tabelle(a.tabelle or [], kopf_farbe=kopf, band_farbe=band)
                a.seite = a.seite_ende = d.seitenzahl
                continue
            if a.art == "bild":
                d.bild(muster_bild(240, 140, a.bild_seed), 200.0)
                a.seite = a.seite_ende = d.seitenzahl
                continue
            kwb: dict[str, Any] = {}
            if a.marker:
                kwb.update(marker=a.marker, marker_stil=_marker_stil(a.markup))
            if a.umbruch:
                kwb["umbruch_nach_zeile"] = a.umbruch
            farben = a.farben
            if farbpaar and farben is None and a.typ and ctx.rng.random() < 0.6:
                farben = (ctx.rng.choice(farbpaar[0]), ctx.rng.choice(farbpaar[1]))
                a.farben = farben
            if farben:
                kwb["farben"] = farben
            markup = a.markup
            if a.fussnote and markup[:1].isalpha():
                m = re.match(r"(\S+)(.*)", markup, re.S)
                if m:
                    markup = m.group(1) + "¹" + m.group(2)
            n0 = len(d.getrennt)
            d.block(markup, **kwb)
            a.seite, a.seite_ende = d.letzter_block
            if a.seite_ende > a.seite and "seitenumbruch_im_absatz" in ctx.gewaehlt:
                ctx.markiere(a, "seitenumbruch_im_absatz")
            if a.fussnote:
                d.fussnote(a.fussnote, "1")
            for wort, stile, strich in d.getrennt[n0:]:
                if strich == "­" and set(stile) - {"k"}:
                    a.soll = schlimmer(a.soll, "pruefen")
                    if "silbentrennung" not in a.stoerer:
                        a.stoerer.append("silbentrennung")
    seite = _rot_seite(ctx)
    if seite:
        d.seite_rotieren(90, seite)
    return d


def _rgb255(f) -> tuple[int, int, int]:
    return tuple(int(round(v * 255)) for v in f)  # type: ignore[return-value]


def _marker_stil(markup: str) -> str:
    m = markup.strip()
    if m.startswith("{+") and m.endswith("+}") and m.count("{+") == 1:
        return "n"
    if m.startswith("{-") and m.endswith("-}") and m.count("{-") == 1:
        return "g"
    return "k"


def _inline(a: Absatz) -> str:
    return ((a.marker + " ") if a.marker else "") + a.markup


def _rot_seite(ctx: Kontext) -> int:
    """Waehlt die Seite, die /Rotate 90 bekommt (Familie gedrehte_seite); 0 = keine."""
    if not ctx.opt["rotiere"]:
        return 0
    mit = sorted({a.seite for s, a in ctx.alle_absaetze() if a.typ and not s.scan})
    if not mit:
        return 0
    seite = mit[len(mit) // 2]
    ctx.opt["rotiert_seite"] = seite
    for s, a in ctx.alle_absaetze():
        if a.seite and a.seite <= seite <= max(a.seite_ende, a.seite):
            ctx.markiere(a, "gedrehte_seite", "pruefen")
    for s in ctx.sektionen:
        if s.seite == seite:
            ctx.markiere(s, "gedrehte_seite", "pruefen")
    return seite


# ====================================================================== Auswertung
def _luecken_sammeln(ctx: Kontext, d: PdfDokument) -> None:
    for s in ctx.sektionen:
        if s.scan:
            w = markierte_woerter(s.kopf_markup) + sum(a.markiert() for a in s.absaetze)
            ctx.luecken.append({"art": "scan_seite", "seite": s.seite, "woerter": w, "mit_aenderung": w > 0})
            ctx.markiere(s, "scan_seite", "nicht_uebernommen")
            for a in s.absaetze:
                ctx.markiere(a, "scan_seite", "nicht_uebernommen")
        elif s.zwei_spalten:
            w = sum(a.markiert() for a in s.absaetze if a.art == "text")
            seiten = sorted({a.seite for a in s.absaetze if a.art == "text"})
            ctx.luecken.append({"art": "mehrspaltig", "seite": seiten[0] if seiten else s.seite, "woerter": w,
                                "mit_aenderung": w > 0})
            ctx.markiere(s, "zwei_spalten", "nicht_uebernommen")
            for a in s.absaetze:
                ctx.markiere(a, "zwei_spalten", "nicht_uebernommen")
        for a in s.absaetze:
            if a.art == "tabelle":
                w = a.markiert()
                ctx.luecken.append({"art": "tabelle", "seite": a.seite, "woerter": w, "mit_aenderung": w > 0})
                if a.typ:
                    ctx.markiere(a, "tabelle_mit_farbkopf", "nicht_uebernommen")
            elif a.art == "bild":
                ctx.luecken.append({"art": "bild", "seite": a.seite, "woerter": 0, "mit_aenderung": False})
    # unbekannte Farbe (gruen) -> je Seite eine Luecke
    je_seite: dict[int, int] = {}
    for s in ctx.sektionen:
        if s.scan:
            continue
        for a in s.absaetze:
            if a.art == "text" and "{@gn:" in a.markup:
                n = sum(len([w for w in t.split() if w]) for t in re.findall(r"\{@gn:(.*?)@\}", a.markup))
                je_seite[a.seite_ende or a.seite] = je_seite.get(a.seite_ende or a.seite, 0) + n
                ctx.markiere(a, "gruene_markierung", "nicht_uebernommen")
    for seite in sorted(je_seite):
        ctx.luecken.append({"art": "unbekannte_farbe", "seite": seite, "woerter": je_seite[seite],
                            "mit_aenderung": True})


def _bewerte_struktur(ctx: Kontext) -> None:
    """Formeln und verbundene Zellen: Zeilen-Aenderungen werden zu 'pruefen'."""
    lay = ctx.layout
    absatzmodus = BASIS[lay["base"]][0] == "absatz"
    for s in ctx.sektionen:
        if s.status:
            if lay["formeln"]:
                ctx.markiere(s, "formeln_merges", "pruefen")
            continue
        for a in s.absaetze:
            if absatzmodus and a.typ in ("absatz_neu", "absatz_geloescht") and (lay["formeln"] or lay["merges"]):
                ctx.markiere(a, "formeln_merges", "pruefen")
    # Absaetze in neuen/geloeschten Sections geben ihre Strenge an die Section weiter
    for s in ctx.sektionen:
        if s.status:
            for a in s.absaetze:
                s.soll = schlimmer(s.soll, a.soll)
                for n in a.stoerer:
                    if n not in s.stoerer:
                        s.stoerer.append(n)


def _anker(ctx: Kontext, pos: _Pos, s: Section) -> int:
    """Zeile in vorher.xlsx, NACH der eine neue Section eingefuegt wird."""
    i = ctx.sektionen.index(s)
    for t in reversed(ctx.sektionen[:i]):
        if t.im_zustand(False):
            return pos.sec_ende[t.id]
    return 1


def _eintraege(ctx: Kontext, pos: _Pos) -> list[dict[str, Any]]:
    modus = BASIS[ctx.layout["base"]][0]
    out: list[dict[str, Any]] = []

    def neu(typ: str, zeile: int, spalte: int, woerter: int, soll: str, stoerer: list[str], seite: int) -> None:
        out.append({"nr": len(out) + 1, "typ": typ, "blatt": pos.blatt, "zeile": zeile, "spalte": spalte,
                    "woerter": woerter, "soll_status": soll, "stoerer": sorted(stoerer), "seite": seite})

    for s in ctx.sektionen:
        if s.status:
            w = markierte_woerter(s.kopf_markup) + sum(a.markiert() for a in s.absaetze)
            if s.status == "neu":
                neu("section_neu", _anker(ctx, pos, s), pos.nr_spalte, w, s.soll, s.stoerer, s.seite)
            else:
                neu("section_geloescht", pos.sec_kopf[s.id], pos.nr_spalte, w, s.soll, s.stoerer, s.seite)
            continue
        vorhergehend = pos.sec_kopf[s.id]
        for a in s.absaetze:
            if not a.typ:
                if a.id in pos.absatz:
                    vorhergehend = pos.absatz[a.id] if modus == "absatz" else vorhergehend
                continue
            if a.typ == "absatz_neu":
                if modus == "absatz":
                    neu(a.typ, vorhergehend, pos.text_spalte, a.woerter, a.soll, a.stoerer, a.seite)
                else:
                    neu(a.typ, pos.sec_text[s.id], pos.text_spalte, a.woerter, a.soll, a.stoerer, a.seite)
            else:
                neu(a.typ, pos.absatz[a.id], pos.text_spalte, a.woerter, a.soll, a.stoerer, a.seite)
                if modus == "absatz":
                    vorhergehend = pos.absatz[a.id]
    return out


# ====================================================================== Hauptfunktion
def erzeuge_paar(seed: int, profil: str | dict = "standard", ziel: Path | str = ".",
                 korpus: str = "entwicklung") -> Path:
    """Erzeugt ein Paar unter ziel/synth_<seed>_<profil>/ und gibt den Paar-Ordner zurueck.

    ziel   Korpus-Ordner (z. B. data/paare/entwicklung); wird bei Bedarf angelegt.
    korpus "entwicklung" oder "test". Nur bei "test" kommen die NUR_TEST-Stoerer-Familien vor.
    Bestehende Paar-Ordner werden nie ueberschrieben (FileExistsError, Regel 29).
    """
    if korpus not in ("entwicklung", "test"):
        raise ValueError("korpus muss 'entwicklung' oder 'test' sein")
    name, cfg = _profil(profil)
    ordner = Path(ziel) / f"synth_{seed:04d}_{name}"
    if ordner.exists():
        raise FileExistsError(f"{ordner} existiert schon (Paare sind append-only)")
    ctx, pdf = _baue(seed, name, cfg, korpus)
    tmp = Path(ziel) / f".tmp_{ordner.name}"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    try:
        pdf.speichern(tmp / "aenderungen.pdf")
        _luecken_sammeln(ctx, pdf)
        _bewerte_struktur(ctx)
        pos_alt = schreibe_excel(tmp / "vorher.xlsx", ctx, False)
        schreibe_excel(tmp / "nachher.xlsx", ctx, True)
        eintraege = _eintraege(ctx, pos_alt)
        markiert = (sum(markierte_woerter(s.kopf_markup) for s in ctx.sektionen)
                    + sum(a.markiert() for _, a in ctx.alle_absaetze()))
        erwartet = {
            "format": FORMAT,
            "generator_version": GENERATOR_VERSION,
            "seed": seed,
            "profil": name,
            "korpus": korpus,
            "stoerer": list(ctx.aktiv),
            "layout": _layout_name(ctx),
            "woerter_markiert": markiert,
            "aenderungen": eintraege,
            "luecken": sorted(ctx.luecken, key=lambda x: (x["seite"], x["art"])),
            "pdf_info": {
                "seiten": pdf.seitenzahl,
                "trennungen": pdf.stats["trennungen"],
                "fussnoten": sum(len(sx.fussnoten) for sx in pdf.seiten),
                "links": pdf.stats["links"],
                "umbrueche_im_absatz": sum(1 for _, a in ctx.alle_absaetze() if a.art == "text" and a.seite_ende > a.seite),
                "gedrehte_seite": ctx.opt.get("rotiert_seite", 0),
            },
        }
        with open(tmp / "erwartet.json", "w", encoding="utf-8") as f:
            json.dump(erwartet, f, ensure_ascii=False, indent=2)
            f.write("\n")
        tmp.rename(ordner)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise
    return ordner


def _baue(seed: int, name: str, cfg: dict[str, Any], korpus: str, versuche: int = 8) -> tuple[Kontext, PdfDokument]:
    """Baut Modell und PDF. Fehlt eine zugesagte Besonderheit (z. B. es wurde nirgends getrennt), gibt es einen
    neuen Versuch mit anderem, aber festem Zufallsstrom -- das Ergebnis bleibt je Seed deterministisch."""
    for versuch in range(versuche):
        ctx = Kontext(seed, name, cfg, korpus, versuch)
        ctx.lauf()
        pdf = rendere_pdf(ctx)
        fehlt = list(ctx.fehlend) + [n for n in ctx.aktiv
                                     if st.REGISTER[n].garantie and not st.REGISTER[n].garantie(ctx, pdf)]
        if not fehlt:
            return ctx, pdf
    ctx.aktiv = [n for n in ctx.aktiv if n not in fehlt]       # nur behaupten, was wirklich drin ist
    return ctx, pdf


def _layout_name(ctx: Kontext) -> str:
    lay = ctx.layout
    teile = [lay["base"]]
    for k in ("lfd", "kommentar", "zahl", "formeln", "merges", "banner", "versteckt"):
        if lay[k]:
            teile.append(k)
    return "+".join(teile)
