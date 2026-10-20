"""Generator v3 `word_markup`: erfundene Paare im Aufbau eines Word-PDFs mit Aenderungsverfolgung und einer CDE-Registerliste
(ADR 0006/0007; Regel 16: nur erfundene Daten -- Projektcodes, Spezifikationsnummern und Woerter sind ausgedacht).

Aus einem LOGISCHEN MODELL (Absaetze mit Gliederung, Aenderungen als Zeichen-Bereiche) entsteht ein Paar-Ordner nach
docs/paar_format.md:  aenderungen.pdf, vorher.xlsx, nachher.xlsx, erwartet.json

  PDF     direkt gezeichnet (pdf_word.py): Times 8,4 pt, Einzugsebenen 55/82/109/137/164/191/218, neu = Autorfarbe + Unterstreichung,
          entfaellt = dieselbe Farbe + Durchstreichung, gruene Listennummern "c.a.", gelbe Platzhalter, Aenderungsbalken, graue
          Markup-Leiste mit Sprechblasen und gestrichelten Verbindungslinien (oder "nur inline"), Kopf-/Fusszeilen mit eigenen
          Aenderungen, Vorspann (Deckblatt, Pruef- und Revisionstabelle, Bearbeiter-Hinweise, ggf. Inhaltsverzeichnis).
  vorher  CDE-Register: Kopfbereich mit Verbuenden und Zaehlformeln, Beispielzeile, eine Zeile je Absatz/Listenpunkt,
          Bedingte Formatierung, Datenpruefung, Autofilter, Druckbereich, Zeichnung, Zusatzteile (customXml, LabelInfo,
          Druckereinstellungen, calcChain) -- per ZIP-Nachbearbeitung nach dem openpyxl-Speichern.
  nachher wie Nico es macht: geaenderte Zelle = PDF-Absatz inline, alt = FF00B0F0 + durchgestrichen, neu = FFFF0000, kein Trenner,
          Grundschrift in jedem Lauf; neue Absaetze als neue Zeile (rot), gestrichene bleiben (hellblau durchgestrichen).

Die Wahrheit (erwartet.json, beide Excel-Dateien) kommt nur aus dem Modell, nie aus einem Leser. Dieses Modul importiert nichts
aus src/. Gleicher Seed => gleiche Dateien (die Excel-Dateien sind bytegleich, das ZIP bekommt feste Zeitstempel).
"""
from __future__ import annotations

import io
import json
import random
import re
import shutil
import zipfile
import zlib
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import pdf_word as pw

GENERATOR_VERSION = "3"
PROFIL = "word_markup"
KORPUS = "nachbau_word"
LEGENDE = "auto"                  # KORPUS.json: Legende, mit der das Tool diese Paare lesen soll (ADR 0006)
NACHHER_FORMAT = "markiert"       # KORPUS.json: nachher.xlsx traegt die Markierung (ADR 0005, Nachtrag 18.10.)
FORMAT = 1
ZEILEN_STANDARD = (300, 550)      # Datenzeilen in vorher.xlsx (wie gemessen); Tests nehmen kleinere Paare
ALT_FARBE, NEU_FARBE = "FF00B0F0", "FFFF0000"           # Markierung dieser Runde (ADR 0005, Nachtrag)
FRUEHER_ALT, FRUEHER_NEU = "FF0070C0", "FFC00000"       # Markierung einer frueheren Runde (bleibt unangetastet)
KENNWORT = "Gliederungszeile"     # Spalte I bei Ueberschriftzeilen (erfundenes Kennwort)
DATEN_START = 8                   # erste Datenzeile (1-4 Kopfbereich, 5/6 Spaltenkoepfe, 7 Beispielzeile)
LETZTE_SPALTE = 28                # AB
G, E_, F_ = 7, 5, 6               # Text-, Section-, Sub-Section-Spalte
HAUPTBLATT_NR = 4                 # Hinweise, Deckblatt, Unterschriften, <Register>, Listen

# ====================================================================== Textbausteine (alles erfunden)
PLURAL = ["Pumpen", "Ventile", "Schieber", "Rohrleitungen", "Behälter", "Armaturen", "Flansche", "Dichtungen", "Antriebe",
          "Lager", "Getriebe", "Kupplungen", "Sensoren", "Messumformer", "Schaltschränke", "Kabeltrassen", "Halterungen",
          "Gehäuse", "Filter", "Lüfter", "Absperrklappen", "Rückschlagventile", "Wärmetauscher", "Zebrafink-Nistkästen"]
ZUSTAND = ["korrosionsgeschützt", "wartungsfreundlich", "frostsicher", "schwingungsarm", "druckfest", "staubdicht",
           "geräuscharm", "austauschbar"]
PRAED = ["regelmäßig gewartet werden", "gegen Überlast geschützt sein", "eindeutig beschriftet sein",
         "ohne Werkzeug zugänglich sein", "vor Feuchtigkeit geschützt werden", "einzeln prüfbar sein",
         "fest verschraubt sein", "frei von Beschädigungen geliefert werden"]
ZUBEHOER = ["Befestigungsmaterial", "Ersatzdichtungen", "Montageanleitung", "Typenschild", "Prüfprotokoll", "Sonderwerkzeug",
            "Wartungsplan"]
GROESSEN = [("Nennweite", "mm"), ("Förderhöhe", "m"), ("Betriebsdruck", "bar"), ("Leistungsaufnahme", "kW"), ("Masse", "kg"),
            ("Fördermenge", "m³/h"), ("Umgebungstemperatur", "°C")]
ZEIT = ["jährlich", "halbjährlich", "monatlich", "wöchentlich"]
PARTIZIP = ["gereinigt", "geprüft", "dokumentiert", "gekennzeichnet", "gespült", "konserviert"]
SCHUTZ = ["IP54", "IP55", "IP65", "IP66"]
FRIST = ["vier Wochen", "zwei Wochen", "zehn Arbeitstage", "sechs Wochen"]
ADVERB = ["mindestens", "zusätzlich", "jeweils", "ausschließlich", "vollständig", "nachweislich", "dauerhaft", "rechtzeitig"]
WERTE = ["10", "16", "25", "40", "63", "80", "100", "125", "150", "250", "400"]
NEUWERTE = ["12,5", "42,9", "48", "65", "90", "110", "160", "320", "DN 80", "DN 150", "PN 16"]
PLATZ_NEU = {"[Wert]": "42,9", "[Anzahl]": "drei", "[Fabrikat]": "Fabrikat QZ-200", "[Norm]": "Werknorm WN 4711",
             "[…]": "gemäß Datenblatt", "[Datum]": "2026_03_01", "[Typ]": "Bauart K7"}
NORMEN = ["Werknorm WN 1001", "Werknorm WN 2040", "Richtlinie QZ 45", "Prüfvorschrift PV 12", "Merkblatt MB 7"]
SUB_TITEL = ["Bezugsnormen:", "Unterlagen:", "Lieferumfang:", "Leistungsumfang:", "Qualitätsnachweise:", "Werkstoffe:",
             "Ersatzteile:", "Kennzeichnung:", "Prüfungen:", "Schnittstellen:", "Abstimmung:", "Lagerung:"]
LISTE_TITEL = ["Allgemeine Anforderungen:", "Ausführung:", "Nachweise:", "Zubehör:", "Grenzwerte:", "Bauteile:",
               "Zusätzliche Angaben:", "Ausnahmen:"]
SECTION_TITEL = {
    1: ["ZUSAMMENFASSUNG", "BEZUGSNORMEN", "DEFINITIONEN", "UNTERLAGEN", "QUALITÄTSSICHERUNG", "LIEFERUNG UND LAGERUNG",
        "GEWÄHRLEISTUNG", "KOORDINATION", "ABNAHME", "SCHULUNG", "BAUSTELLENBEDINGUNGEN", "WARTUNGSVERTRAG"],
    2: ["HERSTELLER", "PUMPEN", "ARMATUREN", "ANTRIEBE", "WERKSTOFFE", "KENNZEICHNUNG", "STEUERUNG", "ZUBEHÖR",
        "OBERFLÄCHENSCHUTZ", "ERSATZTEILE"],
    3: ["VORBEREITUNG", "MONTAGE", "PRÜFUNG VOR ORT", "INBETRIEBNAHME", "REINIGUNG", "ÜBERGABE", "EINWEISUNG",
        "SCHUTZ DER ANLAGE"],
}
PART_TITEL = {1: "ALLGEMEINES", 2: "PRODUKTE", 3: "AUSFÜHRUNG"}
SPEZ_TITEL = ["PUMPENANLAGE", "ARMATUREN UND ROHRLEITUNGEN", "LÜFTUNGSGERÄTE", "WASSERAUFBEREITUNG", "FÖRDERTECHNIK",
              "DRUCKERHÖHUNGSANLAGE"]
PROJEKTE = ["XYZ01", "QRS07", "MNO42", "KLM15", "TUV88"]
ORTE = ["Pumpwerk Nord", "Anlage Süd", "Werk Ostfeld", "Station Mühlbach", "Standort Zebrafink"]
INITIALEN = ["QX", "ZB", "MV", "KR"]
KOMMENTARE = ["Bitte Wert mit dem Hersteller abstimmen.", "Gilt auch für die Reserveeinheit?", "Angabe aus dem Datenblatt übernommen.",
              "Formulierung mit der Fachplanung klären.", "Nachweis liegt noch nicht vor, bitte nachreichen.",
              "Ist hier die Variante mit Edelstahl gemeint? Die Ausschreibung nennt beide Ausführungen."]
FORMATIERUNGEN = ["Nicht Hervorheben", "Schriftart: Fett", "Schriftart: Nicht Fett", "Hervorheben", "Schriftart: Kursiv",
                  "Einzug: Links: 1,9 cm"]
GENERISCH = ["einschließlich", "sowie", "gemeinsam", "einzeln", "getrennt"]
ROLLEN = ["Hersteller", "Lieferant", "Betreiber", "Monteur", "Auftraggeber", "Auftragnehmer", "Prüfer"]

_KATEGORIEN: dict[str, list[str]] = {}
for _liste in (PLURAL, ZUSTAND, ZUBEHOER, ZEIT, PARTIZIP, ADVERB, SCHUTZ):
    for _w in _liste:
        if " " not in _w:
            _KATEGORIEN[_w] = _liste
_WORT = re.compile(r"[A-Za-zÄÖÜäöüß][A-Za-zÄÖÜäöüß\-]{2,}")
_ZAHL = re.compile(r"(?<![\w,\[])\d+(?:,\d+)?(?![\w,\]])")
_PLATZ = re.compile(r"\[[^\]]{1,20}\]")


# ====================================================================== Modell
@dataclass
class Edit:
    """Eine Aenderung im Text eines Absatzes: Zeichen [start, ende) des alten Textes entfallen, `neu` kommt an ihre Stelle."""
    start: int
    ende: int
    neu: str
    autor: int = 0
    art: str = ""                 # Kennung fuer erwartet.json (wort_ersetzt, eingefuegt, ...)


@dataclass
class Absatz:
    id: int
    art: str                      # part | section | sub | sub_titel | liste | liste_titel
    ebene: int                    # Einzugsebene des Labels (part/section 0, sub 1, Liste 2..5)
    text: str = ""                # alter Text (ohne Label); bei neuen Absaetzen ""
    nummer: str = ""              # part "1", section "1.1"
    label_alt: str = ""
    label_neu: str = ""
    sub: str = ""                 # Spalte F (Buchstabe der Sub-Section)
    section: str = ""             # Spalte E
    status: str | None = None     # None | neu | geloescht
    edits: list[Edit] = field(default_factory=list)
    eltern: int | None = None     # Liste: id des Eltern-Absatzes
    stufe: int = 0                # Liste: 1 = "1.", 2 = "a.", 3 = "1)", 4 = "a)"
    fett: bool = False            # ganze Zelle fett (Grundschrift)
    ueberschrift: bool = False    # Ueberschriftzeile (fett, Kennwort in I)
    excel_alt: str | None = None  # Text in vorher.xlsx, wenn die Excel vom PDF abweicht
    abweichung: tuple[int, int, str] | None = None   # (start, ende, Excel-Text) der Abweichung im alten PDF-Text
    trennen: dict[str, int] = field(default_factory=dict)   # nur nachbau_test: Wort -> Trennstelle am Zeilenende (Silbentrennung)
    umbruch_erzwingen: bool = False                   # nur nachbau_test: Absatz moeglichst ueber den Seitenumbruch setzen
    blase_lang: bool = False                          # nur nachbau_test: lange mehrzeilige Sprechblase am Absatz
    autor: int = 0
    stoerer: list[str] = field(default_factory=list)
    seite: int = 0
    seite_ende: int = 0

    @property
    def gelb(self) -> list[tuple[int, int]]:
        return [m.span() for m in _PLATZ.finditer(self.text)]

    def neu_text(self) -> str:
        out, pos = [], 0
        for e in sorted(self.edits, key=lambda x: (x.start, x.ende)):
            out.append(self.text[pos:e.start])
            out.append(e.neu)
            pos = e.ende
        out.append(self.text[pos:])
        return "".join(out)

    def stuecke(self) -> list[tuple[str, str, int, bool]]:
        """Der Absatz im PDF, inline: [(art k|d|a, Text, Autor, gelb)]; Platzhalter (alt) sind gelb hinterlegt."""
        roh: list[tuple[str, int, int, int, str]] = []     # (art, start, ende, autor, neuer Text) im alten Text
        pos = 0
        for e in sorted(self.edits, key=lambda x: (x.start, x.ende)):
            if e.start > pos:
                roh.append(("k", pos, e.start, 0, ""))
            if e.ende > e.start:
                roh.append(("d", e.start, e.ende, e.autor, ""))
            if e.neu:
                roh.append(("a", e.ende, e.ende, e.autor, e.neu))
            pos = max(pos, e.ende)
        if pos < len(self.text):
            roh.append(("k", pos, len(self.text), 0, ""))
        gelb = self.gelb
        out: list[tuple[str, str, int, bool]] = []
        for art, a, b, autor, neu in roh:
            if art == "a":
                out.append(("a", neu, autor, False))
                continue
            grenzen = sorted({a, b} | {x for s, t in gelb for x in (s, t) if a < x < b})
            for x, y in zip(grenzen, grenzen[1:]):
                hinterlegt = any(s <= x and y <= t for s, t in gelb)
                out.append((art, self.text[x:y], autor, hinterlegt))
        return [s for s in out if s[1]]

    def label_stuecke(self) -> list[tuple[str, str, int]]:
        """Label im PDF: [(Text, art k|d|a|gd|ga, Autor)] -- gd/ga = gruene Nummern-Aenderung."""
        if self.art == "part" or not (self.label_alt or self.label_neu):
            return []
        if self.status == "neu":
            return [(self.label_neu, "a", self.autor)]
        if self.status == "geloescht":
            return [(self.label_alt, "d", self.autor)]
        if self.label_alt != self.label_neu:
            return [(self.label_alt, "gd", 0), (self.label_neu, "ga", 0)]
        return [(self.label_alt, "k", 0)]


@dataclass
class Kontext:
    seed: int
    rng: random.Random
    cfg: dict[str, Any]
    absaetze: list[Absatz] = field(default_factory=list)
    bloecke: list[Any] = field(default_factory=list)      # Spec-Bloecke im Register (siehe Block), Ziel-Block = absaetze
    stoerer: list[str] = field(default_factory=list)
    opt: dict[str, Any] = field(default_factory=dict)
    benutzt: set[str] = field(default_factory=set)
    _zaehler: int = 0

    def neue_id(self) -> int:
        self._zaehler += 1
        return self._zaehler

    def markiere(self, name: str) -> None:
        if name not in self.stoerer:
            self.stoerer.append(name)

    # ---- Text -------------------------------------------------------------------------------------------------------
    def _wert(self, platzhalter_p: float = 0.25, art: str = "wert") -> str:
        """Ein Zahlenwert, oft als Platzhalter in eckigen Klammern ("[40]", "[Wert]", "[Anzahl]", "[…]")."""
        r = self.rng
        if r.random() < platzhalter_p:
            if r.random() < 0.6:
                return "[" + r.choice(WERTE) + "]"
            return r.choice(["[Anzahl]", "[…]"] if art == "anzahl" else ["[Wert]", "[…]"])
        return r.choice(WERTE)

    def satz(self) -> str:
        r = self.rng
        p, p2 = r.sample(PLURAL, 2)
        k = r.randint(1, 11)
        if k == 1:
            return f"{p} sind {r.choice(ZUSTAND)} auszuführen."
        if k == 2:
            return f"Alle {p} müssen {r.choice(PRAED)}."
        if k == 3:
            return f"Der Auftragnehmer liefert {p} einschließlich {r.choice(ZUBEHOER)}."
        if k == 4:
            g, e = r.choice(GROESSEN)
            return f"Die {g} der {p} beträgt mindestens {self._wert(0.45)} {e}."
        if k == 5:
            norm = r.choice(NORMEN) if r.random() < 0.7 else "[Norm]"
            return f"Prüfungen der {p} erfolgen {r.choice(ZEIT)} nach {norm}."
        if k == 6:
            return f"{p} werden vor der Montage {r.choice(PARTIZIP)}."
        if k == 7:
            return "Abweichungen von dieser Anforderung bedürfen der schriftlichen Zustimmung des Auftraggebers."
        if k == 8:
            return f"Für {p} und {p2} gilt die Schutzart {r.choice(SCHUTZ)}."
        if k == 9:
            return f"Die Lieferung umfasst {self._wert(0.35, 'anzahl')} {p} sowie {r.choice(ZUBEHOER)}."
        if k == 10:
            return f"Unterlagen zu {p} sind {r.choice(FRIST)} vor der Lieferung einzureichen."
        g, e = r.choice(GROESSEN)
        return f"Die {g} darf {self._wert(0.4)} {e} nicht überschreiten, sofern {p} {r.choice(ZEIT)} {r.choice(PARTIZIP)} werden."

    def absatz_text(self, saetze: tuple[int, int] = (1, 3)) -> str:
        for _ in range(40):
            t = " ".join(self.satz() for _ in range(self.rng.randint(*saetze)))
            if t not in self.benutzt:
                self.benutzt.add(t)
                return t
        t += f" Ergänzung {self.rng.randint(2, 99)}."
        self.benutzt.add(t)
        return t

    def punkt_text(self) -> str:
        r = self.rng
        for _ in range(40):
            k = r.randint(1, 6)
            p, p2 = r.sample(PLURAL, 2)
            if k == 1:
                t = f"{p} und {p2}."
            elif k == 2:
                t = f"Werkstoffnachweise für {p}."
            elif k == 3:
                g, e = r.choice(GROESSEN)
                t = f"{g}: {self._wert(0.4)} {e}."
            elif k == 4:
                t = f"{r.choice(ZUBEHOER)} für {p}."
            elif k == 5:
                t = f"{p} {r.choice(ZEIT)} {r.choice(PARTIZIP)}."
            else:
                t = self.satz()
            if t not in self.benutzt:
                self.benutzt.add(t)
                return t
        t = f"{t[:-1]} {r.randint(2, 99)}."
        self.benutzt.add(t)
        return t


# ====================================================================== Gliederung
def sub_label(i: int) -> str:
    """0 -> A ... 25 -> Z, 26 -> AA, 27 -> BB ... (wie Word)."""
    return chr(65 + i) if i < 26 else chr(65 + (i - 26) % 26) * (2 + (i - 26) // 26)


def listen_label(stufe: int, i: int) -> str:
    if stufe in (1, 3):
        return f"{i + 1}." if stufe == 1 else f"{i + 1})"
    buchstabe = chr(97 + i % 26) * (1 + i // 26)
    return f"{buchstabe}." if stufe == 2 else f"{buchstabe})"


def _baue_struktur(ctx: Kontext, ziel: int) -> None:
    r = ctx.rng
    anteile = (0.45, 0.30, 0.25)
    doppel_section = ctx.opt["doppelbuchstaben"]
    zeilen = 0
    for p in (1, 2, 3):
        ctx.absaetze.append(Absatz(ctx.neue_id(), "part", 0, PART_TITEL[p], nummer=str(p), section=f"Part {p}",
                                   ueberschrift=True))
        zeilen += 1
        ziel_p = zeilen + max(6, int(ziel * anteile[p - 1]))
        titel = list(SECTION_TITEL[p])
        r.shuffle(titel)
        s = 0
        while zeilen < ziel_p or s == 0:
            s += 1
            nummer = f"{p}.{s}"
            ctx.absaetze.append(Absatz(ctx.neue_id(), "section", 0, titel[(s - 1) % len(titel)], nummer=nummer,
                                       label_alt=nummer, label_neu=nummer, section=nummer, ueberschrift=True))
            zeilen += 1
            if doppel_section and p == 1 and s == 2:
                n_sub = r.randint(27, 29)
                ctx.markiere("subsection_doppelbuchstabe")
            else:
                n_sub = r.randint(2, 7)
            for i in range(n_sub):
                zeilen += _baue_sub(ctx, nummer, i, kurz=n_sub > 20)
                if zeilen >= ziel_p and i >= 1:
                    break


def _kuerze(ctx: Kontext, ziel: int) -> None:
    """Schiesst die Gliederung ueber das Ziel hinaus, entfallen letzte Sub-Sections bzw. letzte Listenpunkte (samt Unterpunkten)
    -- immer die letzten, damit die Nummern der uebrigen bleiben; gewaehlt wird, was dem Ziel am naechsten kommt. Jede Section
    behaelt mindestens zwei Sub-Sections (die mit Doppelbuchstaben mindestens 27), jede Liste mindestens zwei Punkte."""
    while True:
        ueber = len(ctx.absaetze) - ziel
        if ueber <= 0:
            return
        subs: dict[str, list[Absatz]] = {}
        for a in ctx.absaetze:
            if a.art in ("sub", "sub_titel"):
                subs.setdefault(a.section, []).append(a)
        gruppen: dict[int, list[Absatz]] = {}
        for a in ctx.absaetze:
            if a.art in ("liste", "liste_titel"):
                gruppen.setdefault(a.eltern or 0, []).append(a)
        letzte = [v[-1] for v in subs.values() if len(v) > (27 if len(v) >= 27 else 2)]
        letzte += [v[-1] for v in gruppen.values() if len(v) > 2]          # auch letzte Listenpunkte (samt Unterpunkten)
        kand = []
        for a in letzte:
            i = ctx.absaetze.index(a)
            j = ctx.absaetze.index(_letzter_nachfahre(ctx, a))
            kand.append((abs(j + 1 - i - ueber), j + 1 - i, i, j))
        if not kand:
            return
        abstand, _groesse, i, j = min(kand)
        if abstand >= ueber:                                   # jede Kuerzung waere weiter vom Ziel weg
            return
        del ctx.absaetze[i:j + 1]


def _baue_sub(ctx: Kontext, section: str, i: int, kurz: bool = False) -> int:
    r = ctx.rng
    lab = sub_label(i)
    titel = (not kurz) and r.random() < 0.35
    if titel:
        a = Absatz(ctx.neue_id(), "sub_titel", 1, r.choice(SUB_TITEL), label_alt=lab + ".", label_neu=lab + ".", sub=lab,
                   section=section, ueberschrift=True)
    else:
        a = Absatz(ctx.neue_id(), "sub", 1, ctx.absatz_text((1, 2) if kurz else (1, 3)), label_alt=lab + ".",
                   label_neu=lab + ".", sub=lab, section=section)
    ctx.absaetze.append(a)
    n = 1
    if titel or (not kurz and r.random() < 0.45):
        n += _baue_liste(ctx, a, 1, r.randint(2, 6) if titel else r.randint(1, 5))
    return n


def _baue_liste(ctx: Kontext, eltern: Absatz, stufe: int, anzahl: int) -> int:
    r = ctx.rng
    n = 0
    for j in range(anzahl):
        kinder = stufe < 4 and r.random() < (0.22, 0.25, 0.3, 0.0)[stufe - 1]
        titel = kinder and r.random() < 0.5
        lab = listen_label(stufe, j)
        a = Absatz(ctx.neue_id(), "liste_titel" if titel else "liste", stufe + 1,
                   r.choice(LISTE_TITEL) if titel else ctx.punkt_text(), label_alt=lab, label_neu=lab, sub=eltern.sub,
                   section=eltern.section, eltern=eltern.id, stufe=stufe, ueberschrift=titel)
        ctx.absaetze.append(a)
        n += 1
        if kinder:
            n += _baue_liste(ctx, a, stufe + 1, r.randint(2, 4))
    return n


def _nummeriere(ctx: Kontext) -> None:
    """Listen-Labels alt (ohne neue Punkte) und neu (ohne gestrichene Punkte) je Geschwister-Gruppe."""
    gruppen: dict[tuple[int, int], list[Absatz]] = {}
    for a in ctx.absaetze:
        if a.art in ("liste", "liste_titel"):
            gruppen.setdefault((a.eltern or 0, a.stufe), []).append(a)
    for (_, stufe), liste in gruppen.items():
        i_alt = i_neu = 0
        for a in liste:
            a.label_alt = listen_label(stufe, i_alt) if a.status != "neu" else ""
            a.label_neu = listen_label(stufe, i_neu) if a.status != "geloescht" else ""
            if a.status != "neu":
                i_alt += 1
            if a.status != "geloescht":
                i_neu += 1


# ====================================================================== Aenderungen
def _woerter(text: str) -> list[tuple[int, int]]:
    return [m.span() for m in _WORT.finditer(text) if not any(s <= m.start() < t for s, t in [p.span() for p in _PLATZ.finditer(text)])]


def _frei(a: Absatz, s: int, e: int, abstand: int = 2) -> bool:
    return all(e + abstand <= x.start or s >= x.ende + abstand for x in a.edits)


def _ersatz(ctx: Kontext, wort: str) -> str:
    r = ctx.rng
    pool = _KATEGORIEN.get(wort)
    if pool:
        andere = [w for w in pool if w != wort and " " not in w]
        if andere:
            return r.choice(andere)
    if wort[:1].isupper():
        return r.choice([w for w in ROLLEN if w != wort])
    return r.choice([w for w in ADVERB + GENERISCH if w != wort])


def _autor(ctx: Kontext) -> int:
    return 1 if ctx.opt["zwei_autoren"] and ctx.rng.random() < 0.35 else 0


def _aendere_text(ctx: Kontext, a: Absatz, art: str) -> bool:
    """Eine Text-Aenderung im Absatz `a` (zeichengenau). False = passt hier nicht."""
    r = ctx.rng
    t = a.text
    autor = _autor(ctx)
    woerter = [w for w in _woerter(t) if w[0] > 0 and _frei(a, *w)]
    if art == "wort_ersetzt":
        if not woerter:
            return False
        s, e = r.choice(woerter)
        neu = _ersatz(ctx, t[s:e])
        if r.random() < 0.3 and e < len(t) and t[e] == " ":
            a.edits.append(Edit(s, e + 1, neu + " ", autor, art))           # mit Leerzeichen gestrichen und neu geschrieben
        else:
            a.edits.append(Edit(s, e, neu, autor, art))                     # direkt hintereinander ("altneu")
        return True
    if art == "zahl_ersetzt":
        zahlen = [m.span() for m in _ZAHL.finditer(t) if _frei(a, *m.span())]
        if not zahlen:
            return False
        s, e = r.choice(zahlen)
        a.edits.append(Edit(s, e, r.choice([w for w in NEUWERTE if w != t[s:e]]), autor, art))
        return True
    if art == "eingefuegt":
        if not woerter:
            return False
        s, e = r.choice(woerter)
        if r.random() < 0.5:
            a.edits.append(Edit(e, e, " " + r.choice(ADVERB), autor, art))
        else:
            a.edits.append(Edit(s, s, r.choice(ADVERB) + " ", autor, art))
        return True
    if art == "satz_eingefuegt":
        if not _frei(a, len(t), len(t)):
            return False
        a.edits.append(Edit(len(t), len(t), " " + ctx.satz(), autor, art))
        return True
    if art == "geloescht":
        kand = [(s, e) for s, e in woerter if s > 0 and t[s - 1] == " " and _frei(a, s - 1, e)]
        if not kand:
            return False
        s, e = r.choice(kand)
        a.edits.append(Edit(s - 1, e, "", autor, art))
        return True
    if art == "satzteil_geloescht":
        saetze = [m.start() for m in re.finditer(r"(?<=[.:]) ", t)]
        if saetze and _frei(a, saetze[-1], len(t)):
            a.edits.append(Edit(saetze[-1], len(t), "", autor, art))        # letzter Satz entfaellt (mit Leerzeichen davor)
            return True
        komma = t.find(", ")
        if komma > 0 and _frei(a, komma, len(t) - 1):
            a.edits.append(Edit(komma, len(t) - 1, "", autor, art))         # halber Satz bis vor den Punkt
            return True
        return False
    if art == "satzzeichen":
        punkte = [m.start() for m in re.finditer(r"[,:]", t) if _frei(a, m.start(), m.start() + 1)]
        if punkte and r.random() < 0.5:
            p = r.choice(punkte)
            a.edits.append(Edit(p, p + 1, ";" if t[p] == "," else "", autor, art))
            return True
        if not woerter:
            return False
        s, e = r.choice(woerter)
        if e < len(t) and t[e] == " ":
            a.edits.append(Edit(e, e, ",", autor, art))
            return True
        return False
    if art in ("klammern", "platzhalter_ersetzt"):
        plaetze = [m.span() for m in _PLATZ.finditer(t) if _frei(a, *m.span())]
        if not plaetze:
            return False
        s, e = r.choice(plaetze)
        inhalt = t[s + 1:e - 1]
        if art == "klammern" and inhalt.isdigit():
            a.edits.append(Edit(s, s + 1, "", autor, art))                  # nur "[" und "]" entfallen
            a.edits.append(Edit(e - 1, e, "", autor, art))
            return True
        neu = PLATZ_NEU.get(t[s:e]) or r.choice([w for w in NEUWERTE if w != inhalt])
        a.edits.append(Edit(s, e, neu, autor, "platzhalter_ersetzt"))       # "[40]" gestrichen, "42,9" direkt dahinter
        return True
    raise ValueError(art)


TEXT_ARTEN = {"wort_ersetzt": 30, "zahl_ersetzt": 8, "eingefuegt": 12, "satz_eingefuegt": 6, "geloescht": 10,
              "satzteil_geloescht": 6, "satzzeichen": 8, "klammern": 8, "platzhalter_ersetzt": 10}
STRUKTUR_ARTEN = {"listenpunkt_weg": 8, "listenpunkt_neu": 7, "absatz_weg": 4, "subsection_neu": 4, "ueberschrift": 3}


def _plane_aenderungen(ctx: Kontext, anzahl: int) -> None:
    r = ctx.rng
    arten = list(TEXT_ARTEN) + list(STRUKTUR_ARTEN)
    gewichte = list(TEXT_ARTEN.values()) + list(STRUKTUR_ARTEN.values())
    erledigt = 0
    versuche = 0
    pflicht = ["klammern", "platzhalter_ersetzt", "listenpunkt_weg", "listenpunkt_neu", "subsection_neu", "absatz_weg"]
    while erledigt < anzahl and versuche < anzahl * 12:
        versuche += 1
        art = pflicht.pop(0) if pflicht else r.choices(arten, gewichte)[0]
        if art in STRUKTUR_ARTEN:
            ok = _struktur_aenderung(ctx, art)
        else:
            kand = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None and len(a.edits) < 3]
            if art in ("klammern", "platzhalter_ersetzt"):
                kand = [a for a in kand if _PLATZ.search(a.text)]
            if not kand:
                continue
            a = r.choice(kand)
            ok = _aendere_text(ctx, a, art)
        erledigt += bool(ok)


def _einfuegen_nach(ctx: Kontext, vorgaenger: Absatz, neu: Absatz) -> None:
    i = ctx.absaetze.index(vorgaenger)
    ctx.absaetze.insert(i + 1, neu)


def _letzter_nachfahre(ctx: Kontext, a: Absatz) -> Absatz:
    """Letzter Absatz im Teilbaum von a (Listen-Kinder und deren Kinder eingeschlossen)."""
    i = ctx.absaetze.index(a)
    ids = {a.id}
    letzt = a
    for b in ctx.absaetze[i + 1:]:
        if b.eltern is None or b.eltern not in ids:
            break
        ids.add(b.id)
        letzt = b
    return letzt


def _struktur_aenderung(ctx: Kontext, art: str) -> bool:
    r = ctx.rng
    autor = _autor(ctx)
    hat_kinder = {b.eltern for b in ctx.absaetze if b.eltern}
    if art == "listenpunkt_weg":
        kand = [a for a in ctx.absaetze if a.art == "liste" and a.status is None and not a.edits and a.id not in hat_kinder]
        if not kand:
            return False
        a = r.choice(kand)
        a.status, a.autor = "geloescht", autor
        a.edits = [Edit(0, len(a.text), "", autor, "absatz_geloescht")]
        return True
    if art == "absatz_weg":
        kand = [a for a in ctx.absaetze if a.art == "sub" and a.status is None and not a.edits and a.id not in hat_kinder]
        # nur die letzte Sub-Section einer Section (sonst wuerden die Buchstaben der Folgenden neu vergeben)
        kand = [a for a in kand if _letzte_sub(ctx, a)]
        neu_in = {b.section for b in ctx.absaetze if b.art == "sub" and b.status == "neu"}
        kand = [a for a in kand if a.section not in neu_in]
        if not kand:
            return False
        a = r.choice(kand)
        a.status, a.autor = "geloescht", autor
        a.edits = [Edit(0, len(a.text), "", autor, "absatz_geloescht")]
        return True
    if art == "listenpunkt_neu":
        kand = [a for a in ctx.absaetze if a.art == "liste" and a.status is None and a.id not in hat_kinder]
        if not kand:
            return False
        vor = r.choice(kand)
        text = ctx.punkt_text()
        neu = Absatz(ctx.neue_id(), "liste", vor.ebene, "", sub=vor.sub, section=vor.section, eltern=vor.eltern,
                     stufe=vor.stufe, status="neu", autor=autor, edits=[Edit(0, 0, text, autor, "absatz_neu")])
        if r.random() < 0.4:                                   # vor dem Punkt einfuegen (die Folgenummern aendern sich)
            i = ctx.absaetze.index(vor)
            ctx.absaetze.insert(i, neu)
        else:
            _einfuegen_nach(ctx, vor, neu)
        return True
    if art == "subsection_neu":
        secs = [a for a in ctx.absaetze if a.art == "section" and a.status is None]
        r.shuffle(secs)
        for sec in secs:
            subs = [b for b in ctx.absaetze if b.section == sec.nummer and b.art in ("sub", "sub_titel")]
            if not subs or any(b.status for b in subs) or len(subs) >= 26:
                continue                                       # keine zweite neue und keine neben einer gestrichenen
            letzte = subs[-1]
            lab = sub_label(len(subs))
            neu = Absatz(ctx.neue_id(), "sub", 1, "", label_alt="", label_neu=lab + ".", sub=lab, section=sec.nummer,
                         status="neu", autor=autor, edits=[Edit(0, 0, ctx.absatz_text((1, 2)), autor, "subsection_neu")])
            ende = _letzter_nachfahre(ctx, letzte)
            _einfuegen_nach(ctx, ende, neu)
            if r.random() < 0.5:                               # mit neuen Listenpunkten darunter
                vor = neu
                for j in range(r.randint(1, 3)):
                    p = Absatz(ctx.neue_id(), "liste", 2, "", sub=lab, section=sec.nummer, eltern=neu.id, stufe=1,
                               status="neu", autor=autor, edits=[Edit(0, 0, ctx.punkt_text(), autor, "absatz_neu")])
                    _einfuegen_nach(ctx, vor, p)
                    vor = p
            ctx.markiere("subsection_neu")
            return True
        return False
    if art == "ueberschrift":
        kand = [a for a in ctx.absaetze if a.art in ("sub_titel", "liste_titel") and a.status is None and not a.edits]
        if not kand:
            return False
        a = r.choice(kand)
        s = 0
        a.edits.append(Edit(s, s, r.choice(["Weitere ", "Zusätzliche ", "Besondere "]), autor, "eingefuegt"))
        ctx.markiere("ueberschrift_aenderung")
        return True
    raise ValueError(art)


def _letzte_sub(ctx: Kontext, a: Absatz) -> bool:
    subs = [b for b in ctx.absaetze if b.section == a.section and b.art in ("sub", "sub_titel")]
    return bool(subs) and subs[-1] is a


def _abweichungen(ctx: Kontext) -> None:
    """Excel weicht vom PDF ab: (1) in einem geaenderten Absatz an ANDERER Stelle (-> pruefen mit fertigem Vorschlag),
    (2) in einem unveraenderten Absatz (-> das Tool darf die Zelle nicht anfassen)."""
    r = ctx.rng
    geaendert = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None and a.edits]
    ruhig = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None and not a.edits
             and a.label_alt == a.label_neu]
    for liste, n, name in ((geaendert, r.randint(1, 2), "excel_abweichung"), (ruhig, r.randint(0, 2), "excel_abweichung_ohne_aenderung")):
        r.shuffle(liste)
        gesetzt = 0
        for a in liste:
            if gesetzt >= n:
                break
            stellen = [w for w in _woerter(a.text) if w[0] > 0 and _frei(a, w[0], w[1], abstand=6)]
            if not stellen:
                continue
            gesetzt += 1
            s, e = r.choice(stellen)
            wort = a.text[s:e]
            neu = wort[:-1] if len(wort) > 4 and r.random() < 0.4 else _ersatz(ctx, wort)
            a.excel_alt = a.text[:s] + neu + a.text[e:]
            a.abweichung = (s, e, neu)
            a.stoerer.append(name)
            ctx.markiere(name)


# ====================================================================== Stoerer nur im ungesehenen Korpus nachbau_test
NUR_TEST_FAMILIEN = ("silbentrennung_zeilenende", "listenpunkt_umbruch_nummer")
KORPUS_TEST = "nachbau_test"
COMPOSITA = [("Wasser", "behälter"), ("Druck", "behälter"), ("Absperr", "klappenantrieb"), ("Kühlwasser", "pumpe"),
             ("Schaltschrank", "gehäuse"), ("Dichtungs", "material"), ("Rohrleitungs", "halterung"), ("Lüftungs", "anlage"),
             ("Antriebs", "welle"), ("Sicherheits", "ventil"), ("Förder", "pumpe"), ("Montage", "anleitung")]
FUELL = ["und", "oder", "sowie", "für", "mit", "an", "zu", "alle", "einzelne", "weitere"] + [w for w in PLURAL if " " not in w]


def _zeilen_von(text: str, ebene: int, trennen: dict[str, int] | None = None) -> tuple[list[str], int]:
    """Zeilen eines Absatzes wie der Setzer sie umbricht (nur Text) und Zahl der Trennungen."""
    x_text = EINZUG[min(ebene + 1, len(EINZUG) - 1)]
    zeilen, n = _umbrechen([(c, ("", pw.SCHWARZ, False)) for c in text], RECHTS - x_text, "Times-Roman", trennen)
    return ["".join(c for tok in z for c, _ in tok) for z in zeilen], n


def _familie_silbentrennung(ctx: Kontext) -> None:
    """Familie "silbentrennung_zeilenende": Ein zusammengesetztes Wort steht am Zeilenende und wird getrennt ("Wasser-" / "behaelter");
    entweder ist genau dieses Wort Teil der Aenderung (gestrichen und ersetzt, der Strich laeuft ueber den Zeilenumbruch) oder
    daneben wird ein Wort eingefuegt. Dazu eine lange, mehrzeilige Sprechblase. Die Excel hat das Wort ungetrennt."""
    r = ctx.rng
    kand = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None and not a.edits and a.excel_alt is None
            and a.label_alt == a.label_neu]
    r.shuffle(kand)
    ziel = r.randint(3, 6)
    for a in kand:
        if ziel <= 0:
            break
        for _versuch in range(80):
            w1, w2 = r.choice(COMPOSITA)
            wort = w1 + w2
            vorne = " ".join([r.choice(["Alle", "Die", "Sämtliche", "Weitere"])] + [r.choice(FUELL) for _ in range(r.randint(2, 9))])
            hinten = ctx.satz()
            text = f"{vorne} {wort} {hinten}"
            s0 = len(vorne) + 1
            e0 = s0 + len(wort)
            probe = Absatz(a.id, a.art, a.ebene, text, trennen={wort: len(w1)})
            if r.random() < 0.5:                                       # das Wort selbst wird ersetzt (Strich ueber den Zeilenumbruch)
                x1, x2 = r.choice([c for c in COMPOSITA if c[0] + c[1] != wort])
                probe.edits = [Edit(s0, e0 + 1, x1 + x2 + " ", _autor(ctx), "wort_ersetzt")]
            else:                                                      # davor wird ein Wort eingefuegt (der Umbruch verschiebt sich)
                probe.edits = [Edit(s0, s0, r.choice(ADVERB) + " ", _autor(ctx), "eingefuegt")]
            anzeige = "".join(t for _art, t, _au, _g in probe.stuecke())
            zeilen, n = _zeilen_von(anzeige, a.ebene, probe.trennen)
            if n and any(z.endswith(w1 + "-") for z in zeilen):
                a.text, a.edits, a.trennen, a.blase_lang = text, probe.edits, probe.trennen, True
                a.stoerer.append("silbentrennung_zeilenende")
                ziel -= 1
                break


def _familie_umbruch_nummer(ctx: Kontext) -> None:
    """Familie "listenpunkt_umbruch_nummer": Ein Listenpunkt, dessen Nummer wechselt (gruen, "c.a."), steht so, dass nur seine erste
    Zeile auf die Seite passt; der Rest und die Textaenderungen (Wort ersetzt, Satz angefuegt) stehen auf der naechsten Seite."""
    r = ctx.rng
    kand = [a for a in ctx.absaetze if a.art == "liste" and a.status is None and a.label_alt != a.label_neu and not a.edits
            and a.excel_alt is None]
    r.shuffle(kand)
    for a in kand[:r.randint(3, 6)]:
        for _ in range(6):
            if len(_zeilen_von(a.text, a.ebene)[0]) >= 3:
                break
            a.text += " " + ctx.satz()
        _aendere_text(ctx, a, "wort_ersetzt")
        a.edits.append(Edit(len(a.text), len(a.text), " " + ctx.satz(), _autor(ctx), "satz_eingefuegt"))
        a.umbruch_erzwingen = True


# ====================================================================== Stoerer ab Seed 401 (nachbau_word)
UMBRUCH_LEERZEICHEN_AB = 401      # Regel 29: Paare der Seeds 301-360 (und alle anderen Korpora) aendern sich dadurch nie
EINHEITEN = ["mm", "kW", "bar", "kg", "cm", "Hz", "ms", "kV"]
STOERER_UMBRUCH_LEERZEICHEN = "umbruch_leerzeichen_ersetzung"


def _familie_umbruch_leerzeichen(ctx: Kontext) -> None:
    """Stoerer "umbruch_leerzeichen_ersetzung" (ab Seed 401, nur Korpus nachbau_word): Ein Wert steht in der Excel und im alten
    Word-Text OHNE Leerzeichen vor der Einheit ("12mm"). Word ersetzt ihn durch "15 " (das Leerzeichen gehoert zur Einfuegung, neu:
    "15 mm"), und der Zeilenumbruch faellt genau hinter die Ersetzung. Am Umbruch zeichnet das PDF kein Leerzeichen, man sieht
    nicht, ob es markiert ist: es sieht aus wie ein gemeinsames Leerzeichen, und der alte Absatz wuerde um ein Leerzeichen zu lang
    ("12 mm"). Die Excel stimmt mit dem alten Word-Text Zeichen fuer Zeichen ueberein (keine Abweichung der Excel): Soll ist
    "uebernommen" (anders als bei excel_abweichung, wo die Excel vom PDF abweicht und nur ein Vorschlag entsteht). Eigener
    Zufallsstrom fuer die Auswahl; der Text kommt mit ctx.satz() aus dem gemeinsamen Strom (die Paare dieser Seeds sind neu)."""
    r = random.Random(ctx.seed * 41 + 11)
    kand = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None and not a.edits and a.excel_alt is None
            and a.label_alt == a.label_neu]
    r.shuffle(kand)
    ziel = r.randint(2, 4)
    for a in kand:
        if ziel <= 0:
            break
        for _versuch in range(300):
            wert, neu_wert = r.sample(["12", "16", "25", "40", "63", "80", "100", "125", "150", "250"], 2)
            einheit = r.choice(EINHEITEN)
            vorne = " ".join([r.choice(["Alle", "Die", "Sämtliche", "Weitere"])] + [r.choice(FUELL) for _ in range(r.randint(3, 11))])
            text = f"{vorne} {wert}{einheit} {ctx.satz()}"
            s0 = len(vorne) + 1
            probe = Absatz(a.id, a.art, a.ebene, text, edits=[Edit(s0, s0 + len(wert), neu_wert + " ", _autor(ctx), "zahl_ersetzt_leerzeichen")])
            anzeige = "".join(t for _art, t, _au, _g in probe.stuecke())
            zeilen, _n = _zeilen_von(anzeige, a.ebene)
            if any(z.endswith(wert + neu_wert) and n.startswith(einheit) for z, n in zip(zeilen, zeilen[1:])):
                a.text, a.edits = text, probe.edits
                a.stoerer.append(STOERER_UMBRUCH_LEERZEICHEN)
                ctx.markiere(STOERER_UMBRUCH_LEERZEICHEN)
                ziel -= 1
                break


# ====================================================================== Excel-Inhalt je Absatz
def _mit_abweichung(a: Absatz) -> Absatz:
    """Kopie des Absatzes mit dem Excel-Text an der Stelle der Abweichung; die Aenderungen dahinter ruecken mit."""
    s, e, neu = a.abweichung                                    # type: ignore[misc]
    d = len(neu) - (e - s)
    edits = [Edit(x.start + (d if x.start >= e else 0), x.ende + (d if x.start >= e else 0), x.neu, x.autor, x.art)
             for x in a.edits]
    assert all(x.ende <= s or x.start >= e for x in a.edits)
    return Absatz(a.id, a.art, a.ebene, a.text[:s] + neu + a.text[e:], edits=edits)



def polster(ctx: Kontext, stufe: int) -> str:
    return ctx.opt["polster"][stufe - 1]


def excel_laeufe(ctx: Kontext, a: Absatz, nachher: bool) -> list[tuple[str, str]]:
    """Inhalt der Textspalte G als Laeufe [(normal|alt|neu, Text)]. vorher: alter Text (ggf. mit Excel-Abweichung)."""
    mit_label = a.art in ("liste", "liste_titel")
    if not nachher:
        if a.status == "neu":
            return []
        text = a.excel_alt if a.excel_alt is not None else a.text
        return [("normal", (a.label_alt + polster(ctx, a.stufe) if mit_label else "") + text)]
    out: list[tuple[str, str]] = []
    if mit_label:
        pol = polster(ctx, a.stufe)
        if a.status == "neu":
            out.append(("neu", a.label_neu + pol))
        elif a.status == "geloescht":
            out.append(("alt", a.label_alt + pol))
        elif a.label_alt != a.label_neu:
            out += [("alt", a.label_alt), ("neu", a.label_neu), ("normal", pol)]
        else:
            out.append(("normal", a.label_alt + pol))
    if not a.edits and a.excel_alt is not None:
        out.append(("normal", a.excel_alt))                       # unveraenderter Absatz: Abweichung bleibt stehen
    else:
        # Weicht die Excel anderswo im Absatz ab, bleibt die Abweichung stehen; nur die PDF-Aenderung wird markiert eingetragen
        # (Nico, 18./19.10.: Vorschlag zum Abhaken). Die Aenderungen liegen mindestens 6 Zeichen von der Abweichung entfernt.
        basis = _mit_abweichung(a) if a.abweichung else a
        for art, text, _autor, _gelb in basis.stuecke():
            out.append(({"k": "normal", "d": "alt", "a": "neu"}[art], text))
    zus: list[tuple[str, str]] = []
    for f, t in out:
        if zus and zus[-1][0] == f:
            zus[-1] = (f, zus[-1][1] + t)
        elif t:
            zus.append((f, t))
    return zus


# ====================================================================== PDF
EINZUG = (55.0, 82.0, 109.0, 137.0, 164.0, 191.0, 218.0)
RECHTS = 395.0                    # rechter Textrand
GROESSE = 8.4                     # Fliesstext (Word skaliert die Seite wegen der Markup-Leiste auf ~70 %)
ZEILE = 10.0                      # Zeilenabstand
GL_VERSATZ = 7.6                  # Grundlinie unter der Oberkante der Zeile
KOERPER_OBEN, KOERPER_UNTEN = 158.0, 620.0
KOPF_GL = (131.0, 142.0)          # Kopfzeile y ~124-145
FUSS_GL = (639.0, 649.0, 659.0)   # Fusszeile y ~630-662
AUTORFARBEN = (pw.WORD_ROT, pw.WORD_BLAU)

Stil = tuple  # (linie "" | "strich" | "unter", farbe, gelb)


def _stil(art: str, autor: int, gelb: bool = False) -> Stil:
    if art == "k":
        return ("", pw.SCHWARZ, gelb)
    if art in ("gd", "ga"):
        return ("strich" if art == "gd" else "unter", pw.WORD_GRUEN, gelb)
    return ("strich" if art == "d" else "unter", AUTORFARBEN[autor], gelb)


class Setzer:
    """Setzt Absaetze zeilenweise auf Seiten (Umbruch wie ein Textprogramm, Seitenumbruch auch mitten im Absatz)."""

    def __init__(self, ctx: Kontext, pdf: pw.WordPdf) -> None:
        self.ctx, self.pdf = ctx, pdf
        self.seite = 0
        self.y = KOERPER_OBEN
        self.balken: dict[int, list[tuple[float, float]]] = {}
        self.woerter_markiert = 0
        self.blasen = 0
        self.kommentare = 0
        self.formatiert = 0
        self.ausserhalb: list[dict[str, Any]] = []      # Aenderungen ausserhalb der CDE (Kopf/Fuss, Vorspann)
        self._kopf_fuss_gemeldet = False

    # ---- Seiten --------------------------------------------------------------------------------------------------
    def neue_seite(self) -> int:
        self.seite = self.pdf.neue_seite()
        self.y = KOERPER_OBEN
        if self.ctx.opt["leiste"]:
            self.pdf.leiste(self.seite)
        self._kopf_fuss()
        return self.seite

    def _kopf_fuss(self) -> None:
        """Kopf- und Fusszeile, auf jeder Seite gleich (bis auf die Seitenzahl), mit eigenen Aenderungen."""
        o = self.ctx.opt
        spez, titel = o["spez"], o["titel"]
        kopf1 = [("k", f"Projekt {o['projekt']} – ", 0, False)]
        if o["kopf_aenderung"]:
            kopf1 += [("d", "[Projektname]", 0, True), ("a", o["ort"], 0, False)]
        else:
            kopf1 += [("k", o["ort"], 0, False)]
        self.zeile(EINZUG[0], KOPF_GL[0], kopf1)
        self.zeile(EINZUG[0], KOPF_GL[1], [("k", f"Spezifikation {spez} – {titel}", 0, False)])
        self.zeile(EINZUG[0], FUSS_GL[0], [("k", titel, 0, False)])
        fuss2 = [("k", "Revision ", 0, False)]
        if o["fuss_aenderung"]:
            fuss2 += [("d", "[Rev]", 0, True), ("a", "A.3.0", 0, False)]
        else:
            fuss2 += [("k", "A.3.0", 0, False)]
        self.zeile(EINZUG[0], FUSS_GL[1], fuss2)
        self.zeile(EINZUG[0], FUSS_GL[2], [("k", f"{spez} - {self.seite}", 0, False)])
        if not self._kopf_fuss_gemeldet:
            self._kopf_fuss_gemeldet = True
            if o["kopf_aenderung"]:
                self.ausserhalb.append({"bereich": "kopf_fuss", "typ": "text_ersetzt", "woerter": 1 + len(o["ort"].split())})
            if o["fuss_aenderung"]:
                self.ausserhalb.append({"bereich": "kopf_fuss", "typ": "text_ersetzt", "woerter": 2})

    def _platz(self) -> None:
        if self.seite == 0 or self.y + ZEILE > KOERPER_UNTEN:
            self.neue_seite()

    # ---- Zeichnen ------------------------------------------------------------------------------------------------
    def _segment(self, x: float, gl: float, text: str, stil: Stil, schrift: str, groesse: float) -> float:
        linie, farbe, gelb = stil
        w = pw.breite(text, schrift, groesse)
        if gelb:
            self.pdf.hinterlegung(self.seite, x, x + w, gl, schrift, groesse)
        self.pdf.text(self.seite, x, gl, text, schrift, groesse, farbe)
        if linie == "strich":
            self.pdf.durchstreichung(self.seite, x, x + w, gl, schrift, groesse, farbe)
        elif linie == "unter":
            self.pdf.unterstreichung(self.seite, x, x + w, gl, schrift, groesse, farbe)
        if linie:
            self.woerter_markiert += len(text.split())
        return x + w

    def zeile(self, x: float, gl: float, stuecke: list[tuple[str, str, int, bool]], schrift: str = "Times-Roman",
              groesse: float = GROESSE) -> bool:
        """Eine Zeile ohne Umbruch (Kopf, Fuss, Vorspann). True = enthaelt eine Aenderung."""
        geaendert = False
        for art, text, autor, gelb in stuecke:
            x = self._segment(x, gl, text, _stil(art, autor, gelb), schrift, groesse)
            geaendert |= art != "k"
        if geaendert:
            self._balken(gl, groesse)
        return geaendert

    def _balken(self, gl: float, groesse: float = GROESSE) -> None:
        self.balken.setdefault(self.seite, []).append((gl - groesse * 0.9, gl + groesse * 0.3))

    def balken_zeichnen(self) -> None:
        for seite, liste in self.balken.items():
            liste.sort()
            zus: list[list[float]] = []
            for a, b in liste:
                if zus and a <= zus[-1][1] + 2.0:
                    zus[-1][1] = max(zus[-1][1], b)
                else:
                    zus.append([a, b])
            for a, b in zus:
                self.pdf.balken(seite, a, b)

    # ---- Absaetze ------------------------------------------------------------------------------------------------
    def absatz(self, x_label: float, x_text: float, label: list[tuple[str, str, int]],
               stuecke: list[tuple[str, str, int, bool]], schrift: str = "Times-Roman", trennen: dict[str, int] | None = None,
               erzwingen: bool = False) -> dict[str, Any]:
        """Setzt einen Absatz mit haengendem Einzug. -> {seite, seite_ende, anker: [(seite, x, gl, geaendert)], getrennt}"""
        zeichen: list[tuple[str, Stil]] = []
        for art, text, autor, gelb in stuecke:
            st = _stil(art, autor, gelb)
            zeichen.extend((c, st) for c in text)
        zeilen, getrennt = _umbrechen(zeichen, RECHTS - x_text, schrift, trennen)
        if erzwingen and len(zeilen) >= 3 and self.seite and self.y + ZEILE <= KOERPER_UNTEN \
                and KOERPER_UNTEN - self.y <= 30 * ZEILE:
            self.y = max(self.y, KOERPER_UNTEN - ZEILE)       # nur die erste Zeile passt auf die Seite, der Rest folgt auf der naechsten
        info: dict[str, Any] = {"seite": 0, "seite_ende": 0, "anker": [], "getrennt": getrennt}
        for nr, z in enumerate(zeilen):
            self._platz()
            gl = self.y + GL_VERSATZ
            if nr == 0:
                info["seite"] = self.seite
                xl = x_label
                for text, art, autor in label:
                    xl = self._segment(xl, gl, text, _stil(art, autor), schrift, GROESSE)
                    if art != "k":
                        self._balken(gl)
            x = x_text
            geaendert = False
            segs: list[tuple[str, Stil]] = []
            for tok in z:
                for c, st in tok:
                    if segs and segs[-1][1] == st:
                        segs[-1] = (segs[-1][0] + c, st)
                    else:
                        segs.append((c, st))
            for text, st in segs:
                x2 = self._segment(x, gl, text, st, schrift, GROESSE)
                if st[0]:
                    geaendert = True
                    info["anker"].append((self.seite, x2, gl, True))
                elif st[2]:
                    info["anker"].append((self.seite, x2, gl, False))
                x = x2
            if geaendert:
                self._balken(gl)
            if nr == 0 and not info["anker"]:
                info["anker"].append((self.seite, x, gl, False))
            info["seite_ende"] = self.seite
            self.y += ZEILE
        return info

    def blase(self, anker: tuple[int, float, float, bool], teile: list[tuple[str, bool]], autor: int = 0) -> bool:
        if not self.ctx.opt["leiste"]:
            return False
        seite, x, gl, _ = anker
        if self.pdf.sprechblase(seite, x, gl, teile, AUTORFARBEN[autor]) is None:
            return False
        self.blasen += 1
        return True


def _umbrechen(zeichen: list[tuple[str, Stil]], maxb: float, schrift: str = "Times-Roman",
               trennen: dict[str, int] | None = None) -> tuple[list[list[list[tuple[str, Stil]]]], int]:
    """Zeilenumbruch eines Absatzes (greedy, wie ein Textprogramm). Passt ein Wort aus `trennen` (Wort -> Trennstelle) nicht mehr auf die
    Zeile, aber sein erster Teil mit Trennstrich, wird es dort getrennt (Silbentrennung am Zeilenende, "Wasser-" / "behaelter").
    -> (Zeilen aus Woertern aus (Zeichen, Stil), Zahl der Trennungen)."""
    token: list[list[tuple[str, Stil]]] = []            # Woerter und Leerzeichen (ein Wort kann mehrere Stile haben)
    for c, st in zeichen:
        if c == " ":
            token.append([(c, st)])
        elif token and token[-1][0][0] != " ":
            token[-1].append((c, st))
        else:
            token.append([(c, st)])

    def w(tok: list[tuple[str, Stil]]) -> float:
        return pw.breite("".join(c for c, _ in tok), schrift, GROESSE)

    zeilen: list[list[list[tuple[str, Stil]]]] = []
    akt: list[list[tuple[str, Stil]]] = []
    breite, luecke, getrennt = 0.0, None, 0
    for tok in token:
        if tok[0][0] == " ":
            if akt:
                luecke = tok if luecke is None else luecke + tok
            continue
        dazu = (w(luecke) if luecke else 0.0) + w(tok)
        if akt and breite + dazu > maxb:
            text = "".join(c for c, _ in tok)
            k = next((k for wort, k in (trennen or {}).items() if text.startswith(wort)), 0)
            if k:
                teil1 = tok[:k] + [("-", tok[k - 1][1])]
                if breite + (w(luecke) if luecke else 0.0) + w(teil1) <= maxb:
                    if luecke:
                        akt.append(luecke)
                    akt.append(teil1)
                    zeilen.append(akt)
                    akt, breite, luecke, getrennt = [tok[k:]], w(tok[k:]), None, getrennt + 1
                    continue
            zeilen.append(akt)                           # Leerzeichen am Zeilenende entfallen (auch gestrichene)
            akt, breite = [tok], w(tok)
        else:
            if luecke:
                akt.append(luecke)
            akt.append(tok)
            breite += dazu
        luecke = None
    if akt or not zeilen:
        zeilen.append(akt)
    return zeilen, getrennt


def _vorspann(s: Setzer) -> None:
    """Deckblatt (alles neu, Helvetica), Pruef- und Revisionstabelle (gestrichen), Bearbeiter-Hinweise, ggf. Inhaltsverzeichnis."""
    ctx, o = s.ctx, s.ctx.opt
    s.neue_seite()
    zeilen = [(f"PROJEKT {o['projekt']}", 14.0), (o["ort"].upper(), 12.0), (f"SPEZIFIKATION {o['spez']}", 12.0),
              (o["titel"], 12.0), ("Revision A.3.0", 10.0), ("Ausgabe 2026_03_01", 10.0)]
    gl = 260.0
    for text, groesse in zeilen:
        s.zeile(120.0, gl, [("a", text, 0, False)], "Helvetica", groesse)
        s.ausserhalb.append({"bereich": "vorspann", "typ": "absatz_neu", "woerter": len(text.split())})
        gl += groesse * 2.2
    # Seite 2: Pruefvermerke (unveraendert) und Revisionstabelle (gestrichen)
    s.neue_seite()
    s.zeile(EINZUG[0], 180.0, [("k", "PRÜFVERMERKE", 0, False)], "Helvetica-Bold", 10.0)
    _tabelle(s, 195.0, [["Schritt", "Stelle", "Datum"], ["Erstellt", "Fachplanung", "2026_02_10"],
                        ["Geprüft", "Qualitätsstelle", "2026_02_20"], ["Freigegeben", "Projektleitung", "2026_03_01"]], "k")
    s.zeile(EINZUG[0], 300.0, [("d", "REVISIONEN", 0, False)], "Helvetica-Bold", 10.0)
    revs = [["Rev.", "Datum", "Beschreibung"], ["A.1.0", "2025_10_01", "Erstausgabe"],
            ["A.2.0", "2025_12_15", "Überarbeitung nach Prüfung"]]
    _tabelle(s, 315.0, revs, "d")
    s.ausserhalb.append({"bereich": "vorspann", "typ": "absatz_geloescht",
                         "woerter": 1 + sum(len(" ".join(z).split()) for z in revs)})
    ctx.markiere("vorspann")
    # Seite 3: Hinweiskaesten fuer Bearbeiter (ganz gestrichen)
    if o["hinweise"]:
        s.neue_seite()
        y = 190.0
        for _ in range(o["hinweise"]):
            texte = ["*" * 40, "HINWEIS FÜR BEARBEITER: Diesen Kasten vor der Ausgabe entfernen.",
                     "Angaben in eckigen Klammern sind projektbezogen zu ergänzen.",
                     "Nicht zutreffende Absätze streichen, Nummerierung prüfen.", "*" * 40]
            s.pdf.rahmen(s.seite, 70.0, y - 12.0, 385.0, y + len(texte) * 11.0, 0.5)
            for i, t in enumerate(texte):
                s.zeile(78.0, y + i * 11.0, [("d", t, 0, False)])
            s.ausserhalb.append({"bereich": "vorspann", "typ": "absatz_geloescht", "woerter": sum(len(t.split()) for t in texte)})
            y += len(texte) * 11.0 + 40.0
        ctx.markiere("bearbeiter_hinweise")
    # Seite 4: Inhaltsverzeichnis (unveraendert, wiederholt die Ueberschriften)
    if o["inhalt"]:
        s.neue_seite()
        s.zeile(EINZUG[0], 180.0, [("k", "INHALT", 0, False)], "Times-Bold")
        gl = 200.0
        for a in ctx.absaetze:
            if a.art == "part":
                t = f"PART {a.nummer} {a.text}"
            elif a.art == "section":
                t = f"{a.nummer} {a.text}"
            else:
                continue
            s.zeile(EINZUG[0] if a.art == "part" else EINZUG[1], gl, [("k", t + " " + "." * 20, 0, False)])
            gl += ZEILE
            if gl > KOERPER_UNTEN - 10:
                break
        ctx.markiere("inhaltsverzeichnis")


def _tabelle(s: Setzer, oben: float, zeilen: list[list[str]], art: str) -> None:
    xs = (EINZUG[0], 150.0, 260.0, 380.0)
    h = 13.0
    for i, z in enumerate(zeilen):
        y = oben + i * h
        for j, t in enumerate(z):
            s.zeile(xs[j] + 3.0, y + 9.5, [(art, t, 0, False)])
    for i in range(len(zeilen) + 1):
        s.pdf.rechteck(s.seite, xs[0], oben + i * h - 0.25, xs[-1], oben + i * h + 0.25, pw.SCHWARZ)
    for x in xs:
        s.pdf.rechteck(s.seite, x - 0.25, oben, x + 0.25, oben + len(zeilen) * h, pw.SCHWARZ)


def rendere_pdf(ctx: Kontext) -> tuple[pw.WordPdf, Setzer]:
    pdf = pw.WordPdf()
    s = Setzer(ctx, pdf)
    r = random.Random(ctx.seed * 31 + 7)               # eigener Strom fuer Sprechblasen (das Modell bleibt unberuehrt)
    o = ctx.opt
    _vorspann(s)
    s.neue_seite()
    s.absatz(EINZUG[0], EINZUG[0], [], [("k", f"SECTION {o['spez']} – {o['titel']}", 0, False)], "Times-Bold")
    s.y += ZEILE
    kommentar_nr = 0
    for a in ctx.absaetze:
        if a.art == "part":
            s.y += 6.0
            info = s.absatz(EINZUG[0], EINZUG[0], [], [("k", f"PART {a.nummer} {a.text}", 0, False)], "Times-Bold")
        elif a.art == "section":
            s.y += 4.0
            info = s.absatz(EINZUG[0], EINZUG[1], a.label_stuecke(), a.stuecke(), "Times-Bold")
        else:
            x_label = EINZUG[a.ebene]
            x_text = EINZUG[min(a.ebene + 1, len(EINZUG) - 1)]
            info = s.absatz(x_label, x_text, a.label_stuecke(), a.stuecke(), trennen=a.trennen or None,
                            erzwingen=a.umbruch_erzwingen)
            if info["getrennt"]:
                ctx.markiere("silbentrennung_zeilenende")
                a.stoerer.append("silbentrennung_zeilenende")
        s.y += 3.0
        a.seite, a.seite_ende = info["seite"], info["seite_ende"]
        if a.seite_ende > a.seite and (a.edits or a.label_alt != a.label_neu):
            ctx.markiere("seitenumbruch_im_absatz")
            a.stoerer.append("seitenumbruch_im_absatz")
            if a.umbruch_erzwingen and a.label_alt != a.label_neu:
                ctx.markiere("listenpunkt_umbruch_nummer")        # Nummer wechselt (gruen) in der ersten Zeile, Aenderung danach
                a.stoerer.append("listenpunkt_umbruch_nummer")
        if a.blase_lang and o["leiste"] and info["anker"]:     # eigener Zufallsstrom: die uebrigen Sprechblasen bleiben gleich
            rb = random.Random(ctx.seed * 37 + a.id)
            kommentar_nr += 1
            text = " ".join(rb.sample(KOMMENTARE, 3))
            if s.blase(info["anker"][-1], [(f"Kommentiert [{o['initialen']}{kommentar_nr}]:", True), (" " + text, False)]):
                s.kommentare += 1
                ctx.markiere("sprechblase_mehrzeilig")
        # Sprechblasen in der Markup-Leiste (nichts davon gehoert in die Excel)
        if not o["leiste"] or not info["anker"]:
            continue
        geaendert = [x for x in info["anker"] if x[3]]
        gelb = [x for x in info["anker"] if not x[3]]
        if o["formatiert"] and (geaendert or a.gelb) and r.random() < 0.22:
            kopf = r.choice(["hat formatiert:", "Formatiert:"])
            if s.blase(r.choice(geaendert or gelb), [(kopf, True), (" " + r.choice(FORMATIERUNGEN), False)], a.autor):
                s.formatiert += 1
        if o["kommentare"] and r.random() < 0.035:
            kommentar_nr += 1
            text = r.choice(KOMMENTARE) + (" " + r.choice(KOMMENTARE) if r.random() < 0.4 else "")
            if s.blase(info["anker"][0], [(f"Kommentiert [{o['initialen']}{kommentar_nr}]:", True), (" " + text, False)]):
                s.kommentare += 1
    s.y += ZEILE
    s.absatz(EINZUG[0], EINZUG[0], [], [("k", "END OF SECTION", 0, False)], "Times-Bold")
    s.balken_zeichnen()
    if s.formatiert:
        ctx.markiere("sprechblase_formatiert")
    if s.kommentare:
        ctx.markiere("sprechblase_kommentar")
    return pdf, s


# ====================================================================== Excel (CDE-Register)
BREITEN = {"A": 7, "B": 12, "C": 10, "D": 12, "E": 9, "F": 11, "G": 96, "H": 8, "I": 115, "J": 4, "K": 4, "L": 4, "M": 4,
           "N": 4, "O": 6, "P": 90, "Q": 3, "R": 11, "S": 18, "AB": 40}
BEGRUENDUNGEN = ["Mit der Fachplanung abgestimmt.", "Gleichwertige Ausführung angeboten.", "Siehe Datenblatt im Anhang.",
                 "Nachweis folgt mit der Werkplanung."]
BEMERKUNGEN = ["siehe Anlage", "Rückfrage offen", "bestätigt", "Klärung im Termin"]
DUENN = Side(style="thin", color="FFBFBFBF")
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)


@dataclass
class Block:
    """Eine Spezifikation im Register (zusammenhaengender Zeilenblock). Nur der Ziel-Block gehoert zum PDF."""
    rolle: str                    # ziel | fremd
    absaetze: list[Absatz]
    spez_excel: str               # Schreibweise in Spalte B ("11_22_33", "11 22 33")
    variante: str                 # ohne | leerzeichen (alle mit Leerzeichen am Ende) | gemischt (einzelne)
    rev: str
    datum: Any


def _excel_werte(ctx: Kontext) -> None:
    """Werte der uebrigen Spalten je Absatz (fest je Seed), je Spec-Block mit eigener Spec-Nr., Revision und Datum. Neue Absaetze
    bekommen nur B-F: B/C/D wie die Zeile davor (ihr Block), E/F aus der Gliederung; alles rechts der Textspalte bleibt leer."""
    o = ctx.opt
    for k, blk in enumerate(ctx.bloecke):
        r = random.Random(ctx.seed * 17 + 3 + 1009 * k * (blk.rolle == "fremd"))
        vorher = None
        for i, a in enumerate(blk.absaetze):
            w: dict[str, Any] = {}
            a.excel = w                                        # type: ignore[attr-defined]
            b = blk.spez_excel + (" " if blk.variante == "leerzeichen"
                                  or (blk.variante == "gemischt" and r.random() < 0.06) else "")
            e = a.section + (" " if o["e_variante"] == "alle" or (o["e_variante"] == "part" and a.art == "part") else "")
            if a.status == "neu" and vorher is not None:
                vw = vorher.excel                              # type: ignore[attr-defined]
                w.update(B=vw["B"], C=vw["C"], D=vw["D"], E=e)
            else:
                w.update(B=b, C=blk.rev, D=blk.datum, E=e)
            if a.sub:
                w["F"] = a.sub
            if a.status != "neu":
                if not a.ueberschrift:
                    w["H"] = r.choices(["D", "E", "C"], [5, 3, 2])[0]
                    w["R"] = r.choice(["Offen", "Offen", "Erledigt"])
                    if r.random() < 0.6:
                        w["JKLMN"[r.randint(0, 4)]] = "X"
                    if r.random() < 0.25:
                        w["O"] = r.choice(["A", "R"])
                    if r.random() < 0.12:
                        w["P"] = r.choice(BEGRUENDUNGEN)
                    if r.random() < 0.08:
                        w["S"] = r.choice(BEMERKUNGEN)
                    if r.random() < 0.05:
                        w["AB"] = r.choice(KOMMENTARE)
                    if r.random() < 0.03:
                        w["A"] = r.randint(1, 40)
                    if o["fette_zellen"] and r.random() < 0.04:
                        a.fett = True
                        ctx.markiere("fette_zelle")
                else:
                    w["I"] = KENNWORT
                naechster = blk.absaetze[i + 1] if i + 1 < len(blk.absaetze) else None
                if blk.rolle == "ziel" and o["alte_markierung"] and r.random() < 0.06 \
                        and not (naechster and naechster.status == "neu"):
                    w["alt_markiert"] = True                   # C/D tragen Markierungen einer frueheren Runde
                    ctx.markiere("alte_markierung_cd")
            vorher = a


SPEZ_REVS = ["A.1.0", "A.2.0", "B.1.0", "C.0.0", "B.2.1"]
SPEZ_DATEN = ["2025_09_12", "2025_11_20", "2026_01_15", "2026_02_03"]


def _fremde_specs(ctx: Kontext, ziel: Block) -> list[Block]:
    """Stoerer `mehrere_specs`: 2-3 weitere Specs als eigene Bloecke vor UND nach der Ziel-Spec, mit eigener Spec-Nr. (uneinheitliche
    Schreibweise), Revision, Datum und Gliederung 1.1 ... Ein Teil ihrer Saetze steht genau so oder fast so auch in der Ziel-Spec
    (Dubletten-Falle) -- bevorzugt Saetze, die das PDF aendert. Laut Soll bleiben diese Bloecke unveraendert."""
    r = random.Random(ctx.seed * 97 + 13)
    n = r.randint(2, 3)
    davor = r.randint(1, n - 1)
    ziel_nr = ctx.opt["spez"]
    kandidaten = [a for a in ctx.absaetze if a.art in ("sub", "liste") and a.status is None]
    kandidaten.sort(key=lambda a: (not a.edits, a.id))
    fremde: list[Block] = []
    for k in range(n):
        nr = ziel_nr
        while nr == ziel_nr:
            nr = f"{r.randint(10, 48)} {r.randint(10, 99)} {r.randint(10, 99)}"
        schreib = r.choice(["unterstrich", "unterstrich", "leer"])
        tmp = Kontext(ctx.seed, random.Random(ctx.seed * 131 + k * 7919), {}, benutzt=ctx.benutzt)
        tmp._zaehler = 100_000 * (k + 1)
        tmp.opt["doppelbuchstaben"] = False
        groesse = r.randint(*ctx.cfg.get("fremd_zeilen", (40, 140)))
        _baue_struktur(tmp, groesse)
        _kuerze(tmp, groesse)
        # Dubletten: gleiche oder fast gleiche Saetze wie in der Ziel-Spec (dieselbe Art: Absatz bzw. Listenpunkt)
        plaetze = [a for a in tmp.absaetze if a.art in ("sub", "liste")]
        r.shuffle(plaetze)
        quelle = [a for a in kandidaten if a.edits][:6] + r.sample(kandidaten, min(6, len(kandidaten)))
        gesetzt = 0
        for a in plaetze:
            passend = [q for q in quelle if q.art == a.art]
            if not passend or gesetzt >= max(2, len(plaetze) // 4):
                continue
            q = passend[gesetzt % len(passend)]
            a.text = q.text if r.random() < 0.6 else _fast_gleich(r, q.text)
            gesetzt += 1
        if gesetzt:
            ctx.markiere("dubletten_fremde_spec")
        fremde.append(Block("fremd", tmp.absaetze, nr.replace(" ", "_") if schreib == "unterstrich" else nr,
                            r.choice(["ohne", "ohne", "leerzeichen", "gemischt"]), r.choice(SPEZ_REVS), r.choice(SPEZ_DATEN)))
    ctx.markiere("mehrere_specs")
    return fremde[:davor] + [ziel] + fremde[davor:]


def _fast_gleich(r: random.Random, text: str) -> str:
    """Fast gleicher Satz: ein Wort (nicht das erste) ausgetauscht."""
    stellen = [m.span() for m in _WORT.finditer(text) if m.start() > 0]
    if not stellen:
        return text + " (Entwurf)"
    s, e = r.choice(stellen)
    wort = text[s:e]
    pool = _KATEGORIEN.get(wort) or (ROLLEN if wort[:1].isupper() else ADVERB)
    neu = r.choice([w for w in pool if w != wort and " " not in w] or ["zusätzlich"])
    return text[:s] + neu + text[e:]


def _schrift(fett: bool = False, groesse: float = 10, **kw: Any) -> Font:
    return Font(name="Arial", size=groesse, bold=fett or None, **kw)


def _laeufe_wert(laeufe: list[tuple[str, str]], fett: bool, einfarbig: str) -> tuple[Any, Font | None]:
    """Zellwert und Zellschrift fuer Laeufe [(normal|alt|neu, Text)] -- Grundschrift (Arial 11, ggf. fett) in jedem Lauf."""
    if not laeufe:
        return None, None
    formate = {f for f, _ in laeufe}
    text = "".join(t for _, t in laeufe)
    if formate == {"normal"}:
        return text, None
    if len(formate) == 1 and einfarbig == "zellschrift":
        f = next(iter(formate))
        return text, _schrift(fett, 11, color=ALT_FARBE if f == "alt" else NEU_FARBE, strike=True if f == "alt" else None)
    bloecke = []
    for f, t in laeufe:
        farbe = {"alt": ALT_FARBE, "neu": NEU_FARBE}.get(f)
        bloecke.append(TextBlock(InlineFont(rFont="Arial", sz=11, b=True if fett else None, color=farbe,
                                            strike=True if f == "alt" else None), t))
    return CellRichText(*bloecke), None


def _frueher(alt: str, neu: str) -> CellRichText:
    """Markierung einer frueheren Runde (dunkelblau alt, Zeilenumbruch, dunkelrot neu) -- bleibt unangetastet."""
    return CellRichText(TextBlock(InlineFont(rFont="Arial", sz=10, color=FRUEHER_ALT), alt), "\n",
                        TextBlock(InlineFont(rFont="Arial", sz=10, color=FRUEHER_NEU), neu))


def _kopfbereich(ws: Any, ctx: Kontext) -> list[str]:
    """Zeilen 1-7: Verbuende, Titel, Zaehlformeln (U2:Z5), Spaltenkoepfe (5/6) und Beispielzeile (7). -> Formelzellen."""
    o = ctx.opt
    for rng in ("A1:A4", "B1:B4", "C1:C4", "D1:F4", "G1:G4", "H1:P4", "G5:G6", "H5:H6", "I5:I6", "J5:N5", "O5:O6", "P5:P6",
                "E5:F5"):
        ws.merge_cells(rng)
    kopf = PatternFill("solid", fgColor="FFDDEBF7")
    ws["B1"] = "CDE"
    ws["C1"] = "Register"
    ws["D1"] = (f"Anforderungsabgleich Spezifikation {o['spez']}" if len(ctx.bloecke) <= 1
                else f"Anforderungsabgleich Projekt {o['projekt']} (mehrere Spezifikationen)")
    ws["G1"] = f"Projekt {o['projekt']} – {o['titel']}"
    ws["H1"] = "Hinweise: Spalten J bis N ankreuzen, Status in Spalte R pflegen, Begründung in Spalte P."
    for z in ("B1", "C1", "D1", "G1", "H1"):
        ws[z].font = _schrift(True, 12 if z == "G1" else 10)
        ws[z].alignment = Alignment(wrap_text=True, vertical="center")
    formeln = []
    for i, k in enumerate(("D", "E", "C")):
        r_ = 2 + i
        ws[f"T{r_}"] = k
        ws[f"U{r_}"] = f'=COUNTIF(H:H,"{k}")'
        ws[f"V{r_}"] = f'=COUNTIFS(H:H,"{k}",R:R,"Offen")'
        ws[f"W{r_}"] = f"=IF(U{r_}=0,0,V{r_}/U{r_})"
        formeln += [f"U{r_}", f"V{r_}", f"W{r_}"]
    ws["T5"] = "Summe"
    for sp in "UVW":
        ws[f"{sp}5"] = f"=SUM({sp}2:{sp}4)"
        formeln.append(f"{sp}5")
    ws["X2"] = '=COUNTIF(R:R,"Erledigt")'
    ws["Y2"] = "=COUNTA(G:G)-2"
    ws["Z2"] = "=IF(Y2=0,0,X2/Y2)"
    formeln += ["X2", "Y2", "Z2"]
    koepfe = {"A5": "Nr.", "B5": "Spezifikation", "C5": "Revision", "D5": "Datum", "E5": "Gliederung", "G5": "Anforderungstext",
              "H5": "Status", "I5": "Beschreibung", "J5": "Bewertung", "O5": "A/R", "P5": "Begründung", "R5": "Bearbeitung",
              "S5": "Bemerkung", "AB5": "Kommentar",
              "A6": "Pos.-Nr.", "B6": "Spez.-Nr.", "C6": "Rev.-Nr.", "D6": "Spez.-Datum", "E6": "Abschnitt:",
              "F6": "Unterabschnitt:", "J6": "K1", "K6": "K2", "L6": "K3", "M6": "K4", "N6": "K5", "R6": "Stand", "S6": "-"}
    for adr, t in koepfe.items():
        ws[adr] = t
        ws[adr].font = _schrift(True, 10)
        ws[adr].fill = kopf
        ws[adr].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws[adr].border = RAHMEN
    beispiel = {"A": "0", "B": "00-00-00", "C": "Rev 0.", "D": "00-00-00", "E": "Section", "F": "Sub-Section",
                "G": "Beispiel: Anforderung wie in der Spezifikation (Text mit Listennummer)", "H": "D",
                "I": "Beispiel für eine Beschreibung", "J": "X", "O": "A", "P": "Beispiel für eine Begründung", "R": "Offen"}
    for sp, t in beispiel.items():
        z = ws[f"{sp}7"]
        z.value = t
        z.font = _schrift(False, 11 if sp == "G" else 10, italic=True, color="FF808080")
        z.border = RAHMEN
    return formeln


def schreibe_excel(pfad: Path, ctx: Kontext, nachher: bool) -> tuple[dict[int, int], list[str]]:
    """Schreibt vorher.xlsx (nachher=False) oder nachher.xlsx. -> (Absatz-id -> Zeile, Formelzellen des Hauptblatts)."""
    o = ctx.opt
    wb = Workbook()
    wb.properties.creator = "erfunden"
    wb.properties.created = wb.properties.modified = datetime(2026, 3, 1)
    hinweise = wb.active
    hinweise.title = "Hinweise"
    hinweise.column_dimensions["A"].width = 150
    for i, t in enumerate(["Hinweise zur Bearbeitung des Registers (erfunden)",
                           "Je Absatz der Spezifikation gibt es eine Zeile. Die Listennummer steht im Text der Spalte G.",
                           "Status D, E oder C in Spalte H wählen (Auswahlliste).",
                           "Änderungen einer neuen Revision werden farbig markiert: neu rot, entfallen blau durchgestrichen.",
                           "Überschriftzeilen tragen in Spalte I ein festes Kennwort."], 1):
        hinweise.cell(i, 1, t).alignment = Alignment(wrap_text=True)
    deck = wb.create_sheet("Deckblatt")
    for i, (k, v) in enumerate([("Projekt", o["projekt"]), ("Standort", o["ort"]), ("Spezifikation", o["spez"]),
                                ("Titel", o["titel"]), ("Revision", "A.3.0")], 2):
        deck.cell(i, 2, k).font = _schrift(True)
        deck.cell(i, 3, v)
    unt = wb.create_sheet("Unterschriften")
    for i, z in enumerate([["Rolle", "Stelle", "Datum", "Unterschrift"], ["Erstellt", "Fachplanung", "", ""],
                           ["Geprüft", "Qualitätsstelle", "", ""], ["Freigegeben", "Projektleitung", "", ""]], 1):
        for j, v in enumerate(z, 1):
            if v:
                unt.cell(i, j, v)
    ws = wb.create_sheet(o["blattname"])
    listen = wb.create_sheet("Listen")
    for i, v in enumerate(["D", "E", "C"], 1):
        listen.cell(i, 1, v)
    for i, v in enumerate(["Offen", "Erledigt"], 1):
        listen.cell(i, 2, v)
    listen.sheet_state = "hidden"

    formeln = _kopfbereich(ws, ctx)
    pos: dict[int, int] = {}
    zeile = DATEN_START - 1
    for a in (a for blk in ctx.bloecke for a in blk.absaetze):
        if not nachher and a.status == "neu":
            continue
        zeile += 1
        pos[a.id] = zeile
        w = a.excel                                            # type: ignore[attr-defined]
        fett = a.ueberschrift or a.fett
        for sp in [get_column_letter(c) for c in range(1, LETZTE_SPALTE + 1)]:
            z = ws[f"{sp}{zeile}"]
            z.font = _schrift(fett if sp == "G" else False, 11 if sp == "G" else 10)
            z.alignment = Alignment(wrap_text=sp in ("G", "I", "P", "AB"), vertical="top")
            z.border = RAHMEN
            if sp == "G":
                wert, schrift = _laeufe_wert(excel_laeufe(ctx, a, nachher), fett, o["einfarbig_als"])
                z.value = wert
                if schrift is not None:
                    z.font = schrift
            elif sp in w:
                z.value = w[sp]
        if w.get("alt_markiert"):
            ws[f"C{zeile}"] = _frueher("A.2.0", "A.3.0")
            if isinstance(o["datum"], str):
                ws[f"D{zeile}"] = _frueher("2025_12_15", o["datum"])
        if not isinstance(o["datum"], str):
            ws[f"D{zeile}"].number_format = "DD.MM.YYYY"
    letzte = zeile
    _bereiche(ws, ctx, letzte, pos, nachher)
    for sp, b in BREITEN.items():
        ws.column_dimensions[sp].width = b
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A6:AB{letzte}"
    ws.print_area = f"A5:P{letzte}"
    ws.print_title_rows = "5:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    wb.active = 3
    wb.save(str(pfad))
    return pos, formeln


def _bereiche(ws: Any, ctx: Kontext, letzte: int, pos: dict[int, int], nachher: bool) -> None:
    """Bedingte Formatierung und Datenpruefung in zerstueckelten Bereichen. Im nachher ziehen sie mit eingefuegten Zeilen mit
    (wie Excel beim Einfuegen von Zeilen)."""
    r = random.Random(ctx.seed * 13 + 5)
    vorher_zeilen = sorted(p for a, p in ((a, pos.get(a.id)) for blk in ctx.bloecke for a in blk.absaetze
                                          if a.status != "neu") if p)
    n = len(vorher_zeilen)
    vletzte = DATEN_START - 1 + n                               # letzte Zeile in vorher

    def z(v: int) -> int:                                      # vorher-Zeile -> Zeile in dieser Datei
        if v < DATEN_START:
            return v
        if v > vletzte:
            return v - vletzte + letzte
        return vorher_zeilen[v - DATEN_START]

    stuecke = sorted(r.sample(range(DATEN_START + 2, max(DATEN_START + 12, vletzte - 2)), 6))
    f_teile = f"F{z(stuecke[0])}:F{z(stuecke[1])} F{z(stuecke[2])}:F{z(stuecke[3])}"
    ws.conditional_formatting.add(f_teile, FormulaRule(formula=[f"LEN(F{z(stuecke[0])})>2"],
                                                       fill=PatternFill("solid", bgColor="FFFCE4D6")))
    ws.conditional_formatting.add(f"O8:O{letzte}", CellIsRule(operator="equal", formula=['"R"'],
                                                             fill=PatternFill("solid", bgColor="FFFFC7CE")))
    ws.conditional_formatting.add("R5 R8:R1048576", FormulaRule(formula=['R5="Offen"'], fill=PatternFill("solid", bgColor="FFFFEB9C")))
    ws.conditional_formatting.add("V2:V4", CellIsRule(operator="greaterThan", formula=["0"],
                                                     fill=PatternFill("solid", bgColor="FFF8CBAD")))
    dv_h = DataValidation(type="list", formula1="Listen!$A$1:$A$3", allow_blank=True)
    dv_h.add(f"H{z(stuecke[1])}:H{z(stuecke[4])}")
    dv_h.add(f"H{z(stuecke[5])}:H{letzte}")
    dv_x = DataValidation(type="list", formula1='"X"', allow_blank=True)
    dv_x.add(f"J7:N{letzte}")
    dv_x.add(f"K{letzte + 6}")
    dv_r = DataValidation(type="list", formula1="Listen!$B$1:$B$2", allow_blank=True)
    dv_r.add(f"R8:R{letzte}")
    for dv in (dv_h, dv_x, dv_r):
        ws.add_data_validation(dv)


# ====================================================================== ZIP-Nachbearbeitung (Zusatzteile wie in Nicos Dateien)
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
FESTE_ZEIT = (2026, 3, 1, 0, 0, 0)
XML_KOPF = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
ZUSATZTEILE = ("xl/sharedStrings.xml", "xl/calcChain.xml", "xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels",
               "xl/media/image1.png", "xl/media/image2.png", "xl/printerSettings/printerSettings1.bin", "customXml/item1.xml",
               "customXml/itemProps1.xml", "customXml/_rels/item1.xml.rels", "customXml/item2.xml", "customXml/itemProps2.xml",
               "customXml/_rels/item2.xml.rels", "docMetadata/LabelInfo.xml")


def _guid(r: random.Random) -> str:
    return "{%08X-%04X-%04X-%04X-%012X}" % (r.getrandbits(32), r.getrandbits(16), r.getrandbits(16), r.getrandbits(16),
                                            r.getrandbits(48))


def _png(r: random.Random, b: int = 24, h: int = 12) -> bytes:
    """Kleines erfundenes Bild (PNG ohne PIL)."""
    import struct
    roh = b"".join(b"\x00" + bytes(v for _ in range(b) for v in (r.randint(0, 255), 120, 200)) for _ in range(h))

    def chunk(typ: bytes, daten: bytes) -> bytes:
        return struct.pack(">I", len(daten)) + typ + daten + struct.pack(">I", zlib.crc32(typ + daten) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", b, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(roh, 9)) + chunk(b"IEND", b""))


def _rel(rid: str, typ: str, ziel: str) -> bytes:
    return f'<Relationship Id="{rid}" Type="{typ}" Target="{ziel}"/>'.encode()


def _rels_dazu(alt: bytes | None, neue: list[bytes]) -> bytes:
    if alt is None:
        return XML_KOPF + f'<Relationships xmlns="{NS_PKG}">'.encode() + b"".join(neue) + b"</Relationships>"
    return alt.replace(b"</Relationships>", b"".join(neue) + b"</Relationships>")


def _zeichnung(r: random.Random, bilder: int) -> bytes:
    """Eine Form und 1-2 Bilder im Kopfbereich (Anker Zeile 3 und 7)."""
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    teile = [f'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" {a}>']

    def anker(c0: int, r0: int, c1: int, r1: int, inhalt: str) -> str:
        return (f'<xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>{c0}</xdr:col><xdr:colOff>19050</xdr:colOff>'
                f'<xdr:row>{r0}</xdr:row><xdr:rowOff>19050</xdr:rowOff></xdr:from><xdr:to><xdr:col>{c1}</xdr:col>'
                f'<xdr:colOff>0</xdr:colOff><xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>{inhalt}'
                f'<xdr:clientData/></xdr:twoCellAnchor>')
    for i in range(bilder):
        teile.append(anker(i * 2, 2, i * 2 + 1, 6,
                           f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{i + 2}" name="Grafik {i + 1}"/><xdr:cNvPicPr>'
                           f'<a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr><xdr:blipFill>'
                           f'<a:blip xmlns:r="{NS_R}" r:embed="rId{i + 1}"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
                           f'<xdr:spPr><a:xfrm><a:off x="{19050 + i * 900000}" y="400000"/><a:ext cx="600000" cy="300000"/></a:xfrm>'
                           f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic>'))
    teile.append(anker(6, 2, 7, 6,
                       '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="9" name="Rechteck 8"/><xdr:cNvSpPr/></xdr:nvSpPr>'
                       '<xdr:spPr><a:xfrm><a:off x="5000000" y="400000"/><a:ext cx="1800000" cy="300000"/></a:xfrm>'
                       '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="DDEBF7"/></a:solidFill>'
                       '</xdr:spPr><xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="de-DE" sz="900"/>'
                       '<a:t>Muster-Stempel</a:t></a:r></a:p></xdr:txBody></xdr:sp>'))
    teile.append("</xdr:wsDr>")
    return XML_KOPF + "".join(teile).encode()


def _blatt_pfad(teile: dict[str, bytes], name: str) -> str:
    wb = teile["xl/workbook.xml"].decode()
    rid = re.search(r'<sheet name="%s" sheetId="\d+" (?:state="\w+" )?r:id="(rId\d+)"' % re.escape(name), wb).group(1)
    rels = teile["xl/_rels/workbook.xml.rels"].decode()
    ziel = re.search(r'<Relationship[^>]*Id="%s"[^>]*/>' % rid, rels).group(0)
    ziel = re.search(r'Target="([^"]+)"', ziel).group(1)
    return ziel.lstrip("/") if ziel.startswith("/") else "xl/" + ziel


def nachbearbeiten(pfad: Path, ctx: Kontext, formeln: list[str]) -> None:
    """Nach dem openpyxl-Speichern: sharedStrings statt Inline-Strings (wie Excel), Blattschutz-Element, Zeichnung mit Form und
    Bildern, Druckereinstellungen, calcChain, customXml und Vertraulichkeits-Kennzeichnung -- mit korrekten Content-Types und
    Rels. Feste Zeitstempel im ZIP: gleiche Eingabe => bytegleiche Datei."""
    r = random.Random(ctx.seed * 7 + 11)
    with zipfile.ZipFile(pfad) as z:
        teile = {n: z.read(n) for n in z.namelist()}
    # 1. Inline-Strings -> sharedStrings
    sst: list[bytes] = []
    index: dict[bytes, int] = {}
    anzahl = 0
    muster = re.compile(rb'<c ([^>]*?)t="inlineStr"([^>]*)><is>(.*?)</is></c>', re.S)

    def ersetze(m: re.Match) -> bytes:
        nonlocal anzahl
        anzahl += 1
        i = index.setdefault(m.group(3), len(sst))
        if i == len(sst):
            sst.append(m.group(3))
        return b"<c " + m.group(1) + b't="s"' + m.group(2) + b"><v>%d</v></c>" % i
    blaetter = sorted((n for n in teile if re.match(r"xl/worksheets/sheet\d+\.xml$", n)), key=lambda n: int(re.findall(r"\d+", n)[0]))
    for n in blaetter:
        teile[n] = muster.sub(ersetze, teile[n])
    teile["xl/sharedStrings.xml"] = (XML_KOPF + f'<sst xmlns="{NS_MAIN}" count="{anzahl}" uniqueCount="{len(sst)}">'.encode()
                                     + b"".join(b"<si>" + s + b"</si>" for s in sst) + b"</sst>")
    # 2. Hauptblatt: Namensraum r, Blattschutz-Element, Druckereinstellungen, Zeichnung
    haupt = _blatt_pfad(teile, ctx.opt["blattname"])
    x = teile[haupt]
    if b"xmlns:r=" not in x[:300]:
        x = x.replace(f'<worksheet xmlns="{NS_MAIN}"'.encode(), f'<worksheet xmlns="{NS_MAIN}" xmlns:r="{NS_R}"'.encode(), 1)
    x = re.sub(rb"<sheetProtection[^>]*/>", b"", x)
    x = x.replace(b"</sheetData>", b'</sheetData><sheetProtection formatCells="0"/>', 1)
    x = x.replace(b"<pageSetup ", b'<pageSetup r:id="rIdP1" ', 1)
    m = re.search(rb"<legacyDrawing|<tableParts|<extLst|</worksheet>", x)
    x = x[:m.start()] + b'<drawing r:id="rIdD1"/>' + x[m.start():]
    teile[haupt] = x
    ordner, datei = haupt.rsplit("/", 1)
    rels_pfad = f"{ordner}/_rels/{datei}.rels"
    teile[rels_pfad] = _rels_dazu(teile.get(rels_pfad), [
        _rel("rIdP1", f"{NS_R}/printerSettings", "../printerSettings/printerSettings1.bin"),
        _rel("rIdD1", f"{NS_R}/drawing", "../drawings/drawing1.xml")])
    bilder = ctx.opt["bilder"]
    teile["xl/drawings/drawing1.xml"] = _zeichnung(r, bilder)
    teile["xl/drawings/_rels/drawing1.xml.rels"] = _rels_dazu(None, [_rel(f"rId{i + 1}", f"{NS_R}/image", f"../media/image{i + 1}.png")
                                                                    for i in range(bilder)])
    for i in range(bilder):
        teile[f"xl/media/image{i + 1}.png"] = _png(r)
    teile["xl/printerSettings/printerSettings1.bin"] = bytes(r.getrandbits(8) for _ in range(220))
    # 3. calcChain (Formelzellen des Hauptblatts)
    blatt_id = re.search(r'<sheet name="%s" sheetId="(\d+)"' % re.escape(ctx.opt["blattname"]),
                         teile["xl/workbook.xml"].decode()).group(1)
    teile["xl/calcChain.xml"] = (XML_KOPF + f'<calcChain xmlns="{NS_MAIN}">'.encode()
                                 + "".join(f'<c r="{z}" i="{blatt_id}"/>' for z in formeln).encode() + b"</calcChain>")
    # 4. customXml und Vertraulichkeits-Kennzeichnung (erfundene Kennungen)
    for i in (1, 2):
        teile[f"customXml/item{i}.xml"] = (XML_KOPF + b'<b:Sources SelectedStyle="" xmlns:b="http://schemas.openxmlformats.org/'
                                           b'officeDocument/2006/bibliography" xmlns="http://schemas.openxmlformats.org/'
                                           b'officeDocument/2006/bibliography"/>' if i == 1 else
                                           XML_KOPF + f'<erfunden xmlns="urn:erfunden:ablage"><Ablage>{ctx.opt["projekt"]}</Ablage>'
                                           f'</erfunden>'.encode())
        teile[f"customXml/itemProps{i}.xml"] = (XML_KOPF + f'<ds:datastoreItem ds:itemID="{_guid(r)}" xmlns:ds="http://schemas.'
                                                f'openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/>'
                                                f'</ds:datastoreItem>'.encode())
        teile[f"customXml/_rels/item{i}.xml.rels"] = _rels_dazu(None, [_rel("rId1", f"{NS_R}/customXmlProps", f"itemProps{i}.xml")])
    teile["docMetadata/LabelInfo.xml"] = (XML_KOPF + f'<clbl:labelList xmlns:clbl="http://schemas.microsoft.com/office/2020/'
                                          f'mipLabelMetadata"><clbl:label id="{_guid(r)}" enabled="1" method="Standard" '
                                          f'siteId="{_guid(r)}" contentBits="0" removed="0"/></clbl:labelList>'.encode())
    # 5. Rels und Content-Types
    teile["_rels/.rels"] = _rels_dazu(teile["_rels/.rels"], [
        _rel("rIdL1", "http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels", "docMetadata/LabelInfo.xml")])
    teile["xl/_rels/workbook.xml.rels"] = _rels_dazu(teile["xl/_rels/workbook.xml.rels"], [
        _rel("rIdS1", f"{NS_R}/sharedStrings", "sharedStrings.xml"), _rel("rIdC1", f"{NS_R}/calcChain", "calcChain.xml"),
        _rel("rIdX1", f"{NS_R}/customXml", "../customXml/item1.xml"), _rel("rIdX2", f"{NS_R}/customXml", "../customXml/item2.xml")])
    ct = teile["[Content_Types].xml"].decode()
    dazu = ['<Default Extension="png" ContentType="image/png"/>',
            '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.printerSettings"/>',
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
            'sharedStrings+xml"/>',
            '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
            'calcChain+xml"/>',
            '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>',
            '<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'customXmlProperties+xml"/>',
            '<Override PartName="/customXml/itemProps2.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'customXmlProperties+xml"/>',
            '<Override PartName="/docMetadata/LabelInfo.xml" ContentType="application/vnd.ms-office.classificationlabels+xml"/>']
    ct = re.sub(r'<Default Extension="(png|bin)"[^>]*/>', "", ct)
    ct = ct.replace("</Types>", "".join(dazu) + "</Types>")
    teile["[Content_Types].xml"] = ct.encode()
    # 6. feste Zeitpunkte in den Dokumenteigenschaften
    teile["docProps/core.xml"] = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-03-01T00:00:00Z",
                                        teile["docProps/core.xml"])
    vorne = ["[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml", "xl/workbook.xml",
             "xl/_rels/workbook.xml.rels"]
    namen = [n for n in vorne if n in teile] + sorted(n for n in teile if n not in vorne)
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
        for n in namen:
            info = zipfile.ZipInfo(n, FESTE_ZEIT)
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, teile[n])
    pfad.write_bytes(puffer.getvalue())


# ====================================================================== Modell bauen, Wahrheit, Hauptfunktion
def baue_kontext(seed: int, cfg: dict[str, Any] | None = None, korpus: str = KORPUS) -> Kontext:
    """Das logische Modell eines Paars (deterministisch je Seed). cfg: zeilen=(min, max), aenderungen=n, leiste=bool ...
    korpus nachbau_test (oder cfg nur_test=True, nur fuer Tests des Generators mit Entwicklungs-Seeds) schaltet die NUR_TEST_FAMILIEN
    zu; in nachbau_word kommen sie nie vor und der Zufallsstrom bleibt unveraendert."""
    cfg = dict(cfg or {})
    rng = random.Random(seed * 1_000_003 + 3_000_017)
    ctx = Kontext(seed, rng, cfg)
    r, o = rng, ctx.opt
    o["leiste"] = cfg.get("leiste", r.random() < 0.8)
    o["formatiert"] = r.random() < 0.8
    o["kommentare"] = r.random() < 0.65
    o["kopf_aenderung"] = r.random() < 0.85
    o["fuss_aenderung"] = r.random() < 0.6
    o["hinweise"] = r.choice([0, 1, 1, 2])
    o["inhalt"] = r.random() < 0.35
    o["zwei_autoren"] = cfg.get("zwei_autoren", r.random() < 0.3)
    o["doppelbuchstaben"] = cfg.get("doppelbuchstaben", r.random() < 0.3 and cfg.get("zeilen", ZEILEN_STANDARD)[0] >= 150)
    o["projekt"] = r.choice(PROJEKTE)
    o["ort"] = r.choice(ORTE)
    o["titel"] = r.choice(SPEZ_TITEL)
    o["spez"] = f"{r.randint(10, 48)} {r.randint(10, 99)} {r.randint(10, 99)}"
    o["initialen"] = r.choice(INITIALEN)
    nb2, nb4 = r.randint(13, 19), r.randint(13, 19)
    o["polster"] = ("\t", "\xa0" * nb2 + " ", r.choice(["\t", "\xa0" * nb4 + " "]), "\xa0" * nb4 + " ")
    o["spez_excel"] = o["spez"].replace(" ", "_")
    o["spez_variante"] = r.choice(["ohne", "ohne", "leerzeichen", "gemischt"])
    o["e_variante"] = r.choice(["keine", "keine", "part", "alle"])
    o["datum"] = "2026_03_01" if r.random() < 0.7 else datetime(2026, 3, 1)
    o["einfarbig_als"] = r.choice(["zellschrift", "rich_text"])
    o["blattname"] = r.choice(["Register", "CDE", "Anforderungen"])
    o["alte_markierung"] = cfg.get("alte_markierung", r.random() < 0.5)
    o["fette_zellen"] = r.random() < 0.4
    o["bilder"] = r.choice([1, 2])
    ziel = r.randint(*cfg.get("zeilen", ZEILEN_STANDARD))
    _baue_struktur(ctx, ziel)
    _kuerze(ctx, ziel)
    anzahl = cfg.get("aenderungen") or max(8, int(ziel * r.uniform(0.10, 0.17)))
    _plane_aenderungen(ctx, anzahl)
    _nummeriere(ctx)
    _abweichungen(ctx)
    if korpus == KORPUS_TEST or cfg.get("nur_test"):
        _familie_silbentrennung(ctx)
        _familie_umbruch_nummer(ctx)
    if (korpus == KORPUS and seed >= UMBRUCH_LEERZEICHEN_AB) or cfg.get("umbruch_leerzeichen"):
        _familie_umbruch_leerzeichen(ctx)             # nur neue Seeds: Paare aelterer Seeds (301-360) bleiben bytegleich
    ziel = Block("ziel", ctx.absaetze, o["spez_excel"], o["spez_variante"], "A.3.0", o["datum"])
    ctx.bloecke = _fremde_specs(ctx, ziel) if cfg.get("mehrere_specs", seed % 2 == 0) else [ziel]
    _excel_werte(ctx)
    # Stoerer-Familien, die aus dem Modell folgen
    o_st = ctx.markiere
    o_st("markup_leiste" if o["leiste"] else "nur_inline")
    if o["kopf_aenderung"] or o["fuss_aenderung"]:
        o_st("kopf_fuss_aenderung")
    alle_edits = [e for a in ctx.absaetze for e in a.edits]
    if any(e.autor == 1 for e in alle_edits):
        o_st("zweite_autorfarbe")
    if any(a.label_alt != a.label_neu and a.status is None for a in ctx.absaetze):
        o_st("listennummer_gruen")
    if any(a.stufe >= 3 for a in ctx.absaetze):
        o_st("listenebenen")
    if any(a.gelb for a in ctx.absaetze):
        o_st("platzhalter_gelb")
    if any(e.art == "klammern" for e in alle_edits):
        o_st("klammer_zeichengenau")
    if any(a.status == "geloescht" for a in ctx.absaetze):
        o_st("absatz_geloescht")
    if any(a.status == "neu" for a in ctx.absaetze):
        o_st("absatz_neu")
    return ctx


def _woerter_von(text: str) -> int:
    return len(text.split())


def _eintraege(ctx: Kontext, pos_v: dict[int, int], pos_n: dict[int, int], setzer: Setzer) -> list[dict[str, Any]]:
    """Die Wahrheit je Aenderung (aus dem Modell). zeile/spalte beziehen sich auf vorher.xlsx (neue Zeile: Zeile davor)."""
    out: list[dict[str, Any]] = []

    def neu(**kw: Any) -> None:
        out.append({"nr": len(out) + 1, **kw})

    vor_zeile = min(pos_v[a.id] for a in ctx.absaetze if a.id in pos_v) - 1
    for a in ctx.absaetze:
        abweichung = a.excel_alt is not None and (bool(a.edits) or a.label_alt != a.label_neu)
        soll = "pruefen" if abweichung else "uebernommen"
        basis = {"blatt": HAUPTBLATT_NR, "spalte": G, "absatz": a.id, "seite": a.seite, "bereich": "cde",
                 "soll_status": soll, "stoerer": sorted(set(a.stoerer))}
        if abweichung:
            basis["vorschlag"] = [list(x) for x in excel_laeufe(ctx, a, True)]
        if a.status == "neu":
            neu(typ="absatz_neu", art=a.edits[0].art, zeile=vor_zeile, nachher_zeile=pos_n[a.id],
                woerter=_woerter_von(a.neu_text()) + 1, **basis)
            continue
        if a.status == "geloescht":
            neu(typ="absatz_geloescht", art="absatz_geloescht", zeile=pos_v[a.id], nachher_zeile=pos_n[a.id],
                woerter=_woerter_von(a.text) + 1, **basis)
        else:
            if a.label_alt != a.label_neu:
                neu(typ="text_ersetzt", art="listennummer", zeile=pos_v[a.id], nachher_zeile=pos_n[a.id], woerter=2, **basis)
            gruppen: list[list[Edit]] = []
            for e in sorted(a.edits, key=lambda x: x.start):
                if gruppen and e.art == "klammern" and gruppen[-1][-1].art == "klammern" and len(gruppen[-1]) == 1:
                    gruppen[-1].append(e)
                else:
                    gruppen.append([e])
            for g in gruppen:
                weg = "".join(a.text[e.start:e.ende] for e in g)
                dazu = "".join(e.neu for e in g)
                typ = "text_ersetzt" if weg and dazu else ("text_geloescht" if weg else "text_eingefuegt")
                neu(typ=typ, art=g[0].art, zeile=pos_v[a.id], nachher_zeile=pos_n[a.id],
                    woerter=max(1, _woerter_von(weg) + _woerter_von(dazu)), **basis)
        vor_zeile = pos_v[a.id]
    for x in setzer.ausserhalb:
        neu(typ=x["typ"], art=x["bereich"], blatt=0, zeile=0, spalte=0, nachher_zeile=0, woerter=x["woerter"],
            seite=0, bereich="ausserhalb", soll_status="ausserhalb", stoerer=[x["bereich"]])
    return out


def erzeuge(seed: int, ziel: Path | str, korpus: str = KORPUS, cfg: dict[str, Any] | None = None, name: str = PROFIL) -> Path:
    """Erzeugt ein Paar unter ziel/synth_<seed>_<name>/ und gibt den Ordner zurueck (append-only, Regel 29)."""
    ordner = Path(ziel) / f"synth_{seed:04d}_{name}"
    if ordner.exists():
        raise FileExistsError(f"{ordner} existiert schon (Paare sind append-only)")
    ctx = baue_kontext(seed, cfg, korpus)
    pdf, setzer = rendere_pdf(ctx)
    tmp = Path(ziel) / f".tmp_{ordner.name}"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    try:
        pdf.speichern(tmp / "aenderungen.pdf")
        pos_v, formeln = schreibe_excel(tmp / "vorher.xlsx", ctx, False)
        nachbearbeiten(tmp / "vorher.xlsx", ctx, formeln)
        pos_n, _ = schreibe_excel(tmp / "nachher.xlsx", ctx, True)
        nachbearbeiten(tmp / "nachher.xlsx", ctx, formeln)
        eintraege = _eintraege(ctx, pos_v, pos_n, setzer)
        abw = [{"blatt": HAUPTBLATT_NR, "zeile": pos_v[a.id], "spalte": G} for a in ctx.absaetze
               if a.excel_alt is not None and not a.edits and a.label_alt == a.label_neu]
        erwartet = {
            "format": FORMAT,
            "generator_version": GENERATOR_VERSION,
            "seed": seed,
            "profil": name,
            "korpus": korpus,
            "erzeuger": "writer",
            "pdf_stil": PROFIL,
            "legende": LEGENDE,
            "nachher_format": NACHHER_FORMAT,
            "stoerer": sorted(ctx.stoerer),
            "layout": "cde_register",
            "hauptblatt": ctx.opt["blattname"],
            "woerter_markiert": setzer.woerter_markiert,
            "aenderungen": eintraege,
            "abweichungen_ohne_aenderung": abw,
            # Spec-Bloecke im Register (Stoerer mehrere_specs): nur die Ziel-Spec (= Spec des PDF) darf sich aendern
            "specs": [{"rolle": blk.rolle, "spez": blk.spez_excel, "revision": blk.rev,
                       "zeilen_vorher": [min(pos_v[a.id] for a in blk.absaetze if a.id in pos_v),
                                         max(pos_v[a.id] for a in blk.absaetze if a.id in pos_v)],
                       "zeilen_nachher": [min(pos_n[a.id] for a in blk.absaetze), max(pos_n[a.id] for a in blk.absaetze)]}
                      for blk in ctx.bloecke],
            "luecken": [],
            "excel": {"datenstart": DATEN_START, "zeilen_vorher": len(pos_v), "zeilen_nachher": len(pos_n),
                      "zusatzteile": list(ZUSATZTEILE[:4]) + [f"xl/media/image{i + 1}.png" for i in range(ctx.opt["bilder"])]
                      + list(ZUSATZTEILE[6:])},
            "pdf_info": {"seiten": len(pdf.seiten), "leiste": ctx.opt["leiste"], "sprechblasen": setzer.blasen,
                         "kommentare": setzer.kommentare, "formatiert": setzer.formatiert,
                         "umbrueche_im_absatz": sum(1 for a in ctx.absaetze if a.seite_ende > a.seite)},
        }
        with open(tmp / "erwartet.json", "w", encoding="utf-8") as f:
            json.dump(erwartet, f, ensure_ascii=False, indent=2)
            f.write("\n")
        tmp.rename(ordner)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise
    return ordner
