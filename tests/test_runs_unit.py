"""Unit-Tests fuer src/runs.py (Laeufe, atomares Schreiben, Archiv, Schema). Nur erfundene Daten (Regel 16)."""
import json
import threading
from pathlib import Path

import pytest

from src import runs
from tests.fixtures.szenario import erzeuge


@pytest.fixture
def fx(tmp_path):
    return erzeuge(tmp_path / "fx")


@pytest.fixture
def lauf(fx):
    rid = runs.neuer_lauf(fx["pdf"], fx["alt"])
    runs.analysiere(rid)
    return rid


# --- Dateinamen -------------------------------------------------------------------------------
@pytest.mark.parametrize("roh, erwartet_stamm", [
    ("CON.pdf", "_CON"), ("nul.xlsx", "_nul"), ("Com1.pdf", "_Com1"), ("lpt9.txt.pdf", "_lpt9.txt"),
    ("normal.pdf", "normal"), ("Bericht Ä.pdf", "Bericht Ä")])
def test_sicherer_name_reservenamen(roh, erwartet_stamm):
    assert Path(runs._sicherer_name(roh, "x.pdf")).stem == erwartet_stamm


def test_sicherer_name_punkte_und_leer():
    assert runs._sicherer_name("..", "fallback.pdf") == "fallback.pdf"
    assert runs._sicherer_name("name. .", "f.pdf") == "name"
    assert runs._sicherer_name("a:b?c.pdf", "f.pdf") == "a_b_c.pdf"


def test_zu_langer_name_legt_nichts_an(fx):
    with pytest.raises(ValueError, match="zu lang"):
        runs.neuer_lauf(fx["pdf"], fx["alt"], "a" * 256 + ".pdf", "cde.xlsx")
    with pytest.raises(ValueError, match="zu lang"):
        runs.neuer_lauf(fx["pdf"], fx["alt"], "a.pdf", "ä" * 130 + ".xlsx")     # 255+ Bytes, aber < 255 Zeichen
    assert list(runs.runs_dir().iterdir()) == []


def test_name_mit_255_zeichen_ist_erlaubt(fx):
    rid = runs.neuer_lauf(fx["pdf"], fx["alt"], "a" * 251 + ".pdf", "cde.xlsx")
    assert (runs.lauf_dir(rid) / "eingabe" / ("a" * 251 + ".pdf")).is_file()


def test_fehler_beim_kopieren_hinterlaesst_keinen_halben_ordner(fx, tmp_path):
    with pytest.raises(FileNotFoundError):
        runs.neuer_lauf(fx["pdf"], tmp_path / "gibt_es_nicht.xlsx")
    assert list(runs.runs_dir().iterdir()) == []


# --- Atomares Schreiben -----------------------------------------------------------------------
def test_atomar_schreiben_ersetzt_ohne_reste(tmp_path):
    ziel = tmp_path / "x.json"
    runs._schreibe_json_atomar(ziel, {"a": 1})
    runs._schreibe_json_atomar(ziel, {"a": 2})
    assert json.loads(ziel.read_text(encoding="utf-8"))["a"] == 2
    assert [p.name for p in tmp_path.iterdir()] == ["x.json"]


def test_atomar_schreiben_bei_fehler_bleibt_alte_datei_und_kein_tmp(tmp_path):
    ziel = tmp_path / "x.json"
    runs._schreibe_json_atomar(ziel, {"a": 1})

    class Kaputt:                                  # nicht als JSON schreibbar
        pass
    with pytest.raises(TypeError):
        runs._schreibe_json_atomar(ziel, {"a": Kaputt()})
    assert json.loads(ziel.read_text(encoding="utf-8"))["a"] == 1
    assert [p.name for p in tmp_path.iterdir()] == ["x.json"]


# --- Ergebnis-Ordner --------------------------------------------------------------------------
def test_ergebnis_nummer_ist_hoechste_plus_eins_nicht_anzahl(lauf):
    d = runs.lauf_dir(lauf)
    (d / "ergebnis_001").mkdir()
    (d / "ergebnis_003").mkdir()                   # Luecke: Anzahl+1 waere 3 und wuerde kollidieren
    assert runs._neuer_ergebnis_ordner(d).name == "ergebnis_004"


def test_ergebnis_ordner_wiederholt_bei_kollision(lauf, monkeypatch):
    d = runs.lauf_dir(lauf)
    echt = Path.mkdir
    zaehler = {"n": 0}

    def mkdir(self, *a, **k):
        if self.name.startswith("ergebnis_") and zaehler["n"] == 0:
            zaehler["n"] += 1
            echt(self)                             # ein anderer Vorgang war schneller
        return echt(self, *a, **k)
    monkeypatch.setattr(Path, "mkdir", mkdir)
    assert runs._neuer_ergebnis_ordner(d).name == "ergebnis_002"


def test_gleichzeitige_uebernahmen_ergeben_verschiedene_ordner(lauf):
    ordner, fehler = [], []

    def arbeit():
        try:
            ordner.append(runs.uebernehmen_lauf(lauf, [])["ordner"])
        except Exception as e:                     # noqa: BLE001
            fehler.append(repr(e))
    ts = [threading.Thread(target=arbeit) for _ in range(4)]
    for t in ts:
        t.start()
    for t in ts:
        t.join(60)
    assert not fehler
    assert sorted(ordner) == [f"ergebnis_{i:03d}" for i in range(1, 5)]
    assert len(runs.ergebnisse(lauf)) == 4


def test_abbruch_zaehlt_nicht_und_hinterlaesst_nichts(lauf, monkeypatch):
    with monkeypatch.context() as m:
        m.setattr(runs, "uebernehmen", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("Abbruch")))
        with pytest.raises(RuntimeError):
            runs.uebernehmen_lauf(lauf, [])
    assert list(runs.lauf_dir(lauf).glob("ergebnis_*")) == []
    assert runs.ergebnisse(lauf) == []
    assert runs.uebernehmen_lauf(lauf, [])["ordner"] == "ergebnis_001"      # Nummer wird wieder frei


def test_halber_ergebnis_ordner_zaehlt_nicht_und_wird_nicht_wiederverwendet(lauf):
    d = runs.lauf_dir(lauf)
    (d / "ergebnis_001").mkdir()                   # z. B. Programm wurde hart beendet
    assert runs.ergebnisse(lauf) == []
    assert runs.liste()[0]["ergebnisse"] == 0
    assert runs.uebernehmen_lauf(lauf, [])["ordner"] == "ergebnis_002"
    assert [e["ordner"] for e in runs.ergebnisse(lauf)] == ["ergebnis_002"]
    assert (d / "ergebnis_001").is_dir()           # nichts wird geloescht


def test_kaputte_ergebnis_json_zaehlt_nicht(lauf):
    runs.uebernehmen_lauf(lauf, [])
    e = runs.lauf_dir(lauf) / "ergebnis_001" / "ergebnis.json"
    e.write_text("{halb", encoding="utf-8")
    assert runs.ergebnisse(lauf) == []
    assert runs.liste()[0]["ergebnisse"] == 0


# --- Schema -----------------------------------------------------------------------------------
@pytest.mark.parametrize("v", [1, 2])
def test_schema_1_und_2_werden_gelesen(v):
    assert runs.pruefe_schema({"schema_version": v}) == v


def test_schema_fehlt_gilt_als_1():
    assert runs.pruefe_schema({}) == 1


def test_schema_zukunft_gibt_klare_meldung():
    with pytest.raises(runs.LaufFehler) as e:
        runs.pruefe_schema({"schema_version": 99})
    assert "neueren Programmversion" in e.value.meldung and e.value.art == "neuere_version"


@pytest.mark.parametrize("v", [0, -1, "2", None, 1.5, True])
def test_schema_unsinn_ist_beschaedigt(v):
    with pytest.raises(runs.LaufFehler) as e:
        runs.pruefe_schema({"schema_version": v})
    assert e.value.art == "beschaedigt"


def test_ergebnis_json_aus_der_zukunft_wird_nicht_gezaehlt(lauf):
    runs.uebernehmen_lauf(lauf, [])
    e = runs.lauf_dir(lauf) / "ergebnis_001" / "ergebnis.json"
    j = json.loads(e.read_text(encoding="utf-8"))
    j["schema_version"] = 99
    e.write_text(json.dumps(j), encoding="utf-8")
    assert runs.ergebnisse(lauf) == []


def test_plan_ohne_pflichtfelder_ist_beschaedigt(lauf):
    pj = runs.lauf_dir(lauf) / "plan.json"
    pj.write_text(json.dumps({"schema_version": 2, "bilanz": {}}), encoding="utf-8")
    with pytest.raises(runs.LaufFehler) as e:
        runs.lade_plan(lauf)
    assert e.value.meldung == "Der Lauf ist beschädigt."


# --- plan.json fehlt / Analyse fehlgeschlagen --------------------------------------------------
def test_fehlende_plan_json_startet_keine_neuanalyse(fx):
    rid = runs.neuer_lauf(fx["pdf"], fx["alt"])
    with pytest.raises(runs.LaufFehler) as e:
        runs.lade_plan(rid)
    assert e.value.status == 409
    assert not (runs.lauf_dir(rid) / "plan.json").exists()


def test_fehlgeschlagene_analyse_bleibt_als_lauf_stehen(fx):
    rid = runs.neuer_lauf(fx["pdf"], fx["alt"])
    runs.markiere_analyse_fehler(rid, "Die Analyse ist fehlgeschlagen. Test.")
    eintrag = runs.liste()[0]
    assert eintrag["run_id"] == rid and eintrag["status"] == "fehlgeschlagen" and "Test" in eintrag["meldung"]
    assert (runs.lauf_dir(rid) / "eingabe").is_dir()
    with pytest.raises(runs.LaufFehler) as e:
        runs.lade_plan(rid)
    assert e.value.art == "fehlgeschlagen" and "Test" in e.value.meldung


# --- Liste ------------------------------------------------------------------------------------
def test_liste_status_der_laeufe(fx, lauf):
    assert runs.liste()[0]["status"] == "analysiert"
    runs.uebernehmen_lauf(lauf, [])
    eintrag = runs.liste()[0]
    assert eintrag["status"] == "uebernommen" and eintrag["ergebnisse"] == 1 and eintrag["bilanz"]["stimmt"]


def test_liste_kaputter_lauf_stoert_die_anderen_nicht(fx, lauf):
    kaputt1 = runs.neuer_lauf(fx["pdf"], fx["alt"])
    kaputt2 = runs.neuer_lauf(fx["pdf"], fx["alt"])
    runs.analysiere(kaputt1)
    runs.analysiere(kaputt2)
    (runs.lauf_dir(kaputt1) / "plan.json").write_text('{"schema_version": 2, "bil', encoding="utf-8")
    (runs.lauf_dir(kaputt2) / "lauf.json").write_text("kein json", encoding="utf-8")
    nach_id = {e["run_id"]: e for e in runs.liste()}
    assert nach_id[lauf]["status"] == "analysiert"
    assert nach_id[kaputt1]["status"] == "beschaedigt" and nach_id[kaputt2]["status"] == "beschaedigt"
    assert nach_id[kaputt1]["meldung"] == "Der Lauf ist beschädigt."


def test_liste_plan_aus_der_zukunft(lauf):
    pj = runs.lauf_dir(lauf) / "plan.json"
    j = json.loads(pj.read_text(encoding="utf-8"))
    j["schema_version"] = 99
    pj.write_text(json.dumps(j), encoding="utf-8")
    assert runs.liste()[0]["status"] == "neuere_version"
    with pytest.raises(runs.LaufFehler, match="neueren Programmversion"):
        runs.uebernehmen_lauf(lauf, [])
    assert list(runs.lauf_dir(lauf).glob("ergebnis_*")) == []


def test_liste_ignoriert_temp_dateien_und_fremde_ordner(lauf):
    (runs.runs_dir() / "kein-lauf").mkdir()
    (runs.runs_dir() / "20260101-000000-abcdef").mkdir()           # ohne lauf.json (z. B. im Aufbau)
    assert [e["run_id"] for e in runs.liste()] == [lauf]


# --- Archiv -----------------------------------------------------------------------------------
def test_archivieren_loescht_nichts(lauf):
    vorher = sorted(str(p.relative_to(runs.lauf_dir(lauf))) for p in runs.lauf_dir(lauf).rglob("*"))
    runs.archiviere(lauf)
    assert runs.liste() == []
    assert [e["run_id"] for e in runs.liste(mit_archiv=True)] == [lauf]
    assert runs.liste(mit_archiv=True)[0]["archiviert"] is True
    nachher = sorted(str(p.relative_to(runs.lauf_dir(lauf))) for p in runs.lauf_dir(lauf).rglob("*"))
    assert nachher == sorted(vorher + ["archiviert.json"])
    assert runs.lade_plan(lauf)["archiviert"] is True               # Lauf bleibt les- und nutzbar


def test_archivieren_ist_wiederholbar_und_ueberschreibt_nichts(lauf):
    runs.archiviere(lauf)
    marker = runs.lauf_dir(lauf) / "archiviert.json"
    inhalt = marker.read_text(encoding="utf-8")
    runs.archiviere(lauf)
    assert marker.read_text(encoding="utf-8") == inhalt


def test_wiederherstellen_behaelt_den_verlauf(lauf):
    runs.archiviere(lauf)
    runs.stelle_wieder_her(lauf)
    assert [e["run_id"] for e in runs.liste()] == [lauf]
    assert runs.liste()[0]["archiviert"] is False
    d = runs.lauf_dir(lauf)
    assert not (d / "archiviert.json").exists()
    assert len(list(d.glob("archiviert_*.json.alt"))) == 1          # alter Marker bleibt als Spur
    runs.stelle_wieder_her(lauf)                                   # nochmal: kein Fehler
    runs.archiviere(lauf)
    runs.stelle_wieder_her(lauf)
    assert len(list(d.glob("archiviert_*.json.alt"))) == 2


def test_archiv_mit_unbekannter_id():
    with pytest.raises(FileNotFoundError):
        runs.archiviere("20260101-000000-abcdef")
    with pytest.raises(ValueError):
        runs.archiviere("../x")


# --- Namen im Plan ----------------------------------------------------------------------------
def test_lauf_json_mit_pfad_im_namen_ist_beschaedigt(lauf):
    (runs.lauf_dir(lauf) / "lauf.json").write_text(
        json.dumps({"pdf": "../x.pdf", "excel": "cde.xlsx", "angelegt": ""}), encoding="utf-8")
    with pytest.raises(runs.LaufFehler):
        runs.uebernehmen_lauf(lauf, [])


# --- Abwahl, Zusammenfassung (Dashboard, Auftrag 003) -----------------------------------------------
@pytest.fixture
def lauf_mit_vorschlag(tmp_path):
    """Erfundenes Paar, in dem genau ein Vorschlag ("pruefen") steckt (teilweise hinterlegtes Wort)."""
    fx = erzeuge(tmp_path / "fx_h", "hintergrund")
    rid = runs.neuer_lauf(fx["pdf"], fx["alt"])
    plan = runs.analysiere(rid)
    (vorschlag,) = [a for a in plan["aenderungen"] if a["status"] == "pruefen"]          # "Rack{+s+}"
    # die automatische Ersetzung "{-zwei-} {+drei+} Wochen" (keine Gruppe, andere Zelle als der Vorschlag)
    (auto,) = [a for a in plan["aenderungen"] if a["status"] == "uebernommen" and (a["alt"], a["neu"]) == ("zwei", "drei")]
    assert not auto["gruppe"] and auto["ziel"] != vorschlag["ziel"]
    return rid, vorschlag, auto


def _cde_text(pfad: Path) -> str:
    from openpyxl import load_workbook
    ws = load_workbook(str(pfad))["CDE"]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_abwahl_wird_nicht_eingetragen_freigabe_schon(lauf_mit_vorschlag):
    rid, vorschlag, auto = lauf_mit_vorschlag
    res = runs.uebernehmen_lauf(rid, [vorschlag["id"]], [auto["id"]])
    nach_id = {a["id"]: a for a in res["aenderungen"]}
    assert nach_id[auto["id"]]["status"] == "nicht_uebernommen" and "abgewaehlt" in nach_id[auto["id"]]["pruefcodes"]
    assert nach_id[vorschlag["id"]]["status"] == "uebernommen" and nach_id[vorschlag["id"]]["bestaetigt"]
    assert res["abwahl"] == [auto["id"]] and res["freigaben"] == [vorschlag["id"]]
    assert res["bilanz"]["stimmt"]                                         # nichts geht still verloren
    text = _cde_text(runs.lauf_dir(rid) / res["ordner"] / res["excel"])
    assert "dauert zwei Wochen" in text and "dauert drei Wochen" not in text      # abgewaehlt: bleibt wie vorher
    assert "Die Racks werden" in text                                              # freigegeben: eingetragen
    gespeichert = json.loads((runs.lauf_dir(rid) / res["ordner"] / "ergebnis.json").read_text(encoding="utf-8"))
    assert gespeichert["abwahl"] == [auto["id"]] and gespeichert["freigaben"] == [vorschlag["id"]]
    assert runs.ergebnisse(rid)[0]["abwahl"] == [auto["id"]]


@pytest.mark.parametrize("fall", ["unbekannt", "nicht_automatisch", "beides"])
def test_abwahl_ungueltig_legt_kein_ergebnis_an(lauf_mit_vorschlag, fall):
    rid, vorschlag, auto = lauf_mit_vorschlag
    freig, abwahl, erwartet = {
        "unbekannt": ([], ["A9999"], "Unbekannte"),
        "nicht_automatisch": ([], [vorschlag["id"]], "nur bei automatisch"),
        "beides": ([auto["id"]], [auto["id"]], "zugleich"),
    }[fall]
    with pytest.raises(ValueError, match=erwartet) as e:
        runs.uebernehmen_lauf(rid, freig, abwahl)
    assert getattr(e.value, "status", None) == 400
    assert list(runs.lauf_dir(rid).glob("ergebnis_*")) == []


def test_aeltere_ergebnisse_ohne_abwahl_bleiben_lesbar(lauf):
    runs.uebernehmen_lauf(lauf, [])
    e = runs.lauf_dir(lauf) / "ergebnis_001" / "ergebnis.json"
    j = json.loads(e.read_text(encoding="utf-8"))
    j.pop("abwahl")
    e.write_text(json.dumps(j), encoding="utf-8")
    assert runs.ergebnisse(lauf)[0]["abwahl"] == []


@pytest.mark.parametrize("status, bereich, erwartet", [
    ("uebernommen", "", "automatisch"), ("uebernommen", "ausserhalb", "automatisch"), ("pruefen", "", "pruefen"),
    ("pruefen", "ausserhalb", "ausserhalb"), ("nicht_uebernommen", "ausserhalb", "ausserhalb"),
    ("nicht_uebernommen", "", "nicht")])
def test_kategorie_fuer_das_dashboard(status, bereich, erwartet):
    assert runs.kategorie({"status": status, "bereich": bereich}) == erwartet


def test_zusammenfassung_zaehlt_jede_aenderung_genau_einmal():
    l = [{"status": "uebernommen"}, {"status": "uebernommen", "bereich": ""}, {"status": "pruefen"},
         {"status": "nicht_uebernommen", "bereich": "ausserhalb"}, {"status": "nicht_uebernommen"}]   # Schema 1: ohne bereich
    z = runs.zusammenfassung(l)
    assert z == {"gesamt": 5, "automatisch": 2, "pruefen": 1, "ausserhalb": 1, "nicht": 1}
    assert z["gesamt"] == z["automatisch"] + z["pruefen"] + z["ausserhalb"] + z["nicht"]


def test_plan_speichert_legende_und_sprechblasen_als_zahl(lauf):
    plan = runs.lade_plan(lauf)
    assert isinstance(plan["legende_modus"], str) and plan["legende_modus"]
    assert plan["sprechblasen"] == 0


# --- Spec (Spec-Filter) -----------------------------------------------------------------------------
@pytest.mark.parametrize("roh, erwartet", [(None, None), ("", None), ("   ", None), ("11 22 33", "11 22 33"), ("  11   22 33 ", "11 22 33"),
                                           ("11_22.33-44", "11_22.33-44"), ("A" * 40, "A" * 40)])
def test_pruefe_spec_gueltig(roh, erwartet):
    assert runs.pruefe_spec(roh) == erwartet


@pytest.mark.parametrize("roh", ["x" * 41, "11;22", "<b>", "11\n22; rm", "-11", "../x", "11 22 33'", 5, ["1"]])
def test_pruefe_spec_ungueltig(roh):
    with pytest.raises(ValueError) as e:
        runs.pruefe_spec(roh)
    assert getattr(e.value, "status", None) == 400


def test_neu_analysieren_laesst_den_alten_lauf_unberuehrt(lauf):
    d = runs.lauf_dir(lauf)
    vorher = {str(p.relative_to(d)): p.read_bytes() for p in d.rglob("*") if p.is_file()}
    neu, plan = runs.neu_analysieren(lauf, "11 22 33")
    assert neu != lauf and plan["run_id"] == neu and runs.lade_plan(neu)["spec"] is None      # CDE ohne Spec-Spalte
    assert {str(p.relative_to(d)): p.read_bytes() for p in d.rglob("*") if p.is_file()} == vorher
    assert (runs.lauf_dir(neu) / "eingabe" / "cde_alt.xlsx").read_bytes() == (d / "eingabe" / "cde_alt.xlsx").read_bytes()
