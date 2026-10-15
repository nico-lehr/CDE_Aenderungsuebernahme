"""Red-Team R3 (Analyse der stillen Fehler im Entwicklungskorpus), Teil 3: SCHWAECHEN DER MESSUNG (Sabotage-Tests).

Frage: Wo koennte `python -m src.cli messen` ein Team taeuschen, weil die Zahlen zu GUT aussehen? Methode: Ein sauberes
Basis-Paar (das Tool macht alles richtig, 0 stille Fehler) wird so SABOTIERT, wie ein fehlerhaftes Tool es tun koennte -- die
Sabotage greift nur im Messwerkzeug an (das Ergebnis des Tools wird verfaelscht), nichts in src/ wird geaendert.

  * Teil A (WAECHTER, normale Tests, muessen gruen bleiben): Sabotage, die die Messung bemerkt. Das schuetzt davor, dass
    ein spaeterer Umbau von messen.py diese Faelle verliert.
  * Teil B (xfail strict, Besitzer Mess-Bauer): Sabotage, die die Messung NICHT bemerkt. Wird die Messung besser,
    wird der Test XPASS (rot) und der Mess-Bauer entfernt den Marker.

Alles ERFUNDEN (Regel 16). Vorschlaege zur Reparatur: tests/redteam/README_korpus.md.
"""
from __future__ import annotations

import json
from pathlib import Path
from unittest import mock

import pytest
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from src import config, messen
from src.models import PRUEFEN
from tests.fixtures.pdf_writer import PdfDokument, baue_pdf, scan_bild
from tests.fixtures.szenario import baue_excel, zustaende
from tests.redteam._bau import aufbau, scheitern
from tests.redteam._bau_korpus import (BASIS, KOPF, auto_lauf, baue_basis_paar, finde_zelle, messe_mit_fertigem_ergebnis,
                                       messe_mit_sabotage, paar)


def _bemerkt(res: dict) -> bool:
    """Hat die Messung IRGENDETWAS bemerkt (stiller Fehler, Kollateral, falsche Zelle, nicht erreichte Soll-Zelle)?"""
    return bool(res["stille_fehler"] or res["kollateral"] or res["auto_richtig"] != res["auto_gesamt"]
                or res["erreicht_auto"] != res["soll_zellen"])


def _basis_sauber(tmp_path) -> Path:
    ordner = baue_basis_paar(tmp_path / "basis")
    sauber = messen.messe_paar(ordner, False)
    aufbau(not _bemerkt(sauber) and sauber["soll_zellen"] >= 3 and sauber["idempotent"],
           f"Das Basis-Paar ist nicht sauber: {sauber['stille_fehler']} stille Fehler")
    return ordner


# ---- Sabotage-Funktionen (Workbook -> Workbook) ----------------------------------------------------------------------
def _cde(wb):
    return wb["CDE"]


def _ersetze_in(wb, suche: str, alt: str, neu: str) -> None:
    c = finde_zelle(_cde(wb), suche)
    aufbau(alt in str(c.value), f"'{alt}' steht nicht in der Zelle")
    c.value = str(c.value).replace(alt, neu)


SABOTAGE_ERKANNT = {
    "falsches_wort": lambda wb: _ersetze_in(wb, "monatlich", "monatlich", "woechentlich"),
    "zahl_vertauscht": lambda wb: _ersetze_in(wb, "ms nicht", "120", "12"),
    "aenderung_zurueckgenommen": lambda wb: _ersetze_in(wb, "monatlich", "monatlich", "jaehrlich"),
    "absaetze_vertauscht": lambda wb: _vertausche(wb),
    "text_in_bemerkungsspalte": lambda wb: setattr(_cde(wb).cell(3, 4), "value", "XXX"),
    "zeile_doppelt_eingefuegt": lambda wb: _doppelte_zeile(wb),
    "zeile_geloescht": lambda wb: _cde(wb).delete_rows(finde_zelle(_cde(wb), "Lieferung umfasst").row),
    "anderes_blatt_veraendert": lambda wb: wb["Deckblatt"].__setitem__("A2", "Geaendert"),
}


def _vertausche(wb) -> None:
    a, b = finde_zelle(_cde(wb), "Messprotokoll"), finde_zelle(_cde(wb), "Alle Steckverbindungen")
    a.value, b.value = b.value, a.value


def _doppelte_zeile(wb) -> None:
    ws = _cde(wb)
    r = finde_zelle(ws, "Ersatzteile").row
    ws.insert_rows(r + 1)
    for c in range(1, 5):
        ws.cell(r + 1, c).value = ws.cell(r, c).value


# ==================================================================================================
# Teil A: WAECHTER -- diese Sabotage bemerkt die Messung heute
# ==================================================================================================
def test_waechter_basis_paar_ist_sauber(tmp_path):
    res = messen.messe_paar(baue_basis_paar(tmp_path / "basis"), False)
    assert not _bemerkt(res) and res["idempotent"] and res["original_unveraendert"]


@pytest.mark.parametrize("name", sorted(SABOTAGE_ERKANNT))
def test_waechter_messung_bemerkt_sabotage(tmp_path, name):
    ordner = _basis_sauber(tmp_path)
    res = messe_mit_sabotage(ordner, SABOTAGE_ERKANNT[name])
    assert _bemerkt(res), f"Sabotage '{name}' wurde von der Messung nicht bemerkt: {res['stille_fehler']} stille Fehler"
    assert res["stille_fehler"] >= 1, f"Sabotage '{name}' ist kein stiller Fehler in der Messung"


def test_waechter_zurueckgenommene_aenderung_heisst_keine_wirkung(tmp_path):
    res = messe_mit_sabotage(_basis_sauber(tmp_path), SABOTAGE_ERKANNT["aenderung_zurueckgenommen"])
    assert "keine_wirkung" in res["stille_fehler_nach_art"]


# ==================================================================================================
# Teil B1: Die Messung vergleicht nur NORMALISIERTEN TEXT (messen._k: NFKC, unsichtbare Zeichen weg, Leerraum zusammen)
# ==================================================================================================
SABOTAGE_TEXT = {
    "absaetze_zu_einem_verschmolzen": lambda wb: _ersetze_in(wb, "Ersatzteile", "\n", " "),
    "doppelte_leerzeichen": lambda wb: _ersetze_in(wb, "Messprotokoll", " ", "  "),
    "geschuetzte_leerzeichen": lambda wb: _ersetze_in(wb, "Messprotokoll", " ", " "),
    "weicher_trennstrich_im_wort": lambda wb: _ersetze_in(wb, "Messprotokoll", "Messprotokoll", "Mess­protokoll"),
    "unsichtbares_zeichen_im_wort": lambda wb: _ersetze_in(wb, "Messprotokoll", "Messprotokoll", "Mess​protokoll"),
    "ligatur_fi": lambda wb: _ersetze_in(wb, "findet", "findet", "ﬁndet"),
}


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-135 MESSLUECKE: Verfaelschungen, die nur Leerraum, Absatzgrenzen (Zeilenumbruch), weiche/unsichtbare "
                          "Zeichen oder Ligaturen betreffen, sind fuer die Messung unsichtbar (_k normalisiert alles weg); "
                          "u. a. RT-06/07/12 (doppeltes Leerzeichen, Ligatur) und verschmolzene Absaetze bleiben unbemerkt; "
                          "Besitzer Mess-Bauer (messen.py, _k)")
@pytest.mark.parametrize("name", sorted(SABOTAGE_TEXT))
def test_rt135_messung_bemerkt_unsichtbare_textverfaelschung(tmp_path, name):
    res = messe_mit_sabotage(_basis_sauber(tmp_path), SABOTAGE_TEXT[name])
    if not _bemerkt(res):
        scheitern(f"Verfaelschung '{name}' bleibt unbemerkt: stille Fehler {res['stille_fehler']}, Kollateral "
                  f"{res['kollateral']}, Treffergenauigkeit {res['auto_richtig']} von {res['auto_gesamt']} = 100 %")


# ==================================================================================================
# Teil B2: Die Messung vergleicht nur WERTE -- Struktur und Format sind unsichtbar (vgl. R1: RT-18..26, RT-35/36)
# ==================================================================================================
def _verbinden(wb) -> None:
    _cde(wb).merge_cells("G2:H3")


def _pruefung(wb) -> None:
    dv = DataValidation(type="list", formula1='"a,b"')
    dv.add("D2:D6")
    _cde(wb).add_data_validation(dv)


def _fett_rot(wb) -> None:
    for zeile in _cde(wb).iter_rows(min_row=2):
        for c in zeile:
            c.font = Font(bold=True, color="FF0000")


SABOTAGE_STRUKTUR = {
    "blatt_umbenannt": lambda wb: setattr(_cde(wb), "title", "CDE_neu"),
    "blatt_ausgeblendet": lambda wb: setattr(wb["Deckblatt"], "sheet_state", "hidden"),
    "zeile_ausgeblendet": lambda wb: setattr(_cde(wb).row_dimensions[3], "hidden", True),
    "spalte_ausgeblendet": lambda wb: setattr(_cde(wb).column_dimensions["D"], "hidden", True),
    "zellen_verbunden": _verbinden,
    "datenpruefung_dazu": _pruefung,
    "schrift_fett_rot": _fett_rot,
}


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-136 MESSLUECKE: umbenanntes oder ausgeblendetes Blatt, ausgeblendete Zeile/Spalte, verbundene Zellen, "
                          "Datenvalidierung und Schriftformat (R1: RT-18..26, RT-35/36) sind fuer die Messung unsichtbar, "
                          "weil nur Zellwerte verglichen werden; Besitzer Mess-Bauer (messen.py, _lese_excel)")
@pytest.mark.parametrize("name", sorted(SABOTAGE_STRUKTUR))
def test_rt136_messung_bemerkt_strukturschaden(tmp_path, name):
    res = messe_mit_sabotage(_basis_sauber(tmp_path), SABOTAGE_STRUKTUR[name])
    if not _bemerkt(res):
        scheitern(f"Strukturschaden '{name}' bleibt unbemerkt (stille Fehler {res['stille_fehler']}, Kollateral "
                  f"{res['kollateral']}, Hinweise {res['hinweise']})")


# ==================================================================================================
# Teil B3: Regel 27 -- ein Tool, das "pruefen"-Aenderungen OHNE Freigabe einarbeitet, wird von der Messung BELOHNT
# ==================================================================================================
_EL_OFFEN = [("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme."),
             ("h", "2 Zwei"), ("p", "Die Reaktionszeit darf 150 ms nicht ueberschreiten."),
             ("h", "3 Drei"),
             ("p", "Die Wartung erfolgt {-jaehrlich-} {+monatlich+} durch den Betreiber und {-den-} {+das+} Personal."),
             ("h", "4 Vier"), ("p", "Alle Steckverbindungen sind eindeutig beschriftet.")]


def _alle_pruefen_ohne_freigabe_anwenden():
    echt = messen.uebernehmen

    def bruch(aend, vorher, out, freigaben=None):
        if Path(out).name == "ergebnis_auto.xlsx":                  # der Automatik-Lauf bekommt heimlich alle Freigaben
            freigaben = {a.id for a in aend if a.status == PRUEFEN}
        return echt(aend, vorher, out, freigaben)

    return mock.patch.object(messen, "uebernehmen", bruch)


def _plane_mit_offener_aenderung():
    echt = messen.plane

    def offen(pdf, xl, *a, **k):
        p = echt(pdf, xl, *a, **k)
        p.aenderungen[-1].status, p.aenderungen[-1].konfidenz = PRUEFEN, "mittel"
        return p

    return mock.patch.object(messen, "plane", offen)


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-137 MESSLUECKE (Regel 27): wendet ein Tool eine 'pruefen'-Aenderung ohne Freigabe an, sieht die "
                          "Messung besser aus (stille Fehler sinken, Treffergenauigkeit und Abdeckung steigen) und meldet "
                          "den Verstoss nicht; Besitzer Mess-Bauer (messen.py: Ist-Zelle ohne uebernommene Aenderung als Fehler)")
def test_rt137_pruefen_ohne_freigabe_angewendet_wird_gemeldet(tmp_path):
    ordner = tmp_path / "paar"
    ordner.mkdir()
    baue_pdf(ordner / "aenderungen.pdf", _EL_OFFEN, "farbe")
    baue_excel(ordner / "vorher.xlsx", zustaende(_EL_OFFEN, False), "L1", True)
    baue_excel(ordner / "nachher.xlsx", zustaende(_EL_OFFEN, True), "L1", True)
    with _plane_mit_offener_aenderung():
        ehrlich = messen.messe_paar(ordner, False)
        with _alle_pruefen_ohne_freigabe_anwenden():
            bruch = messen.messe_paar(ordner, False)
    aufbau(ehrlich["aenderungen_nach_status"].get(PRUEFEN) == 1, "Erwartet: genau eine offene Aenderung")
    aufbau(bruch["auto_richtig"] == bruch["auto_gesamt"] == 1 and bruch["erreicht_auto"] == bruch["soll_zellen"] == 1,
           "Durch den Verstoss muss das Ergebnis dem Soll entsprechen (sonst gilt der Fall nicht)")
    if bruch["stille_fehler"] == 0 and bruch["erreicht_auto"] >= ehrlich["erreicht_auto"]:
        scheitern(f"Der Verstoss gegen Regel 27 bleibt unbemerkt und sieht BESSER aus: stille Fehler {ehrlich['stille_fehler']} "
                  f"-> {bruch['stille_fehler']}, Abdeckung ohne Nachfrage {ehrlich['erreicht_auto']} -> {bruch['erreicht_auto']} "
                  f"von {bruch['soll_zellen']} Soll-Zellen")


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-138 MESSLUECKE (Regel 27), Korpus Seeds 11 und 58: Gleicher Verstoss wie RT-137 an echten Korpus-Paaren; "
                          "die Messung meldet 0 stille Fehler, obwohl eine 'pruefen'-Aenderung ohne Freigabe angewendet wurde; "
                          "Besitzer Mess-Bauer")
@pytest.mark.parametrize("seed", [11, 58])
def test_rt138_korpus_pruefen_ohne_freigabe_angewendet(seed):
    ordner = paar(seed, "standard")
    offen = [a for a in auto_lauf(seed, "standard").plan.aenderungen if a.status == PRUEFEN]
    aufbau(bool(offen), "Das Paar hat keine offene Aenderung")
    with _alle_pruefen_ohne_freigabe_anwenden():
        bruch = messen.messe_paar(ordner, False)
    if bruch["stille_fehler"] == 0:
        scheitern(f"Seed {seed}: {len(offen)} 'pruefen'-Aenderung(en) ohne Freigabe angewendet -- die Messung zeigt 0 stille "
                  f"Fehler, Treffergenauigkeit {bruch['auto_richtig']} von {bruch['auto_gesamt']}")


# ==================================================================================================
# Teil B4: Ein ABSTURZ im Wiederholungslauf (Idempotenz) macht die Funde des ersten Laufs ungueltig, und Kriterium 2 bleibt "ok"
# ==================================================================================================
def _starte_messung(eltern: Path, name: str) -> dict:
    """Wie `python -m src.cli messen --paar NAME --ordner ELTERN`; liest den Bericht (landet im Temp-Datenordner)."""
    code = messen.main(["--korpus", "entwicklung", "--ordner", str(eltern), "--paar", name])
    aufbau(code == 0, f"messen.main lieferte Code {code}")
    neu = sorted(config.pfad("messung_dir").iterdir())[-1]
    return json.loads((neu / "bericht.json").read_text(encoding="utf-8"))


def _ampel(bericht: dict, nr: int, name_teil: str) -> str:
    return next(k["ampel"] for k in bericht["kriterien"] if k["nr"] == nr and name_teil in k["name"])


def _stuerzt_im_wiederholungslauf(sabotage=None):
    echt = messen.uebernehmen

    def mit(aend, vorher, out, freigaben=None):
        if Path(out).name == "ergebnis_auto2.xlsx":
            raise ValueError("Absturz im Wiederholungslauf")
        r = echt(aend, vorher, out, freigaben)
        if sabotage and Path(out).name == "ergebnis_auto.xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(out))
            sabotage(wb)
            wb.save(str(out))
        return r

    return mock.patch.object(messen, "uebernehmen", mit)


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-139 MESSLUECKE: Stuerzt der Wiederholungslauf (Idempotenz-Pruefung) eines Paares ab, wird das ganze "
                          "Paar 'nicht messbar' mit 0 stillen Fehlern -- auch ein stiller Fehler des ERSTEN Laufs, der schon "
                          "gemessen werden konnte, geht verloren; Besitzer Mess-Bauer (messen.py, messe_paar/_fehlgeschlagen)")
def test_rt139_absturz_im_wiederholungslauf_versteckt_stillen_fehler(tmp_path):
    ordner = baue_basis_paar(tmp_path / "eltern" / "synth_9001_basis")
    einzeln = messe_mit_sabotage(ordner, SABOTAGE_ERKANNT["falsches_wort"])
    aufbau(einzeln["stille_fehler"] >= 1, "Ohne Absturz muss die Sabotage ein stiller Fehler sein")
    with _stuerzt_im_wiederholungslauf(SABOTAGE_ERKANNT["falsches_wort"]):
        bericht = _starte_messung(tmp_path / "eltern", "synth_9001_basis")
    g = bericht["gesamt"]
    aufbau(g["paare_abgestuerzt"] == 1, "Das Paar muss als abgestuerzt gelten")
    if g["stille_fehler"] == 0:
        scheitern(f"Der stille Fehler des ersten Laufs ({einzeln['stille_fehler']}) ist weg: Bericht zeigt {g['stille_fehler']} "
                  f"stille Fehler, Paar 'nicht messbar'")


@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-140 MESSLUECKE: Bei einem abgestuerzten Paar steht Kriterium 2 'Stille Fehler' auf 'ok' (Wert 0), "
                          "ebenso 'Kollateralschaden'; nur Kriterium 1 wird rot; Besitzer Mess-Bauer (messen.py, fasse_zusammen)")
def test_rt140_abgestuerztes_paar_laesst_kriterium_stille_fehler_gruen(tmp_path):
    ordner = baue_basis_paar(tmp_path / "eltern" / "synth_9002_basis")
    with _stuerzt_im_wiederholungslauf():
        bericht = _starte_messung(tmp_path / "eltern", "synth_9002_basis")
    aufbau(bericht["gesamt"]["paare_abgestuerzt"] == 1, "Das Paar muss als abgestuerzt gelten")
    ampel2 = _ampel(bericht, 2, "Stille Fehler")
    if ampel2 == "ok":
        scheitern(f"Kriterium 2 'Stille Fehler' steht auf '{ampel2}' (Wert {bericht['gesamt']['stille_fehler']}), obwohl "
                  f"{bericht['gesamt']['paare_abgestuerzt']} von {bericht['gesamt']['paare']} Paaren nicht gemessen werden konnten")


# ==================================================================================================
# Teil B5: Ohne erwartet.json (= bei ALLEN echten Paaren) kann die Messung eine fehlende Luecken-Meldung nie bemerken
# ==================================================================================================
@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-141 MESSLUECKE: Die Pruefung 'erwartete Luecke nicht gemeldet' liest nur erwartet.json (nur bei "
                          "erzeugten Paaren). Bei echten Paaren (Nicos Beispiele) bleibt eine verschluckte Scan-Seite "
                          "unbemerkt, Kriterium 1 'Bilanz stimmig' steht auf 'ok'; Besitzer Mess-Bauer "
                          "(messen.py: Luecken auch aus dem PDF selbst ableiten, z. B. Seiten ohne Textebene)")
def test_rt141_scan_seite_ohne_erwartet_json_bleibt_unbemerkt(tmp_path):
    ordner = tmp_path / "paar"
    ordner.mkdir()
    d = PdfDokument(variante="farbe")
    for art, markup in BASIS:
        d.ueberschrift(markup) if art == "h" else d.block(markup)
    d.scan_seite(scan_bild([("h", "6 Scan"), ("p", "Auf der Scan-Seite steht {+neuer Text+} und {-alter Text-} nur als Pixel.")]))
    d.speichern(ordner / "aenderungen.pdf")
    baue_excel(ordner / "vorher.xlsx", zustaende(BASIS, False), "L1", True)
    baue_excel(ordner / "nachher.xlsx", zustaende(BASIS, True), "L1", True)
    echt = messen.analysiere_pdf

    def ohne_luecken(pfad):                      # so verhaelt sich der PDF-Leser heute (R2: RT-98); unabhaengig von seiner Reparatur
        a = echt(pfad)
        a.luecken = []
        return a

    with mock.patch.object(messen, "analysiere_pdf", ohne_luecken):
        res = messen.messe_paar(ordner, False)
    aufbau(not (ordner / "erwartet.json").exists() and res["luecken_gemeldet"] == 0, "Testaufbau: keine Luecke erwartet/gemeldet")
    if res["stille_fehler"] == 0 and res["luecken_nicht_gemeldet"] == 0:
        scheitern("Eine Scan-Seite mit Aenderungen wurde nicht als Luecke gemeldet, die Messung zeigt 0 stille Fehler "
                  f"(Bilanz stimmt: {res['bilanz_stimmt']}, vollstaendig: {res['vollstaendig']})")


# ==================================================================================================
# Teil B6: "Nicht verwaltete" Spalten: Soll-Zellen dort zaehlen weder bei der Abdeckung noch als Fehler
# ==================================================================================================
@pytest.mark.redteam
@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="RT-142 MESSLUECKE: Soll-Zellen in nicht verwalteten Spalten (Korpus: 305 von 1004 = 30 %) fehlen in "
                          "Abdeckung und Fehlerzaehlung; hier 100 % Abdeckung und lauter ok, obwohl 6 von 7 geaenderten Zellen "
                          "im Ergebnis fehlen; die Zahl steht nur im md-Bericht, in keinem Kriterium; Besitzer Mess-Bauer")
def test_rt142_unverwaltete_soll_zellen_tauchen_in_keinem_kriterium_auf(tmp_path):
    zusatz = " Dieser lange Zusatz steht nur in der ersten Zeile und macht die Inhaltsspalte zur Spalte mit den meisten Aenderungen."
    text = "Der Inhalt dieser Section steht hier und ist ein ganz normaler Satz ueber die Anlage Nummer {}."
    vorher = [KOPF] + [[str(i), f"Titel {i}", text.format(i), "A"] for i in range(1, 7)]
    nachher = [KOPF] + [[str(i), f"Titel {i}", text.format(i) + (zusatz if i == 1 else ""), "B"] for i in range(1, 7)]
    ergebnis = [KOPF] + [[str(i), f"Titel {i}", text.format(i) + (zusatz if i == 1 else ""), "A"] for i in range(1, 7)]
    res = messe_mit_fertigem_ergebnis(tmp_path, vorher, nachher, ergebnis)
    aufbau(res["soll_zellen_unverwaltet"] == 6 and res["soll_zellen"] == 1 and res["erreicht_auto"] == 1,
           f"Testaufbau: erwartet 1 verwaltete (erreicht) und 6 unverwaltete Soll-Zellen, gemessen {res['soll_zellen']} / "
           f"{res['soll_zellen_unverwaltet']}")
    bericht = messen.fasse_zusammen([dict(res, name="synth_9003")], "entwicklung", "1", False)
    # Kriterium 5 (Lese-Luecke) bleibt aussen vor: das PDF dieses Testpaars ist leer, die Lese-Luecke daher absichtlich 100 %
    alle_ok = all(k["ampel"] in ("ok", "nicht gemessen") for k in bericht["kriterien"] if k["nr"] != 5)
    erwaehnt = any("verwaltet" in k["name"].lower() for k in bericht["kriterien"])
    if alle_ok and not erwaehnt:
        scheitern(f"Abdeckung {bericht['gesamt']['abdeckung_ohne_nachfrage']} und alle Kriterien 'ok', obwohl "
                  f"{res['soll_zellen_unverwaltet']} geaenderte Soll-Zellen (nicht verwaltete Spalte) im Ergebnis fehlen; "
                  f"kein Kriterium nennt sie")
