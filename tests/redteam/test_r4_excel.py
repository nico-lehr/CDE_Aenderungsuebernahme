"""Red-Team R4, Teil ZEILEN/APPLY: neue Zeilen in Excel-Dateien mit Formeln und Zahlen, Gegenprobe A3
(erfundene Daten, Regel 16). Pruefregel: "richtig ODER sichtbar unsicher, nie still falsch".
"""
import pytest
from openpyxl import load_workbook

from src.applier import uebernehmen
from src.models import UEBERNOMMEN, WIRKUNG_OK, Absatz, Aenderung, Lauf, Operation
from src.gegenprobe import zustand
from tests.redteam._bau_r4 import aufbau, ende_zu_ende, excel_sections, scheitern, unsicher, zelle

A = "Absatz mit genug Text darin für den Test."


# ---------------------------------------------------------------------------------------------
# RT-211: =LEN(Cn) in jeder Zeile und =SUM(D2:D5) darunter; neue Section in der Mitte (Schutz-Test, besteht schon)
# ---------------------------------------------------------------------------------------------
S12 = [("1", "Eins", ["Erster " + A]), ("2", "Zwei", ["Zweiter " + A]), ("3", "Drei", ["Dritter " + A]),
       ("4", "Vier", ["Vierter " + A])]
EL12 = [("h", "1 Eins"), ("p", "Erster " + A), ("h", "2 Zwei"), ("p", "Zweiter " + A),
        ("h", "{+2.1 Neu+}"), ("p", "{+Neuer Absatz der neuen Section mit genug Text.+}"),
        ("h", "3 Drei"), ("p", "Dritter " + A), ("h", "4 Vier"), ("p", "Vierter " + A)]


def _formeln(wb):
    ws = wb["CDE"]
    ws["D1"] = "Länge"
    for r in range(2, 6):
        ws.cell(r, 4).value = f"=LEN(C{r})"
    ws.cell(7, 3).value = "Summe"
    ws.cell(7, 4).value = "=SUM(D2:D5)"


@pytest.mark.redteam
def test_rt211_formeln_und_summe_bei_neuer_section(tmp_path):
    alt = excel_sections(tmp_path, S12, "L1", _formeln)
    l = ende_zu_ende(tmp_path, EL12, alt)
    betroffen = [a for a in l.aen if a.typ in ("section_neu", "absatz_neu")]
    aufbau(bool(betroffen), f"neue Section fehlt im Plan: {l.kurz()}")
    ws = load_workbook(str(l.out))["CDE"]
    if not any(a.status == UEBERNOMMEN for a in betroffen):
        if not all(unsicher(a) for a in betroffen):
            scheitern(f"RT-211: nicht uebernommen ohne Grund: {l.kurz()}")
        if ws.max_row != 7 or ws["D7"].value != "=SUM(D2:D5)":
            scheitern("RT-211: Excel veraendert, obwohl die neue Section nicht eingetragen ist")
        return
    neu = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == "Neu"), None)
    aufbau(neu is not None, "neue Section-Zeile nicht gefunden")
    fehler = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value and ws.cell(r, 2).value != "Titel" and ws.cell(r, 4).value != f"=LEN(C{r})":
            fehler.append(f"D{r}={ws.cell(r, 4).value!r}")
    summe = next((ws.cell(r, 4).value for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == "Summe"), None)
    letzte = max(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value)
    if summe != f"=SUM(D2:D{letzte})":
        fehler.append(f"Summe={summe!r} (soll =SUM(D2:D{letzte}))")
    if fehler:
        scheitern(f"RT-211: neue Section automatisch eingetragen, aber Formeln falsch: {fehler} ({l.kurz()})")


# ---------------------------------------------------------------------------------------------
# RT-212: Section-Nummern als Zahl (Format "0.0"); neue Section 2.3 muss eine Zahl mit gleichem Format werden
# ---------------------------------------------------------------------------------------------
S13 = [("2", "Zwei", ["Zweiter " + A]), ("2.1", "Eins", ["Erster Unter" + A.lower()]),
       ("2.2", "Zwei", ["Zweiter Unter" + A.lower()]), ("3", "Drei", ["Dritter " + A])]
EL13 = [("h", "2 Zwei"), ("p", "Zweiter " + A), ("h", "2.1 Eins"), ("p", "Erster Unter" + A.lower()),
        ("h", "2.2 Zwei"), ("p", "Zweiter Unter" + A.lower()),
        ("h", "{+2.3 Drei+}"), ("p", "{+Dritter Unterabsatz der neuen Section mit Text.+}"),
        ("h", "3 Drei"), ("p", "Dritter " + A)]


def _zahlen(wb):
    ws = wb["CDE"]
    for r in range(2, 6):
        v = str(ws.cell(r, 1).value)
        ws.cell(r, 1).value = float(v) if "." in v else int(v)
        ws.cell(r, 1).number_format = "0.0" if "." in v else "0"


@pytest.mark.redteam
def test_rt212_neue_section_nummer_bleibt_zahl(tmp_path):
    alt = excel_sections(tmp_path, S13, "L1", _zahlen)
    aufbau(zelle(alt, 4, 1) == 2.2, "Section 2.2 muss als Zahl in A4 stehen")
    vorschau = ende_zu_ende(tmp_path / "vorschau", EL13, alt)
    freigaben = {a.id for a in vorschau.aen}            # ein Mensch gibt alles frei: auch dann muss das Format stimmen
    l = ende_zu_ende(tmp_path, EL13, alt, freigaben)
    kopf = [a for a in l.aen if a.typ == "section_neu"]
    aufbau(bool(kopf), f"neue Section fehlt im Plan: {l.kurz()}")
    if all(a.status != UEBERNOMMEN for a in kopf):
        if not all(unsicher(a) for a in kopf):
            scheitern(f"RT-212: nicht eingetragen ohne Grund: {l.kurz()}")
        return
    ws = load_workbook(str(l.out))["CDE"]
    zeile = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == "Drei" and ws.cell(r, 1).value
                  not in (3, "3")), None)
    aufbau(zeile is not None, "Zeile der neuen Section 2.3 nicht gefunden")
    c = ws.cell(zeile, 1)
    if not (isinstance(c.value, (int, float)) and abs(c.value - 2.3) < 1e-9 and c.number_format == "0.0"):
        scheitern(f"RT-212: neue Section-Nummer ist {c.value!r} ({type(c.value).__name__}, Format {c.number_format!r}), "
                  f"soll die Zahl 2.3 mit Format '0.0' sein ({l.kurz()})")


# ---------------------------------------------------------------------------------------------
# RT-216: Gegenprobe A3 -- von Hand gebauter, absichtlich falscher Plan
# ---------------------------------------------------------------------------------------------
TEXT = "Die Steuerung meldet alle Werte an die alte Leitstelle."


def _laeufe() -> list[Lauf]:
    out: list[Lauf] = []
    for i, w in enumerate(TEXT.split(" ")):
        if i:
            out.append(Lauf("s", " "))
        if w == "alte":
            out += [Lauf("d", "alte", "A0001"), Lauf("s", " "), Lauf("a", "neue", "A0001")]
        else:
            out.append(Lauf("k", w))
    return out


@pytest.mark.redteam
def test_rt216_falscher_plan_besteht_die_gegenprobe_nicht(tmp_path):
    alt = excel_sections(tmp_path, [("1", "Eins", [TEXT]), ("2", "Zwei", ["Zweiter " + A])])
    laeufe = _laeufe()
    aufbau(zustand(laeufe, set()) == TEXT, "Absatz (vorher) muss genau dem Excel-Text entsprechen")
    aufbau(zustand(laeufe, {"A0001"}) == TEXT.replace("alte", "neue"), "Absatz (nachher) falsch gebaut")
    absatz = Absatz("P0001", "CDE", [(2, 3, 0, len(TEXT))], laeufe)
    s = TEXT.index("Werte")
    op = Operation(art="ersetzen", blatt="CDE", zeile=2, spalte=3, start=s, ende=s + len("Werte"), text="neue",
                   erwartet="Werte", aenderung="A0001")          # absichtlich das FALSCHE Wort
    ae = Aenderung(id="A0001", typ="text_ersetzt", seite=1, bbox=(56.0, 90.0, 300.0, 101.0), section="1", alt="alte",
                   neu="neue", woerter=2, status=UEBERNOMMEN, konfidenz="hoch", ziel="CDE!C2", ops=[op])
    uebernehmen([ae], alt, tmp_path / "out.xlsx", set(), absaetze=[absatz])
    ist = zelle(tmp_path / "out.xlsx", 2, 3)
    aufbau(ist != TEXT.replace("alte", "neue"), "der falsche Plan muss ein falsches Ergebnis erzeugen")
    if ae.status == UEBERNOMMEN and ae.wirkung == WIRKUNG_OK:
        scheitern(f"RT-216: falscher Plan bleibt 'uebernommen' mit Wirkung 'ok' (Nachkontrolle {ae.nachkontrolle!r}); "
                  f"Zelle {ist!r}, PDF verlangt {TEXT.replace('alte', 'neue')!r}")
