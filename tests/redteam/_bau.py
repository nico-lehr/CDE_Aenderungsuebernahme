"""Hilfsfunktionen fuer die Red-Team-Tests (nur erfundene Daten, Regel 16).

Alles lokal in tests/redteam/; es wird nichts in src/ oder tests/fixtures/ veraendert.
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from src.applier import UebernahmeErgebnis, uebernehmen
from src.excel_model import ExcelModell
from src.models import Aenderung
from src.pdf_reader import analysiere_pdf
from src.planner import Plan, plane
from src.protokoll import bilanz
from tests.fixtures.pdf_writer import baue_pdf
from tests.fixtures.szenario import baue_excel, zustaende


def scheitern(text: str) -> None:
    """Meldet einen Fehlerfall als AssertionError (xfail(raises=AssertionError) faengt nur diesen)."""
    raise AssertionError(text)


def sha(pfad: Path) -> str:
    return hashlib.sha256(Path(pfad).read_bytes()).hexdigest()


@dataclass
class Lauf:
    plan: Plan
    xl: ExcelModell
    alt: Path
    out: Path | None = None
    res: UebernahmeErgebnis | None = None
    bil: dict[str, Any] = field(default_factory=dict)

    @property
    def aen(self) -> list[Aenderung]:
        return self.plan.aenderungen

    def finde(self, typ: str | None = None, enthaelt: str = "") -> list[Aenderung]:
        return [a for a in self.aen if (typ is None or a.typ == typ) and enthaelt in (a.neu + " " + a.alt)]


def baue_alt(ordner: Path, elemente, layout: str = "L1", nachher: Callable[[Path], None] | None = None,
             deckblatt: bool = True, name: str = "alt.xlsx") -> Path:
    """Alte Excel (Stand vor den Aenderungen) aus dem Szenario-Markup bauen, optional nachbearbeiten."""
    ordner.mkdir(parents=True, exist_ok=True)
    p = ordner / name
    baue_excel(p, zustaende(elemente, False), layout, deckblatt)
    if nachher:
        nachher(p)
    return p


def bearbeite(pfad: Path, fn: Callable[[Any], None]) -> None:
    """Excel mit openpyxl oeffnen, fn(workbook) aufrufen, speichern."""
    wb = load_workbook(str(pfad))
    fn(wb)
    wb.save(str(pfad))


def schreibe_excel(pfad: Path, zeilen: list[list[Any]], blatt: str = "CDE") -> Path:
    """Einfache Excel aus einer Liste von Zeilen (erste Zeile = Kopf)."""
    wb = Workbook()
    ws = wb.active
    ws.title = blatt
    for z in zeilen:
        ws.append(z)
    wb.save(str(pfad))
    return pfad


def analysiere(ordner: Path, elemente, alt: Path, stil: str = "farbe") -> Lauf:
    ordner.mkdir(parents=True, exist_ok=True)
    pdf = ordner / "a.pdf"
    baue_pdf(pdf, elemente, stil)
    xl = ExcelModell(alt)
    plan = plane(analysiere_pdf(pdf), xl)
    return Lauf(plan=plan, xl=xl, alt=alt)


def uebernimm(l: Lauf, freigaben: set[str] | None = None, name: str = "out.xlsx") -> Lauf:
    l.out = l.alt.parent / name
    l.res = uebernehmen(l.plan.aenderungen, l.alt, l.out, freigaben)
    l.bil = bilanz(l.plan.aenderungen, l.plan.markierte_woerter, l.plan.ignoriert)
    return l


def lauf(ordner: Path, elemente, layout: str = "L1", nachher=None, freigaben: set[str] | None = None,
         alt: Path | None = None, stil: str = "farbe", deckblatt: bool = True) -> Lauf:
    alt = alt or baue_alt(ordner, elemente, layout, nachher, deckblatt)
    return uebernimm(analysiere(ordner, elemente, alt, stil), freigaben)


def blatt(pfad: Path, name: str = "CDE"):
    return load_workbook(str(pfad), rich_text=True)[name]


def zelle(pfad: Path, zeile: int, spalte: int, name: str = "CDE") -> str:
    v = blatt(pfad, name).cell(zeile, spalte).value
    return "" if v is None else str(v)


def alle_texte(pfad: Path, name: str = "CDE") -> list[str]:
    ws = blatt(pfad, name)
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


def aufbau(bedingung: bool, text: str) -> None:
    """Pruefung des TESTAUFBAUS: Fehler hier ist kein erwarteter Tool-Fehler (RuntimeError, nicht AssertionError)."""
    if not bedingung:
        raise RuntimeError("Testaufbau fehlerhaft: " + text)


def _xref_neu(daten: bytes) -> bytes:
    """Baut die xref-Tabelle eines (von pdf_writer erzeugten) PDFs neu auf, nachdem Bytes ersetzt wurden."""
    import re
    pos = daten.rindex(b"xref\n0 ")
    kopf = daten[:pos]
    offs = {int(m.group(1)): m.start() for m in re.finditer(rb"(?m)^(\d+) 0 obj", kopf)}
    n = max(offs) + 1
    xr = b"xref\n0 %d\n0000000000 65535 f \n" % n
    for i in range(1, n):
        xr += b"%010d 00000 n \n" % offs[i]
    root = re.search(rb"/Root (\d+) 0 R", daten[pos:]).group(1)
    return kopf + xr + b"trailer\n<< /Size %d /Root %s 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (n, root, len(kopf))


def baue_pdf_ligatur(pfad: Path, elemente, stil: str = "farbe") -> None:
    """Wie baue_pdf, aber das Zeichen '€' im Markup erscheint beim Auslesen als Ligatur 'ﬁ'
    (Schrift-Kodierung 128 -> Glyph 'fi'; so liefern echte PDFs oft Ligaturen)."""
    baue_pdf(pfad, elemente, stil)
    daten = Path(pfad).read_bytes()
    alt = b"/BaseFont /Helvetica /Encoding /WinAnsiEncoding"
    aufbau(alt in daten, "Schrift-Objekt im Test-PDF nicht gefunden")
    daten = daten.replace(alt, b"/BaseFont /Helvetica /Encoding << /Type /Encoding /BaseEncoding /WinAnsiEncoding "
                               b"/Differences [128 /fi] >>")
    Path(pfad).write_bytes(_xref_neu(daten))


def analysiere_mit_pdf(pdf: Path, alt: Path) -> Lauf:
    """Wie analysiere(), aber mit einem schon vorhandenen PDF."""
    xl = ExcelModell(alt)
    return Lauf(plan=plane(analysiere_pdf(pdf), xl), xl=xl, alt=alt)


def zeile_mit(pfad: Path, text: str, name: str = "CDE") -> int | None:
    """Erste Zeile, in der irgendeine Zelle `text` enthaelt."""
    ws = blatt(pfad, name)
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and text in str(c.value):
                return c.row
    return None
