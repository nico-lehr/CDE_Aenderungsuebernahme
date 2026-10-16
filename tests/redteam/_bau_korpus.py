"""Hilfen fuer die Red-Team-Tests R3 (Analyse der stillen Fehler im Entwicklungskorpus).

Nur erfundene Daten (Regel 16). Es wird nichts in src/, config/ oder tests/fixtures/ veraendert.
Der Testkorpus (Seeds ab 1000) wird NIE erzeugt oder gemessen (Regel 11): `paar()` bricht bei Seeds >= 1000 ab.

Paare werden mit dem festen Generator (Version 1) aus dem Seed neu gebaut und in einem Temp-Ordner abgelegt,
der beim Beenden des Prozesses wieder verschwindet. Jedes Paar wird hoechstens einmal erzeugt und einmal gemessen.
"""
from __future__ import annotations

import atexit
import copy
import json
import shutil
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any, Callable
from unittest import mock

from openpyxl import Workbook, load_workbook

from src import messen
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.messen import _diff, _k, _lese_excel, _wort_diff, _wort_liste, _zeilen
from src.models import Aenderung, PRUEFEN, UEBERNOMMEN
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests.fixtures import generator as g
from tests.fixtures.pdf_writer import PdfDokument, baue_pdf
from tests.fixtures.szenario import baue_excel, zustaende
from tests.redteam._bau import aufbau

# ---------------------------------------------------------------------------------------------
# Temp-Wurzel und Korpus-Paare (Seeds < 1000)
# ---------------------------------------------------------------------------------------------
_WURZEL: Path | None = None
_PAARE: dict[tuple[int, str], Path] = {}
_MESS: dict[tuple[int, str], dict[str, Any]] = {}
_AUTO: dict[tuple[int, str, bool], "AutoLauf"] = {}


def wurzel() -> Path:
    global _WURZEL
    if _WURZEL is None:
        _WURZEL = Path(tempfile.mkdtemp(prefix="cde_r3_"))
        atexit.register(shutil.rmtree, _WURZEL, ignore_errors=True)
    return _WURZEL


def paar(seed: int, profil: str) -> Path:
    """Erzeugt (einmal) das Entwicklungs-Paar zu Seed/Profil und gibt den Ordner zurueck."""
    # Die R3-Faelle sind an Generator-Version 1 gebunden; v2 kann v1-Paare weiter nachbauen (version="1").
    aufbau("1" in getattr(g, "VERSIONEN", (g.GENERATOR_VERSION,)), "Generator kann Version 1 nicht mehr nachbauen")
    aufbau(seed < g.SEED_GRENZE, "Nur Entwicklungskorpus (Seeds < 1000, Regel 11)")
    if (seed, profil) not in _PAARE:
        _PAARE[(seed, profil)] = g.erzeuge_paar(seed, profil, wurzel() / "paare", "entwicklung", version="1")
    return _PAARE[(seed, profil)]


def erwartet(seed: int, profil: str) -> dict[str, Any]:
    return json.loads((paar(seed, profil) / "erwartet.json").read_text(encoding="utf-8"))


def messung(seed: int, profil: str) -> dict[str, Any]:
    """Messwerkzeug-Ergebnis fuer ein Korpus-Paar (nur Automatik, ohne Vorschlaege)."""
    if (seed, profil) not in _MESS:
        _MESS[(seed, profil)] = messen.messe_paar(paar(seed, profil), False)
    return _MESS[(seed, profil)]


def stumme(res: dict[str, Any]) -> list[dict[str, Any]]:
    return [f for f in res["fehler"] if f.get("stumm")]


def stumm_bei(res: dict[str, Any], blatt: int, zeile: int, spalte: int) -> list[dict[str, Any]]:
    return [f for f in stumme(res) if (f["blatt"], f["zeile"], f["spalte"]) == (blatt, zeile, spalte)]


# ---------------------------------------------------------------------------------------------
# Der Automatik-Lauf eines Korpus-Paares (wie im Messwerkzeug) mit Zugriff auf Plan und Ergebnis
# ---------------------------------------------------------------------------------------------
class AutoLauf:
    def __init__(self, ordner: Path, ergebnis: Path, plan, aend: list[Aenderung], namen: list[str]) -> None:
        self.ordner, self.ergebnis, self.plan, self.aend, self.namen = ordner, ergebnis, plan, aend, namen

    def aenderungen_bei(self, blatt: int, zeile: int, spalte: int) -> list[Aenderung]:
        """Aenderungen mit einer Operation genau auf dieser Zelle (blatt = 0-basiert)."""
        name = self.namen[blatt]
        return [a for a in self.aend if any(o.blatt == name and o.zeile == zeile and o.spalte == spalte for o in a.ops)]


def auto_lauf(seed: int, profil: str, alle_pruefen_frei: bool = False) -> AutoLauf:
    key = (seed, profil, alle_pruefen_frei)
    if key not in _AUTO:
        ordner = paar(seed, profil)
        plan = plane(analysiere_pdf(ordner / "aenderungen.pdf"), ExcelModell(ordner / "vorher.xlsx"))
        aend = copy.deepcopy(plan.aenderungen)
        ziel = wurzel() / "ergebnisse"
        ziel.mkdir(exist_ok=True)
        out = ziel / f"{ordner.name}_{'alle' if alle_pruefen_frei else 'auto'}.xlsx"
        frei = {a.id for a in aend if a.status == PRUEFEN} if alle_pruefen_frei else None
        uebernehmen(aend, ordner / "vorher.xlsx", out, frei)
        namen = load_workbook(str(ordner / "vorher.xlsx"), read_only=True).sheetnames
        _AUTO[key] = AutoLauf(ordner, out, plan, aend, list(namen))
    return _AUTO[key]


def zelle_vorher_soll_ist(ordner: Path, ergebnis: Path, blatt: int, zeile: int, spalte: int) -> tuple[str, str, str]:
    """(vorher, soll, ist) einer Zelle. Soll = Zelle laut nachher, nach derselben Zeilenausrichtung wie im
    Messwerkzeug; Ist = Zelle im Ergebnis. blatt = 0-basiert."""
    vb, nb, eb = (_lese_excel(p) for p in (ordner / "vorher.xlsx", ordner / "nachher.xlsx", ergebnis))
    v, n, e = _zeilen(vb[blatt]), _zeilen(nb[blatt]), _zeilen(eb[blatt])
    zeile_v = next((z for z in v if z.nr == zeile), None)
    vt = zeile_v.zellen[spalte - 1] if zeile_v and spalte - 1 < len(zeile_v.zellen) else ""
    return vt, _diff(v, n).zellen.get((zeile, spalte), vt), _diff(v, e).zellen.get((zeile, spalte), vt)


def zeile_im_ergebnis(ergebnis: Path, blatt: int, zeile: int) -> list[str]:
    zeilen = _lese_excel(ergebnis)[blatt]
    return list(zeilen[zeile - 1]) if zeile - 1 < len(zeilen) else []


def nur_durch_offene_aenderungen_verschieden(ist: str, soll: str, offene: list[Aenderung]) -> bool:
    """True, wenn jedes Wort, in dem sich ist und soll unterscheiden, zu einer der OFFENEN Aenderungen gehoert."""
    entf, dazu = _wort_diff(ist, soll)
    erlaubt: Counter[str] = Counter()
    for a in offene:
        erlaubt.update(_wort_liste(a.alt))
        erlaubt.update(_wort_liste(a.neu))
    return not ((Counter(entf) + Counter(dazu)) - erlaubt)


# ---------------------------------------------------------------------------------------------
# Handgebaute Paare (Messwerkzeug-Tests, in denen das Tool selbst keine Rolle spielen soll)
# ---------------------------------------------------------------------------------------------
KOPF = ["Nr.", "Titel", "Inhalt", "Bemerkung"]


def schreibe_xlsx(pfad: Path, zeilen: list[list[Any]], blatt: str = "CDE") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = blatt
    for z in zeilen:
        ws.append(z)
    wb.save(str(pfad))
    return pfad


def leeres_pdf(pfad: Path) -> None:
    """PDF ganz ohne Markierung: das Tool erkennt keine Aenderung."""
    d = PdfDokument(kopf_fuss=False)
    d.ueberschrift("1 Eins")
    d.block("Ein unveraenderter Absatz ganz ohne Markierung.")
    d.speichern(pfad)


def messe_mit_fertigem_ergebnis(ordner: Path, vorher: list[list[Any]], nachher: list[list[Any]],
                                ergebnis: list[list[Any]]) -> dict[str, Any]:
    """Misst ein Paar, bei dem das Ergebnis von Hand vorgegeben ist (das Tool aendert nichts; `uebernehmen` wird
    nur im Messwerkzeug durch 'kopiere das vorbereitete Ergebnis' ersetzt). So prueft man NUR das Messwerkzeug."""
    ordner.mkdir(parents=True, exist_ok=True)
    leeres_pdf(ordner / "aenderungen.pdf")
    schreibe_xlsx(ordner / "vorher.xlsx", vorher)
    schreibe_xlsx(ordner / "nachher.xlsx", nachher)
    fertig = schreibe_xlsx(ordner / "fertig.xlsx", ergebnis)

    def kopiere(aend, vor, out, freigaben=None, **_):
        shutil.copyfile(fertig, out)

    with mock.patch.object(messen, "uebernehmen", kopiere):
        return messen.messe_paar(ordner, False)


# ---------------------------------------------------------------------------------------------
# Sauberes Basis-Paar fuer Sabotage-Tests: das Tool macht alles richtig (stille Fehler 0, Kollateral 0)
# ---------------------------------------------------------------------------------------------
BASIS = [
    ("h", "1 Eins"), ("p", "Der Nachweis erfolgt durch ein Messprotokoll der Abnahme. {+Die Pruefung findet jaehrlich statt.+}"),
    ("h", "2 Zwei"), ("p", "Die Reaktionszeit darf {-150-} {+120+} ms nicht ueberschreiten."),
    ("h", "3 Drei"), ("p", "Alle Steckverbindungen sind eindeutig beschriftet."),
    ("p", "{+Ersatzteile werden fuenf Jahre vorgehalten.+}"),
    ("h", "4 Vier"), ("p", "Die Wartung erfolgt {-jaehrlich-} {+monatlich+} durch den Betreiber."),
    ("h", "5 Fuenf"), ("p", "Die Lieferung umfasst alles Zubehoer fuer die Montage."),
]


def baue_basis_paar(ordner: Path, layout: str = "L1", elemente=BASIS) -> Path:
    ordner.mkdir(parents=True, exist_ok=True)
    baue_pdf(ordner / "aenderungen.pdf", elemente, "farbe")
    baue_excel(ordner / "vorher.xlsx", zustaende(elemente, False), layout, True)
    baue_excel(ordner / "nachher.xlsx", zustaende(elemente, True), layout, True)
    return ordner


def finde_zelle(ws, teil: str):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and teil in str(c.value):
                return c
    raise RuntimeError(f"Testaufbau fehlerhaft: Zelle mit '{teil}' nicht gefunden")


def messe_mit_sabotage(ordner: Path, sabotage: Callable[[Any], None]) -> dict[str, Any]:
    """Misst ein Paar, nachdem das ERGEBNIS des ersten Automatik-Laufs von `sabotage(workbook)` verfaelscht wurde
    (nur im Messwerkzeug; das Tool selbst bleibt unberuehrt). Der Wiederholungslauf (Idempotenz) bleibt echt."""
    echt = messen.uebernehmen

    def mit(aend, vorher, out, freigaben=None, **_):
        r = echt(aend, vorher, out, freigaben)
        if Path(out).name == "ergebnis_auto.xlsx":
            wb = load_workbook(str(out))
            sabotage(wb)
            wb.save(str(out))
        return r

    with mock.patch.object(messen, "uebernehmen", mit):
        return messen.messe_paar(ordner, False)
