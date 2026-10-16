"""Waechter: Eigenschaften, die heute gelten und weiter gelten muessen (kein xfail).

Bricht einer dieser Tests, wurde etwas kaputt repariert. Nur erfundene Daten (Regel 16).
"""
import hashlib
import json
import os
import random
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src import runs
from src.api.app import app
from src.applier import uebernehmen
from src.models import Aenderung, Operation
from tests.fixtures.gross import erzeuge_gross, satz
from tests.fixtures.pdf_writer import baue_pdf
from tests.fixtures.szenario import ELEMENTE, erzeuge
from tests.helfer import durchlauf
from tests.redteam._bau import (alle_texte, analysiere, aufbau, baue_alt, bearbeite, blatt, lauf, schreibe_excel, sha,
                                uebernimm, zelle, zeile_mit)


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


def _hochladen(client, d, name_pdf="a.pdf", name_xlsx="cde.xlsx"):
    return client.post("/api/lauf", data={"pdf": (open(d["pdf"], "rb"), name_pdf), "excel": (open(d["alt"], "rb"), name_xlsx)},
                       content_type="multipart/form-data")


# ---------------------------------------------------------------------------------------------
# Original bleibt unveraendert
# ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("layout", ["L1", "L2", "L3"])
def test_waechter_original_excel_bleibt_unveraendert(tmp_path, layout):
    d = erzeuge(tmp_path, "farbe", layout)
    vorher = sha(d["alt"])
    l = lauf(tmp_path, ELEMENTE, layout, alt=d["alt"])
    assert sha(d["alt"]) == vorher, "Original-Excel wurde veraendert"
    assert l.out.exists() and sha(l.out) != vorher


def test_waechter_original_excel_im_lauf_ordner_unveraendert(client, tmp_path):
    d = erzeuge(tmp_path / "fx")
    vorher = sha(d["alt"])
    rid = _hochladen(client, d).get_json()["run_id"]
    antwort = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []})
    assert antwort.status_code == 200
    assert sha(runs.lauf_dir(rid) / "eingabe" / "cde.xlsx") == vorher
    assert sha(d["alt"]) == vorher


# ---------------------------------------------------------------------------------------------
# Bilanz: nichts geht still verloren
# ---------------------------------------------------------------------------------------------
def _fall_bullet(p):
    return lauf(p, [("h", "1 Eins"), ("p", "Das System läuft stabil {-•-} und liefert Daten."), ("h", "2 Zwei"),
                    ("p", "Zweiter Absatz mit genug Text darin.")], "L1")


def _fall_zelle_leeren(p):
    alt = baue_alt(p, [("h", "1 Eins"), ("p", "Erster Absatz mit Text."),
                       ("p", "Alle Racks werden zusätzlich jährlich noch vor dem Versand geprüft."), ("h", "2 Zwei"),
                       ("p", "Zweiter Absatz mit genug Text darin.")], "L1")
    return uebernimm(analysiere(p, [("h", "1 Eins"), ("p", "Erster Absatz mit Text."),
                                    ("p", "{-Alle Racks werden jährlich vor dem Versand geprüft.-}"), ("h", "2 Zwei"),
                                    ("p", "Zweiter Absatz mit genug Text darin.")], alt))


def _fall_doppelt(p):
    return lauf(p, [("h", "1 Eins"), ("p", "Die Anlage wird vor dem Versand einzeln geprüft."),
                    ("p", "Die Anlage wird vor dem Versand {-einzeln-} {+gemeinsam+} geprüft."), ("h", "2 Zwei"),
                    ("p", "Zweiter Absatz mit genug Text darin.")], "L1")


def _fall_cid(p):
    return lauf(p, [("h", "1 Eins"), ("p", "Erster Absatz."), ("p", "{+Teil (cid:12) hier unten.+}"), ("h", "2 Zwei"),
                    ("p", "Zweiter Absatz mit genug Text darin.")], "L3")


def _fall_unbekannt(p):
    return lauf(p, [("h", "1 Eins"), ("p", "Dieser Absatz steht gar nicht in der Excel {+und wird ergaenzt+}."),
                    ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")], "L1",
                alt=baue_alt(p, [("h", "1 Eins"), ("p", "Voellig anderer Inhalt hier."), ("h", "2 Zwei"),
                                 ("p", "Zweiter Absatz mit genug Text darin.")], "L1"))


@pytest.mark.parametrize("fall", [_fall_bullet, _fall_zelle_leeren, _fall_doppelt, _fall_cid, _fall_unbekannt],
                         ids=lambda f: f.__name__)
def test_waechter_bilanz_stimmt_in_schwierigen_faellen(tmp_path, fall):
    l = fall(tmp_path)
    assert l.bil["stimmt"], l.bil
    zaehler = l.bil["uebernommen"] + l.bil["pruefen"] + l.bil["nicht_uebernommen"]
    assert zaehler == l.bil["aenderungen"] == len(l.aen)
    for a in l.aen:
        if a.status == "nicht_uebernommen":
            assert a.grund, f"{a.id}: nicht uebernommen ohne Grund"


def test_waechter_unbekannter_absatz_wird_nicht_automatisch_geschrieben(tmp_path):
    l = _fall_unbekannt(tmp_path)
    assert all(a.status != "uebernommen" for a in l.aen)
    assert "ergaenzt" not in " ".join(alle_texte(l.out))


# ---------------------------------------------------------------------------------------------
# Determinismus und kleine Zufallsfaelle mit festen Seeds
# ---------------------------------------------------------------------------------------------
def _zellen(pfad):
    wb = load_workbook(str(pfad))
    return {(ws.title, c.row, c.column): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.value is not None}


def test_waechter_determinismus(tmp_path):
    el = erzeuge_gross(seed=3, sections=5, absaetze=4, aenderungen=8)
    l1 = lauf(tmp_path / "a", el, "L3")
    l2 = lauf(tmp_path / "b", el, "L3")
    assert [a.zu_dict() for a in l1.aen] == [a.zu_dict() for a in l2.aen]
    assert _zellen(l1.out) == _zellen(l2.out)


@pytest.mark.parametrize("seed", [1, 2, 3, 4, 5])
def test_waechter_zufallsdokument_bilanz_und_loesung(tmp_path, seed):
    rng = random.Random(seed)
    layout = rng.choice(["L1", "L2", "L3"])
    el = erzeuge_gross(seed=seed, sections=rng.randint(4, 7), absaetze=rng.randint(3, 5), aenderungen=rng.randint(5, 9))
    plan, bil, gleich = durchlauf(tmp_path, el, layout)
    assert bil["stimmt"]
    assert gleich, f"Ergebnis weicht von der Loesung ab (Seed {seed}, Layout {layout})"
    for a in plan.aenderungen:
        assert a.status != "uebernommen" or a.nachkontrolle == "ok"


# ---------------------------------------------------------------------------------------------
# Zweiter Lauf: Ersetzungen werden nicht doppelt gemacht (Gegenstueck zu RT-03..05)
# ---------------------------------------------------------------------------------------------
def test_waechter_zweiter_lauf_ersetzen_aendert_nichts_mehr(tmp_path):
    el = [("h", "1 Eins"), ("p", "Das System {-muss-} {+soll+} die Daten speichern."), ("h", "2 Zwei"),
          ("p", "Zweiter Absatz mit genug Text darin.")]
    l1 = lauf(tmp_path / "eins", el, "L1")
    l2 = uebernimm(analysiere(tmp_path / "zwei", el, l1.out), name="out2.xlsx")
    assert _zellen(l2.out) == _zellen(l1.out)
    assert l2.aen[0].status == "nicht_uebernommen" and l2.aen[0].grund


# ---------------------------------------------------------------------------------------------
# Zeilen einfuegen: was heute richtig funktioniert
# ---------------------------------------------------------------------------------------------
def test_waechter_zeilenhoehe_wandert_mit(tmp_path):
    el = [("h", "1 Eins"), ("p", "Absatz eins mit genug Text darin."), ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin."),
          ("h", "{+2.5 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}"),
          ("h", "3 Drei"), ("p", "Absatz drei mit genug Text darin.")]
    l = lauf(tmp_path, el, "L1", nachher=lambda p: bearbeite(p, lambda wb: setattr(wb["CDE"].row_dimensions[4], "height", 55)))
    ws = blatt(l.out)
    r = zeile_mit(l.out, "Absatz drei")
    assert ws.row_dimensions[r].height == 55
    assert ws.row_dimensions[r - 1].height != 55


def test_waechter_verbundene_zelle_unter_der_einfuegestelle_wandert_mit(tmp_path):
    el = [("h", "1 Eins"), ("p", "Absatz eins mit genug Text darin."), ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin."),
          ("h", "{+2.5 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}"),
          ("h", "3 Drei"), ("p", "Absatz drei mit genug Text darin."), ("h", "4 Vier"), ("p", "Absatz vier mit genug Text darin.")]
    l = lauf(tmp_path, el, "L1", nachher=lambda p: bearbeite(p, lambda wb: wb["CDE"].merge_cells("D4:D5")))
    ws = load_workbook(str(l.out))["CDE"]
    rng = next(iter(ws.merged_cells.ranges))
    assert rng.min_row == zeile_mit(l.out, "Absatz drei") and rng.max_row == zeile_mit(l.out, "Absatz vier")


# ---------------------------------------------------------------------------------------------
# Reihenfolge bei gleichem Start (die Richtung, die heute stimmt)
# ---------------------------------------------------------------------------------------------
def test_waechter_einfuegen_vor_ersetzen_gleicher_start_richtige_reihenfolge(tmp_path):
    p = schreibe_excel(tmp_path / "a.xlsx", [["Nr.", "Titel", "Inhalt"],
                                              ["1", "Eins", "Das System muss laufen und stabil bleiben."],
                                              ["2", "Zwei", "Zweiter Absatz mit genug Text darin."]])
    ins = Operation("einfuegen", "CDE", 2, 3, 4, 4, "neue ", 1, aenderung="A0001")
    rep = Operation("ersetzen", "CDE", 2, 3, 4, 10, "Gerät", 2, aenderung="A0002", erwartet="System")
    a1 = Aenderung(id="A0001", typ="text_eingefuegt", seite=1, bbox=(0, 0, 1, 1), section="1", alt="", neu="neue",
                   woerter=1, status="uebernommen", konfidenz="hoch", ops=[ins])
    a2 = Aenderung(id="A0002", typ="text_ersetzt", seite=1, bbox=(0, 0, 1, 1), section="1", alt="System", neu="Gerät",
                   woerter=1, status="uebernommen", konfidenz="hoch", ops=[rep])
    uebernehmen([a1, a2], p, tmp_path / "out.xlsx")
    assert zelle(tmp_path / "out.xlsx", 2, 3) == "Das neue Gerät muss laufen und stabil bleiben."


# ---------------------------------------------------------------------------------------------
# Mehrere Blaetter: gleiche Section-Nummer in beiden -> keine automatische Entscheidung
# ---------------------------------------------------------------------------------------------
def test_waechter_gleicher_text_in_zwei_blaettern_gleiche_section_nur_vorschlag(tmp_path):
    satz_ = "Das System muss alle Daten speichern und sichern."
    wb = Workbook()
    for name, pre in (("Blatt A", "A"), ("Blatt B", "B")):
        ws = wb.active if name == "Blatt A" else wb.create_sheet()
        ws.title = name
        ws.append(["Nr.", "Titel", "Inhalt"])
        ws.append(["1", f"Titel {pre}1", f"Alpha {pre} text eins steht hier im Dokument."])
        ws.append(["2", f"Titel {pre}2", satz_])
        ws.append(["3", f"Titel {pre}3", f"Gamma {pre} text drei steht hier im Dokument."])
    p = tmp_path / "m.xlsx"
    wb.save(str(p))
    el = [("h", "1 Eins"), ("p", "Alpha A text eins steht hier im Dokument."), ("h", "2 Zwei"),
          ("p", "Das System {-muss-} {+soll+} alle Daten speichern und sichern."), ("h", "3 Drei"),
          ("p", "Gamma A text drei steht hier im Dokument.")]
    l = uebernimm(analysiere(tmp_path, el, p))
    ae = next(a for a in l.aen if a.typ == "text_ersetzt")
    # Welle 1 (ADR 0004, Reissverschluss): die Nachbarabsaetze stehen exakt nur in Blatt A -> die Reihenfolge entscheidet
    # eindeutig fuer Blatt A (automatisch), Blatt B bleibt unberuehrt.
    assert ae.status == "uebernommen" and ae.ziel.startswith("Blatt A!")
    from openpyxl import load_workbook
    wb2 = load_workbook(str(l.out))
    assert wb2["Blatt B"]["C3"].value == satz_ and wb2["Blatt A"]["C3"].value == satz_.replace("muss", "soll")
    # ohne eindeutige Nachbarn entscheidet die Reihenfolge nicht: nur ein Vorschlag
    el2 = [("h", "2 Zwei"), ("p", "Das System {-muss-} {+soll+} alle Daten speichern und sichern.")]
    l2 = analysiere(tmp_path / "zwei", el2, p)
    ae2 = next(a for a in l2.aen if a.typ == "text_ersetzt")
    assert ae2.status == "pruefen" and ae2.konfidenz == "mittel"


# ---------------------------------------------------------------------------------------------
# Groessenordnung: 20 000 gefuellte Zeilen
# ---------------------------------------------------------------------------------------------
def test_waechter_20000_zeilen_gesamtlauf_unter_60_sekunden(tmp_path):
    rng = random.Random(5)
    wb = Workbook()
    ws = wb.active
    ws.title = "CDE"
    ws.append(["Nr", "Titel", "Text", "Prio"])
    texte = {}
    for s in range(1, 2001):
        ws.append([str(s), f"Abschnitt {s}", None, "A"])
        for a in range(9):
            tx = f"Punkt{s}x{a} " + " ".join(satz(rng, rng.randint(6, 10)) for _ in range(2))
            texte[(s, a)] = tx
            ws.append([None, None, tx, "B"])
    aufbau(ws.max_row == 20001, "20 000 gefuellte Zeilen erwartet")
    alt = tmp_path / "gross.xlsx"
    wb.save(str(alt))
    el = []
    for s in (3, 1000, 1999):
        el.append(("h", f"{s} Abschnitt {s}"))
        w = texte[(s, 0)].split()
        w[4] = "{-" + w[4] + "-} {+NEUWORT+}"
        el.append(("p", " ".join(w)))
        el.append(("p", f"{{+Ein neuer Absatz im Abschnitt {s} ganz am Ende.+}}"))
    t0 = time.time()
    l = uebernimm(analysiere(tmp_path, el, alt))
    dauer = time.time() - t0
    assert dauer < 60, f"Gesamtlauf dauerte {dauer:.0f} s"
    # Welle 1 (ADR 0004): der neue Absatz steht im PDF direkt hinter Absatz 0, in der Excel folgen darauf aber noch 8 Absaetze,
    # die das PDF nicht zeigt -> die Einfuegestelle ist nicht sicher (Code anker_luecke, nur Vorschlag). Die Text-Aenderungen
    # bleiben automatisch.
    assert l.bil["stimmt"] and all(a.status == "uebernommen" and a.nachkontrolle == "ok" for a in l.aen if a.typ != "absatz_neu")
    assert all(a.status == "pruefen" and "anker_luecke" in a.pruefcodes for a in l.aen if a.typ == "absatz_neu")


# ---------------------------------------------------------------------------------------------
# API: append-only, Schema 1, Host/Origin, Dateinamen
# ---------------------------------------------------------------------------------------------
def test_waechter_zweite_uebernahme_legt_neuen_ordner_an_und_laesst_den_ersten_unberuehrt(client, tmp_path):
    d = erzeuge(tmp_path / "fx")
    rid = _hochladen(client, d).get_json()["run_id"]
    u1 = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []}).get_json()
    ordner1 = runs.lauf_dir(rid) / u1["ordner"]
    hashes = {p.name: sha(p) for p in ordner1.iterdir()}
    plan_vorher = sha(runs.lauf_dir(rid) / "plan.json")
    u2 = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []}).get_json()
    assert u2["ordner"] != u1["ordner"]
    assert {p.name: sha(p) for p in ordner1.iterdir()} == hashes
    assert sha(runs.lauf_dir(rid) / "plan.json") == plan_vorher


def test_waechter_schema_1_plan_json_laedt_und_uebernimmt(client, tmp_path):
    d = erzeuge(tmp_path / "fx")
    rid = _hochladen(client, d).get_json()["run_id"]
    pj = runs.lauf_dir(rid) / "plan.json"
    j = json.loads(pj.read_text(encoding="utf-8"))
    j["schema_version"] = 1                          # so sah ein Plan vor Schema 2 aus
    j.pop("luecken", None)
    for a in j["aenderungen"]:                       # aeltere Plaene hatten weniger optionale Felder
        for feld in ("kontext", "excel_text", "score", "bestaetigt", "nachkontrolle", "hinweise", "wirkung"):
            a.pop(feld, None)
    pj.write_text(json.dumps(j), encoding="utf-8")
    assert client.get(f"/api/lauf/{rid}").status_code == 200
    antwort = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []})
    assert antwort.status_code == 200, antwort.get_json()


def test_waechter_eigener_host_und_origin_werden_angenommen(client, tmp_path):
    assert client.get("/api/health", base_url="http://127.0.0.1:5100").status_code == 200
    assert client.get("/api/health", base_url="http://localhost:5100").status_code == 200
    d = erzeuge(tmp_path / "fx")
    rid = _hochladen(client, d).get_json()["run_id"]
    r = client.post(f"/api/lauf/{rid}/uebernehmen", json={"freigaben": []}, base_url="http://127.0.0.1:5100",
                    headers={"Origin": "http://127.0.0.1:5100"})
    assert r.status_code == 200


@pytest.mark.parametrize("pdf_name, xlsx_name", [
    ("../../evil.pdf", "..\\..\\evil.xlsx"), ("/etc/passwd.pdf", "C:\\Windows\\x.xlsx"),
    ("a/b/c.pdf", "x:y?z*.xlsx"), ("na\u0000me.pdf", "böse name.xlsx")])
def test_waechter_pfad_tricks_im_dateinamen_bleiben_im_lauf_ordner(fx_dir, pdf_name, xlsx_name):
    rid = runs.neuer_lauf(fx_dir["pdf"], fx_dir["alt"], pdf_name, xlsx_name)
    d = runs.lauf_dir(rid)
    for p in d.rglob("*"):
        assert d in p.parents or p == d
    gespeichert = {p.name for p in (d / "eingabe").iterdir()}
    assert len(gespeichert) == 2
    for n in gespeichert:                           # ein einzelner Dateiname, kein Pfad, liegt direkt in eingabe/
        assert "/" not in n and "\\" not in n and n not in (".", "..")
        assert (d / "eingabe" / n).resolve().parent == (d / "eingabe").resolve()
    # nichts ausserhalb des Datenordners angelegt
    wurzel = runs.runs_dir()
    assert [p.name for p in wurzel.iterdir()] == [rid]
    assert not (wurzel.parent / "evil.pdf").exists() and not (wurzel.parent.parent / "evil.pdf").exists()


@pytest.fixture
def fx_dir(tmp_path):
    return erzeuge(tmp_path / "fx")


@pytest.mark.parametrize("pfad", [
    "eingabe/..%2f..%2fplan.json", "eingabe/..%5c..%5cplan.json", "eingabe/%2e%2e%2f%2e%2e%2fplan.json",
    "ergebnis_001/../../plan.json", "../eingabe/cde.xlsx"])
def test_waechter_download_pfad_tricks(client, fx_dir, pfad):
    rid = _hochladen(client, fx_dir).get_json()["run_id"]
    antwort = client.get(f"/api/lauf/{rid}/download/{pfad}")
    assert antwort.status_code in (400, 404)
