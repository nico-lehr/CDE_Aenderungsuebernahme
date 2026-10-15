"""Red-Team K-XL (excel_model / matcher): falsche Treffer, falsch erkannter Aufbau, mehrere Blaetter, Groesse."""
import json
import os
import subprocess
import sys
import textwrap
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from src.excel_model import ExcelModell
from tests.redteam._bau import (alle_texte, analysiere, aufbau, baue_alt, blatt, lauf, scheitern, schreibe_excel,
                                uebernimm, zelle)


# ---------------------------------------------------------------------------------------------
# RT-27 / RT-28: Zuordnung innerhalb einer Zelle
# ---------------------------------------------------------------------------------------------
@pytest.mark.redteam
def test_rt27_gleicher_absatz_zweimal_in_derselben_zelle(tmp_path):
    el = [("h", "1 Eins"),
          ("p", "Die Anlage wird vor dem Versand einzeln geprüft."),
          ("p", "Die Anlage wird vor dem Versand {-einzeln-} {+gemeinsam+} geprüft."),     # der ZWEITE aendert sich
          ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]
    l = lauf(tmp_path, el, "L1")
    ae = l.aen[0]
    erste, zweite = zelle(l.out, 2, 3).split("\n")
    falsch = "gemeinsam" in erste or "gemeinsam" not in zweite
    if ae.status == "uebernommen" and falsch:
        scheitern(f"Status {ae.status}/{ae.konfidenz}: erster Absatz {erste!r}, zweiter Absatz {zweite!r}")


@pytest.mark.redteam
def test_rt28_pdf_block_nur_anfang_des_excel_absatzes(tmp_path):
    el = [("h", "1 Eins"),
          ("p", "Die Anlage wird vor dem Versand einzeln geprüft und danach"),
          ("p", "{+Neuer Absatz nach dem Anfang.+}"),
          ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]
    ex = [("h", "1 Eins"),
          ("p", "Die Anlage wird vor dem Versand einzeln geprüft und danach dokumentiert; anschließend wird ein "
                "Prüfbericht erstellt und dem Kunden übergeben."),
          ("h", "2 Zwei"), ("p", "Zweiter Absatz mit genug Text darin.")]
    alt = baue_alt(tmp_path, ex, "L1")
    l = uebernimm(analysiere(tmp_path, el, alt))
    ae = next(a for a in l.aen if a.typ == "absatz_neu")
    text = zelle(l.out, 2, 3)
    mitten_drin = "Neuer Absatz nach dem Anfang." in text and text.split("Neuer Absatz nach dem Anfang.")[1].strip() != ""
    if ae.status == "uebernommen" and mitten_drin:
        scheitern(f"Status {ae.status}/{ae.konfidenz}: neuer Absatz steht mitten im alten Absatz: {text!r}")


# ---------------------------------------------------------------------------------------------
# RT-29 / RT-30: Section-Nummern als Zahl
# ---------------------------------------------------------------------------------------------
def _zahlen_excel(pfad):
    wb = Workbook()
    ws = wb.active
    ws.title = "CDE"
    ws.append(["Nr.", "Titel", "Inhalt"])
    zeilen = [(1, "Eins", "Erster Abschnitt mit genug Text hier drin."),
              (2.1, "Zwei Eins", "Abschnitt zwei Punkt eins mit viel Text hier."),
              (2.1, "Zwei Zehn", "Abschnitt zwei Punkt zehn mit anderem Text hier."),     # Excel zeigt 2.10 (Format 0.00)
              (2.2, "Zwei Zwei", "Abschnitt zwei Punkt zwei mit noch mehr Text hier."),
              (3, "Drei", "Dritter Abschnitt mit genug Text hier drin.")]
    for z in zeilen:
        ws.append(list(z))
    ws["A3"].number_format = "0.0"
    ws["A4"].number_format = "0.00"        # angezeigt: 2.10
    ws["A5"].number_format = "0.0"
    wb.save(str(pfad))
    return pfad


@pytest.mark.redteam
def test_rt29_section_2_1_und_2_10_als_zahl_bleiben_verschieden(tmp_path):
    xl = ExcelModell(_zahlen_excel(tmp_path / "z.xlsx"))
    s3, s4 = xl.zeilen_section[("CDE", 3)], xl.zeilen_section[("CDE", 4)]
    assert s3 != s4, f"Zeile 3 (angezeigt 2.1) und Zeile 4 (angezeigt 2.10) haben dieselbe Section {s3!r}"


@pytest.mark.redteam
def test_rt30_zahl_bleibt_beim_ersetzen_eine_zahl(tmp_path):
    p = _zahlen_excel(tmp_path / "z.xlsx")
    el = [("h", "1 Eins"), ("p", "Erster Abschnitt mit genug Text hier drin."),
          ("h", "{-2.2-} {+2.3+} Zwei Zwei"), ("p", "Abschnitt zwei Punkt zwei mit noch mehr Text hier."),
          ("h", "3 Drei"), ("p", "Dritter Abschnitt mit genug Text hier drin.")]
    l = uebernimm(analysiere(tmp_path, el, p))
    ae = next(a for a in l.aen if a.typ == "text_ersetzt")
    z = blatt(l.out)["A5"]
    aufbau(str(z.value) == "2.3", f"Ersetzung nicht ausgefuehrt: {z.value!r}")
    if ae.status == "uebernommen" and z.data_type != "n":
        scheitern(f"Zelle A5 war eine Zahl, ist jetzt {z.data_type!r} ({z.value!r}), Format {z.number_format}")


# ---------------------------------------------------------------------------------------------
# RT-31..33: Aufbau falsch erkannt -- nicht still falsch
# ---------------------------------------------------------------------------------------------
_WARN_WORTE = ("Aufbau", "Layout", "Spalte", "unsicher")


def _aufbau_ok_oder_gewarnt(xl, blatt_name, **erwartet):
    lay = xl.layouts[blatt_name]
    ist = {"section_spalte": lay.section_spalte, "inhalt_spalte": lay.inhalt_spalte}
    if all(ist[k] == v for k, v in erwartet.items()):
        return
    if any(w for w in xl.warnungen if any(x in w for x in _WARN_WORTE)):
        return
    scheitern(f"Aufbau falsch erkannt: {ist}, richtig waere {erwartet}; keine Warnung ({xl.warnungen})")


@pytest.mark.redteam
def test_rt31_kommentarspalte_ist_nicht_die_inhaltsspalte(tmp_path):
    lang = "Das ist ein langer interner Kommentar des Bearbeiters zu diesem Punkt mit vielen Details und Hinweisen darin."
    p = schreibe_excel(tmp_path / "a.xlsx", [
        ["Nr.", "Titel", "Inhalt", "Kommentar"],
        ["1", "Eins", "Erster Abschnitt mit genug Text hier drin.", lang],
        ["2", "Zwei", "Zweiter Abschnitt mit genug Text hier drin.", lang],
        ["3", "Drei", "Dritter Abschnitt mit genug Text hier drin.", lang]])
    el = [("h", "1 Eins"), ("p", "Erster Abschnitt mit genug Text hier drin."), ("h", "2 Zwei"),
          ("p", "Zweiter Abschnitt mit genug Text hier drin."), ("p", "{+Neuer Absatz in Abschnitt zwei.+}"),
          ("h", "3 Drei"), ("p", "Dritter Abschnitt mit genug Text hier drin.")]
    l = uebernimm(analysiere(tmp_path, el, p))
    ae = next(a for a in l.aen if a.typ == "absatz_neu")
    if ae.status == "uebernommen" and "Neuer Absatz" in zelle(l.out, 3, 4):
        scheitern(f"Neuer Absatz steht in der Kommentar-Zelle D3 ({zelle(l.out, 3, 4)[-60:]!r}), Status {ae.status}/{ae.konfidenz}")


@pytest.mark.redteam
def test_rt32_lfd_spalte_ist_nicht_die_section_spalte(tmp_path):
    zeilen = [["Lfd", "Nr.", "Titel", "Text"]]
    lfd = 0
    for nr, titel in (("1", "Einleitung"), ("1.1", "Ziel"), ("2", "Technik"), ("2.1", "Leistung")):
        lfd += 1
        zeilen.append([lfd, nr, titel, None])
        for k in (1, 2):
            lfd += 1
            zeilen.append([lfd, None, None, f"Absatz {k} zu {titel} mit genug Text damit er als Inhalt zaehlt."])
    p = schreibe_excel(tmp_path / "b.xlsx", zeilen)
    _aufbau_ok_oder_gewarnt(ExcelModell(p), "CDE", section_spalte=2)


@pytest.mark.redteam
def test_rt33_aufzaehlung_im_text_ist_keine_section(tmp_path):
    p = schreibe_excel(tmp_path / "c.xlsx", [
        ["Nr", "Text", "Prio"],
        ["1", "Erster Abschnitt mit genug Text hier drin.", "A"],
        [None, "1. Das System muss dies tun und noch mehr.", "B"],
        [None, "2. Das System muss das tun und noch mehr.", "B"],
        [None, "3. Das System muss jenes tun und noch mehr.", "B"],
        ["2", "Zweiter Abschnitt mit genug Text hier drin.", "A"],
        [None, "1. Das System soll etwas anderes tun und mehr.", "B"],
        [None, "2. Das System soll es wieder tun und mehr.", "B"]])
    _aufbau_ok_oder_gewarnt(ExcelModell(p), "CDE", section_spalte=1, inhalt_spalte=2)


# ---------------------------------------------------------------------------------------------
# RT-34..36: mehrere / versteckte Blaetter und Zeilen
# ---------------------------------------------------------------------------------------------
def _mehrblatt(pfad, blaetter: dict, versteckt=()):
    wb = Workbook()
    erstes = True
    for name, rows in blaetter.items():
        ws = wb.active if erstes else wb.create_sheet()
        erstes = False
        ws.title = name
        for r in rows:
            ws.append(r)
    for h in versteckt:
        wb[h].sheet_state = "hidden"
    wb.save(str(pfad))
    return pfad


def _sections(prefix, absaetze):
    rows = [["Nr.", "Titel", "Inhalt"]]
    for i, tx in enumerate(absaetze, 1):
        rows.append([str(i), f"Titel {prefix}{i}", tx])
    return rows


EL_AEND = [("h", "1 Eins"), ("p", "Alpha text eins steht hier im Dokument."), ("h", "2 Zwei"),
           ("p", "Das System {-muss-} {+soll+} alle Daten speichern und sichern."), ("h", "3 Drei"),
           ("p", "Gamma text drei steht hier im Dokument.")]
SATZ = "Das System muss alle Daten speichern und sichern."


@pytest.mark.redteam
def test_rt34_section_nummer_gilt_nur_innerhalb_eines_blattes(tmp_path):
    a = ["Alpha text eins steht hier im Dokument.", "Anderer Inhalt zwei nur hier.", SATZ]
    b = ["Beta andere eins steht dort im Dokument.", SATZ, "Delta text drei steht dort im Dokument."]
    p = _mehrblatt(tmp_path / "m.xlsx", {"Blatt A": _sections("A", a), "Blatt B": _sections("B", b)})
    l = uebernimm(analysiere(tmp_path, EL_AEND, p))
    ae = l.aen[0]
    if ae.status == "uebernommen":
        scheitern(f"Der Satz steht in zwei Blaettern; das PDF sagt nicht, welches gemeint ist, trotzdem {ae.status}/{ae.konfidenz} "
                  f"in {ae.ziel} ({ae.grund!r})")


@pytest.mark.redteam
def test_rt35_verstecktes_blatt_wird_nicht_automatisch_geaendert(tmp_path):
    sichtbar = ["Sichtbar eins steht hier irgendwo.", "Sichtbar zwei steht hier irgendwo anders.",
                "Sichtbar drei steht hier irgendwo noch anders."]
    archiv = ["Alpha text eins steht hier im Dokument.", SATZ, "Gamma text drei steht hier im Dokument."]
    p = _mehrblatt(tmp_path / "h.xlsx", {"CDE": _sections("S", sichtbar), "Archiv": _sections("V", archiv)},
                   versteckt=("Archiv",))
    l = uebernimm(analysiere(tmp_path, EL_AEND, p))
    ae = l.aen[0]
    warnt = any("versteckt" in w.lower() or "ausgeblendet" in w.lower() for w in l.plan.warnungen)
    if ae.status == "uebernommen" and ae.ziel.startswith("Archiv!") and not warnt:
        scheitern(f"Aenderung automatisch ({ae.status}/{ae.konfidenz}) in das versteckte Blatt {ae.ziel} geschrieben, ohne Warnung")


@pytest.mark.redteam
def test_rt36_versteckte_zeile_wird_nicht_automatisch_geaendert(tmp_path):
    alt = ["Alpha text eins steht hier im Dokument.", SATZ, "Gamma text drei steht hier im Dokument."]
    p = _mehrblatt(tmp_path / "hz.xlsx", {"CDE": _sections("V", alt)})
    wb = load_workbook(str(p))
    wb["CDE"].row_dimensions[3].hidden = True
    wb.save(str(p))
    l = uebernimm(analysiere(tmp_path, EL_AEND, p))
    ae = l.aen[0]
    warnt = any("versteckt" in w.lower() or "ausgeblendet" in w.lower() for w in l.plan.warnungen)
    if ae.status == "uebernommen" and not warnt:
        scheitern(f"Aenderung automatisch ({ae.status}/{ae.konfidenz}) in die versteckte Zeile {ae.ziel} geschrieben, ohne Warnung")


# ---------------------------------------------------------------------------------------------
# RT-37 / RT-38: Groesse -- eine einzige formatierte Zelle weit unten/rechts
# ---------------------------------------------------------------------------------------------
_MESSEN = textwrap.dedent("""
    import json, resource, sys, time
    from src.excel_model import ExcelModell
    t0 = time.time()
    ExcelModell(sys.argv[1])
    print(json.dumps({"sek": time.time() - t0, "mb": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024}))
""")


@pytest.fixture(scope="module")
def messung_grosses_blatt(tmp_path_factory):
    """Blatt mit 5 Zeilen Inhalt und EINER formatierten Zelle bei Zeile 60000 / Spalte 40, in frischem Prozess gemessen."""
    d = tmp_path_factory.mktemp("gross")
    wb = Workbook()
    ws = wb.active
    ws.title = "CDE"
    ws.append(["Nr.", "Titel", "Inhalt"])
    for i in range(1, 6):
        ws.append([str(i), f"Titel {i}", f"Absatz {i} mit genug Text damit er als Inhalt gilt."])
    ws.cell(60000, 40).fill = PatternFill("solid", fgColor="FFFF00")
    p = d / "gross.xlsx"
    wb.save(str(p))
    wurzel = Path(__file__).resolve().parents[2]
    r = subprocess.run([sys.executable, "-c", _MESSEN, str(p)], capture_output=True, text=True, timeout=28,
                       cwd=str(wurzel), env={**os.environ, "PYTHONPATH": str(wurzel)})
    aufbau(r.returncode == 0, f"Messprozess fehlgeschlagen: {r.stderr[-300:]}")
    return json.loads(r.stdout.strip().splitlines()[-1])


@pytest.mark.redteam
def test_rt37_leeres_riesiges_blatt_wird_schnell_gelesen(messung_grosses_blatt):
    assert messung_grosses_blatt["sek"] < 3, f"Einlesen dauerte {messung_grosses_blatt['sek']:.1f} s (Grenze 3 s)"


@pytest.mark.redteam
def test_rt38_leeres_riesiges_blatt_braucht_wenig_speicher(messung_grosses_blatt):
    assert messung_grosses_blatt["mb"] < 250, f"Spitzenspeicher {messung_grosses_blatt['mb']:.0f} MB (Grenze 250 MB)"
