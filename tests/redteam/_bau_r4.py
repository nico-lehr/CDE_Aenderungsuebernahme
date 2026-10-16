"""Hilfsfunktionen fuer Red-Team R4 (nur erfundene Daten, Regel 16).

Alles lokal in tests/redteam/; es wird nichts in src/, tests/fixtures/ oder den anderen _bau*-Dateien veraendert
(die werden nur importiert).

Ende-zu-Ende-Lauf wie das Tool selbst:
    analysiere_pdf -> ExcelModell -> plane -> uebernehmen(aend, vorher, ziel, freigaben, absaetze=plan.absaetze)

Pruefregel aller R4-Tests: "richtig ODER sichtbar unsicher, nie still falsch".
  * `richtig_oder_unsicher(...)`: die betroffene Aenderung ist "pruefen" oder "nicht_uebernommen" (mit Grund) und die
    Zelle ist unveraendert, ODER die Zelle ist exakt das erwartete Ergebnis.
  * `erkannt_oder_gemeldet(...)`: eine Markierung, die der Leser evtl. nicht versteht, ist entweder als Aenderung erkannt
    (und dann richtig oder unsicher) ODER der Lauf meldet sichtbar etwas (Luecke im Plan, Aenderung "pruefen").

Roh-PDF-Schreiber (`RohDok`): fuer Faelle, die tests/fixtures/pdf_writer.py nicht kann (Separation, DeviceN,
Pfad-Flaechen, Render-Modus 1, Hochzahlen, Zeilen genau wie gewuenscht umbrochen). Standard-Schrift Helvetica, A4.
`wirkt(...)` rendert das Test-PDF mit pdfium und prueft im AUFBAU, dass eine Stelle wirklich blau/rot aussieht.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from src.applier import UebernahmeErgebnis, uebernehmen
from src.excel_model import ExcelModell
from src.models import NICHT, PRUEFEN, UEBERNOMMEN, Aenderung
from src.pdf_reader import PdfAnalyse, analysiere_pdf
from src.planner import Plan, plane
from tests.fixtures.pdf_writer import SEITE_B, SEITE_H, PdfDokument, breite
from tests.fixtures.szenario import baue_excel

# ---------------------------------------------------------------- allgemeine Pruefhelfer


def scheitern(text: str) -> None:
    """Gefundener Tool-Fehler (nur diesen faengt xfail(raises=AssertionError))."""
    raise AssertionError(text)


def aufbau(bedingung: bool, text: str) -> None:
    """Pruefung des TESTAUFBAUS: Fehler hier ist kein erwarteter Tool-Fehler (RuntimeError)."""
    if not bedingung:
        raise RuntimeError("Testaufbau fehlerhaft: " + text)


# ---------------------------------------------------------------- Excel bauen
def excel_sections(ordner: Path, sections: list[tuple[str, str, list[str]]], layout: str = "L1",
                   nachher: Callable[[Any], None] | None = None, name: str = "alt.xlsx") -> Path:
    """Alte Excel aus [(nummer, titel, [absaetze])] (Szenario-Layouts L1/L2/L3, Blatt 'CDE' hinter einem Deckblatt).
    `nachher(workbook)` darf das Workbook vor dem Speichern noch veraendern."""
    ordner.mkdir(parents=True, exist_ok=True)
    p = ordner / name
    baue_excel(p, sections, layout, True)
    if nachher:
        wb = load_workbook(str(p))
        nachher(wb)
        wb.save(str(p))
    return p


def zelle(pfad: Path, zeile: int, spalte: int, blatt: str = "CDE") -> Any:
    return load_workbook(str(pfad))[blatt].cell(zeile, spalte).value


def zelltext(pfad: Path, zeile: int, spalte: int, blatt: str = "CDE") -> str:
    v = zelle(pfad, zeile, spalte, blatt)
    return "" if v is None else str(v)


def alle_texte(pfad: Path, blatt: str = "CDE") -> list[str]:
    ws = load_workbook(str(pfad))[blatt]
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


# ---------------------------------------------------------------- Ende-zu-Ende
@dataclass
class R4Lauf:
    pdf: PdfAnalyse
    plan: Plan
    xl: ExcelModell
    alt: Path
    out: Path
    res: UebernahmeErgebnis
    freigaben: set[str] = field(default_factory=set)

    @property
    def aen(self) -> list[Aenderung]:
        return self.plan.aenderungen

    def kurz(self) -> str:
        """Kurzer Ueberblick fuer Fehlermeldungen (nur Status/Typ/Grund; die Daten sind erfunden)."""
        return "; ".join(f"{a.id}:{a.typ}/{a.status}/{a.konfidenz}/wirkung={a.wirkung or '-'}"
                         + (f" grund={a.grund[:60]!r}" if a.grund else "") for a in self.aen) or "keine Aenderungen"


def pdf_schreiben(ordner: Path, quelle: "PdfDokument | RohDok | bytes | list[tuple[str, str]]", name: str = "a.pdf") -> Path:
    """PDF ablegen: PdfDokument, RohDok, rohe Bytes oder Szenario-Elemente [('h'|'p', markup)] (ohne Kopf-/Fusszeile)."""
    ordner.mkdir(parents=True, exist_ok=True)
    pfad = ordner / name
    if isinstance(quelle, RohDok):
        pfad.write_bytes(quelle.bytes())
    elif isinstance(quelle, (bytes, bytearray)):
        pfad.write_bytes(bytes(quelle))
    elif isinstance(quelle, PdfDokument):
        quelle.speichern(pfad)
    else:
        d = PdfDokument(variante="farbe", kopf_fuss=False)
        for art, markup in quelle:
            (d.ueberschrift if art == "h" else d.block)(markup)
        d.speichern(pfad)
    return pfad


def ende_zu_ende(ordner: Path, quelle: PdfDokument | bytes | list[tuple[str, str]] | Path, alt: Path,
                 freigaben: set[str] | None = None, name: str = "out.xlsx") -> R4Lauf:
    """Der ganze Weg ueber die oeffentlichen Funktionen, genau wie das Tool."""
    pdf_pfad = quelle if isinstance(quelle, Path) else pdf_schreiben(ordner, quelle)
    analyse = analysiere_pdf(pdf_pfad)
    xl = ExcelModell(alt)
    plan = plane(analyse, xl)
    out = ordner / name
    freig = set(freigaben or ())
    res = uebernehmen(plan.aenderungen, alt, out, freig, absaetze=plan.absaetze)
    return R4Lauf(pdf=analyse, plan=plan, xl=xl, alt=alt, out=out, res=res, freigaben=freig)


def automatisch(ae: Aenderung) -> bool:
    return ae.status == UEBERNOMMEN and not ae.bestaetigt


def unsicher(ae: Aenderung) -> bool:
    """Sichtbar unsicher: 'pruefen' (ist immer ein sichtbarer Vorschlag, auch bei Konfidenz 'mittel' ohne Grundtext)
    oder 'nicht_uebernommen' MIT Grund."""
    return ae.status == PRUEFEN or (ae.status == NICHT and bool(ae.grund.strip()))


def richtig_oder_unsicher(l: R4Lauf, betroffen: list[Aenderung], zeile: int, spalte: int, erwartet: Any,
                          vorher: Any, was: str, blatt: str = "CDE") -> str:
    """Prueft die Regel fuer EINE Zelle. Gibt 'richtig' oder 'unsicher' zurueck, sonst AssertionError.

    * ist irgendeine betroffene Aenderung automatisch uebernommen -> die Zelle muss exakt `erwartet` sein
    * sonst muessen alle betroffenen sichtbar unsicher sein (mit Grund) und die Zelle unveraendert (`vorher`) oder richtig
    """
    aufbau(bool(betroffen), f"{was}: keine betroffene Aenderung gefunden ({l.kurz()})")
    ist = zelle(l.out, zeile, spalte, blatt)
    if ist == erwartet:
        return "richtig"
    if any(a.status == UEBERNOMMEN for a in betroffen):
        scheitern(f"{was}: still falsch -- Status {[a.status for a in betroffen]}, Wirkung "
                  f"{[a.wirkung for a in betroffen]}, Zelle {ist!r} statt {erwartet!r}")
    if not all(unsicher(a) for a in betroffen):
        scheitern(f"{was}: nicht uebernommen, aber ohne Grund ({l.kurz()})")
    if ist != vorher:
        scheitern(f"{was}: Aenderung steht auf {[a.status for a in betroffen]}, aber die Zelle wurde trotzdem "
                  f"veraendert: {ist!r} (vorher {vorher!r}, richtig waere {erwartet!r})")
    return "unsicher"


def sichtbare_meldungen(l: R4Lauf, fremd: Callable[[Aenderung], bool] = lambda a: False) -> list[str]:
    """Was ein Mensch im Ergebnis als 'unsicher' sieht: Luecken des Plans, nicht automatische Aenderungen
    (ohne die `fremd`en, die zu einer ANDEREN, absichtlich eingebauten Aenderung gehoeren)."""
    out = [f"luecke:{x.get('art')}" for x in l.plan.luecken]
    out += [f"aenderung:{a.id}:{a.status}" for a in l.aen if unsicher(a) and not fremd(a)]
    return out


def erkannt_oder_gemeldet(l: R4Lauf, zeile: int, spalte: int, erwartet: Any, vorher: Any, was: str,
                          passt: Callable[[Aenderung], bool], blatt: str = "CDE",
                          fremd: Callable[[Aenderung], bool] = lambda a: False) -> str:
    """Wie richtig_oder_unsicher, aber die Markierung darf auch gar nicht als Aenderung erkannt sein -- dann muss der
    Lauf sie sichtbar melden (Luecke oder 'pruefen'), sonst gilt der Absatz still als 'unveraendert'."""
    betroffen = [a for a in l.aen if passt(a)]
    if betroffen:
        return richtig_oder_unsicher(l, betroffen, zeile, spalte, erwartet, vorher, was, blatt)
    ist = zelle(l.out, zeile, spalte, blatt)
    if ist == erwartet:
        return "richtig"
    meld = sichtbare_meldungen(l, fremd)
    if not meld:
        scheitern(f"{was}: nicht erkannt und nicht gemeldet -- Zelle bleibt still {ist!r} "
                  f"(richtig waere {erwartet!r}); Aenderungen: {l.kurz()}; Luecken: {l.plan.luecken}")
    if ist != vorher:
        scheitern(f"{was}: Zelle wurde ohne passende Aenderung veraendert: {ist!r}")
    return "gemeldet"


# ---------------------------------------------------------------- Roh-PDF-Schreiber
# Eine Zeile ist eine Liste von Stuecken: (text, farbe) oder (text, farbe, optionen).
#   farbe "k"          schwarz (0 g)
#   farbe "<befehle>"  rohe Inhaltsstrom-Befehle VOR dem Text, z. B. "/Rot cs 1 scn" oder "0 0 1 RG 1 Tr"
#   farbe ("flaeche", farbbefehl, form)  Text schwarz, dahinter eine Flaeche (form: "re" Rechteck | "kurve" Bezier-Pfad)
#   optionen {"groesse": 7, "hoch": 4}   kleinere, angehobene Schrift (Hochzahl); "luecke": False = direkt angeklebt
# Zwischen zwei Stuecken steht KEIN Leerzeichen; Leerzeichen gehoeren in den Text (" und ").
GROESSE = 11.0
LINKS = 56.0
RECHTS = SEITE_B - 56.0


@dataclass
class RohDok:
    """Roh-PDF mit Ueberschriften und Absaetzen (Helvetica, WinAnsi), eigene Objekte (z. B. Funktions-Stroeme) moeglich.

    `ressourcen`: zusaetzliche /Resources-Eintraege; darin wird "{OBJ<n>}" durch die Nummer des n-ten Zusatzobjekts ersetzt.
    `objekte`: Zusatzobjekte als Bytes (komplett, z. B. b"<< /FunctionType 4 ... >>\\nstream\\n...\\nendstream")."""
    ressourcen: str = ""
    objekte: list[bytes] = field(default_factory=list)
    seiten: list[list[str]] = field(default_factory=lambda: [[]])
    top: float = 64.0

    @property
    def inhalt(self) -> list[str]:
        return self.seiten[-1]

    def neue_seite(self) -> None:
        self.seiten.append([])
        self.top = 64.0

    def _text(self, x: float, top: float, text: str, groesse: float, fett: bool, vorher: str, hoch: float = 0.0) -> None:
        basis = SEITE_H - (top + GROESSE * 0.8) + hoch if not fett else SEITE_H - (top + groesse * 0.8)
        roh = text.encode("cp1252").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)").decode("latin-1")
        self.inhalt.append(f"q BT /{'F2' if fett else 'F1'} {groesse} Tf {vorher} {x:.2f} {basis:.2f} Td ({roh}) Tj ET Q\n")

    def ueberschrift(self, text: str) -> None:
        self.top += 6
        self._text(LINKS, self.top, text, 14.0, True, "0 g")
        self.top += 18 + 8

    def absatz(self, zeilen: list[list[tuple]], abstand_danach: float = 9.0) -> list[tuple[float, float, float, float]]:
        """Schreibt die Zeilen (je Zeile 15 pt); gibt die Kaesten (x0, top, x1, bottom) aller Zeilen zurueck."""
        kaesten = []
        for z in zeilen:
            x = LINKS
            for stueck in z:
                text, farbe = stueck[0], stueck[1]
                opt = stueck[2] if len(stueck) > 2 else {}
                groesse = float(opt.get("groesse", GROESSE))
                w = breite(text, groesse)
                if isinstance(farbe, tuple) and farbe[0] == "flaeche":
                    _, fb, form = farbe
                    self.flaeche(x, self.top - 1, x + w, self.top + groesse + 1, fb, form)
                    vor = "0 g"
                else:
                    vor = "0 g" if farbe == "k" else farbe
                self._text(x, self.top, text, groesse, False, vor, float(opt.get("hoch", 0.0)))
                x += w
            kaesten.append((LINKS, self.top, x, self.top + GROESSE))
            self.top += 15
        self.top += abstand_danach
        return kaesten

    def flaeche(self, x0: float, top: float, x1: float, bot: float, farbbefehl: str, form: str = "re") -> None:
        y0, y1 = SEITE_H - bot, SEITE_H - top
        if form == "re":
            pfad = f"{x0:.2f} {y0:.2f} {x1 - x0:.2f} {y1 - y0:.2f} re"
        else:   # geschlossene Kurve mit gerundeten Ecken (Bezier), KEIN Rechteck-Operator
            r = min(3.0, (y1 - y0) / 3)
            pfad = (f"{x0 + r:.2f} {y0:.2f} m {x1 - r:.2f} {y0:.2f} l {x1:.2f} {y0:.2f} {x1:.2f} {y0:.2f} {x1:.2f} {y0 + r:.2f} c "
                    f"{x1:.2f} {y1 - r:.2f} l {x1:.2f} {y1:.2f} {x1:.2f} {y1:.2f} {x1 - r:.2f} {y1:.2f} c "
                    f"{x0 + r:.2f} {y1:.2f} l {x0:.2f} {y1:.2f} {x0:.2f} {y1:.2f} {x0:.2f} {y1 - r:.2f} c "
                    f"{x0:.2f} {y0 + r:.2f} l {x0:.2f} {y0:.2f} {x0:.2f} {y0:.2f} {x0 + r:.2f} {y0:.2f} c h")
        self.inhalt.insert(0, f"q {farbbefehl} {pfad} f Q\n")     # Flaechen zuerst (hinter dem Text)

    def linie(self, x0: float, x1: float, top: float, farbe: str = "0 0 0 RG", staerke: float = 0.8) -> None:
        y = SEITE_H - top
        self.inhalt.append(f"q {farbe} {staerke} w {x0:.2f} {y:.2f} m {x1:.2f} {y:.2f} l S Q\n")

    def bytes(self) -> bytes:
        objekte: list[bytes] = [b"", b""]               # 1 = Katalog, 2 = Seitenbaum

        def neu(inhalt: bytes) -> int:
            objekte.append(inhalt)
            return len(objekte)

        f1 = neu(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
        f2 = neu(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
        zusatz = [neu(o) for o in self.objekte]
        res = self.ressourcen
        for i, nr in enumerate(zusatz):
            res = res.replace("{OBJ%d}" % i, str(nr))
        ids = []
        for s in self.seiten:
            strom = "".join(s).encode("latin-1")
            inh = neu(b"<< /Length %d >>\nstream\n" % len(strom) + strom + b"\nendstream")
            ids.append(neu(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Contents %d 0 R "
                           b"/Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> %s >> >>"
                           % (SEITE_B, SEITE_H, inh, f1, f2, res.encode("latin-1"))))
        objekte[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
        objekte[1] = b"<< /Type /Pages /Kids [" + b" ".join(b"%d 0 R" % i for i in ids) + b"] /Count %d >>" % len(ids)
        out = bytearray(b"%PDF-1.4\n")
        offs = []
        for i, o in enumerate(objekte, 1):
            offs.append(len(out))
            out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objekte) + 1)
        for o in offs:
            out += b"%010d 00000 n \n" % o
        out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objekte) + 1, xref)
        return bytes(out)


def funktion_ps(code: str, eingaenge: int, ausgaenge: int) -> bytes:
    """PDF-Funktion Typ 4 (PostScript-Rechner) als Strom-Objekt, z. B. fuer DeviceN mit zwei Kanaelen."""
    dom = " ".join(["0 1"] * eingaenge)
    rng = " ".join(["0 1"] * ausgaenge)
    strom = code.encode("latin-1")
    return (b"<< /FunctionType 4 /Domain [%s] /Range [%s] /Length %d >>\nstream\n"
            % (dom.encode(), rng.encode(), len(strom)) + strom + b"\nendstream")


def separation(name: str, cmyk_voll: tuple[float, float, float, float]) -> str:
    """Separation-Farbraum (Ton 0..1 -> CMYK, Funktion Typ 2) als Eintrag fuer /ColorSpace."""
    c1 = " ".join(f"{v:g}" for v in cmyk_voll)
    return (f"/{name} [/Separation /{name} /DeviceCMYK "
            f"<< /FunctionType 2 /Domain [0 1] /C0 [0 0 0 0] /C1 [{c1}] /N 1 >>]")


def wirkt(dok: RohDok, x0: float, top: float, x1: float, bot: float, seite: int = 0) -> str:
    """Wie wirkt der Bereich im gerenderten PDF (pdfium)? 'blau' | 'rot' | 'gruen' | 'neutral' -- nur fuer die
    Aufbau-Pruefung, dass das Test-PDF wirklich so aussieht, wie der Test es behauptet."""
    import colorsys
    from collections import Counter

    import pypdfium2 as pdfium
    skala = 2.0
    pdf = pdfium.PdfDocument(dok.bytes())
    try:
        bild = pdf[seite].render(scale=skala).to_pil().convert("RGB")
    finally:
        pdf.close()
    ausschnitt = bild.crop((int(x0 * skala), int(top * skala), int(x1 * skala), int(bot * skala)))
    zaehler: Counter = Counter()
    roh = ausschnitt.tobytes()
    for i in range(0, len(roh) - 2, 3):
        r, g, b = roh[i], roh[i + 1], roh[i + 2]
        h, s, v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
        if s < 0.15 or v < 0.25:
            continue
        grad = h * 360
        zaehler["rot" if (grad <= 25 or grad >= 330) else "blau" if 185 <= grad <= 265 else "gruen" if 80 <= grad <= 160
                else "sonst"] += 1
    if not zaehler:
        return "neutral"
    return zaehler.most_common(1)[0][0]
