"""Red-Team K-PL: Zeilen einfuegen und Excel-Eigenschaften, die openpyxl nicht mitverschiebt.

Erwartung laut Auftrag: entweder bleibt alles konsistent, ODER die Aenderung steht auf "pruefen"
(mit Warnung) -- nie still beschaedigt. Gemeinsame Pruefung: `_konsistent_oder_pruefen`.
"""
import pytest
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from tests.redteam._bau import aufbau, baue_alt, bearbeite, analysiere, scheitern, uebernimm, zeile_mit

# Neue Section "2.5" (mit Folgeabsatz) kommt nach Section 2 -> in der Excel wird EINE Zeile bei Zeile 4 eingefuegt.
EL = [("h", "1 Eins"), ("p", "Absatz eins mit genug Text darin."),
      ("h", "2 Zwei"), ("p", "Absatz zwei mit genug Text darin."),
      ("h", "{+2.5 Neue Section+}"), ("p", "{+Folgeabsatz der neuen Section steht hier.+}"),
      ("h", "3 Drei"), ("p", "Absatz drei mit genug Text darin."),
      ("h", "4 Vier"), ("p", "Absatz vier mit genug Text darin."),
      ("h", "5 Fuenf"), ("p", "Absatz fuenf mit genug Text darin."),
      ("h", "6 Sechs"), ("p", "Absatz sechs mit genug Text darin.")]


def _lauf(tmp_path, eigenschaft):
    """Excel mit einer Eigenschaft bauen, Tool laufen lassen (ohne Freigaben)."""
    alt = baue_alt(tmp_path, EL, "L1", nachher=lambda p: bearbeite(p, eigenschaft))
    l = uebernimm(analysiere(tmp_path, EL, alt))
    neue = [a for a in l.aen if a.typ == "section_neu"]
    aufbau(len(neue) == 1, "genau eine neue Section erwartet")
    return l, neue[0]


def _konsistent_oder_pruefen(l, neue, ist_konsistent, was: str) -> None:
    """Fehler nur, wenn die Eigenschaft nach dem Einfuegen NICHT mehr stimmt, die Zeile aber automatisch
    eingetragen wurde ('uebernommen'). Steht die Aenderung auf 'pruefen', wird nichts automatisch geschrieben;
    dann muss sie aber auch erklaeren, warum (Hinweise nicht leer)."""
    if neue.status == "pruefen" and not neue.hinweise:
        scheitern(f"{was} -- die Aenderung steht auf 'pruefen', hat aber keinen Hinweis fuer den Menschen (Grund: {neue.grund!r})")
    if ist_konsistent():
        return
    if neue.status != "uebernommen":
        return
    scheitern(f"{was} -- neue Zeile wurde trotzdem automatisch eingetragen (Status {neue.status}/{neue.konfidenz}); "
              f"Warnungen: {l.plan.warnungen}")


def _zeile(l, text: str) -> int:
    r = zeile_mit(l.out, text)
    aufbau(r is not None, f"{text!r} muss in der Ergebnis-Excel stehen")
    return r


# ----------------------------------------------------------------------------------------------
@pytest.mark.redteam
def test_rt18_verbundene_zelle_ueber_der_einfuegestelle(tmp_path):
    def merge(wb):
        wb["CDE"].merge_cells("D2:D4")       # Bemerkung gilt fuer die Sections 1 bis 3
    l, neue = _lauf(tmp_path, merge)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        rng = next(r for r in ws.merged_cells.ranges if r.min_row == 2 and r.min_col == 4)
        return rng.max_row >= _zeile(l, "Absatz drei")          # Verbindung reicht weiter bis Section 3
    _konsistent_oder_pruefen(l, neue, ok, "Verbundene Zelle D2:D4 deckt Section 3 nicht mehr ab")


@pytest.mark.redteam
def test_rt19_versteckte_zeile_bleibt_versteckt(tmp_path):
    def verstecke(wb):
        wb["CDE"].row_dimensions[6].hidden = True       # Section 5 ("Absatz fuenf")
    l, neue = _lauf(tmp_path, verstecke)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        return bool(ws.row_dimensions[_zeile(l, "Absatz fuenf")].hidden) and not ws.row_dimensions[_zeile(l, "Absatz vier")].hidden
    _konsistent_oder_pruefen(l, neue, ok, "Die versteckte Zeile (Section 5) ist nicht mehr versteckt / eine andere Zeile ist versteckt")


@pytest.mark.redteam
def test_rt20_hyperlink_wandert_mit(tmp_path):
    def link(wb):
        wb["CDE"]["C6"].hyperlink = "https://example.org/x"      # Zelle von Section 5
    l, neue = _lauf(tmp_path, link)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        r = _zeile(l, "Absatz fuenf")
        return ws.cell(r, 3).hyperlink is not None and ws.cell(r - 1, 3).hyperlink is None
    _konsistent_oder_pruefen(l, neue, ok, "Der Hyperlink steht nicht mehr an der Zelle von Section 5")


@pytest.mark.redteam
def test_rt21_bedingte_formatierung_wird_verschoben(tmp_path):
    def cf(wb):
        wb["CDE"].conditional_formatting.add(
            "D2:D7", CellIsRule(operator="equal", formula=['"Status offen"'], fill=PatternFill("solid", fgColor="FFFF00")))
    l, neue = _lauf(tmp_path, cf)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        letzte = _zeile(l, "Absatz sechs")
        return any(f"D{letzte}" in rng.sqref for rng in ws.conditional_formatting)
    _konsistent_oder_pruefen(l, neue, ok, "Bedingte Formatierung D2:D7 deckt die letzte Zeile (Section 6) nicht mehr ab")


@pytest.mark.redteam
def test_rt22_datenvalidierung_wird_verschoben(tmp_path):
    def dv(wb):
        ws = wb["CDE"]
        d = DataValidation(type="list", formula1='"Status offen,erledigt"')
        ws.add_data_validation(d)
        d.add("D2:D7")
    l, neue = _lauf(tmp_path, dv)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        letzte = _zeile(l, "Absatz sechs")
        return any(f"D{letzte}" in v.sqref for v in ws.data_validations.dataValidation)
    _konsistent_oder_pruefen(l, neue, ok, "Datenvalidierung D2:D7 deckt die letzte Zeile (Section 6) nicht mehr ab")


@pytest.mark.redteam
def test_rt23_autofilter_wird_erweitert(tmp_path):
    def af(wb):
        wb["CDE"].auto_filter.ref = "A1:D7"
    l, neue = _lauf(tmp_path, af)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        return CellRange(ws.auto_filter.ref).max_row >= _zeile(l, "Absatz sechs")
    _konsistent_oder_pruefen(l, neue, ok, "Autofilter A1:D7 endet vor der letzten Zeile")


@pytest.mark.redteam
def test_rt24_excel_tabelle_wird_erweitert(tmp_path):
    def tab(wb):
        wb["CDE"].add_table(Table(displayName="Tab1", ref="A1:D7"))
    l, neue = _lauf(tmp_path, tab)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        return all(CellRange(t.ref).max_row >= _zeile(l, "Absatz sechs") for t in ws._tables.values())
    _konsistent_oder_pruefen(l, neue, ok, "Excel-Tabelle A1:D7 endet vor der letzten Zeile")


@pytest.mark.redteam
def test_rt25_benannter_bereich_wird_erweitert(tmp_path):
    def name(wb):
        wb.defined_names["Daten"] = DefinedName("Daten", attr_text="CDE!$A$2:$D$7")
    l, neue = _lauf(tmp_path, name)

    def ok():
        wb = load_workbook(str(l.out))
        bis = max(CellRange(rng.replace("$", "")).max_row for _, rng in wb.defined_names["Daten"].destinations)
        return bis >= _zeile(l, "Absatz sechs")
    _konsistent_oder_pruefen(l, neue, ok, "Benannter Bereich 'Daten' (A2:D7) endet vor der letzten Zeile")


@pytest.mark.redteam
def test_rt26_formel_nach_unten_wird_angepasst(tmp_path):
    def formel(wb):
        wb["CDE"]["F2"] = "=C7"            # zeigt auf Section 6
    l, neue = _lauf(tmp_path, formel)

    def ok():
        ws = load_workbook(str(l.out))["CDE"]
        return ws["F2"].value == f"=C{_zeile(l, 'Absatz sechs')}"
    _konsistent_oder_pruefen(l, neue, ok, "Formel =C7 zeigt nach dem Einfuegen nicht mehr auf Section 6")
