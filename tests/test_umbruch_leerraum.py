"""Leerzeichen am Zeilenumbruch neben einer Ersetzung (Auftrag ZUORDNEN, Welle 6): Am Zeilen- oder Seitenumbruch zeichnet das PDF
kein Leerzeichen, also sieht man nicht, ob es markiert ist. Steht es hinter einer Ersetzung ("12" gestrichen, "15" neu, dann der
Umbruch, dann "mm"), kann es dreierlei sein: gemeinsam (alt und neu), nur neu (Teil der Einfuegung "15 ") oder nur alt (Teil der
Streichung). Die Excel entscheidet nur bei echter Mehrdeutigkeit: Steht der alte Absatz dort OHNE Leerzeichen ("12mm"), kann es nur
zur Einfuegung gehoeren; steht er MIT Leerzeichen, bleibt es beim gemeinsamen (bisher) -- die Excel kann "gemeinsam" und "nur alt"
nicht unterscheiden. Ohne Umbruch (sichtbarer, unmarkierter Leerraum) und ohne passenden Excel-Absatz aendert sich nichts.

Nur erfundene Daten (Regel 16). Die Mappe und das Word-PDF entstehen im Test (tests/test_register.py, tests/_word_pdf.py).
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from src import markierung as mk
from src.applier import uebernehmen
from src.excel_model import ExcelModell
from src.models import PRUEFEN, UEBERNOMMEN
from src.pdf_reader import analysiere_pdf
from src.planner import plane
from tests._word_pdf import EBENEN, RECHTS, WordDok, breite
from tests.test_register import GRUND, G, _EBENE, _label_art, register_mappe, stuecke

VOKABULAR = "Das Gehaeuse steht auf einem Sockel aus Beton und ist gegen Feuchte und Frost gut geschuetzt".split()
EBENE_TEXT = 3                                  # Text eines Listenpunkts "1." steht auf Einzugsebene 3 (Nummer auf 2)
ZIEL = 4                                        # Index in GRUND: der Listenpunkt "1." unter "B."
ZEILE_B = 8 + ZIEL                              # Zeile in der Mappe (Daten ab Zeile 8)
REST = "mm im Betrieb und darf nicht ueberschritten werden."


def _fuelltext(platz: float) -> str:
    """Fuelltext mit Leerzeichen dahinter, der in `platz` pt passt (Zeile fast voll: der Leser haelt die naechste Zeile dann fuer
    die Fortsetzung desselben Absatzes)."""
    text, i = "", 0
    while breite(text + VOKABULAR[i % len(VOKABULAR)] + " ") <= platz - 1.0 and i < 60:
        text += VOKABULAR[i % len(VOKABULAR)] + " "
        i += 1
    return text


def _fall(tmp_path, vorn: str, mitte: list[tuple], rest: str, excel_mitte: str, umbruch: bool = True,
          excel_rest: str | None = None):
    """Das kleine Register GRUND (tests/test_register.py); der Listenpunkt "1." unter "B." (Zeile 12) ist der Zielabsatz mit
    Umbruch hinter der Aenderung `mitte` (Zeile 1 endet mit ihr, Zeile 2 beginnt mit `rest`).
    -> (Plan, ExcelModell, Excel-Datei, PdfAnalyse, Excel-Text des Listenpunkts ohne Nummer)"""
    fuell = _fuelltext(RECHTS - EBENEN[EBENE_TEXT] - breite(vorn) - sum(breite(st[0]) for st in mitte)) if umbruch else ""
    zeile1 = [(fuell + vorn, "k")] + mitte
    zeilen = [zeile1, [(rest, "k")]] if umbruch else [zeile1 + [(rest, "k")]]
    excel_b = fuell + vorn + excel_mitte + (rest if excel_rest is None else excel_rest)
    d = WordDok()
    for i, (art, label, text) in enumerate(GRUND):
        if d.top > 600:
            d.neue_seite()
        if art == "part":
            d.kapitel(stuecke(f"PART {label} ") + stuecke(text))
        elif art == "section":
            d.ueberschrift(stuecke(label), stuecke(text))
        elif art in ("sub", "sub_titel"):
            d.listenpunkt(1, stuecke(label), [stuecke(text)], abstand=0)
        elif i == ZIEL:
            d.listenpunkt(2, stuecke(label), zeilen, abstand=0)
        else:
            d.listenpunkt(_EBENE[_label_art(label)], stuecke(label), [stuecke(text)], abstand=0)
    pdf = analysiere_pdf(d.speichern(tmp_path / "w.pdf"))
    absaetze = [(art, label.rstrip(".") if art in ("sub", "sub_titel") else label, excel_b if i == ZIEL else text)
                for i, (art, label, text) in enumerate(GRUND)]
    xlsx = register_mappe(tmp_path / "alt.xlsx", absaetze)
    xl = ExcelModell(xlsx)
    return plane(pdf, xl), xl, xlsx, pdf, excel_b


def _zelle(excel_b: str) -> str:
    """Zellinhalt der Zeile: Listennummer, Tabulator (Polster der Mappe) und Text."""
    return "1.\t" + excel_b


def _ops(ae):
    return [(o.start, o.ende, o.text) for o in ae.ops]


def _a1(ae) -> list[str]:
    return [c for c in ae.pruefcodes if c.startswith("a1_")]


@pytest.mark.legende_neu
def test_ersetzung_am_zeilenende_excel_ohne_leerzeichen_wird_exakt(tmp_path):
    """Das Leerzeichen hinter "15" gehoert zur Einfuegung (Excel: "12mm"): der alte Absatz stimmt Zeichen fuer Zeichen, die
    Aenderung laeuft automatisch (vorher: A1 `a1_leerraum` + `excel_weicht_ab`, nur ein Vorschlag)."""
    plan, xl, xlsx, pdf, excel_b = _fall(tmp_path, "betraegt ", [("12", "d"), ("15", "n")], REST, "12")
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and not _a1(ae) and "excel_weicht_ab" not in ae.pruefcodes, (ae.status, ae.pruefcodes)
    s = _zelle(excel_b).index("12")
    assert _ops(ae) == [(s, s + 2, "15 ")]                       # "12mm" -> "15 mm": das Leerzeichen kommt mit der Einfuegung


@pytest.mark.markierung_an
@pytest.mark.legende_neu
def test_ersetzung_am_zeilenende_ende_zu_ende_mit_markierung_und_zweitem_pruefweg(tmp_path):
    """Excel-Absatz im alten Zustand unveraendert, neuer Zustand = PDF neu, das fehlende Leerzeichen ist neu (rot); die Gegenprobe
    (A3, `wirkung`) und der zweite Pruefweg (Konsens-Text) sind gruen."""
    plan, xl, xlsx, pdf, excel_b = _fall(tmp_path, "betraegt ", [("12", "d"), ("15", "n")], REST, "12")
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze, pdf=pdf)
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and ae.wirkung == "ok", (ae.status, ae.wirkung, ae.pruefcodes)
    assert not any(c.startswith(("gegenprobe", "konsens")) for c in ae.pruefcodes)
    st = mk.stuecke_aus_zelle(load_workbook(str(out), rich_text=True)["Abgleich"].cell(ZEILE_B, G))
    assert mk.ansicht(st, "alt") == _zelle(excel_b)                            # alter Zustand unveraendert
    assert mk.ansicht(st, "neu") == _zelle(excel_b).replace("12mm", "15 mm")     # neuer Zustand = PDF neu
    assert [[s.format, s.text] for s in st if s.format != mk.NORMAL] == [["alt", "12"], ["neu", "15 "]]   # Leerzeichen: neu (rot)
    assert load_workbook(str(xlsx))["Abgleich"].cell(ZEILE_B, G).value == _zelle(excel_b)    # die Original-Excel bleibt
    assert ae.stuecke and ["neu", "15 "] in ae.stuecke                           # Vorschau = spaeteres Ergebnis


@pytest.mark.markierung_an
@pytest.mark.legende_neu
def test_ersetzung_mit_markiertem_leerzeichen_vorn_und_neuer_wortgruppe(tmp_path):
    """Form aus dem Befund: gestrichen " 12" (beginnt mit dem markierten Leerzeichen), neu " ca. 15" (mit markiertem Leerzeichen
    in der Mitte), dann der Umbruch und "mm" (die Excel: "12mm")."""
    plan, xl, xlsx, pdf, excel_b = _fall(tmp_path, "betraegt", [(" 12", "d"), (" ca. 15", "n")], REST, " 12")
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and not _a1(ae) and "excel_weicht_ab" not in ae.pruefcodes, (ae.status, ae.pruefcodes)
    out = tmp_path / "out.xlsx"
    uebernehmen(plan.aenderungen, xlsx, out, xl=xl, absaetze=plan.absaetze, pdf=pdf)
    assert ae.status == UEBERNOMMEN and ae.wirkung == "ok", (ae.status, ae.wirkung, ae.pruefcodes)
    st = mk.stuecke_aus_zelle(load_workbook(str(out), rich_text=True)["Abgleich"].cell(ZEILE_B, G))
    assert mk.ansicht(st, "alt") == _zelle(excel_b)
    assert mk.ansicht(st, "neu") == _zelle(excel_b).replace(" 12mm", " ca. 15 mm")


@pytest.mark.legende_neu
def test_excel_mit_leerzeichen_bleibt_gemeinsam_wie_bisher(tmp_path):
    """Steht das Leerzeichen in der Excel ("12 mm"), gehoert es wie bisher beiden Fassungen: nur "12" -> "15". Die Excel kann
    "gemeinsam" und "nur alt" nicht unterscheiden -- dann aendert sich nichts."""
    plan, xl, xlsx, pdf, excel_b = _fall(tmp_path, "betraegt ", [("12", "d"), ("15", "n")], REST, "12 ")
    (ae,) = plan.aenderungen
    assert ae.status == UEBERNOMMEN and not _a1(ae) and "excel_weicht_ab" not in ae.pruefcodes, (ae.status, ae.pruefcodes)
    s = _zelle(excel_b).index("12")
    assert _ops(ae) == [(s, s + 2, "15")]


@pytest.mark.legende_neu
def test_ohne_umbruch_bleibt_fehlendes_leerzeichen_eine_abweichung(tmp_path):
    """Steht das Leerzeichen sichtbar (unmarkiert) hinter der Ersetzung, gehoert es beiden Fassungen. Hat die Excel es nicht, ist
    das eine Abweichung der Excel (kein Umbruch, also keine Mehrdeutigkeit): nur ein Vorschlag, nie automatisch."""
    plan, xl, xlsx, pdf, excel_b = _fall(tmp_path, "betraegt ", [("12", "d"), ("15", "n")], " mm im Betrieb.", "12", umbruch=False,
                                         excel_rest="mm im Betrieb.")
    (ae,) = plan.aenderungen
    assert ae.status == PRUEFEN and "excel_weicht_ab" in ae.pruefcodes and "a1_leerraum" in ae.pruefcodes, (ae.status, ae.pruefcodes)
